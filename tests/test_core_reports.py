from __future__ import annotations

import base64
import ftplib
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from time import perf_counter
from typing import Any
from unittest import mock

from src.core import database as database_module
from src.core.database import Database
from src.core.services import (
    AppError,
    BatchExportService,
    ExportService,
    PairingService,
    PhotoAlbumService,
)
from tests.fixtures import load_tournament_fixture
from tests.support.core_service_base import CoreServiceTestCase


class _RecordingFTP:
    """Cliente FTP falso (sem rede) para testar ftp_publish/PhotoAlbumService."""

    def __init__(self) -> None:
        self.stored: dict[str, bytes] = {}
        self.made: list[str] = []
        self.cwds: list[str] = []
        self.quit_called = False

    def cwd(self, path: str) -> None:
        self.cwds.append(path)
        if path not in self.made:
            raise ftplib.error_perm("550 No such directory")

    def mkd(self, path: str) -> None:
        self.made.append(path)

    def storbinary(self, command: str, handle: Any) -> None:
        name = command.split(" ", 1)[1]
        self.stored[name] = handle.read()

    def voidcmd(self, command: str) -> str:
        return "200 OK"

    def quit(self) -> None:
        self.quit_called = True


class FtpPhotoAlbumTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.service = PhotoAlbumService(self.db)
        self.photos = base / "fotos"
        self.photos.mkdir()
        (self.photos / "a.jpg").write_bytes(b"jpgdata")
        (self.photos / "b.PNG").write_bytes(b"pngdata")
        (self.photos / "notas.txt").write_text("ignore", encoding="utf-8")

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_list_images_and_gallery(self) -> None:
        from src.services.ftp_publish import build_gallery_html, list_images

        images = list_images(self.photos)
        self.assertEqual(["a.jpg", "b.PNG"], [img.name for img in images])
        html_out = build_gallery_html([img.name for img in images], "Torneio")
        self.assertIn("a.jpg", html_out)
        self.assertIn("<title>Torneio</title>", html_out)

    def test_upload_files_creates_remote_dir(self) -> None:
        from src.services.ftp_publish import list_images, upload_files

        fake = _RecordingFTP()
        uploaded = upload_files(fake, list_images(self.photos), "public/fotos")
        self.assertEqual(["a.jpg", "b.PNG"], uploaded)
        self.assertEqual(["public", "fotos"], fake.made)
        self.assertIn("a.jpg", fake.stored)

    def test_publish_album_with_fake_client(self) -> None:
        self.service.save_config(
            {"host": "ftp.exemplo.com", "port": 21, "user": "u", "password": "secret", "remote_dir": "fotos"}
        )
        fake = _RecordingFTP()
        result = self.service.publish_album(self.photos, client_factory=lambda _cfg: fake)
        self.assertEqual(3, result["total"])  # 2 imagens + index.html
        self.assertTrue(result["gallery"])
        self.assertIn("index.html", fake.stored)
        self.assertTrue(fake.quit_called)

    def test_password_protected_roundtrip(self) -> None:
        self.service.save_config({"host": "h", "port": 21, "user": "u", "password": "topsecret"})
        # get_config desprotege; o valor cru no banco não é o texto puro em Windows.
        self.assertEqual("topsecret", self.service.get_config()["password"])


class AccessExportTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.pairing_service = PairingService(self.db)
        self.export_service = ExportService(self.db, self.pairing_service)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_write_csv_bundle(self) -> None:
        from src.services.access_export import write_csv_bundle

        tables = {"Jogadores": (["id", "nome"], [[1, "Ana"], [2, "Joao"]])}
        bundle = write_csv_bundle(tables, self.temp_dir.name)
        self.assertTrue(Path(bundle["schema_ini"]).exists())
        self.assertTrue(any(p.endswith("Jogadores.csv") for p in bundle["csv_paths"]))
        schema = Path(bundle["schema_ini"]).read_text(encoding="utf-8")
        self.assertIn("[Jogadores.csv]", schema)
        self.assertIn("CharacterSet=65001", schema)

    def test_export_access_service(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        result = self.export_service.export_access(tournament_id, self.temp_dir.name)
        self.assertIn("Jogadores", result["tables"])
        self.assertIn("Classificacao", result["tables"])
        self.assertTrue(result["csv_paths"])
        self.assertTrue(Path(result["schema_ini"]).exists())
        # Sem driver ACE no ambiente de teste, o .accdb real não é gerado.
        self.assertFalse(result["driver_available"])
        self.assertIsNone(result["accdb"])


class BatchExportPhaseJTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.pairing_service = PairingService(self.db)
        self.export_service = ExportService(self.db, self.pairing_service)
        self.service = BatchExportService(self.export_service)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_available_report_keys(self) -> None:
        from src.services.batch_export import available_report_keys

        individual = available_report_keys(False)
        self.assertIn("classificacao", individual)
        self.assertNotIn("equipes", individual)
        self.assertIn("equipes", available_report_keys(True))

    def test_run_batch_generates_files_and_isolates(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        dest = Path(self.temp_dir.name) / "lote"
        result = self.service.run_batch(
            tournament_id,
            ["classificacao", "jogadores", "equipes"],
            ["csv", "xlsx"],
            dest,
        )
        # classificacao+jogadores × csv+xlsx = 4 arquivos; equipes é só de equipes -> ignorado.
        self.assertEqual(4, result["count"])
        self.assertEqual(4, len(result["generated"]))
        self.assertTrue(all(os.path.exists(path) for path in result["generated"]))
        self.assertTrue(any("equipes" in item.lower() for item in result["skipped"]))

    def test_run_batch_skips_incompatible_format(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        dest = Path(self.temp_dir.name) / "lote2"
        # 'jogadores' não suporta HTML -> nada gerado, reportado em skipped.
        result = self.service.run_batch(tournament_id, ["jogadores"], ["html"], dest)
        self.assertEqual(0, result["count"])
        self.assertTrue(result["skipped"])

    def test_run_batch_requires_selection(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        with self.assertRaises(AppError):
            self.service.run_batch(tournament_id, [], ["csv"], self.temp_dir.name)


class ReportsTest(CoreServiceTestCase):
    def test_export_club_report_includes_administrative_summary(self) -> None:
        self.db.save_club("Clube Teste", city="Teresina", email="contato@example.com")
        self.member_service.create_member(
            {
                "name": "Aluno Ativo",
                "rating": "1700",
                "category": "Sub-18",
                "member_type": "aluno",
                "status": "active",
            }
        )
        output_path = Path(self.temp_dir.name) / "clube.csv"

        self.export_service.export_club_report(output_path)
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertIn("Clube Teste", content)
        self.assertIn("Indicadores", content)
        self.assertIn("Membros ativos", content)
        self.assertIn("Ranking interno", content)
        self.assertIn("Aluno Ativo", content)

    def test_export_tournaments_period_report_filters_by_date(self) -> None:
        self.db.create_tournament("Torneio maio", start_date="2026-05-10")
        self.db.create_tournament("Torneio junho", start_date="2026-06-10")
        output_path = Path(self.temp_dir.name) / "torneios_periodo.csv"

        self.export_service.export_tournaments_period_report(
            output_path,
            start_date="2026-05-01",
            end_date="2026-05-31",
        )
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertIn("Torneios por periodo", content)
        self.assertIn("Torneio maio", content)
        self.assertNotIn("Torneio junho", content)

    def test_export_site_creates_static_html_package(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Jogador A",
            rating=1800,
            club="Clube",
            category="ABS",
        )
        self.db.create_player(
            self.tournament_id,
            name="Jogador B",
            rating=1700,
            club="Clube",
            category="ABS",
        )
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(round_data["id"])[0]
        self.db.update_pairing_result(pairing["id"], "1-0")
        self.service.close_round(self.tournament_id, round_data["id"])

        output_dir = Path(self.temp_dir.name) / "site"
        index_path = self.export_service.export_site(self.tournament_id, output_dir)
        html = index_path.read_text(encoding="utf-8")

        self.assertTrue(index_path.exists())
        self.assertTrue((output_dir / "styles.css").exists())
        self.assertIn("Torneio teste", html)
        self.assertIn("Classificacao", html)
        self.assertIn("Rodada 1", html)

    def test_game_statistics_counts_results(self) -> None:
        self._create_players(8)  # 4 partidas na rodada 1
        self._play_first_round(["1-0", "0-1", "1/2-1/2", "1F-0F"])

        sections = self.export_service._game_statistics_sections(self.tournament_id)
        summary = {row[0]: row[1] for row in sections[0][2]}
        distribution = {row[0]: row[1] for row in sections[1][2]}

        self.assertEqual(summary["Partidas jogadas (tabuleiro)"], 3)  # WO nao conta
        self.assertEqual(summary["WO / forfait"], 1)
        self.assertEqual(summary["Byes"], 0)
        self.assertEqual(summary["Total de pareamentos fechados"], 4)
        self.assertEqual(distribution["Vitorias de brancas"], 1)
        self.assertEqual(distribution["Empates"], 1)
        self.assertEqual(distribution["Vitorias de pretas"], 1)

    def test_player_cards_summary_and_round_by_round(self) -> None:
        self._create_players(8)
        self._play_first_round(["1-0", "0-1", "1/2-1/2", "1-0"])

        sections = self.export_service._player_cards_sections(self.tournament_id)
        summary_rows = sections[0][2]
        game_rows = sections[1][2]

        self.assertEqual(len(summary_rows), 8)   # uma linha de resumo por jogador
        self.assertEqual(len(game_rows), 8)       # 8 jogadores, 1 partida cada na rodada 1
        # V + E + D somados entre todos = 8 jogadores com 1 jogo cada.
        total_results = sum(int(row[6]) + int(row[7]) + int(row[8]) for row in summary_rows)
        self.assertEqual(total_results, 8)

    def test_export_player_cards_writes_file(self) -> None:
        self._create_players(8)
        self._play_first_round(["1-0", "0-1", "1/2-1/2", "1-0"])

        path = Path(self.temp_dir.name) / "fichas.csv"
        self.export_service.export_player_cards(self.tournament_id, path)
        self.assertTrue(path.exists())
        content = path.read_text(encoding="utf-8-sig")
        self.assertIn("Resumo por jogador", content)
        self.assertIn("Resultados rodada a rodada", content)

    def test_standings_default_columns_unchanged(self) -> None:
        self._create_players(4)
        _title, headers, rows = self.export_service._standings_section(self.tournament_id)
        self.assertEqual(
            headers,
            [
                "Pos", "Nome", "Categoria", "Categoria idade", "Categoria rating",
                "Tags premiacao", "Pts", "Buchholz", "Buchholz M", "SB",
                "Vitorias", "Performance", "Rating", "Clube",
            ],
        )
        self.assertEqual(len(rows), 4)
        self.assertTrue(all(len(row) == 14 for row in rows))

    def test_standings_layout_changes_columns(self) -> None:
        self._create_players(4)
        self.list_layout_service.save_columns(self.tournament_id, ["position", "name", "points"], "standings")

        _title, headers, rows = self.export_service._standings_section(self.tournament_id)
        self.assertEqual(headers, ["Pos", "Nome", "Pts"])
        self.assertTrue(all(len(row) == 3 for row in rows))

    def test_list_layout_reset_returns_to_default(self) -> None:
        self._create_players(2)
        self.list_layout_service.save_columns(self.tournament_id, ["position", "name"], "standings")
        self.list_layout_service.reset(self.tournament_id, "standings")
        # Apos reset, a classificacao volta a usar as colunas padrao.
        _title, headers, _rows = self.export_service._standings_section(self.tournament_id)
        self.assertEqual(len(headers), 14)

    def test_list_layout_persists_column_widths(self) -> None:
        self._create_players(2)
        saved = self.list_layout_service.save_columns(
            self.tournament_id,
            [{"key": "position", "width": 8}, {"key": "name", "width": 30}],
            "standings",
        )
        self.assertEqual(saved, [{"key": "position", "width": 8}, {"key": "name", "width": 30}])
        selected = self.list_layout_service.get_columns(self.tournament_id, "standings")["selected"]
        self.assertEqual({spec["key"]: spec["width"] for spec in selected}, {"position": 8, "name": 30})

    def test_standings_column_widths_applied_in_xlsx(self) -> None:
        from openpyxl import load_workbook

        self._create_players(4)
        self.list_layout_service.save_columns(
            self.tournament_id,
            [{"key": "position", "width": 7}, {"key": "name", "width": 40}],
            "standings",
        )
        path = Path(self.temp_dir.name) / "classificacao.xlsx"
        self.export_service.export_standings(self.tournament_id, path)

        worksheet = load_workbook(path).active
        self.assertEqual(worksheet.column_dimensions["A"].width, 7)
        self.assertEqual(worksheet.column_dimensions["B"].width, 40)

    def test_rating_fee_report_counts_bases_rated_unrated_and_exports(self) -> None:
        self.db.save_tournament_settings(
            self.tournament_id,
            {
                "rating_fee_fide": "2",
                "rating_fee_cbx": "3",
                "rating_fee_lbx": "4",
            },
        )
        self.db.create_player(
            self.tournament_id,
            name="Multibase rated",
            fide_id="1001",
            cbx_id="2001",
            lbx_id="3001",
            rating=1800,
            national_rating=1900,
            international_rating=2000,
        )
        self.db.create_player(
            self.tournament_id,
            name="FIDE sem rating",
            fide_id="1002",
            cbx_id="2002",
            national_rating=1600,
        )
        self.db.create_player(self.tournament_id, name="LBX sem rating", lbx_id="3002")
        self.db.create_player(self.tournament_id, name="Rated sem ID", rating=1400)

        summary = self.export_service.rating_fee_summary(self.tournament_id)
        bases = {item["base"]: item for item in summary["bases"]}

        self.assertEqual(4, summary["players_count"])
        self.assertEqual(3, summary["players_with_base"])
        self.assertEqual(1, summary["players_without_base"])
        self.assertEqual(3, summary["rated_players"])
        self.assertEqual(1, summary["unrated_players"])
        self.assertEqual({"identified": 2, "rated": 1, "unrated": 1}, {
            key: bases["FIDE"][key] for key in ("identified", "rated", "unrated")
        })
        self.assertEqual(6.0, bases["CBX"]["subtotal"])
        self.assertEqual(8.0, bases["LBX"]["subtotal"])
        self.assertEqual(18.0, summary["total"])

        xlsx_path = Path(self.temp_dir.name) / "taxas_rating.xlsx"
        pdf_path = Path(self.temp_dir.name) / "taxas_rating.pdf"
        self.export_service.export_rating_fee_report(self.tournament_id, xlsx_path)
        self.export_service.export_rating_fee_report(self.tournament_id, pdf_path)

        from openpyxl import load_workbook
        from pypdf import PdfReader

        workbook = load_workbook(xlsx_path, read_only=True)
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)
        self.assertEqual(["Resumo de taxas de rating", "Taxas por base", "Inscritos por base"], workbook.sheetnames)
        workbook.close()
        self.assertIn("Total das taxas", pdf_text)
        self.assertIn("R$ 18,00", pdf_text)
        with self.assertRaisesRegex(AppError, "XLSX ou PDF"):
            self.export_service.export_rating_fee_report(
                self.tournament_id,
                Path(self.temp_dir.name) / "taxas_rating.csv",
            )

    def test_export_club_portal_creates_static_html_package(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Escola Portal",
                "kind": "school",
                "city": "Curitiba",
                "active": 1,
            },
            club_id=None,
        )
        class_id = self.club_service.save_class(
            {
                "club_id": club_id,
                "name": "Turma Portal",
                "teacher": "Professora Portal",
                "weekday": "Sabado",
                "time": "09:00",
                "location": "Sala 1",
                "active": 1,
            }
        )
        self.member_service.create_member(
            {
                "name": "Aluno Portal",
                "club_id": club_id,
                "class_id": class_id,
                "rating": "1500",
                "category": "Sub-12",
                "member_type": "aluno",
                "status": "active",
            }
        )
        self.training_service.save_session(
            {
                "club_id": club_id,
                "class_id": class_id,
                "title": "Aula de estrategia",
                "session_type": "aula",
                "session_date": "2999-08-05",
                "start_time": "09:00",
                "instructor": "Professora Portal",
                "location": "Sala 1",
                "objective": "Plano de meio-jogo",
                "homework": "Resolver dois diagramas.",
                "status": "planned",
            }
        )
        self.event_service.save_event(
            {
                "club_id": club_id,
                "title": "Festival Portal",
                "event_type": "social",
                "event_date": "2999-08-10",
                "start_time": "10:00",
                "location": "Salao",
                "status": "confirmed",
                "notes": "Comunicado aos responsaveis.",
            }
        )
        self.db.create_tournament(
            "Torneio Portal",
            club_id=club_id,
            class_id=class_id,
            start_date="2999-08-20",
            location="Salao",
        )

        output_dir = Path(self.temp_dir.name) / "portal"
        index_path = self.export_service.export_club_portal(output_dir, club_id=club_id, class_id=class_id)
        html = index_path.read_text(encoding="utf-8")

        self.assertTrue(index_path.exists())
        self.assertTrue((output_dir / "styles.css").exists())
        self.assertIn("Turma Portal", html)
        self.assertIn("Comunicados", html)
        self.assertIn("Festival Portal", html)
        self.assertIn("Aula de estrategia", html)
        self.assertIn("Resolver dois diagramas.", html)
        self.assertIn("Aluno Portal", html)
        self.assertIn("Ranking interno", html)
        self.assertIn("Torneio Portal", html)

    def test_certificate_service_exports_tournament_pdf(self) -> None:
        player_a = self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube A",
            category="Sub-18",
        )
        player_c = self.db.create_player(
            self.tournament_id,
            name="Carla",
            surname="Lima",
            rating=1700,
            club="Clube B",
            category="Absoluto",
        )
        self.db.create_player(
            self.tournament_id,
            name="Diego",
            surname="Rocha",
            rating=1600,
            club="Clube B",
            category="Absoluto",
        )

        category_recipients = self.certificate_service.tournament_recipients(
            self.tournament_id,
            certificate_type="category_award",
            top_n=1,
        )
        selected_path = Path(self.temp_dir.name) / "diplomas.pdf"
        result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            selected_path,
            certificate_type="participation",
            player_ids=[player_a, player_c],
        )

        self.assertEqual(result["exported"], 2)
        self.assertEqual(len(result["verification_codes"]), 2)
        issuance = self.certificate_service.verify_issuance(result["verification_codes"][0])
        self.assertEqual(issuance["context_type"], "tournament")
        self.assertEqual(issuance["source_title"], "Torneio teste")
        self.assertTrue(selected_path.exists())
        self.assertEqual(selected_path.read_bytes()[:4], b"%PDF")
        self.assertEqual({item["category"] for item in category_recipients}, {"Sub-18", "Absoluto"})
        self.assertEqual(len(category_recipients), 2)

    def test_certificate_export_without_issuance_record_does_not_leave_pdf(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube A",
            category="Sub-18",
        )
        output_path = Path(self.temp_dir.name) / "falha_registro.pdf"

        with mock.patch.object(
            self.db,
            "create_certificate_issuances",
            side_effect=sqlite3.OperationalError("registro falhou"),
        ):
            with self.assertRaises(sqlite3.OperationalError):
                self.certificate_service.export_tournament_certificates(
                    self.tournament_id,
                    output_path,
                    certificate_type="participation",
                )

        self.assertFalse(output_path.exists())
        self.assertEqual([], list(output_path.parent.glob("*.tmp.pdf")))

    def test_certificate_export_replace_failure_removes_issuance_records_and_temp_pdf(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        output_path = Path(self.temp_dir.name) / "falha_substituir.pdf"

        with mock.patch.object(Path, "replace", side_effect=OSError("replace falhou")):
            with self.assertRaises(OSError):
                self.certificate_service.export_tournament_certificates(
                    self.tournament_id,
                    output_path,
                    certificate_type="participation",
                )
        with self.db.connect() as connection:
            issuance_count = connection.execute("SELECT COUNT(*) FROM certificate_issuances").fetchone()[0]

        self.assertEqual(issuance_count, 0)
        self.assertFalse(output_path.exists())
        self.assertEqual([], list(output_path.parent.glob("*.tmp.pdf")))

    def test_certificate_templates_can_be_saved_previewed_and_used(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube A",
            category="Sub-18",
        )
        logo_path = Path(self.temp_dir.name) / "logo.png"
        logo_path.write_bytes(
            base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
            )
        )
        background_path = Path(self.temp_dir.name) / "fundo.png"
        background_path.write_bytes(logo_path.read_bytes())
        secondary_logo_path = Path(self.temp_dir.name) / "logo_secundario.png"
        secondary_logo_path.write_bytes(logo_path.read_bytes())
        template_id = self.certificate_service.create_template(
            {
                "name": "Modelo personalizado",
                "certificate_type": "overall_award",
                "title_template": "Diploma {torneio}",
                "body_template": "{nome} ficou em {posicao} com {pontos} ponto(s).",
                "footer_template": "{local} {periodo}",
                "orientation": "portrait",
                "signature_left": "Direcao",
                "signature_right": "Arbitro",
                "logo_path": str(logo_path),
                "background_image_path": str(background_path),
                "background_opacity": "35",
                "secondary_logo_path": str(secondary_logo_path),
                "primary_color": "#0F766E",
                "accent_color": "#F59E0B",
                "title_font_size": "30",
                "body_font_size": "16",
                "footer_font_size": "9",
            }
        )

        self.certificate_service.update_template(
            template_id,
            {
                "name": "Modelo personalizado",
                "certificate_type": "overall_award",
                "title_template": "Diploma especial {torneio}",
                "body_template": "{nome} ficou em {posicao} com {pontos} ponto(s).",
                "footer_template": "{local} {periodo}",
                "orientation": "portrait",
                "signature_left": "Direcao",
                "signature_right": "Arbitro",
                "logo_path": str(logo_path),
                "background_image_path": str(background_path),
                "background_opacity": "0.42",
                "secondary_logo_path": str(secondary_logo_path),
                "primary_color": "#0F766E",
                "accent_color": "#F59E0B",
                "title_font_size": "30",
                "body_font_size": "16",
                "footer_font_size": "9",
            },
        )
        preview = self.certificate_service.preview_template(
            self.tournament_id,
            template_id,
            top_n=1,
        )
        output_path = Path(self.temp_dir.name) / "modelo_personalizado.pdf"
        result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            output_path,
            template_id=template_id,
            top_n=1,
        )

        template = self.db.get_certificate_template(template_id)
        self.assertEqual(template["orientation"], "portrait")
        self.assertEqual(template["primary_color"], "#0F766E")
        self.assertEqual(template["accent_color"], "#F59E0B")
        self.assertEqual(template["background_image_path"], str(background_path))
        self.assertEqual(template["background_opacity"], 0.42)
        self.assertEqual(template["secondary_logo_path"], str(secondary_logo_path))
        self.assertEqual(template["title_font_size"], 30)
        self.assertEqual(result["exported"], 1)
        self.assertIn("Diploma especial Torneio teste", preview["title"])
        self.assertIn("Ana Silva", preview["body"])
        self.assertTrue(output_path.exists())
        self.assertEqual(output_path.read_bytes()[:4], b"%PDF")

    def test_default_certificate_background_assets_and_templates_export(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube B",
            category="Absoluto",
        )
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Fundo",
                "rating": "1700",
                "category": "Sub-18",
                "member_type": "aluno",
                "status": "active",
            }
        )

        for _label, relative_path, _opacity in database_module.CERTIFICATE_BACKGROUND_PRESETS:
            path = database_module.BASE_DIR / relative_path
            self.assertTrue(path.exists(), relative_path)
            self.assertEqual(path.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")

        templates_by_name = {
            str(template["name"]): template
            for template in self.certificate_service.list_templates(active_only=True)
        }
        expected_templates = {
            "Xadrez classico - Participacao",
            "Xadrez escolar - Membro aluno",
            "Xadrez premium - Premiacao geral",
            "Tabuleiro sutil - Premiacao por categoria",
            "Pecas marca d'agua - Ranking interno",
        }
        self.assertTrue(expected_templates.issubset(templates_by_name))

        participation_path = Path(self.temp_dir.name) / "participacao_fundo.pdf"
        overall_path = Path(self.temp_dir.name) / "premiacao_fundo.pdf"
        category_path = Path(self.temp_dir.name) / "categoria_fundo.pdf"
        member_path = Path(self.temp_dir.name) / "membro_fundo.pdf"
        ranking_path = Path(self.temp_dir.name) / "ranking_fundo.pdf"

        participation_result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            participation_path,
            template_id=int(templates_by_name["Xadrez classico - Participacao"]["id"]),
        )
        overall_result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            overall_path,
            template_id=int(templates_by_name["Xadrez premium - Premiacao geral"]["id"]),
            top_n=1,
        )
        category_result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            category_path,
            template_id=int(templates_by_name["Tabuleiro sutil - Premiacao por categoria"]["id"]),
            top_n=1,
        )
        member_result = self.certificate_service.export_member_certificates(
            member_path,
            template_id=int(templates_by_name["Xadrez escolar - Membro aluno"]["id"]),
            member_ids=[member_id],
        )
        ranking_result = self.certificate_service.export_ranking_certificates(
            ranking_path,
            template_id=int(templates_by_name["Pecas marca d'agua - Ranking interno"]["id"]),
            top_n=1,
        )

        self.assertEqual(participation_result["exported"], 2)
        self.assertEqual(overall_result["exported"], 1)
        self.assertEqual(category_result["exported"], 2)
        self.assertEqual(member_result["exported"], 1)
        self.assertEqual(ranking_result["exported"], 1)
        for path in (participation_path, overall_path, category_path, member_path, ranking_path):
            self.assertTrue(path.exists())
            self.assertEqual(path.read_bytes()[:4], b"%PDF")

    def test_default_certificate_logo_assets_export_as_primary_and_secondary(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        logo_dir = database_module.BASE_DIR / "assets" / "certificates" / "logos"
        primary_logo = logo_dir / "albericus_knight.png"
        secondary_logo = logo_dir / "clube_rook.png"
        expected_logos = {
            "albericus_knight.png",
            "clube_rook.png",
            "escola_pawn.png",
            "torneio_trophy.png",
            "selo_queen.png",
        }

        for filename in expected_logos:
            path = logo_dir / filename
            self.assertTrue(path.exists(), filename)
            content = path.read_bytes()
            self.assertEqual(content[:8], b"\x89PNG\r\n\x1a\n")
            self.assertEqual(content[25], 6)

        template_id = self.certificate_service.create_template(
            {
                "name": "Modelo com logos padrao",
                "certificate_type": "participation",
                "title_template": "Certificado",
                "body_template": "Certificamos que {nome} participou do torneio {torneio}.",
                "footer_template": "{local} - {periodo}",
                "orientation": "landscape",
                "signature_left": "Direcao",
                "signature_right": "Arbitragem",
                "logo_path": str(primary_logo),
                "secondary_logo_path": str(secondary_logo),
                "primary_color": "#1E3A8A",
                "accent_color": "#C8A24A",
                "title_font_size": "30",
                "body_font_size": "16",
                "footer_font_size": "9",
            }
        )
        output_path = Path(self.temp_dir.name) / "logos_padrao.pdf"
        result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            output_path,
            template_id=template_id,
        )

        self.assertEqual(result["exported"], 1)
        self.assertTrue(output_path.exists())
        self.assertEqual(output_path.read_bytes()[:4], b"%PDF")

    def test_certificate_service_exports_local_verification_site(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube B",
            category="Absoluto",
        )
        pdf_path = Path(self.temp_dir.name) / "diplomas.pdf"
        result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            pdf_path,
        )
        revoked_issuance = self.certificate_service.verify_issuance(result["verification_codes"][0])
        self.certificate_service.revoke_issuance(int(revoked_issuance["id"]), notes="Reemitido")

        site_path = Path(self.temp_dir.name) / "verificador"
        exported_path = self.certificate_service.export_verification_site(site_path)
        html = exported_path.read_text(encoding="utf-8")

        self.assertEqual(exported_path, site_path.with_suffix(".html"))
        self.assertIn("Verificador de diplomas Albericus", html)
        self.assertIn(result["verification_codes"][0], html)
        self.assertIn(result["verification_codes"][1], html)
        self.assertIn("Ana Silva", html)
        self.assertIn("Bruno Souza", html)
        self.assertIn("Revogado", html)
        self.assertIn("Ativo", html)
        self.assertIn("Torneio teste", html)
        self.assertNotIn(str(pdf_path), html)

    def test_scoresheets_pdf_exports_121_players_under_five_seconds(self) -> None:
        for index in range(121):
            self.db.create_player(
                self.tournament_id,
                name=f"Jogador {index + 1:03d}",
                rating=2400 - index,
                club="Clube",
            )
        round_data = self.service.generate_next_round(self.tournament_id)
        output_path = Path(self.temp_dir.name) / "sumulas_121.pdf"

        started_at = perf_counter()
        self.export_service.export_scoresheets(int(round_data["id"]), output_path)
        elapsed = perf_counter() - started_at

        from pypdf import PdfReader

        self.assertEqual(60, len(PdfReader(output_path).pages))
        self.assertLess(elapsed, 5.0)

    def test_table_cards_pdf_exports_200_cards_under_five_seconds(self) -> None:
        output_path = Path(self.temp_dir.name) / "cartoes_200.pdf"

        started_at = perf_counter()
        self.export_service.export_table_cards(output_path, 1, 200)
        elapsed = perf_counter() - started_at

        from pypdf import PdfReader

        self.assertEqual(50, len(PdfReader(output_path).pages))
        self.assertLess(elapsed, 5.0)

    def test_initial_player_list_pdf_uses_trf_start_rank_and_signature_column(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana Ausente",
            rating=1800,
            international_rating=2200,
            club="Clube A",
            category="ABS",
            player_status="absent",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno Presente",
            rating=2100,
            club="Clube B",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Carla Presente",
            rating=1900,
            club="Clube C",
            category="FEM",
        )
        output_path = Path(self.temp_dir.name) / "lista_chamada.pdf"

        self.export_service.export_initial_player_list(self.tournament_id, output_path)

        from pypdf import PdfReader

        text = "\n".join(page.extract_text() or "" for page in PdfReader(output_path).pages)
        self.assertEqual(output_path.read_bytes()[:4], b"%PDF")
        self.assertIn("Inicial", text)
        self.assertIn("Jogador", text)
        self.assertIn("Rating", text)
        self.assertIn("Clube", text)
        self.assertIn("Categoria", text)
        self.assertIn("Assinatura", text)
        self.assertLess(text.index("Ana Ausente"), text.index("Bruno Presente"))
        self.assertLess(text.index("Bruno Presente"), text.index("Carla Presente"))

    def test_initial_player_list_pdf_exports_801_players_under_five_seconds(self) -> None:
        for index in range(801):
            self.db.create_player(
                self.tournament_id,
                name=f"Jogador {index + 1:03d}",
                rating=2600 - index,
                club="Clube",
                category="ABS",
            )
        output_path = Path(self.temp_dir.name) / "lista_chamada_801.pdf"

        started_at = perf_counter()
        self.export_service.export_initial_player_list(self.tournament_id, output_path)
        elapsed = perf_counter() - started_at

        from pypdf import PdfReader

        reader = PdfReader(output_path)
        self.assertEqual(22, len(reader.pages))
        self.assertLess(elapsed, 5.0)

if __name__ == "__main__":
    unittest.main()
