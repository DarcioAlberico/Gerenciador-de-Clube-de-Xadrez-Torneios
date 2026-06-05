from __future__ import annotations

import base64
import ftplib
import hashlib
import json
import os
import sqlite3
import tempfile
import unittest
from datetime import date
from pathlib import Path
from time import perf_counter
from typing import Any
from unittest import mock
from urllib.parse import parse_qs, urlparse

from src.core import database as database_module
from src.core.database import APP_DATA_DIR_ENV_VAR, Database, resolve_app_data_dir
from src.core.services import (
    AppError,
    BatchExportService,
    CertificateService,
    ChessResultsService,
    ClockIntegrationService,
    ClubService,
    DashboardService,
    EventService,
    ExerciseService,
    ExportService,
    FinanceService,
    GuardianService,
    ImportService,
    InternalRatingService,
    InventoryService,
    LearningLevelService,
    ListLayoutService,
    MemberService,
    NormAssistantService,
    OfficialRatingService,
    PairingService,
    PhotoAlbumService,
    PrizeService,
    QRResultService,
    SecurityService,
    SyncService,
    TeamService,
    TournamentService,
    TrainingService,
)
from src.services.federation_exporters import FederationExporterRegistry, TRF16Exporter
from src.services.result_server import LocalResultServer
from tests.fixtures import load_tournament_fixture


class RolePermissionMatrixTest(unittest.TestCase):
    """Spec §14.1 — perfis capitao/jogador e matriz de permissões."""

    def test_new_roles_exist_in_operator_roles(self) -> None:
        from src.services.constants import OPERATOR_ROLES
        self.assertIn("capitao", OPERATOR_ROLES)
        self.assertIn("jogador", OPERATOR_ROLES)
        # 'publico' propositalmente NÃO é login role (portal anônimo).
        self.assertNotIn("publico", OPERATOR_ROLES)

    def test_capitao_can_submit_lineup_and_request_substitution(self) -> None:
        from src.config.permissions import get_permissions_for_role
        perms = get_permissions_for_role("capitao")
        self.assertIn("team_lineup_submit", perms)
        self.assertIn("team_substitution_request", perms)
        self.assertIn("own_data_read", perms)
        # NÃO pode mexer em config nem em outros membros
        self.assertNotIn("settings_write", perms)
        self.assertNotIn("member_write", perms)
        self.assertNotIn("tournament_write", perms)

    def test_jogador_has_only_read_own_and_presence(self) -> None:
        from src.config.permissions import get_permissions_for_role
        perms = get_permissions_for_role("jogador")
        self.assertEqual({"own_data_read", "presence_confirm"}, set(perms))

    def test_admin_inherits_all_new_permissions(self) -> None:
        from src.config.permissions import get_permissions_for_role
        perms = set(get_permissions_for_role("admin"))
        for action in (
            "team_lineup_submit",
            "team_substitution_request",
            "own_data_read",
            "presence_confirm",
        ):
            self.assertIn(action, perms)

    def test_arbiter_can_act_on_lineups_and_substitutions(self) -> None:
        from src.config.permissions import get_permissions_for_role
        perms = get_permissions_for_role("arbiter")
        self.assertIn("team_lineup_submit", perms)
        self.assertIn("team_substitution_request", perms)


class TeamRosterPolicyValidatorTest(unittest.TestCase):
    """Spec §6.2 — validador de policy de escalação retorna warnings."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base_path = Path(self.temp_dir.name)
        self.db = Database(base_path / "albericus.db", backup_dir=base_path / "backups")
        self.team_service = TeamService(self.db)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_validator_returns_empty_when_no_team_tournament(self) -> None:
        # Torneio individual recém-criado — sem lineups, sem warnings.
        tournament_id = self.db.create_tournament("Vazio")
        issues = self.team_service.validate_roster_policy(tournament_id)
        self.assertEqual([], issues)

    def test_validator_signature_accepts_round_filter(self) -> None:
        tournament_id = self.db.create_tournament("Filtrado")
        # Não deve levantar exceção mesmo sem dados.
        self.assertEqual(
            [],
            self.team_service.validate_roster_policy(tournament_id, round_id=999),
        )

    def test_migration_added_team_rating_tolerance_column(self) -> None:
        tournament_id = self.db.create_tournament("Tolerancia")
        settings = self.db.get_tournament_settings(tournament_id) or {}
        # Coluna existe e default é 0.
        self.assertEqual(0, int(settings.get("team_rating_tolerance", -1) or 0))


class TournamentFixtureTest(unittest.TestCase):
    """Smoke da pasta tests/fixtures/tournaments/ — garante que o cenário
    carrega e o motor produz um estado consistente."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base_path = Path(self.temp_dir.name)
        self.db = Database(base_path / "albericus.db", backup_dir=base_path / "backups")
        self.service = PairingService(self.db)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_individual_8_players_3_rounds_fixture(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.service
        )

        players = self.db.list_players(tournament_id, active_only=False)
        rounds = self.db.list_rounds(tournament_id)
        standings = self.service.standings(tournament_id)

        self.assertEqual(8, len(players))
        self.assertEqual(3, len(rounds))
        self.assertTrue(
            all(r["status"] == "closed" for r in rounds),
            "Todas as rodadas da fixture devem ficar fechadas.",
        )
        self.assertEqual(8, len(standings))
        # Soma de pontos = total de mesas (4 por rodada) × 3 rodadas × 1.0
        total_points = sum(float(row.get("points") or 0) for row in standings)
        self.assertAlmostEqual(12.0, total_points, places=2)

        # result_states_summary cobre todas as mesas como "locked" (round closed,
        # com resultado) — exercita o derive_pairing_state na ponta.
        states = self.service.result_states_summary(tournament_id)
        self.assertEqual(12, states["locked"])
        self.assertEqual(0, states["empty"])
        self.assertEqual(0, states["published"])


class DerivePairingStateTest(unittest.TestCase):
    """Função pura — exercita a matriz de estados sem precisar de banco."""

    def test_empty_when_no_result(self) -> None:
        self.assertEqual(
            PairingService.derive_pairing_state(result="", round_closed=False), "empty"
        )

    def test_submitted_when_pending_submission_no_result(self) -> None:
        self.assertEqual(
            PairingService.derive_pairing_state(
                result="", round_closed=False, has_pending_submission=True
            ),
            "submitted",
        )

    def test_published_when_result_open_round(self) -> None:
        self.assertEqual(
            PairingService.derive_pairing_state(result="1-0", round_closed=False),
            "published",
        )

    def test_locked_when_round_closed_with_result(self) -> None:
        self.assertEqual(
            PairingService.derive_pairing_state(result="1-0", round_closed=True),
            "locked",
        )

    def test_corrected_trumps_locked(self) -> None:
        self.assertEqual(
            PairingService.derive_pairing_state(
                result="0-1", round_closed=True, has_correction=True
            ),
            "corrected",
        )

    def test_corrected_trumps_submitted(self) -> None:
        self.assertEqual(
            PairingService.derive_pairing_state(
                result="1-0", round_closed=False,
                has_pending_submission=True, has_correction=True,
            ),
            "corrected",
        )

    def test_closed_round_without_result_is_empty(self) -> None:
        # Mesa sem resultado em rodada fechada ainda é "empty" (não locked).
        self.assertEqual(
            PairingService.derive_pairing_state(result="", round_closed=True),
            "empty",
        )


class PairingServiceTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base_path = Path(self.temp_dir.name)
        self.backup_dir = base_path / "backups"
        self.db = Database(base_path / "albericus.db", backup_dir=self.backup_dir)
        self.service = PairingService(self.db)
        self.club_service = ClubService(self.db)
        self.clock_integration_service = ClockIntegrationService(self.db)
        self.dashboard_service = DashboardService(self.db)
        self.guardian_service = GuardianService(self.db)
        self.member_service = MemberService(self.db)
        self.learning_level_service = LearningLevelService(self.db)
        self.training_service = TrainingService(self.db)
        self.exercise_service = ExerciseService(self.db)
        self.tournament_service = TournamentService(self.db)
        self.team_service = TeamService(self.db)
        self.event_service = EventService(self.db)
        self.import_service = ImportService(self.db)
        self.official_rating_service = OfficialRatingService(self.db)
        self.internal_rating_service = InternalRatingService(self.db)
        self.inventory_service = InventoryService(self.db)
        self.security_service = SecurityService(self.db)
        self.sync_service = SyncService(self.db)
        self.qr_result_service = QRResultService(self.db, self.service)
        self.export_service = ExportService(self.db, self.service)
        self.certificate_service = CertificateService(self.db, self.service)
        self.finance_service = FinanceService(self.db)
        self.prize_service = PrizeService(self.db)
        self.norm_assistant_service = NormAssistantService(self.db)
        self.list_layout_service = ListLayoutService(self.db)
        self.tournament_id = self.db.create_tournament("Torneio teste", rounds_count=5)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_user_data_path_resolver_supports_env_and_platform_defaults(self) -> None:
        base_path = Path(self.temp_dir.name)

        self.assertEqual(
            resolve_app_data_dir(
                {APP_DATA_DIR_ENV_VAR: str(base_path / "custom")},
                platform="win32",
                home=base_path / "home",
            ),
            base_path / "custom",
        )
        self.assertEqual(
            resolve_app_data_dir(
                {"LOCALAPPDATA": str(base_path / "local")},
                platform="win32",
                home=base_path / "home",
            ),
            base_path / "local" / "Albericus",
        )
        self.assertEqual(
            resolve_app_data_dir(
                {"XDG_DATA_HOME": str(base_path / "xdg")},
                platform="linux",
                home=base_path / "home",
            ),
            base_path / "xdg" / "Albericus",
        )
        self.assertEqual(
            resolve_app_data_dir({}, platform="darwin", home=base_path / "home"),
            base_path / "home" / "Library" / "Application Support" / "Albericus",
        )

    def test_default_database_uses_user_data_dir_and_copies_legacy_database(self) -> None:
        base_path = Path(self.temp_dir.name)
        legacy_db_path = base_path / "legacy" / "data" / "albericus.db"
        legacy_backup_dir = base_path / "legacy" / "backups"
        legacy_db = Database(legacy_db_path, backup_dir=legacy_backup_dir)
        legacy_db.create_tournament("Torneio legado")
        legacy_backup_path = legacy_db.backup("manual")

        user_data_dir = base_path / "user-data"
        with (
            mock.patch.dict(os.environ, {APP_DATA_DIR_ENV_VAR: str(user_data_dir)}),
            mock.patch.object(database_module, "LEGACY_DB_PATH", legacy_db_path),
            mock.patch.object(database_module, "LEGACY_BACKUP_DIR", legacy_backup_dir),
        ):
            migrated_db = Database()

        self.assertEqual(migrated_db.db_path, user_data_dir / "data" / "albericus.db")
        self.assertEqual(migrated_db.backup_dir, user_data_dir / "backups")
        self.assertTrue(migrated_db.db_path.exists())
        self.assertTrue((migrated_db.backup_dir / legacy_backup_path.name).exists())
        self.assertTrue(
            any(tournament["name"] == "Torneio legado" for tournament in migrated_db.list_tournaments())
        )

    def _create_players(self, total: int) -> list[int]:
        player_ids = []
        for index in range(total):
            player_ids.append(
                self.db.create_player(
                    self.tournament_id,
                    name=f"Jogador {index + 1}",
                    rating=2000 - index * 50,
                    club="Clube",
                    category="Absoluto",
                )
            )
        return player_ids

    def _create_team_tournament(
        self,
        teams_count: int = 4,
        boards_count: int = 2,
    ) -> tuple[int, list[int]]:
        tournament_id = self.tournament_service.create_tournament(
            {
                "name": "Torneio por equipes",
                "competition_type": "team",
                "rounds_count": "5",
                "bye_points": "1",
            }
        )
        self.tournament_service.save_profile(
            tournament_id,
            {
                "name": "Torneio por equipes",
                "competition_type": "team",
                "scope": "standalone",
                "rounds_count": "5",
                "bye_points": "1",
            },
            {
                "team_boards_count": str(boards_count),
                "team_match_win_points": "2",
                "team_match_draw_points": "1",
                "team_match_loss_points": "0",
                "team_pairing_method": "swiss",
                "team_standing_primary": "match_points",
                "team_standing_secondary": "game_points",
            },
            [],
        )
        team_ids = []
        for team_index in range(teams_count):
            team_id = self.team_service.create_team(
                tournament_id,
                {
                    "name": f"Equipe {team_index + 1}",
                    "club": f"Clube {team_index + 1}",
                    "captain": f"Capitao {team_index + 1}",
                },
            )
            team_ids.append(team_id)
            for board_number in range(1, boards_count + 1):
                player_id = self.db.create_player(
                    tournament_id,
                    name=f"Equipe {team_index + 1} Jogador {board_number}",
                    rating=2200 - team_index * 100 - board_number * 10,
                    club=f"Clube {team_index + 1}",
                )
                self.team_service.add_player(team_id, player_id, board_number=str(board_number), role="starter")
        return tournament_id, team_ids

    def _fill_decisive_results(self, round_id: int) -> None:
        for pairing in self.db.get_pairings_for_round(round_id):
            if not pairing["is_bye"]:
                self.db.update_pairing_result(pairing["id"], "1-0")

    def _create_member_win(self, rating: int = 1600) -> tuple[int, int]:
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Rating",
                "rating": str(rating),
                "member_type": "aluno",
                "status": "active",
            }
        )
        member_player_id = self.member_service.register_member_in_tournament(
            self.tournament_id,
            member_id,
        )
        self.db.create_player(
            self.tournament_id,
            name="Oponente",
            rating=1500,
            club="Clube",
            category="Absoluto",
        )
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(round_data["id"])[0]
        result = "1-0" if pairing["white_player_id"] == member_player_id else "0-1"
        self.db.update_pairing_result(pairing["id"], result)
        self.service.close_round(self.tournament_id, round_data["id"])
        return member_id, member_player_id

    def test_tournament_service_validates_creation_payload(self) -> None:
        with self.assertRaises(AppError):
            self.tournament_service.create_tournament(
                {"name": "Invalido", "rounds_count": "0", "bye_points": "1"}
            )
        with self.assertRaises(AppError):
            self.tournament_service.create_tournament(
                {"name": "Invalido", "rounds_count": "3", "bye_points": "-1"}
            )
        with self.assertRaises(AppError):
            self.tournament_service.create_tournament(
                {
                    "name": "Invalido",
                    "rounds_count": "3",
                    "bye_points": "1",
                    "start_date": "2026-05-14",
                    "end_date": "2026-05-13",
                }
            )

        tournament_id = self.tournament_service.create_tournament(
            {
                "name": "Valido",
                "scope": "standalone",
                "rounds_count": "3",
                "bye_points": "0.5",
                "start_date": "2026-05-13",
                "end_date": "2026-05-14",
            }
        )
        tournament = self.db.get_tournament(tournament_id)
        self.assertIsNotNone(tournament)
        self.assertEqual(tournament["rounds_count"], 3)
        self.assertEqual(tournament["bye_points"], 0.5)

    def test_tournament_creation_accepts_brazilian_date_format(self) -> None:
        tournament_id = self.tournament_service.create_tournament(
            {
                "name": "Data brasileira",
                "scope": "standalone",
                "rounds_count": "3",
                "bye_points": "1",
                "start_date": "13/05/2026",
                "end_date": "14/05/2026",
            }
        )

        tournament = self.db.get_tournament(tournament_id)
        self.assertIsNotNone(tournament)
        self.assertEqual(tournament["start_date"], "2026-05-13")
        self.assertEqual(tournament["end_date"], "2026-05-14")

    def test_player_competition_categories_use_tournament_year_and_rating(self) -> None:
        tournament_id = self.tournament_service.create_tournament(
            {
                "name": "Categorias automaticas",
                "scope": "standalone",
                "location": "Sao Paulo",
                "start_date": "2024-05-01",
                "rounds_count": "3",
                "bye_points": "1",
            }
        )

        player_id = self.db.create_player(
            tournament_id,
            name="Jovem Local",
            rating=1399,
            club="Sao Paulo",
            birth_date="2008-01-01",
            sex="F",
        )
        player = self.db.get_player(player_id)

        self.assertIsNotNone(player)
        self.assertEqual(player["category"], "Sub-16")
        self.assertEqual(player["age_category"], "Sub-16")
        self.assertEqual(player["rating_category"], "Sub-1400")
        self.assertEqual(player["prize_tags"], "Feminino; Melhor Local")

    def test_member_registration_recalculates_age_category_for_tournament_year(self) -> None:
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Sub",
                "birth_date": "2008-01-01",
                "rating": "1500",
                "member_type": "socio",
                "status": "active",
            }
        )
        tournament_id = self.tournament_service.create_tournament(
            {
                "name": "Torneio Sub",
                "scope": "club",
                "club_id": "1",
                "start_date": "2024-06-01",
                "rounds_count": "3",
                "bye_points": "1",
            }
        )

        player_id = self.member_service.register_member_in_tournament(tournament_id, member_id)
        player = self.db.get_player(player_id)

        self.assertIsNotNone(player)
        self.assertEqual(player["category"], "Sub-16")
        self.assertEqual(player["age_category"], "Sub-16")
        self.assertEqual(player["rating_category"], "Sub-1800")
        self.assertEqual(player["prize_tags"], "Socio do Clube")

    def test_learning_levels_can_be_assigned_to_members(self) -> None:
        default_levels = self.db.list_learning_levels(active_only=True)
        self.assertGreaterEqual(len(default_levels), 5)
        self.assertEqual(default_levels[0]["name"], "Iniciante")

        level_id = self.learning_level_service.create_level(
            {
                "name": "Pre-competitivo",
                "description": "Pronto para torneios internos.",
                "display_order": "45",
            }
        )
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Nivel",
                "member_type": "aluno",
                "status": "active",
                "learning_level_id": str(level_id),
            }
        )

        member = self.db.get_member(member_id)
        levels = self.db.list_learning_levels(active_only=False)
        custom_level = self.db.get_learning_level(level_id)

        self.assertIsNotNone(member)
        self.assertEqual(member["learning_level_id"], level_id)
        self.assertEqual(member["learning_level_name"], "Pre-competitivo")
        self.assertEqual(custom_level["members_count"], 1)
        self.assertIn("Pre-competitivo", {level["name"] for level in levels})

        self.member_service.update_member(
            member_id,
            {
                "name": "Aluno Nivel",
                "member_type": "aluno",
                "status": "active",
                "learning_level_id": "",
            },
        )
        self.assertIsNone(self.db.get_member(member_id)["learning_level_id"])

        with self.assertRaisesRegex(AppError, "Nivel de aprendizagem"):
            self.member_service.update_member(
                member_id,
                {
                    "name": "Aluno Nivel",
                    "member_type": "aluno",
                    "status": "active",
                    "learning_level_id": "99999",
                },
            )

    def test_database_schema_version_and_pairing_indexes_exist(self) -> None:
        with self.db.connect() as connection:
            user_version = connection.execute("PRAGMA user_version").fetchone()[0]
            tournament_columns = {
                row["name"]
                for row in connection.execute("PRAGMA table_info(tournaments)").fetchall()
            }
            settings_columns = {
                row["name"]
                for row in connection.execute("PRAGMA table_info(tournament_settings)").fetchall()
            }
            round_columns = {
                row["name"]
                for row in connection.execute("PRAGMA table_info(rounds)").fetchall()
            }
            member_columns = {
                row["name"]
                for row in connection.execute("PRAGMA table_info(members)").fetchall()
            }
            training_session_columns = {
                row["name"]
                for row in connection.execute("PRAGMA table_info(training_sessions)").fetchall()
            }
            pairing_indexes = {
                row["name"]
                for row in connection.execute("PRAGMA index_list(pairings)").fetchall()
            }
            learning_tables = {
                row["name"]
                for row in connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table' AND name = 'learning_levels'
                    """
                ).fetchall()
            }
            learning_indexes = {
                row["name"]
                for row in connection.execute("PRAGMA index_list(learning_levels)").fetchall()
            }
            team_tables = {
                row["name"]
                for row in connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table'
                        AND name IN (
                            'teams', 'team_players', 'team_matches', 'team_boards',
                            'team_lineups', 'team_lineup_boards', 'team_substitution_events'
                        )
                    """
                ).fetchall()
            }
            team_indexes = {
                row["name"]
                for table_name in (
                    "teams", "team_players", "team_matches", "team_boards",
                    "team_lineups", "team_lineup_boards", "team_substitution_events"
                )
                for row in connection.execute(f"PRAGMA index_list({table_name})").fetchall()
            }
            certificate_tables = {
                row["name"]
                for row in connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table'
                        AND name IN ('certificate_templates', 'certificate_issuances')
                    """
                ).fetchall()
            }
            certificate_indexes = {
                row["name"]
                for table_name in ("certificate_templates", "certificate_issuances")
                for row in connection.execute(f"PRAGMA index_list({table_name})").fetchall()
            }
            exercise_tables = {
                row["name"]
                for row in connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table'
                        AND name IN (
                            'exercise_library',
                            'training_lists',
                            'training_list_exercises',
                            'exercise_attempts'
                        )
                    """
                ).fetchall()
            }
            exercise_indexes = {
                row["name"]
                for table_name in (
                    "exercise_library",
                    "training_lists",
                    "training_list_exercises",
                    "exercise_attempts",
                    "training_sessions",
                )
                for row in connection.execute(f"PRAGMA index_list({table_name})").fetchall()
            }
            inventory_tables = {
                row["name"]
                for row in connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table'
                        AND name IN (
                            'inventory_items',
                            'inventory_loans',
                            'inventory_maintenance'
                        )
                    """
                ).fetchall()
            }
            inventory_indexes = {
                row["name"]
                for table_name in ("inventory_items", "inventory_loans", "inventory_maintenance")
                for row in connection.execute(f"PRAGMA index_list({table_name})").fetchall()
            }
            audit_table = connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'audit_log'
                """
            ).fetchone()
            phase0_tables = {
                row["name"]
                for row in connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table'
                        AND name IN (
                            'audit_events', 'pairing_snapshots', 'standings_snapshots',
                            'tiebreak_components', 'public_tokens', 'result_submissions'
                        )
                    """
                ).fetchall()
            }
            audit_indexes = {
                row["name"]
                for row in connection.execute("PRAGMA index_list(audit_log)").fetchall()
            }
            phase0_indexes = {
                row["name"]
                for table_name in (
                    "audit_events", "pairing_snapshots", "standings_snapshots",
                    "tiebreak_components", "public_tokens", "result_submissions"
                )
                for row in connection.execute(f"PRAGMA index_list({table_name})").fetchall()
            }
            integration_tables = {
                row["name"]
                for row in connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table'
                        AND name IN ('devices', 'sync_outbox', 'clock_events')
                    """
                ).fetchall()
            }
            integration_indexes = {
                row["name"]
                for table_name in ("devices", "sync_outbox", "clock_events")
                for row in connection.execute(f"PRAGMA index_list({table_name})").fetchall()
            }
            certificate_template_count = connection.execute(
                "SELECT COUNT(*) FROM certificate_templates"
            ).fetchone()[0]
            certificate_template_columns = {
                row["name"]
                for row in connection.execute("PRAGMA table_info(certificate_templates)").fetchall()
            }

        self.assertEqual(user_version, Database.SCHEMA_VERSION)
        self.assertIn("competition_type", tournament_columns)
        self.assertIn("tournament_profile", settings_columns)
        self.assertIn("pairing_system", settings_columns)
        self.assertIn("acceleration_method", settings_columns)
        self.assertIn("team_boards_count", settings_columns)
        self.assertIn("team_match_win_points", settings_columns)
        self.assertIn("team_standing_primary", settings_columns)
        self.assertIn("team_standing_secondary", settings_columns)
        self.assertIn("team_board_order_policy", settings_columns)
        self.assertIn("team_reserve_policy", settings_columns)
        self.assertIn("team_lineup_deadline", settings_columns)
        self.assertIn("team_max_substitutions", settings_columns)
        self.assertIn("rating_fee_fide", settings_columns)
        self.assertIn("rating_fee_cbx", settings_columns)
        self.assertIn("rating_fee_lbx", settings_columns)
        self.assertIn("learning_level_id", member_columns)
        self.assertIn("training_list_id", training_session_columns)
        self.assertIn("pairing_engine_version", round_columns)
        self.assertIn("ruleset_version", round_columns)
        self.assertIn("closed_at", round_columns)
        self.assertEqual({"learning_levels"}, learning_tables)
        self.assertIn("idx_learning_levels_order", learning_indexes)
        self.assertIn("idx_pairings_white_player", pairing_indexes)
        self.assertIn("idx_pairings_black_player", pairing_indexes)
        self.assertEqual(
            {
                "teams", "team_players", "team_matches", "team_boards",
                "team_lineups", "team_lineup_boards", "team_substitution_events",
            },
            team_tables,
        )
        self.assertIn("idx_teams_tournament", team_indexes)
        self.assertIn("idx_team_matches_round", team_indexes)
        self.assertIn("idx_team_lineups_round", team_indexes)
        self.assertIn("idx_team_substitutions_round", team_indexes)
        self.assertEqual({"certificate_templates", "certificate_issuances"}, certificate_tables)
        self.assertIn("idx_certificate_templates_type", certificate_indexes)
        self.assertIn("idx_certificate_issuances_context", certificate_indexes)
        self.assertEqual(
            {"exercise_library", "training_lists", "training_list_exercises", "exercise_attempts"},
            exercise_tables,
        )
        self.assertIn("idx_exercise_library_filters", exercise_indexes)
        self.assertIn("idx_training_lists_filters", exercise_indexes)
        self.assertIn("idx_training_sessions_training_list", exercise_indexes)
        self.assertIn("idx_exercise_attempts_member", exercise_indexes)
        self.assertEqual({"inventory_items", "inventory_loans", "inventory_maintenance"}, inventory_tables)
        self.assertIn("idx_inventory_items_filters", inventory_indexes)
        self.assertIn("idx_inventory_loans_item", inventory_indexes)
        self.assertIn("idx_inventory_maintenance_item", inventory_indexes)
        self.assertIsNotNone(audit_table)
        self.assertEqual(
            {
                "audit_events", "pairing_snapshots", "standings_snapshots",
                "tiebreak_components", "public_tokens", "result_submissions",
            },
            phase0_tables,
        )
        self.assertIn("idx_audit_log_created", audit_indexes)
        self.assertIn("idx_audit_log_entity", audit_indexes)
        self.assertIn("idx_audit_events_tournament", phase0_indexes)
        self.assertIn("idx_pairing_snapshots_round", phase0_indexes)
        self.assertIn("idx_standings_snapshots_round", phase0_indexes)
        self.assertIn("idx_tiebreak_components_player", phase0_indexes)
        self.assertIn("idx_public_tokens_hash", phase0_indexes)
        self.assertIn("idx_result_submissions_status", phase0_indexes)
        self.assertEqual({"devices", "sync_outbox", "clock_events"}, integration_tables)
        self.assertIn("idx_devices_status", integration_indexes)
        self.assertIn("idx_sync_outbox_status", integration_indexes)
        self.assertIn("idx_clock_events_tournament", integration_indexes)
        self.assertGreaterEqual(certificate_template_count, 7)
        self.assertIn("logo_path", certificate_template_columns)
        self.assertIn("background_image_path", certificate_template_columns)
        self.assertIn("background_opacity", certificate_template_columns)
        self.assertIn("secondary_logo_path", certificate_template_columns)
        self.assertIn("primary_color", certificate_template_columns)
        self.assertIn("accent_color", certificate_template_columns)
        self.assertIn("title_font_size", certificate_template_columns)

    def test_phase0_records_pairing_snapshots_standings_snapshot_and_audit_events(self) -> None:
        self._create_players(4)
        round_data = self.service.generate_next_round(self.tournament_id)
        round_id = int(round_data["id"])
        pairings = self.db.get_pairings_for_round(round_id)
        for pairing in pairings:
            self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")

        self.service.close_round(self.tournament_id, round_id)

        snapshots = self.db.list_pairing_snapshots(self.tournament_id, round_id=round_id)
        input_snapshots = self.db.list_pairing_snapshots(self.tournament_id, round_number=1)
        standings_snapshots = self.db.list_standings_snapshots(self.tournament_id)
        audit_events = self.db.list_audit_events(self.tournament_id, limit=20)
        current_round = self.db.get_round(round_id)

        self.assertTrue(any(item["stage"] == "input" for item in input_snapshots))
        self.assertTrue(any(item["stage"] == "output" for item in snapshots))
        self.assertEqual(1, len(standings_snapshots))
        self.assertTrue(standings_snapshots[0]["snapshot_hash"])
        self.assertEqual("albericus-swiss-1", current_round["pairing_engine_version"])
        self.assertEqual("albericus-2026-phase0", current_round["ruleset_version"])
        self.assertIn("round_generated", {item["action"] for item in audit_events})
        self.assertIn("round_closed", {item["action"] for item in audit_events})
        self.assertIn("backup_created", {item["action"] for item in audit_events})

    def test_phase1_preview_next_round_does_not_persist(self) -> None:
        self._create_players(4)

        preview = self.service.preview_next_round(self.tournament_id)

        self.assertEqual("individual", preview["competition_type"])
        self.assertEqual(1, preview["round_number"])
        self.assertEqual(2, len(preview["pairings"]))
        self.assertEqual([], self.db.list_rounds(self.tournament_id))
        self.assertEqual([], self.db.list_pairing_snapshots(self.tournament_id))
        self.assertEqual([], self.db.list_audit_events(self.tournament_id))
        self.assertEqual([], list(self.backup_dir.glob("*.db")))

    def test_phase1_preview_flags_bye_alert(self) -> None:
        self._create_players(3)

        preview = self.service.preview_next_round(self.tournament_id)

        self.assertEqual(2, len(preview["pairings"]))
        self.assertGreaterEqual(preview["alerts_count"], 1)
        self.assertTrue(any("Bye" in pairing["alerts"] for pairing in preview["pairings"]))

    def test_phase2_arbitration_dashboard_reports_pending_ready_and_corrections(self) -> None:
        self._create_players(2)
        empty_dashboard = self.service.arbitration_dashboard(self.tournament_id)
        self.assertEqual(0, empty_dashboard["metrics"]["generated_rounds"])
        self.assertIn("Nenhuma rodada gerada", empty_dashboard["alerts"][0])

        round_data = self.service.generate_next_round(self.tournament_id)
        pending_dashboard = self.service.arbitration_dashboard(self.tournament_id)
        self.assertEqual(1, pending_dashboard["metrics"]["pending_results"])
        self.assertFalse(pending_dashboard["metrics"]["ready_to_close"])
        self.assertEqual(1, len(pending_dashboard["pending_items"]))
        self.assertEqual(1, pending_dashboard["pending_items"][0]["board"])
        self.assertTrue(pending_dashboard["pending_items"][0]["pairing_id"])

        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        ready_dashboard = self.service.arbitration_dashboard(self.tournament_id)
        self.assertEqual(0, ready_dashboard["metrics"]["pending_results"])
        self.assertTrue(ready_dashboard["metrics"]["ready_to_close"])
        self.assertEqual([], ready_dashboard["pending_items"])

        self.service.close_round(self.tournament_id, int(round_data["id"]))
        settings = self.db.get_tournament_settings(self.tournament_id) or {}
        settings["allow_dangerous_changes"] = 1
        self.db.save_tournament_settings(self.tournament_id, settings)
        self.service.update_result(self.tournament_id, int(pairing["id"]), "0-1")

        corrected_dashboard = self.service.arbitration_dashboard(self.tournament_id)
        self.assertEqual(1, corrected_dashboard["metrics"]["corrections"])
        self.assertTrue(corrected_dashboard["metrics"]["can_preview_next_round"])

    def test_phase2_arbitration_dashboard_respects_pending_items_limit(self) -> None:
        self._create_players(24)
        self.service.generate_next_round(self.tournament_id)

        dashboard = self.service.arbitration_dashboard(self.tournament_id, pending_limit=10)

        self.assertEqual(12, dashboard["metrics"]["pending_results"])
        self.assertEqual(10, len(dashboard["pending_items"]))
        self.assertEqual(list(range(1, 11)), [item["board"] for item in dashboard["pending_items"]])

    def test_phase2_arbitration_dashboard_searches_exact_pending_board_before_limit(self) -> None:
        self._create_players(120)
        self.service.generate_next_round(self.tournament_id)

        dashboard = self.service.arbitration_dashboard(
            self.tournament_id,
            pending_limit=50,
            pending_query="60",
        )

        self.assertEqual(60, dashboard["metrics"]["pending_results"])
        self.assertEqual([60], [item["board"] for item in dashboard["pending_items"]])

    def test_phase2_arbitration_dashboard_round_clock_freezes_when_round_closes(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        with self.db.connect() as connection:
            connection.execute(
                "UPDATE rounds SET created_at = '2026-05-31 10:00:00' WHERE id = ?",
                (int(round_data["id"]),),
            )

        with mock.patch.object(self.db, "now", return_value="2026-05-31 11:30:00"):
            open_dashboard = self.service.arbitration_dashboard(self.tournament_id)
        self.assertEqual("em_andamento", open_dashboard["metrics"]["round_clock_status"])
        self.assertEqual("31/05 10:00", open_dashboard["metrics"]["round_started_label"])
        self.assertEqual("01:30:00", open_dashboard["metrics"]["round_duration_label"])

        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        with mock.patch.object(self.db, "now", return_value="2026-05-31 12:00:00"):
            self.service.close_round(self.tournament_id, int(round_data["id"]))
        with mock.patch.object(self.db, "now", return_value="2026-05-31 14:00:00"):
            closed_dashboard = self.service.arbitration_dashboard(self.tournament_id)

        self.assertEqual("2026-05-31 12:00:00", self.db.get_round(int(round_data["id"]))["closed_at"])
        self.assertEqual("fechada", closed_dashboard["metrics"]["round_clock_status"])
        self.assertEqual("02:00:00", closed_dashboard["metrics"]["round_duration_label"])

    def test_phase2_round_close_blocks_arbitration_decision_issues(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        token_payload = self.qr_result_service.result_url_for_pairing(self.tournament_id, int(pairing["id"]))
        submission = self.qr_result_service.submit_result(token_payload["token"], "1-0", submitter="Mesa 1")
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")

        dashboard = self.service.arbitration_dashboard(self.tournament_id)
        self.assertFalse(dashboard["metrics"]["ready_to_close"])
        self.assertEqual(1, dashboard["metrics"]["blocking_issues"])
        with self.assertRaisesRegex(AppError, "pendencias de arbitragem bloqueantes"):
            self.service.close_round(self.tournament_id, int(round_data["id"]))

        self.qr_result_service.reject_submission(int(submission["id"]), reviewer="Arbitro", reason="Resultado lancado no desktop.")
        self.service.close_round(self.tournament_id, int(round_data["id"]))
        self.assertEqual("closed", self.db.get_round(int(round_data["id"]))["status"])

    def test_phase2_round_close_allows_attention_only_clock_issue(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.clock_integration_service.record_manual_event(
            self.tournament_id,
            "flag_fall",
            pairing_id=int(pairing["id"]),
            side="white",
            seconds_remaining=0,
            note="Apenas alerta conferido em mesa.",
        )

        dashboard = self.service.arbitration_dashboard(self.tournament_id)
        self.assertTrue(dashboard["metrics"]["ready_to_close"])
        self.assertEqual(0, dashboard["metrics"]["blocking_issues"])

        self.service.close_round(self.tournament_id, int(round_data["id"]))
        self.assertEqual("closed", self.db.get_round(int(round_data["id"]))["status"])

    def test_phase2_arbitration_issues_aggregate_qr_sync_and_clock_alerts(self) -> None:
        self._create_players(2)
        remote_device = self.sync_service.register_device("Mesa 1", role="assistant", device_id="device-board-1")
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        token_payload = self.qr_result_service.result_url_for_pairing(self.tournament_id, int(pairing["id"]))
        self.qr_result_service.submit_result(token_payload["token"], "1-0", submitter="Mesa 1")
        self.sync_service.apply_remote_event(
            {
                "event_id": "remote-invalid-result",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id,
                "round_id": int(round_data["id"]),
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": {"result": "BYE"},
                "payload_hash": self.db._hash_payload({"result": "BYE"}),
            }
        )
        self.clock_integration_service.record_manual_event(
            self.tournament_id,
            "flag_fall",
            pairing_id=int(pairing["id"]),
            side="white",
            seconds_remaining=0,
            note="Conferir mesa.",
        )

        issues = self.service.arbitration_issues(self.tournament_id)
        sources = {item["source"] for item in issues["issues"]}

        self.assertEqual(3, issues["metrics"]["total"])
        self.assertEqual(1, issues["metrics"]["qr_pending"])
        self.assertEqual(1, issues["metrics"]["sync_conflicts"])
        self.assertEqual(1, issues["metrics"]["clock_alerts"])
        self.assertEqual({"qr", "sync", "clock"}, sources)
        self.assertTrue(all(item.get("payload") for item in issues["issues"]))
        self.assertTrue(all(item.get("issue_key") for item in issues["issues"]))

    def test_phase2_arbitration_issue_filters_are_pure_and_search_payload_board(self) -> None:
        issues = [
            {
                "severity": "decision",
                "source": "qr",
                "kind": "result_submission",
                "title": "Resultado QR pendente - mesa 17",
                "detail": "Resultado enviado: 1-0",
                "payload": {"board_number": 17},
            },
            {
                "severity": "decision",
                "source": "sync",
                "kind": "remote_rejected",
                "title": "Evento remoto rejeitado",
                "detail": "Conferir conflito",
                "payload": {"board_number": 42},
            },
            {
                "severity": "attention",
                "source": "clock",
                "kind": "flag_fall",
                "title": "Queda de seta registrada",
                "detail": "Mesa decisiva",
                "payload": {"board_number": 3},
            },
        ]

        self.assertEqual(issues, self.service.filter_arbitration_issues(issues))
        self.assertEqual([issues[0], issues[1]], self.service.filter_arbitration_issues(issues, "decision"))
        self.assertEqual([issues[0]], self.service.filter_arbitration_issues(issues, "qr"))
        self.assertEqual([issues[1]], self.service.filter_arbitration_issues(issues, query="42"))
        self.assertEqual([issues[2]], self.service.filter_arbitration_issues(issues, query="mesa decisiva"))
        self.assertEqual(3, len(issues))

    def test_phase2_arbitration_issue_acknowledgement_hides_non_qr_issue(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        self.clock_integration_service.record_manual_event(
            self.tournament_id,
            "flag_fall",
            pairing_id=int(pairing["id"]),
            side="black",
            seconds_remaining=0,
            note="Mesa conferida.",
        )

        issues = self.service.arbitration_issues(self.tournament_id)
        clock_issue = next(item for item in issues["issues"] if item["source"] == "clock")
        self.service.acknowledge_arbitration_issue(self.tournament_id, str(clock_issue["issue_key"]))

        refreshed = self.service.arbitration_issues(self.tournament_id)
        audit_events = self.db.list_audit_events(
            self.tournament_id,
            action="arbitration_issue_acknowledged",
            entity_type="arbitration_issue",
        )

        self.assertEqual(0, refreshed["metrics"]["total"])
        self.assertEqual(1, len(audit_events))
        self.assertIn(str(clock_issue["issue_key"]), audit_events[0]["after_json"])

    def test_phase2_arbitration_issue_acknowledgement_rejects_qr_issue(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        token_payload = self.qr_result_service.result_url_for_pairing(self.tournament_id, int(pairing["id"]))
        self.qr_result_service.submit_result(token_payload["token"], "1-0", submitter="Mesa 1")

        issues = self.service.arbitration_issues(self.tournament_id)
        qr_issue = next(item for item in issues["issues"] if item["source"] == "qr")

        with self.assertRaisesRegex(AppError, "Use Aprovar QR ou Rejeitar QR"):
            self.service.acknowledge_arbitration_issue(self.tournament_id, str(qr_issue["issue_key"]))

    def test_phase0_audits_dangerous_result_correction(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, int(round_data["id"]))
        settings = self.db.get_tournament_settings(self.tournament_id) or {}
        settings["allow_dangerous_changes"] = 1
        self.db.save_tournament_settings(self.tournament_id, settings)

        self.service.update_result(self.tournament_id, int(pairing["id"]), "0-1")

        corrections = self.db.list_audit_events(
            self.tournament_id,
            action="result_corrected",
            entity_type="pairing",
        )
        self.assertEqual(1, len(corrections))
        self.assertIn("1-0", corrections[0]["before_json"])
        self.assertIn("0-1", corrections[0]["after_json"])
        self.assertTrue(corrections[0]["before_hash"])
        self.assertTrue(corrections[0]["after_hash"])

    def test_phase8_audit_events_enqueue_sync_outbox_and_network_failure_keeps_pending(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        outbox = self.db.list_sync_outbox(status="pending", limit=20)
        round_generated = next(item for item in outbox if item["action"] == "round_generated")

        with mock.patch("src.services.sync_service.urlopen", side_effect=OSError("offline")):
            summary = self.sync_service.sync_pending("https://sync.example.test", limit=1)

        pending = self.db.list_sync_outbox(status="pending", limit=20)
        failed_event = next(item for item in pending if int(item["attempts"]) == 1)
        self.assertEqual(int(round_data["id"]), int(round_generated["round_id"]))
        self.assertEqual({"sent": 1, "synced": 0, "rejected": 0, "failed": 1}, summary)
        self.assertEqual("pending", failed_event["status"])
        self.assertEqual(1, failed_event["attempts"])
        self.assertIn("offline", failed_event["last_error"])

    def test_phase8_device_registration_validates_role_name_and_revoke_target(self) -> None:
        with self.assertRaisesRegex(AppError, "nome do dispositivo"):
            self.sync_service.register_device("   ", role="assistant")
        with self.assertRaisesRegex(AppError, "Perfil"):
            self.sync_service.register_device("Mesa 1", role="admin")

        device = self.sync_service.register_device(" Mesa 1 ", role="assistant", device_id=" device-board-1 ")
        self.sync_service.revoke_device("device-board-1")

        revoked = next(item for item in self.sync_service.list_devices() if item["device_id"] == "device-board-1")
        self.assertEqual("Mesa 1", device["name"])
        self.assertEqual("revoked", revoked["status"])
        with self.assertRaisesRegex(AppError, "nao encontrado"):
            self.sync_service.revoke_device("missing-device")

    def test_phase8_sync_pending_validates_target_url_limit_and_timeout(self) -> None:
        with self.assertRaisesRegex(AppError, "URL"):
            self.sync_service.sync_pending("sync.example.test")
        with self.assertRaisesRegex(AppError, "Limite"):
            self.sync_service.sync_pending("https://sync.example.test", limit=0)
        with self.assertRaisesRegex(AppError, "Timeout"):
            self.sync_service.sync_pending("https://sync.example.test", timeout_seconds=0)

        self.assertEqual([], self.db.list_sync_outbox(status="rejected", limit=20))

    def test_phase8_remote_rejection_ignores_invalid_audit_scope(self) -> None:
        response = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-invalid-scope",
                "device_id": "unknown-device",
                "tournament_id": 999999,
                "round_id": 999999,
                "entity_type": "pairing",
                "entity_id": 999999,
                "action": "result_updated",
                "payload": {"result": "1-0"},
            }
        )

        rejected = self.db.list_audit_events(action="sync_remote_rejected")
        self.assertEqual("rejected", response["status"])
        self.assertTrue(rejected)
        self.assertIsNone(rejected[0]["tournament_id"])
        self.assertIsNone(rejected[0]["round_id"])

    def test_phase8_remote_result_rejects_malformed_ids_without_exception(self) -> None:
        self._create_players(2)
        remote_device = self.sync_service.register_device("Mesa 1", role="assistant", device_id="device-board-1")
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        bad_pairing = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-bad-pairing",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id,
                "round_id": int(round_data["id"]),
                "entity_type": "pairing",
                "entity_id": "abc",
                "action": "result_updated",
                "payload": {"result": "1-0"},
            }
        )
        bad_tournament = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-bad-tournament",
                "device_id": remote_device["device_id"],
                "tournament_id": "abc",
                "round_id": int(round_data["id"]),
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": {"result": "1-0"},
            }
        )
        bad_round = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-bad-round",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id,
                "round_id": "abc",
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": {"result": "1-0"},
            }
        )

        unchanged = self.db.get_pairing(int(pairing["id"]))
        self.assertEqual("rejected", bad_pairing["status"])
        self.assertEqual("rejected", bad_tournament["status"])
        self.assertEqual("rejected", bad_round["status"])
        self.assertEqual("", unchanged["result"])

    def test_phase8_remote_result_rejects_malformed_payload_without_exception(self) -> None:
        self._create_players(2)
        remote_device = self.sync_service.register_device("Mesa 1", role="assistant", device_id="device-board-1")
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        response = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-bad-payload",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id,
                "round_id": int(round_data["id"]),
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": ["1-0"],
            }
        )

        unchanged = self.db.get_pairing(int(pairing["id"]))
        rejected = self.db.list_audit_events(self.tournament_id, action="sync_remote_rejected")
        self.assertEqual("rejected", response["status"])
        self.assertEqual("", unchanged["result"])
        self.assertTrue(any("Payload remoto invalido" in item["reason"] for item in rejected))

    def test_phase8_remote_result_conflict_does_not_change_closed_round(self) -> None:
        self._create_players(2)
        remote_device = self.sync_service.register_device("Mesa 1", role="assistant", device_id="device-board-1")
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, int(round_data["id"]))

        response = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-event-1",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id,
                "round_id": int(round_data["id"]),
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": {"result": "0-1"},
            }
        )

        unchanged = self.db.get_pairing(int(pairing["id"]))
        rejected = self.db.list_audit_events(self.tournament_id, action="sync_remote_rejected")
        self.assertEqual("conflict", response["status"])
        self.assertEqual("1-0", unchanged["result"])
        self.assertTrue(rejected)
        self.assertIn("Rodada fechada", rejected[0]["reason"])

    def test_phase8_remote_result_rejects_payload_hash_mismatch(self) -> None:
        self._create_players(2)
        remote_device = self.sync_service.register_device("Mesa 1", role="assistant", device_id="device-board-1")
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        response = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-event-tampered",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id,
                "round_id": int(round_data["id"]),
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": {"result": "0-1"},
                "payload_hash": self.db._hash_payload({"result": "1-0"}),
            }
        )

        unchanged = self.db.get_pairing(int(pairing["id"]))
        rejected = self.db.list_audit_events(self.tournament_id, action="sync_remote_rejected")
        self.assertEqual("rejected", response["status"])
        self.assertEqual("", unchanged["result"])
        self.assertTrue(rejected)
        self.assertIn("Hash do payload", rejected[0]["reason"])

    def test_phase8_remote_result_cannot_clear_pairing_result(self) -> None:
        self._create_players(2)
        remote_device = self.sync_service.register_device("Mesa 1", role="assistant", device_id="device-board-1")
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")

        response = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-event-clear",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id,
                "round_id": int(round_data["id"]),
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": {"result": ""},
                "payload_hash": self.db._hash_payload({"result": ""}),
            }
        )

        unchanged = self.db.get_pairing(int(pairing["id"]))
        rejected = self.db.list_audit_events(self.tournament_id, action="sync_remote_rejected")
        self.assertEqual("rejected", response["status"])
        self.assertEqual("1-0", unchanged["result"])
        self.assertTrue(rejected)
        self.assertIn("Resultado remoto invalido", rejected[0]["reason"])

    def test_phase8_remote_result_conflict_does_not_override_open_desktop_result(self) -> None:
        self._create_players(2)
        remote_device = self.sync_service.register_device("Mesa 1", role="assistant", device_id="device-board-1")
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")

        response = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-open-conflict",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id,
                "round_id": int(round_data["id"]),
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": {"result": "0-1"},
                "payload_hash": self.db._hash_payload({"result": "0-1"}),
            }
        )

        unchanged = self.db.get_pairing(int(pairing["id"]))
        rejected = self.db.list_audit_events(self.tournament_id, action="sync_remote_rejected")
        self.assertEqual("conflict", response["status"])
        self.assertEqual("1-0", unchanged["result"])
        self.assertTrue(any("outro resultado" in item["reason"] for item in rejected))

    def test_phase8_remote_result_rejects_tournament_or_round_mismatch(self) -> None:
        self._create_players(2)
        remote_device = self.sync_service.register_device("Mesa 1", role="assistant", device_id="device-board-1")
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        wrong_tournament = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-event-wrong-tournament",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id + 999,
                "round_id": int(round_data["id"]),
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": {"result": "1-0"},
                "payload_hash": self.db._hash_payload({"result": "1-0"}),
            }
        )
        wrong_round = self.sync_service.apply_remote_event(
            {
                "event_id": "remote-event-wrong-round",
                "device_id": remote_device["device_id"],
                "tournament_id": self.tournament_id,
                "round_id": int(round_data["id"]) + 999,
                "entity_type": "pairing",
                "entity_id": int(pairing["id"]),
                "action": "result_updated",
                "payload": {"result": "0-1"},
                "payload_hash": self.db._hash_payload({"result": "0-1"}),
            }
        )

        unchanged = self.db.get_pairing(int(pairing["id"]))
        rejected = self.db.list_audit_events(self.tournament_id, action="sync_remote_rejected")
        self.assertEqual("rejected", wrong_tournament["status"])
        self.assertEqual("rejected", wrong_round["status"])
        self.assertEqual("", unchanged["result"])
        self.assertTrue(any("outra rodada" in item["reason"] for item in rejected))

    def test_phase8_pending_event_syncs_when_server_returns(self) -> None:
        event_id = self.db.create_sync_outbox_event(
            action="heartbeat",
            tournament_id=self.tournament_id,
            entity_type="tournament",
            entity_id=self.tournament_id,
            payload={"status": "local_ok"},
        )
        event = self.db.list_sync_outbox(status="pending", limit=1)[0]

        class FakeResponse:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, traceback):
                return False

            @staticmethod
            def read() -> bytes:
                return b'{"status": "accepted"}'

        with mock.patch("src.services.sync_service.urlopen", return_value=FakeResponse()):
            summary = self.sync_service.sync_pending("https://sync.example.test", limit=1)

        synced = self.db.list_sync_outbox(status="synced", limit=20)
        self.assertEqual({"sent": 1, "synced": 1, "rejected": 0, "failed": 0}, summary)
        self.assertEqual(event["event_id"], next(item["event_id"] for item in synced if item["id"] == event_id))

    def test_phase9_manual_clock_event_is_audited_without_changing_result(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        event = self.clock_integration_service.record_manual_event(
            self.tournament_id,
            "flag_fall",
            pairing_id=int(pairing["id"]),
            side="white",
            seconds_remaining=0,
            note="Seta caiu, aguardando decisao do arbitro.",
        )

        unchanged = self.db.get_pairing(int(pairing["id"]))
        audit_events = self.db.list_audit_events(self.tournament_id, action="clock_event_logged")
        alerts = self.clock_integration_service.anomaly_alerts(self.tournament_id)
        self.assertEqual("flag_fall", event["event_type"])
        self.assertEqual("", unchanged["result"])
        self.assertTrue(audit_events)
        self.assertTrue(any("Apenas alerta" in alert["recommendation"] for alert in alerts))

    def test_phase9_clock_event_rejects_invalid_side_and_negative_time(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        with self.assertRaisesRegex(AppError, "Lado"):
            self.clock_integration_service.record_manual_event(
                self.tournament_id,
                "time_warning",
                pairing_id=int(pairing["id"]),
                side="red",
                seconds_remaining=30,
            )
        with self.assertRaisesRegex(AppError, "Segundos"):
            self.clock_integration_service.record_manual_event(
                self.tournament_id,
                "time_warning",
                pairing_id=int(pairing["id"]),
                side="white",
                seconds_remaining=-1,
            )

        self.assertEqual([], self.clock_integration_service.list_clock_events(self.tournament_id))

    def test_phase9_clock_event_validates_player_tournament_and_pairing(self) -> None:
        player_ids = self._create_players(2)
        outside_tournament_id = self.db.create_tournament("Outro torneio", rounds_count=1)
        outside_player_id = self.db.create_player(outside_tournament_id, name="Jogador externo", rating=1500)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        with self.assertRaisesRegex(AppError, "torneio selecionado"):
            self.clock_integration_service.record_manual_event(
                self.tournament_id,
                "manual_note",
                player_id=outside_player_id,
            )
        with self.assertRaisesRegex(AppError, "mesa informada"):
            self.clock_integration_service.record_manual_event(
                self.tournament_id,
                "manual_note",
                pairing_id=int(pairing["id"]),
                player_id=outside_player_id,
            )

        event = self.clock_integration_service.record_manual_event(
            self.tournament_id,
            "manual_note",
            pairing_id=int(pairing["id"]),
            player_id=player_ids[0],
        )

        self.assertEqual(player_ids[0], event["player_id"])
        self.assertEqual(1, len(self.clock_integration_service.list_clock_events(self.tournament_id)))

    def test_phase9_plugin_events_are_optional_and_never_apply_result(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        class FakeClockPlugin:
            plugin_id = "fake_clock"
            label = "Relogio falso"

            def normalize_event(self, raw_event: dict[str, object]) -> dict[str, object]:
                return {**raw_event, "event_type": "absence", "note": "Ausencia detectada pelo dispositivo."}

        self.clock_integration_service.register_plugin(FakeClockPlugin())
        disabled = self.clock_integration_service.record_plugin_event(
            "fake_clock",
            {
                "tournament_id": self.tournament_id,
                "pairing_id": int(pairing["id"]),
                "device_id": "clock-1",
                "result": "0-1",
            },
        )
        self.db.save_app_settings({"device_integrations_enabled": "1"})
        accepted = self.clock_integration_service.record_plugin_event(
            "fake_clock",
            {
                "tournament_id": self.tournament_id,
                "pairing_id": int(pairing["id"]),
                "device_id": "clock-1",
                "result": "0-1",
            },
        )

        unchanged = self.db.get_pairing(int(pairing["id"]))
        self.assertEqual("ignored", disabled["status"])
        self.assertEqual("absence", accepted["event_type"])
        self.assertEqual("", unchanged["result"])

    def test_phase9_notifications_are_optional_audited_queue_only(self) -> None:
        skipped = self.clock_integration_service.queue_notification(
            "EMAIL",
            " arbitro@example.com ",
            " Rodada publicada. ",
            tournament_id=self.tournament_id,
        )
        self.db.save_app_settings({"notifications_enabled": "1"})
        queued = self.clock_integration_service.queue_notification(
            "sms",
            "+5500000000000",
            "Mesa 1 requer atencao.",
            tournament_id=self.tournament_id,
        )

        skipped_events = self.db.list_audit_events(self.tournament_id, action="notification_skipped")
        queued_events = self.db.list_audit_events(self.tournament_id, action="notification_queued")
        self.assertEqual("skipped", skipped["status"])
        self.assertEqual("queued", queued["status"])
        self.assertTrue(skipped_events)
        self.assertTrue(queued_events)

    def test_phase9_notification_queue_validates_channel_recipient_and_message(self) -> None:
        self.db.save_app_settings({"notifications_enabled": "1"})

        with self.assertRaisesRegex(AppError, "Canal"):
            self.clock_integration_service.queue_notification("telegram", "arbitro@example.com", "Rodada publicada")
        with self.assertRaisesRegex(AppError, "destinatario"):
            self.clock_integration_service.queue_notification("email", "", "Rodada publicada")
        with self.assertRaisesRegex(AppError, "mensagem"):
            self.clock_integration_service.queue_notification("email", "arbitro@example.com", " ")

        self.assertEqual([], self.db.list_audit_events(self.tournament_id, action="notification_queued"))

    def test_phase0_exports_tournament_audit_report(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, int(round_data["id"]))

        output_path = Path(self.temp_dir.name) / "auditoria.csv"
        self.export_service.export_tournament_audit(self.tournament_id, output_path)
        content = output_path.read_text(encoding="utf-8")

        self.assertIn("Data/hora;Acao;Entidade", content)
        self.assertIn("round_generated", content)
        self.assertIn("round_closed", content)

    def test_phase3_tiebreak_components_are_persisted_and_exported(self) -> None:
        self._create_players(4)
        first_round = self.service.generate_next_round(self.tournament_id)
        self._fill_decisive_results(first_round["id"])
        self.service.close_round(self.tournament_id, first_round["id"])

        second_round = self.service.generate_next_round(self.tournament_id)
        self._fill_decisive_results(second_round["id"])
        self.service.close_round(self.tournament_id, second_round["id"])

        standings = self.service.tiebreak_report(self.tournament_id)
        first = standings[0]
        components = first["tiebreak_components"]
        self.assertIn("buchholz", components)
        self.assertIn("buchholz_median", components)
        self.assertIn("sonneborn_berger", components)
        self.assertIn("direct_encounter", components)
        self.assertIn("performance", components)
        self.assertEqual(first["buchholz"], components["buchholz"]["total"])
        self.assertEqual(first["sonneborn_berger"], components["sonneborn_berger"]["total"])
        self.assertTrue(components["buchholz"]["opponents"])

        persisted = self.db.list_tiebreak_components(
            self.tournament_id,
            player_id=int(first["player_id"]),
            round_id=int(second_round["id"]),
        )
        self.assertGreaterEqual(len(persisted), 6)
        self.assertIn("buchholz", {item["criterion"] for item in persisted})
        self.assertIn("opponents", persisted[0]["components_json"])

        report_path = Path(self.temp_dir.name) / "desempates.csv"
        self.export_service.export_tiebreak_report(self.tournament_id, report_path)
        report_content = report_path.read_text(encoding="utf-8-sig")
        self.assertIn("Buchholz", report_content)
        self.assertIn("Sonneborn-Berger", report_content)

        site_path = self.export_service.export_site(self.tournament_id, Path(self.temp_dir.name) / "site")
        html = site_path.read_text(encoding="utf-8")
        self.assertIn("Componentes de desempate", html)
        self.assertIn("Buchholz", html)

    def test_phase4_qr_result_submission_requires_approval_and_audits_review(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        token_payload = self.qr_result_service.result_url_for_pairing(
            self.tournament_id,
            int(pairing["id"]),
            base_url="http://localhost:8765",
        )
        submission = self.qr_result_service.submit_result(token_payload["token"], "1-0", submitter="Mesa 1")

        self.assertEqual("submitted", submission["status"])
        self.assertEqual("", self.db.get_pairing(int(pairing["id"]))["result"])
        self.assertEqual(1, len(self.qr_result_service.pending_submissions(self.tournament_id)))

        self.qr_result_service.approve_submission(int(submission["id"]), reviewer="Arbitro")

        self.assertEqual("1-0", self.db.get_pairing(int(pairing["id"]))["result"])
        self.assertEqual([], self.qr_result_service.pending_submissions(self.tournament_id))
        actions = {event["action"] for event in self.db.list_audit_events(self.tournament_id, limit=20)}
        self.assertIn("result_submitted", actions)
        self.assertIn("result_submission_approved", actions)

    def test_phase4_qr_result_rejects_duplicate_pending_submission(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        token_payload = self.qr_result_service.result_url_for_pairing(self.tournament_id, int(pairing["id"]))

        first = self.qr_result_service.submit_result(token_payload["token"], "1-0", submitter="Mesa 1")

        with self.assertRaisesRegex(AppError, "envio pendente"):
            self.qr_result_service.submit_result(token_payload["token"], "0-1", submitter="Mesa 1")
        self.assertEqual(1, len(self.qr_result_service.pending_submissions(self.tournament_id)))

        self.qr_result_service.reject_submission(int(first["id"]), reviewer="Arbitro", reason="Conferir novamente")
        second = self.qr_result_service.submit_result(token_payload["token"], "0-1", submitter="Mesa 1")

        self.assertEqual("submitted", second["status"])
        self.assertEqual(1, len(self.qr_result_service.pending_submissions(self.tournament_id)))

    def test_phase4_qr_new_link_revokes_previous_active_token_for_pairing(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]

        first = self.qr_result_service.result_url_for_pairing(self.tournament_id, int(pairing["id"]))
        second = self.qr_result_service.result_url_for_pairing(self.tournament_id, int(pairing["id"]))

        first_row = self.db.get_public_token_by_hash(self.qr_result_service._token_hash(first["token"]))
        second_row = self.db.get_public_token_by_hash(self.qr_result_service._token_hash(second["token"]))
        with self.assertRaisesRegex(AppError, "utilizado ou cancelado"):
            self.qr_result_service.submit_result(first["token"], "1-0")
        submission = self.qr_result_service.submit_result(second["token"], "1-0")

        self.assertEqual("revoked", first_row["status"])
        self.assertEqual("active", second_row["status"])
        self.assertEqual("submitted", submission["status"])

    def test_phase4_qr_batch_links_revoke_previous_token_and_audit_round(self) -> None:
        self._create_players(4)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairings = self.db.get_pairings_for_round(int(round_data["id"]))
        first_pairing = pairings[0]
        previous = self.qr_result_service.result_url_for_pairing(
            self.tournament_id,
            int(first_pairing["id"]),
        )

        urls = self.qr_result_service.result_urls_for_pairings(
            self.tournament_id,
            [{**pairing, "tournament_id": self.tournament_id} for pairing in pairings],
        )

        previous_row = self.db.get_public_token_by_hash(self.qr_result_service._token_hash(previous["token"]))
        active_tokens = []
        for pairing_id, url in urls.items():
            token = parse_qs(urlparse(url).query)["token"][0]
            active_tokens.append(self.qr_result_service._validate_token(token))
            self.assertEqual(pairing_id, int(active_tokens[-1]["pairing_id"]))
        audit_events = self.db.list_audit_events(self.tournament_id, limit=20)
        batch_event = next(event for event in audit_events if event["action"] == "public_result_tokens_created")

        self.assertEqual("revoked", previous_row["status"])
        self.assertEqual(2, len(active_tokens))
        self.assertIn('"tokens_count":2', batch_event["after_json"])

    def test_phase4_qr_does_not_override_desktop_result(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        token_payload = self.qr_result_service.result_url_for_pairing(self.tournament_id, int(pairing["id"]))

        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")

        with self.assertRaisesRegex(AppError, "resultado registrado no desktop"):
            self.qr_result_service.submit_result(token_payload["token"], "0-1")

        self.service.update_result(self.tournament_id, int(pairing["id"]), "")
        submission = self.qr_result_service.submit_result(token_payload["token"], "0-1")
        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")

        with self.assertRaisesRegex(AppError, "outro resultado registrado"):
            self.qr_result_service.approve_submission(int(submission["id"]), reviewer="Arbitro")
        self.assertEqual("1-0", self.db.get_pairing(int(pairing["id"]))["result"])

    def test_phase4_qr_token_rejects_tampering_expiration_closed_round_and_rejection(self) -> None:
        self._create_players(2)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        token_payload = self.qr_result_service.result_url_for_pairing(self.tournament_id, int(pairing["id"]))

        with self.assertRaisesRegex(AppError, "Assinatura"):
            self.qr_result_service.submit_result(token_payload["token"][:-1] + "x", "1-0")

        submission = self.qr_result_service.submit_result(token_payload["token"], "0-1")
        self.qr_result_service.reject_submission(int(submission["id"]), reviewer="Arbitro", reason="Conferido na mesa")
        self.assertEqual("", self.db.get_pairing(int(pairing["id"]))["result"])
        rejected = self.db.get_result_submission(int(submission["id"]))
        self.assertEqual("rejected", rejected["status"])

        expired = self.qr_result_service.result_url_for_pairing(self.tournament_id, int(pairing["id"]), expires_minutes=1)
        token_row = self.db.get_public_token_by_hash(self.qr_result_service._token_hash(expired["token"]))
        with self.db.connect() as connection:
            connection.execute(
                "UPDATE public_tokens SET expires_at = '2000-01-01 00:00:00' WHERE id = ?",
                (int(token_row["id"]),),
            )
        with self.assertRaisesRegex(AppError, "expirado"):
            self.qr_result_service.submit_result(expired["token"], "1-0")

        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, int(round_data["id"]))
        closed_token = self.qr_result_service.result_url_for_pairing
        with self.assertRaisesRegex(AppError, "rodada fechada"):
            closed_token(self.tournament_id, int(pairing["id"]))

    def test_phase4_local_result_server_escapes_mobile_html(self) -> None:
        form_html = LocalResultServer._result_form_html("abc'><script>alert(1)</script>")
        response_html = LocalResultServer._result_response_html("Erro <script>alert(1)</script>")

        self.assertNotIn("<script>", form_html)
        self.assertIn("abc&#x27;&gt;&lt;script&gt;alert(1)&lt;/script&gt;", form_html)
        self.assertNotIn("<script>", response_html)
        self.assertIn("Erro &lt;script&gt;alert(1)&lt;/script&gt;", response_html)

    def test_team_tournament_storage_foundation(self) -> None:
        tournament_id = self.tournament_service.create_tournament(
            {
                "name": "Interclubes",
                "competition_type": "team",
                "rounds_count": "3",
                "bye_points": "1",
            }
        )
        tournament = self.db.get_tournament(tournament_id)

        self.assertIsNotNone(tournament)
        self.assertEqual(tournament["competition_type"], "team")

        players = [
            self.db.create_player(tournament_id, name="Equipe A 1", rating=2100, club="A"),
            self.db.create_player(tournament_id, name="Equipe A 2", rating=2000, club="A"),
            self.db.create_player(tournament_id, name="Equipe B 1", rating=2050, club="B"),
            self.db.create_player(tournament_id, name="Equipe B 2", rating=1950, club="B"),
        ]
        team_a = self.db.create_team(tournament_id, "Equipe A", club="A", captain="Capitao A")
        team_b = self.db.create_team(tournament_id, "Equipe B", club="B", captain="Capitao B")
        self.db.add_player_to_team(team_a, players[0], board_number=1)
        self.db.add_player_to_team(team_a, players[1], board_number=2)
        self.db.add_player_to_team(team_b, players[2], board_number=1)
        self.db.add_player_to_team(team_b, players[3], board_number=2)

        teams = self.db.list_teams(tournament_id)
        team_players = self.db.list_team_players(team_a)

        self.assertEqual([team["name"] for team in teams], ["Equipe A", "Equipe B"])
        self.assertEqual(teams[0]["starters_count"], 2)
        self.assertEqual([player["board_number"] for player in team_players], [1, 2])

        round_id = self.db.create_round_with_team_matches(
            tournament_id,
            1,
            [
                {
                    "match_number": 1,
                    "white_team_id": team_a,
                    "black_team_id": team_b,
                    "boards": [
                        {
                            "board_number": 1,
                            "white_player_id": players[0],
                            "black_player_id": players[2],
                        },
                        {
                            "board_number": 2,
                            "white_player_id": players[3],
                            "black_player_id": players[1],
                        },
                    ],
                }
            ],
        )
        matches = self.db.list_team_matches_for_round(round_id)
        boards = self.db.list_team_boards(matches[0]["id"])

        self.assertEqual(matches[0]["white_team_name"], "Equipe A")
        self.assertEqual(matches[0]["black_team_name"], "Equipe B")
        self.assertEqual([board["board_number"] for board in boards], [1, 2])
        self.assertEqual(boards[0]["white_player_name"], "Equipe A 1")

    def test_team_service_validates_team_roster(self) -> None:
        tournament_id = self.tournament_service.create_tournament(
            {
                "name": "Equipes validacao",
                "competition_type": "team",
                "rounds_count": "3",
                "bye_points": "1",
            }
        )
        player_a = self.db.create_player(tournament_id, name="Titular A", rating=1900)
        player_b = self.db.create_player(tournament_id, name="Titular B", rating=1800)
        team_id = self.team_service.create_team(
            tournament_id,
            {"name": "Equipe Azul", "club": "Clube", "captain": "Capitao"},
        )

        assignment_id = self.team_service.add_player(team_id, player_a, board_number="1", role="starter")
        assignment = self.db.get_team_player(assignment_id)

        self.assertIsNotNone(assignment)
        self.assertEqual(assignment["board_number"], 1)
        self.assertEqual(assignment["role"], "starter")

        with self.assertRaisesRegex(AppError, "tabuleiro"):
            self.team_service.add_player(team_id, player_b, board_number="1", role="starter")

        with self.assertRaisesRegex(AppError, "equipe Equipe Azul"):
            self.team_service.add_player(team_id, player_a, board_number="2", role="starter")

        self.team_service.add_player(team_id, player_b, board_number="", role="reserve")
        players = self.db.list_team_players(team_id)
        self.assertEqual([player["role"] for player in players], ["starter", "reserve"])

    def test_team_service_rejects_individual_tournament(self) -> None:
        with self.assertRaisesRegex(AppError, "formato Equipes"):
            self.team_service.create_team(self.tournament_id, {"name": "Equipe indevida"})

    def test_team_service_deletes_only_unused_teams(self) -> None:
        tournament_id, team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        unused_team_id = self.team_service.create_team(tournament_id, {"name": "Equipe sem rodada"})

        self.team_service.delete_team(unused_team_id)

        self.assertIsNone(self.db.get_team(unused_team_id))

        self.service.generate_next_round(tournament_id)

        with self.assertRaisesRegex(AppError, "preservar o historico"):
            self.team_service.delete_team(team_ids[0])

        self.assertIsNotNone(self.db.get_team(team_ids[0]))
        self.assertGreater(self.db.count_team_matches(team_ids[0]), 0)

    def test_team_tournament_settings_are_saved_and_validate_board_limit(self) -> None:
        tournament_id = self.tournament_service.create_tournament(
            {
                "name": "Equipes regras",
                "competition_type": "team",
                "rounds_count": "3",
                "bye_points": "1",
            }
        )
        player_id = self.db.create_player(tournament_id, name="Segundo tabuleiro", rating=1800)
        team_id = self.team_service.create_team(tournament_id, {"name": "Equipe Regras"})
        self.team_service.add_player(team_id, player_id, board_number="2", role="starter")

        self.tournament_service.save_profile(
            tournament_id,
            {
                "name": "Equipes regras",
                "competition_type": "team",
                "scope": "standalone",
                "rounds_count": "3",
                "bye_points": "1",
            },
            {
                "team_boards_count": "2",
                "team_match_win_points": "3",
                "team_match_draw_points": "1",
                "team_match_loss_points": "0",
                "team_pairing_method": "swiss",
                "team_standing_primary": "match_points",
                "team_standing_secondary": "game_points",
            },
            [],
        )
        settings = self.db.get_tournament_settings(tournament_id)

        self.assertEqual(settings["team_boards_count"], 2)
        self.assertEqual(settings["team_match_win_points"], 3.0)
        self.assertEqual(settings["team_match_draw_points"], 1.0)
        self.assertEqual(settings["team_match_loss_points"], 0.0)
        self.assertEqual(settings["team_pairing_method"], "swiss")
        self.assertEqual(settings["team_standing_primary"], "match_points")
        self.assertEqual(settings["team_standing_secondary"], "game_points")

        with self.assertRaisesRegex(AppError, "tabuleiro 2"):
            self.tournament_service.save_profile(
                tournament_id,
                {
                    "name": "Equipes regras",
                    "competition_type": "team",
                    "scope": "standalone",
                    "rounds_count": "3",
                    "bye_points": "1",
                },
                {
                    "team_boards_count": "1",
                    "team_match_win_points": "3",
                    "team_match_draw_points": "1",
                    "team_match_loss_points": "0",
                    "team_pairing_method": "swiss",
                    "team_standing_primary": "match_points",
                    "team_standing_secondary": "game_points",
                },
                [],
            )

    def test_team_standing_criteria_must_be_different(self) -> None:
        tournament_id = self.tournament_service.create_tournament(
            {
                "name": "Equipes criterios",
                "competition_type": "team",
                "rounds_count": "3",
                "bye_points": "1",
            }
        )

        with self.assertRaisesRegex(AppError, "criterios diferentes"):
            self.tournament_service.save_profile(
                tournament_id,
                {
                    "name": "Equipes criterios",
                    "competition_type": "team",
                    "scope": "standalone",
                    "rounds_count": "3",
                    "bye_points": "1",
                },
                {
                    "team_boards_count": "4",
                    "team_standing_primary": "match_points",
                    "team_standing_secondary": "match_points",
                },
                [],
            )

    def test_invalid_competition_type_is_rejected(self) -> None:
        with self.assertRaisesRegex(AppError, "Formato do torneio invalido"):
            self.tournament_service.create_tournament(
                {
                    "name": "Formato invalido",
                    "competition_type": "duplas",
                    "rounds_count": "3",
                    "bye_points": "1",
                }
            )

    def test_close_round_creates_backup_and_updates_standings(self) -> None:
        self._create_players(4)
        round_data = self.service.generate_next_round(self.tournament_id)
        self._fill_decisive_results(round_data["id"])

        self.service.close_round(self.tournament_id, round_data["id"])

        backups = list(self.backup_dir.glob("*.db"))
        standings = self.service.standings(self.tournament_id)
        closed_round = self.db.get_round(round_data["id"])

        self.assertEqual(closed_round["status"], "closed")
        self.assertGreaterEqual(len(backups), 2)
        self.assertTrue(any("before_close_round" in backup.name for backup in backups))
        self.assertEqual(standings[0]["points"], 1.0)

    def test_team_pairing_generates_matches_and_boards(self) -> None:
        tournament_id, team_ids = self._create_team_tournament(teams_count=4, boards_count=2)

        round_data = self.service.generate_next_round(tournament_id)
        matches = self.db.list_team_matches_for_round(round_data["id"])
        individual_pairings = self.db.get_pairings_for_round(round_data["id"])

        self.assertEqual(len(matches), 2)
        self.assertEqual(individual_pairings, [])
        scheduled_players = []
        for match in matches:
            boards = self.db.list_team_boards(int(match["id"]))
            self.assertEqual(len(boards), 2)
            self.assertFalse(match["is_bye"])
            scheduled_players.extend(
                player_id
                for board in boards
                for player_id in (board["white_player_id"], board["black_player_id"])
                if player_id
            )
            white_roster = self.db.list_team_players(int(match["white_team_id"]))
            black_roster = self.db.list_team_players(int(match["black_team_id"]))
            white_board_1 = next(item for item in white_roster if item["board_number"] == 1)
            white_board_2 = next(item for item in white_roster if item["board_number"] == 2)
            black_board_1 = next(item for item in black_roster if item["board_number"] == 1)
            black_board_2 = next(item for item in black_roster if item["board_number"] == 2)
            self.assertEqual(boards[0]["white_player_id"], white_board_1["player_id"])
            self.assertEqual(boards[0]["black_player_id"], black_board_1["player_id"])
            self.assertEqual(boards[1]["white_player_id"], black_board_2["player_id"])
            self.assertEqual(boards[1]["black_player_id"], white_board_2["player_id"])

        all_team_players = [
            item["player_id"]
            for team_id in team_ids
            for item in self.db.list_team_players(team_id)
        ]
        self.assertCountEqual(scheduled_players, all_team_players)

    def test_team_swiss_pairing_avoids_repeat_when_possible(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=4, boards_count=2)

        first_round = self.service.generate_next_round(tournament_id)
        first_matches = self.db.list_team_matches_for_round(first_round["id"])
        first_pairs = {
            frozenset((match["white_team_id"], match["black_team_id"]))
            for match in first_matches
            if not match["is_bye"]
        }
        for match in first_matches:
            for board in self.db.list_team_boards(int(match["id"])):
                self.service.update_result(tournament_id, int(board["id"]), "1-0")
        self.service.close_round(tournament_id, first_round["id"])

        second_round = self.service.generate_next_round(tournament_id)
        second_pairs = {
            frozenset((match["white_team_id"], match["black_team_id"]))
            for match in self.db.list_team_matches_for_round(second_round["id"])
            if not match["is_bye"]
        }

        self.assertTrue(first_pairs.isdisjoint(second_pairs))

    def test_odd_team_count_gets_team_bye(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=3, boards_count=2)

        round_data = self.service.generate_next_round(tournament_id)
        matches = self.db.list_team_matches_for_round(round_data["id"])
        bye_matches = [match for match in matches if match["is_bye"]]

        self.assertEqual(len(matches), 2)
        self.assertEqual(len(bye_matches), 1)
        self.assertEqual(bye_matches[0]["result"], "BYE")
        self.assertEqual(self.db.list_team_boards(int(bye_matches[0]["id"])), [])

    def test_team_round_close_calculates_match_summary_and_standings(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        boards = self.db.list_team_boards(int(match["id"]))

        self.service.update_result(tournament_id, int(boards[0]["id"]), "1-0")
        self.service.update_result(tournament_id, int(boards[1]["id"]), "0-1")
        self.service.close_round(tournament_id, round_data["id"])

        updated_match = self.db.list_team_matches_for_round(round_data["id"])[0]
        standings = self.service.team_standings(tournament_id)

        self.assertEqual(updated_match["result"], "1-0")
        self.assertEqual(updated_match["white_game_points"], 2.0)
        self.assertEqual(updated_match["black_game_points"], 0.0)
        self.assertEqual(updated_match["white_match_points"], 2.0)
        self.assertEqual(updated_match["black_match_points"], 0.0)
        self.assertEqual(standings[0]["team_id"], updated_match["white_team_id"])
        self.assertEqual(standings[0]["match_points"], 2.0)
        self.assertEqual(standings[0]["game_points"], 2.0)
        self.assertEqual(standings[0]["wins"], 1)
        self.assertEqual(self.db.get_round(round_data["id"])["status"], "closed")

    def test_team_round_cannot_close_with_pending_board_result(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        boards = self.db.list_team_boards(int(match["id"]))

        self.service.update_result(tournament_id, int(boards[0]["id"]), "1-0")

        with self.assertRaisesRegex(AppError, "resultados dos tabuleiros"):
            self.service.close_round(tournament_id, round_data["id"])

    def test_team_board_manual_color_swap(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        board = self.db.list_team_boards(int(match["id"]))[0]

        self.service.swap_team_board_colors(tournament_id, round_data["id"], int(board["id"]))

        updated = self.db.list_team_boards(int(match["id"]))[0]
        self.assertEqual(updated["white_player_id"], board["black_player_id"])
        self.assertEqual(updated["black_player_id"], board["white_player_id"])

    def test_team_board_manual_player_swap_keeps_player_in_same_team(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        boards = self.db.list_team_boards(int(match["id"]))

        source_player_id = int(boards[0]["white_player_id"])
        replacement_player_id = int(boards[1]["black_player_id"])
        self.service.adjust_team_board_player(
            tournament_id,
            round_data["id"],
            int(boards[0]["id"]),
            "white",
            replacement_player_id,
        )

        updated = self.db.list_team_boards(int(match["id"]))
        self.assertEqual(updated[0]["white_player_id"], replacement_player_id)
        self.assertEqual(updated[1]["black_player_id"], source_player_id)

    def test_team_board_manual_player_swap_rejects_other_team_player(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        boards = self.db.list_team_boards(int(match["id"]))

        with self.assertRaisesRegex(AppError, "mesma equipe"):
            self.service.adjust_team_board_player(
                tournament_id,
                round_data["id"],
                int(boards[0]["id"]),
                "white",
                int(boards[0]["black_player_id"]),
            )

    def test_phase6_team_lineups_substitutions_and_export_are_recorded(self) -> None:
        tournament_id, team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        reserve_id = self.db.create_player(tournament_id, name="Reserva Equipe 1", rating=1600, club="Clube 1")
        self.team_service.add_player(team_ids[0], reserve_id, board_number="", role="reserve")
        settings = self.db.get_tournament_settings(tournament_id) or {}
        settings.update(
            {
                "team_max_substitutions": 2,
                "team_board_order_policy": "fixed",
                "team_reserve_policy": "same_team",
            }
        )
        self.db.save_tournament_settings(tournament_id, settings)

        round_data = self.service.generate_next_round(tournament_id)
        lineups = self.db.list_team_lineups(tournament_id, round_id=int(round_data["id"]))
        self.assertEqual(2, len(lineups))
        self.assertEqual(4, sum(len(self.db.list_team_lineup_boards(int(lineup["id"]))) for lineup in lineups))

        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        board = self.db.list_team_boards(int(match["id"]))[0]
        team_player_ids = {int(item["player_id"]) for item in self.db.list_team_players(team_ids[0], active_only=True)}
        color = "white" if int(board["white_player_id"]) in team_player_ids else "black"

        self.service.adjust_team_board_player(
            tournament_id,
            int(round_data["id"]),
            int(board["id"]),
            color,
            reserve_id,
            reason="Titular chegou atrasado",
        )

        substitutions = self.db.list_team_substitution_events(tournament_id, round_id=int(round_data["id"]))
        self.assertEqual(1, len(substitutions))
        self.assertEqual("Titular chegou atrasado", substitutions[0]["reason"])
        self.assertEqual(reserve_id, substitutions[0]["in_player_id"])
        lineup_player_ids = {
            int(board_row["player_id"])
            for lineup in self.db.list_team_lineups(tournament_id, round_id=int(round_data["id"]))
            for board_row in self.db.list_team_lineup_boards(int(lineup["id"]))
            if board_row.get("player_id")
        }
        self.assertIn(reserve_id, lineup_player_ids)

        export_path = Path(self.temp_dir.name) / "escalacoes.csv"
        self.export_service.export_team_lineups(tournament_id, export_path)
        content = export_path.read_text(encoding="utf-8-sig")
        self.assertIn("Escalacoes por equipes", content)
        self.assertIn("Substituicoes por equipes", content)
        self.assertIn("Reserva Equipe 1", content)

    def test_phase6_team_substitution_after_result_requires_formal_correction(self) -> None:
        tournament_id, team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        reserve_id = self.db.create_player(tournament_id, name="Reserva Bloqueio", rating=1500)
        self.team_service.add_player(team_ids[0], reserve_id, board_number="", role="reserve")
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        board = self.db.list_team_boards(int(match["id"]))[0]
        self.service.update_result(tournament_id, int(board["id"]), "1-0")
        team_player_ids = {int(item["player_id"]) for item in self.db.list_team_players(team_ids[0], active_only=True)}
        color = "white" if int(board["white_player_id"]) in team_player_ids else "black"

        with self.assertRaisesRegex(AppError, "correcao formal"):
            self.service.adjust_team_board_player(
                tournament_id,
                int(round_data["id"]),
                int(board["id"]),
                color,
                reserve_id,
            )

    def test_odd_player_count_does_not_repeat_bye_next_round(self) -> None:
        self._create_players(5)
        first_round = self.service.generate_next_round(self.tournament_id)
        first_bye = [
            pairing["white_player_id"]
            for pairing in self.db.get_pairings_for_round(first_round["id"])
            if pairing["is_bye"]
        ][0]
        self._fill_decisive_results(first_round["id"])
        self.service.close_round(self.tournament_id, first_round["id"])

        second_round = self.service.generate_next_round(self.tournament_id)
        second_bye = [
            pairing["white_player_id"]
            for pairing in self.db.get_pairings_for_round(second_round["id"])
            if pairing["is_bye"]
        ][0]

        self.assertNotEqual(first_bye, second_bye)

    def test_swiss_pairing_uses_global_matching_to_avoid_repeat_when_possible(self) -> None:
        tournament_id = self.tournament_service.create_tournament(
            {"name": "Pareamento global", "rounds_count": "20", "bye_points": "1"}
        )
        player_ids = [
            self.db.create_player(
                tournament_id,
                name=f"Jogador {index + 1}",
                rating=2000 - index * 10,
                club="Clube",
                category="Absoluto",
            )
            for index in range(6)
        ]
        allowed_pairs = {
            frozenset((player_ids[0], player_ids[1])),
            frozenset((player_ids[0], player_ids[2])),
            frozenset((player_ids[1], player_ids[3])),
            frozenset((player_ids[4], player_ids[5])),
        }
        played_pairs = [
            (white_id, black_id)
            for index, white_id in enumerate(player_ids)
            for black_id in player_ids[index + 1 :]
            if frozenset((white_id, black_id)) not in allowed_pairs
        ]
        for round_number, (white_id, black_id) in enumerate(played_pairs, start=1):
            round_id = self.db.create_round_with_pairings(
                tournament_id,
                round_number,
                [
                    {
                        "board_number": 1,
                        "white_player_id": white_id,
                        "black_player_id": black_id,
                        "result": "0F-0F",
                        "is_bye": 0,
                    }
                ],
            )
            self.db.close_round(round_id)

        next_round = self.service.generate_next_round(tournament_id)
        generated_pairs = {
            frozenset((pairing["white_player_id"], pairing["black_player_id"]))
            for pairing in self.db.get_pairings_for_round(next_round["id"])
            if not pairing["is_bye"]
        }

        self.assertTrue(generated_pairs.issubset(allowed_pairs))
        self.assertEqual(
            {
                frozenset((player_ids[0], player_ids[2])),
                frozenset((player_ids[1], player_ids[3])),
                frozenset((player_ids[4], player_ids[5])),
            },
            generated_pairs,
        )

    def test_swiss_pairing_keeps_score_groups_when_possible(self) -> None:
        tournament_id = self.tournament_service.create_tournament(
            {"name": "Grupos de pontuacao", "rounds_count": "5", "bye_points": "1"}
        )
        player_ids = [
            self.db.create_player(
                tournament_id,
                name=f"Jogador {index + 1}",
                rating=2000 - index * 10,
                club="Clube",
                category="Absoluto",
                starting_points=starting_points,
            )
            for index, starting_points in enumerate([2.0, 2.0, 1.0, 1.0, 0.0, 0.0])
        ]
        previous_round_id = self.db.create_round_with_pairings(tournament_id, 1, [])
        self.db.close_round(previous_round_id)

        next_round = self.service.generate_next_round(tournament_id)
        points_by_player = {
            item["player_id"]: item["points"]
            for item in self.service.standings(tournament_id)
        }

        for pairing in self.db.get_pairings_for_round(next_round["id"]):
            if pairing["is_bye"]:
                continue
            self.assertEqual(
                points_by_player[pairing["white_player_id"]],
                points_by_player[pairing["black_player_id"]],
            )
        self.assertCountEqual(
            [
                player_id
                for pairing in self.db.get_pairings_for_round(next_round["id"])
                for player_id in (pairing["white_player_id"], pairing["black_player_id"])
                if player_id
            ],
            player_ids,
        )

    def test_swiss_pairing_avoids_repeating_same_float_when_possible(self) -> None:
        tournament_id = self.tournament_service.create_tournament(
            {"name": "Historico de flutuacao", "rounds_count": "5", "bye_points": "1"}
        )
        player_ids = [
            self.db.create_player(
                tournament_id,
                name=f"Jogador {index + 1}",
                rating=2000 - index * 10,
                club="Clube",
                category="Absoluto",
                starting_points=starting_points,
            )
            for index, starting_points in enumerate([2.0, 2.0, 2.0, 1.0, 1.0, 1.0])
        ]
        prior_floater = player_ids[3]
        prior_round_id = self.db.create_round_with_pairings(
            tournament_id,
            1,
            [
                {
                    "board_number": 1,
                    "white_player_id": player_ids[0],
                    "black_player_id": prior_floater,
                    "result": "0F-0F",
                    "is_bye": 0,
                }
            ],
        )
        self.db.close_round(prior_round_id)

        next_round = self.service.generate_next_round(tournament_id)
        high_score_players = set(player_ids[:3])
        lower_score_players = set(player_ids[3:])
        cross_pairs = []
        for pairing in self.db.get_pairings_for_round(next_round["id"]):
            if pairing["is_bye"]:
                continue
            pair = {pairing["white_player_id"], pairing["black_player_id"]}
            if pair & high_score_players and pair & lower_score_players:
                cross_pairs.append(pair)

        self.assertEqual(1, len(cross_pairs))
        lower_floater = next(iter(cross_pairs[0] & lower_score_players))
        self.assertNotEqual(prior_floater, lower_floater)

    def test_adjust_pairing_player_swaps_players_between_slots(self) -> None:
        player_ids = self._create_players(4)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairings = self.db.get_pairings_for_round(round_data["id"])
        first_pairing = pairings[0]
        second_pairing = pairings[1]

        source_player_id = first_pairing["white_player_id"]
        replacement_player_id = second_pairing["black_player_id"]
        self.service.adjust_pairing_player(
            self.tournament_id,
            round_data["id"],
            first_pairing["id"],
            "white",
            replacement_player_id,
        )

        updated_pairings = self.db.get_pairings_for_round(round_data["id"])
        updated_first = next(pairing for pairing in updated_pairings if pairing["id"] == first_pairing["id"])
        updated_second = next(pairing for pairing in updated_pairings if pairing["id"] == second_pairing["id"])
        scheduled_player_ids = []
        for pairing in updated_pairings:
            scheduled_player_ids.append(pairing["white_player_id"])
            if pairing["black_player_id"]:
                scheduled_player_ids.append(pairing["black_player_id"])

        self.assertEqual(updated_first["white_player_id"], replacement_player_id)
        self.assertEqual(updated_second["black_player_id"], source_player_id)
        self.assertCountEqual(scheduled_player_ids, player_ids)

    def test_member_registration_creates_linked_tournament_player(self) -> None:
        self.db.save_club("Clube Teste", city="Curitiba")
        member_id = self.member_service.create_member(
            {
                "name": "Ana Membro",
                "rating": "1800",
                "category": "Sub-18",
                "member_type": "aluno",
                "status": "active",
            }
        )

        player_id = self.member_service.register_member_in_tournament(self.tournament_id, member_id)
        player = self.db.get_player(player_id)

        self.assertEqual(player["member_id"], member_id)
        self.assertEqual(player["name"], "Ana Membro")
        self.assertEqual(player["club"], "Clube Teste")

    def test_member_surname_formats_tournament_player_name(self) -> None:
        member_id = self.member_service.create_member(
            {
                "name": "Lucas",
                "surname": "Costa",
                "rating": "1700",
                "member_type": "aluno",
                "status": "active",
            }
        )

        player_id = self.member_service.register_member_in_tournament(self.tournament_id, member_id)
        member = self.db.get_member(member_id)
        player = self.db.get_player(player_id)

        self.assertEqual(member["surname"], "Costa")
        self.assertEqual(player["name"], "Lucas Costa")
        self.assertEqual(player["surname"], "Costa")
        self.assertEqual(player["given_name"], "Lucas")

    def test_pairings_use_inverted_names_and_standings_use_full_names(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Lucas",
            surname="Lima",
            given_name="Lucas",
            rating=1800,
        )
        self.db.create_player(
            self.tournament_id,
            name="Rafael",
            surname="Cruz",
            given_name="Rafael",
            rating=1700,
        )

        standings_names = [item["name"] for item in self.service.standings(self.tournament_id)]
        self.assertIn("Lucas Lima", standings_names)
        self.assertIn("Rafael Cruz", standings_names)
        self.assertNotIn("Lima, Lucas", standings_names)

        round_data = self.service.generate_next_round(self.tournament_id)
        _title, _headers, rows = self.export_service._pairings_section(round_data["id"])
        pairing_names = {str(row[1]) for row in rows} | {str(row[4]) for row in rows if row[4] != "BYE"}

        self.assertIn("Lima, Lucas", pairing_names)
        self.assertIn("Cruz, Rafael", pairing_names)
        self.assertNotIn("Lucas Lima", pairing_names)

    def test_delete_unpaired_player_removes_registration(self) -> None:
        player_id = self.db.create_player(
            self.tournament_id,
            name="Cadastro errado",
            rating=1200,
        )

        self.service.delete_player_if_unpaired(self.tournament_id, player_id)

        self.assertIsNone(self.db.get_player(player_id))

    def test_delete_paired_player_is_blocked_to_preserve_history(self) -> None:
        player_ids = self._create_players(2)
        self.service.generate_next_round(self.tournament_id)

        with self.assertRaises(AppError):
            self.service.delete_player_if_unpaired(self.tournament_id, player_ids[0])

        self.assertIsNotNone(self.db.get_player(player_ids[0]))

    def test_member_statuses_control_tournament_registration_eligibility(self) -> None:
        active_id = self.member_service.create_member(
            {
                "name": "Ativo",
                "member_type": "socio",
                "status": "active",
            }
        )
        visitor_id = self.member_service.create_member(
            {
                "name": "Visitante",
                "member_type": "visitante",
                "status": "visitor",
            }
        )
        guest_id = self.member_service.create_member(
            {
                "name": "Convidado",
                "member_type": "convidado",
                "status": "guest",
            }
        )
        withdrawn_id = self.member_service.create_member(
            {
                "name": "Desistente",
                "member_type": "aluno",
                "status": "withdrawn",
            }
        )

        eligible = self.db.list_members_for_tournament(self.tournament_id, active_only=True)
        eligible_ids = {member["id"] for member in eligible}

        self.assertIn(active_id, eligible_ids)
        self.assertNotIn(visitor_id, eligible_ids)
        self.assertNotIn(guest_id, eligible_ids)
        self.assertNotIn(withdrawn_id, eligible_ids)

        with self.assertRaises(AppError):
            self.member_service.register_member_in_tournament(self.tournament_id, visitor_id)

    def test_school_classes_filter_members_for_tournament_registration(self) -> None:
        school_id = self.club_service.save_profile(
            {
                "name": "Escola Alpha",
                "kind": "school",
                "city": "Teresina",
                "active": 1,
            },
            club_id=None,
        )
        other_club_id = self.club_service.save_profile(
            {
                "name": "Clube Beta",
                "kind": "club",
                "active": 1,
            },
            club_id=None,
        )
        class_id = self.club_service.save_class(
            {
                "club_id": school_id,
                "name": "Turma A",
                "teacher": "Professora",
                "weekday": "Segunda",
                "time": "14:00",
                "active": 1,
            }
        )
        school_member_id = self.member_service.create_member(
            {
                "name": "Aluno Escola",
                "rating": "1500",
                "club_id": school_id,
                "class_id": class_id,
                "member_type": "aluno",
                "status": "active",
            }
        )
        other_member_id = self.member_service.create_member(
            {
                "name": "Aluno Outro Clube",
                "rating": "1450",
                "club_id": other_club_id,
                "member_type": "aluno",
                "status": "active",
            }
        )
        school_tournament_id = self.db.create_tournament(
            "Interclasse",
            club_id=school_id,
            rounds_count=3,
        )

        eligible = self.db.list_members_for_tournament(school_tournament_id, active_only=True)
        eligible_ids = {member["id"] for member in eligible}
        player_id = self.member_service.register_member_in_tournament(
            school_tournament_id,
            school_member_id,
        )
        player = self.db.get_player(player_id)
        tournament_players = self.db.list_players(school_tournament_id, active_only=False)
        member = self.db.get_member(school_member_id)
        class_data = self.db.list_classes(club_id=school_id)[0]

        self.assertIn(school_member_id, eligible_ids)
        self.assertNotIn(other_member_id, eligible_ids)
        self.assertEqual(player["club"], "Escola Alpha")
        self.assertEqual(player["active_class_name"], "Turma A")
        self.assertEqual(tournament_players[0]["active_class_name"], "Turma A")
        self.assertEqual(member["club_name"], "Escola Alpha")
        self.assertEqual(member["active_class_name"], "Turma A")
        self.assertEqual(class_data["active_members_count"], 1)

    def test_standalone_tournament_accepts_members_from_any_club(self) -> None:
        club_a_id = self.club_service.save_profile(
            {
                "name": "Clube A",
                "kind": "club",
                "active": 1,
            },
            club_id=None,
        )
        club_b_id = self.club_service.save_profile(
            {
                "name": "Clube B",
                "kind": "club",
                "active": 1,
            },
            club_id=None,
        )
        member_a_id = self.member_service.create_member(
            {
                "name": "Jogador A",
                "club_id": club_a_id,
                "status": "active",
            }
        )
        member_b_id = self.member_service.create_member(
            {
                "name": "Jogador B",
                "club_id": club_b_id,
                "status": "active",
            }
        )
        standalone_id = self.db.create_tournament(
            "Avulso",
            club_id=None,
            class_id=None,
            rounds_count=3,
        )

        eligible = self.db.list_members_for_tournament(standalone_id, active_only=True)
        eligible_ids = {member["id"] for member in eligible}
        player_id = self.member_service.register_member_in_tournament(standalone_id, member_b_id)

        self.assertIn(member_a_id, eligible_ids)
        self.assertIn(member_b_id, eligible_ids)
        self.assertIsNotNone(self.db.get_player(player_id))
        self.assertIsNone(self.db.get_tournament(standalone_id)["club_id"])

    def test_class_tournament_filters_members_by_class(self) -> None:
        school_id = self.club_service.save_profile(
            {
                "name": "Escola Gamma",
                "kind": "school",
                "active": 1,
            },
            club_id=None,
        )
        class_a_id = self.club_service.save_class(
            {
                "club_id": school_id,
                "name": "Turma A",
                "active": 1,
            }
        )
        class_b_id = self.club_service.save_class(
            {
                "club_id": school_id,
                "name": "Turma B",
                "active": 1,
            }
        )
        class_member_id = self.member_service.create_member(
            {
                "name": "Aluno Turma A",
                "club_id": school_id,
                "class_id": class_a_id,
                "status": "active",
            }
        )
        other_class_member_id = self.member_service.create_member(
            {
                "name": "Aluno Turma B",
                "club_id": school_id,
                "class_id": class_b_id,
                "status": "active",
            }
        )
        no_class_member_id = self.member_service.create_member(
            {
                "name": "Aluno Sem Turma",
                "club_id": school_id,
                "status": "active",
            }
        )
        class_tournament_id = self.db.create_tournament(
            "Interturma A",
            club_id=school_id,
            class_id=class_a_id,
            rounds_count=3,
        )

        eligible = self.db.list_members_for_tournament(class_tournament_id, active_only=True)
        eligible_ids = {member["id"] for member in eligible}
        player_id = self.member_service.register_member_in_tournament(
            class_tournament_id,
            class_member_id,
        )

        self.assertEqual(eligible_ids, {class_member_id})
        self.assertIsNotNone(self.db.get_player(player_id))
        self.assertEqual(self.db.get_tournament(class_tournament_id)["class_name"], "Turma A")
        self.assertNotIn(other_class_member_id, eligible_ids)
        self.assertNotIn(no_class_member_id, eligible_ids)
        with self.assertRaises(AppError):
            self.member_service.register_member_in_tournament(class_tournament_id, other_class_member_id)

    def test_class_tournament_can_import_members_from_other_classes_when_allowed(self) -> None:
        school_id = self.club_service.save_profile(
            {
                "name": "Escola Delta",
                "kind": "school",
                "active": 1,
            },
            club_id=None,
        )
        class_a_id = self.club_service.save_class(
            {
                "club_id": school_id,
                "name": "Turma A",
                "active": 1,
            }
        )
        class_b_id = self.club_service.save_class(
            {
                "club_id": school_id,
                "name": "Turma B",
                "active": 1,
            }
        )
        class_member_id = self.member_service.create_member(
            {
                "name": "Aluno Turma A",
                "club_id": school_id,
                "class_id": class_a_id,
                "status": "active",
            }
        )
        other_class_member_id = self.member_service.create_member(
            {
                "name": "Aluno Turma B",
                "club_id": school_id,
                "class_id": class_b_id,
                "status": "active",
            }
        )
        class_tournament_id = self.db.create_tournament(
            "Interturma com convidados",
            club_id=school_id,
            class_id=class_a_id,
            rounds_count=3,
        )

        scoped = self.db.list_members_for_tournament(class_tournament_id, active_only=True)
        all_available = self.db.list_members_for_tournament(
            class_tournament_id,
            active_only=True,
            include_out_of_scope=True,
        )
        imported_player_id = self.member_service.register_member_in_tournament(
            class_tournament_id,
            other_class_member_id,
            allow_out_of_scope=True,
        )
        result = self.member_service.register_active_members_in_tournament(
            class_tournament_id,
            include_out_of_scope=True,
        )

        self.assertEqual({member["id"] for member in scoped}, {class_member_id})
        self.assertEqual({member["id"] for member in all_available}, {class_member_id, other_class_member_id})
        self.assertEqual(self.db.get_player(imported_player_id)["member_id"], other_class_member_id)
        self.assertEqual(result["registered"], 1)
        self.assertEqual(result["skipped"], 1)

    def test_guardian_links_multiple_members_and_keeps_one_primary_contact(self) -> None:
        member_id = self.member_service.create_member(
            {
                "name": "Aluno menor",
                "birth_date": f"{date.today().year - 12}-01-01",
                "member_type": "aluno",
                "status": "active",
            }
        )
        sibling_id = self.member_service.create_member(
            {
                "name": "Irmao menor",
                "birth_date": f"{date.today().year - 10}-01-01",
                "member_type": "aluno",
                "status": "active",
            }
        )
        mother_id = self.guardian_service.create_guardian(
            {"name": "Maria Responsavel", "phone": "1111", "email": "maria@example.com"}
        )
        father_id = self.guardian_service.create_guardian(
            {"name": "Jose Responsavel", "phone": "2222"}
        )

        self.guardian_service.link_guardian_to_member(
            {
                "member_id": member_id,
                "guardian_id": mother_id,
                "relationship": "Mae",
                "primary_contact": 1,
                "emergency_contact": 1,
            }
        )
        self.guardian_service.link_guardian_to_member(
            {
                "member_id": member_id,
                "guardian_id": father_id,
                "relationship": "Pai",
                "primary_contact": 1,
            }
        )
        self.guardian_service.link_guardian_to_member(
            {
                "member_id": sibling_id,
                "guardian_id": mother_id,
                "relationship": "Mae",
                "primary_contact": 1,
            }
        )

        member_guardians = self.db.list_member_guardians(member_id)
        mother_members = self.db.list_guardian_members(mother_id)

        self.assertEqual(len(member_guardians), 2)
        self.assertEqual(member_guardians[0]["guardian_id"], father_id)
        self.assertEqual(member_guardians[0]["primary_contact"], 1)
        self.assertEqual(member_guardians[1]["primary_contact"], 0)
        self.assertCountEqual(
            [item["member_id"] for item in mother_members],
            [member_id, sibling_id],
        )

    def test_minor_members_without_guardians_filters_only_active_minors(self) -> None:
        minor_id = self.member_service.create_member(
            {
                "name": "Menor sem responsavel",
                "birth_date": f"{date.today().year - 9}-01-01",
                "member_type": "aluno",
                "status": "active",
            }
        )
        adult_id = self.member_service.create_member(
            {
                "name": "Adulto",
                "birth_date": f"{date.today().year - 30}-01-01",
                "member_type": "socio",
                "status": "active",
            }
        )
        inactive_minor_id = self.member_service.create_member(
            {
                "name": "Menor inativo",
                "birth_date": f"{date.today().year - 11}-01-01",
                "member_type": "aluno",
                "status": "inactive",
            }
        )
        guardian_id = self.guardian_service.create_guardian({"name": "Responsavel"})

        missing_before = self.guardian_service.minor_members_without_guardians()
        self.guardian_service.link_guardian_to_member(
            {
                "member_id": minor_id,
                "guardian_id": guardian_id,
                "relationship": "Responsavel",
                "primary_contact": 1,
            }
        )
        missing_after = self.guardian_service.minor_members_without_guardians()

        self.assertIn(minor_id, {member["id"] for member in missing_before})
        self.assertNotIn(adult_id, {member["id"] for member in missing_before})
        self.assertNotIn(inactive_minor_id, {member["id"] for member in missing_before})
        self.assertNotIn(minor_id, {member["id"] for member in missing_after})

    def test_training_session_records_attendance_and_exports_report(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Escola de Xadrez",
                "kind": "school",
                "active": 1,
            },
            club_id=None,
        )
        class_id = self.club_service.save_class(
            {
                "club_id": club_id,
                "name": "Turma Iniciantes",
                "teacher": "Instrutora",
                "weekday": "Quarta",
                "time": "15:00",
                "active": 1,
            }
        )
        present_member_id = self.member_service.create_member(
            {
                "name": "Aluno Presente",
                "club_id": club_id,
                "class_id": class_id,
                "member_type": "aluno",
                "status": "active",
            }
        )
        absent_member_id = self.member_service.create_member(
            {
                "name": "Aluno Ausente",
                "club_id": club_id,
                "class_id": class_id,
                "member_type": "aluno",
                "status": "active",
            }
        )
        outside_member_id = self.member_service.create_member(
            {
                "name": "Aluno Outra Turma",
                "club_id": club_id,
                "member_type": "aluno",
                "status": "active",
            }
        )

        session_id = self.training_service.save_session(
            {
                "club_id": club_id,
                "class_id": class_id,
                "title": "Aula de finais",
                "session_type": "aula",
                "session_date": "2026-05-20",
                "start_time": "15:00",
                "end_time": "16:30",
                "instructor": "Instrutora",
                "location": "Sala 1",
                "status": "planned",
            }
        )
        session_members = self.db.list_session_members_for_attendance(session_id)

        result = self.training_service.record_attendance(
            session_id,
            [
                {"member_id": present_member_id, "status": "present"},
                {"member_id": absent_member_id, "status": "absent", "notes": "Avisado"},
            ],
        )
        summary = self.training_service.member_attendance_summary(present_member_id)
        sessions = self.db.list_training_sessions()
        report_rows = self.training_service.attendance_report("2026-05-01", "2026-05-31")
        output_path = Path(self.temp_dir.name) / "presencas.csv"
        self.export_service.export_attendance_report(
            output_path,
            start_date="2026-05-01",
            end_date="2026-05-31",
        )
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertCountEqual(
            [member["id"] for member in session_members],
            [present_member_id, absent_member_id],
        )
        self.assertNotIn(outside_member_id, {member["id"] for member in session_members})
        self.assertEqual(result["saved"], 2)
        self.assertEqual(summary["present"], 1)
        self.assertEqual(summary["attendance_rate"], 100.0)
        self.assertEqual(sessions[0]["present_count"], 1)
        self.assertEqual(sessions[0]["absent_count"], 1)
        self.assertEqual(len(report_rows), 2)
        self.assertIn("Aula de finais", content)
        self.assertIn("Aluno Presente", content)
        self.assertIn("Aluno Ausente", content)

    def test_training_sessions_store_pedagogical_plan_fields(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Escola Pedagogica",
                "kind": "school",
                "active": 1,
            },
            club_id=None,
        )
        class_id = self.club_service.save_class(
            {
                "club_id": club_id,
                "name": "Turma Estrategia",
                "teacher": "Instrutora",
                "active": 1,
            }
        )
        level_id = self.learning_level_service.create_level(
            {
                "name": "Plano Tatico",
                "description": "Reconhece cravadas, garfos e ataques duplos.",
                "display_order": 42,
                "active": 1,
            }
        )

        session_id = self.training_service.save_session(
            {
                "club_id": club_id,
                "class_id": class_id,
                "learning_level_id": level_id,
                "title": "Taticas de ataque duplo",
                "session_type": "aula",
                "session_date": "2026-05-21",
                "start_time": "14:00",
                "end_time": "15:30",
                "instructor": "Instrutora",
                "location": "Sala 3",
                "objective": "Identificar garfos em posicoes simples.",
                "content": "Exercicios de ataque duplo com dama, cavalo e torre.",
                "homework": "Resolver 10 diagramas de garfo.",
                "status": "planned",
                "notes": "Usar tabuleiro mural.",
            }
        )

        session = self.db.get_training_session(session_id)
        sessions = self.db.list_training_sessions(class_id=class_id)

        self.assertIsNotNone(session)
        assert session is not None
        self.assertEqual(session["learning_level_id"], level_id)
        self.assertEqual(session["learning_level_name"], "Plano Tatico")
        self.assertEqual(session["objective"], "Identificar garfos em posicoes simples.")
        self.assertEqual(session["content"], "Exercicios de ataque duplo com dama, cavalo e torre.")
        self.assertEqual(session["homework"], "Resolver 10 diagramas de garfo.")
        self.assertEqual(sessions[0]["learning_level_name"], "Plano Tatico")

        self.training_service.save_session(
            {
                **session,
                "objective": "Aplicar garfos em partidas de treino.",
                "homework": "Anotar dois exemplos encontrados em casa.",
            },
            session_id=session_id,
        )
        updated = self.db.get_training_session(session_id)

        self.assertIsNotNone(updated)
        assert updated is not None
        self.assertEqual(updated["objective"], "Aplicar garfos em partidas de treino.")
        self.assertEqual(updated["homework"], "Anotar dois exemplos encontrados em casa.")
        with self.assertRaisesRegex(AppError, "Nivel pedagogico"):
            self.training_service.save_session(
                {
                    "club_id": club_id,
                    "class_id": class_id,
                    "learning_level_id": 99999,
                    "title": "Aula invalida",
                    "session_type": "aula",
                    "status": "planned",
                }
            )

    def test_exercise_library_lists_attempts_and_training_session_link(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Escola Taticas",
                "kind": "school",
                "active": 1,
            },
            club_id=None,
        )
        class_id = self.club_service.save_class(
            {
                "club_id": club_id,
                "name": "Turma Taticas",
                "teacher": "Instrutor",
                "active": 1,
            }
        )
        level_id = self.learning_level_service.create_level(
            {
                "name": "Tatica Basica",
                "description": "Garfos, cravadas e mates simples.",
                "display_order": 50,
                "active": 1,
            }
        )
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Tatico",
                "club_id": club_id,
                "class_id": class_id,
                "member_type": "aluno",
                "status": "active",
            }
        )
        exercise_id = self.exercise_service.save_exercise(
            {
                "club_id": club_id,
                "learning_level_id": level_id,
                "title": "Mate em 1 com dama",
                "theme": "Mate em 1",
                "difficulty": "easy",
                "fen": "6k1/5ppp/8/8/8/8/5PPP/5QK1 w - - 0 1",
                "solution": "Df7#",
                "objective": "Reconhecer mate direto.",
                "tags": "mate,dama",
            }
        )
        second_exercise_id = self.exercise_service.save_exercise(
            {
                "club_id": club_id,
                "learning_level_id": level_id,
                "title": "Garfo de cavalo",
                "theme": "Garfo",
                "difficulty": "basic",
            }
        )
        list_id = self.exercise_service.save_training_list(
            {
                "club_id": club_id,
                "class_id": class_id,
                "learning_level_id": level_id,
                "name": "Lista de mates e garfos",
                "target_date": "2026-05-22",
                "status": "ready",
                "description": "Treino tatico da turma.",
            }
        )
        self.exercise_service.set_training_list_exercises(list_id, [exercise_id, second_exercise_id])
        session_id = self.training_service.save_session(
            {
                "club_id": club_id,
                "class_id": class_id,
                "learning_level_id": level_id,
                "training_list_id": list_id,
                "title": "Treino tatico guiado",
                "session_type": "treino",
                "session_date": "2026-05-22",
                "status": "planned",
            }
        )
        attempt_id = self.exercise_service.record_attempt(
            {
                "exercise_id": exercise_id,
                "member_id": member_id,
                "list_id": list_id,
                "session_id": session_id,
                "attempt_date": "2026-05-22",
                "result": "correct",
            }
        )

        exercise = self.db.get_exercise(exercise_id)
        training_list = self.db.get_training_list(list_id)
        list_items = self.db.list_training_list_exercises(list_id)
        session = self.db.get_training_session(session_id)
        attempts = self.db.list_exercise_attempts(member_id=member_id)

        self.assertIsNotNone(exercise)
        self.assertIsNotNone(training_list)
        self.assertIsNotNone(session)
        assert exercise is not None
        assert training_list is not None
        assert session is not None
        self.assertEqual(exercise["learning_level_name"], "Tatica Basica")
        self.assertEqual(training_list["exercise_count"], 2)
        self.assertEqual([item["exercise_id"] for item in list_items], [exercise_id, second_exercise_id])
        self.assertEqual(session["training_list_id"], list_id)
        self.assertEqual(session["training_list_name"], "Lista de mates e garfos")
        self.assertEqual(attempts[0]["id"], attempt_id)
        self.assertEqual(attempts[0]["score"], 1.0)
        self.assertEqual(attempts[0]["exercise_title"], "Mate em 1 com dama")

        self.exercise_service.remove_exercise_from_training_list(list_id, exercise_id)
        remaining = self.db.list_training_list_exercises(list_id)

        self.assertEqual([item["exercise_id"] for item in remaining], [second_exercise_id])
        with self.assertRaisesRegex(AppError, "Dificuldade"):
            self.exercise_service.save_exercise({"title": "Invalido", "difficulty": "impossivel"})

    def test_inventory_tracks_items_loans_and_maintenance(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Clube Material",
                "kind": "club",
                "active": 1,
            },
            club_id=None,
        )
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Emprestimo",
                "club_id": club_id,
                "member_type": "aluno",
                "status": "active",
            }
        )
        item_id = self.inventory_service.save_item(
            {
                "club_id": club_id,
                "code": "REL-001",
                "name": "Relogio digital",
                "item_type": "clocks",
                "quantity_total": "2",
                "condition_status": "good",
                "storage_location": "Armario 1",
                "acquisition_date": "2026-05-01",
                "acquisition_value": "150,50",
            }
        )
        loan_id = self.inventory_service.save_loan(
            {
                "item_id": item_id,
                "member_id": member_id,
                "quantity": "1",
                "loan_date": "2026-05-10",
                "due_date": "2026-05-20",
                "status": "open",
                "notes": "Treino em casa",
            }
        )

        item = self.db.get_inventory_item(item_id)
        loans = self.db.list_inventory_loans(item_id=item_id)
        summary = self.inventory_service.inventory_summary()

        self.assertIsNotNone(item)
        assert item is not None
        self.assertEqual(item["borrowed_quantity"], 1)
        self.assertEqual(item["available_quantity"], 1)
        self.assertEqual(loans[0]["member_name"], "Aluno Emprestimo")
        self.assertEqual(summary["borrowed_quantity"], 1)
        self.assertEqual(summary["open_loans"], 1)

        with self.assertRaisesRegex(AppError, "Quantidade indisponivel"):
            self.inventory_service.save_loan(
                {
                    "item_id": item_id,
                    "member_id": member_id,
                    "quantity": "2",
                    "status": "open",
                }
            )
        with self.assertRaisesRegex(AppError, "emprestimos em aberto"):
            self.inventory_service.toggle_item_active(item_id)

        self.inventory_service.return_loan(loan_id, return_date="2026-05-18")
        returned = self.db.get_inventory_loan(loan_id)
        item_after_return = self.db.get_inventory_item(item_id)

        self.assertIsNotNone(returned)
        self.assertIsNotNone(item_after_return)
        assert returned is not None
        assert item_after_return is not None
        self.assertEqual(returned["status"], "returned")
        self.assertEqual(returned["return_date"], "2026-05-18")
        self.assertEqual(item_after_return["available_quantity"], 2)

        maintenance_id = self.inventory_service.save_maintenance(
            {
                "item_id": item_id,
                "opened_date": "2026-05-19",
                "description": "Trocar pilha",
                "cost": "12,75",
                "vendor": "Loja local",
                "status": "open",
            }
        )
        maintenance_open = self.db.get_inventory_maintenance(maintenance_id)
        self.inventory_service.close_maintenance(maintenance_id, resolved_date="2026-05-20")
        maintenance_closed = self.db.get_inventory_maintenance(maintenance_id)

        self.assertIsNotNone(maintenance_open)
        self.assertIsNotNone(maintenance_closed)
        assert maintenance_open is not None
        assert maintenance_closed is not None
        self.assertEqual(maintenance_open["cost"], 12.75)
        self.assertEqual(maintenance_closed["status"], "done")
        self.assertEqual(maintenance_closed["resolved_date"], "2026-05-20")

    def test_financial_plans_payments_status_and_export_report(self) -> None:
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Financeiro",
                "member_type": "aluno",
                "status": "active",
            }
        )
        plan_id = self.finance_service.save_plan(
            {
                "name": "Mensalidade",
                "amount": "120,50",
                "billing_cycle": "monthly",
                "active": 1,
            }
        )

        overdue_payment_id = self.finance_service.save_payment(
            {
                "member_id": member_id,
                "plan_id": plan_id,
                "reference_period": "2000-01",
                "due_date": "2000-01-10",
                "status": "pending",
            }
        )
        paid_payment_id = self.finance_service.save_payment(
            {
                "member_id": member_id,
                "description": "Aula avulsa",
                "reference_period": "2999-01",
                "due_date": "2999-01-10",
                "payment_date": "2999-01-09",
                "amount": "50",
                "status": "paid",
                "method": "Pix",
            }
        )

        payments = self.finance_service.payments_report()
        status = self.finance_service.member_financial_status(member_id)
        summary = self.finance_service.finance_summary()
        output_path = Path(self.temp_dir.name) / "financeiro.csv"
        self.export_service.export_financial_report(output_path)
        content = output_path.read_text(encoding="utf-8-sig")

        overdue_payment = next(payment for payment in payments if payment["id"] == overdue_payment_id)
        paid_payment = next(payment for payment in payments if payment["id"] == paid_payment_id)
        self.assertEqual(overdue_payment["amount"], 120.5)
        self.assertEqual(overdue_payment["effective_status"], "late")
        self.assertEqual(paid_payment["effective_status"], "paid")
        self.assertEqual(status["status"], "late")
        self.assertEqual(status["open_amount"], 120.5)
        self.assertEqual(summary["late"], 1)
        self.assertEqual(summary["paid"], 1)
        self.assertEqual(summary["paid_amount"], 50.0)
        self.assertIn("Aluno Financeiro", content)
        self.assertIn("Mensalidade", content)
        self.assertIn("Atrasado", content)

    def test_finance_generates_recurring_payments_and_receipts(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Escola Financeira",
                "kind": "school",
                "active": 1,
            },
            club_id=None,
        )
        billable_member_id = self.member_service.create_member(
            {
                "name": "Aluno Mensalidade",
                "club_id": club_id,
                "member_type": "aluno",
                "status": "active",
            }
        )
        second_billable_member_id = self.member_service.create_member(
            {
                "name": "Socio Mensalidade",
                "club_id": club_id,
                "member_type": "socio",
                "status": "active",
            }
        )
        visitor_id = self.member_service.create_member(
            {
                "name": "Visitante Sem Cobranca",
                "club_id": club_id,
                "member_type": "visitante",
                "status": "active",
            }
        )
        plan_id = self.finance_service.save_plan(
            {
                "name": "Mensalidade recorrente",
                "amount": "95",
                "billing_cycle": "monthly",
                "active": 1,
            }
        )

        result = self.finance_service.generate_recurring_payments(
            plan_id,
            "2999-06",
            "2999-06-10",
            club_id=club_id,
        )
        duplicate_result = self.finance_service.generate_recurring_payments(
            plan_id,
            "2999-06",
            "2999-06-10",
            club_id=club_id,
        )
        payments = self.finance_service.payments_report(
            start_date="2999-06-01",
            end_date="2999-06-30",
        )
        generated_member_ids = {int(payment["member_id"]) for payment in payments}

        self.assertEqual(result["created"], 2)
        self.assertEqual(result["skipped"], 0)
        self.assertEqual(duplicate_result["created"], 0)
        self.assertEqual(duplicate_result["skipped"], 2)
        self.assertIn(billable_member_id, generated_member_ids)
        self.assertIn(second_billable_member_id, generated_member_ids)
        self.assertNotIn(visitor_id, generated_member_ids)
        for payment in payments:
            self.assertEqual(payment["description"], "Mensalidade recorrente - 2999-06")
            self.assertEqual(payment["reference_period"], "2999-06")
            self.assertEqual(payment["due_date"], "2999-06-10")
            self.assertEqual(payment["amount"], 95.0)

        first_payment_id = int(result["payment_ids"][0])
        receipt_path = Path(self.temp_dir.name) / "recibo.pdf"
        with self.assertRaisesRegex(AppError, "Recibo"):
            self.export_service.export_payment_receipt(first_payment_id, receipt_path)

        self.finance_service.mark_payment_paid(first_payment_id, payment_date="2999-06-09", method="Pix")
        exported_receipt = self.export_service.export_payment_receipt(first_payment_id, receipt_path)

        self.assertEqual(exported_receipt, receipt_path)
        self.assertTrue(receipt_path.exists())
        self.assertEqual(receipt_path.read_bytes()[:4], b"%PDF")

    def test_calendar_events_sync_tournament_and_export_report(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Clube Calendario",
                "kind": "club",
                "active": 1,
            },
            club_id=None,
        )
        tournament_id = self.db.create_tournament(
            "Aberto do Clube",
            club_id=club_id,
            location="Salao",
            start_date="2999-01-10",
            rounds_count=3,
        )

        tournament_event_id = self.event_service.sync_tournament_event(tournament_id)
        meeting_event_id = self.event_service.save_event(
            {
                "club_id": club_id,
                "title": "Reuniao de pais",
                "event_type": "meeting",
                "event_date": "2999-01-05",
                "start_time": "19:00",
                "location": "Sala 2",
                "status": "confirmed",
            }
        )
        canceled_event_id = self.event_service.save_event(
            {
                "club_id": club_id,
                "title": "Evento cancelado",
                "event_type": "social",
                "event_date": "2999-01-03",
                "status": "canceled",
            }
        )

        tournament_event = self.db.get_club_event(tournament_event_id)
        events = self.event_service.events_report("2999-01-01", "2999-01-31")
        upcoming = self.event_service.upcoming_events(limit=10)
        output_path = Path(self.temp_dir.name) / "eventos.csv"
        self.export_service.export_events_report(
            output_path,
            start_date="2999-01-01",
            end_date="2999-01-31",
        )
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertEqual(tournament_event["title"], "Aberto do Clube")
        self.assertEqual(tournament_event["event_type"], "tournament")
        self.assertEqual(tournament_event["tournament_id"], tournament_id)
        self.assertCountEqual([event["id"] for event in events], [tournament_event_id, meeting_event_id, canceled_event_id])
        self.assertNotIn("Evento cancelado", [event["title"] for event in upcoming])
        self.assertIn("Aberto do Clube", content)
        self.assertIn("Reuniao de pais", content)
        self.assertIn("Torneio", content)

    def test_dashboard_overview_and_administrative_package(self) -> None:
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Painel",
                "rating": "1500",
                "member_type": "aluno",
                "status": "active",
            }
        )
        self.finance_service.save_payment(
            {
                "member_id": member_id,
                "description": "Mensalidade painel",
                "due_date": "2999-02-10",
                "amount": "100",
                "status": "pending",
            }
        )
        self.event_service.save_event(
            {
                "title": "Encontro do clube",
                "event_type": "social",
                "event_date": "2999-02-05",
                "status": "confirmed",
            }
        )
        self.db.create_tournament("Torneio painel", start_date="2999-02-20")
        output_path = Path(self.temp_dir.name) / "pacote.csv"

        overview = self.dashboard_service.overview()
        self.export_service.export_administrative_package(
            output_path,
            start_date="2999-02-01",
            end_date="2999-02-28",
        )
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertEqual(overview["summary"]["active_members"], 1)
        self.assertEqual(overview["finance_summary"]["pending_amount"], 100.0)
        self.assertEqual(overview["upcoming_events"][0]["title"], "Encontro do clube")
        self.assertEqual(overview["recent_tournaments"][0]["name"], "Torneio painel")
        self.assertIn("Pacote administrativo", content)
        self.assertIn("Mensalidade painel", content)
        self.assertIn("Encontro do clube", content)
        self.assertIn("Ranking interno", content)

    def test_register_active_members_in_tournament_skips_existing_and_ineligible(self) -> None:
        active_id = self.member_service.create_member(
            {
                "name": "Ativo 1",
                "member_type": "socio",
                "status": "active",
            }
        )
        second_active_id = self.member_service.create_member(
            {
                "name": "Ativo 2",
                "member_type": "aluno",
                "status": "active",
            }
        )
        visitor_id = self.member_service.create_member(
            {
                "name": "Visitante",
                "member_type": "visitante",
                "status": "visitor",
            }
        )
        self.member_service.register_member_in_tournament(self.tournament_id, active_id)

        result = self.member_service.register_active_members_in_tournament(self.tournament_id)
        players = self.db.list_players(self.tournament_id, active_only=False)
        registered_member_ids = {player["member_id"] for player in players}

        self.assertEqual(result["registered"], 1)
        self.assertEqual(result["skipped"], 1)
        self.assertIn(active_id, registered_member_ids)
        self.assertIn(second_active_id, registered_member_ids)
        self.assertNotIn(visitor_id, registered_member_ids)

    def test_member_tournament_history_summarizes_results(self) -> None:
        member_id = self.member_service.create_member(
            {
                "name": "Ana Membro",
                "rating": "1800",
                "member_type": "aluno",
                "status": "active",
            }
        )
        member_player_id = self.member_service.register_member_in_tournament(
            self.tournament_id,
            member_id,
        )
        self.db.create_player(
            self.tournament_id,
            name="Oponente",
            rating=1500,
            club="Clube",
            category="Absoluto",
        )
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(round_data["id"])[0]
        result = "1-0" if pairing["white_player_id"] == member_player_id else "0-1"
        self.db.update_pairing_result(pairing["id"], result)
        self.service.close_round(self.tournament_id, round_data["id"])

        history = self.member_service.tournament_history(member_id)
        results = self.member_service.tournament_results(member_id, self.tournament_id)

        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["name"], "Torneio teste")
        self.assertEqual(history[0]["points"], 1.0)
        self.assertEqual(history[0]["wins"], 1)
        self.assertEqual(history[0]["rounds_played"], 1)
        self.assertEqual(history[0]["last_result"], "Vitoria")
        self.assertEqual(results[0]["opponent"], "Oponente")
        self.assertEqual(results[0]["outcome"], "Vitoria")
        self.assertEqual(results[0]["points"], 1.0)
        self.assertNotEqual(self.service.standings(self.tournament_id)[0]["performance"], "")

    def test_internal_rating_update_creates_history_and_updates_member_rating(self) -> None:
        member_id, _member_player_id = self._create_member_win(rating=1600)

        result = self.internal_rating_service.apply_tournament_ratings(self.tournament_id)
        member = self.db.get_member(member_id)
        history = self.internal_rating_service.member_rating_history(member_id)

        self.assertEqual(result["updated"], 1)
        self.assertEqual(result["duplicates"], 0)
        self.assertEqual(result["external"], 1)
        self.assertGreater(member["rating"], 1600)
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["old_rating"], 1600)
        self.assertEqual(history[0]["new_rating"], member["rating"])
        self.assertEqual(history[0]["games"], 1)
        self.assertEqual(history[0]["performance"], 2300)

    def test_internal_rating_update_is_idempotent(self) -> None:
        member_id, _member_player_id = self._create_member_win(rating=1600)

        first = self.internal_rating_service.apply_tournament_ratings(self.tournament_id)
        rating_after_first = self.db.get_member(member_id)["rating"]
        second = self.internal_rating_service.apply_tournament_ratings(self.tournament_id)
        rating_after_second = self.db.get_member(member_id)["rating"]
        history = self.internal_rating_service.member_rating_history(member_id)

        self.assertEqual(first["updated"], 1)
        self.assertEqual(second["updated"], 0)
        self.assertEqual(second["duplicates"], 1)
        self.assertEqual(rating_after_second, rating_after_first)
        self.assertEqual(len(history), 1)

    def test_next_tournament_registration_uses_updated_internal_rating(self) -> None:
        member_id, _member_player_id = self._create_member_win(rating=1600)
        self.internal_rating_service.apply_tournament_ratings(self.tournament_id)
        updated_rating = self.db.get_member(member_id)["rating"]
        next_tournament_id = self.db.create_tournament("Torneio seguinte", rounds_count=3)

        next_player_id = self.member_service.register_member_in_tournament(
            next_tournament_id,
            member_id,
        )
        next_player = self.db.get_player(next_player_id)

        self.assertEqual(next_player["rating"], updated_rating)

    def test_internal_ranking_summarizes_rating_results_and_exports(self) -> None:
        member_id, _member_player_id = self._create_member_win(rating=1600)
        self.internal_rating_service.apply_tournament_ratings(self.tournament_id)
        self.member_service.update_member(
            member_id,
            {
                "name": "Aluno Rating",
                "rating": str(self.db.get_member(member_id)["rating"]),
                "category": "Sub-18",
                "member_type": "aluno",
                "status": "active",
            },
        )
        self.member_service.create_member(
            {
                "name": "Aluno Sem Jogos",
                "rating": "1200",
                "category": "Sub-18",
                "member_type": "aluno",
                "status": "active",
            }
        )
        output_path = Path(self.temp_dir.name) / "ranking.csv"

        ranking = self.internal_rating_service.ranking(category="Sub-18")
        self.export_service.export_internal_ranking_report(output_path, category="Sub-18")
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertEqual(ranking[0]["member_id"], member_id)
        self.assertEqual(ranking[0]["games"], 1)
        self.assertEqual(ranking[0]["wins"], 1)
        self.assertEqual(ranking[0]["score_rate"], 100.0)
        self.assertGreater(ranking[0]["last_delta"], 0)
        self.assertIn("Ranking interno", content)
        self.assertIn("Aluno Rating", content)
        self.assertIn("Aproveitamento", content)

    def test_internal_ranking_filters_by_season_club_and_class(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Escola Ranking",
                "kind": "school",
                "active": 1,
            },
            club_id=None,
        )
        class_id = self.club_service.save_class(
            {
                "club_id": club_id,
                "name": "Turma Ranking",
                "teacher": "Professor",
                "active": 1,
            }
        )
        member_id, _member_player_id = self._create_member_win(rating=1600)
        self.db.update_tournament_details(
            self.tournament_id,
            name="Torneio temporada",
            club_id=club_id,
            class_id=class_id,
            location="Sala 1",
            rounds_count=5,
            start_date="2026-05-10",
        )
        self.internal_rating_service.apply_tournament_ratings(self.tournament_id)
        updated_rating = self.db.get_member(member_id)["rating"]
        self.member_service.update_member(
            member_id,
            {
                "name": "Aluno Temporada",
                "club_id": club_id,
                "class_id": class_id,
                "rating": str(updated_rating),
                "category": "Sub-18",
                "member_type": "aluno",
                "status": "active",
            },
        )
        outside_member_id = self.member_service.create_member(
            {
                "name": "Aluno Fora",
                "rating": "1900",
                "category": "Sub-18",
                "member_type": "aluno",
                "status": "active",
            }
        )
        output_path = Path(self.temp_dir.name) / "ranking_temporada.csv"

        ranking = self.internal_rating_service.ranking(
            category="Sub-18",
            club_id=club_id,
            class_id=class_id,
            start_date="2026-01-01",
            end_date="2026-12-31",
        )
        empty_season = self.internal_rating_service.ranking(
            category="Sub-18",
            club_id=club_id,
            class_id=class_id,
            start_date="2027-01-01",
            end_date="2027-12-31",
        )
        self.export_service.export_internal_ranking_report(
            output_path,
            category="Sub-18",
            club_id=club_id,
            class_id=class_id,
            start_date="2026-01-01",
            end_date="2026-12-31",
        )
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertEqual([item["member_id"] for item in ranking], [member_id])
        self.assertNotIn(outside_member_id, {item["member_id"] for item in ranking})
        self.assertEqual(ranking[0]["games"], 1)
        self.assertGreater(ranking[0]["season_delta"], 0)
        self.assertGreater(ranking[0]["last_delta"], 0)
        self.assertEqual(empty_season[0]["games"], 0)
        self.assertEqual(empty_season[0]["last_delta"], 0)
        self.assertIn("Escola Ranking", content)
        self.assertIn("Turma Ranking", content)
        self.assertIn("Temporada inicial", content)
        self.assertIn("2026-01-01", content)

        with self.assertRaisesRegex(AppError, "temporada"):
            self.internal_rating_service.ranking(start_date="2026-13-01")

    def test_export_member_evolution_report_includes_rating_and_results(self) -> None:
        member_id, _member_player_id = self._create_member_win(rating=1600)
        self.internal_rating_service.apply_tournament_ratings(self.tournament_id)
        guardian_id = self.guardian_service.create_guardian(
            {"name": "Maria Responsavel", "phone": "1111"}
        )
        self.guardian_service.link_guardian_to_member(
            {
                "member_id": member_id,
                "guardian_id": guardian_id,
                "relationship": "Mae",
                "primary_contact": 1,
            }
        )
        output_path = Path(self.temp_dir.name) / "evolucao.csv"

        self.export_service.export_member_evolution(member_id, output_path)
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertIn("Aluno Rating", content)
        self.assertIn("Maria Responsavel", content)
        self.assertIn("Mae", content)
        self.assertIn("Evolucao por torneio", content)
        self.assertIn("Rating interno", content)
        self.assertIn("Resultados", content)
        self.assertIn("Torneio teste", content)
        self.assertIn("2300", content)
        self.assertIn("Vitoria", content)

    def test_export_club_report_includes_administrative_summary(self) -> None:
        self.db.save_club("Clube Teste", city="Teresina", email="contato@example.com")
        self.member_service.create_member(
            {
                "name": "Aluno Ativo",
                "rating": "1700",
                "category": "Sub-18",
                "member_type": "aluno",
                "status": "active",
            }
        )
        output_path = Path(self.temp_dir.name) / "clube.csv"

        self.export_service.export_club_report(output_path)
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertIn("Clube Teste", content)
        self.assertIn("Indicadores", content)
        self.assertIn("Membros ativos", content)
        self.assertIn("Ranking interno", content)
        self.assertIn("Aluno Ativo", content)

    def test_export_tournaments_period_report_filters_by_date(self) -> None:
        self.db.create_tournament("Torneio maio", start_date="2026-05-10")
        self.db.create_tournament("Torneio junho", start_date="2026-06-10")
        output_path = Path(self.temp_dir.name) / "torneios_periodo.csv"

        self.export_service.export_tournaments_period_report(
            output_path,
            start_date="2026-05-01",
            end_date="2026-05-31",
        )
        content = output_path.read_text(encoding="utf-8-sig")

        self.assertIn("Torneios por periodo", content)
        self.assertIn("Torneio maio", content)
        self.assertNotIn("Torneio junho", content)

    def test_app_settings_and_backup_restore_roundtrip(self) -> None:
        export_dir = Path(self.temp_dir.name) / "exports"
        self.db.save_app_settings(
            {
                "appearance_mode": "Dark",
                "default_export_dir": str(export_dir),
                "backup_dir": str(self.backup_dir),
                "arbitration_auto_refresh_enabled": "0",
                "arbitration_refresh_interval_seconds": "60",
                "arbitration_inline_tables_limit": "40",
            }
        )
        before_member_id = self.member_service.create_member(
            {
                "name": "Antes do backup",
                "rating": "1500",
                "member_type": "aluno",
                "status": "active",
            }
        )
        backup_path = self.db.backup("roundtrip")
        self.member_service.create_member(
            {
                "name": "Depois do backup",
                "rating": "1600",
                "member_type": "aluno",
                "status": "active",
            }
        )

        safety_backup = self.db.restore_backup(backup_path)
        settings = self.db.get_app_settings()
        members = self.db.list_members(active_only=False)
        member_names = {member["name"] for member in members}
        backups = self.db.list_backups()

        self.assertEqual(settings["appearance_mode"], "Dark")
        self.assertEqual(settings["default_export_dir"], str(export_dir))
        self.assertEqual(settings["arbitration_auto_refresh_enabled"], "0")
        self.assertEqual(settings["arbitration_refresh_interval_seconds"], "60")
        self.assertEqual(settings["arbitration_inline_tables_limit"], "40")
        self.assertTrue(safety_backup.exists())
        self.assertTrue(any(backup["name"] == backup_path.name for backup in backups))
        self.assertIn("Antes do backup", member_names)
        self.assertNotIn("Depois do backup", member_names)
        self.assertIsNotNone(self.db.get_member(before_member_id))

    def test_backup_copies_to_configured_cloud_dir(self) -> None:
        cloud_dir = Path(self.temp_dir.name) / "nuvem"
        cloud_dir.mkdir()
        self.db.save_app_settings({"cloud_sync_dir": str(cloud_dir)})

        backup_path = self.db.backup("com_nuvem")

        self.assertEqual(self.db.last_cloud_backup["status"], "success")
        self.assertTrue((cloud_dir / backup_path.name).exists())
        self.assertEqual(self.db.last_cloud_backup["path"], str(cloud_dir / backup_path.name))

    def test_backup_without_cloud_dir_reports_not_configured(self) -> None:
        self.db.save_app_settings({"cloud_sync_dir": ""})

        self.db.backup("sem_nuvem")

        self.assertEqual(self.db.last_cloud_backup["status"], "not_configured")

    def test_backup_with_invalid_cloud_dir_does_not_raise(self) -> None:
        missing = Path(self.temp_dir.name) / "nao_existe"
        self.db.save_app_settings({"cloud_sync_dir": str(missing)})

        # Falha de nuvem nao pode derrubar o backup local.
        backup_path = self.db.backup("nuvem_invalida")

        self.assertTrue(backup_path.exists())
        self.assertEqual(self.db.last_cloud_backup["status"], "invalid_directory")

    def test_create_backup_reports_cloud_status_from_db(self) -> None:
        cloud_dir = Path(self.temp_dir.name) / "nuvem2"
        cloud_dir.mkdir()
        self.db.save_app_settings({"cloud_sync_dir": str(cloud_dir)})

        result = self.security_service.create_backup("manual")

        self.assertEqual(result["cloud_status"], "success")

    def test_send_bulk_email_collects_per_recipient_results(self) -> None:
        from src.services.message_service import MessageService

        class _RecordingMailer(MessageService):
            def __init__(self, fail=()):
                super().__init__()
                self.fail = set(fail)
                self.sent: list[str] = []

            def send_email(self, to_email, subject, body):
                if to_email in self.fail:
                    raise RuntimeError("smtp down")
                self.sent.append(to_email)
                return True

        mailer = _RecordingMailer(fail={"erro@x.com"})
        recipients = [
            {"member_id": 1, "email": "ok@x.com"},
            {"member_id": 2, "email": "erro@x.com"},
            {"member_id": 3, "email": ""},  # sem e-mail -> falha sem tentar enviar
        ]

        summary = mailer.send_bulk_email(recipients, "Assunto", "Corpo")

        self.assertEqual(summary["sent_count"], 1)
        self.assertEqual(summary["failed_count"], 2)
        self.assertEqual(mailer.sent, ["ok@x.com"])
        self.assertEqual(summary["total"], 3)

    def test_bulk_email_members_targets_audience_and_logs(self) -> None:
        from src.services.dashboard_service import CommunicationService

        class _FakeMailer:
            def __init__(self):
                self.calls: list[tuple[str, str, str]] = []

            def send_bulk_email(self, recipients, subject, body):
                sent = []
                for r in recipients:
                    self.calls.append((r["email"], subject, body))
                    sent.append(r)
                return {"total": len(recipients), "sent": sent, "failed": [],
                        "sent_count": len(sent), "failed_count": 0}

        ana = self.member_service.create_member(
            {"name": "Ana", "rating": "1500", "member_type": "aluno",
             "status": "active", "email": "ana@x.com"}
        )
        self.member_service.create_member(
            {"name": "Beto", "rating": "1600", "member_type": "socio",
             "status": "active", "email": "beto@x.com"}
        )
        self.member_service.create_member(
            {"name": "Carla", "rating": "1400", "member_type": "aluno",
             "status": "active", "email": ""}  # sem e-mail -> ignorada
        )

        comm = CommunicationService(self.db)
        mailer = _FakeMailer()

        summary = comm.bulk_email_members(
            mailer, "Aviso", "Corpo do aviso", active_only=True, member_type="aluno"
        )

        # So Ana (aluno com e-mail). Beto e socio; Carla nao tem e-mail.
        self.assertEqual(summary["sent_count"], 1)
        self.assertEqual(summary["recipients_total"], 1)
        self.assertEqual(summary["skipped_no_email"], 1)
        self.assertEqual(summary["failed_count"], 0)
        self.assertEqual([call[0] for call in mailer.calls], ["ana@x.com"])
        # Envio bem-sucedido foi registrado em communication_logs.
        logs = comm.list_member_communications(ana)
        self.assertEqual(len(logs), 1)
        self.assertEqual(logs[0]["subject"], "Aviso")

    def test_bulk_email_members_requires_subject_and_body(self) -> None:
        from src.services.dashboard_service import CommunicationService

        comm = CommunicationService(self.db)
        with self.assertRaises(ValueError):
            comm.bulk_email_members(object(), "", "Corpo")

    def test_schedule_email_normalizes_and_persists(self) -> None:
        from src.services.dashboard_service import CommunicationService

        comm = CommunicationService(self.db)
        msg_id = comm.schedule_email(
            "Assunto", "Corpo", "2026-06-01 09:30", audience_kind="all_active"
        )

        pending = comm.list_scheduled_messages(status="pending")
        self.assertEqual(len(pending), 1)
        self.assertEqual(pending[0]["id"], msg_id)
        self.assertEqual(pending[0]["scheduled_at"], "2026-06-01 09:30:00")

    def test_schedule_email_rejects_bad_datetime(self) -> None:
        from src.services.dashboard_service import CommunicationService

        comm = CommunicationService(self.db)
        with self.assertRaises(ValueError):
            comm.schedule_email("A", "B", "amanha de manha")

    def test_dispatch_sends_due_and_skips_future(self) -> None:
        from src.services.dashboard_service import CommunicationService

        class _FakeMailer:
            def __init__(self):
                self.calls: list[str] = []

            def send_bulk_email(self, recipients, subject, body):
                for r in recipients:
                    self.calls.append(r["email"])
                return {"total": len(recipients), "sent": list(recipients), "failed": [],
                        "sent_count": len(recipients), "failed_count": 0}

        self.member_service.create_member(
            {"name": "Ana", "rating": "1500", "member_type": "aluno",
             "status": "active", "email": "ana@x.com"}
        )
        comm = CommunicationService(self.db)
        due_id = comm.schedule_email("Vencida", "Corpo", "2000-01-01 00:00")
        comm.schedule_email("Futura", "Corpo", "2999-01-01 00:00")

        mailer = _FakeMailer()
        dispatched = comm.dispatch_due_scheduled_messages(mailer)

        self.assertEqual(len(dispatched), 1)
        self.assertEqual(dispatched[0]["id"], due_id)
        self.assertEqual(dispatched[0]["status"], "sent")
        self.assertEqual(mailer.calls, ["ana@x.com"])
        pending = comm.list_scheduled_messages(status="pending")
        self.assertEqual(len(pending), 1)
        self.assertEqual(pending[0]["subject"], "Futura")

    def test_cancel_scheduled_message_prevents_dispatch(self) -> None:
        from src.services.dashboard_service import CommunicationService

        class _FakeMailer:
            def send_bulk_email(self, recipients, subject, body):
                raise AssertionError("nao deveria enviar uma mensagem cancelada")

        comm = CommunicationService(self.db)
        msg_id = comm.schedule_email("X", "Y", "2000-01-01 00:00")
        comm.cancel_scheduled_message(msg_id)

        dispatched = comm.dispatch_due_scheduled_messages(_FakeMailer())

        self.assertEqual(dispatched, [])
        self.assertEqual(comm.list_scheduled_messages(status="pending"), [])

    def test_app_settings_normalizes_legacy_default_paths(self) -> None:
        base_path = Path(self.temp_dir.name)
        legacy_export_dir = base_path / "repo" / "exports"
        legacy_backup_dir = base_path / "repo" / "backups"
        user_data_dir = base_path / "user-data"
        self.db.save_app_settings(
            {
                "default_export_dir": str(legacy_export_dir),
                "backup_dir": str(legacy_backup_dir),
            }
        )

        with (
            mock.patch.dict(os.environ, {APP_DATA_DIR_ENV_VAR: str(user_data_dir)}),
            mock.patch.object(database_module, "LEGACY_EXPORTS_DIR", legacy_export_dir),
            mock.patch.object(database_module, "LEGACY_BACKUP_DIR", legacy_backup_dir),
        ):
            settings = self.db.get_app_settings()

        self.assertEqual(settings["default_export_dir"], str(user_data_dir / "exports"))
        self.assertEqual(settings["backup_dir"], str(user_data_dir / "backups"))
        self.assertEqual(self.db.backup_dir, user_data_dir / "backups")

    def test_security_settings_audit_and_backup_retention(self) -> None:
        self.security_service.save_security_settings(
            {
                "backup_retention_count": "2",
            }
        )
        first = self.security_service.create_backup("seguranca_1")
        second = self.security_service.create_backup("seguranca_2")
        third = self.security_service.create_backup("seguranca_3")

        settings = self.db.get_app_settings()
        audit_rows = self.security_service.list_audit_logs(limit=20)
        backups = self.db.list_backups()
        backup_names = {backup["name"] for backup in backups}

        self.assertEqual(settings["backup_retention_count"], "2")
        self.assertTrue(first["path"].exists() or first["path"].name not in backup_names)
        self.assertTrue(second["path"].exists() or second["path"].name not in backup_names)
        self.assertTrue(third["path"].exists() or third["path"].name not in backup_names)
        self.assertLessEqual(len(backups), 2)
        self.assertTrue(any(row["action"] == "settings_saved" for row in audit_rows))
        self.assertTrue(any(row["action"] == "backup_created" for row in audit_rows))
        self.assertTrue(any(row["action"] == "backup_retention_applied" for row in audit_rows))

    def test_restore_corrupted_backup_keeps_current_database_usable(self) -> None:
        before_member_id = self.member_service.create_member(
            {
                "name": "Antes da falha",
                "rating": "1500",
                "member_type": "aluno",
                "status": "active",
            }
        )
        corrupted_backup = self.backup_dir / "corrompido.db"
        corrupted_backup.parent.mkdir(parents=True, exist_ok=True)
        corrupted_backup.write_text("nao e um sqlite valido", encoding="utf-8")

        with self.assertRaises(sqlite3.DatabaseError):
            self.db.restore_backup(corrupted_backup)
        after_member_id = self.member_service.create_member(
            {
                "name": "Depois da falha",
                "rating": "1510",
                "member_type": "aluno",
                "status": "active",
            }
        )
        member_names = {member["name"] for member in self.db.list_members(active_only=False)}

        self.assertIsNotNone(self.db.get_member(before_member_id))
        self.assertIsNotNone(self.db.get_member(after_member_id))
        self.assertIn("Antes da falha", member_names)
        self.assertIn("Depois da falha", member_names)

    def test_restore_legacy_backup_migrates_restored_database(self) -> None:
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        legacy_path = self.backup_dir / "legacy_restore.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE tournaments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    location TEXT DEFAULT '',
                    start_date TEXT DEFAULT '',
                    end_date TEXT DEFAULT '',
                    system TEXT NOT NULL DEFAULT 'Suico',
                    rounds_count INTEGER NOT NULL DEFAULT 5,
                    time_control TEXT DEFAULT '',
                    bye_points REAL NOT NULL DEFAULT 1.0,
                    status TEXT NOT NULL DEFAULT 'draft',
                    created_at TEXT NOT NULL
                );

                CREATE TABLE players (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    tournament_id INTEGER NOT NULL,
                    name TEXT NOT NULL,
                    club TEXT DEFAULT '',
                    federation_id TEXT DEFAULT '',
                    fide_id TEXT DEFAULT '',
                    rating INTEGER NOT NULL DEFAULT 0,
                    category TEXT DEFAULT '',
                    birth_date TEXT DEFAULT '',
                    active INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL
                );
                """
            )
            connection.commit()
        finally:
            connection.close()

        safety_backup = self.db.restore_backup(legacy_path)
        with self.db.connect() as restored_connection:
            user_version = restored_connection.execute("PRAGMA user_version").fetchone()[0]
            issuance_table = restored_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'certificate_issuances'
                """
            ).fetchone()
            template_count = restored_connection.execute(
                "SELECT COUNT(*) FROM certificate_templates"
            ).fetchone()[0]

        self.assertTrue(safety_backup.exists())
        self.assertEqual(Database.SCHEMA_VERSION, user_version)
        self.assertIsNotNone(issuance_table)
        self.assertGreaterEqual(template_count, 7)
        self.assertEqual([], self.db.list_certificate_issuances(limit=1))

    def test_withdrawn_player_keeps_history_but_is_not_paired_again(self) -> None:
        player_ids = self._create_players(4)
        first_round = self.service.generate_next_round(self.tournament_id)
        self._fill_decisive_results(first_round["id"])
        self.service.close_round(self.tournament_id, first_round["id"])

        self.db.set_player_status(player_ids[0], "withdrawn")
        second_round = self.service.generate_next_round(self.tournament_id)
        scheduled_player_ids = []
        for pairing in self.db.get_pairings_for_round(second_round["id"]):
            scheduled_player_ids.append(pairing["white_player_id"])
            if pairing["black_player_id"]:
                scheduled_player_ids.append(pairing["black_player_id"])
        withdrawn_standing = next(
            item for item in self.service.standings(self.tournament_id)
            if item["player_id"] == player_ids[0]
        )

        self.assertNotIn(player_ids[0], scheduled_player_ids)
        self.assertEqual(withdrawn_standing["player_status"], "withdrawn")
        self.assertGreaterEqual(withdrawn_standing["points"], 0.0)

    def test_late_entry_points_are_added_to_new_players(self) -> None:
        self.tournament_service.save_profile(
            self.tournament_id,
            {
                "name": "Torneio teste",
                "location": "",
                "rounds_count": "3",
                "time_control": "",
                "start_date": "",
                "end_date": "",
                "bye_points": "1",
            },
            {
                "initial_order": "rating",
                "tournament_type": "real",
                "late_entry_points": "0.5",
            },
            [],
        )
        self._create_players(2)
        first_round = self.service.generate_next_round(self.tournament_id)
        self._fill_decisive_results(first_round["id"])
        self.service.close_round(self.tournament_id, first_round["id"])

        late_player_id = self.db.create_player(
            self.tournament_id,
            name="Entrada Tardia",
            rating=1200,
        )
        standing = next(
            item for item in self.service.standings(self.tournament_id)
            if item["player_id"] == late_player_id
        )

        self.assertEqual(self.db.get_player(late_player_id)["starting_points"], 0.5)
        self.assertEqual(standing["points"], 0.5)

    def test_disable_bye_requires_even_active_players(self) -> None:
        self.tournament_service.save_profile(
            self.tournament_id,
            {
                "name": "Torneio teste",
                "location": "",
                "rounds_count": "3",
                "time_control": "",
                "start_date": "",
                "end_date": "",
                "bye_points": "1",
            },
            {
                "initial_order": "rating",
                "tournament_type": "real",
                "disable_bye": 1,
            },
            [],
        )
        self._create_players(3)

        with self.assertRaises(AppError):
            self.service.generate_next_round(self.tournament_id)

    def test_closed_result_edit_requires_dangerous_changes_flag(self) -> None:
        self._create_players(2)
        first_round = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(first_round["id"])[0]
        self.db.update_pairing_result(pairing["id"], "1-0")
        self.service.close_round(self.tournament_id, first_round["id"])

        with self.assertRaises(AppError):
            self.service.update_result(self.tournament_id, pairing["id"], "0-1")

        self.tournament_service.save_profile(
            self.tournament_id,
            {
                "name": "Torneio teste",
                "location": "",
                "rounds_count": "5",
                "time_control": "",
                "start_date": "",
                "end_date": "",
                "bye_points": "1",
            },
            {
                "initial_order": "rating",
                "tournament_type": "real",
                "allow_dangerous_changes": 1,
            },
            [],
        )
        self.service.update_result(self.tournament_id, pairing["id"], "0-1")

        self.assertEqual(self.db.get_pairing(pairing["id"])["result"], "0-1")

    def test_tournament_settings_and_schedule_are_saved(self) -> None:
        self.tournament_service.save_profile(
            self.tournament_id,
            {
                "name": "Aberto Escolar",
                "location": "Clube Central",
                "start_date": "2026-06-01",
                "end_date": "2026-06-03",
                "rounds_count": "3",
                "time_control": "Rapido 15+10",
                "bye_points": "1",
            },
            {
                "fide_event_id": "123456",
                "organizer": "Clube Central",
                "website": "https://example.com",
                "contact_email": "contato@example.com",
                "director": "Diretor",
                "chief_arbiter": "Arbitro",
                "arbiters": "Auxiliar",
                "federation": "BRA",
                "state": "PI",
                "categories": "S10,S12,ABS",
                "cutoff_date": "2026-01-01",
                "comments": "Premiacao escolar",
                "prizes": "Medalhas",
                "initial_order": "international_then_national",
                "tournament_type": "real",
                "tournament_profile": "club",
                "allow_public_registration": 1,
                "calculate_performance": 1,
                "late_entry_points": "0.5",
                "rating_fee_fide": "2.50",
                "rating_fee_cbx": "3.75",
                "rating_fee_lbx": "1.25",
            },
            [
                {"round_number": 1, "date": "2026-06-01", "time": "09:00"},
                {"round_number": 2, "date": "2026-06-01", "time": "14:00"},
                {"round_number": 3, "date": "2026-06-02", "time": "09:00"},
            ],
        )

        tournament = self.db.get_tournament(self.tournament_id)
        settings = self.db.get_tournament_settings(self.tournament_id)
        schedule = self.db.list_round_schedule(self.tournament_id)

        self.assertEqual(tournament["name"], "Aberto Escolar")
        self.assertEqual(tournament["rounds_count"], 3)
        self.assertEqual(settings["fide_event_id"], "123456")
        self.assertEqual(settings["initial_order"], "international_then_national")
        self.assertEqual(settings["tournament_profile"], "club")
        self.assertEqual(settings["allow_public_registration"], 1)
        self.assertEqual(settings["calculate_performance"], 1)
        self.assertEqual(settings["late_entry_points"], 0.5)
        self.assertEqual(settings["rating_fee_fide"], 2.5)
        self.assertEqual(settings["rating_fee_cbx"], 3.75)
        self.assertEqual(settings["rating_fee_lbx"], 1.25)
        self.assertEqual(len(schedule), 3)
        self.assertEqual(schedule[1]["date"], "2026-06-01")
        self.assertEqual(schedule[1]["time"], "14:00")

    def test_tournament_profile_defaults_to_free_and_rejects_invalid(self) -> None:
        settings = self.db.get_tournament_settings(self.tournament_id)

        self.assertEqual(settings["tournament_profile"], "free")

        with self.assertRaisesRegex(AppError, "Perfil do torneio"):
            self.tournament_service.save_profile(
                self.tournament_id,
                {
                    "name": "Perfil invalido",
                    "location": "",
                    "rounds_count": "3",
                    "time_control": "",
                    "start_date": "",
                    "end_date": "",
                    "bye_points": "1",
                },
                {"tournament_profile": "oficial-obrigatorio"},
                [],
            )

    def test_tournament_schedule_generator_can_fill_all_rounds_same_day(self) -> None:
        generated = TournamentService.generate_round_schedule(
            rounds_count=5,
            start_date="2026-06-01",
            first_time="08:00",
            round_duration_minutes="45",
            break_minutes="15",
            rounds_per_day=5,
        )

        self.assertEqual([item["date"] for item in generated], ["2026-06-01"] * 5)
        self.assertEqual([item["time"] for item in generated], ["08:00", "09:00", "10:00", "11:00", "12:00"])

    def test_tournament_schedule_generator_splits_rounds_by_day(self) -> None:
        generated = TournamentService.generate_round_schedule(
            rounds_count=5,
            start_date="2026-06-01",
            first_time="08:30",
            round_duration_minutes=30,
            break_minutes=10,
            rounds_per_day=2,
        )

        self.assertEqual(
            [item["date"] for item in generated],
            ["2026-06-01", "2026-06-01", "2026-06-02", "2026-06-02", "2026-06-03"],
        )
        self.assertEqual([item["time"] for item in generated], ["08:30", "09:10", "08:30", "09:10", "08:30"])

    def test_tournament_profile_saves_standalone_and_class_scope(self) -> None:
        school_id = self.club_service.save_profile(
            {
                "name": "Escola Delta",
                "kind": "school",
                "active": 1,
            },
            club_id=None,
        )
        class_id = self.club_service.save_class(
            {
                "club_id": school_id,
                "name": "Turma Delta",
                "active": 1,
            }
        )

        self.tournament_service.save_profile(
            self.tournament_id,
            {
                "name": "Torneio avulso",
                "location": "",
                "rounds_count": "3",
                "time_control": "",
                "start_date": "",
                "end_date": "",
                "bye_points": "1",
                "scope": "standalone",
                "club_id": None,
                "class_id": None,
            },
            {
                "initial_order": "rating",
                "tournament_type": "real",
            },
            [],
        )
        standalone_tournament = self.db.get_tournament(self.tournament_id)

        self.assertIsNone(standalone_tournament["club_id"])
        self.assertIsNone(standalone_tournament["class_id"])

        self.tournament_service.save_profile(
            self.tournament_id,
            {
                "name": "Torneio da turma",
                "location": "",
                "rounds_count": "3",
                "time_control": "",
                "start_date": "",
                "end_date": "",
                "bye_points": "1",
                "scope": "class",
                "club_id": school_id,
                "class_id": class_id,
            },
            {
                "initial_order": "rating",
                "tournament_type": "real",
            },
            [],
        )
        class_tournament = self.db.get_tournament(self.tournament_id)

        self.assertEqual(class_tournament["club_id"], school_id)
        self.assertEqual(class_tournament["class_id"], class_id)
        self.assertEqual(class_tournament["class_name"], "Turma Delta")

    def test_player_official_fields_and_csv_import(self) -> None:
        player_id = self.db.create_player(
            self.tournament_id,
            name="Francisco Ximenes",
            surname="Ximenes",
            given_name="Francisco",
            title="NM",
            sex="M",
            club="Clube",
            fide_id="1896",
            cbx_id="1928",
            rating=1928,
            national_rating=1928,
            international_rating=1896,
            category="ABS",
            birth_date="2003-01-01",
        )

        self.db.update_player(
            player_id,
            name="Francisco E Ximenes",
            club="Clube",
            rating=1950,
            category="ABS",
            active=1,
            fide_id="1896",
            cbx_id="1928",
            surname="Ximenes",
            given_name="Francisco E",
            title="NM",
            sex="M",
            birth_date="2003-01-01",
            national_rating=1950,
            international_rating=1900,
        )

        player = self.db.get_player(player_id)
        self.assertEqual(player["cbx_id"], "1928")
        self.assertEqual(player["national_rating"], 1950)
        self.assertEqual(player["international_rating"], 1900)

        csv_path = Path(self.temp_dir.name) / "players.csv"
        csv_path.write_text(
            "nome,sobrenome,titulo,fide,cbx,rating_nacional,rating_internacional,categoria\n"
            "Ana Silva,Silva,WFM,222,333,1800,1750,Sub-18\n",
            encoding="utf-8",
        )
        result = self.import_service.import_players_csv(self.tournament_id, csv_path)
        imported = next(
            player
            for player in self.db.list_players(self.tournament_id)
            if player["name"] == "Ana Silva"
        )

        self.assertEqual(result["imported"], 1)
        self.assertEqual(imported["title"], "WFM")
        self.assertEqual(imported["fide_id"], "222")
        self.assertEqual(imported["cbx_id"], "333")
        self.assertEqual(imported["rating"], 1800)

        semicolon_csv_path = Path(self.temp_dir.name) / "players_semicolon.csv"
        semicolon_csv_path.write_text(
            "nome;clube;elo;categoria;id_fide;id_cbx\n"
            "Bruna Costa;Clube B;1675;ABS;444;555\n",
            encoding="utf-8",
        )
        semicolon_result = self.import_service.import_players_csv(self.tournament_id, semicolon_csv_path)
        semicolon_imported = next(
            player
            for player in self.db.list_players(self.tournament_id)
            if player["name"] == "Bruna Costa"
        )

        self.assertEqual(semicolon_result["imported"], 1)
        self.assertEqual(semicolon_imported["rating"], 1675)
        self.assertEqual(semicolon_imported["fide_id"], "444")
        self.assertEqual(semicolon_imported["cbx_id"], "555")

        from openpyxl import Workbook

        xlsx_path = Path(self.temp_dir.name) / "players.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["nome", "clube", "elo", "categoria", "id_fide", "id_cbx"])
        sheet.append(["Carlos Lima", "Clube C", 1610, "ABS", "666", "777"])
        workbook.save(xlsx_path)

        xlsx_result = self.import_service.import_players(self.tournament_id, xlsx_path)
        xlsx_imported = next(
            player
            for player in self.db.list_players(self.tournament_id)
            if player["name"] == "Carlos Lima"
        )

        self.assertEqual(xlsx_result["imported"], 1)
        self.assertEqual(xlsx_imported["rating"], 1610)
        self.assertEqual(xlsx_imported["fide_id"], "666")
        self.assertEqual(xlsx_imported["cbx_id"], "777")

    def test_online_registration_import_previews_duplicates_and_imports_ready_rows(self) -> None:
        self.tournament_service.save_profile(
            self.tournament_id,
            {
                "name": "Torneio teste",
                "scope": "standalone",
                "location": "Sao Paulo",
                "rounds_count": "5",
                "time_control": "",
                "start_date": "2024-05-01",
                "end_date": "",
                "bye_points": "1",
            },
            {
                "initial_order": "rating",
                "tournament_type": "real",
            },
            [],
        )
        self.db.create_player(
            self.tournament_id,
            name="Duplicado Existente",
            fide_id="999",
            rating=1700,
        )
        csv_path = Path(self.temp_dir.name) / "inscricoes.csv"
        csv_path.write_text(
            "Carimbo de data/hora;Nome completo;Data de nascimento;Sexo;Clube / Cidade;"
            "Rating nacional;Rating internacional;FIDE ID;CBX ID;Categoria\n"
            "2026-05-14 10:00;Ana Silva;2008-01-01;Feminino;Sao Paulo;1390;;111;222;\n"
            "2026-05-14 10:01;Duplicado Existente;1990-01-01;M;Rio;1700;;999;;ABS\n"
            "2026-05-14 10:02;Ana Silva;2008-01-01;Feminino;Sao Paulo;1390;;111;222;\n"
            "2026-05-14 10:03;;2009-01-01;M;Sao Paulo;1200;;;;\n",
            encoding="utf-8",
        )

        preview = self.import_service.preview_online_registrations_csv(self.tournament_id, csv_path)
        result = self.import_service.import_online_registrations_csv(self.tournament_id, csv_path)
        imported = next(
            player
            for player in self.db.list_players(self.tournament_id, active_only=False)
            if player["name"] == "Ana Silva"
        )

        self.assertEqual(preview["total"], 4)
        self.assertEqual(preview["ready"], 1)
        self.assertEqual(preview["duplicate"], 2)
        self.assertEqual(preview["error"], 1)
        self.assertEqual(result["imported"], 1)
        self.assertEqual(result["skipped"], 3)
        self.assertEqual(imported["category"], "Sub-16")
        self.assertEqual(imported["age_category"], "Sub-16")
        self.assertEqual(imported["rating_category"], "Sub-1400")
        self.assertEqual(imported["prize_tags"], "Feminino; Melhor Local")

    def test_online_registration_import_accepts_google_sheets_link(self) -> None:
        csv_content = (
            "Nome completo;Data de nascimento;Sexo;Clube / Cidade;Rating nacional;FIDE ID;Categoria\n"
            "Carla Forms;2007-02-03;Feminino;Curitiba;1510;333;Sub-18\n"
        ).encode("utf-8")

        class FakeResponse:
            def __enter__(self) -> "FakeResponse":
                return self

            def __exit__(self, *_args: object) -> None:
                return None

            @staticmethod
            def read() -> bytes:
                return csv_content

        requested_urls: list[str] = []

        def fake_urlopen(request: object, timeout: int = 0) -> FakeResponse:
            requested_urls.append(request.full_url)
            self.assertEqual(timeout, 20)
            return FakeResponse()

        source_url = "https://docs.google.com/spreadsheets/d/abc123/edit#gid=987"
        with mock.patch("src.services.import_service.urlopen", side_effect=fake_urlopen):
            preview = self.import_service.preview_online_registrations(self.tournament_id, source_url)
            result = self.import_service.import_online_registrations(self.tournament_id, source_url)

        imported = next(
            player
            for player in self.db.list_players(self.tournament_id, active_only=False)
            if player["name"] == "Carla Forms"
        )

        self.assertEqual(
            requested_urls[0],
            "https://docs.google.com/spreadsheets/d/abc123/export?format=csv&gid=987",
        )
        self.assertEqual(preview["ready"], 1)
        self.assertEqual(result["imported"], 1)
        self.assertEqual(imported["club"], "Curitiba")
        self.assertEqual(imported["fide_id"], "333")
        self.assertEqual(imported["category"], "Sub-20")

    def test_official_rating_import_updates_tournament_players(self) -> None:
        player_id = self.db.create_player(
            self.tournament_id,
            name="Ana Silva",
            rating=1500,
            fide_id="222",
            cbx_id="333",
            club="Clube antigo",
        )
        fide_csv = Path(self.temp_dir.name) / "fide.csv"
        fide_csv.write_text(
            "name,fide,title,fide_rating,federation,birth_date\n"
            "Ana Silva,222,WFM,1810,BRA,2008-01-01\n",
            encoding="utf-8",
        )
        cbx_csv = Path(self.temp_dir.name) / "cbx.csv"
        cbx_csv.write_text(
            "nome,cbx,rating_nacional,clube\n"
            "Ana Silva,333,1850,Clube novo\n",
            encoding="utf-8",
        )

        fide_result = self.official_rating_service.import_official_csv(fide_csv, "FIDE", "2026-05")
        cbx_result = self.official_rating_service.import_official_csv(cbx_csv, "CBX", "2026-05")
        update_result = self.official_rating_service.update_tournament_players(self.tournament_id)
        player = self.db.get_player(player_id)

        self.assertEqual(fide_result["imported"], 1)
        self.assertEqual(cbx_result["imported"], 1)
        self.assertEqual(update_result["updated"], 1)
        self.assertEqual(player["title"], "WFM")
        self.assertEqual(player["international_rating"], 1810)
        self.assertEqual(player["national_rating"], 1850)
        self.assertEqual(player["rating"], 1850)
        self.assertEqual(player["club"], "Clube novo")

    def test_official_rating_preview_does_not_persist_before_confirmation(self) -> None:
        matched_player_id = self.db.create_player(
            self.tournament_id,
            name="Nome antigo",
            rating=1500,
            fide_id="222",
            club="Clube antigo",
        )
        unmatched_player_id = self.db.create_player(
            self.tournament_id,
            name="Sem cadastro oficial",
            rating=1400,
            fide_id="999",
        )
        fide_csv = Path(self.temp_dir.name) / "fide_preview.csv"
        fide_csv.write_text(
            "name,fide,title,fide_rating,club\n"
            "Nome oficial,222,FM,1810,Clube novo\n",
            encoding="utf-8",
        )
        self.official_rating_service.import_official_csv(fide_csv, "FIDE", "2026-05")

        preview = self.official_rating_service.preview_tournament_player_updates(self.tournament_id)
        matched_before_confirmation = self.db.get_player(matched_player_id)
        unmatched = self.db.get_player(unmatched_player_id)

        self.assertEqual(preview["total"], 2)
        self.assertEqual(preview["matched"], 1)
        self.assertEqual(preview["changed"], 1)
        self.assertEqual(preview["unchanged"], 0)
        self.assertEqual(preview["unmatched_count"], 1)
        self.assertEqual(preview["rows"][0]["status"], "changed")
        self.assertEqual(
            preview["rows"][0]["changed_fields"],
            ["name", "club", "title", "rating", "international_rating"],
        )
        self.assertEqual(preview["unmatched"][0]["name"], unmatched["name"])
        self.assertEqual(matched_before_confirmation["name"], "Nome antigo")
        self.assertEqual(matched_before_confirmation["club"], "Clube antigo")
        self.assertEqual(matched_before_confirmation["rating"], 1500)

        result = self.official_rating_service.apply_tournament_player_updates(
            self.tournament_id,
            [matched_player_id],
        )
        matched_after_confirmation = self.db.get_player(matched_player_id)

        self.assertEqual(result["updated"], 1)
        self.assertEqual(result["unmatched"], ["Sem cadastro oficial"])
        self.assertEqual(matched_after_confirmation["name"], "Nome oficial")
        self.assertEqual(matched_after_confirmation["club"], "Clube novo")
        self.assertEqual(matched_after_confirmation["title"], "FM")
        self.assertEqual(matched_after_confirmation["rating"], 1810)

    def test_official_rating_preview_confirmation_applies_only_selected_players(self) -> None:
        first_player_id = self.db.create_player(
            self.tournament_id,
            name="Primeiro antigo",
            rating=1500,
            fide_id="101",
        )
        second_player_id = self.db.create_player(
            self.tournament_id,
            name="Segundo antigo",
            rating=1400,
            fide_id="202",
        )
        fide_csv = Path(self.temp_dir.name) / "fide_selected.csv"
        fide_csv.write_text(
            "name,fide,fide_rating\n"
            "Primeiro oficial,101,1800\n"
            "Segundo oficial,202,1700\n",
            encoding="utf-8",
        )
        self.official_rating_service.import_official_csv(fide_csv, "FIDE", "2026-05")

        preview = self.official_rating_service.preview_tournament_player_updates(self.tournament_id)
        result = self.official_rating_service.apply_tournament_player_updates(
            self.tournament_id,
            [first_player_id],
        )

        self.assertEqual(preview["changed"], 2)
        self.assertEqual(result["updated"], 1)
        self.assertEqual(result["skipped"], 1)
        self.assertEqual(self.db.get_player(first_player_id)["name"], "Primeiro oficial")
        self.assertEqual(self.db.get_player(second_player_id)["name"], "Segundo antigo")

    def test_official_rating_import_accepts_semicolon_csv_and_normalized_headers(self) -> None:
        player_id = self.db.create_player(
            self.tournament_id,
            name="Carla Lima",
            rating=1300,
            fide_id="777",
        )
        official_csv = Path(self.temp_dir.name) / "fide_semicolon.csv"
        official_csv.write_text(
            "Nome;FIDE ID;Titulo;Rating internacional;Federacao;Data nascimento\n"
            "Carla Lima;777;WCM;1888;BRA;2009-02-03\n",
            encoding="utf-8",
        )

        result = self.official_rating_service.import_official_csv(official_csv, "FIDE", "2026-05")
        update_result = self.official_rating_service.update_tournament_players(self.tournament_id)
        player = self.db.get_player(player_id)

        self.assertEqual(result["imported"], 1)
        self.assertEqual(result["errors"], [])
        self.assertIsNotNone(result["snapshot_id"])
        self.assertEqual(update_result["updated"], 1)
        self.assertEqual(player["title"], "WCM")
        self.assertEqual(player["international_rating"], 1888)
        self.assertEqual(player["rating"], 1888)

    def test_lbx_online_import_combines_lists_and_updates_player_by_lbx_id(self) -> None:
        player_id = self.db.create_player(
            self.tournament_id,
            name="Nome antigo",
            rating=1200,
            lbx_id="350004",
        )
        header = "ID_No;Name;Fed;Sex;Clubnumber;ClubName;Birthday;Rtg_Nat;Fide_No;Rtg_Int;Title;Type;Status;K;\n"
        ratings = {
            "LBXstandard.php": 2348,
            "LBXrapid.php": 2438,
            "LBXblitz.php": 2436,
        }

        class FakeResponse:
            def __init__(self, content: str) -> None:
                self.content = content.encode("utf-8")

            def __enter__(self) -> FakeResponse:
                return self

            def __exit__(self, *_args: object) -> None:
                return None

            def read(self) -> bytes:
                return self.content

        def fake_urlopen(request: object, timeout: int = 0) -> FakeResponse:
            url = str(getattr(request, "full_url", ""))
            rating = next(value for suffix, value in ratings.items() if url.endswith(suffix))
            row = f"350004;Reis, Paulo F Jatoba de Oliveira;BRA;;29;Salvador (BA);1973-01-01;{rating};350004;{rating};FM;;;30;\n"
            return FakeResponse(header + row)

        with mock.patch("src.services.rating_service.urllib.request.urlopen", side_effect=fake_urlopen):
            result = self.official_rating_service.import_lbx_lists_from_url()
        update_result = self.official_rating_service.update_tournament_players(self.tournament_id)
        player = self.db.get_player(player_id)
        official = self.db.find_latest_official_player(lbx_id="350004")

        self.assertEqual(result["imported"], 1)
        self.assertEqual(result["errors"], [])
        self.assertEqual(update_result["updated"], 1)
        self.assertEqual(official["standard_rating"], 2348)
        self.assertEqual(official["rapid_rating"], 2438)
        self.assertEqual(official["blitz_rating"], 2436)
        self.assertEqual(official["fide_id"], "")
        self.assertEqual(player["lbx_id"], "350004")
        self.assertEqual(player["fide_id"], "")
        self.assertEqual(player["national_rating"], 2348)
        self.assertEqual(player["rating"], 2348)
        self.assertEqual(player["club"], "Salvador (BA)")

    def test_official_rating_import_without_header_does_not_create_snapshot(self) -> None:
        bad_csv = Path(self.temp_dir.name) / "official_empty.csv"
        bad_csv.write_text("", encoding="utf-8")

        with self.assertRaises(AppError):
            self.official_rating_service.import_official_csv(bad_csv, "FIDE")
        with self.db.connect() as connection:
            snapshots_count = connection.execute(
                "SELECT COUNT(*) FROM official_rating_snapshots"
            ).fetchone()[0]

        self.assertEqual(snapshots_count, 0)

    def test_official_rating_import_rolls_back_snapshot_when_player_insert_fails(self) -> None:
        official_csv = Path(self.temp_dir.name) / "official_rollback.csv"
        official_csv.write_text(
            "name,fide,rating\n"
            "Ana Silva,222,1810\n"
            "Bruno Souza,333,1720\n",
            encoding="utf-8",
        )
        original_insert = self.db._insert_official_player
        call_count = 0

        def flaky_insert(*args: object, **kwargs: object) -> int:
            nonlocal call_count
            call_count += 1
            if call_count == 2:
                raise sqlite3.OperationalError("falha simulada")
            return original_insert(*args, **kwargs)

        with mock.patch.object(self.db, "_insert_official_player", side_effect=flaky_insert):
            with self.assertRaises(sqlite3.OperationalError):
                self.official_rating_service.import_official_csv(official_csv, "FIDE")
        with self.db.connect() as connection:
            snapshots_count = connection.execute(
                "SELECT COUNT(*) FROM official_rating_snapshots"
            ).fetchone()[0]
            players_count = connection.execute("SELECT COUNT(*) FROM official_players").fetchone()[0]

        self.assertEqual(snapshots_count, 0)
        self.assertEqual(players_count, 0)

    def test_official_rating_update_respects_initial_order(self) -> None:
        self.tournament_service.save_profile(
            self.tournament_id,
            {
                "name": "Torneio teste",
                "location": "",
                "rounds_count": "5",
                "time_control": "",
                "start_date": "",
                "end_date": "",
                "bye_points": "1",
            },
            {
                "initial_order": "international_then_national",
                "tournament_type": "real",
            },
            [],
        )
        player_id = self.db.create_player(
            self.tournament_id,
            name="Bruno Souza",
            rating=1200,
            fide_id="444",
            cbx_id="555",
        )
        official_csv = Path(self.temp_dir.name) / "official.csv"
        official_csv.write_text(
            "name,fide,cbx,rating_nacional,rating_internacional\n"
            "Bruno Souza,444,555,1900,1750\n",
            encoding="utf-8",
        )

        self.official_rating_service.import_official_csv(official_csv, "FIDE")
        self.official_rating_service.update_tournament_players(self.tournament_id)
        player = self.db.get_player(player_id)

        self.assertEqual(player["rating"], 1750)

    def test_export_site_creates_static_html_package(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Jogador A",
            rating=1800,
            club="Clube",
            category="ABS",
        )
        self.db.create_player(
            self.tournament_id,
            name="Jogador B",
            rating=1700,
            club="Clube",
            category="ABS",
        )
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(round_data["id"])[0]
        self.db.update_pairing_result(pairing["id"], "1-0")
        self.service.close_round(self.tournament_id, round_data["id"])

        output_dir = Path(self.temp_dir.name) / "site"
        index_path = self.export_service.export_site(self.tournament_id, output_dir)
        html = index_path.read_text(encoding="utf-8")

        self.assertTrue(index_path.exists())
        self.assertTrue((output_dir / "styles.css").exists())
        self.assertIn("Torneio teste", html)
        self.assertIn("Classificacao", html)
        self.assertIn("Rodada 1", html)

    def test_phase5_public_portal_payload_json_and_html_hide_sensitive_data(self) -> None:
        first_player = self.db.create_player(
            self.tournament_id,
            name="Jogador Publico",
            rating=1800,
            club="Clube",
            category="ABS",
            birth_date="2010-01-02",
            fide_id="123456",
        )
        self.db.create_player(
            self.tournament_id,
            name="Jogador Visitante",
            rating=1700,
            club="Clube",
            category="ABS",
        )
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(round_data["id"])[0]
        self.db.update_pairing_result(pairing["id"], "1-0")
        self.service.close_round(self.tournament_id, round_data["id"])
        self.db.save_app_settings({"live_portal_notice": "Resultados sujeitos a homologacao."})
        self.db.save_tournament_settings(
            self.tournament_id,
            {"contact_email": "arbitro@example.com", "comments": "Comentario interno"},
        )

        payload = self.export_service.public_tournament_payload(self.tournament_id, mode="publico")
        serialized = json.dumps(payload, ensure_ascii=False)

        self.assertEqual("publico", payload["mode"])
        self.assertEqual("Resultados sujeitos a homologacao.", payload["notice"])
        self.assertEqual("Jogador Publico", payload["players"][0]["name"])
        self.assertNotIn("arbitro@example.com", serialized)
        self.assertNotIn("Comentario interno", serialized)
        self.assertNotIn("2010-01-02", serialized)
        self.assertNotIn("123456", serialized)
        self.assertEqual(first_player, payload["players"][0]["id"])

        output_path = Path(self.temp_dir.name) / "publico.json"
        self.export_service.export_public_json(self.tournament_id, output_path)
        exported = output_path.read_text(encoding="utf-8")
        self.assertIn("Jogador Publico", exported)
        self.assertNotIn("arbitro@example.com", exported)

        html = self.export_service.live_portal_html(self.tournament_id, mode="publico")
        self.assertIn("Albericus Live", html)
        self.assertIn("Rodada atual", html)
        self.assertIn("Fichas publicas", html)
        self.assertNotIn("arbitro@example.com", html)

    def test_export_chess_results_trf16_includes_required_fields(self) -> None:
        tournament_id = self.db.create_tournament(
            "Aberto Sao Paulo",
            location="Sao Paulo",
            rounds_count=1,
            time_control="90 min + 30 sec",
            start_date="2026-05-24",
            end_date="2026-05-24",
        )
        self.db.save_tournament_settings(
            tournament_id,
            {
                "fide_event_id": "12345",
                "federation": "BRA",
                "chief_arbiter": "Arbitro Chefe",
                "arbiters": "Adjunto Um",
                "tournament_profile": "fide",
            },
        )
        self.db.save_round_schedule(
            tournament_id,
            [{"round_number": 1, "date": "2026-05-24", "time": "10:00"}],
        )
        first_player = self.db.create_player(
            tournament_id,
            name="Silva, Ana",
            surname="Silva",
            given_name="Ana",
            title="WFM",
            sex="w",
            federation_id="BRA",
            fide_id="1234567",
            rating=2100,
            international_rating=2100,
            birth_date="2000-01-02",
        )
        second_player = self.db.create_player(
            tournament_id,
            name="Souza, Bruno",
            surname="Souza",
            given_name="Bruno",
            title="FM",
            sex="m",
            federation_id="BRA",
            fide_id="7654321",
            rating=2000,
            international_rating=2000,
            birth_date="1999-03-04",
        )
        round_id = self.db.create_round_with_pairings(
            tournament_id,
            1,
            [
                {
                    "board_number": 1,
                    "white_player_id": first_player,
                    "black_player_id": second_player,
                    "result": "1-0",
                }
            ],
        )
        self.service.close_round(tournament_id, round_id)

        output_path = Path(self.temp_dir.name) / "chess_results.trf"
        warnings = self.export_service.export_chess_results_trf(tournament_id, output_path)
        content = output_path.read_text(encoding="utf-8")
        lines = content.splitlines()
        player_lines = [line for line in lines if line.startswith("001 ")]

        self.assertEqual(warnings, [])
        self.assertIn("012 Aberto Sao Paulo", content)
        self.assertIn("032 BRA", content)
        self.assertIn("042 2026/05/24", content)
        self.assertIn("052 2026/05/24", content)
        self.assertIn("062 2", content)
        self.assertIn("072 2", content)
        self.assertIn("082 0", content)
        self.assertIn("092 Individual: Suico (FIDE-rated)", content)
        self.assertIn("102 Arbitro Chefe", content)
        self.assertIn("112 Adjunto Um", content)
        self.assertIn("122 90 min + 30 sec", content)
        self.assertEqual(len(player_lines), 2)
        self.assertIn("wWFM Silva, Ana", player_lines[0])
        self.assertIn("2100 BRA", player_lines[0])
        self.assertIn("1234567", player_lines[0])
        self.assertIn("2000/01/02", player_lines[0])
        self.assertTrue(player_lines[0].endswith("2 w 1"))
        self.assertTrue(player_lines[1].endswith("1 b 0"))

    def test_validate_chess_results_trf16_warns_about_open_rounds_and_invalid_fide_data(self) -> None:
        tournament_id = self.db.create_tournament(
            "Aberto Validacao",
            location="Curitiba",
            rounds_count=2,
            time_control="60 min",
            start_date="2026-05-24",
            end_date="2026-05-25",
        )
        self.db.save_tournament_settings(
            tournament_id,
            {
                "federation": "BR",
                "chief_arbiter": "Arbitro Chefe",
                "tournament_profile": "fide",
            },
        )
        self.db.save_round_schedule(
            tournament_id,
            [
                {"round_number": 1, "date": "2026-05-24", "time": "10:00"},
                {"round_number": 2, "date": "data ruim", "time": "15:00"},
            ],
        )
        first_player = self.db.create_player(
            tournament_id,
            name="Jogador Um",
            fide_id="ABC123",
            federation_id="BR",
            rating=1800,
            international_rating=1800,
            birth_date="data ruim",
        )
        second_player = self.db.create_player(
            tournament_id,
            name="Jogador Dois",
            fide_id="222",
            federation_id="BRA",
            rating=1700,
            international_rating=1700,
            birth_date="2000-01-02",
        )
        self.db.create_round_with_pairings(
            tournament_id,
            1,
            [
                {
                    "board_number": 1,
                    "white_player_id": first_player,
                    "black_player_id": second_player,
                    "result": "",
                }
            ],
        )

        warnings = self.export_service.validate_chess_results_trf(tournament_id)

        self.assertIn("Federacao FIDE do torneio deve ter 3 letras.", warnings)
        self.assertIn("Torneio FIDE-rated sem FIDE Event-ID.", warnings)
        self.assertIn("Rodadas ainda nao fechadas: 1.", warnings)
        self.assertIn("Rodadas com resultados pendentes/incompletos: 1.", warnings)
        self.assertIn("Rodadas com data invalida no calendario: 2.", warnings)
        self.assertTrue(any("FIDE ID numerico" in warning for warning in warnings))
        self.assertTrue(any("data de nascimento valida" in warning for warning in warnings))
        self.assertTrue(any("federacao FIDE com 3 letras" in warning for warning in warnings))

    def test_export_chess_results_trf16_supports_team_tournaments(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        boards = self.db.list_team_boards(int(match["id"]))
        self.service.update_result(tournament_id, int(boards[0]["id"]), "1-0")
        self.service.update_result(tournament_id, int(boards[1]["id"]), "1/2-1/2")
        self.service.close_round(tournament_id, round_data["id"])

        output_path = Path(self.temp_dir.name) / "team_chess_results.trf"
        warnings = self.export_service.export_chess_results_trf(tournament_id, output_path)
        content = output_path.read_text(encoding="utf-8")
        lines = content.splitlines()
        player_lines = [line for line in lines if line.startswith("001 ")]
        team_lines = [line for line in lines if line.startswith("013 ")]

        self.assertIn("082 2", content)
        self.assertIn("092 Team: Suico (Standard)", content)
        self.assertEqual(len(player_lines), 4)
        self.assertEqual(len(team_lines), 2)
        self.assertTrue(any("Equipe 1" in line and "1" in line and "2" in line for line in team_lines))
        self.assertTrue(any("   3 w 1" in line for line in player_lines))
        self.assertTrue(any("   4 b =" in line for line in player_lines))
        self.assertTrue(any("Jogadores sem FIDE ID" in warning for warning in warnings))

    def test_trf25_export_emits_310_and_header_extensions_for_teams(self) -> None:
        from src.services.federation_exporters import (
            TRF25_SCAFFOLD_WARNING,
            TRF25Exporter,
        )

        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        boards = self.db.list_team_boards(int(match["id"]))
        self.service.update_result(tournament_id, int(boards[0]["id"]), "1-0")
        self.service.update_result(tournament_id, int(boards[1]["id"]), "1/2-1/2")
        self.service.close_round(tournament_id, round_data["id"])

        output_path = Path(self.temp_dir.name) / "team_trf25.trf"
        warnings = TRF25Exporter(self.export_service).export(tournament_id, output_path)
        content = output_path.read_text(encoding="utf-8")
        lines = content.splitlines()
        player_lines = [line for line in lines if line.startswith("001 ")]
        team_310 = [line for line in lines if line.startswith("310 ")]
        team_802 = [line for line in lines if line.startswith("802 ")]

        # Cabeçalho TRF25: nº de rodadas (142), tipo codificado (192) e
        # sequência de cores dos tabuleiros (352) para torneio por equipes.
        self.assertTrue(any(line.startswith("142 ") for line in lines))
        self.assertIn("192 FIDE_TEAM_TYPEA_MP_GP", content)
        self.assertIn("352 WB", content)
        # Tie-breaks de classificação (212), sempre começando por PTS.
        self.assertIn("212 PTS,BH:MP,WIN", content)
        # Pontuação padrão (TW=2/TD=1/TL=0) → 362 omitido.
        self.assertFalse(any(line.startswith("362 ") for line in lines))
        # Equipes saem como 310 (substitui o 013); o 013 não é mais emitido.
        self.assertEqual(len(team_310), 2)
        self.assertFalse(any(line.startswith("013 ") for line in lines))
        # Linhas 001 dos jogadores continuam idênticas ao TRF16.
        self.assertEqual(len(player_lines), 4)
        # Registro informativo 802: um por equipe, espelhando o TPN do 310.
        self.assertEqual(len(team_802), 2)
        # Warning de scaffold continua presente.
        self.assertIn(TRF25_SCAFFOLD_WARNING, warnings)

    def test_trf25_export_emits_330_for_double_forfeit_match(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        # Match inteiro W.O.: todos os tabuleiros 0F-0F → duplo forfeit (--).
        for board in self.db.list_team_boards(int(match["id"])):
            self.service.update_result(tournament_id, int(board["id"]), "0F-0F")
        self.service.close_round(tournament_id, round_data["id"])

        output_path = Path(self.temp_dir.name) / "team_trf25_ff.trf"
        TRF25Exporter(self.export_service).export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()
        forfeit_lines = [line for line in lines if line.startswith("330 ")]
        self.assertEqual(len(forfeit_lines), 1)
        self.assertEqual(forfeit_lines[0][4:6], "--")

    def test_trf25_export_emits_330_directional_walkover(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        # Equipe branca do match vence por W.O. em ambos os tabuleiros.
        # Tabuleiro 1 (ímpar): brancas = equipe branca → 1F-0F.
        # Tabuleiro 2 (par, cores invertidas): pretas = equipe branca → 0F-1F.
        boards = self.db.list_team_boards(int(match["id"]))
        self.service.update_result(tournament_id, int(boards[0]["id"]), "1F-0F")
        self.service.update_result(tournament_id, int(boards[1]["id"]), "0F-1F")
        self.service.close_round(tournament_id, round_data["id"])

        output_path = Path(self.temp_dir.name) / "team_trf25_wo.trf"
        TRF25Exporter(self.export_service).export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()
        forfeit_lines = [line for line in lines if line.startswith("330 ")]
        self.assertEqual(len(forfeit_lines), 1)
        self.assertEqual(forfeit_lines[0][4:6], "+-")

    def test_trf25_export_skips_330_for_normal_and_partial_results(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        # Forfeit parcial: 1 tabuleiro W.O., outro jogado → NÃO é um 330.
        boards = self.db.list_team_boards(int(match["id"]))
        self.service.update_result(tournament_id, int(boards[0]["id"]), "1F-0F")
        self.service.update_result(tournament_id, int(boards[1]["id"]), "1-0")
        self.service.close_round(tournament_id, round_data["id"])

        output_path = Path(self.temp_dir.name) / "team_trf25_partial.trf"
        TRF25Exporter(self.export_service).export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()
        self.assertFalse(any(line.startswith("330 ") for line in lines))

    def test_trf25_export_emits_320_pab_for_team_bye(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        # 3 equipes → uma recebe bye (pairing-allocated-bye) a cada rodada.
        tournament_id, _team_ids = self._create_team_tournament(teams_count=3, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        matches = self.db.list_team_matches_for_round(round_data["id"])
        played = [m for m in matches if not m.get("is_bye")]
        bye = [m for m in matches if m.get("is_bye")]
        self.assertEqual(len(bye), 1)
        for board in self.db.list_team_boards(int(played[0]["id"])):
            self.service.update_result(tournament_id, int(board["id"]), "1-0")
        self.service.close_round(tournament_id, round_data["id"])

        output_path = Path(self.temp_dir.name) / "team_trf25_bye.trf"
        TRF25Exporter(self.export_service).export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()
        pab_lines = [line for line in lines if line.startswith("320 ")]
        # Um único registro 320 por torneio, com MP/GP do bye e o TPN na rodada 1.
        self.assertEqual(len(pab_lines), 1)
        self.assertEqual(pab_lines[0][4:8], " 2.0")
        self.assertEqual(pab_lines[0][9:13], " 2.0")

    def test_trf25_scoring_362_and_individual_212_codes(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        # Pontuação de match não-padrão (3-1-0) → emite o 362; padrão → None.
        self.assertIsNone(
            TRF25Exporter._scoring_system_362(
                {"team_match_win_points": "2", "team_match_draw_points": "1", "team_match_loss_points": "0"}
            )
        )
        line_362 = TRF25Exporter._scoring_system_362(
            {"team_match_win_points": "3", "team_match_draw_points": "1", "team_match_loss_points": "0"}
        )
        self.assertEqual(line_362[:6], "362 TW")
        self.assertIn(" 3.0", line_362)
        # Ordem de desempate individual espelha pairing/tiebreaks.py.
        self.assertEqual(
            TRF25Exporter._tiebreak_codes_212(is_team=False),
            ["PTS", "BH", "BH/M1", "SB", "WIN"],
        )

    def test_federation_exporter_registry_keeps_trf16_flow_extensible(self) -> None:
        registry = FederationExporterRegistry()
        exporter = TRF16Exporter(self.export_service)
        registry.register(exporter)

        self.assertEqual(["trf16"], registry.list_formats())
        self.assertIs(registry.get("trf16"), exporter)

    def test_trf25_scaffold_registers_and_warns_about_pending_extensions(self) -> None:
        from src.services.federation_exporters import (
            TRF25_SCAFFOLD_WARNING,
            TRF25Exporter,
        )

        registry = FederationExporterRegistry()
        trf16 = TRF16Exporter(self.export_service)
        trf25 = TRF25Exporter(self.export_service)
        registry.register(trf16)
        registry.register(trf25)

        # Códigos distintos no registry, sem colidir.
        self.assertEqual(["trf16", "trf25"], registry.list_formats())
        self.assertIs(registry.get("trf25"), trf25)
        # Scaffold herda comportamento de validação, mas sempre prefixa
        # o warning explícito de "extensões não implementadas".
        self._create_players(2)
        warnings = trf25.validate(self.tournament_id)
        self.assertEqual(warnings[0], TRF25_SCAFFOLD_WARNING)
        # Os warnings subsequentes são os do TRF16 herdado (sem federacao,
        # sem datas, etc) — não devem ser duplicados.
        self.assertEqual(1, sum(1 for w in warnings if w == TRF25_SCAFFOLD_WARNING))

    def test_export_chess_results_trf25_routes_to_trf25_exporter(self) -> None:
        from src.services.federation_exporters import TRF25_SCAFFOLD_WARNING

        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        output_path = Path(self.temp_dir.name) / "fide.trf"

        warnings = self.export_service.export_chess_results_trf25(tournament_id, output_path)

        # O método novo deve devolver o aviso de scaffold do TRF25 e emitir
        # registros 310 (equipe), não o 013 herdado do TRF16.
        self.assertIn(TRF25_SCAFFOLD_WARNING, warnings)
        lines = output_path.read_text(encoding="utf-8").splitlines()
        self.assertTrue(any(line.startswith("310 ") for line in lines))
        self.assertFalse(any(line.startswith("013 ") for line in lines))

    def test_trf25_export_emits_152_and_222_for_individual(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        self.db.update_tournament_details(
            self.tournament_id, name="Torneio teste", time_control="90 min + 30 s"
        )
        self._create_players(4)
        self.service.generate_next_round(self.tournament_id)

        output_path = Path(self.temp_dir.name) / "ind_trf25.trf"
        TRF25Exporter(self.export_service).export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        # 222: ritmo codificado (90 min = 5400 s, +30 s de incremento).
        self.assertIn("222 5400+30", lines)
        # 152: cor do top seed (rank 1) na rodada 1 — W ou B, exatamente um.
        colour_lines = [line for line in lines if line.startswith("152 ")]
        self.assertEqual(len(colour_lines), 1)
        self.assertIn(colour_lines[0], ("152 W", "152 B"))

    def test_trf25_omits_222_when_time_control_unparseable(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        self.db.update_tournament_details(
            self.tournament_id, name="Torneio teste", time_control="ritmo livre"
        )
        self._create_players(2)

        output_path = Path(self.temp_dir.name) / "ind_no222.trf"
        TRF25Exporter(self.export_service).export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()
        self.assertFalse(any(line.startswith("222 ") for line in lines))

    def test_trf25_emits_299_for_individual_point_adjustment(self) -> None:
        from src.services.constants import player_pairing_name
        from src.services.federation_exporters import TRF25Exporter

        player_ids = self._create_players(4)
        # Penalidade de meio ponto ao 2º jogador por rating (vira o start-rank
        # exato no export, qualquer que seja a ordem de seeding).
        self.db.add_point_adjustment(
            self.tournament_id,
            round_number=0,
            player_id=player_ids[1],
            aat_type="",
            game_points=-0.5,
            reason="Penalidade de comportamento",
        )

        output_path = Path(self.temp_dir.name) / "ind_299.trf"
        TRF25Exporter(self.export_service).export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        players = sorted(
            self.db.list_players(self.tournament_id, active_only=False),
            key=lambda p: (
                -self.export_service._trf_rating(p),
                player_pairing_name(p).casefold(),
                int(p.get("id") or 0),
            ),
        )
        expected_rank = next(
            i for i, p in enumerate(players, start=1) if int(p["id"]) == player_ids[1]
        )

        adj_lines = [line for line in lines if line.startswith("299 ")]
        self.assertEqual(len(adj_lines), 1)
        line = adj_lines[0]
        self.assertEqual(line[13:17], "-0.5")
        self.assertEqual(line[23:27].strip(), str(expected_rank))

    def test_trf25_emits_299_for_team_point_adjustment(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        # Penalidade de 2 match points e 1 game point a uma equipe.
        self.db.add_point_adjustment(
            tournament_id,
            round_number=0,
            team_id=team_ids[0],
            aat_type="",
            match_points=-2.0,
            game_points=-1.0,
            reason="Penalidade de equipe",
        )

        output_path = Path(self.temp_dir.name) / "team_299.trf"
        TRF25Exporter(self.export_service).export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        # Mapeia nome da equipe (cols 9-40) -> TPN (cols 5-7) a partir dos 310.
        target_name = next(
            t["name"] for t in self.db.list_teams(tournament_id, active_only=False)
            if int(t["id"]) == team_ids[0]
        )
        tpn_by_name = {
            line[8:40].strip(): line[4:7].strip()
            for line in lines
            if line.startswith("310 ")
        }
        expected_tpn = tpn_by_name[target_name]

        adj_lines = [line for line in lines if line.startswith("299 ")]
        self.assertEqual(len(adj_lines), 1)
        line = adj_lines[0]
        self.assertEqual(line[7:11], "-2.0")
        self.assertEqual(line[13:17], "-1.0")
        self.assertEqual(line[23:27].strip(), expected_tpn)

    def test_trf25_emits_300_only_for_out_of_order_team(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        team_a, team_b = team_ids[0], team_ids[1]

        def roster(team_id: int) -> list[int]:
            assignments = sorted(
                self.db.list_team_players(team_id, active_only=False),
                key=lambda a: int(a.get("board_number") or 0),
            )
            return [int(a["player_id"]) for a in assignments]

        a_board1, a_board2 = roster(team_a)
        b_board1, b_board2 = roster(team_b)

        # Cores: tab.1 (impar) branco=A/preto=B; tab.2 (par) branco=B/preto=A.
        # Equipe A escala invertida (tab.1 com o jogador do tab.2 e vice-versa);
        # equipe B mantém a ordem do roster.
        round_id = self.db.create_round_with_team_matches(
            tournament_id,
            1,
            [
                {
                    "match_number": 1,
                    "white_team_id": team_a,
                    "black_team_id": team_b,
                    "is_bye": 0,
                    "boards": [
                        {
                            "board_number": 1,
                            "white_player_id": a_board2,
                            "black_player_id": b_board1,
                            "result": "1-0",
                        },
                        {
                            "board_number": 2,
                            "white_player_id": b_board2,
                            "black_player_id": a_board1,
                            "result": "0-1",
                        },
                    ],
                }
            ],
        )
        self.db.close_round(round_id)

        output_path = Path(self.temp_dir.name) / "team_300.trf"
        TRF25Exporter(self.export_service).export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        name_by_id = {
            int(t["id"]): t["name"]
            for t in self.db.list_teams(tournament_id, active_only=False)
        }
        tpn_by_name = {
            line[8:40].strip(): line[4:7].strip()
            for line in lines
            if line.startswith("310 ")
        }
        a_tpn = tpn_by_name[name_by_id[team_a]]
        start_rank = {
            int(p["id"]): index
            for index, p in enumerate(
                sorted(
                    self.db.list_players(tournament_id, active_only=False),
                    key=lambda p: -self.export_service._trf_rating(p),
                ),
                start=1,
            )
        }

        order_lines = [line for line in lines if line.startswith("300 ")]
        self.assertEqual(len(order_lines), 1)
        line = order_lines[0]
        self.assertEqual(line[4:7].strip(), a_tpn)
        # Tab.1 jogado pelo jogador do tab.2 do roster; tab.2 pelo do tab.1.
        self.assertEqual(line[16:20].strip(), str(start_rank[a_board2]))
        self.assertEqual(line[21:25].strip(), str(start_rank[a_board1]))

    def test_trf25_emits_250_for_classic_acceleration(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        self._create_players(8)
        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "accel.trf"

        # Sem aceleração: nenhum registro 250.
        self.db.save_tournament_settings(self.tournament_id, {"acceleration_method": "none"})
        exporter.export(self.tournament_id, output_path)
        plain = output_path.read_text(encoding="utf-8").splitlines()
        self.assertFalse([line for line in plain if line.startswith("250 ")])

        # Aceleração clássica: 250 com bônus 1.0, rodadas 1-2, ranks 1..4 (N//2).
        self.db.save_tournament_settings(
            self.tournament_id, {"acceleration_method": "accelerated"}
        )
        exporter.export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()
        accel_lines = [line for line in lines if line.startswith("250 ")]
        self.assertEqual(len(accel_lines), 1)
        line = accel_lines[0]
        self.assertEqual(line[4:8].strip(), "")  # match points em branco (individual)
        self.assertEqual(float(line[9:13]), 1.0)  # game points
        self.assertEqual(line[14:17].strip(), "1")  # primeira rodada
        self.assertEqual(line[18:21].strip(), "2")  # última rodada
        self.assertEqual(line[22:26].strip(), "1")  # primeiro jogador (rank 1)
        self.assertEqual(line[27:31].strip(), "4")  # último jogador (N//2 = 4)

    def test_prohibited_pairing_is_never_paired(self) -> None:
        self._create_players(8)
        players = self.db.list_players(self.tournament_id, active_only=True)
        seeding = [
            int(p["id"])
            for p in sorted(players, key=lambda p: -int(p.get("rating") or 0))
        ]

        def pair_set(pairings: list[dict[str, Any]]) -> set[frozenset[int]]:
            return {
                frozenset({int(p["white_player_id"]), int(p["black_player_id"])})
                for p in pairings
                if p.get("black_player_id") is not None
            }

        # Sem proibição: seed 1 pareia com seed 5 (topo vs base do mesmo grupo).
        plain = pair_set(self.service._swiss_pairings(self.tournament_id, players, 2))
        self.assertIn(frozenset({seeding[0], seeding[4]}), plain)

        # Proibindo seed 1 x seed 5, eles nunca podem ser pareados.
        self.db.add_prohibited_pairing(self.tournament_id, seeding[0], seeding[4])
        guarded = pair_set(self.service._swiss_pairings(self.tournament_id, players, 2))
        self.assertNotIn(frozenset({seeding[0], seeding[4]}), guarded)
        self.assertEqual(len(guarded), 4)  # 8 jogadores → 4 jogos íntegros

    def test_prohibition_respects_round_window(self) -> None:
        self._create_players(8)
        players = self.db.list_players(self.tournament_id, active_only=True)
        seeding = [
            int(p["id"])
            for p in sorted(players, key=lambda p: -int(p.get("rating") or 0))
        ]

        def pair_set(pairings: list[dict[str, Any]]) -> set[frozenset[int]]:
            return {
                frozenset({int(p["white_player_id"]), int(p["black_player_id"])})
                for p in pairings
                if p.get("black_player_id") is not None
            }

        # Proibição válida só na rodada 3.
        self.db.add_prohibited_pairing(
            self.tournament_id, seeding[0], seeding[4], first_round=3, last_round=3
        )
        # Rodada 2 fora da janela: o par ainda ocorre.
        self.assertIn(
            frozenset({seeding[0], seeding[4]}),
            pair_set(self.service._swiss_pairings(self.tournament_id, players, 2)),
        )
        # Rodada 3 dentro da janela: bloqueado.
        self.assertNotIn(
            frozenset({seeding[0], seeding[4]}),
            pair_set(self.service._swiss_pairings(self.tournament_id, players, 3)),
        )

    def test_prohibited_team_pairing_is_never_paired(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=4, boards_count=2)
        teams = self.db.list_teams(tournament_id, active_only=True)
        settings = self.db.get_tournament_settings(tournament_id) or {}
        boards_count = int(settings.get("team_boards_count") or 2)
        rosters, seed_ratings = self.service._team_starter_rosters(teams, boards_count)

        def match_set(matches: list[dict[str, Any]]) -> set[frozenset[int]]:
            return {
                frozenset((int(m["white_team_id"]), int(m["black_team_id"])))
                for m in matches
                if m.get("black_team_id") is not None and not m.get("is_bye")
            }

        # Rodada 2 sem resultados: grupo único, dobra por seed.
        plain = match_set(
            self.service._swiss_team_matches(
                tournament_id, teams, rosters, seed_ratings, boards_count, settings, 2
            )
        )
        self.assertEqual(len(plain), 2)
        sample = next(iter(plain))
        team_a, team_b = tuple(sample)

        # Proibindo o par, ele nunca pode ser pareado.
        self.db.add_prohibited_team_pairing(tournament_id, team_a, team_b)
        guarded = match_set(
            self.service._swiss_team_matches(
                tournament_id, teams, rosters, seed_ratings, boards_count, settings, 2
            )
        )
        self.assertIn(sample, plain)
        self.assertNotIn(sample, guarded)
        self.assertEqual(len(guarded), 2)  # 4 equipes → 2 confrontos íntegros

    def test_trf25_emits_260_for_prohibited_team_pairing(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, team_ids = self._create_team_tournament(teams_count=4, boards_count=2)
        # Proíbe a equipe seed 1 (TPN 1) x seed 3 (TPN 3), janela aberta.
        self.db.add_prohibited_team_pairing(tournament_id, team_ids[0], team_ids[2])
        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "team_prohib.trf"
        exporter.export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        prohibition_lines = [line for line in lines if line.startswith("260 ")]
        self.assertEqual(len(prohibition_lines), 1)
        line = prohibition_lines[0]
        self.assertEqual(line[4:7].strip(), "1")  # primeira rodada
        self.assertEqual(line[8:11].strip(), "5")  # última rodada (rounds_count=5)
        self.assertEqual(line[12:16].strip(), "1")  # TPN da seed 1
        self.assertEqual(line[17:21].strip(), "3")  # TPN da seed 3

    def test_requested_team_bye_excludes_team_from_pairing(self) -> None:
        # 5 equipes, 1 bye solicitado → 4 a parear (par), sem bye alocado extra.
        tournament_id, team_ids = self._create_team_tournament(teams_count=5, boards_count=2)
        self.db.add_requested_team_bye(tournament_id, team_ids[4], 1, "H")

        round_data = self.service.generate_next_round(tournament_id)
        matches = self.db.list_team_matches_for_round(int(round_data["id"]))

        byes = [m for m in matches if m["is_bye"]]
        self.assertEqual(len(byes), 1)
        self.assertEqual(int(byes[0]["white_team_id"]), team_ids[4])
        self.assertEqual(str(byes[0]["result"]).upper(), "H")

        normal = [m for m in matches if not m["is_bye"]]
        self.assertEqual(len(normal), 2)  # 4 equipes → 2 confrontos
        paired = {int(m["white_team_id"]) for m in normal}
        paired |= {int(m["black_team_id"]) for m in normal if m["black_team_id"]}
        self.assertNotIn(team_ids[4], paired)

    def test_requested_team_bye_scores_by_type(self) -> None:
        # H num torneio com win=2/draw=1/loss=0 e 2 tabuleiros:
        # match = draw (1.0), game = boards/2 (1.0) — distinto da vitória cheia.
        tournament_id, team_ids = self._create_team_tournament(teams_count=5, boards_count=2)
        self.db.add_requested_team_bye(tournament_id, team_ids[4], 1, "H")
        round_data = self.service.generate_next_round(tournament_id)

        for match in self.db.list_team_matches_for_round(int(round_data["id"])):
            if match["is_bye"]:
                continue
            for board in self.db.list_team_boards(int(match["id"])):
                self.service.update_result(tournament_id, int(board["id"]), "1-0")
        self.service.close_round(tournament_id, int(round_data["id"]))

        bye = next(
            m for m in self.db.list_team_matches_for_round(int(round_data["id"])) if m["is_bye"]
        )
        self.assertEqual(float(bye["white_match_points"]), 1.0)  # draw points
        self.assertEqual(float(bye["white_game_points"]), 1.0)  # boards/2

    def test_trf25_emits_240_for_requested_team_bye(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, team_ids = self._create_team_tournament(teams_count=5, boards_count=2)
        # Equipe seed 5 (TPN 5, rating mais baixo) com bye zero-point na rodada 1.
        self.db.add_requested_team_bye(tournament_id, team_ids[4], 1, "Z")
        self.service.generate_next_round(tournament_id)

        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "team_bye240.trf"
        exporter.export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        bye_lines = [line for line in lines if line.startswith("240 ")]
        self.assertEqual(len(bye_lines), 1)
        line = bye_lines[0]
        self.assertEqual(line[4:5], "Z")           # tipo (col 5)
        self.assertEqual(line[6:9].strip(), "1")   # rodada (col 7-9)
        self.assertEqual(line[10:14].strip(), "5")  # TPN da equipe (col 11-14)

    def _close_team_round_with_decisive_boards(self, tournament_id: int, round_id: int) -> None:
        for match in self.db.list_team_matches_for_round(int(round_id)):
            if match["is_bye"]:
                continue
            for board in self.db.list_team_boards(int(match["id"])):
                self.service.update_result(tournament_id, int(board["id"]), "1-0")
        self.service.close_round(tournament_id, int(round_id))

    def test_trf25_requested_team_bye_is_not_pab_and_uses_zpb_in_802(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, team_ids = self._create_team_tournament(teams_count=5, boards_count=2)
        self.db.add_requested_team_bye(tournament_id, team_ids[4], 1, "Z")
        round_data = self.service.generate_next_round(tournament_id)
        self._close_team_round_with_decisive_boards(tournament_id, int(round_data["id"]))

        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "team_req_bye.trf"
        exporter.export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        # Bye solicitado nao e pairing-allocated: nenhum registro 320.
        self.assertEqual([line for line in lines if line.startswith("320 ")], [])
        # Sai como 240 (tipo Z).
        self.assertTrue(any(line.startswith("240 ") and line[4:5] == "Z" for line in lines))
        # O 802 da equipe TPN 5 mostra ZPB (zero-point-bye) na rodada 1.
        team5_802 = next(line for line in lines if line.startswith("802 ") and line[4:7].strip() == "5")
        self.assertEqual(team5_802[28:31], "ZPB")

    def test_trf25_allocated_team_bye_is_pab_in_320_and_802(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, _team_ids = self._create_team_tournament(teams_count=3, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        self._close_team_round_with_decisive_boards(tournament_id, int(round_data["id"]))

        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "team_pab.trf"
        exporter.export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        # Bye alocado pelo pareamento: ha um 320 (PAB) e nenhum 240 (nada solicitado).
        self.assertEqual(len([line for line in lines if line.startswith("320 ")]), 1)
        self.assertFalse(any(line.startswith("240 ") for line in lines))
        # O 802 do bye alocado usa o codigo PAB.
        self.assertTrue(any(line.startswith("802 ") and "PAB" in line for line in lines))

    def test_trf25_802_marks_full_match_forfeit(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(int(round_data["id"]))[0]
        boards = self.db.list_team_boards(int(match["id"]))
        # W.O. consistente: uma equipe vence o match inteiro por forfeit.
        # Tab. 1 (impar): brancas vencem; tab. 2 (par, cores invertidas): pretas
        # vencem — ambos os pontos vao para a mesma equipe.
        self.service.update_result(tournament_id, int(boards[0]["id"]), "1F-0F")
        self.service.update_result(tournament_id, int(boards[1]["id"]), "0F-1F")
        self.service.close_round(tournament_id, int(round_data["id"]))

        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "team_ff.trf"
        exporter.export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        # 330 do match forfeitado e indicadores f/F no 802 (col 39 = indice 38;
        # nao some no rstrip por ser caractere nao-branco).
        self.assertTrue(any(line.startswith("330 ") for line in lines))
        markers = sorted(
            line[38] for line in lines if line.startswith("802 ") and len(line) >= 39
        )
        self.assertEqual(markers, ["F", "f"])  # vencedor por W.O. e perdedor

    def test_trf25_802_no_forfeit_marker_on_normal_match(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        self._close_team_round_with_decisive_boards(tournament_id, int(round_data["id"]))

        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "team_normal.trf"
        exporter.export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        # Resultado normal (sem W.O.): nenhum 330 e nenhum indicador f/F no 802
        # (a posicao de forfeit fica em branco, somindo no rstrip).
        self.assertFalse(any(line.startswith("330 ") for line in lines))
        self.assertFalse(
            any(
                line.startswith("802 ") and len(line) >= 39 and line[38] in "fF"
                for line in lines
            )
        )

    def test_trf25_310_lists_starters_before_reserves(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        tournament_id, team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        # Reserva com rating altissimo → start_rank 1; sem fix, viria como 1o jogador.
        reserve_id = self.db.create_player(
            tournament_id, name="Reserva Forte", rating=3000, club="Clube 1"
        )
        self.team_service.add_player(team_ids[0], reserve_id, role="reserve")

        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "team_reserve.trf"
        exporter.export(tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        def ranks_of(line: str) -> list[str]:
            # Jogadores do 310: a partir da col 74 (indice 73), 4 chars, passo 5.
            return [
                line[i:i + 4].strip()
                for i in range(73, len(line), 5)
                if line[i:i + 4].strip()
            ]

        team310 = [line for line in lines if line.startswith("310 ")]
        # A equipe da reserva e a unica cujo 310 contem o start-rank 1.
        target = next(line for line in team310 if "1" in ranks_of(line))
        ranks = ranks_of(target)
        self.assertEqual(len(ranks), 3)         # 2 titulares + 1 reserva
        self.assertEqual(ranks[-1], "1")        # reserva (rank 1) por ultimo
        self.assertNotEqual(ranks[0], "1")      # 1o jogador e um titular (board 1)

    def test_trf25_emits_national_rating_records(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        self.tournament_service.save_profile(
            self.tournament_id,
            {"name": "Torneio teste", "rounds_count": "5", "bye_points": "1"},
            {"federation": "BRA"},
            [],
        )
        # p1 (rating maior) tem rating nacional; p2 nao tem.
        self.db.create_player(
            self.tournament_id, name="Com Nacional", rating=2000,
            national_rating=1928, cbx_id="55501",
        )
        self.db.create_player(self.tournament_id, name="Sem Nacional", rating=1900)

        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "national.trf"
        exporter.export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        nat_lines = [line for line in lines if line.startswith("BRA ")]
        self.assertEqual(len(nat_lines), 1)  # so o jogador com rating nacional
        line = nat_lines[0]
        self.assertEqual(line[4:8].strip(), "1")     # start-rank do p1 (maior rating)
        self.assertEqual(line[48:52].strip(), "1928")  # rating nacional
        self.assertIn("55501", line)                  # nº nacional

        # 172 (Encoded Starting Rank Method) e obrigatorio quando ha registros NRS.
        header_172 = [ln for ln in lines if ln.startswith("172 ")]
        self.assertEqual(len(header_172), 1)
        self.assertEqual(header_172[0][4:7], "BRA")          # federacao (cols 5-7)
        self.assertEqual(header_172[0][8:13].strip(), "FIDE")  # metodo (cols 9-13)

    def test_trf25_omits_national_rating_without_federation(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        # Sem federacao configurada, mesmo com rating nacional nada e emitido.
        self.db.create_player(
            self.tournament_id, name="Com Nacional", rating=2000,
            national_rating=1928, cbx_id="55501",
        )
        self.db.create_player(self.tournament_id, name="Outro", rating=1900)

        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "no_fed.trf"
        exporter.export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        self.assertFalse(any("55501" in line for line in lines))
        # Sem registros NRS, o 172 tambem nao deve aparecer.
        self.assertFalse(any(line.startswith("172 ") for line in lines))

    def test_trf25_validate_signals_data_complete_when_no_pending(self) -> None:
        from src.services.federation_exporters import TRF25Exporter
        from src.services.federation_exporters.trf16 import TRF16Exporter
        from src.services.federation_exporters.trf25 import (
            TRF25_DATA_COMPLETE_NOTE,
            TRF25_SCAFFOLD_WARNING,
        )

        exporter = TRF25Exporter(self.export_service)
        # Sem pendencias herdadas → caveat de formato + nota de prontidao de dados.
        with mock.patch.object(TRF16Exporter, "validate", return_value=[]):
            ready = exporter.validate(self.tournament_id)
        self.assertIn(TRF25_SCAFFOLD_WARNING, ready)
        self.assertIn(TRF25_DATA_COMPLETE_NOTE, ready)
        # Com pendencia herdada → sem nota de prontidao (so o caveat + a pendencia).
        with mock.patch.object(
            TRF16Exporter, "validate", return_value=["Torneio sem cidade/local."]
        ):
            pending = exporter.validate(self.tournament_id)
        self.assertIn(TRF25_SCAFFOLD_WARNING, pending)
        self.assertNotIn(TRF25_DATA_COMPLETE_NOTE, pending)
        self.assertIn("Torneio sem cidade/local.", pending)

    def test_trf25_emits_260_for_prohibited_pairing(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        player_ids = self._create_players(4)  # ratings decrescentes → rank = ordem
        # Proíbe rank 1 x rank 3, janela aberta (last_round=0 → última rodada=5).
        self.db.add_prohibited_pairing(
            self.tournament_id, player_ids[0], player_ids[2]
        )
        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "prohib.trf"
        exporter.export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        prohibition_lines = [line for line in lines if line.startswith("260 ")]
        self.assertEqual(len(prohibition_lines), 1)
        line = prohibition_lines[0]
        self.assertEqual(line[4:7].strip(), "1")  # primeira rodada
        self.assertEqual(line[8:11].strip(), "5")  # última rodada (rounds_count=5)
        self.assertEqual(line[12:16].strip(), "1")  # entidade 1 (rank 1)
        self.assertEqual(line[17:21].strip(), "3")  # entidade 2 (rank 3)

    def test_requested_bye_excludes_player_from_pairing(self) -> None:
        # 5 jogadores, 1 bye solicitado → 4 a parear (par), sem bye alocado extra.
        player_ids = self._create_players(5)
        self.db.add_requested_bye(self.tournament_id, player_ids[4], 1, "H")

        round_data = self.service.generate_next_round(self.tournament_id)
        pairings = self.db.get_pairings_for_round(int(round_data["id"]))

        byes = [p for p in pairings if p["is_bye"]]
        self.assertEqual(len(byes), 1)
        self.assertEqual(int(byes[0]["white_player_id"]), player_ids[4])
        self.assertEqual(str(byes[0]["result"]).upper(), "H")

        normal = [p for p in pairings if not p["is_bye"]]
        self.assertEqual(len(normal), 2)  # 4 jogadores → 2 confrontos
        paired = {int(p["white_player_id"]) for p in normal}
        paired |= {int(p["black_player_id"]) for p in normal if p["black_player_id"]}
        self.assertNotIn(player_ids[4], paired)

    def _enable_disable_bye(self) -> None:
        self.tournament_service.save_profile(
            self.tournament_id,
            {"name": "Torneio teste", "rounds_count": "5", "bye_points": "1"},
            {"initial_order": "rating", "tournament_type": "real", "disable_bye": 1},
            [],
        )

    def test_requested_bye_makes_odd_field_pairable_under_disable_bye(self) -> None:
        # disable_bye + 5 ativos (impar): o bye solicitado deixa 4 a parear, par.
        # A guarda de paridade deve incidir sobre to_pair, nao sobre todos.
        self._enable_disable_bye()
        player_ids = self._create_players(5)
        self.db.add_requested_bye(self.tournament_id, player_ids[4], 1, "H")

        round_data = self.service.generate_next_round(self.tournament_id)
        pairings = self.db.get_pairings_for_round(int(round_data["id"]))

        byes = [p for p in pairings if p["is_bye"]]
        self.assertEqual(len(byes), 1)  # apenas o bye solicitado, nenhum alocado
        self.assertEqual(int(byes[0]["white_player_id"]), player_ids[4])
        self.assertEqual(len([p for p in pairings if not p["is_bye"]]), 2)

    def test_disable_bye_still_blocks_odd_field_without_requested_bye(self) -> None:
        # Sem bye solicitado, disable_bye + impar continua bloqueando.
        self._enable_disable_bye()
        self._create_players(5)
        with self.assertRaises(AppError):
            self.service.generate_next_round(self.tournament_id)

    def _set_individual_pairing_method(self, method: str) -> None:
        self.db.get_tournament_settings(self.tournament_id)  # garante a linha
        with self.db.connect() as connection:
            connection.execute(
                "UPDATE tournament_settings SET pairing_method = ? WHERE tournament_id = ?",
                (method, self.tournament_id),
            )

    def test_requested_bye_rejected_in_round_robin(self) -> None:
        player_ids = self._create_players(4)
        self._set_individual_pairing_method("round_robin")
        self.db.add_requested_bye(self.tournament_id, player_ids[0], 1, "H")
        with self.assertRaisesRegex(AppError, "exclusivos do sistema Suico"):
            self.service.generate_next_round(self.tournament_id)

    def test_requested_bye_rejected_in_knockout(self) -> None:
        player_ids = self._create_players(4)
        self._set_individual_pairing_method("knockout")
        self.db.add_requested_bye(self.tournament_id, player_ids[0], 1, "Z")
        with self.assertRaisesRegex(AppError, "exclusivos do sistema Suico"):
            self.service.generate_next_round(self.tournament_id)

    def test_round_robin_still_generates_without_requested_bye(self) -> None:
        # Regressao: sem bye solicitado, round-robin continua gerando normalmente.
        self._create_players(4)
        self._set_individual_pairing_method("round_robin")
        round_data = self.service.generate_next_round(self.tournament_id)
        self.assertTrue(int(round_data["id"]))
        pairings = self.db.get_pairings_for_round(int(round_data["id"]))
        self.assertEqual(len(pairings), 2)  # 4 jogadores → 2 confrontos

    def test_requested_bye_scores_by_type(self) -> None:
        from src.services.pairing import calculate_player_standings

        player_ids = self._create_players(4)  # bye_points padrão = 1.0
        players = self.db.list_players(self.tournament_id, active_only=False)
        tournament = self.db.get_tournament(self.tournament_id)
        closed_pairings = [
            {"white_player_id": player_ids[0], "black_player_id": None,
             "result": "F", "is_bye": 1, "round_number": 1},
            {"white_player_id": player_ids[1], "black_player_id": None,
             "result": "H", "is_bye": 1, "round_number": 1},
            {"white_player_id": player_ids[2], "black_player_id": None,
             "result": "Z", "is_bye": 1, "round_number": 1},
            {"white_player_id": player_ids[3], "black_player_id": None,
             "result": "BYE", "is_bye": 1, "round_number": 1},
        ]
        standings = calculate_player_standings(tournament, players, closed_pairings)
        points = {int(s["player_id"]): float(s["points"]) for s in standings}
        self.assertEqual(points[player_ids[0]], 1.0)  # F = ponto inteiro
        self.assertEqual(points[player_ids[1]], 0.5)  # H = meio ponto
        self.assertEqual(points[player_ids[2]], 0.0)  # Z = zero ponto
        self.assertEqual(points[player_ids[3]], 1.0)  # bye alocado usa bye_points

    # --- Fase A: desempates configuraveis e completos (E1/E2) -------------

    def _phase_a_fixture(self) -> tuple[dict, list[dict], list[int]]:
        """Cenario deterministico de 4 jogadores e 3 rodadas.

        R1: 1>4, 2>3 ; R2: 1>2, 3>4 ; R3: 1=3, 2>4.
        Pontos finais: P1=2.5, P2=2.0, P3=1.5, P4=0.0.
        """
        ids = self._create_players(4)  # ratings 2000,1950,1900,1850
        tournament = self.db.get_tournament(self.tournament_id)
        closed = [
            {"white_player_id": ids[0], "black_player_id": ids[3], "result": "1-0", "is_bye": 0, "round_number": 1},
            {"white_player_id": ids[1], "black_player_id": ids[2], "result": "1-0", "is_bye": 0, "round_number": 1},
            {"white_player_id": ids[0], "black_player_id": ids[1], "result": "1-0", "is_bye": 0, "round_number": 2},
            {"white_player_id": ids[2], "black_player_id": ids[3], "result": "1-0", "is_bye": 0, "round_number": 2},
            {"white_player_id": ids[0], "black_player_id": ids[2], "result": "1/2-1/2", "is_bye": 0, "round_number": 3},
            {"white_player_id": ids[1], "black_player_id": ids[3], "result": "1-0", "is_bye": 0, "round_number": 3},
        ]
        return tournament, closed, ids

    def test_tiebreak_default_sequence_reproduces_legacy_order(self) -> None:
        from src.services.pairing import DEFAULT_PLAYER_TIEBREAKS, calculate_player_standings

        tournament, closed, ids = self._phase_a_fixture()
        players = self.db.list_players(self.tournament_id, active_only=False)

        standings = calculate_player_standings(tournament, players, closed)

        # Ordem por pontos (P1>P2>P3>P4) e a sequencia padrao (sem config) e usada.
        self.assertEqual([s["player_id"] for s in standings], [ids[0], ids[1], ids[2], ids[3]])
        self.assertEqual(standings[0]["tiebreak_order"], DEFAULT_PLAYER_TIEBREAKS)
        # Componentes historicos seguem presentes (compat. com exportacoes).
        self.assertLessEqual(
            {"buchholz", "buchholz_median", "sonneborn_berger", "direct_encounter", "wins", "performance"},
            set(standings[0]["tiebreak_components"].keys()),
        )

    def test_tiebreak_new_criteria_values_match_manual(self) -> None:
        from src.services.pairing import calculate_player_standings, parse_player_tiebreak_sequence

        tournament, closed, ids = self._phase_a_fixture()
        players = self.db.list_players(self.tournament_id, active_only=False)
        sequence = parse_player_tiebreak_sequence(
            [{"code": "buchholz_cut1"}, {"code": "cumulative"}, {"code": "koya"}, {"code": "aro"}, {"code": "black_games"}]
        )

        standings = calculate_player_standings(tournament, players, closed, sequence=sequence)
        leader = next(s for s in standings if s["player_id"] == ids[0])
        components = leader["tiebreak_components"]

        # P1 enfrentou P4(0 pts), P2(2.0 pts), P3(1.5 pts); ratings 1850/1950/1900.
        self.assertEqual(components["buchholz_cut1"]["value"], 3.5)   # descarta o 0 → 2.0+1.5
        self.assertEqual(components["cumulative"]["value"], 5.5)      # 1+2+2.5
        self.assertEqual(components["koya"]["value"], 1.5)            # vs P2 (1.0) + vs P3 (0.5)
        self.assertEqual(components["aro"]["value"], 1900.0)          # (1850+1950+1900)/3
        self.assertEqual(components["black_games"]["value"], 0.0)     # P1 jogou sempre de brancas

    def test_tiebreak_buchholz_cut_handles_bye_via_unplayed_param(self) -> None:
        from src.services.pairing import calculate_player_standings

        ids = self._create_players(3)  # X=ids[0], Y=ids[1], Z=ids[2]
        tournament = self.db.get_tournament(self.tournament_id)
        players = self.db.list_players(self.tournament_id, active_only=False)
        closed = [
            {"white_player_id": ids[0], "black_player_id": ids[1], "result": "1-0", "is_bye": 0, "round_number": 1},
            {"white_player_id": ids[0], "black_player_id": ids[2], "result": "1-0", "is_bye": 0, "round_number": 2},
            {"white_player_id": ids[0], "black_player_id": None, "result": "BYE", "is_bye": 1, "round_number": 3},
        ]

        def cut1(unplayed: str) -> float:
            standings = calculate_player_standings(
                tournament, players, closed,
                sequence=[{"code": "buchholz_cut1", "params": {"cut_low": 1, "unplayed": unplayed}}],
            )
            leader = next(s for s in standings if s["player_id"] == ids[0])
            return leader["tiebreak_values"]["buchholz_cut1"]

        # Adversarios Y e Z terminam com 0 pts. "real" ignora o bye → descarta um 0 → 0.0.
        self.assertEqual(cut1("real"), 0.0)
        # "self" injeta a propria pontuacao do jogador (3.0) para a rodada de bye.
        self.assertEqual(cut1("self"), 3.0)

    def test_tiebreak_sequence_changes_final_order(self) -> None:
        from src.services.pairing import calculate_player_standings

        ids = self._create_players(8)  # A=ids[0] (2000), B=ids[1] (1950)
        tournament = self.db.get_tournament(self.tournament_id)
        players = self.db.list_players(self.tournament_id, active_only=False)
        closed = [
            # A: 2 vitorias de brancas e 1 derrota → 2 pts, 2 vitorias, 0 partidas de pretas.
            {"white_player_id": ids[0], "black_player_id": ids[2], "result": "1-0", "is_bye": 0, "round_number": 1},
            {"white_player_id": ids[0], "black_player_id": ids[3], "result": "1-0", "is_bye": 0, "round_number": 2},
            {"white_player_id": ids[0], "black_player_id": ids[4], "result": "0-1", "is_bye": 0, "round_number": 3},
            # B: vitoria de pretas + 2 empates de pretas → 2 pts, 1 vitoria, 3 partidas de pretas.
            {"white_player_id": ids[5], "black_player_id": ids[1], "result": "0-1", "is_bye": 0, "round_number": 1},
            {"white_player_id": ids[6], "black_player_id": ids[1], "result": "1/2-1/2", "is_bye": 0, "round_number": 2},
            {"white_player_id": ids[7], "black_player_id": ids[1], "result": "1/2-1/2", "is_bye": 0, "round_number": 3},
        ]

        def position_of(player_id: int, sequence: list[dict]) -> int:
            standings = calculate_player_standings(tournament, players, closed, sequence=sequence)
            return next(s["position"] for s in standings if s["player_id"] == player_id)

        # Empatados em pontos (2.0). Por vitorias, A (2) fica a frente de B (1).
        self.assertLess(position_of(ids[0], [{"code": "wins"}]), position_of(ids[1], [{"code": "wins"}]))
        # Por partidas com pretas, B (3) fica a frente de A (0) — ordem invertida.
        self.assertLess(position_of(ids[1], [{"code": "black_games"}]), position_of(ids[0], [{"code": "black_games"}]))

    def test_tiebreak_sequence_persists_and_threads_into_standings(self) -> None:
        from src.services.pairing import serialize_tiebreak_sequence

        self._create_players(4)
        sequence = [{"code": "direct_encounter", "params": {}}, {"code": "cumulative", "params": {}}, {"code": "wins", "params": {}}]
        self.db.save_tournament_settings(
            self.tournament_id, {"tiebreak_sequence": serialize_tiebreak_sequence(sequence)}
        )

        round_data = self.service.generate_next_round(self.tournament_id)
        round_id = int(round_data["id"])
        for pairing in self.db.get_pairings_for_round(round_id):
            self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, round_id)

        # A sequencia gravada chega ate o calculo da classificacao.
        standings = self.service.standings(self.tournament_id)
        self.assertEqual(standings[0]["tiebreak_order"], ["direct_encounter", "cumulative", "wins"])
        settings = self.db.get_tournament_settings(self.tournament_id)
        self.assertEqual(settings["tiebreak_sequence"], serialize_tiebreak_sequence(sequence))

    def test_validated_settings_normalizes_tiebreak_sequence(self) -> None:
        from src.services.pairing import serialize_tiebreak_sequence

        payload = self.tournament_service._validated_settings(
            {"tiebreak_sequence": [{"code": "points"}, {"code": "buchholz"}, {"code": "xxx"}, {"code": "buchholz"}]}
        )

        # 'points' (implicito), criterio desconhecido e duplicata sao descartados.
        self.assertEqual(payload["tiebreak_sequence"], serialize_tiebreak_sequence([{"code": "buchholz", "params": {}}]))
        self.assertEqual(payload["team_tiebreak_sequence"], "[]")

    def test_v35_database_adds_tiebreak_sequence_columns(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v35_tiebreaks.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE tournament_settings (
                    tournament_id INTEGER PRIMARY KEY,
                    updated_at TEXT NOT NULL DEFAULT ''
                );
                INSERT INTO tournament_settings (tournament_id, updated_at)
                VALUES (1, '2026-06-01 00:00:00');
                PRAGMA user_version = 34;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(tournament_settings)").fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("tiebreak_sequence", columns)
        self.assertIn("team_tiebreak_sequence", columns)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    # --- Fase B: relatorio de variacao de rating FIDE (E3) ----------------

    @staticmethod
    def _fide_player(pid: int, rating: int, *, intl: int | None = None, birth: str = "") -> dict:
        return {
            "id": pid, "name": f"P{pid}", "surname": "", "given_name": "",
            "international_rating": rating if intl is None else intl,
            "national_rating": 0, "rating": rating, "birth_date": birth, "k_factor": None,
        }

    def test_fide_expected_score_applies_400_rule(self) -> None:
        from src.services.fide_rating import fide_expected_score

        self.assertEqual(fide_expected_score(2000, 2000), 0.50)
        self.assertEqual(fide_expected_score(2000, 1800), 0.76)
        self.assertEqual(fide_expected_score(1800, 2000), 0.24)
        # Diferenca de 500 e tratada como 400 (regra dos 400).
        self.assertEqual(fide_expected_score(2000, 1500), 0.92)
        self.assertEqual(fide_expected_score(1500, 2000), 0.08)

    def test_fide_k_factor_rules(self) -> None:
        from src.services.fide_rating import fide_k_factor

        self.assertEqual(fide_k_factor(2500), 10)
        self.assertEqual(fide_k_factor(2000), 20)
        self.assertEqual(fide_k_factor(2000, k_override=40), 40)
        self.assertEqual(fide_k_factor(2200, birth_year=2012, tournament_year=2026), 40)  # sub-18 <2300
        self.assertEqual(fide_k_factor(2350, birth_year=2012, tournament_year=2026), 20)  # sub-18 mas >=2300

    def test_fide_report_delta_matches_manual(self) -> None:
        from src.services.fide_rating import build_fide_report_rows

        rows = build_fide_report_rows(
            [self._fide_player(1, 2000), self._fide_player(2, 1800)],
            [{"white_player_id": 1, "black_player_id": 2, "result": "1-0", "is_bye": 0}],
            "fide",
        )
        by_id = {row["player_id"]: row for row in rows}
        # We(A)=0.76, K=20, ΔA=20*(1-0.76)=4.8, Rc=2005.
        self.assertEqual(by_id[1]["we"], 0.76)
        self.assertEqual(by_id[1]["k"], 20)
        self.assertEqual(by_id[1]["delta"], 4.8)
        self.assertEqual(by_id[1]["rc"], 2005)
        # We(B)=0.24, ΔB=-4.8, Rc=1795.
        self.assertEqual(by_id[2]["delta"], -4.8)
        self.assertEqual(by_id[2]["rc"], 1795)

    def test_fide_report_400_rule_and_unrated_opponent(self) -> None:
        from src.services.fide_rating import build_fide_report_rows

        # 400: A=2000 vence C=1500 → We=0.92, n_over_400=1, Δ=1.6.
        rows = build_fide_report_rows(
            [self._fide_player(1, 2000), self._fide_player(3, 1500)],
            [{"white_player_id": 1, "black_player_id": 3, "result": "1-0", "is_bye": 0}],
            "fide",
        )
        leader = next(row for row in rows if row["player_id"] == 1)
        self.assertEqual(leader["we"], 0.92)
        self.assertEqual(leader["n_over_400"], 1)
        self.assertEqual(leader["delta"], 1.6)

        # Adversario sem rating e ignorado no calculo; jogador sem rating recebe so Rp.
        rows = build_fide_report_rows(
            [self._fide_player(1, 2000), self._fide_player(2, 1800), self._fide_player(4, 0, intl=0)],
            [
                {"white_player_id": 1, "black_player_id": 2, "result": "1-0", "is_bye": 0},
                {"white_player_id": 1, "black_player_id": 4, "result": "1-0", "is_bye": 0},
            ],
            "fide",
        )
        by_id = {row["player_id"]: row for row in rows}
        self.assertEqual(by_id[1]["games_rated"], 1)   # D (sem rating) nao conta para A
        self.assertEqual(by_id[1]["we"], 0.76)
        self.assertEqual(by_id[4]["ro"], 0)
        self.assertIsNone(by_id[4]["delta"])
        self.assertEqual(by_id[4]["rp"], 1200)         # 2000 + dp(0.0) = 2000 - 800

    def test_fide_report_excludes_byes_and_walkovers(self) -> None:
        from src.services.fide_rating import build_fide_report_rows

        rows = build_fide_report_rows(
            [self._fide_player(1, 2000), self._fide_player(2, 1800)],
            [
                {"white_player_id": 1, "black_player_id": 2, "result": "1F-0F", "is_bye": 0},  # WO: nao conta
                {"white_player_id": 1, "black_player_id": None, "result": "BYE", "is_bye": 1},  # bye: nao conta
            ],
            "fide",
        )
        # Nenhuma partida valida para rating → ninguem entra no relatorio.
        self.assertEqual(rows, [])

    def test_fide_rating_report_persists_idempotently(self) -> None:
        from src.services.rating_service import FideRatingService

        self._create_players(4)
        round_data = self.service.generate_next_round(self.tournament_id)
        round_id = int(round_data["id"])
        for pairing in self.db.get_pairings_for_round(round_id):
            self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, round_id)

        service = FideRatingService(self.db)
        service.save_report(self.tournament_id, "fide")
        service.save_report(self.tournament_id, "fide")  # recomputar nao duplica

        saved = self.db.get_fide_rating_report(self.tournament_id, "fide")
        self.assertEqual(len(saved), 4)                       # 4 jogadores, 1 partida ranqueada cada
        self.assertEqual(len({row["player_id"] for row in saved}), 4)
        self.assertTrue(all(row["games_rated"] == 1 for row in saved))

    def test_export_fide_rating_report_writes_file(self) -> None:
        self._create_players(4)
        round_data = self.service.generate_next_round(self.tournament_id)
        round_id = int(round_data["id"])
        for pairing in self.db.get_pairings_for_round(round_id):
            self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, round_id)

        path = Path(self.temp_dir.name) / "rating_fide.csv"
        self.export_service.export_fide_rating_report(self.tournament_id, path, "fide")

        self.assertTrue(path.exists())
        content = path.read_text(encoding="utf-8-sig")
        self.assertIn("Variacao de rating FIDE", content)
        self.assertIn("Rp", content)
        # Exportar persiste um snapshot para auditoria/reimpressao.
        self.assertEqual(len(self.db.get_fide_rating_report(self.tournament_id, "fide")), 4)

    def test_v36_database_adds_fide_rating_schema(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v36_fide.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE players (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    tournament_id INTEGER,
                    name TEXT NOT NULL
                );
                PRAGMA user_version = 35;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            player_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(players)").fetchall()
            }
            tables = {
                row["name"]
                for row in migrated_connection.execute(
                    "SELECT name FROM sqlite_master WHERE type = 'table'"
                ).fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("k_factor", player_columns)
        self.assertIn("fide_rating_reports", tables)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    # --- Fase C: distribuicao de premios (E4) -----------------------------

    @staticmethod
    def _prize_standings() -> list[dict]:
        # P1 e P2 empatados em pontos (3.0) ocupam as posicoes 1-2.
        return [
            {"player_id": 1, "name": "P1", "position": 1, "points": 3.0, "category": "A"},
            {"player_id": 2, "name": "P2", "position": 2, "points": 3.0, "category": "B"},
            {"player_id": 3, "name": "P3", "position": 3, "points": 2.0, "category": "A"},
            {"player_id": 4, "name": "P4", "position": 4, "points": 1.0, "category": "B"},
        ]

    @staticmethod
    def _prize_rows() -> list[dict]:
        return [
            {"kind": "overall", "label": "1o", "rank_from": 1, "rank_to": 1, "amount": 1000},
            {"kind": "overall", "label": "2o", "rank_from": 2, "rank_to": 2, "amount": 600},
            {"kind": "overall", "label": "3o", "rank_from": 3, "rank_to": 3, "amount": 200},
            {"kind": "category", "label": "Cat A 1o", "category": "A", "rank_from": 1, "rank_to": 1, "amount": 300},
            {"kind": "category", "label": "Cat B 1o", "category": "B", "rank_from": 1, "rank_to": 1, "amount": 300},
        ]

    def test_prize_best_only_splits_ties_without_double_counting(self) -> None:
        from src.services.prizes import allocate_prizes

        result = allocate_prizes(self._prize_standings(), self._prize_rows(), "best_only", 0.0)
        by_id = {row["player_id"]: row for row in result["allocations"]}
        # Empate 1-2 divide (1000+600)/2 = 800 cada; sem somar a categoria.
        self.assertEqual(by_id[1]["gross"], 800.0)
        self.assertEqual(by_id[2]["gross"], 800.0)
        self.assertEqual(by_id[3]["gross"], 200.0)
        self.assertNotIn(4, by_id)
        self.assertEqual(result["total_gross"], 1800.0)

    def test_prize_cumulative_adds_overall_and_category(self) -> None:
        from src.services.prizes import allocate_prizes

        result = allocate_prizes(self._prize_standings(), self._prize_rows(), "cumulative", 0.0)
        by_id = {row["player_id"]: row for row in result["allocations"]}
        self.assertEqual(by_id[1]["gross"], 1100.0)  # 800 + 300
        self.assertEqual(by_id[2]["gross"], 1100.0)
        self.assertEqual(by_id[3]["gross"], 200.0)
        self.assertEqual(result["total_gross"], 2400.0)

    def test_prize_hort_policy_combines_overall_and_category(self) -> None:
        from src.services.prizes import allocate_prizes

        standings = [
            {"player_id": 1, "name": "P1", "position": 1, "points": 5.0, "category": "Y"},
            {"player_id": 2, "name": "P2", "position": 2, "points": 4.0, "category": "Y"},
            {"player_id": 3, "name": "P3", "position": 3, "points": 3.0, "category": "Y"},
            {"player_id": 10, "name": "P10", "position": 4, "points": 1.0, "category": "X"},
        ]
        prizes = [
            {"kind": "overall", "label": "4o", "rank_from": 4, "rank_to": 4, "amount": 100},
            {"kind": "category", "label": "Cat X 1o", "category": "X", "rank_from": 1, "rank_to": 1, "amount": 500},
        ]

        def gross(policy: str) -> float:
            return allocate_prizes(standings, prizes, policy, 0.0)["allocations"][0]["gross"]

        self.assertEqual(gross("best_only"), 500.0)        # max(100, 500)
        self.assertEqual(gross("cumulative"), 600.0)       # 100 + 500
        self.assertEqual(gross("hort"), 300.0)             # max(100, (100+500)/2)

    def test_prize_tax_deduction(self) -> None:
        from src.services.prizes import allocate_prizes

        result = allocate_prizes(self._prize_standings(), self._prize_rows(), "best_only", 10.0)
        by_id = {row["player_id"]: row for row in result["allocations"]}
        self.assertEqual(by_id[1]["net"], 720.0)           # 800 * 0.9
        self.assertEqual(result["total_net"], 1620.0)      # 1800 * 0.9
        self.assertEqual(result["total_tax"], 180.0)

    def test_prize_special_listed_as_manual(self) -> None:
        from src.services.prizes import allocate_prizes

        prizes = self._prize_rows() + [{"kind": "special", "label": "Melhor feminino", "amount": 150}]
        result = allocate_prizes(self._prize_standings(), prizes, "best_only", 0.0)
        self.assertEqual(len(result["manual_prizes"]), 1)
        self.assertEqual(result["manual_prizes"][0]["label"], "Melhor feminino")
        self.assertEqual(result["total_gross"], 1800.0)    # especial nao entra no total automatico

    def test_prize_service_persists_validates_and_allocates(self) -> None:
        self._create_players(4)  # categoria "Absoluto"
        round_data = self.service.generate_next_round(self.tournament_id)
        round_id = int(round_data["id"])
        for pairing in self.db.get_pairings_for_round(round_id):
            self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, round_id)

        self.prize_service.replace_prizes(
            self.tournament_id,
            [
                {"kind": "overall", "label": "1o", "rank_from": 1, "rank_to": 1, "amount": 1000},
                {"kind": "overall", "label": "2o", "rank_from": 2, "rank_to": 2, "amount": 600},
            ],
        )
        # Persistencia idempotente.
        self.assertEqual(len(self.db.list_tournament_prizes(self.tournament_id)), 2)

        result = self.prize_service.allocate(self.tournament_id)
        # 2 vencedores empatados em 1.0 dividem (1000+600)/2 = 800 cada.
        self.assertEqual(result["winners"], 2)
        self.assertEqual(result["total_gross"], 1600.0)
        self.assertTrue(all(row["gross"] == 800.0 for row in result["allocations"]))

        with self.assertRaisesRegex(AppError, "categoria"):
            self.prize_service.replace_prizes(
                self.tournament_id,
                [{"kind": "category", "label": "Sub-12", "amount": 100}],  # sem categoria
            )

    def test_export_prize_report_writes_file(self) -> None:
        self._create_players(4)
        round_data = self.service.generate_next_round(self.tournament_id)
        round_id = int(round_data["id"])
        for pairing in self.db.get_pairings_for_round(round_id):
            self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, round_id)
        self.prize_service.replace_prizes(
            self.tournament_id,
            [{"kind": "overall", "label": "1o", "rank_from": 1, "rank_to": 1, "amount": 1000}],
        )

        path = Path(self.temp_dir.name) / "premiacao.csv"
        self.export_service.export_prize_report(self.tournament_id, path)
        self.assertTrue(path.exists())
        content = path.read_text(encoding="utf-8-sig")
        self.assertIn("Premiacao por jogador", content)
        self.assertIn("Total liquido distribuido", content)

    def test_v37_database_adds_prize_schema(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v37_prizes.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE tournament_settings (
                    tournament_id INTEGER PRIMARY KEY,
                    updated_at TEXT NOT NULL DEFAULT ''
                );
                INSERT INTO tournament_settings (tournament_id, updated_at)
                VALUES (1, '2026-06-01 00:00:00');
                PRAGMA user_version = 36;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            settings_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(tournament_settings)").fetchall()
            }
            tables = {
                row["name"]
                for row in migrated_connection.execute(
                    "SELECT name FROM sqlite_master WHERE type = 'table'"
                ).fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("prize_policy", settings_columns)
        self.assertIn("prize_tax_percent", settings_columns)
        self.assertIn("tournament_prizes", tables)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    # --- Fase D: estatisticas e fichas no padrao FIDE (E6) ----------------

    def _play_first_round(self, results: list[str]) -> None:
        round_data = self.service.generate_next_round(self.tournament_id)
        round_id = int(round_data["id"])
        pairings = self.db.get_pairings_for_round(round_id)
        for pairing, result in zip(pairings, results):
            self.service.update_result(self.tournament_id, int(pairing["id"]), result)
        self.service.close_round(self.tournament_id, round_id)

    def test_game_statistics_counts_results(self) -> None:
        self._create_players(8)  # 4 partidas na rodada 1
        self._play_first_round(["1-0", "0-1", "1/2-1/2", "1F-0F"])

        sections = self.export_service._game_statistics_sections(self.tournament_id)
        summary = {row[0]: row[1] for row in sections[0][2]}
        distribution = {row[0]: row[1] for row in sections[1][2]}

        self.assertEqual(summary["Partidas jogadas (tabuleiro)"], 3)  # WO nao conta
        self.assertEqual(summary["WO / forfait"], 1)
        self.assertEqual(summary["Byes"], 0)
        self.assertEqual(summary["Total de pareamentos fechados"], 4)
        self.assertEqual(distribution["Vitorias de brancas"], 1)
        self.assertEqual(distribution["Empates"], 1)
        self.assertEqual(distribution["Vitorias de pretas"], 1)

    def test_federation_statistics_groups_players(self) -> None:
        ids = self._create_players(8)
        with self.db.connect() as connection:
            for pid in ids[:3]:
                connection.execute("UPDATE players SET federation_id = ? WHERE id = ?", ("BRA", pid))
        self._play_first_round(["1-0", "1-0", "1-0", "1-0"])

        sections = self.export_service._federation_sections(self.tournament_id)
        summary = {row[0]: row[1] for row in sections[0][2]}
        fed_rows = {row[0]: row[1] for row in sections[1][2]}

        self.assertEqual(summary["Jogadores"], 8)
        self.assertEqual(summary["Federacoes"], 2)   # BRA e "—"
        self.assertEqual(fed_rows["BRA"], 3)
        self.assertEqual(fed_rows["—"], 5)

    def test_player_cards_summary_and_round_by_round(self) -> None:
        self._create_players(8)
        self._play_first_round(["1-0", "0-1", "1/2-1/2", "1-0"])

        sections = self.export_service._player_cards_sections(self.tournament_id)
        summary_rows = sections[0][2]
        game_rows = sections[1][2]

        self.assertEqual(len(summary_rows), 8)   # uma linha de resumo por jogador
        self.assertEqual(len(game_rows), 8)       # 8 jogadores, 1 partida cada na rodada 1
        # V + E + D somados entre todos = 8 jogadores com 1 jogo cada.
        total_results = sum(int(row[6]) + int(row[7]) + int(row[8]) for row in summary_rows)
        self.assertEqual(total_results, 8)

    def test_export_player_cards_writes_file(self) -> None:
        self._create_players(8)
        self._play_first_round(["1-0", "0-1", "1/2-1/2", "1-0"])

        path = Path(self.temp_dir.name) / "fichas.csv"
        self.export_service.export_player_cards(self.tournament_id, path)
        self.assertTrue(path.exists())
        content = path.read_text(encoding="utf-8-sig")
        self.assertIn("Resumo por jogador", content)
        self.assertIn("Resultados rodada a rodada", content)

    def test_fide_statistics_reject_team_tournaments(self) -> None:
        team_tournament_id, _team_ids = self._create_team_tournament()
        with self.assertRaisesRegex(AppError, "individuais"):
            self.export_service._game_statistics_sections(team_tournament_id)

    # --- Fase E: assistente de normas/titulos FIDE (E5) -------------------

    def test_norm_verdict_meets_and_lists_missing(self) -> None:
        from src.services.fide_norms import evaluate_titles

        meets = next(
            item for item in evaluate_titles({"sex": "M"}, 2650, 9, 3, 3, 2400) if item["title"] == "GM"
        )
        self.assertTrue(meets["meets"])
        self.assertEqual(meets["missing"], [])

        fails = next(
            item for item in evaluate_titles({"sex": "M"}, 2500, 9, 2, 1, 2200) if item["title"] == "GM"
        )
        self.assertFalse(fails["meets"])
        joined = " | ".join(fails["missing"])
        self.assertIn("Performance", joined)
        self.assertIn("Federacoes", joined)

    def test_norm_report_identifies_gm_candidate(self) -> None:
        from src.services.fide_norms import build_norm_report

        def player(pid: int, rating: int, title: str = "", federation: str = "") -> dict:
            return {
                "id": pid, "name": f"P{pid}", "surname": "", "given_name": "",
                "international_rating": rating, "national_rating": 0, "rating": rating,
                "title": title, "federation_id": federation, "sex": "M",
            }

        players = [player(1, 2700, federation="A")]
        federations = ["A", "B", "C"]
        for opponent in range(2, 11):
            players.append(player(opponent, 2500, "GM" if opponent % 2 else "IM", federations[opponent % 3]))
        pairings = [
            {"white_player_id": 1, "black_player_id": opponent, "result": "1-0" if index < 6 else "0-1", "is_bye": 0}
            for index, opponent in enumerate(range(2, 11))
        ]

        report = build_norm_report(players, pairings, "fide")
        leader = next(item for item in report if item["player_id"] == 1)
        self.assertEqual(leader["games"], 9)
        self.assertGreaterEqual(leader["performance"], 2600)
        self.assertTrue(any(title["title"] == "GM" and title["meets"] for title in leader["titles"]))

    def test_norm_service_and_export(self) -> None:
        self._create_players(8)
        self._play_first_round(["1-0", "0-1", "1/2-1/2", "1-0"])

        report = self.norm_assistant_service.evaluate_tournament(self.tournament_id)
        self.assertTrue(report["players"])  # todos jogaram 1 partida

        path = Path(self.temp_dir.name) / "normas.csv"
        self.export_service.export_norm_report(self.tournament_id, path)
        self.assertTrue(path.exists())
        content = path.read_text(encoding="utf-8-sig")
        self.assertIn("Indicadores de norma por jogador", content)
        self.assertIn("NAO concede norma", content)

    def test_norm_report_rejects_team_tournaments(self) -> None:
        team_tournament_id, _team_ids = self._create_team_tournament()
        with self.assertRaisesRegex(AppError, "individuais"):
            self.norm_assistant_service.evaluate_tournament(team_tournament_id)

    def test_export_arbiter_norm_report_writes_file(self) -> None:
        self._create_players(2)
        settings = self.db.get_tournament_settings(self.tournament_id) or {}
        settings["chief_arbiter"] = "Joao Arbitro"
        settings["arbiters"] = "Maria Aux, Pedro Aux"
        settings["federation"] = "BRA"
        self.db.save_tournament_settings(self.tournament_id, settings)

        path = Path(self.temp_dir.name) / "arbitro.csv"
        self.export_service.export_arbiter_norm_report(self.tournament_id, path)

        self.assertTrue(path.exists())
        content = path.read_text(encoding="utf-8-sig")
        self.assertIn("Dados do torneio (norma de arbitro)", content)
        self.assertIn("Arbitros designados", content)
        self.assertIn("Joao Arbitro", content)
        self.assertIn("Maria Aux", content)

    # --- Fase F: editor de colunas da classificacao (E7) ------------------

    def test_standings_default_columns_unchanged(self) -> None:
        self._create_players(4)
        _title, headers, rows = self.export_service._standings_section(self.tournament_id)
        self.assertEqual(
            headers,
            [
                "Pos", "Nome", "Categoria", "Categoria idade", "Categoria rating",
                "Tags premiacao", "Pts", "Buchholz", "Buchholz M", "SB",
                "Vitorias", "Performance", "Rating", "Clube",
            ],
        )
        self.assertEqual(len(rows), 4)
        self.assertTrue(all(len(row) == 14 for row in rows))

    def test_standings_layout_changes_columns(self) -> None:
        self._create_players(4)
        self.list_layout_service.save_columns(self.tournament_id, ["position", "name", "points"], "standings")

        _title, headers, rows = self.export_service._standings_section(self.tournament_id)
        self.assertEqual(headers, ["Pos", "Nome", "Pts"])
        self.assertTrue(all(len(row) == 3 for row in rows))

    def test_list_layout_normalizes_and_validates(self) -> None:
        self._create_players(2)
        # Coluna desconhecida e duplicata sao descartadas; ordem preservada.
        saved = self.list_layout_service.save_columns(
            self.tournament_id, ["position", "xxx", "name", "name"], "standings"
        )
        self.assertEqual([spec["key"] for spec in saved], ["position", "name"])
        selected = self.list_layout_service.get_columns(self.tournament_id, "standings")["selected"]
        self.assertEqual([spec["key"] for spec in selected], ["position", "name"])
        # Selecao vazia / so codigos invalidos sao rejeitadas.
        with self.assertRaisesRegex(AppError, "ao menos uma coluna"):
            self.list_layout_service.save_columns(self.tournament_id, [], "standings")
        with self.assertRaisesRegex(AppError, "ao menos uma coluna"):
            self.list_layout_service.save_columns(self.tournament_id, ["xxx"], "standings")

    def test_list_layout_reset_returns_to_default(self) -> None:
        self._create_players(2)
        self.list_layout_service.save_columns(self.tournament_id, ["position", "name"], "standings")
        self.list_layout_service.reset(self.tournament_id, "standings")
        # Apos reset, a classificacao volta a usar as colunas padrao.
        _title, headers, _rows = self.export_service._standings_section(self.tournament_id)
        self.assertEqual(len(headers), 14)

    def test_v38_database_adds_report_layouts(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v38_layouts.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript("PRAGMA user_version = 37;")
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            tables = {
                row["name"]
                for row in migrated_connection.execute(
                    "SELECT name FROM sqlite_master WHERE type = 'table'"
                ).fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("report_layouts", tables)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_list_layout_persists_column_widths(self) -> None:
        self._create_players(2)
        saved = self.list_layout_service.save_columns(
            self.tournament_id,
            [{"key": "position", "width": 8}, {"key": "name", "width": 30}],
            "standings",
        )
        self.assertEqual(saved, [{"key": "position", "width": 8}, {"key": "name", "width": 30}])
        selected = self.list_layout_service.get_columns(self.tournament_id, "standings")["selected"]
        self.assertEqual({spec["key"]: spec["width"] for spec in selected}, {"position": 8, "name": 30})

    def test_standings_column_widths_applied_in_xlsx(self) -> None:
        from openpyxl import load_workbook

        self._create_players(4)
        self.list_layout_service.save_columns(
            self.tournament_id,
            [{"key": "position", "width": 7}, {"key": "name", "width": 40}],
            "standings",
        )
        path = Path(self.temp_dir.name) / "classificacao.xlsx"
        self.export_service.export_standings(self.tournament_id, path)

        worksheet = load_workbook(path).active
        self.assertEqual(worksheet.column_dimensions["A"].width, 7)
        self.assertEqual(worksheet.column_dimensions["B"].width, 40)

    # --- Fase G: sistema Scheveningen (E8) --------------------------------

    def test_scheveningen_requires_even_field(self) -> None:
        from src.services.pairing import scheveningen_pairings

        players = [{"id": index, "name": f"P{index}", "rating": 2000 - index} for index in range(5)]
        with self.assertRaisesRegex(AppError, "par"):
            scheveningen_pairings(players, 1, {})

    def test_scheveningen_pairs_every_cross_group_match(self) -> None:
        ids = self._create_players(6)  # A = 3 de maior rating, B = 3 de menor
        group_a, group_b = set(ids[:3]), set(ids[3:])
        self._set_individual_pairing_method("scheveningen")

        seen: set[tuple[int, int]] = set()
        for _round in range(3):  # 3 jogadores por grupo => 3 rodadas
            round_data = self.service.generate_next_round(self.tournament_id)
            round_id = int(round_data["id"])
            pairings = self.db.get_pairings_for_round(round_id)
            self.assertEqual(len(pairings), 3)
            for pairing in pairings:
                white, black = int(pairing["white_player_id"]), int(pairing["black_player_id"])
                self.assertFalse(pairing["is_bye"])
                # Cada confronto cruza os grupos (exatamente um lado em A).
                self.assertTrue((white in group_a) != (black in group_a))
                a_player = white if white in group_a else black
                b_player = black if white in group_a else white
                seen.add((a_player, b_player))
                self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
            self.service.close_round(self.tournament_id, round_id)

        # Todos os 9 cruzamentos A×B ocorreram exatamente uma vez.
        self.assertEqual(seen, {(a, b) for a in group_a for b in group_b})

    def test_scheveningen_uses_manual_groups(self) -> None:
        ids = self._create_players(4)  # ratings 2000,1950,1900,1850
        # Grupos manuais que NAO sao as metades por ranking: A = mais forte + mais fraco.
        self.db.set_player_scheveningen_group(ids[0], "A")
        self.db.set_player_scheveningen_group(ids[3], "A")
        self.db.set_player_scheveningen_group(ids[1], "B")
        self.db.set_player_scheveningen_group(ids[2], "B")
        self._set_individual_pairing_method("scheveningen")
        group_a, group_b = {ids[0], ids[3]}, {ids[1], ids[2]}

        seen: set[tuple[int, int]] = set()
        for _round in range(2):
            round_data = self.service.generate_next_round(self.tournament_id)
            round_id = int(round_data["id"])
            for pairing in self.db.get_pairings_for_round(round_id):
                white, black = int(pairing["white_player_id"]), int(pairing["black_player_id"])
                self.assertTrue((white in group_a) != (black in group_a))
                a_player = white if white in group_a else black
                b_player = black if white in group_a else white
                seen.add((a_player, b_player))
                self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
            self.service.close_round(self.tournament_id, round_id)

        self.assertEqual(seen, {(a, b) for a in group_a for b in group_b})

    def test_scheveningen_manual_groups_must_match_size(self) -> None:
        from src.services.pairing import scheveningen_pairings

        players = [
            {"id": 1, "name": "A1", "rating": 2000, "scheveningen_group": "A"},
            {"id": 2, "name": "A2", "rating": 1900, "scheveningen_group": "A"},
            {"id": 3, "name": "B1", "rating": 1800, "scheveningen_group": "B"},
        ]
        with self.assertRaisesRegex(AppError, "mesmo numero"):
            scheveningen_pairings(players, 1, {})

    def test_v40_database_adds_scheveningen_group_column(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v40_schev.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE players (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    tournament_id INTEGER,
                    name TEXT NOT NULL
                );
                PRAGMA user_version = 39;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(players)").fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("scheveningen_group", columns)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    # --- Fase H: importar torneio do Swiss-Manager via TRF (E9) -----------

    def test_trf_import_round_trips_players_and_header(self) -> None:
        self.db.create_player(
            self.tournament_id, name="Ana Silva", surname="Silva", given_name="Ana",
            rating=2000, international_rating=2000, federation_id="BRA",
            fide_id="123456", birth_date="2008-01-01", sex="w", title="WFM",
        )
        self.db.create_player(
            self.tournament_id, name="Bruno Souza", surname="Souza", given_name="Bruno",
            rating=1800, international_rating=1800, federation_id="BRA", fide_id="234567",
        )
        self.db.create_player(
            self.tournament_id, name="Carla Dias", surname="Dias", given_name="Carla",
            rating=1700, international_rating=1700, federation_id="ARG", fide_id="345678",
        )

        trf_path = Path(self.temp_dir.name) / "torneio.trf"
        self.export_service.export_chess_results_trf(self.tournament_id, trf_path)

        result = self.import_service.import_trf(trf_path)
        self.assertEqual(result["players_imported"], 3)

        imported = self.db.list_players(result["tournament_id"], active_only=False)
        by_fide = {str(player["fide_id"]): player for player in imported}
        self.assertEqual(set(by_fide), {"123456", "234567", "345678"})
        self.assertEqual(int(by_fide["123456"]["international_rating"]), 2000)
        self.assertEqual(by_fide["123456"]["federation_id"], "BRA")
        self.assertEqual(by_fide["123456"]["name"], "Ana Silva")
        self.assertEqual(by_fide["345678"]["federation_id"], "ARG")
        # Cabecalho: novo torneio herda nome e federacao do TRF.
        new_tournament = self.db.get_tournament(result["tournament_id"])
        self.assertEqual(new_tournament["name"], "Torneio teste")

    def test_trf_import_rejects_file_without_players(self) -> None:
        empty_path = Path(self.temp_dir.name) / "vazio.trf"
        empty_path.write_text("012 Torneio sem jogadores\r\n022 Cidade\r\n", encoding="utf-8")
        with self.assertRaisesRegex(AppError, "sem jogadores"):
            self.import_service.import_trf(empty_path)

    def test_trf_import_reconstructs_rounds_and_results(self) -> None:
        self._create_players(4)
        round_data = self.service.generate_next_round(self.tournament_id)
        round_id = int(round_data["id"])
        for pairing in self.db.get_pairings_for_round(round_id):
            self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, round_id)
        original_points = sorted(
            float(item["points"]) for item in self.service.standings(self.tournament_id)
        )

        trf_path = Path(self.temp_dir.name) / "jogado.trf"
        self.export_service.export_chess_results_trf(self.tournament_id, trf_path)
        result = self.import_service.import_trf(trf_path)

        self.assertEqual(result["rounds_imported"], 1)
        rounds = self.db.list_rounds(result["tournament_id"])
        self.assertEqual(len(rounds), 1)
        self.assertEqual(rounds[0]["status"], "closed")
        # A classificacao reconstruida bate com a original (2x1.0 e 2x0.0).
        imported_points = sorted(
            float(item["points"]) for item in self.service.standings(result["tournament_id"])
        )
        self.assertEqual(imported_points, original_points)

    # --- Fase I: mudar tipo / dividir torneio (E10) -----------------------

    def test_change_tournament_type_blocked_after_first_round(self) -> None:
        self._create_players(4)
        self.tournament_service.change_tournament_type(self.tournament_id, "round_robin")
        settings = self.db.get_tournament_settings(self.tournament_id)
        self.assertEqual(settings["pairing_method"], "round_robin")

        self.service.generate_next_round(self.tournament_id)
        with self.assertRaisesRegex(AppError, "primeira rodada"):
            self.tournament_service.change_tournament_type(self.tournament_id, "swiss")

    def test_split_tournament_partitions_by_ranking(self) -> None:
        self._create_players(6)  # ratings 2000..1750 decrescentes
        children = self.tournament_service.split_tournament(self.tournament_id, 2)
        self.assertEqual(len(children), 2)
        for child_id in children:
            self.assertEqual(
                self.db.get_tournament(child_id)["parent_tournament_id"], self.tournament_id
            )
        group_a = self.db.list_players(children[0], active_only=False)
        group_b = self.db.list_players(children[1], active_only=False)
        self.assertEqual(len(group_a), 3)
        self.assertEqual(len(group_b), 3)
        # Grupo A reune os tres maiores ratings (divisao por ranking).
        self.assertGreaterEqual(
            min(int(player["rating"]) for player in group_a),
            max(int(player["rating"]) for player in group_b),
        )

    def test_split_tournament_requires_clean_state(self) -> None:
        self._create_players(4)
        self.service.generate_next_round(self.tournament_id)
        with self.assertRaisesRegex(AppError, "antes de gerar"):
            self.tournament_service.split_tournament(self.tournament_id, 2)

    def test_split_tournament_rejects_insufficient_players(self) -> None:
        self._create_players(2)
        with self.assertRaisesRegex(AppError, "insuficientes"):
            self.tournament_service.split_tournament(self.tournament_id, 3)

    def test_save_profile_blocks_pairing_change_after_round(self) -> None:
        self._create_players(4)
        self.service.generate_next_round(self.tournament_id)
        with self.assertRaisesRegex(AppError, "primeira rodada"):
            self.tournament_service.save_profile(
                self.tournament_id,
                {
                    "name": "Torneio teste",
                    "scope": "standalone",
                    "competition_type": "individual",
                    "rounds_count": "5",
                    "bye_points": "1",
                },
                {"pairing_method": "round_robin"},
                [],
            )

    def test_v39_database_adds_parent_tournament_column(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v39_split.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE tournaments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    club_id INTEGER,
                    class_id INTEGER,
                    created_at TEXT
                );
                PRAGMA user_version = 38;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(tournaments)").fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("parent_tournament_id", columns)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_trf25_emits_240_for_requested_bye(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        player_ids = self._create_players(5)  # ratings decrescentes → rank = ordem
        self.db.add_requested_bye(self.tournament_id, player_ids[4], 1, "H")
        self.service.generate_next_round(self.tournament_id)

        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "bye240.trf"
        exporter.export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()

        bye_lines = [line for line in lines if line.startswith("240 ")]
        self.assertEqual(len(bye_lines), 1)
        line = bye_lines[0]
        self.assertEqual(line[4:5], "H")          # tipo (col 5)
        self.assertEqual(line[6:9].strip(), "1")  # rodada (col 7-9)
        self.assertEqual(line[10:14].strip(), "5")  # start-rank do jogador (col 11-14)
        # A célula do 001 mostra o mesmo bye solicitado, nunca divergindo do 240.
        self.assertTrue(any("0000 - H" in line for line in lines))

    def test_classic_acceleration_bonus_boundaries(self) -> None:
        from src.services.pairing import classic_acceleration_bonus

        # 8 jogadores: metade superior = ranks 1..4.
        self.assertEqual(classic_acceleration_bonus(1, 8, 1), 1.0)
        self.assertEqual(classic_acceleration_bonus(4, 8, 2), 1.0)
        self.assertEqual(classic_acceleration_bonus(5, 8, 1), 0.0)
        # Fora das rodadas 1 e 2 não há bônus, mesmo para o topo.
        self.assertEqual(classic_acceleration_bonus(1, 8, 3), 0.0)
        # 7 jogadores: piso de N/2 = 3 (ranks 1..3).
        self.assertEqual(classic_acceleration_bonus(3, 7, 1), 1.0)
        self.assertEqual(classic_acceleration_bonus(4, 7, 1), 0.0)
        # start_rank inválido não recebe bônus.
        self.assertEqual(classic_acceleration_bonus(0, 8, 1), 0.0)

    def test_accelerated_standings_adds_bonus_only_to_top_half(self) -> None:
        from src.services.pairing import accelerated_standings

        standings = {pid: {"points": 0.0} for pid in range(1, 5)}
        seeding = [1, 2, 3, 4]

        # Método não acelerado devolve o standings original (sem cópia).
        self.assertIs(accelerated_standings(standings, seeding, 2, "none"), standings)

        effective = accelerated_standings(standings, seeding, 2, "accelerated")
        self.assertEqual(effective[1]["points"], 1.0)
        self.assertEqual(effective[2]["points"], 1.0)
        self.assertEqual(effective[3]["points"], 0.0)
        self.assertEqual(effective[4]["points"], 0.0)
        # O standings real permanece intacto.
        self.assertEqual(standings[1]["points"], 0.0)
        # Rodada 3 não recebe bônus.
        round3 = accelerated_standings(standings, seeding, 3, "accelerated")
        self.assertEqual(round3[1]["points"], 0.0)

    def test_classic_acceleration_reshapes_round_two_score_groups(self) -> None:
        self._create_players(8)
        players = self.db.list_players(self.tournament_id, active_only=True)

        def pair_set(pairings: list[dict[str, Any]]) -> set[frozenset[int]]:
            return {
                frozenset({int(p["white_player_id"]), int(p["black_player_id"])})
                for p in pairings
                if p.get("black_player_id") is not None
            }

        seeding = [
            int(p["id"])
            for p in sorted(players, key=lambda p: -int(p.get("rating") or 0))
        ]

        # Sem aceleração: todos no mesmo grupo (0 pontos) → topo vs base.
        self.db.save_tournament_settings(self.tournament_id, {"acceleration_method": "none"})
        plain = pair_set(self.service._swiss_pairings(self.tournament_id, players, 2))
        self.assertIn(frozenset({seeding[0], seeding[4]}), plain)

        # Com aceleração clássica: top-metade (seeds 1..4) ganha +1 fictício,
        # formando dois grupos de pontuação. Seed 1 pareia dentro do topo (seed 3),
        # não cruza para a base (seed 5).
        self.db.save_tournament_settings(
            self.tournament_id, {"acceleration_method": "accelerated"}
        )
        accel = pair_set(self.service._swiss_pairings(self.tournament_id, players, 2))
        self.assertIn(frozenset({seeding[0], seeding[2]}), accel)
        self.assertNotIn(frozenset({seeding[0], seeding[4]}), accel)

    def test_custom_acceleration_applies_configured_params(self) -> None:
        from src.services.pairing import accelerated_standings, acceleration_spec

        spec = acceleration_spec("custom:rounds=3;bonus=2.0;upper=0.25")
        self.assertEqual(spec["scheme"], "custom")
        self.assertEqual(spec["round_count"], 3)
        self.assertEqual(spec["bonus"], 2.0)
        self.assertEqual(spec["upper_fraction"], 0.25)

        standings = {pid: {"points": 0.0} for pid in range(1, 9)}
        seeding = list(range(1, 9))
        # upper=0.25 de 8 → só ranks 1..2; rodada 3 ainda dentro de round_count=3.
        effective = accelerated_standings(
            standings, seeding, 3, "custom:rounds=3;bonus=2.0;upper=0.25"
        )
        self.assertEqual(effective[1]["points"], 2.0)
        self.assertEqual(effective[2]["points"], 2.0)
        self.assertEqual(effective[3]["points"], 0.0)
        # Rodada 4 fora da janela.
        round4 = accelerated_standings(
            standings, seeding, 4, "custom:rounds=3;bonus=2.0;upper=0.25"
        )
        self.assertEqual(round4[1]["points"], 0.0)

    def test_trf25_emits_250_for_custom_acceleration(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        self._create_players(8)
        self.db.save_tournament_settings(
            self.tournament_id,
            {"acceleration_method": "custom:rounds=3;bonus=2.0;upper=0.25"},
        )
        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "custom.trf"
        exporter.export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()
        accel_lines = [line for line in lines if line.startswith("250 ")]
        self.assertEqual(len(accel_lines), 1)
        line = accel_lines[0]
        self.assertEqual(float(line[9:13]), 2.0)  # bônus custom
        self.assertEqual(line[14:17].strip(), "1")  # primeira rodada
        self.assertEqual(line[18:21].strip(), "3")  # round_count=3
        self.assertEqual(line[22:26].strip(), "1")  # rank 1
        self.assertEqual(line[27:31].strip(), "2")  # upper 0.25 de 8 = 2

    def test_custom_acceleration_ignores_nonpositive_params(self) -> None:
        from src.services.federation_exporters import TRF25Exporter
        from src.services.pairing import acceleration_bonus, acceleration_spec

        # Bônus zero: sem efeito no motor.
        spec_zero = acceleration_spec("custom:rounds=2;bonus=0;upper=0.5")
        self.assertEqual(acceleration_bonus(1, 8, 1, spec_zero), 0.0)
        # Bônus negativo nunca "desacelera".
        spec_neg = acceleration_spec("custom:rounds=2;bonus=-1;upper=0.5")
        self.assertEqual(acceleration_bonus(1, 8, 1, spec_neg), 0.0)
        # Zero rodadas: sem efeito.
        spec_no_round = acceleration_spec("custom:rounds=0;bonus=1;upper=0.5")
        self.assertEqual(acceleration_bonus(1, 8, 1, spec_no_round), 0.0)

        self._create_players(8)
        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "noop.trf"
        for method in (
            "custom:rounds=2;bonus=0;upper=0.5",
            "custom:rounds=0;bonus=1;upper=0.5",
            "custom:rounds=2;bonus=1;upper=0",
        ):
            self.db.save_tournament_settings(self.tournament_id, {"acceleration_method": method})
            exporter.export(self.tournament_id, output_path)
            lines = output_path.read_text(encoding="utf-8").splitlines()
            self.assertFalse(
                [line for line in lines if line.startswith("250 ")],
                msg=f"250 indevido para {method}",
            )

    def test_baku_does_not_apply_bonus_or_emit_250(self) -> None:
        from src.services.federation_exporters import TRF25Exporter
        from src.services.federation_exporters.trf25 import BAKU_NOT_IMPLEMENTED
        from src.services.pairing import accelerated_standings

        # Motor: Baku não soma bônus (sem fórmula oficial).
        standings = {pid: {"points": 0.0} for pid in range(1, 5)}
        self.assertIs(accelerated_standings(standings, [1, 2, 3, 4], 2, "baku"), standings)

        self._create_players(8)
        self.db.save_tournament_settings(self.tournament_id, {"acceleration_method": "baku"})
        exporter = TRF25Exporter(self.export_service)
        output_path = Path(self.temp_dir.name) / "baku.trf"
        warnings = exporter.export(self.tournament_id, output_path)
        lines = output_path.read_text(encoding="utf-8").splitlines()
        # Nenhum 250 e nenhum sufixo _BAKU no registro 192.
        self.assertFalse([line for line in lines if line.startswith("250 ")])
        type_line = next(line for line in lines if line.startswith("192 "))
        self.assertNotIn("_BAKU", type_line)
        # Aviso ao árbitro de que Baku não está implementado.
        self.assertIn(BAKU_NOT_IMPLEMENTED, warnings)

    def test_trf25_dutch_192_is_dated_by_tournament_date(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        exporter = TRF25Exporter(self.export_service)
        settings = {"pairing_method": "swiss"}
        # Antes do corte (2025-07-01) → regras 2017.
        self.assertEqual(
            exporter._type_code_192({"start_date": "2025-06-30"}, settings),
            "FIDE_DUTCH_2017",
        )
        # No corte ou depois → regras 2025.
        self.assertEqual(
            exporter._type_code_192({"start_date": "2025-07-01"}, settings),
            "FIDE_DUTCH_2025",
        )
        # Sem start_date, usa o end_date como fallback.
        self.assertEqual(
            exporter._type_code_192({"end_date": "2024-01-10"}, settings),
            "FIDE_DUTCH_2017",
        )
        # Sem data parseável → FIDE_DUTCH puro (default-por-data, nunca chuta versão).
        self.assertEqual(exporter._type_code_192({}, settings), "FIDE_DUTCH")
        self.assertEqual(
            exporter._type_code_192({"start_date": "data invalida"}, settings),
            "FIDE_DUTCH",
        )

    def test_trf25_team_192_score_code_reflects_standing_criteria(self) -> None:
        from src.services.federation_exporters import TRF25Exporter

        exporter = TRF25Exporter(self.export_service)
        team = {"competition_type": "team"}

        def code(primary: str, secondary: str) -> str:
            return exporter._type_code_192(
                team,
                {"team_standing_primary": primary, "team_standing_secondary": secondary},
            )

        # Padrao: match points primario, game points secundario.
        self.assertEqual(code("match_points", "game_points"), "FIDE_TEAM_TYPEA_MP_GP")
        # Ordem invertida.
        self.assertEqual(code("game_points", "match_points"), "FIDE_TEAM_TYPEA_GP_MP")
        # 'wins' e desempate, nao codigo de pontuacao: e ignorado.
        self.assertEqual(code("match_points", "wins"), "FIDE_TEAM_TYPEA_MP")
        self.assertEqual(code("wins", "game_points"), "FIDE_TEAM_TYPEA_GP")
        # Criterios repetidos deduplicam.
        self.assertEqual(code("match_points", "match_points"), "FIDE_TEAM_TYPEA_MP")
        # Sem MP/GP configurado, cai no padrao FIDE MP_GP.
        self.assertEqual(code("wins", "wins"), "FIDE_TEAM_TYPEA_MP_GP")

    def test_validate_chess_results_trf16_reports_special_result_statuses(self) -> None:
        tournament_id = self.db.create_tournament(
            "Aberto Pendencias",
            location="Sao Paulo",
            rounds_count=1,
            time_control="15 min",
            start_date="2026-05-24",
            end_date="2026-05-24",
        )
        self.db.save_tournament_settings(
            tournament_id,
            {
                "fide_event_id": "12345",
                "federation": "BRA",
                "chief_arbiter": "Arbitro Chefe",
                "tournament_profile": "fide",
            },
        )
        self.db.save_round_schedule(
            tournament_id,
            [{"round_number": 1, "date": "2026-05-24", "time": "10:00"}],
        )

        def player(index: int) -> int:
            return self.db.create_player(
                tournament_id,
                name=f"Jogador {index}",
                federation_id="BRA",
                fide_id=str(1000 + index),
                rating=1800 + index,
                international_rating=1800 + index,
                birth_date="2000-01-02",
            )

        players = [player(index) for index in range(1, 7)]
        round_id = self.db.create_round_with_pairings(
            tournament_id,
            1,
            [
                {
                    "board_number": 1,
                    "white_player_id": players[0],
                    "black_player_id": players[1],
                    "result": "1F-0F",
                    "is_bye": 0,
                },
                {
                    "board_number": 2,
                    "white_player_id": players[2],
                    "black_player_id": players[3],
                    "result": "0F-0F",
                    "is_bye": 0,
                },
                {
                    "board_number": 3,
                    "white_player_id": players[4],
                    "black_player_id": None,
                    "result": "BYE",
                    "is_bye": 1,
                },
            ],
        )
        self.service.close_round(tournament_id, round_id)

        warnings = self.export_service.validate_chess_results_trf(tournament_id)

        self.assertIn("TRF16 contem 1 bye(s).", warnings)
        self.assertIn("TRF16 contem 1 resultado(s) por WO.", warnings)
        self.assertIn("TRF16 contem 1 dupla(s) ausencia(s).", warnings)
        self.assertIn("TRF16 contem 1 jogador(es) nao emparceirado(s) em rodadas geradas.", warnings)

    def test_export_chess_results_trf16_validation_report_and_encoding(self) -> None:
        tournament_id = self.db.create_tournament(
            "Aberto Acentuacao",
            location="Sao Paulo",
            rounds_count=1,
            time_control="90 min",
            start_date="2026-05-24",
            end_date="2026-05-24",
        )
        self.db.save_tournament_settings(
            tournament_id,
            {
                "fide_event_id": "12345",
                "federation": "BRA",
                "chief_arbiter": "Arbitro Chefe",
                "tournament_profile": "fide",
            },
        )
        first_player = self.db.create_player(
            tournament_id,
            name="Ávila, José",
            surname="Ávila",
            given_name="José",
            federation_id="BRA",
            fide_id="1234",
            rating=2100,
            international_rating=2100,
            birth_date="2000-01-02",
        )
        second_player = self.db.create_player(
            tournament_id,
            name="Núñez, Maria",
            surname="Núñez",
            given_name="Maria",
            federation_id="BRA",
            fide_id="1235",
            rating=2000,
            international_rating=2000,
            birth_date="2000-01-02",
        )
        round_id = self.db.create_round_with_pairings(
            tournament_id,
            1,
            [
                {
                    "board_number": 1,
                    "white_player_id": first_player,
                    "black_player_id": second_player,
                    "result": "1-0",
                }
            ],
        )
        self.service.close_round(tournament_id, round_id)

        trf_path = Path(self.temp_dir.name) / "acentuacao.trf"
        report_path = Path(self.temp_dir.name) / "pendencias_trf.csv"
        self.export_service.export_chess_results_trf(tournament_id, trf_path)
        self.export_service.export_chess_results_trf_validation_report(tournament_id, report_path)

        trf_content = trf_path.read_text(encoding="utf-8")
        report_content = report_path.read_text(encoding="utf-8")
        self.assertIn("Avila, Jose", trf_content)
        self.assertIn("Nunez, Maria", trf_content)
        self.assertIn("Partidas jogadas", report_content)
        self.assertIn("1", report_content)

    def test_rating_fee_report_counts_bases_rated_unrated_and_exports(self) -> None:
        self.db.save_tournament_settings(
            self.tournament_id,
            {
                "rating_fee_fide": "2",
                "rating_fee_cbx": "3",
                "rating_fee_lbx": "4",
            },
        )
        self.db.create_player(
            self.tournament_id,
            name="Multibase rated",
            fide_id="1001",
            cbx_id="2001",
            lbx_id="3001",
            rating=1800,
            national_rating=1900,
            international_rating=2000,
        )
        self.db.create_player(
            self.tournament_id,
            name="FIDE sem rating",
            fide_id="1002",
            cbx_id="2002",
            national_rating=1600,
        )
        self.db.create_player(self.tournament_id, name="LBX sem rating", lbx_id="3002")
        self.db.create_player(self.tournament_id, name="Rated sem ID", rating=1400)

        summary = self.export_service.rating_fee_summary(self.tournament_id)
        bases = {item["base"]: item for item in summary["bases"]}

        self.assertEqual(4, summary["players_count"])
        self.assertEqual(3, summary["players_with_base"])
        self.assertEqual(1, summary["players_without_base"])
        self.assertEqual(3, summary["rated_players"])
        self.assertEqual(1, summary["unrated_players"])
        self.assertEqual({"identified": 2, "rated": 1, "unrated": 1}, {
            key: bases["FIDE"][key] for key in ("identified", "rated", "unrated")
        })
        self.assertEqual(6.0, bases["CBX"]["subtotal"])
        self.assertEqual(8.0, bases["LBX"]["subtotal"])
        self.assertEqual(18.0, summary["total"])

        xlsx_path = Path(self.temp_dir.name) / "taxas_rating.xlsx"
        pdf_path = Path(self.temp_dir.name) / "taxas_rating.pdf"
        self.export_service.export_rating_fee_report(self.tournament_id, xlsx_path)
        self.export_service.export_rating_fee_report(self.tournament_id, pdf_path)

        from openpyxl import load_workbook
        from pypdf import PdfReader

        workbook = load_workbook(xlsx_path, read_only=True)
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)
        self.assertEqual(["Resumo de taxas de rating", "Taxas por base", "Inscritos por base"], workbook.sheetnames)
        workbook.close()
        self.assertIn("Total das taxas", pdf_text)
        self.assertIn("R$ 18,00", pdf_text)
        with self.assertRaisesRegex(AppError, "XLSX ou PDF"):
            self.export_service.export_rating_fee_report(
                self.tournament_id,
                Path(self.temp_dir.name) / "taxas_rating.csv",
            )

    def test_export_club_portal_creates_static_html_package(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Escola Portal",
                "kind": "school",
                "city": "Curitiba",
                "active": 1,
            },
            club_id=None,
        )
        class_id = self.club_service.save_class(
            {
                "club_id": club_id,
                "name": "Turma Portal",
                "teacher": "Professora Portal",
                "weekday": "Sabado",
                "time": "09:00",
                "location": "Sala 1",
                "active": 1,
            }
        )
        self.member_service.create_member(
            {
                "name": "Aluno Portal",
                "club_id": club_id,
                "class_id": class_id,
                "rating": "1500",
                "category": "Sub-12",
                "member_type": "aluno",
                "status": "active",
            }
        )
        self.training_service.save_session(
            {
                "club_id": club_id,
                "class_id": class_id,
                "title": "Aula de estrategia",
                "session_type": "aula",
                "session_date": "2999-08-05",
                "start_time": "09:00",
                "instructor": "Professora Portal",
                "location": "Sala 1",
                "objective": "Plano de meio-jogo",
                "homework": "Resolver dois diagramas.",
                "status": "planned",
            }
        )
        self.event_service.save_event(
            {
                "club_id": club_id,
                "title": "Festival Portal",
                "event_type": "social",
                "event_date": "2999-08-10",
                "start_time": "10:00",
                "location": "Salao",
                "status": "confirmed",
                "notes": "Comunicado aos responsaveis.",
            }
        )
        self.db.create_tournament(
            "Torneio Portal",
            club_id=club_id,
            class_id=class_id,
            start_date="2999-08-20",
            location="Salao",
        )

        output_dir = Path(self.temp_dir.name) / "portal"
        index_path = self.export_service.export_club_portal(output_dir, club_id=club_id, class_id=class_id)
        html = index_path.read_text(encoding="utf-8")

        self.assertTrue(index_path.exists())
        self.assertTrue((output_dir / "styles.css").exists())
        self.assertIn("Turma Portal", html)
        self.assertIn("Comunicados", html)
        self.assertIn("Festival Portal", html)
        self.assertIn("Aula de estrategia", html)
        self.assertIn("Resolver dois diagramas.", html)
        self.assertIn("Aluno Portal", html)
        self.assertIn("Ranking interno", html)
        self.assertIn("Torneio Portal", html)

    def test_certificate_service_exports_tournament_pdf(self) -> None:
        player_a = self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube A",
            category="Sub-18",
        )
        player_c = self.db.create_player(
            self.tournament_id,
            name="Carla",
            surname="Lima",
            rating=1700,
            club="Clube B",
            category="Absoluto",
        )
        self.db.create_player(
            self.tournament_id,
            name="Diego",
            surname="Rocha",
            rating=1600,
            club="Clube B",
            category="Absoluto",
        )

        category_recipients = self.certificate_service.tournament_recipients(
            self.tournament_id,
            certificate_type="category_award",
            top_n=1,
        )
        selected_path = Path(self.temp_dir.name) / "diplomas.pdf"
        result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            selected_path,
            certificate_type="participation",
            player_ids=[player_a, player_c],
        )

        self.assertEqual(result["exported"], 2)
        self.assertEqual(len(result["verification_codes"]), 2)
        issuance = self.certificate_service.verify_issuance(result["verification_codes"][0])
        self.assertEqual(issuance["context_type"], "tournament")
        self.assertEqual(issuance["source_title"], "Torneio teste")
        self.assertTrue(selected_path.exists())
        self.assertEqual(selected_path.read_bytes()[:4], b"%PDF")
        self.assertEqual({item["category"] for item in category_recipients}, {"Sub-18", "Absoluto"})
        self.assertEqual(len(category_recipients), 2)

    def test_certificate_export_without_issuance_record_does_not_leave_pdf(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube A",
            category="Sub-18",
        )
        output_path = Path(self.temp_dir.name) / "falha_registro.pdf"

        with mock.patch.object(
            self.db,
            "create_certificate_issuances",
            side_effect=sqlite3.OperationalError("registro falhou"),
        ):
            with self.assertRaises(sqlite3.OperationalError):
                self.certificate_service.export_tournament_certificates(
                    self.tournament_id,
                    output_path,
                    certificate_type="participation",
                )

        self.assertFalse(output_path.exists())
        self.assertEqual([], list(output_path.parent.glob("*.tmp.pdf")))

    def test_certificate_export_replace_failure_removes_issuance_records_and_temp_pdf(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        output_path = Path(self.temp_dir.name) / "falha_substituir.pdf"

        with mock.patch.object(Path, "replace", side_effect=OSError("replace falhou")):
            with self.assertRaises(OSError):
                self.certificate_service.export_tournament_certificates(
                    self.tournament_id,
                    output_path,
                    certificate_type="participation",
                )
        with self.db.connect() as connection:
            issuance_count = connection.execute("SELECT COUNT(*) FROM certificate_issuances").fetchone()[0]

        self.assertEqual(issuance_count, 0)
        self.assertFalse(output_path.exists())
        self.assertEqual([], list(output_path.parent.glob("*.tmp.pdf")))

    def test_certificate_templates_can_be_saved_previewed_and_used(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube A",
            category="Sub-18",
        )
        logo_path = Path(self.temp_dir.name) / "logo.png"
        logo_path.write_bytes(
            base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
            )
        )
        background_path = Path(self.temp_dir.name) / "fundo.png"
        background_path.write_bytes(logo_path.read_bytes())
        secondary_logo_path = Path(self.temp_dir.name) / "logo_secundario.png"
        secondary_logo_path.write_bytes(logo_path.read_bytes())
        template_id = self.certificate_service.create_template(
            {
                "name": "Modelo personalizado",
                "certificate_type": "overall_award",
                "title_template": "Diploma {torneio}",
                "body_template": "{nome} ficou em {posicao} com {pontos} ponto(s).",
                "footer_template": "{local} {periodo}",
                "orientation": "portrait",
                "signature_left": "Direcao",
                "signature_right": "Arbitro",
                "logo_path": str(logo_path),
                "background_image_path": str(background_path),
                "background_opacity": "35",
                "secondary_logo_path": str(secondary_logo_path),
                "primary_color": "#0F766E",
                "accent_color": "#F59E0B",
                "title_font_size": "30",
                "body_font_size": "16",
                "footer_font_size": "9",
            }
        )

        self.certificate_service.update_template(
            template_id,
            {
                "name": "Modelo personalizado",
                "certificate_type": "overall_award",
                "title_template": "Diploma especial {torneio}",
                "body_template": "{nome} ficou em {posicao} com {pontos} ponto(s).",
                "footer_template": "{local} {periodo}",
                "orientation": "portrait",
                "signature_left": "Direcao",
                "signature_right": "Arbitro",
                "logo_path": str(logo_path),
                "background_image_path": str(background_path),
                "background_opacity": "0.42",
                "secondary_logo_path": str(secondary_logo_path),
                "primary_color": "#0F766E",
                "accent_color": "#F59E0B",
                "title_font_size": "30",
                "body_font_size": "16",
                "footer_font_size": "9",
            },
        )
        preview = self.certificate_service.preview_template(
            self.tournament_id,
            template_id,
            top_n=1,
        )
        output_path = Path(self.temp_dir.name) / "modelo_personalizado.pdf"
        result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            output_path,
            template_id=template_id,
            top_n=1,
        )

        template = self.db.get_certificate_template(template_id)
        self.assertEqual(template["orientation"], "portrait")
        self.assertEqual(template["primary_color"], "#0F766E")
        self.assertEqual(template["accent_color"], "#F59E0B")
        self.assertEqual(template["background_image_path"], str(background_path))
        self.assertEqual(template["background_opacity"], 0.42)
        self.assertEqual(template["secondary_logo_path"], str(secondary_logo_path))
        self.assertEqual(template["title_font_size"], 30)
        self.assertEqual(result["exported"], 1)
        self.assertIn("Diploma especial Torneio teste", preview["title"])
        self.assertIn("Ana Silva", preview["body"])
        self.assertTrue(output_path.exists())
        self.assertEqual(output_path.read_bytes()[:4], b"%PDF")

    def test_default_certificate_background_assets_and_templates_export(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube B",
            category="Absoluto",
        )
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Fundo",
                "rating": "1700",
                "category": "Sub-18",
                "member_type": "aluno",
                "status": "active",
            }
        )

        for _label, relative_path, _opacity in database_module.CERTIFICATE_BACKGROUND_PRESETS:
            path = database_module.BASE_DIR / relative_path
            self.assertTrue(path.exists(), relative_path)
            self.assertEqual(path.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")

        templates_by_name = {
            str(template["name"]): template
            for template in self.certificate_service.list_templates(active_only=True)
        }
        expected_templates = {
            "Xadrez classico - Participacao",
            "Xadrez escolar - Membro aluno",
            "Xadrez premium - Premiacao geral",
            "Tabuleiro sutil - Premiacao por categoria",
            "Pecas marca d'agua - Ranking interno",
        }
        self.assertTrue(expected_templates.issubset(templates_by_name))

        participation_path = Path(self.temp_dir.name) / "participacao_fundo.pdf"
        overall_path = Path(self.temp_dir.name) / "premiacao_fundo.pdf"
        category_path = Path(self.temp_dir.name) / "categoria_fundo.pdf"
        member_path = Path(self.temp_dir.name) / "membro_fundo.pdf"
        ranking_path = Path(self.temp_dir.name) / "ranking_fundo.pdf"

        participation_result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            participation_path,
            template_id=int(templates_by_name["Xadrez classico - Participacao"]["id"]),
        )
        overall_result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            overall_path,
            template_id=int(templates_by_name["Xadrez premium - Premiacao geral"]["id"]),
            top_n=1,
        )
        category_result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            category_path,
            template_id=int(templates_by_name["Tabuleiro sutil - Premiacao por categoria"]["id"]),
            top_n=1,
        )
        member_result = self.certificate_service.export_member_certificates(
            member_path,
            template_id=int(templates_by_name["Xadrez escolar - Membro aluno"]["id"]),
            member_ids=[member_id],
        )
        ranking_result = self.certificate_service.export_ranking_certificates(
            ranking_path,
            template_id=int(templates_by_name["Pecas marca d'agua - Ranking interno"]["id"]),
            top_n=1,
        )

        self.assertEqual(participation_result["exported"], 2)
        self.assertEqual(overall_result["exported"], 1)
        self.assertEqual(category_result["exported"], 2)
        self.assertEqual(member_result["exported"], 1)
        self.assertEqual(ranking_result["exported"], 1)
        for path in (participation_path, overall_path, category_path, member_path, ranking_path):
            self.assertTrue(path.exists())
            self.assertEqual(path.read_bytes()[:4], b"%PDF")

    def test_default_certificate_logo_assets_export_as_primary_and_secondary(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        logo_dir = database_module.BASE_DIR / "assets" / "certificates" / "logos"
        primary_logo = logo_dir / "albericus_knight.png"
        secondary_logo = logo_dir / "clube_rook.png"
        expected_logos = {
            "albericus_knight.png",
            "clube_rook.png",
            "escola_pawn.png",
            "torneio_trophy.png",
            "selo_queen.png",
        }

        for filename in expected_logos:
            path = logo_dir / filename
            self.assertTrue(path.exists(), filename)
            content = path.read_bytes()
            self.assertEqual(content[:8], b"\x89PNG\r\n\x1a\n")
            self.assertEqual(content[25], 6)

        template_id = self.certificate_service.create_template(
            {
                "name": "Modelo com logos padrao",
                "certificate_type": "participation",
                "title_template": "Certificado",
                "body_template": "Certificamos que {nome} participou do torneio {torneio}.",
                "footer_template": "{local} - {periodo}",
                "orientation": "landscape",
                "signature_left": "Direcao",
                "signature_right": "Arbitragem",
                "logo_path": str(primary_logo),
                "secondary_logo_path": str(secondary_logo),
                "primary_color": "#1E3A8A",
                "accent_color": "#C8A24A",
                "title_font_size": "30",
                "body_font_size": "16",
                "footer_font_size": "9",
            }
        )
        output_path = Path(self.temp_dir.name) / "logos_padrao.pdf"
        result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            output_path,
            template_id=template_id,
        )

        self.assertEqual(result["exported"], 1)
        self.assertTrue(output_path.exists())
        self.assertEqual(output_path.read_bytes()[:4], b"%PDF")

    def test_certificate_service_exports_member_training_event_and_ranking_pdfs(self) -> None:
        club_id = self.club_service.save_profile(
            {
                "name": "Escola Diplomas",
                "kind": "school",
                "active": 1,
            },
            club_id=None,
        )
        class_id = self.club_service.save_class(
            {
                "club_id": club_id,
                "name": "Turma Diplomas",
                "teacher": "Professora",
                "active": 1,
            }
        )
        member_id = self.member_service.create_member(
            {
                "name": "Aluno Diploma",
                "club_id": club_id,
                "class_id": class_id,
                "category": "Sub-14",
                "rating": "1500",
                "member_type": "aluno",
                "status": "active",
            }
        )
        absent_member_id = self.member_service.create_member(
            {
                "name": "Aluno Ausente",
                "club_id": club_id,
                "class_id": class_id,
                "category": "Sub-14",
                "rating": "1200",
                "member_type": "aluno",
                "status": "active",
            }
        )
        ranking_member_id = self.member_service.create_member(
            {
                "name": "Aluno Ranking",
                "club_id": club_id,
                "category": "Sub-14",
                "rating": "2200",
                "member_type": "aluno",
                "status": "active",
            }
        )
        session_id = self.training_service.save_session(
            {
                "club_id": club_id,
                "class_id": class_id,
                "title": "Aula de calculo",
                "session_type": "aula",
                "session_date": "2026-05-20",
                "instructor": "Professora",
                "location": "Sala 2",
                "status": "done",
            }
        )
        self.training_service.record_attendance(
            session_id,
            [
                {"member_id": member_id, "status": "present"},
                {"member_id": absent_member_id, "status": "absent"},
            ],
        )
        event_id = self.event_service.save_event(
            {
                "club_id": club_id,
                "title": "Festival Escolar",
                "event_type": "social",
                "event_date": "2026-05-21",
                "location": "Auditorio",
                "status": "done",
            }
        )

        member_path = Path(self.temp_dir.name) / "membro.pdf"
        training_path = Path(self.temp_dir.name) / "aula.pdf"
        event_path = Path(self.temp_dir.name) / "evento.pdf"
        ranking_path = Path(self.temp_dir.name) / "ranking.pdf"
        member_result = self.certificate_service.export_member_certificates(
            member_path,
            member_ids=[member_id],
        )
        training_result = self.certificate_service.export_training_certificates(
            session_id,
            training_path,
            present_only=False,
            member_ids=[member_id, absent_member_id],
        )
        event_result = self.certificate_service.export_event_certificates(
            event_id,
            event_path,
            member_ids=[member_id],
        )
        ranking_result = self.certificate_service.export_ranking_certificates(
            ranking_path,
            category="Sub-14",
            top_n=1,
        )
        member_recipient = self.certificate_service.member_recipients(member_ids=[member_id])[0]
        training_recipients = self.certificate_service.training_recipients(
            session_id,
            present_only=False,
        )
        event_recipient = self.certificate_service.event_recipients(event_id, member_ids=[member_id])[0]
        ranking_recipient = self.certificate_service.ranking_recipients(category="Sub-14", top_n=1)[0]

        self.assertEqual(member_result["exported"], 1)
        self.assertEqual(training_result["exported"], 2)
        self.assertEqual(event_result["exported"], 1)
        self.assertEqual(ranking_result["exported"], 1)
        all_codes = (
            member_result["verification_codes"]
            + training_result["verification_codes"]
            + event_result["verification_codes"]
            + ranking_result["verification_codes"]
        )
        issued_records = self.certificate_service.list_issuances(limit=10)
        verified_member = self.certificate_service.verify_issuance(member_result["verification_codes"][0])
        self.assertEqual(len(all_codes), 5)
        self.assertEqual(len(set(all_codes)), 5)
        self.assertEqual(len(issued_records), 5)
        self.assertEqual(verified_member["context_type"], "members")
        self.assertEqual(verified_member["recipient_name"], "Aluno Diploma")
        self.assertEqual(member_recipient["class_name"], "Turma Diplomas")
        self.assertEqual({item["status"] for item in training_recipients}, {"Presente", "Falta"})
        self.assertEqual(event_recipient["event"], "Festival Escolar")
        self.assertEqual(ranking_recipient["member_id"], ranking_member_id)
        for path in (member_path, training_path, event_path, ranking_path):
            self.assertTrue(path.exists())
            self.assertEqual(path.read_bytes()[:4], b"%PDF")

    def test_certificate_service_exports_local_verification_site(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana",
            surname="Silva",
            rating=1900,
            club="Clube A",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno",
            surname="Souza",
            rating=1800,
            club="Clube B",
            category="Absoluto",
        )
        pdf_path = Path(self.temp_dir.name) / "diplomas.pdf"
        result = self.certificate_service.export_tournament_certificates(
            self.tournament_id,
            pdf_path,
        )
        revoked_issuance = self.certificate_service.verify_issuance(result["verification_codes"][0])
        self.certificate_service.revoke_issuance(int(revoked_issuance["id"]), notes="Reemitido")

        site_path = Path(self.temp_dir.name) / "verificador"
        exported_path = self.certificate_service.export_verification_site(site_path)
        html = exported_path.read_text(encoding="utf-8")

        self.assertEqual(exported_path, site_path.with_suffix(".html"))
        self.assertIn("Verificador de diplomas Albericus", html)
        self.assertIn(result["verification_codes"][0], html)
        self.assertIn(result["verification_codes"][1], html)
        self.assertIn("Ana Silva", html)
        self.assertIn("Bruno Souza", html)
        self.assertIn("Revogado", html)
        self.assertIn("Ativo", html)
        self.assertIn("Torneio teste", html)
        self.assertNotIn(str(pdf_path), html)

    def test_team_tournament_exports_reports_and_site(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(round_data["id"])[0]
        boards = self.db.list_team_boards(int(match["id"]))
        self.service.update_result(tournament_id, int(boards[0]["id"]), "1-0")
        self.service.update_result(tournament_id, int(boards[1]["id"]), "0-1")
        self.service.close_round(tournament_id, round_data["id"])

        standings_path = Path(self.temp_dir.name) / "team_standings.csv"
        pairings_path = Path(self.temp_dir.name) / "team_pairings.csv"
        teams_path = Path(self.temp_dir.name) / "teams.csv"
        complete_path = Path(self.temp_dir.name) / "team_complete.csv"
        scoresheets_path = Path(self.temp_dir.name) / "team_scoresheets.pdf"
        site_dir = Path(self.temp_dir.name) / "team_site"

        self.export_service.export_standings(tournament_id, standings_path)
        self.export_service.export_pairings(round_data["id"], pairings_path)
        self.export_service.export_teams(tournament_id, teams_path)
        self.export_service.export_complete(tournament_id, complete_path)
        self.export_service.export_scoresheets(round_data["id"], scoresheets_path)
        index_path = self.export_service.export_site(tournament_id, site_dir)

        standings_csv = standings_path.read_text(encoding="utf-8-sig")
        pairings_csv = pairings_path.read_text(encoding="utf-8-sig")
        teams_csv = teams_path.read_text(encoding="utf-8-sig")
        complete_csv = complete_path.read_text(encoding="utf-8-sig")
        html = index_path.read_text(encoding="utf-8")

        self.assertIn("Match points", standings_csv)
        self.assertIn("Equipe 1", standings_csv)
        self.assertIn("Confronto", pairings_csv)
        self.assertIn("2 x 0", pairings_csv)
        self.assertIn("Escalacoes", teams_csv)
        self.assertIn("Titular", teams_csv)
        self.assertIn("Classificacao por equipes", complete_csv)
        self.assertIn("Rodada 1", complete_csv)
        self.assertIn("Classificacao por equipes", html)
        self.assertIn("Equipes", html)
        self.assertIn("Escalacoes", html)
        self.assertIn("2 x 0", html)
        self.assertEqual(scoresheets_path.read_bytes()[:4], b"%PDF")
        from pypdf import PdfReader

        self.assertEqual(2, len(PdfReader(scoresheets_path).pages))

    def test_team_crosstable_preserves_board_history_totals_and_exports(self) -> None:
        tournament_id, _team_ids = self._create_team_tournament(teams_count=2, boards_count=2)
        round_data = self.service.generate_next_round(tournament_id)
        match = self.db.list_team_matches_for_round(int(round_data["id"]))[0]
        boards = self.db.list_team_boards(int(match["id"]))
        source_team = self.db.get_team_player_by_player(int(boards[0]["white_player_id"]))
        reserve_id = self.db.create_player(tournament_id, name="Reserva escalado", rating=1750)
        self.team_service.add_player(int(source_team["team_id"]), reserve_id, board_number="", role="reserve")
        self.service.adjust_team_board_player(
            tournament_id,
            int(round_data["id"]),
            int(boards[0]["id"]),
            "white",
            reserve_id,
        )
        for board in boards:
            white_player_team = self.db.get_team_player_by_player(int(board["white_player_id"]))
            result = "1-0" if int(white_player_team["team_id"]) == int(match["white_team_id"]) else "0-1"
            self.service.update_result(tournament_id, int(board["id"]), result)
        self.service.close_round(tournament_id, int(round_data["id"]))

        payload = self.service.team_crosstable(tournament_id)
        standings = self.service.team_standings(tournament_id)
        white_row = next(item for item in payload["rows"] if int(item["team_id"]) == int(match["white_team_id"]))
        black_row = next(item for item in payload["rows"] if int(item["team_id"]) == int(match["black_team_id"]))
        white_cell = white_row["rounds"][1]
        black_cell = black_row["rounds"][1]

        self.assertEqual("team", payload["competition_type"])
        self.assertIn("B 1-0 MP 2 GP 2", white_cell["label"])
        self.assertIn("P 0-1 MP 0 GP 0", black_cell["label"])
        self.assertEqual("Reserva escalado", white_cell["boards"][0]["white_player_name"])
        self.assertEqual(
            sum(float(item["match_points"]) for item in standings),
            sum(float(item["match_points"]) for item in payload["rows"]),
        )
        self.assertEqual(
            sum(float(item["game_points"]) for item in standings),
            sum(float(item["game_points"]) for item in payload["rows"]),
        )

        for extension in ("csv", "xlsx", "pdf", "html"):
            output_path = Path(self.temp_dir.name) / f"tabela_cruzada_equipes.{extension}"
            self.export_service.export_crosstable(tournament_id, output_path)
            self.assertTrue(output_path.exists())
        csv_content = (Path(self.temp_dir.name) / "tabela_cruzada_equipes.csv").read_text(encoding="utf-8-sig")
        html_content = (Path(self.temp_dir.name) / "tabela_cruzada_equipes.html").read_text(encoding="utf-8")
        self.assertIn("Equipe", csv_content)
        self.assertIn("MP 2 GP 2", csv_content)
        self.assertIn("Tabela cruzada por equipes", html_content)

    def test_scoresheets_pdf_skips_bye_and_includes_player_data(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana Silva",
            rating=1810,
            club="Clube A",
            fide_id="1234567",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno Souza",
            rating=1720,
            club="Clube B",
            cbx_id="7654",
        )
        self.db.create_player(self.tournament_id, name="Carla Lima", rating=1650, club="Clube C")
        round_data = self.service.generate_next_round(self.tournament_id)
        output_path = Path(self.temp_dir.name) / "sumulas.pdf"

        self.export_service.export_scoresheets(int(round_data["id"]), output_path)

        from pypdf import PdfReader

        reader = PdfReader(output_path)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        self.assertEqual(output_path.read_bytes()[:4], b"%PDF")
        self.assertEqual(1, len(reader.pages))
        self.assertIn("Resultado:", text)
        self.assertIn("Assinatura das brancas", text)
        self.assertTrue("Ana Silva" in text or "Bruno Souza" in text or "Carla Lima" in text)

    def test_scoresheets_pdf_exports_121_players_under_five_seconds(self) -> None:
        for index in range(121):
            self.db.create_player(
                self.tournament_id,
                name=f"Jogador {index + 1:03d}",
                rating=2400 - index,
                club="Clube",
            )
        round_data = self.service.generate_next_round(self.tournament_id)
        output_path = Path(self.temp_dir.name) / "sumulas_121.pdf"

        started_at = perf_counter()
        self.export_service.export_scoresheets(int(round_data["id"]), output_path)
        elapsed = perf_counter() - started_at

        from pypdf import PdfReader

        self.assertEqual(60, len(PdfReader(output_path).pages))
        self.assertLess(elapsed, 5.0)

    def test_pairings_wall_pdf_includes_valid_qr_only_for_open_round(self) -> None:
        self.db.create_player(self.tournament_id, name="Ana Silva", rating=1810)
        self.db.create_player(self.tournament_id, name="Bruno Souza", rating=1720)
        round_data = self.service.generate_next_round(self.tournament_id)
        pairing = self.db.get_pairings_for_round(int(round_data["id"]))[0]
        open_path = Path(self.temp_dir.name) / "mural_aberto.pdf"
        captured_urls: list[str] = []

        original_qr_urls = self.export_service._qr_result_urls

        def capture_qr_urls(tournament_id: int, pairings: list[dict[str, object]]) -> dict[int, str]:
            urls = original_qr_urls(tournament_id, pairings)
            captured_urls.extend(urls.values())
            return urls

        with mock.patch.object(self.export_service, "_qr_result_urls", side_effect=capture_qr_urls):
            self.export_service.export_pairings(int(round_data["id"]), open_path)

        from urllib.parse import parse_qs, urlparse

        from pypdf import PdfReader

        open_text = "\n".join(page.extract_text() or "" for page in PdfReader(open_path).pages)
        token = parse_qs(urlparse(captured_urls[0]).query)["token"][0]
        token_row = self.qr_result_service._validate_token(token)
        self.assertEqual(1, len(captured_urls))
        self.assertEqual(int(pairing["id"]), int(token_row["pairing_id"]))
        self.assertIn("Rodada aberta - QR para envio de resultado", open_text)
        self.assertIn("Ana Silva", open_text)
        self.assertIn("Bruno Souza", open_text)

        self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, int(round_data["id"]))
        closed_path = Path(self.temp_dir.name) / "mural_fechado.pdf"
        with mock.patch.object(self.export_service, "_qr_result_urls") as qr_urls:
            self.export_service.export_pairings(int(round_data["id"]), closed_path)

        closed_text = "\n".join(page.extract_text() or "" for page in PdfReader(closed_path).pages)
        qr_urls.assert_called_once_with(self.tournament_id, [])
        self.assertIn("Rodada fechada - QR desativado", closed_text)

    def test_pairings_wall_pdf_exports_60_tables_under_five_seconds(self) -> None:
        for index in range(120):
            self.db.create_player(
                self.tournament_id,
                name=f"Jogador {index + 1:03d}",
                rating=2400 - index,
            )
        round_data = self.service.generate_next_round(self.tournament_id)
        output_path = Path(self.temp_dir.name) / "mural_60_mesas.pdf"

        started_at = perf_counter()
        self.export_service.export_pairings(int(round_data["id"]), output_path)
        elapsed = perf_counter() - started_at

        from pypdf import PdfReader

        self.assertEqual(4, len(PdfReader(output_path).pages))
        self.assertLess(elapsed, 5.0)

    def test_table_cards_pdf_exports_configured_range_with_optional_qr(self) -> None:
        for index in range(6):
            self.db.create_player(
                self.tournament_id,
                name=f"Jogador {index + 1}",
                rating=1800 - index,
            )
        round_data = self.service.generate_next_round(self.tournament_id)
        output_path = Path(self.temp_dir.name) / "cartoes.pdf"
        captured_pairing_ids: list[int] = []
        original_qr_urls = self.export_service._qr_result_urls

        def capture_qr_urls(tournament_id: int, pairings: list[dict[str, object]]) -> dict[int, str]:
            captured_pairing_ids.extend(int(pairing["id"]) for pairing in pairings)
            return original_qr_urls(tournament_id, pairings)

        with mock.patch.object(self.export_service, "_qr_result_urls", side_effect=capture_qr_urls):
            self.export_service.export_table_cards(
                output_path,
                1,
                6,
                round_id=int(round_data["id"]),
                include_qr=True,
            )

        from pypdf import PdfReader

        reader = PdfReader(output_path)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        pairings = self.db.get_pairings_for_round(int(round_data["id"]))
        self.assertEqual(2, len(reader.pages))
        self.assertEqual([int(pairing["id"]) for pairing in pairings], captured_pairing_ids)
        self.assertIn("MESA", text)
        self.assertIn("1", text)
        self.assertIn("6", text)
        self.assertIn("Rodada 1", text)
        self.assertIn("QR para enviar resultado", text)

        for pairing in pairings:
            self.service.update_result(self.tournament_id, int(pairing["id"]), "1-0")
        self.service.close_round(self.tournament_id, int(round_data["id"]))
        closed_path = Path(self.temp_dir.name) / "cartoes_fechados.pdf"
        with mock.patch.object(self.export_service, "_qr_result_urls") as qr_urls:
            self.export_service.export_table_cards(
                closed_path,
                1,
                2,
                round_id=int(round_data["id"]),
                include_qr=True,
            )
        closed_text = "\n".join(page.extract_text() or "" for page in PdfReader(closed_path).pages)
        qr_urls.assert_called_once_with(self.tournament_id, [])
        self.assertIn("QR indisponivel para esta mesa", closed_text)

    def test_table_cards_pdf_exports_200_cards_under_five_seconds(self) -> None:
        output_path = Path(self.temp_dir.name) / "cartoes_200.pdf"

        started_at = perf_counter()
        self.export_service.export_table_cards(output_path, 1, 200)
        elapsed = perf_counter() - started_at

        from pypdf import PdfReader

        self.assertEqual(50, len(PdfReader(output_path).pages))
        self.assertLess(elapsed, 5.0)

    def test_individual_crosstable_round_robin_fixture_and_exports(self) -> None:
        player_ids = [
            self.db.create_player(self.tournament_id, name=f"Jogador {index}", rating=2100 - index * 50)
            for index in range(1, 7)
        ]
        schedules = [
            [(0, 5), (1, 4), (2, 3)],
            [(5, 3), (4, 2), (0, 1)],
            [(1, 5), (2, 0), (3, 4)],
            [(5, 4), (0, 3), (1, 2)],
            [(2, 5), (3, 1), (4, 0)],
        ]
        for round_number, schedule in enumerate(schedules, start=1):
            round_id = self.db.create_round_with_pairings(
                self.tournament_id,
                round_number,
                [
                    {
                        "board_number": board_number,
                        "white_player_id": player_ids[white],
                        "black_player_id": player_ids[black],
                        "result": "1-0",
                    }
                    for board_number, (white, black) in enumerate(schedule, start=1)
                ],
            )
            self.db.close_round(round_id)

        payload = self.service.crosstable(self.tournament_id)
        rows_by_player = {int(row["player_id"]): row for row in payload["rows"]}
        first = rows_by_player[player_ids[0]]

        self.assertEqual([1, 2, 3, 4, 5], payload["rounds"])
        self.assertEqual(6, len(payload["rows"]))
        self.assertEqual("6B 1-0", first["rounds"][1]["label"])
        self.assertEqual("2B 1-0", first["rounds"][2]["label"])
        self.assertEqual("3P 1-0", first["rounds"][3]["label"])
        self.assertAlmostEqual(15.0, sum(float(row["points"]) for row in payload["rows"]))

        for extension in ("csv", "xlsx", "pdf", "html"):
            output_path = Path(self.temp_dir.name) / f"tabela_cruzada.{extension}"
            self.export_service.export_crosstable(self.tournament_id, output_path)
            self.assertTrue(output_path.exists())
        csv_content = (Path(self.temp_dir.name) / "tabela_cruzada.csv").read_text(encoding="utf-8-sig")
        html_content = (Path(self.temp_dir.name) / "tabela_cruzada.html").read_text(encoding="utf-8")
        self.assertIn("R5", csv_content)
        self.assertIn("6B 1-0", csv_content)
        self.assertIn("Tabela cruzada", html_content)
        self.assertIn("B = brancas", html_content)

    def test_individual_crosstable_distinguishes_bye_wo_and_absence(self) -> None:
        player_ids = [
            self.db.create_player(self.tournament_id, name=f"Jogador {index}", rating=1800 - index)
            for index in range(1, 5)
        ]
        round_id = self.db.create_round_with_pairings(
            self.tournament_id,
            1,
            [
                {
                    "board_number": 1,
                    "white_player_id": player_ids[0],
                    "black_player_id": player_ids[1],
                    "result": "1F-0F",
                },
                {
                    "board_number": 2,
                    "white_player_id": player_ids[2],
                    "black_player_id": None,
                    "result": "H",
                    "is_bye": 1,
                },
            ],
        )
        self.db.close_round(round_id)

        payload = self.service.crosstable(self.tournament_id)
        rows_by_player = {int(row["player_id"]): row for row in payload["rows"]}
        self.assertEqual("1F-0F", rows_by_player[player_ids[0]]["rounds"][1]["result"])
        self.assertEqual("bye", rows_by_player[player_ids[2]]["rounds"][1]["kind"])
        self.assertEqual("BYE H", rows_by_player[player_ids[2]]["rounds"][1]["label"])
        self.assertEqual("absent", rows_by_player[player_ids[3]]["rounds"][1]["kind"])
        self.assertEqual("-", rows_by_player[player_ids[3]]["rounds"][1]["label"])

    def test_initial_player_list_pdf_uses_trf_start_rank_and_signature_column(self) -> None:
        self.db.create_player(
            self.tournament_id,
            name="Ana Ausente",
            rating=1800,
            international_rating=2200,
            club="Clube A",
            category="ABS",
            player_status="absent",
        )
        self.db.create_player(
            self.tournament_id,
            name="Bruno Presente",
            rating=2100,
            club="Clube B",
            category="Sub-18",
        )
        self.db.create_player(
            self.tournament_id,
            name="Carla Presente",
            rating=1900,
            club="Clube C",
            category="FEM",
        )
        output_path = Path(self.temp_dir.name) / "lista_chamada.pdf"

        self.export_service.export_initial_player_list(self.tournament_id, output_path)

        from pypdf import PdfReader

        text = "\n".join(page.extract_text() or "" for page in PdfReader(output_path).pages)
        self.assertEqual(output_path.read_bytes()[:4], b"%PDF")
        self.assertIn("Inicial", text)
        self.assertIn("Jogador", text)
        self.assertIn("Rating", text)
        self.assertIn("Clube", text)
        self.assertIn("Categoria", text)
        self.assertIn("Assinatura", text)
        self.assertLess(text.index("Ana Ausente"), text.index("Bruno Presente"))
        self.assertLess(text.index("Bruno Presente"), text.index("Carla Presente"))

    def test_initial_player_list_pdf_exports_801_players_under_five_seconds(self) -> None:
        for index in range(801):
            self.db.create_player(
                self.tournament_id,
                name=f"Jogador {index + 1:03d}",
                rating=2600 - index,
                club="Clube",
                category="ABS",
            )
        output_path = Path(self.temp_dir.name) / "lista_chamada_801.pdf"

        started_at = perf_counter()
        self.export_service.export_initial_player_list(self.tournament_id, output_path)
        elapsed = perf_counter() - started_at

        from pypdf import PdfReader

        reader = PdfReader(output_path)
        self.assertEqual(22, len(reader.pages))
        self.assertLess(elapsed, 5.0)

    def test_existing_database_gets_member_column_on_initialize(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE tournaments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    location TEXT DEFAULT '',
                    start_date TEXT DEFAULT '',
                    end_date TEXT DEFAULT '',
                    system TEXT NOT NULL DEFAULT 'Suico',
                    rounds_count INTEGER NOT NULL DEFAULT 5,
                    time_control TEXT DEFAULT '',
                    bye_points REAL NOT NULL DEFAULT 1.0,
                    status TEXT NOT NULL DEFAULT 'draft',
                    created_at TEXT NOT NULL
                );

                CREATE TABLE players (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    tournament_id INTEGER NOT NULL,
                    name TEXT NOT NULL,
                    club TEXT DEFAULT '',
                    federation_id TEXT DEFAULT '',
                    fide_id TEXT DEFAULT '',
                    rating INTEGER NOT NULL DEFAULT 0,
                    category TEXT DEFAULT '',
                    birth_date TEXT DEFAULT '',
                    active INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL
                );
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(players)").fetchall()
            }
            club_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(clubs)").fetchall()
            }
            member_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(members)").fetchall()
            }
            tournament_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(tournaments)").fetchall()
            }
            settings_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(tournament_settings)").fetchall()
            }
            settings_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'tournament_settings'
                """
            ).fetchone()
            schedule_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'round_schedule'
                """
            ).fetchone()
            official_snapshots_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'official_rating_snapshots'
                """
            ).fetchone()
            official_players_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'official_players'
                """
            ).fetchone()
            internal_rating_history_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'internal_rating_history'
                """
            ).fetchone()
            app_settings_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'app_settings'
                """
            ).fetchone()
            audit_log_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'audit_log'
                """
            ).fetchone()
            classes_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'classes'
                """
            ).fetchone()
            enrollments_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'member_class_enrollments'
                """
            ).fetchone()
            guardians_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'guardians'
                """
            ).fetchone()
            member_guardians_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'member_guardians'
                """
            ).fetchone()
            learning_levels_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'learning_levels'
                """
            ).fetchone()
            learning_levels_count = migrated_connection.execute(
                "SELECT COUNT(*) FROM learning_levels"
            ).fetchone()[0]
            training_sessions_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'training_sessions'
                """
            ).fetchone()
            training_session_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(training_sessions)").fetchall()
            }
            exercise_library_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'exercise_library'
                """
            ).fetchone()
            training_lists_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'training_lists'
                """
            ).fetchone()
            training_list_exercises_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'training_list_exercises'
                """
            ).fetchone()
            exercise_attempts_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'exercise_attempts'
                """
            ).fetchone()
            attendance_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'attendance'
                """
            ).fetchone()
            membership_plans_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'membership_plans'
                """
            ).fetchone()
            payments_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'payments'
                """
            ).fetchone()
            club_events_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'club_events'
                """
            ).fetchone()
            inventory_items_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'inventory_items'
                """
            ).fetchone()
            inventory_loans_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'inventory_loans'
                """
            ).fetchone()
            inventory_maintenance_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'inventory_maintenance'
                """
            ).fetchone()
            teams_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'teams'
                """
            ).fetchone()
            team_players_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'team_players'
                """
            ).fetchone()
            team_matches_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'team_matches'
                """
            ).fetchone()
            team_boards_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'team_boards'
                """
            ).fetchone()
            team_lineups_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'team_lineups'
                """
            ).fetchone()
            team_lineup_boards_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'team_lineup_boards'
                """
            ).fetchone()
            team_substitution_events_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'team_substitution_events'
                """
            ).fetchone()
            certificate_templates_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'certificate_templates'
                """
            ).fetchone()
            certificate_issuances_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'certificate_issuances'
                """
            ).fetchone()
            certificate_templates_count = migrated_connection.execute(
                "SELECT COUNT(*) FROM certificate_templates"
            ).fetchone()[0]
            certificate_template_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(certificate_templates)").fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("member_id", columns)
        self.assertIn("cbx_id", columns)
        self.assertIn("lbx_id", columns)
        self.assertIn("national_rating", columns)
        self.assertIn("international_rating", columns)
        self.assertIn("player_status", columns)
        self.assertIn("starting_points", columns)
        self.assertIn("age_category", columns)
        self.assertIn("rating_category", columns)
        self.assertIn("prize_tags", columns)
        self.assertIn("kind", club_columns)
        self.assertIn("active", club_columns)
        self.assertIn("surname", member_columns)
        self.assertIn("age_category", member_columns)
        self.assertIn("rating_category", member_columns)
        self.assertIn("prize_tags", member_columns)
        self.assertIn("learning_level_id", member_columns)
        self.assertIn("club_id", tournament_columns)
        self.assertIn("class_id", tournament_columns)
        self.assertIn("competition_type", tournament_columns)
        self.assertIn("tournament_profile", settings_columns)
        self.assertIn("team_boards_count", settings_columns)
        self.assertIn("team_match_win_points", settings_columns)
        self.assertIn("team_match_draw_points", settings_columns)
        self.assertIn("team_match_loss_points", settings_columns)
        self.assertIn("team_pairing_method", settings_columns)
        self.assertIn("team_standing_primary", settings_columns)
        self.assertIn("team_standing_secondary", settings_columns)
        self.assertIn("team_fixed_board_order", settings_columns)
        self.assertIn("team_board_order_policy", settings_columns)
        self.assertIn("team_reserve_policy", settings_columns)
        self.assertIn("team_lineup_deadline", settings_columns)
        self.assertIn("team_max_substitutions", settings_columns)
        self.assertIsNotNone(settings_table)
        self.assertIsNotNone(schedule_table)
        self.assertIsNotNone(official_snapshots_table)
        self.assertIsNotNone(official_players_table)
        self.assertIsNotNone(internal_rating_history_table)
        self.assertIsNotNone(app_settings_table)
        self.assertIsNotNone(audit_log_table)
        self.assertIsNotNone(classes_table)
        self.assertIsNotNone(enrollments_table)
        self.assertIsNotNone(guardians_table)
        self.assertIsNotNone(member_guardians_table)
        self.assertIsNotNone(learning_levels_table)
        self.assertGreaterEqual(learning_levels_count, 5)
        self.assertIsNotNone(training_sessions_table)
        self.assertIn("training_list_id", training_session_columns)
        self.assertIn("learning_level_id", training_session_columns)
        self.assertIn("objective", training_session_columns)
        self.assertIn("content", training_session_columns)
        self.assertIn("homework", training_session_columns)
        self.assertIsNotNone(exercise_library_table)
        self.assertIsNotNone(training_lists_table)
        self.assertIsNotNone(training_list_exercises_table)
        self.assertIsNotNone(exercise_attempts_table)
        self.assertIsNotNone(attendance_table)
        self.assertIsNotNone(membership_plans_table)
        self.assertIsNotNone(payments_table)
        self.assertIsNotNone(club_events_table)
        self.assertIsNotNone(inventory_items_table)
        self.assertIsNotNone(inventory_loans_table)
        self.assertIsNotNone(inventory_maintenance_table)
        self.assertIsNotNone(teams_table)
        self.assertIsNotNone(team_players_table)
        self.assertIsNotNone(team_matches_table)
        self.assertIsNotNone(team_boards_table)
        self.assertIsNotNone(team_lineups_table)
        self.assertIsNotNone(team_lineup_boards_table)
        self.assertIsNotNone(team_substitution_events_table)
        self.assertIsNotNone(certificate_templates_table)
        self.assertIsNotNone(certificate_issuances_table)
        self.assertGreaterEqual(certificate_templates_count, 7)
        self.assertIn("logo_path", certificate_template_columns)
        self.assertIn("background_image_path", certificate_template_columns)
        self.assertIn("background_opacity", certificate_template_columns)
        self.assertIn("secondary_logo_path", certificate_template_columns)
        self.assertIn("primary_color", certificate_template_columns)
        self.assertIn("title_font_size", certificate_template_columns)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_v3_database_adds_learning_level_column_before_indexes(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v3.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE members (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    club_id INTEGER DEFAULT 1,
                    name TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'active',
                    surname TEXT DEFAULT '',
                    age_category TEXT DEFAULT '',
                    rating_category TEXT DEFAULT '',
                    prize_tags TEXT DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                PRAGMA user_version = 3;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            member_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(members)").fetchall()
            }
            member_indexes = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA index_list(members)").fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("learning_level_id", member_columns)
        self.assertIn("idx_members_learning_level", member_indexes)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_v7_database_repairs_missing_learning_column_during_v8_migration(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v7_missing_learning.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE members (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    club_id INTEGER DEFAULT 1,
                    name TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'active',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                PRAGMA user_version = 7;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            member_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(members)").fetchall()
            }
            settings_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(tournament_settings)").fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("learning_level_id", member_columns)
        self.assertIn("tournament_profile", settings_columns)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_v8_database_adds_certificate_visual_asset_columns(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v8_certificate_assets.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE certificate_templates (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL UNIQUE,
                    certificate_type TEXT NOT NULL DEFAULT 'participation',
                    title_template TEXT NOT NULL DEFAULT '',
                    body_template TEXT NOT NULL DEFAULT '',
                    footer_template TEXT DEFAULT '',
                    orientation TEXT NOT NULL DEFAULT 'landscape',
                    signature_left TEXT DEFAULT '',
                    signature_right TEXT DEFAULT '',
                    logo_path TEXT DEFAULT '',
                    primary_color TEXT NOT NULL DEFAULT '#1E3A8A',
                    accent_color TEXT NOT NULL DEFAULT '#93C5FD',
                    title_font_size INTEGER NOT NULL DEFAULT 32,
                    body_font_size INTEGER NOT NULL DEFAULT 18,
                    footer_font_size INTEGER NOT NULL DEFAULT 10,
                    active INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                INSERT INTO certificate_templates (
                    name, certificate_type, title_template, body_template,
                    created_at, updated_at
                ) VALUES (
                    'Modelo antigo', 'participation', 'Titulo', 'Texto', '2026-01-01', '2026-01-01'
                );

                PRAGMA user_version = 8;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            certificate_template_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(certificate_templates)").fetchall()
            }
            template = migrated_connection.execute(
                """
                SELECT background_image_path, background_opacity, secondary_logo_path
                FROM certificate_templates
                WHERE name = 'Modelo antigo'
                """
            ).fetchone()
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("background_image_path", certificate_template_columns)
        self.assertIn("background_opacity", certificate_template_columns)
        self.assertIn("secondary_logo_path", certificate_template_columns)
        self.assertEqual(template["background_image_path"], "")
        self.assertEqual(template["background_opacity"], 0.18)
        self.assertEqual(template["secondary_logo_path"], "")
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_v10_database_adds_exercise_library_schema(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v10_exercises.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE training_sessions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    club_id INTEGER DEFAULT 1,
                    class_id INTEGER,
                    title TEXT NOT NULL,
                    session_type TEXT NOT NULL DEFAULT 'aula',
                    session_date TEXT DEFAULT '',
                    start_time TEXT DEFAULT '',
                    end_time TEXT DEFAULT '',
                    instructor TEXT DEFAULT '',
                    location TEXT DEFAULT '',
                    learning_level_id INTEGER,
                    objective TEXT DEFAULT '',
                    content TEXT DEFAULT '',
                    homework TEXT DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'planned',
                    notes TEXT DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                PRAGMA user_version = 10;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            training_session_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(training_sessions)").fetchall()
            }
            exercise_tables = {
                row["name"]
                for row in migrated_connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table'
                        AND name IN (
                            'exercise_library',
                            'training_lists',
                            'training_list_exercises',
                            'exercise_attempts'
                        )
                    """
                ).fetchall()
            }
            exercise_indexes = {
                row["name"]
                for table_name in (
                    "exercise_library",
                    "training_lists",
                    "training_list_exercises",
                    "exercise_attempts",
                    "training_sessions",
                )
                for row in migrated_connection.execute(f"PRAGMA index_list({table_name})").fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("training_list_id", training_session_columns)
        self.assertEqual(
            {"exercise_library", "training_lists", "training_list_exercises", "exercise_attempts"},
            exercise_tables,
        )
        self.assertIn("idx_exercise_library_filters", exercise_indexes)
        self.assertIn("idx_training_lists_filters", exercise_indexes)
        self.assertIn("idx_training_sessions_training_list", exercise_indexes)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_v11_database_adds_inventory_schema(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v11_inventory.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                PRAGMA user_version = 11;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            inventory_tables = {
                row["name"]
                for row in migrated_connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table'
                        AND name IN (
                            'inventory_items',
                            'inventory_loans',
                            'inventory_maintenance'
                        )
                    """
                ).fetchall()
            }
            inventory_indexes = {
                row["name"]
                for table_name in ("inventory_items", "inventory_loans", "inventory_maintenance")
                for row in migrated_connection.execute(f"PRAGMA index_list({table_name})").fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertEqual({"inventory_items", "inventory_loans", "inventory_maintenance"}, inventory_tables)
        self.assertIn("idx_inventory_items_filters", inventory_indexes)
        self.assertIn("idx_inventory_loans_item", inventory_indexes)
        self.assertIn("idx_inventory_maintenance_item", inventory_indexes)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_v12_database_adds_audit_log_schema(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v12_audit.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                PRAGMA user_version = 12;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            audit_table = migrated_connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'audit_log'
                """
            ).fetchone()
            audit_indexes = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA index_list(audit_log)").fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIsNotNone(audit_table)
        self.assertIn("idx_audit_log_created", audit_indexes)
        self.assertIn("idx_audit_log_entity", audit_indexes)
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_v22_database_adds_phase0_arbitration_schema(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v22_phase0.db"
        Database(legacy_path, backup_dir=self.backup_dir)
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                DROP TABLE IF EXISTS audit_events;
                DROP TABLE IF EXISTS pairing_snapshots;
                DROP TABLE IF EXISTS standings_snapshots;
                DROP TABLE IF EXISTS tiebreak_components;
                DROP TABLE IF EXISTS public_tokens;
                DROP TABLE IF EXISTS result_submissions;
                PRAGMA user_version = 22;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            settings_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(tournament_settings)").fetchall()
            }
            round_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(rounds)").fetchall()
            }
            phase0_tables = {
                row["name"]
                for row in migrated_connection.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type = 'table'
                        AND name IN (
                            'audit_events', 'pairing_snapshots', 'standings_snapshots',
                            'tiebreak_components', 'public_tokens', 'result_submissions'
                        )
                    """
                ).fetchall()
            }
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("pairing_system", settings_columns)
        self.assertIn("acceleration_method", settings_columns)
        self.assertIn("pairing_engine_version", round_columns)
        self.assertIn("ruleset_version", round_columns)
        self.assertEqual(
            {
                "audit_events", "pairing_snapshots", "standings_snapshots",
                "tiebreak_components", "public_tokens", "result_submissions",
            },
            phase0_tables,
        )
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_v32_database_adds_round_closed_at_and_backfills_closed_rounds(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_v32_round_clock.db"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                """
                CREATE TABLE rounds (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    tournament_id INTEGER NOT NULL,
                    number INTEGER NOT NULL,
                    status TEXT NOT NULL DEFAULT 'generated',
                    pairing_engine_version TEXT NOT NULL DEFAULT 'albericus-swiss-1',
                    ruleset_version TEXT NOT NULL DEFAULT 'albericus-2026-phase0',
                    created_at TEXT NOT NULL
                );
                INSERT INTO rounds (tournament_id, number, status, created_at)
                VALUES (1, 1, 'closed', '2026-05-31 09:00:00');
                PRAGMA user_version = 32;
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        with migrated.connect() as migrated_connection:
            columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(rounds)").fetchall()
            }
            round_data = migrated_connection.execute("SELECT * FROM rounds WHERE id = 1").fetchone()
            user_version = migrated_connection.execute("PRAGMA user_version").fetchone()[0]

        self.assertIn("closed_at", columns)
        self.assertEqual("2026-05-31 09:00:00", round_data["closed_at"])
        self.assertEqual(Database.SCHEMA_VERSION, user_version)

    def test_current_database_repairs_legacy_users_table_before_login(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy_current_users.db"
        legacy_hash = hashlib.sha256("admin".encode()).hexdigest()
        connection = sqlite3.connect(legacy_path)
        try:
            connection.executescript(
                f"""
                CREATE TABLE users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT NOT NULL UNIQUE,
                    password_hash TEXT NOT NULL,
                    role TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );

                INSERT INTO users (username, password_hash, role, created_at)
                VALUES ('admin', '{legacy_hash}', 'admin', '2026-01-01 00:00:00');

                PRAGMA user_version = {Database.SCHEMA_VERSION};
                """
            )
            connection.commit()
        finally:
            connection.close()

        migrated = Database(legacy_path, backup_dir=self.backup_dir)
        security_service = SecurityService(migrated)

        self.assertTrue(security_service.login("admin", "admin"))
        with migrated.connect() as migrated_connection:
            user_columns = {
                row["name"]
                for row in migrated_connection.execute("PRAGMA table_info(users)").fetchall()
            }
            user = migrated_connection.execute(
                "SELECT password_hash, updated_at FROM users WHERE username = 'admin'"
            ).fetchone()

        self.assertIn("updated_at", user_columns)
        self.assertTrue(str(user["password_hash"]).startswith("pbkdf2_sha256$"))
        self.assertTrue(user["updated_at"])

    def test_future_database_version_is_rejected(self) -> None:
        future_path = Path(self.temp_dir.name) / "future.db"
        connection = sqlite3.connect(future_path)
        try:
            connection.execute(f"PRAGMA user_version = {Database.SCHEMA_VERSION + 1}")
            connection.commit()
        finally:
            connection.close()

        with self.assertRaisesRegex(RuntimeError, "versao mais nova"):
            Database(future_path, backup_dir=self.backup_dir)


# ----------------------------------------------------------------------
# Fase J — integrações externas e publicação
# ----------------------------------------------------------------------


class _RecordingFTP:
    """Cliente FTP falso (sem rede) para testar ftp_publish/PhotoAlbumService."""

    def __init__(self) -> None:
        self.stored: dict[str, bytes] = {}
        self.made: list[str] = []
        self.cwds: list[str] = []
        self.quit_called = False

    def cwd(self, path: str) -> None:
        self.cwds.append(path)
        if path not in self.made:
            raise ftplib.error_perm("550 No such directory")

    def mkd(self, path: str) -> None:
        self.made.append(path)

    def storbinary(self, command: str, handle: Any) -> None:
        name = command.split(" ", 1)[1]
        self.stored[name] = handle.read()

    def voidcmd(self, command: str) -> str:
        return "200 OK"

    def quit(self) -> None:
        self.quit_called = True


class ChessResultsBridgeTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.pairing_service = PairingService(self.db)
        self.export_service = ExportService(self.db, self.pairing_service)
        self.service = ChessResultsService(self.db, self.export_service)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_normalize_results_url(self) -> None:
        from src.services.chess_results import normalize_results_url

        self.assertEqual(
            "https://chess-results.com/tnr123.aspx",
            normalize_results_url("chess-results.com/tnr123.aspx"),
        )
        self.assertEqual(
            "https://chess-results.com/tnr1.aspx",
            normalize_results_url("http://chess-results.com/tnr1.aspx"),
        )
        self.assertEqual("", normalize_results_url("https://exemplo.com/tnr1"))
        self.assertEqual("", normalize_results_url(""))

    def test_parse_entries_csv(self) -> None:
        from src.services.chess_results import parse_entries_csv

        content = "Name,Rtg,FED,FideID,Sex\nAna Silva,1500,BRA,12345,F\nJoao Souza,1800,POR,999,M\n"
        payloads, errors = parse_entries_csv(content)
        self.assertEqual(2, len(payloads))
        self.assertEqual([], errors)
        self.assertEqual("Ana Silva", payloads[0]["name"])
        self.assertEqual(1500, payloads[0]["rating"])
        self.assertEqual("BRA", payloads[0]["federation_id"])
        self.assertEqual("12345", payloads[0]["fide_id"])
        self.assertEqual("F", payloads[0]["sex"])

    def test_parse_entries_csv_reports_empty_name(self) -> None:
        from src.services.chess_results import parse_entries_csv

        payloads, errors = parse_entries_csv("Name,Rtg\n,1500\nBia,1600\n")
        self.assertEqual(1, len(payloads))
        self.assertTrue(any("nome vazio" in err.lower() for err in errors))

    def test_import_entries_creates_players(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        before = len(self.db.list_players(tournament_id, active_only=False))
        csv_path = Path(self.temp_dir.name) / "entries.csv"
        csv_path.write_text(
            "Name,Rtg,FED,FideID\nNovo Um,1234,BRA,1\nNovo Dois,1300,POR,2\n",
            encoding="utf-8",
        )
        result = self.service.import_entries(tournament_id, csv_path)
        self.assertEqual(2, result["imported"])
        after = len(self.db.list_players(tournament_id, active_only=False))
        self.assertEqual(before + 2, after)

    def test_prepare_upload_and_published_url(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        result = self.service.prepare_upload(tournament_id, self.temp_dir.name)
        self.assertTrue(os.path.exists(result["trf_path"]))
        self.assertTrue(result["steps"])
        self.assertIn("chess-results.com", result["register_url"])

        saved = self.service.set_published_url(tournament_id, "chess-results.com/tnr9.aspx")
        self.assertEqual("https://chess-results.com/tnr9.aspx", saved)
        self.assertEqual(saved, self.service.get_published_url(tournament_id))
        with self.assertRaises(AppError):
            self.service.set_published_url(tournament_id, "http://exemplo.com/x")


class ForeignRatingListTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.service = OfficialRatingService(self.db)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_map_row_uses_aliases(self) -> None:
        from src.services.foreign_rating_lists import map_row

        payload = map_row({"Nome": "Maria Lima", "Rating": "1700", "ID": "77", "Sexo": "F"}, "POR")
        assert payload is not None
        self.assertEqual("Maria Lima", payload["name"])
        self.assertEqual(1700, payload["national_rating"])
        self.assertEqual(1700, payload["standard_rating"])
        self.assertEqual(0, payload["international_rating"])
        self.assertEqual("POR", payload["federation"])
        self.assertEqual("77", payload["external_id"])

    def test_map_row_without_name_returns_none(self) -> None:
        from src.services.foreign_rating_lists import map_row

        self.assertIsNone(map_row({"Rating": "1500"}, "ESP"))

    def test_import_foreign_list_creates_snapshot(self) -> None:
        csv_path = Path(self.temp_dir.name) / "por.csv"
        csv_path.write_text(
            "Nome,Rating,ID,Titulo\nMaria Lima,1700,77,\nPedro Alves,1650,88,FM\n",
            encoding="utf-8",
        )
        result = self.service.import_foreign_list(csv_path, "POR")
        self.assertEqual("POR", result["source"])
        self.assertEqual(2, result["imported"])
        self.assertIsNotNone(result["snapshot_id"])
        with self.db.connect() as connection:
            count = connection.execute(
                "SELECT COUNT(*) FROM official_players WHERE source = ?", ("POR",)
            ).fetchone()[0]
        self.assertEqual(2, count)

    def test_invalid_federation_rejected(self) -> None:
        csv_path = Path(self.temp_dir.name) / "x.csv"
        csv_path.write_text("Nome,Rating\nA,1\n", encoding="utf-8")
        with self.assertRaises(AppError):
            self.service.import_foreign_list(csv_path, "1")

    def test_add_and_list_foreign_federation(self) -> None:
        seeded = self.service.list_foreign_federations()
        self.assertIn("POR", seeded)
        code = self.service.add_foreign_federation("bol", "Bolivia (FEBODA)")
        self.assertEqual("BOL", code)
        self.assertIn("BOL", self.service.list_foreign_federations())


class FtpPhotoAlbumTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.service = PhotoAlbumService(self.db)
        self.photos = base / "fotos"
        self.photos.mkdir()
        (self.photos / "a.jpg").write_bytes(b"jpgdata")
        (self.photos / "b.PNG").write_bytes(b"pngdata")
        (self.photos / "notas.txt").write_text("ignore", encoding="utf-8")

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_list_images_and_gallery(self) -> None:
        from src.services.ftp_publish import build_gallery_html, list_images

        images = list_images(self.photos)
        self.assertEqual(["a.jpg", "b.PNG"], [img.name for img in images])
        html_out = build_gallery_html([img.name for img in images], "Torneio")
        self.assertIn("a.jpg", html_out)
        self.assertIn("<title>Torneio</title>", html_out)

    def test_upload_files_creates_remote_dir(self) -> None:
        from src.services.ftp_publish import list_images, upload_files

        fake = _RecordingFTP()
        uploaded = upload_files(fake, list_images(self.photos), "public/fotos")
        self.assertEqual(["a.jpg", "b.PNG"], uploaded)
        self.assertEqual(["public", "fotos"], fake.made)
        self.assertIn("a.jpg", fake.stored)

    def test_publish_album_with_fake_client(self) -> None:
        self.service.save_config(
            {"host": "ftp.exemplo.com", "port": 21, "user": "u", "password": "secret", "remote_dir": "fotos"}
        )
        fake = _RecordingFTP()
        result = self.service.publish_album(self.photos, client_factory=lambda _cfg: fake)
        self.assertEqual(3, result["total"])  # 2 imagens + index.html
        self.assertTrue(result["gallery"])
        self.assertIn("index.html", fake.stored)
        self.assertTrue(fake.quit_called)

    def test_password_protected_roundtrip(self) -> None:
        self.service.save_config({"host": "h", "port": 21, "user": "u", "password": "topsecret"})
        # get_config desprotege; o valor cru no banco não é o texto puro em Windows.
        self.assertEqual("topsecret", self.service.get_config()["password"])


class AccessExportTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.pairing_service = PairingService(self.db)
        self.export_service = ExportService(self.db, self.pairing_service)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_write_csv_bundle(self) -> None:
        from src.services.access_export import write_csv_bundle

        tables = {"Jogadores": (["id", "nome"], [[1, "Ana"], [2, "Joao"]])}
        bundle = write_csv_bundle(tables, self.temp_dir.name)
        self.assertTrue(Path(bundle["schema_ini"]).exists())
        self.assertTrue(any(p.endswith("Jogadores.csv") for p in bundle["csv_paths"]))
        schema = Path(bundle["schema_ini"]).read_text(encoding="utf-8")
        self.assertIn("[Jogadores.csv]", schema)
        self.assertIn("CharacterSet=65001", schema)

    def test_export_access_service(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        result = self.export_service.export_access(tournament_id, self.temp_dir.name)
        self.assertIn("Jogadores", result["tables"])
        self.assertIn("Classificacao", result["tables"])
        self.assertTrue(result["csv_paths"])
        self.assertTrue(Path(result["schema_ini"]).exists())
        # Sem driver ACE no ambiente de teste, o .accdb real não é gerado.
        self.assertFalse(result["driver_available"])
        self.assertIsNone(result["accdb"])


class BatchExportPhaseJTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.pairing_service = PairingService(self.db)
        self.export_service = ExportService(self.db, self.pairing_service)
        self.service = BatchExportService(self.export_service)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_available_report_keys(self) -> None:
        from src.services.batch_export import available_report_keys

        individual = available_report_keys(False)
        self.assertIn("classificacao", individual)
        self.assertNotIn("equipes", individual)
        self.assertIn("equipes", available_report_keys(True))

    def test_run_batch_generates_files_and_isolates(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        dest = Path(self.temp_dir.name) / "lote"
        result = self.service.run_batch(
            tournament_id,
            ["classificacao", "jogadores", "equipes"],
            ["csv", "xlsx"],
            dest,
        )
        # classificacao+jogadores × csv+xlsx = 4 arquivos; equipes é só de equipes -> ignorado.
        self.assertEqual(4, result["count"])
        self.assertEqual(4, len(result["generated"]))
        self.assertTrue(all(os.path.exists(path) for path in result["generated"]))
        self.assertTrue(any("equipes" in item.lower() for item in result["skipped"]))

    def test_run_batch_skips_incompatible_format(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        dest = Path(self.temp_dir.name) / "lote2"
        # 'jogadores' não suporta HTML -> nada gerado, reportado em skipped.
        result = self.service.run_batch(tournament_id, ["jogadores"], ["html"], dest)
        self.assertEqual(0, result["count"])
        self.assertTrue(result["skipped"])

    def test_run_batch_requires_selection(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        with self.assertRaises(AppError):
            self.service.run_batch(tournament_id, [], ["csv"], self.temp_dir.name)


# ----------------------------------------------------------------------
# Aprofundamento do Painel do Arbitro (UX + docs)
# ----------------------------------------------------------------------


class ArbiterPanelMetricsProgressTest(unittest.TestCase):
    FINAL = {"1-0", "0-1", "1/2-1/2"}

    def test_individual_progress(self) -> None:
        from src.services.pairing.arbitration import individual_round_dashboard_metrics

        pairings = [{"result": "1-0"}, {"result": ""}, {"is_bye": True, "result": ""}]
        metrics = individual_round_dashboard_metrics(pairings, self.FINAL)
        self.assertEqual(2, metrics["total_results"])  # exclui o bye
        self.assertEqual(1, metrics["resolved_results"])
        self.assertEqual(1, metrics["byes"])

    def test_team_progress(self) -> None:
        from src.services.pairing.arbitration import team_round_dashboard_metrics

        matches = [{"id": 1, "is_bye": False}, {"id": 2, "is_bye": True}]
        boards = {1: [{"result": "1-0"}, {"result": ""}]}
        metrics = team_round_dashboard_metrics(matches, boards, self.FINAL)
        self.assertEqual(2, metrics["total_results"])
        self.assertEqual(1, metrics["resolved_results"])
        self.assertEqual(1, metrics["pending_results"])
        self.assertEqual(1, metrics["byes"])


class ArbiterPanelDeepeningTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.pairing_service = PairingService(self.db)
        self.export_service = ExportService(self.db, self.pairing_service)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _open_round_tournament(self) -> int:
        tournament_id = self.db.create_tournament(
            name="Aberto",
            location="Sala",
            rounds_count=3,
            time_control="90+30",
            start_date="2026-06-02",
            end_date="2026-06-02",
        )
        for index in range(4):
            self.db.create_player(tournament_id=tournament_id, name=f"Jogador {index}", rating=1500 - index)
        self.pairing_service.generate_next_round(tournament_id)
        return tournament_id

    def test_dashboard_structured_alerts_and_progress(self) -> None:
        tournament_id = self._open_round_tournament()
        dashboard = self.pairing_service.arbitration_dashboard(tournament_id)
        metrics = dashboard["metrics"]
        self.assertEqual(2, metrics["total_results"])
        self.assertEqual(0, metrics["resolved_results"])
        self.assertEqual(0, metrics["round_progress_percent"])
        detailed = dashboard["alerts_detailed"]
        self.assertTrue(detailed)
        for alert in detailed:
            self.assertIn("text", alert)
            self.assertIn("severity", alert)
            self.assertIn("action", alert)
        self.assertTrue(any(alert["action"] == "pending_results" for alert in detailed))
        # `alerts` (strings) continua em sincronia.
        self.assertEqual([alert["text"] for alert in detailed], dashboard["alerts"])

    def test_progress_updates_after_result(self) -> None:
        tournament_id = self._open_round_tournament()
        rounds = self.db.list_rounds(tournament_id)
        pairings = self.db.get_pairings_for_round(int(rounds[-1]["id"]))
        self.pairing_service.update_result(tournament_id, int(pairings[0]["id"]), "1-0")
        metrics = self.pairing_service.arbitration_dashboard(tournament_id)["metrics"]
        self.assertEqual(1, metrics["resolved_results"])
        self.assertEqual(50, metrics["round_progress_percent"])

    def test_export_round_package(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        latest_round = sorted(self.db.list_rounds(tournament_id), key=lambda item: int(item["number"]))[-1]
        dest = Path(self.temp_dir.name) / "pacote"
        result = self.export_service.export_round_package(int(latest_round["id"]), dest)
        self.assertEqual([], result["errors"])
        self.assertEqual(3, result["count"])  # mural + sumulas + cartoes
        self.assertTrue(all(os.path.exists(path) for path in result["generated"]))

    def test_export_tournament_minutes(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        sections = self.export_service._tournament_minutes_sections(tournament_id)
        titles = [title for title, _headers, _rows in sections]
        self.assertEqual("Ata final do torneio", titles[0])
        self.assertTrue(any("assinatura" in title.lower() for title in titles))
        out_path = Path(self.temp_dir.name) / "ata.xlsx"
        self.export_service.export_tournament_minutes(tournament_id, out_path)
        self.assertTrue(out_path.exists())
        self.assertGreater(out_path.stat().st_size, 0)


class ArbiterPanelNetNewTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.pairing_service = PairingService(self.db)
        self.tournament_service = TournamentService(self.db)
        self.team_service = TeamService(self.db)
        self.export_service = ExportService(self.db, self.pairing_service)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _open_round_tournament(self) -> int:
        tournament_id = self.db.create_tournament(
            name="Aberto",
            location="Sala",
            rounds_count=3,
            time_control="90+30",
            start_date="2026-06-02",
            end_date="2026-06-02",
        )
        for index in range(4):
            self.db.create_player(
                tournament_id=tournament_id,
                name=f"Jogador {index}",
                rating=1500 - index,
                category="Absoluto" if index % 2 == 0 else "Sub-12",
            )
        self.pairing_service.generate_next_round(tournament_id)
        return tournament_id

    def _team_tournament(self) -> tuple[int, int]:
        tournament_id = self.tournament_service.create_tournament(
            {"name": "Equipes", "competition_type": "team", "rounds_count": "3", "bye_points": "1"}
        )
        self.tournament_service.save_profile(
            tournament_id,
            {
                "name": "Equipes",
                "competition_type": "team",
                "scope": "standalone",
                "rounds_count": "3",
                "bye_points": "1",
            },
            {
                "team_boards_count": "4",
                "team_match_win_points": "2",
                "team_match_draw_points": "1",
                "team_match_loss_points": "0",
                "team_pairing_method": "swiss",
                "team_standing_primary": "match_points",
                "team_standing_secondary": "game_points",
            },
            [],
        )
        for team_index in range(4):
            team_id = self.team_service.create_team(
                tournament_id,
                {"name": f"Equipe {team_index + 1}", "club": f"Clube {team_index + 1}", "captain": f"Cap {team_index + 1}"},
            )
            for board_number in range(1, 5):
                player_id = self.db.create_player(
                    tournament_id,
                    name=f"E{team_index + 1} J{board_number}",
                    rating=2200 - team_index * 100 - board_number * 10,
                    club=f"Clube {team_index + 1}",
                )
                self.team_service.add_player(team_id, player_id, board_number=str(board_number), role="starter")
        round_data = self.pairing_service.generate_next_round(tournament_id)
        return tournament_id, int(round_data["id"])

    # N1 - Ata por categoria
    def test_category_winners_section(self) -> None:
        tournament_id = self._open_round_tournament()
        title, headers, rows = self.export_service._category_winners_section(tournament_id)
        categories = {row[0] for row in rows}
        self.assertIn("Absoluto", categories)
        self.assertIn("Sub-12", categories)
        self.assertNotIn("", categories)

    def test_minutes_includes_category_section(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        titles = [title for title, _h, _r in self.export_service._tournament_minutes_sections(tournament_id)]
        self.assertIn("Vencedores por categoria", titles)

    # N2 - Boletim da rodada
    def test_round_bulletin(self) -> None:
        tournament_id = load_tournament_fixture(
            "individual_8_players_3_rounds", self.db, self.pairing_service
        )
        latest = sorted(self.db.list_rounds(tournament_id), key=lambda item: int(item["number"]))[-1]
        sections = self.export_service._round_bulletin_sections(int(latest["id"]))
        titles = [title for title, _h, _r in sections]
        self.assertTrue(titles[0].startswith("Boletim da rodada"))
        self.assertTrue(any("Classificacao apos a rodada" in title for title in titles))
        self.assertIn("Destaques", titles)
        out_path = Path(self.temp_dir.name) / "boletim.xlsx"
        self.export_service.export_round_bulletin(int(latest["id"]), out_path)
        self.assertTrue(out_path.exists())

    # N3 - Checklist de fechamento
    def test_closing_checklist_open_round(self) -> None:
        tournament_id = self._open_round_tournament()
        checklist = self.pairing_service.closing_checklist(tournament_id)
        results_item = next(item for item in checklist if item["action"] == "pending_results")
        self.assertFalse(results_item["ok"])
        ready_item = next(item for item in checklist if item["action"] == "ready_to_close")
        self.assertFalse(ready_item["ok"])
        rounds = self.db.list_rounds(tournament_id)
        for pairing in self.db.get_pairings_for_round(int(rounds[-1]["id"])):
            if not pairing.get("is_bye"):
                self.pairing_service.update_result(tournament_id, int(pairing["id"]), "1-0")
        checklist_after = self.pairing_service.closing_checklist(tournament_id)
        self.assertTrue(next(item for item in checklist_after if item["action"] == "pending_results")["ok"])
        self.assertTrue(next(item for item in checklist_after if item["action"] == "ready_to_close")["ok"])

    def test_closing_checklist_without_rounds(self) -> None:
        tournament_id = self.db.create_tournament(
            name="Vazio", location="", rounds_count=3, time_control="", start_date="", end_date=""
        )
        checklist = self.pairing_service.closing_checklist(tournament_id)
        self.assertEqual(1, len(checklist))
        self.assertEqual("initial_call", checklist[0]["action"])
        self.assertFalse(checklist[0]["ok"])

    # N4 - Pacote da rodada para equipes
    def test_export_round_package_team(self) -> None:
        _tournament_id, round_id = self._team_tournament()
        dest = Path(self.temp_dir.name) / "pacote_equipes"
        result = self.export_service.export_round_package(round_id, dest)
        self.assertTrue(all(os.path.exists(path) for path in result["generated"]))
        self.assertTrue(any("cartoes" in path for path in result["generated"]))
        self.assertTrue(any("sumulas" in path for path in result["generated"]))


class ArbiterPanelExtrasTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        self.db = Database(base / "albericus.db", backup_dir=base / "backups")
        self.pairing_service = PairingService(self.db)
        self.tournament_service = TournamentService(self.db)
        self.team_service = TeamService(self.db)
        self.export_service = ExportService(self.db, self.pairing_service)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _individual_with_categories(self) -> int:
        tournament_id = self.db.create_tournament(
            name="Podio", location="Sala", rounds_count=3, time_control="", start_date="2026-06-02", end_date="2026-06-02"
        )
        for index in range(4):
            self.db.create_player(
                tournament_id=tournament_id,
                name=f"Jogador {index}",
                rating=1600 - index * 10,
                category="Absoluto" if index % 2 == 0 else "Sub-12",
            )
        self.pairing_service.generate_next_round(tournament_id)
        return tournament_id

    def _team_tournament(self) -> int:
        tournament_id = self.tournament_service.create_tournament(
            {"name": "Equipes", "competition_type": "team", "rounds_count": "3", "bye_points": "1"}
        )
        self.tournament_service.save_profile(
            tournament_id,
            {
                "name": "Equipes",
                "competition_type": "team",
                "scope": "standalone",
                "rounds_count": "3",
                "bye_points": "1",
            },
            {
                "team_boards_count": "4",
                "team_match_win_points": "2",
                "team_match_draw_points": "1",
                "team_match_loss_points": "0",
                "team_pairing_method": "swiss",
                "team_standing_primary": "match_points",
                "team_standing_secondary": "game_points",
            },
            [],
        )
        for team_index in range(4):
            team_id = self.team_service.create_team(
                tournament_id,
                {"name": f"Equipe {team_index + 1}", "club": f"Clube {team_index + 1}", "captain": "Cap"},
            )
            for board_number in range(1, 5):
                player_id = self.db.create_player(
                    tournament_id,
                    name=f"E{team_index + 1} J{board_number}",
                    rating=2200 - team_index * 100 - board_number * 10,
                )
                self.team_service.add_player(team_id, player_id, board_number=str(board_number), role="starter")
        self.pairing_service.generate_next_round(tournament_id)
        return tournament_id

    # M1 - Poster do podio
    def test_podium_data(self) -> None:
        tournament_id = self._individual_with_categories()
        data = self.export_service._podium_data(tournament_id)
        self.assertFalse(data["is_team"])
        self.assertEqual(3, len(data["top"]))
        self.assertEqual(1, data["top"][0]["rank"])
        category_labels = {row["category"] for row in data["categories"]}
        self.assertIn("Absoluto", category_labels)
        self.assertIn("Sub-12", category_labels)

    def test_export_podium_pdf(self) -> None:
        tournament_id = self._individual_with_categories()
        out_path = Path(self.temp_dir.name) / "podio.pdf"
        self.export_service.export_podium(tournament_id, out_path)
        self.assertTrue(out_path.exists())
        self.assertGreater(out_path.stat().st_size, 0)
        with self.assertRaises(AppError):
            self.export_service.export_podium(tournament_id, Path(self.temp_dir.name) / "podio.csv")

    # M3 - Checklist especifico para equipes
    def test_team_closing_checklist_has_lineups(self) -> None:
        tournament_id = self._team_tournament()
        checklist = self.pairing_service.closing_checklist(tournament_id)
        actions = [item["action"] for item in checklist]
        self.assertIn("lineups", actions)
        self.assertIn("ready_to_close", actions)


if __name__ == "__main__":
    unittest.main()
