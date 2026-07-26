from __future__ import annotations
import csv
import html
import io
import json
import logging
import math
import re
import secrets
import shutil
import sqlite3
import unicodedata
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Mapping, TYPE_CHECKING
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse
from urllib.request import Request, urlopen

from src.core.database import BASE_DIR, DEFAULT_CERTIFICATE_TEMPLATES, Database
from src.services.access_export import access_driver_available, write_accdb, write_csv_bundle
from src.services.constants import *
from src.services.fide_norms import build_norm_report
from src.services.fide_rating import build_fide_report_rows
from src.services.list_layouts import STANDINGS_COLUMNS, resolve_column_specs, resolve_columns
from src.services.prizes import PRIZE_KINDS, PRIZE_POLICIES, allocate_prizes
from src.services.trf_import import build_trf_rounds, parse_trf

if TYPE_CHECKING:
    from src.services.club_service import ClubService
    from src.services.member_service import MemberService, GuardianService
    from src.services.tournament_service import TournamentService, RefereeService, TeamService
    from src.services.pairing_service import PairingService
    from src.services.finance_service import FinanceService
    from src.services.event_service import EventService, CalendarService
    from src.services.education_service import TrainingService, ExerciseService, LibraryService, LearningLevelService
    from src.services.inventory_service import InventoryService
    from src.services.security_service import SecurityService
    from src.services.export_service import ExportService, ImportService, CertificateService
    from src.services.rating_service import OfficialRatingService, InternalRatingService
    from src.services.dashboard_service import DashboardService, CommunicationService

logger = logging.getLogger(__name__)


class ReportWritersMixin:
    @staticmethod
    def _format_report_number(value: Any) -> str:
        try:
            numeric = float(value or 0)
        except (TypeError, ValueError):
            return str(value or "")
        if numeric.is_integer():
            return str(int(numeric))
        return f"{numeric:.2f}".rstrip("0").rstrip(".")

    def _write_report(
        self,
        file_path: str | Path,
        title: str,
        headers: list[str],
        rows: list[list[Any]],
        widths: list[int] | None = None,
    ) -> None:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension == ".csv":
            self._write_csv(path, headers, rows)
            logger.info("Relatório exportado em CSV: %s", path)
            return
        if extension == ".xlsx":
            self._write_xlsx(path, title, headers, rows, widths)
            logger.info("Relatório exportado em XLSX: %s", path)
            return
        if extension == ".pdf":
            self._write_pdf(path, title, headers, rows)
            logger.info("Relatório exportado em PDF: %s", path)
            return
        raise AppError("Formato não suportado. Use .csv, .xlsx ou .pdf.")

    def _write_multi_report(
        self,
        file_path: str | Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension == ".csv":
            self._write_multi_csv(path, sections)
            logger.info("Relatório composto exportado em CSV: %s", path)
            return
        if extension == ".xlsx":
            self._write_multi_xlsx(path, sections)
            logger.info("Relatório composto exportado em XLSX: %s", path)
            return
        if extension == ".pdf":
            self._write_multi_pdf(path, sections)
            logger.info("Relatório composto exportado em PDF: %s", path)
            return
        raise AppError("Formato não suportado. Use .csv, .xlsx ou .pdf.")

    @staticmethod
    def _write_csv(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
        with path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file, delimiter=";")
            writer.writerow(headers)
            writer.writerows(rows)

    @staticmethod
    def _write_multi_csv(
        path: Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        with path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file, delimiter=";")
            for index, (title, headers, rows) in enumerate(sections):
                if index:
                    writer.writerow([])
                writer.writerow([title])
                writer.writerow(headers)
                writer.writerows(rows)

    @staticmethod
    def _write_xlsx(
        path: Path,
        title: str,
        headers: list[str],
        rows: list[list[Any]],
        widths: list[int] | None = None,
    ) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise AppError("Instale openpyxl para exportar Excel.") from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title[:31]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for column_cells in sheet.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)
        # Larguras explicitas do layout sobrescrevem o auto-dimensionamento.
        if widths:
            for index, width in enumerate(widths):
                if width and width > 0:
                    sheet.column_dimensions[get_column_letter(index + 1)].width = int(width)
        workbook.save(path)

    @staticmethod
    def _write_multi_xlsx(
        path: Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as exc:
            raise AppError("Instale openpyxl para exportar Excel.") from exc

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        used_names: set[str] = set()

        for title, headers, rows in sections:
            sheet_name = ReportWritersMixin._unique_sheet_name(title, used_names)
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in rows:
                sheet.append(row)
            for column_cells in sheet.columns:
                length = max(len(str(cell.value or "")) for cell in column_cells)
                sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 42)

        workbook.save(path)

    @staticmethod
    def _unique_sheet_name(title: str, used_names: set[str]) -> str:
        invalid = set("[]:*?/\\")
        base = "".join("_" if character in invalid else character for character in title).strip()
        base = (base or "Relatório")[:31]
        name = base
        counter = 2
        while name in used_names:
            suffix = f" {counter}"
            name = f"{base[:31 - len(suffix)]}{suffix}"
            counter += 1
        used_names.add(name)
        return name

    @staticmethod
    def _write_pdf(path: Path, title: str, headers: list[str], rows: list[list[Any]]) -> None:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise AppError("Instale reportlab para exportar PDF.") from exc

        document = SimpleDocTemplate(str(path), pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        data = [headers] + [[str(value) for value in row] for row in rows]
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9CA3AF")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        document.build([Paragraph(title, styles["Title"]), Spacer(1, 12), table])

    @staticmethod
    def _write_multi_pdf(
        path: Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise AppError("Instale reportlab para exportar PDF.") from exc

        document = SimpleDocTemplate(str(path), pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        story = []
        for index, (title, headers, rows) in enumerate(sections):
            if index:
                story.append(PageBreak())
            story.append(Paragraph(title, styles["Title"]))
            story.append(Spacer(1, 12))
            data = [headers] + [[str(value) for value in row] for row in rows]
            table = Table(data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9CA3AF")),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ]
                )
            )
            story.append(table)
        document.build(story)
