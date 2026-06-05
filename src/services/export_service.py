from __future__ import annotations
import csv
import html
import io
import json
import logging
import math
import re
import secrets
import shutil
import sqlite3
import unicodedata
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Mapping, TYPE_CHECKING
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse
from urllib.request import Request, urlopen

from src.core.database import BASE_DIR, DEFAULT_CERTIFICATE_TEMPLATES, Database
from src.services.access_export import access_driver_available, write_accdb, write_csv_bundle
from src.services.constants import *
from src.services.fide_norms import build_norm_report
from src.services.fide_rating import build_fide_report_rows
from src.services.list_layouts import STANDINGS_COLUMNS, resolve_column_specs, resolve_columns
from src.services.prizes import PRIZE_KINDS, PRIZE_POLICIES, allocate_prizes
from src.services.trf_import import build_trf_rounds, parse_trf

# ImportService e CertificateService foram extraidos para modulos proprios
# (import_service.py / certificate_service.py). Reexportados aqui para
# preservar a API publica historica deste modulo:
#     from src.services.export_service import ImportService, CertificateService
from src.services.certificate_service import CertificateService
from src.services.import_service import (
    REGISTRATION_FORM_QUESTIONS,
    REGISTRATION_IMPORT_FIELDS,
    ImportService,
)

if TYPE_CHECKING:
    from src.services.club_service import ClubService
    from src.services.member_service import MemberService, GuardianService
    from src.services.tournament_service import TournamentService, RefereeService, TeamService
    from src.services.pairing_service import PairingService
    from src.services.finance_service import FinanceService
    from src.services.event_service import EventService, CalendarService
    from src.services.education_service import TrainingService, ExerciseService, LibraryService, LearningLevelService
    from src.services.inventory_service import InventoryService
    from src.services.security_service import SecurityService
    from src.services.export_service import ExportService, ImportService, CertificateService
    from src.services.rating_service import OfficialRatingService, InternalRatingService
    from src.services.dashboard_service import DashboardService, CommunicationService

logger = logging.getLogger(__name__)


class ExportService:
    def __init__(self, db: Database, pairing_service: PairingService) -> None:
        self.db = db
        self.pairing_service = pairing_service

    def export_players(self, tournament_id: int, file_path: str | Path) -> None:
        title, headers, rows = self._players_section(tournament_id)
        self._write_report(file_path, title, headers, rows)

    def export_player_import_template(self, file_path: str | Path) -> None:
        self._write_report(
            file_path,
            "Modelo de importacao de jogadores",
            [
                "nome",
                "sobrenome",
                "rating",
                "clube",
                "categoria",
                "fide_id",
                "cbx_id",
                "lbx_id",
                "nascimento",
                "sexo",
            ],
            [["Ana", "Silva", "1500", "Clube A", "ABS", "", "", "", "2012-05-10", "F"]],
        )

    def export_online_registration_template(self, file_path: str | Path) -> None:
        self._write_report(
            file_path,
            "Modelo de inscricoes online",
            [
                "Nome completo do jogador",
                "Rating",
                "Clube/Cidade",
                "Categoria",
                "Data de nascimento",
                "FIDE ID",
                "CBX ID",
                "LBX ID",
                "Sexo",
            ],
            [["Ana Silva", "1500", "Clube A", "ABS", "2012-05-10", "", "", "", "F"]],
        )

    def export_registration_form(
        self, file_path: str | Path, tournament_id: int | None = None
    ) -> dict[str, str]:
        """Gera o formulario de inscricao padronizado (REG-01).

        Produz dois artefatos com o mesmo nome base:
        - ``.gs``: script Google Apps Script colavel em script.google.com que
          cria o formulario com as perguntas padronizadas;
        - ``.json``: definicao reutilizavel dos campos.

        Os titulos das perguntas sao escolhidos para que o CSV de respostas
        importe direto pelo fluxo de inscricoes online, sem mapeamento manual.
        """
        title = "Inscricao no torneio"
        if tournament_id is not None:
            tournament = self.db.get_tournament(int(tournament_id))
            if tournament and tournament.get("name"):
                title = f"Inscricao - {tournament['name']}"

        path = Path(file_path)
        if path.suffix.lower() not in {".gs", ".js"}:
            path = path.with_suffix(".gs")
        path.parent.mkdir(parents=True, exist_ok=True)

        path.write_text(self._registration_form_script(title), encoding="utf-8")
        definition_path = path.with_suffix(".json")
        definition_path.write_text(
            json.dumps(
                {"title": title, "questions": REGISTRATION_FORM_QUESTIONS},
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        logger.info("Formulario de inscricao padronizado gerado em %s e %s", path, definition_path)
        return {"script_path": str(path), "definition_path": str(definition_path)}

    @staticmethod
    def _registration_form_script(title: str) -> str:
        def js_text(value: str) -> str:
            return value.replace("\\", "\\\\").replace("'", "\\'")

        lines: list[str] = [
            "/**",
            " * Albericus - Formulario de inscricao padronizado (REG-01).",
            " *",
            " * Passo a passo:",
            " *  1. Acesse https://script.google.com e crie um novo projeto.",
            " *  2. Cole todo este codigo e salve.",
            " *  3. Execute a funcao criarFormularioInscricao e autorize o acesso.",
            " *  4. O link do formulario aparece no registro de execucao (Ver > Registros).",
            " *  5. No formulario, vincule as respostas a uma planilha.",
            " *  6. Baixe a planilha como CSV (ou publique como CSV) e importe em",
            " *     Jogadores > Importar inscricoes online / Importar link Forms/Sheets.",
            " *",
            " * Os titulos das perguntas ja seguem o padrao do importador; nao os",
            " * renomeie para manter a importacao automatica sem mapeamento.",
            " */",
            "function criarFormularioInscricao() {",
            f"  var form = FormApp.create('{js_text(title)}');",
            "  form.setDescription('Inscricao gerada pelo Albericus. Preencha os dados do jogador.');",
            "  form.setCollectEmail(false);",
            "",
        ]
        for index, question in enumerate(REGISTRATION_FORM_QUESTIONS):
            title_js = js_text(str(question["title"]))
            required = "true" if question.get("required") else "false"
            qtype = question.get("type")
            if qtype == "date":
                lines.append(
                    f"  form.addDateItem().setTitle('{title_js}').setRequired({required});"
                )
            elif qtype == "choice":
                choices = ", ".join(f"'{js_text(str(choice))}'" for choice in question.get("choices", []))
                lines.append(
                    f"  form.addMultipleChoiceItem().setTitle('{title_js}')"
                    f".setChoiceValues([{choices}]).setRequired({required});"
                )
            elif qtype == "integer":
                var = f"q{index}"
                lines.append(
                    f"  var {var} = form.addTextItem().setTitle('{title_js}').setRequired({required});"
                )
                lines.append(
                    f"  {var}.setValidation(FormApp.createTextValidation()"
                    ".setHelpText('Informe um numero.').requireNumber().build());"
                )
            else:
                lines.append(
                    f"  form.addTextItem().setTitle('{title_js}').setRequired({required});"
                )
        lines.append("")
        lines.append("  Logger.log('Formulario criado: ' + form.getPublishedUrl());")
        lines.append("  Logger.log('Edicao: ' + form.getEditUrl());")
        lines.append("}")
        lines.append("")
        return "\n".join(lines)

    # --- REG-01: link pre-preenchido do Google Forms ----------------------

    @staticmethod
    def parse_prefill_link(link: str) -> dict[str, Any]:
        """Extrai a URL base e os campos ``entry.*`` de um link pre-preenchido
        do Google Forms (obtido em 'Receber link preenchido automaticamente')."""
        text = (link or "").strip()
        parsed = urlparse(text)
        if parsed.scheme not in {"http", "https"} or "docs.google.com" not in parsed.netloc:
            raise AppError("Cole um link valido do Google Forms (pre-preenchido).")
        base_url = urlunparse((parsed.scheme, parsed.netloc, parsed.path, "", "", ""))
        query = parse_qs(parsed.query, keep_blank_values=True)
        entries = [
            (key, values[0] if values else "")
            for key, values in query.items()
            if key.startswith("entry.")
        ]
        if not entries:
            raise AppError(
                "Nenhum campo encontrado no link. Use 'Receber link preenchido "
                "automaticamente', preencha o campo do torneio e cole o link gerado."
            )
        return {"base_url": base_url, "entries": entries}

    @staticmethod
    def build_registration_prefill_url(
        base_url: str, tournament_entry: str, tournament_name: str
    ) -> str:
        """Monta o link de inscricao com o nome do torneio ja preenchido e os
        campos do jogador em branco."""
        base = (base_url or "").strip()
        if not base:
            raise AppError("Formulario de inscricao nao configurado.")
        params = {"usp": "pp_url"}
        entry = (tournament_entry or "").strip()
        name = (tournament_name or "").strip()
        if entry and name:
            params[entry] = name
        separator = "&" if urlparse(base).query else "?"
        return base + separator + urlencode(params)

    def registration_form_config(self) -> dict[str, str]:
        settings = self.db.get_app_settings()
        return {
            "base_url": str(settings.get("registration_form_prefill_base") or ""),
            "tournament_entry": str(settings.get("registration_form_tournament_entry") or ""),
        }

    def save_registration_form_config(self, prefill_link: str, tournament_entry: str) -> dict[str, str]:
        parsed = self.parse_prefill_link(prefill_link)
        entry = (tournament_entry or "").strip()
        self.db.save_app_settings(
            {
                "registration_form_prefill_base": parsed["base_url"],
                "registration_form_tournament_entry": entry,
            }
        )
        return {"base_url": parsed["base_url"], "tournament_entry": entry}

    def registration_prefill_url(self, tournament_id: int) -> str:
        config = self.registration_form_config()
        if not config["base_url"]:
            raise AppError(
                "Formulario de inscricao nao configurado. Use 'Configurar formulario "
                "(link)' e cole o link pre-preenchido do Google Forms."
            )
        tournament = self.db.get_tournament(int(tournament_id))
        name = str(tournament.get("name") or "") if tournament else ""
        return self.build_registration_prefill_url(
            config["base_url"], config["tournament_entry"], name
        )

    def export_member_import_template(self, file_path: str | Path) -> None:
        self._write_report(
            file_path,
            "Modelo de importacao de membros",
            [
                "nome",
                "sobrenome",
                "rating",
                "escola",
                "turma",
                "tipo",
                "status",
                "nascimento",
                "email",
                "telefone",
                "responsavel",
                "telefone responsavel",
            ],
            [["Ana", "Silva", "1500", "Escola A", "Turma 1", "aluno", "ativo", "2012-05-10", "", "", "", ""]],
        )

    def export_initial_player_list(self, tournament_id: int, file_path: str | Path) -> None:
        title, headers, rows = self._initial_player_list_section(tournament_id)
        path = Path(file_path)
        if path.suffix.lower() == ".pdf":
            self._write_initial_player_list_pdf(path, title, headers, rows)
            logger.info("%s jogadores exportados na lista de chamada PDF: %s", len(rows), path)
            return
        self._write_report(file_path, title, headers, rows)

    def export_pairings(self, round_id: int, file_path: str | Path) -> None:
        path = Path(file_path)
        round_data = self.db.get_round(round_id)
        if not round_data:
            raise AppError("Rodada nao encontrada.")
        tournament = self.db.get_tournament(int(round_data["tournament_id"]))
        if path.suffix.lower() == ".pdf" and tournament and tournament.get("competition_type") != "team":
            self._write_pairings_wall_pdf(path, tournament, round_data)
            logger.info("Folha de mural da rodada %s exportada em PDF: %s", round_data["number"], path)
            return
        title, headers, rows = self._pairings_section(round_id)
        self._write_report(file_path, title, headers, rows)

    def export_scoresheets(self, round_id: int, file_path: str | Path) -> None:
        path = Path(file_path)
        if path.suffix.lower() != ".pdf":
            raise AppError("Sumulas de mesa devem ser exportadas em PDF.")
        round_data = self.db.get_round(round_id)
        if not round_data:
            raise AppError("Rodada nao encontrada.")
        tournament = self.db.get_tournament(int(round_data["tournament_id"]))
        if not tournament:
            raise AppError("Torneio nao encontrado.")
        scoresheets = self._scoresheet_rows(round_data, tournament)
        if not scoresheets:
            raise AppError("Nao ha mesas validas para gerar sumulas.")
        self._write_scoresheets_pdf(path, tournament, round_data, scoresheets)
        logger.info("%s sumulas exportadas em PDF: %s", len(scoresheets), path)

    def export_table_cards(
        self,
        file_path: str | Path,
        start_board: int,
        end_board: int,
        round_id: int | None = None,
        include_qr: bool = False,
    ) -> None:
        path = Path(file_path)
        if path.suffix.lower() != ".pdf":
            raise AppError("Cartoes de mesa devem ser exportados em PDF.")
        if int(start_board) <= 0 or int(end_board) < int(start_board):
            raise AppError("Informe um intervalo valido de mesas.")
        if int(end_board) - int(start_board) + 1 > 500:
            raise AppError("Gere no maximo 500 cartoes de mesa por arquivo.")
        round_data = self.db.get_round(int(round_id)) if round_id else None
        tournament = self.db.get_tournament(int(round_data["tournament_id"])) if round_data else None
        if include_qr and not round_data:
            raise AppError("Selecione uma rodada para incluir QR nos cartoes.")
        if include_qr and tournament and tournament.get("competition_type") == "team":
            raise AppError("QR nos cartoes esta disponivel apenas para torneios individuais.")
        self._write_table_cards_pdf(path, int(start_board), int(end_board), tournament, round_data, include_qr)
        logger.info("%s cartoes de mesa exportados em PDF: %s", int(end_board) - int(start_board) + 1, path)

    def export_teams(self, tournament_id: int, file_path: str | Path) -> None:
        self._write_multi_report(
            file_path,
            [
                self._teams_section(tournament_id),
                self._team_rosters_section(tournament_id),
            ],
        )

    def export_team_lineups(self, tournament_id: int, file_path: str | Path) -> None:
        self._write_multi_report(
            file_path,
            [
                self._team_lineups_section(tournament_id),
                self._team_substitutions_section(tournament_id),
            ],
        )

    def export_all_rounds(self, tournament_id: int, file_path: str | Path) -> None:
        sections = []
        for round_data in sorted(self.db.list_rounds(tournament_id), key=lambda item: item["number"]):
            sections.append(self._pairings_section(round_data["id"]))
        if not sections:
            raise AppError("Nao ha rodadas para exportar.")
        self._write_multi_report(file_path, sections)

    def export_standings(self, tournament_id: int, file_path: str | Path) -> None:
        title, headers, rows = self._standings_section(tournament_id)
        self._write_report(
            file_path, title, headers, rows, widths=self._standings_column_widths(tournament_id)
        )

    def _standings_column_widths(self, tournament_id: int) -> list[int] | None:
        tournament = self.db.get_tournament(tournament_id)
        if tournament and tournament.get("competition_type") == "team":
            return None
        specs = resolve_column_specs(
            self.db.get_report_layout_columns(tournament_id, "standings"), "standings"
        )
        widths = [int(spec.get("width") or 0) for spec in specs]
        return widths if any(width > 0 for width in widths) else None

    def export_crosstable(self, tournament_id: int, file_path: str | Path) -> None:
        path = Path(file_path)
        title, headers, rows = self._crosstable_section(tournament_id)
        if path.suffix.lower() == ".html":
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(self._crosstable_html(title, headers, rows), encoding="utf-8")
            logger.info("Tabela cruzada exportada em HTML: %s", path)
            return
        self._write_report(path, title, headers, rows)

    def rating_fee_summary(self, tournament_id: int) -> dict[str, Any]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        settings = self.db.get_tournament_settings(tournament_id) or {}
        players = self.db.list_players(tournament_id, active_only=False)
        base_specs = [
            ("FIDE", "fide_id", "international_rating", "rating_fee_fide"),
            ("CBX", "cbx_id", "national_rating", "rating_fee_cbx"),
            ("LBX", "lbx_id", "rating", "rating_fee_lbx"),
        ]
        bases = []
        for label, id_field, rating_field, fee_field in base_specs:
            identified = [player for player in players if str(player.get(id_field) or "").strip()]
            rated = [player for player in identified if int(player.get(rating_field) or 0) > 0]
            unit_fee = float(settings.get(fee_field) or 0.0)
            bases.append(
                {
                    "base": label,
                    "id_field": id_field,
                    "rating_field": rating_field,
                    "identified": len(identified),
                    "rated": len(rated),
                    "unrated": len(identified) - len(rated),
                    "unit_fee": unit_fee,
                    "subtotal": round(len(identified) * unit_fee, 2),
                }
            )
        players_with_base = sum(
            1
            for player in players
            if any(str(player.get(field) or "").strip() for field in ("fide_id", "cbx_id", "lbx_id"))
        )
        rated_players = sum(
            1
            for player in players
            if max(
                int(player.get("rating") or 0),
                int(player.get("national_rating") or 0),
                int(player.get("international_rating") or 0),
            )
            > 0
        )
        return {
            "tournament_id": int(tournament_id),
            "tournament_name": tournament["name"],
            "players_count": len(players),
            "players_with_base": players_with_base,
            "players_without_base": len(players) - players_with_base,
            "rated_players": rated_players,
            "unrated_players": len(players) - rated_players,
            "bases": bases,
            "total": round(sum(float(item["subtotal"]) for item in bases), 2),
            "players": players,
        }

    def export_rating_fee_report(self, tournament_id: int, file_path: str | Path) -> None:
        path = Path(file_path)
        if path.suffix.lower() not in {".xlsx", ".pdf"}:
            raise AppError("Relatorio de taxas de rating deve ser exportado em XLSX ou PDF.")
        self._write_multi_report(path, self._rating_fee_sections(tournament_id))

    def export_fide_rating_report(
        self, tournament_id: int, file_path: str | Path, rating_type: str = "fide"
    ) -> None:
        self._write_multi_report(
            Path(file_path), self._fide_rating_sections(tournament_id, rating_type)
        )

    def export_prize_report(self, tournament_id: int, file_path: str | Path) -> None:
        self._write_multi_report(Path(file_path), self._prize_sections(tournament_id))

    def export_federation_statistics(self, tournament_id: int, file_path: str | Path) -> None:
        self._write_multi_report(Path(file_path), self._federation_sections(tournament_id))

    def export_game_statistics(self, tournament_id: int, file_path: str | Path) -> None:
        self._write_multi_report(Path(file_path), self._game_statistics_sections(tournament_id))

    def export_player_cards(self, tournament_id: int, file_path: str | Path) -> None:
        self._write_multi_report(Path(file_path), self._player_cards_sections(tournament_id))

    def export_norm_report(self, tournament_id: int, file_path: str | Path) -> None:
        self._write_multi_report(Path(file_path), self._norm_sections(tournament_id))

    def export_arbiter_norm_report(self, tournament_id: int, file_path: str | Path) -> None:
        self._write_multi_report(Path(file_path), self._arbiter_norm_sections(tournament_id))

    def export_access(self, tournament_id: int, dest_dir: str | Path) -> dict[str, Any]:
        """Exporta o torneio para Access (Fase J).

        Gera sempre um pacote importável (CSV por tabela + ``schema.ini``) e, quando
        o driver ACE estiver instalado, também um ``.accdb`` real. Reaproveita os
        construtores de seção (jogadores, classificação, emparceiramentos).
        """
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        directory = Path(dest_dir)

        tables: dict[str, tuple[list[str], list[list[Any]]]] = {}
        _, player_headers, player_rows = self._players_section(tournament_id)
        tables["Jogadores"] = (player_headers, [list(row) for row in player_rows])
        _, standings_headers, standings_rows = self._standings_section(tournament_id)
        tables["Classificacao"] = (standings_headers, [list(row) for row in standings_rows])

        pairing_headers: list[str] | None = None
        pairing_rows: list[list[Any]] = []
        for round_data in sorted(self.db.list_rounds(tournament_id), key=lambda item: item["number"]):
            try:
                _, headers, rows = self._pairings_section(round_data["id"])
            except Exception:
                continue
            if pairing_headers is None:
                pairing_headers = ["Rodada", *headers]
            for row in rows:
                pairing_rows.append([round_data["number"], *list(row)])
        if pairing_headers is not None:
            tables["Emparceiramentos"] = (pairing_headers, pairing_rows)

        bundle = write_csv_bundle(tables, directory)

        slug = re.sub(r"[^A-Za-z0-9._-]+", "_", str(tournament.get("name") or "torneio").strip()).strip("_") or "torneio"
        accdb_path = directory / f"{slug}_access.accdb"
        accdb = str(accdb_path) if write_accdb(tables, accdb_path) else None

        logger.info(
            "Exportacao Access do torneio %s: %s tabelas, accdb=%s",
            tournament_id,
            len(tables),
            bool(accdb),
        )
        return {
            "csv_paths": bundle["csv_paths"],
            "schema_ini": bundle["schema_ini"],
            "accdb": accdb,
            "driver_available": access_driver_available(),
            "tables": list(tables.keys()),
        }

    def export_round_package(
        self,
        round_id: int,
        dest_dir: str | Path,
        *,
        wall: bool = True,
        scoresheets: bool = True,
        cards: bool = True,
    ) -> dict[str, Any]:
        """Pacote da rodada em um passo: mural + sumulas + cartoes numa pasta (PDF).

        Reaproveita os geradores existentes e isola erros (um documento que falha
        nao impede os demais). Pensado para o arbitro afixar/distribuir de uma vez.
        """
        round_data = self.db.get_round(round_id)
        if not round_data:
            raise AppError("Rodada nao encontrada.")
        tournament = self.db.get_tournament(int(round_data["tournament_id"]))
        if not tournament:
            raise AppError("Torneio nao encontrado.")
        directory = Path(dest_dir)
        directory.mkdir(parents=True, exist_ok=True)
        slug = re.sub(r"[^A-Za-z0-9._-]+", "_", str(tournament.get("name") or "torneio").strip()).strip("_") or "torneio"
        base = f"{slug}_rodada_{round_data['number']}"

        generated: list[str] = []
        errors: list[str] = []
        if wall:
            try:
                path = directory / f"{base}_mural.pdf"
                self.export_pairings(round_id, path)
                generated.append(str(path))
            except Exception as exc:
                errors.append(f"Mural: {exc}")
        if scoresheets:
            try:
                path = directory / f"{base}_sumulas.pdf"
                self.export_scoresheets(round_id, path)
                generated.append(str(path))
            except Exception as exc:
                errors.append(f"Sumulas: {exc}")
        if cards:
            try:
                if tournament.get("competition_type") == "team":
                    # Em equipes os tabuleiros ficam em team_boards (nao na tabela pairings):
                    # o total de cartoes = soma dos tabuleiros dos confrontos sem bye.
                    total_boards = sum(
                        len(self.db.list_team_boards(int(match["id"])))
                        for match in self.db.list_team_matches_for_round(round_id)
                        if not match.get("is_bye")
                    )
                else:
                    pairings = self.db.get_pairings_for_round(round_id)
                    total_boards = max((int(item.get("board_number") or 0) for item in pairings), default=0)
                if total_boards <= 0:
                    raise AppError("Rodada sem mesas para cartoes.")
                path = directory / f"{base}_cartoes.pdf"
                self.export_table_cards(path, 1, total_boards, round_id, False)
                generated.append(str(path))
            except Exception as exc:
                errors.append(f"Cartoes: {exc}")

        logger.info(
            "Pacote da rodada %s do torneio %s: %s documentos, %s erros",
            round_data["number"],
            tournament["id"],
            len(generated),
            len(errors),
        )
        return {"generated": generated, "errors": errors, "count": len(generated)}

    def export_tournament_minutes(self, tournament_id: int, file_path: str | Path) -> None:
        """Ata final do torneio: documento unico de encerramento (Fase painel)."""
        self._write_multi_report(Path(file_path), self._tournament_minutes_sections(tournament_id))

    def export_round_bulletin(self, round_id: int, file_path: str | Path) -> None:
        """Boletim/press-release da rodada: resultados + classificacao + destaques."""
        self._write_multi_report(Path(file_path), self._round_bulletin_sections(round_id))

    def export_podium(self, tournament_id: int, file_path: str | Path) -> None:
        """Poster/diploma do podio (PDF A4): top 3 + campeoes por categoria."""
        path = Path(file_path)
        if path.suffix.lower() != ".pdf":
            raise AppError("Poster do podio deve ser exportado em PDF.")
        self._write_podium_poster_pdf(path, self._podium_data(tournament_id))
        logger.info("Poster do podio gerado: %s", path)

    def export_tiebreak_report(self, tournament_id: int, file_path: str | Path) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") == "team":
            raise AppError("Relatorio de desempates por jogador disponivel apenas para torneios individuais.")
        rows = []
        for standing in self.pairing_service.tiebreak_report(tournament_id):
            components = dict(standing.get("tiebreak_components") or {})
            for criterion in (
                "buchholz",
                "buchholz_median",
                "sonneborn_berger",
                "direct_encounter",
                "wins",
                "performance",
            ):
                item = dict(components.get(criterion) or {})
                rows.append(
                    [
                        standing["position"],
                        standing["name"],
                        standing["points"],
                        item.get("label", criterion),
                        item.get("value", ""),
                        item.get("formula", ""),
                        self._tiebreak_component_summary(item),
                    ]
                )
        self._write_report(
            file_path,
            f"Desempates - {tournament['name']}",
            ["Pos", "Jogador", "Pts", "Criterio", "Valor", "Formula", "Componentes"],
            rows,
        )

    def export_tournament_audit(self, tournament_id: int, file_path: str | Path) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        events = self.db.list_audit_events(tournament_id, limit=5000)
        rows = [
            [
                event.get("created_at", ""),
                event.get("action", ""),
                event.get("entity_type", ""),
                event.get("entity_id", ""),
                event.get("round_id", ""),
                event.get("actor", ""),
                event.get("role", ""),
                event.get("reason", ""),
                event.get("before_hash", ""),
                event.get("after_hash", ""),
            ]
            for event in events
        ]
        self._write_report(
            file_path,
            f"Auditoria do torneio - {tournament['name']}",
            [
                "Data/hora",
                "Acao",
                "Entidade",
                "ID entidade",
                "Rodada ID",
                "Operador",
                "Perfil",
                "Motivo",
                "Hash anterior",
                "Hash posterior",
            ],
            rows,
        )

    def export_complete(self, tournament_id: int, file_path: str | Path) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        sections = [
            self._tournament_section(tournament_id),
        ]
        if tournament.get("competition_type") == "team":
            sections.append(self._teams_section(tournament_id))
            sections.append(self._team_rosters_section(tournament_id))
            sections.append(self._team_lineups_section(tournament_id))
            sections.append(self._team_substitutions_section(tournament_id))
        sections.append(self._players_section(tournament_id))
        for round_data in sorted(self.db.list_rounds(tournament_id), key=lambda item: item["number"]):
            sections.append(self._pairings_section(round_data["id"]))
        sections.append(self._standings_section(tournament_id))
        self._write_multi_report(file_path, sections)

    def export_pgn(self, tournament_id: int, file_path: str | Path) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
            
        path = Path(file_path)
        with path.open("w", encoding="utf-8") as f:
            for round_data in sorted(self.db.list_rounds(tournament_id), key=lambda r: r["number"]):
                pairings = self.db.get_pairings_for_round(round_data["id"])
                for p in pairings:
                    if p["is_bye"]:
                        continue
                    
                    white_name = p.get("white_name") or "Unknown"
                    black_name = p.get("black_name") or "Unknown"
                    result = p.get("result") or "*"
                    if result == "1-0":
                        res_str = "1-0"
                    elif result == "0-1":
                        res_str = "0-1"
                    elif result == "1/2-1/2":
                        res_str = "1/2-1/2"
                    else:
                        res_str = "*"
                        
                    f.write(f'[Event "{tournament["name"]}"]\n')
                    f.write(f'[Site "{tournament.get("location") or ""}"]\n')
                    f.write(f'[Date "{tournament.get("start_date") or ""}"]\n')
                    f.write(f'[Round "{round_data["number"]}"]\n')
                    f.write(f'[White "{white_name}"]\n')
                    f.write(f'[Black "{black_name}"]\n')
                    f.write(f'[Result "{res_str}"]\n')
                    f.write(f'\n{res_str}\n\n')

    def export_trf(self, tournament_id: int, file_path: str | Path) -> list[str]:
        return self.export_chess_results_trf(tournament_id, file_path)

    def export_chess_results_trf(self, tournament_id: int, file_path: str | Path) -> list[str]:
        return self._federation_exporter("trf16").export(tournament_id, file_path)

    def export_chess_results_trf25(self, tournament_id: int, file_path: str | Path) -> list[str]:
        return self._federation_exporter("trf25").export(tournament_id, file_path)

    def validate_chess_results_trf(self, tournament_id: int) -> list[str]:
        return self._federation_exporter("trf16").validate(tournament_id)

    def export_chess_results_trf_validation_report(self, tournament_id: int, file_path: str | Path) -> None:
        rows = self._federation_exporter("trf16").validation_report_rows(tournament_id)
        self._write_report(
            file_path,
            "Pendencias TRF16",
            ["Tipo", "Item", "Valor"],
            rows,
        )

    def _federation_exporter(self, code: str):
        from src.services.federation_exporters import (
            FederationExporterRegistry,
            TRF16Exporter,
            TRF25Exporter,
        )

        registry = FederationExporterRegistry()
        registry.register(TRF16Exporter(self))
        registry.register(TRF25Exporter(self))
        return registry.get(code)

    def _trf_pending_result_rounds(
        self,
        tournament: Mapping[str, Any],
        rounds: list[dict[str, Any]],
    ) -> list[str]:
        pending_rounds = []
        for round_data in rounds:
            if tournament.get("competition_type") == "team":
                has_pending = any(
                    not board.get("result") or str(board.get("result")) not in RESULT_POINTS
                    for match in self.db.list_team_matches_for_round(int(round_data["id"]))
                    if not match.get("is_bye")
                    for board in self.db.list_team_boards(int(match["id"]))
                )
            else:
                has_pending = any(
                    not pairing.get("result") or str(pairing.get("result")) not in FINAL_RESULTS
                    for pairing in self.db.get_pairings_for_round(int(round_data["id"]))
                )
            if has_pending:
                pending_rounds.append(str(round_data["number"]))
        return pending_rounds

    @staticmethod
    def _trf_clean(value: Any, width: int | None = None) -> str:
        text = unicodedata.normalize("NFKD", str(value or ""))
        text = "".join(char for char in text if not unicodedata.combining(char))
        text = " ".join(text.replace("\r", " ").replace("\n", " ").split())
        if width is not None:
            return text[:width].ljust(width)
        return text

    @staticmethod
    def _trf_date(value: Any, *, long_year: bool) -> str:
        raw = str(value or "").strip()
        if not raw:
            return ""
        parsed: datetime | None = None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d/%m/%Y", "%d. %m. %Y", "%d.%m.%Y"):
            try:
                parsed = datetime.strptime(raw, fmt)
                break
            except ValueError:
                continue
        if not parsed:
            return raw.replace("-", "/")[:10]
        return parsed.strftime("%Y/%m/%d" if long_year else "%y/%m/%d")

    @classmethod
    def _trf_date_is_valid(cls, value: Any) -> bool:
        raw = str(value or "").strip()
        if not raw:
            return False
        normalized = cls._trf_date(raw, long_year=True)
        try:
            datetime.strptime(normalized, "%Y/%m/%d")
            return True
        except ValueError:
            return False

    @staticmethod
    def _trf_valid_federation_code(value: Any) -> bool:
        code = str(value or "").strip()
        return len(code) == 3 and code.isalpha()

    @staticmethod
    def _trf_rating(player: Mapping[str, Any]) -> int:
        return int(
            player.get("international_rating")
            or player.get("rating")
            or player.get("national_rating")
            or 0
        )

    @staticmethod
    def _trf_points(value: Any) -> str:
        points = float(value or 0.0)
        if points >= 100:
            return f"{points:4.0f}"[:4]
        return f"{points:4.1f}"[-4:]

    @staticmethod
    def _trf_missing_field_warnings(field_name: str, player_names: list[str]) -> list[str]:
        if not player_names:
            return []
        shown = ", ".join(player_names[:5])
        suffix = f" e mais {len(player_names) - 5}" if len(player_names) > 5 else ""
        return [f"Jogadores sem {field_name}: {shown}{suffix}."]

    @staticmethod
    def _trf_player_result(result: str, *, is_white: bool, is_bye: bool) -> str:
        if is_bye:
            return "U"
        if result == "1-0":
            return "1" if is_white else "0"
        if result == "0-1":
            return "0" if is_white else "1"
        if result == "1/2-1/2":
            return "="
        if result == "1F-0F":
            return "+" if is_white else "-"
        if result == "0F-1F":
            return "-" if is_white else "+"
        if result == "0F-0F":
            return "-"
        return "Z"

    @staticmethod
    def _trf_tournament_line(code: str, value: Any) -> str:
        return f"{code} {ExportService._trf_clean(value)}\r\n"

    def _trf_tournament_type(self, tournament: Mapping[str, Any], settings: Mapping[str, Any]) -> str:
        profile = str(settings.get("tournament_profile") or "").strip()
        suffix = "FIDE-rated" if profile == "fide" else "Standard"
        system = str(tournament.get("system") or "Suico").strip()
        if tournament.get("competition_type") == "team":
            return f"Team: {system} ({suffix})"
        return f"Individual: {system} ({suffix})"

    def _trf_pairings_by_round(
        self,
        tournament: Mapping[str, Any],
        rounds: list[dict[str, Any]],
    ) -> dict[int, list[dict[str, Any]]]:
        if tournament.get("competition_type") != "team":
            return {
                int(round_data["number"]): self.db.get_pairings_for_round(round_data["id"])
                for round_data in rounds
            }

        pairings_by_round: dict[int, list[dict[str, Any]]] = {}
        for round_data in rounds:
            board_pairings = []
            for match in self.db.list_team_matches_for_round(int(round_data["id"])):
                if match.get("is_bye"):
                    continue
                for board in self.db.list_team_boards(int(match["id"])):
                    if not board.get("white_player_id") or not board.get("black_player_id"):
                        continue
                    board_pairings.append(
                        {
                            "white_player_id": int(board["white_player_id"]),
                            "black_player_id": int(board["black_player_id"]),
                            "result": str(board.get("result") or ""),
                            "is_bye": 0,
                        }
                    )
            pairings_by_round[int(round_data["number"])] = board_pairings
        return pairings_by_round

    def _trf_player_standings(
        self,
        tournament: Mapping[str, Any],
        players: list[dict[str, Any]],
    ) -> dict[int, dict[str, Any]]:
        if tournament.get("competition_type") != "team":
            return {
                int(item["player_id"]): item
                for item in self.pairing_service.standings(int(tournament["id"]))
            }

        stats = {
            int(player["id"]): {
                "player_id": int(player["id"]),
                "points": float(player.get("starting_points", 0.0) or 0.0),
                "position": 0,
                "rating": self._trf_rating(player),
                "name": player_pairing_name(player),
            }
            for player in players
        }
        for round_data in self.db.list_rounds(int(tournament["id"])):
            for match in self.db.list_team_matches_for_round(int(round_data["id"])):
                for board in self.db.list_team_boards(int(match["id"])):
                    white_id = int(board.get("white_player_id") or 0)
                    black_id = int(board.get("black_player_id") or 0)
                    result = str(board.get("result") or "")
                    if not white_id or not black_id or result not in RESULT_POINTS:
                        continue
                    white_points, black_points = RESULT_POINTS[result]
                    if white_id in stats:
                        stats[white_id]["points"] += white_points
                    if black_id in stats:
                        stats[black_id]["points"] += black_points

        ordered_stats = sorted(
            stats.values(),
            key=lambda item: (-float(item["points"]), -int(item["rating"]), str(item["name"]).casefold()),
        )
        for index, item in enumerate(ordered_stats, start=1):
            item["position"] = index
        return stats

    def _trf_chief_arbiter(self, tournament_id: int, settings: Mapping[str, Any]) -> str:
        chief = str(settings.get("chief_arbiter") or "").strip()
        if chief:
            return chief
        for referee in self.db.list_tournament_referees(tournament_id):
            role = str(referee.get("role") or "").casefold()
            if "chief" in role or "principal" in role or "arbitro chefe" in role:
                return str(referee.get("name") or "").strip()
        return str(settings.get("director") or settings.get("organizer") or "").strip()

    def _trf_deputy_arbiters(self, tournament_id: int, settings: Mapping[str, Any]) -> list[str]:
        deputies = [
            item.strip()
            for item in str(settings.get("arbiters") or "").replace(";", "\n").splitlines()
            if item.strip()
        ]
        chief = self._trf_chief_arbiter(tournament_id, settings).casefold()
        for referee in self.db.list_tournament_referees(tournament_id):
            name = str(referee.get("name") or "").strip()
            if name and name.casefold() != chief and name not in deputies:
                deputies.append(name)
        return deputies

    def _trf_round_dates_line(self, round_count: int, schedule: Mapping[int, str]) -> str:
        line = "132" + (" " * 88)
        for round_number in range(1, round_count + 1):
            line += f"{self._trf_date(schedule.get(round_number), long_year=False):<8}  "
        return f"{line.rstrip()}\r\n"

    def _trf_player_line(
        self,
        player: Mapping[str, Any],
        player_id_to_start_rank: Mapping[int, int],
        standings_by_player: Mapping[int, Mapping[str, Any]],
        pairings_by_round: Mapping[int, list[dict[str, Any]]],
        round_count: int,
        settings: Mapping[str, Any],
    ) -> str:
        player_id = int(player["id"])
        start_rank = player_id_to_start_rank[player_id]
        standing = standings_by_player.get(player_id, {})
        sex = self._trf_clean(str(player.get("sex") or "")[:1].lower(), 1)
        title = self._trf_clean(str(player.get("title") or "").upper(), 3)
        name = self._trf_clean(player_pairing_name(player), 33)
        rating = f"{self._trf_rating(player):4d}"[-4:]
        federation = self._trf_clean(
            str(player.get("federation_id") or settings.get("federation") or "").upper(),
            3,
        )
        fide_id = self._trf_clean(player.get("fide_id"))[:11].rjust(11)
        birth_date = self._trf_clean(self._trf_date(player.get("birth_date"), long_year=True), 10)
        points = self._trf_points(standing.get("points", player.get("starting_points", 0.0)))
        rank = int(standing.get("position") or 0)
        line = (
            f"001 {start_rank:4d} {sex}{title} {name} {rating} {federation} "
            f"{fide_id} {birth_date} {points} {rank:4d}  "
        )
        for round_number in range(1, round_count + 1):
            line += self._trf_round_cell(player_id, player_id_to_start_rank, pairings_by_round.get(round_number, []))
        return f"{line.rstrip()}\r\n"

    def _trf_round_cell(
        self,
        player_id: int,
        player_id_to_start_rank: Mapping[int, int],
        pairings: list[dict[str, Any]],
    ) -> str:
        pairing = next(
            (
                item
                for item in pairings
                if int(item["white_player_id"]) == player_id
                or int(item.get("black_player_id") or 0) == player_id
            ),
            None,
        )
        if not pairing:
            return "0000 - Z  "
        is_bye = bool(pairing.get("is_bye"))
        is_white = int(pairing["white_player_id"]) == player_id
        if is_bye:
            code = str(pairing.get("result") or "").strip().upper()
            bye_code = code if code in {"F", "H", "Z"} else "U"
            return f"0000 - {bye_code}  "
        opponent_id = int(pairing["black_player_id"] if is_white else pairing["white_player_id"])
        opponent_rank = player_id_to_start_rank.get(opponent_id, 0)
        color = "w" if is_white else "b"
        result = self._trf_player_result(str(pairing.get("result") or ""), is_white=is_white, is_bye=is_bye)
        return f"{opponent_rank:4d} {color} {result}  "

    def _trf_team_line(
        self,
        team: Mapping[str, Any],
        player_id_to_start_rank: Mapping[int, int],
    ) -> str:
        team_name = self._trf_clean(team.get("name"), 32)
        line = f"013 {team_name}"
        for assignment in self.db.list_team_players(int(team["id"]), active_only=False):
            start_rank = player_id_to_start_rank.get(int(assignment["player_id"]))
            if start_rank:
                line += f"{start_rank:4d} "
        return f"{line.rstrip()}\r\n"

    def export_member_evolution(self, member_id: int, file_path: str | Path) -> None:
        sections = self._member_evolution_sections(member_id)
        self._write_multi_report(file_path, sections)

    def export_club_report(self, file_path: str | Path) -> None:
        sections = self._club_report_sections()
        self._write_multi_report(file_path, sections)

    def export_member_report(self, member_id: int, file_path: str | Path) -> None:
        sections = self._member_evolution_sections(member_id)
        self._write_multi_report(file_path, sections)

    def export_tournaments_period_report(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._tournaments_period_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_attendance_report(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._attendance_report_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_financial_report(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._financial_report_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_events_report(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._events_report_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_internal_ranking_report(
        self,
        file_path: str | Path,
        category: str = "",
        club_id: int | None = None,
        class_id: int | None = None,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._internal_ranking_sections(
            category,
            club_id=club_id,
            class_id=class_id,
            start_date=start_date,
            end_date=end_date,
        )
        self._write_multi_report(file_path, sections)

    def export_payment_receipt(self, payment_id: int, file_path: str | Path) -> Path:
        payment = self.db.get_payment(payment_id)
        if not payment:
            raise AppError("Lancamento financeiro nao encontrado.")
        finance_service = __import__('src.services.finance_service', fromlist=['FinanceService']).FinanceService(self.db)
        payment = finance_service._with_effective_status(payment)
        if payment["effective_status"] not in {"paid", "exempt"}:
            raise AppError("Recibo disponivel apenas para lancamentos pagos ou isentos.")

        path = Path(file_path)
        if path.suffix.lower() != ".pdf":
            # Para CSV/Excel continua usando o padrão genérico
            rows = [
                ["Recibo", f"REC-{int(payment['id']):06d}"],
                ["Membro", payment.get("member_name") or ""],
                ["Plano", payment.get("plan_name") or ""],
                ["Descricao", payment.get("description") or ""],
                ["Referencia", payment.get("reference_period") or ""],
                ["Valor", self._format_report_number(payment.get("amount") or 0.0)],
                ["Pagamento", payment.get("payment_date") or ""],
            ]
            self._write_report(path, "Recibo financeiro", ["Campo", "Valor"], rows)
            return path
            
        try:
            from reportlab.pdfgen import canvas
            from reportlab.lib.pagesizes import A4
            
            c = canvas.Canvas(str(path), pagesize=A4)
            c.setFont("Helvetica-Bold", 18)
            c.drawString(50, 800, "RECIBO DE PAGAMENTO")
            
            c.setFont("Helvetica-Bold", 12)
            c.drawString(400, 800, f"N. REC-{int(payment['id']):06d}")
            
            c.setFont("Helvetica", 12)
            c.drawString(50, 750, f"Recebemos de: {payment.get('member_name') or 'N/A'}")
            
            amount = float(payment.get("amount") or 0.0)
            c.drawString(50, 720, f"A quantia de: R$ {amount:,.2f}")
            c.drawString(50, 690, f"Referente a: {payment.get('description') or ''} - Ref: {payment.get('reference_period') or ''}")
            
            c.drawString(50, 660, f"Data do Pagamento: {payment.get('payment_date') or ''}")
            
            c.drawString(50, 600, "Por ser verdade, firmamos o presente recibo.")
            
            c.line(50, 520, 300, 520)
            c.drawString(50, 500, "Assinatura do Tesoureiro / Diretoria")
            c.drawString(50, 480, "Clube de Xadrez Albericus")
            
            c.save()
            return path
        except ImportError as exc:
            raise AppError("Instale reportlab para gerar Recibos em PDF (pip install reportlab).") from exc

    def export_administrative_package(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._administrative_package_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_club_portal(
        self,
        output_dir: str | Path,
        club_id: int | None = None,
        class_id: int | None = None,
    ) -> Path:
        club, class_data = self._club_portal_scope(club_id, class_id)
        resolved_club_id = int(club["id"]) if club else None
        resolved_class_id = int(class_data["id"]) if class_data else None
        path = Path(output_dir)
        path.mkdir(parents=True, exist_ok=True)
        html_content = self._club_portal_html(
            club=club,
            class_data=class_data,
            members=self.db.list_members(active_only=True, club_id=resolved_club_id, class_id=resolved_class_id),
            classes=[class_data]
            if class_data
            else self.db.list_classes(club_id=resolved_club_id, active_only=True),
            events=self.db.list_club_events(
                start_date=date.today().isoformat(),
                club_id=resolved_club_id,
                include_canceled=False,
            )[:10],
            sessions=self.db.list_training_sessions(
                start_date=date.today().isoformat(),
                club_id=resolved_club_id,
                class_id=resolved_class_id,
            )[:10],
            tournaments=[
                tournament
                for tournament in self.db.list_tournaments()
                if self._portal_tournament_in_scope(tournament, resolved_club_id, resolved_class_id)
            ][:8],
            ranking=__import__('src.services.rating_service', fromlist=['InternalRatingService']).InternalRatingService(self.db).ranking(
                active_only=True,
                club_id=resolved_club_id,
                class_id=resolved_class_id,
            )[:10],
        )
        (path / "index.html").write_text(html_content, encoding="utf-8")
        (path / "styles.css").write_text(self._club_portal_css(), encoding="utf-8")
        logger.info("Portal estatico do clube exportado em %s", path)
        return path / "index.html"

    def export_site(self, tournament_id: int, output_dir: str | Path) -> Path:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")

        path = Path(output_dir)
        path.mkdir(parents=True, exist_ok=True)
        settings = self.db.get_tournament_settings(tournament_id) or {}
        schedule = self.db.list_round_schedule(tournament_id)
        players = self.db.list_players(tournament_id, active_only=False)
        standings = (
            []
            if settings.get("hide_standings")
            else self._standings_for_tournament(tournament_id, tournament)
        )
        rounds = sorted(self.db.list_rounds(tournament_id), key=lambda item: item["number"])

        index_html = self._site_html(tournament, settings, schedule, players, rounds, standings)
        (path / "index.html").write_text(index_html, encoding="utf-8")
        (path / "styles.css").write_text(self._site_css(), encoding="utf-8")
        logger.info("Site estatico exportado em %s", path)
        return path / "index.html"

    def export_public_json(self, tournament_id: int, file_path: str | Path, mode: str = "publico") -> Path:
        payload = self.public_tournament_payload(tournament_id, mode=mode)
        path = Path(file_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        return path

    def public_tournament_payload(self, tournament_id: int, mode: str = "publico") -> dict[str, Any]:
        mode = self._portal_mode(mode)
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        settings = self.db.get_tournament_settings(tournament_id) or {}
        app_settings = self.db.get_app_settings()
        is_team_tournament = tournament.get("competition_type") == "team"
        standings = [] if settings.get("hide_standings") else self._standings_for_tournament(tournament_id, tournament)
        rounds = sorted(self.db.list_rounds(tournament_id), key=lambda item: item["number"])
        latest_round = rounds[-1] if rounds else None
        players = self.db.list_players(tournament_id, active_only=False)
        payload = {
            "mode": mode,
            "generated_at": Database.now(),
            "notice": str(app_settings.get("live_portal_notice") or ""),
            "tournament": self._public_tournament_info(tournament, settings, mode),
            "current_round": self._public_round_payload(latest_round, is_team_tournament) if latest_round else None,
            "rounds": [self._public_round_payload(round_data, is_team_tournament) for round_data in rounds],
            "standings": self._public_standings_payload(standings, is_team_tournament, mode),
            "players": [self._public_player_card(player, mode) for player in players],
        }
        if is_team_tournament:
            payload["teams"] = [self._public_team_card(team, mode) for team in self.db.list_teams(tournament_id, active_only=False)]
        return payload

    @staticmethod
    def _portal_mode(mode: str) -> str:
        normalized = str(mode or "publico").strip().casefold()
        if normalized in {"privado", "private"}:
            return "privado"
        if normalized in {"clube", "club"}:
            return "clube"
        return "publico"

    def _public_tournament_info(
        self,
        tournament: dict[str, Any],
        settings: dict[str, Any],
        mode: str,
    ) -> dict[str, Any]:
        info = {
            "id": int(tournament["id"]),
            "name": tournament.get("name", ""),
            "location": tournament.get("location", ""),
            "start_date": tournament.get("start_date", ""),
            "end_date": tournament.get("end_date", ""),
            "rounds_count": int(tournament.get("rounds_count") or 0),
            "time_control": tournament.get("time_control", ""),
            "status": tournament.get("status", ""),
            "competition_type": tournament.get("competition_type", "individual"),
            "chief_arbiter": settings.get("chief_arbiter", ""),
            "organizer": settings.get("organizer", ""),
        }
        if mode == "privado":
            info["contact_email"] = settings.get("contact_email", "")
            info["comments"] = settings.get("comments", "")
        return info

    def _public_round_payload(self, round_data: dict[str, Any], is_team_tournament: bool) -> dict[str, Any]:
        if is_team_tournament:
            matches = []
            for match in self.db.list_team_matches_for_round(int(round_data["id"])):
                matches.append(
                    {
                        "match_number": int(match.get("match_number") or 0),
                        "white_team": match.get("white_team_name", ""),
                        "black_team": "BYE" if match.get("is_bye") else match.get("black_team_name", ""),
                        "result": self._team_match_score(match),
                        "boards": [
                            {
                                "board": int(board.get("board_number") or 0),
                                "white": self._team_board_player_name(board, "white"),
                                "black": self._team_board_player_name(board, "black"),
                                "result": board.get("result", ""),
                            }
                            for board in self.db.list_team_boards(int(match["id"]))
                        ],
                    }
                )
            return {"id": int(round_data["id"]), "number": int(round_data["number"]), "status": round_data["status"], "matches": matches}
        pairings = []
        for pairing in self.db.get_pairings_for_round(int(round_data["id"])):
            pairings.append(
                {
                    "board": int(pairing.get("board_number") or 0),
                    "white": pairing_player_name(pairing, "white"),
                    "black": "BYE" if pairing.get("is_bye") else pairing_player_name(pairing, "black"),
                    "result": pairing.get("result", ""),
                    "is_bye": bool(pairing.get("is_bye")),
                }
            )
        return {"id": int(round_data["id"]), "number": int(round_data["number"]), "status": round_data["status"], "pairings": pairings}

    @staticmethod
    def _public_standings_payload(
        standings: list[dict[str, Any]],
        is_team_tournament: bool,
        mode: str,
    ) -> list[dict[str, Any]]:
        if is_team_tournament:
            return [
                {
                    "position": item["position"],
                    "team": item["name"],
                    "club": item.get("club", ""),
                    "match_points": item.get("match_points", 0),
                    "game_points": item.get("game_points", 0),
                    "wins": item.get("wins", 0),
                    "draws": item.get("draws", 0),
                    "losses": item.get("losses", 0),
                    "buchholz": item.get("buchholz", 0),
                }
                for item in standings
            ]
        rows = []
        for item in standings:
            row = {
                "position": item["position"],
                "name": item["name"],
                "category": item.get("category", ""),
                "points": item.get("points", 0),
                "buchholz": item.get("buchholz", 0),
                "buchholz_median": item.get("buchholz_median", 0),
                "sonneborn_berger": item.get("sonneborn_berger", 0),
                "wins": item.get("wins", 0),
            }
            if mode in {"clube", "privado"}:
                row["rating"] = item.get("rating", 0)
                row["club"] = item.get("club", "")
            rows.append(row)
        return rows

    @staticmethod
    def _public_player_card(player: dict[str, Any], mode: str) -> dict[str, Any]:
        card = {
            "id": int(player.get("id") or 0),
            "name": player_full_name(player),
            "club": player.get("club", ""),
            "category": player.get("category", ""),
            "rating": int(player.get("rating") or 0),
            "status": player.get("player_status", "active"),
        }
        if mode in {"clube", "privado"}:
            card["fide_id"] = player.get("fide_id", "")
            card["cbx_id"] = player.get("cbx_id", "")
        return card

    @staticmethod
    def _public_team_card(team: dict[str, Any], mode: str) -> dict[str, Any]:
        card = {
            "id": int(team.get("id") or 0),
            "name": team.get("name", ""),
            "club": team.get("club", ""),
            "captain": team.get("captain", "") if mode in {"clube", "privado"} else "",
            "active": bool(team.get("active")),
        }
        return card

    def live_portal_html(self, tournament_id: int, mode: str = "publico") -> str:
        payload = self.public_tournament_payload(tournament_id, mode=mode)
        tournament = payload["tournament"]
        standings_rows = "".join(
            "<tr>"
            f"<td>{item.get('position', '')}</td>"
            f"<td>{self._escape(item.get('name') or item.get('team') or '')}</td>"
            f"<td>{self._escape(item.get('category') or item.get('club') or '')}</td>"
            f"<td>{self._escape(item.get('points', item.get('match_points', '')))}</td>"
            f"<td>{self._escape(item.get('buchholz', ''))}</td>"
            "</tr>"
            for item in payload["standings"]
        )
        current = payload.get("current_round") or {}
        pairings = current.get("matches") or current.get("pairings") or []
        current_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('match_number', item.get('board', '')))}</td>"
            f"<td>{self._escape(item.get('white_team', item.get('white', '')))}</td>"
            f"<td>{self._escape(item.get('result', ''))}</td>"
            f"<td>{self._escape(item.get('black_team', item.get('black', '')))}</td>"
            "</tr>"
            for item in pairings
        )
        player_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('name', ''))}</td>"
            f"<td>{self._escape(item.get('club', ''))}</td>"
            f"<td>{self._escape(item.get('category', ''))}</td>"
            f"<td>{self._escape(item.get('rating', ''))}</td>"
            "</tr>"
            for item in payload["players"]
        )
        rounds_rows = "".join(
            "<tr>"
            f"<td>{round_data.get('number', '')}</td>"
            f"<td>{self._escape(round_data.get('status', ''))}</td>"
            f"<td>{len(round_data.get('matches') or round_data.get('pairings') or [])}</td>"
            "</tr>"
            for round_data in payload["rounds"]
        )
        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{self._escape(tournament['name'])}</title>
  <style>{self._site_css()}</style>
</head>
<body>
  <header><p class="eyebrow">Albericus Live</p><h1>{self._escape(tournament['name'])}</h1><p>{self._escape(tournament.get('location', ''))}</p></header>
  <main>
    <section class="summary">
      <div><strong>Modo</strong><span>{self._escape(payload['mode'])}</span></div>
      <div><strong>Rodada atual</strong><span>{self._escape(current.get('number', 'Sem rodada'))}</span></div>
      <div><strong>Status</strong><span>{self._escape(tournament.get('status', ''))}</span></div>
    </section>
    <section><h2>Avisos do arbitro</h2><p>{self._escape(payload.get('notice') or 'Nenhum aviso publicado.')}</p></section>
    <section><h2>Rodada atual</h2><table><thead><tr><th>Mesa</th><th>Brancas/Equipe A</th><th>Resultado</th><th>Pretas/Equipe B</th></tr></thead><tbody>{current_rows}</tbody></table></section>
    <section><h2>Classificacao</h2><table><thead><tr><th>Pos</th><th>Nome</th><th>Categoria/Clube</th><th>Pts</th><th>Buchholz</th></tr></thead><tbody>{standings_rows}</tbody></table></section>
    <section><h2>Historico de rodadas</h2><table><thead><tr><th>Rodada</th><th>Status</th><th>Mesas</th></tr></thead><tbody>{rounds_rows}</tbody></table></section>
    <section><h2>Fichas publicas</h2><table><thead><tr><th>Nome</th><th>Clube</th><th>Categoria</th><th>Rating</th></tr></thead><tbody>{player_rows}</tbody></table></section>
  </main>
  <footer>Atualizado em {self._escape(payload['generated_at'])}</footer>
</body>
</html>"""

    def _club_portal_scope(
        self,
        club_id: int | None,
        class_id: int | None,
    ) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
        class_data = self.db.get_class(class_id) if class_id else None
        if class_id and not class_data:
            raise AppError("Turma do portal nao encontrada.")
        if class_data and club_id and int(class_data["club_id"]) != int(club_id):
            raise AppError("A turma selecionada nao pertence ao clube/escola do portal.")

        if class_data and not club_id:
            club_id = int(class_data["club_id"])
        club = self.db.get_club(club_id) if club_id else None
        if club_id and not club:
            raise AppError("Clube/escola do portal nao encontrado.")
        if not club:
            active_clubs = self.db.list_clubs(active_only=True)
            club = active_clubs[0] if active_clubs else self.db.get_club(1)
        return club, class_data

    @staticmethod
    def _portal_tournament_in_scope(
        tournament: Mapping[str, Any],
        club_id: int | None,
        class_id: int | None,
    ) -> bool:
        if class_id and int(tournament.get("class_id") or 0) != class_id:
            return False
        if club_id and int(tournament.get("club_id") or 0) != club_id:
            return False
        return True

    def _club_portal_html(
        self,
        club: dict[str, Any] | None,
        class_data: dict[str, Any] | None,
        members: list[dict[str, Any]],
        classes: list[dict[str, Any]],
        events: list[dict[str, Any]],
        sessions: list[dict[str, Any]],
        tournaments: list[dict[str, Any]],
        ranking: list[dict[str, Any]],
    ) -> str:
        title = class_data.get("name") if class_data else club.get("name") if club else "Portal do clube"
        club_kind_labels = {
            "club": "Clube",
            "school": "Escola",
            "project": "Projeto",
            "partner": "Parceiro",
        }
        subtitle = (
            class_data.get("club_name")
            if class_data
            else club_kind_labels.get(str(club.get("kind") or "club"), "") if club else ""
        )
        generated_at = Database.now()
        notice_rows = self._portal_notice_rows(events, sessions)
        class_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('name') or '')}</td>"
            f"<td>{self._escape(item.get('teacher') or '')}</td>"
            f"<td>{self._escape(item.get('weekday') or '')}</td>"
            f"<td>{self._escape(item.get('time') or '')}</td>"
            f"<td>{self._escape(item.get('location') or '')}</td>"
            f"<td>{item.get('active_members_count', '')}</td>"
            "</tr>"
            for item in classes
        )
        event_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('event_date') or '')}</td>"
            f"<td>{self._escape(item.get('start_time') or '')}</td>"
            f"<td>{self._escape(EVENT_TYPES.get(item.get('event_type', ''), item.get('event_type') or ''))}</td>"
            f"<td>{self._escape(item.get('title') or '')}</td>"
            f"<td>{self._escape(item.get('location') or '')}</td>"
            f"<td>{self._escape(EVENT_STATUSES.get(item.get('status', ''), item.get('status') or ''))}</td>"
            "</tr>"
            for item in events
        )
        session_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('session_date') or '')}</td>"
            f"<td>{self._escape(item.get('start_time') or '')}</td>"
            f"<td>{self._escape(item.get('title') or '')}</td>"
            f"<td>{self._escape(item.get('class_name') or '')}</td>"
            f"<td>{self._escape(item.get('learning_level_name') or '')}</td>"
            f"<td>{self._escape(item.get('objective') or '')}</td>"
            "</tr>"
            for item in sessions
        )
        ranking_rows = "".join(
            "<tr>"
            f"<td>{item['position']}</td>"
            f"<td>{self._escape(item.get('name') or '')}</td>"
            f"<td>{item.get('rating', '')}</td>"
            f"<td>{self._escape(item.get('category') or '')}</td>"
            f"<td>{self._escape(item.get('class_name') or '')}</td>"
            f"<td>{self._escape(str(item.get('score_rate') or 0))}%</td>"
            "</tr>"
            for item in ranking
        )
        tournament_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('name') or '')}</td>"
            f"<td>{self._escape(item.get('start_date') or '')}</td>"
            f"<td>{self._escape(item.get('location') or '')}</td>"
            f"<td>{self._escape(item.get('status') or '')}</td>"
            "</tr>"
            for item in tournaments
        )
        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{self._escape(title)}</title>
  <link rel="stylesheet" href="styles.css">
</head>
<body>
  <header>
    <p class="eyebrow">Albericus</p>
    <h1>{self._escape(title)}</h1>
    <p>{self._escape(subtitle)}</p>
  </header>
  <main>
    <section class="summary">
      <div><strong>Membros ativos</strong><span>{len(members)}</span></div>
      <div><strong>Turmas</strong><span>{len(classes)}</span></div>
      <div><strong>Eventos futuros</strong><span>{len(events)}</span></div>
      <div><strong>Aulas futuras</strong><span>{len(sessions)}</span></div>
    </section>
    <section>
      <h2>Comunicados</h2>
      <table><thead><tr><th>Data</th><th>Comunicado</th></tr></thead><tbody>{notice_rows}</tbody></table>
    </section>
    <section>
      <h2>Calendario</h2>
      <h3>Eventos</h3>
      <table><thead><tr><th>Data</th><th>Hora</th><th>Tipo</th><th>Evento</th><th>Local</th><th>Status</th></tr></thead><tbody>{event_rows}</tbody></table>
      <h3>Aulas e treinos</h3>
      <table><thead><tr><th>Data</th><th>Hora</th><th>Aula/Treino</th><th>Turma</th><th>Nivel</th><th>Objetivo</th></tr></thead><tbody>{session_rows}</tbody></table>
    </section>
    <section>
      <h2>Turmas</h2>
      <table><thead><tr><th>Turma</th><th>Professor</th><th>Dia</th><th>Horario</th><th>Local</th><th>Alunos</th></tr></thead><tbody>{class_rows}</tbody></table>
    </section>
    <section>
      <h2>Ranking interno</h2>
      <table><thead><tr><th>Pos</th><th>Membro</th><th>Rating</th><th>Categoria</th><th>Turma</th><th>Aproveitamento</th></tr></thead><tbody>{ranking_rows}</tbody></table>
    </section>
    <section>
      <h2>Torneios recentes</h2>
      <table><thead><tr><th>Torneio</th><th>Data</th><th>Local</th><th>Status</th></tr></thead><tbody>{tournament_rows}</tbody></table>
    </section>
  </main>
  <footer>Gerado em {self._escape(generated_at)}</footer>
</body>
</html>
"""

    def _portal_notice_rows(
        self,
        events: list[dict[str, Any]],
        sessions: list[dict[str, Any]],
    ) -> str:
        notices: list[tuple[str, str]] = []
        for event in events[:5]:
            details = " - ".join(
                part
                for part in [
                    str(event.get("title") or ""),
                    str(event.get("location") or ""),
                    str(event.get("notes") or ""),
                ]
                if part
            )
            notices.append((str(event.get("event_date") or ""), details))
        for session in sessions[:5]:
            details = " - ".join(
                part
                for part in [
                    str(session.get("title") or ""),
                    str(session.get("class_name") or ""),
                    str(session.get("homework") or session.get("objective") or ""),
                ]
                if part
            )
            notices.append((str(session.get("session_date") or ""), details))
        notices.sort(key=lambda item: item[0])
        if not notices:
            return "<tr><td colspan=\"2\">Nenhum comunicado futuro cadastrado.</td></tr>"
        return "".join(
            "<tr>"
            f"<td>{self._escape(day)}</td>"
            f"<td>{self._escape(text)}</td>"
            "</tr>"
            for day, text in notices
        )

    @staticmethod
    def _club_portal_css() -> str:
        return """* { box-sizing: border-box; }
body {
  margin: 0;
  font-family: Arial, Helvetica, sans-serif;
  color: #18212f;
  background: #f5f7fb;
}
header {
  background: #17403a;
  color: #fff;
  padding: 30px 40px;
}
header h1 { margin: 4px 0 8px; font-size: 32px; }
header p { margin: 0; color: #dbe8e5; }
.eyebrow { text-transform: uppercase; letter-spacing: 0; font-size: 12px; }
main { max-width: 1180px; margin: 0 auto; padding: 24px; }
section {
  background: #fff;
  border: 1px solid #dfe5ed;
  border-radius: 8px;
  margin-bottom: 16px;
  padding: 18px;
  overflow-x: auto;
}
.summary {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(170px, 1fr));
  gap: 12px;
}
.summary div {
  border: 1px solid #dfe5ed;
  border-radius: 8px;
  padding: 12px;
}
.summary strong, .summary span { display: block; }
.summary span { margin-top: 4px; color: #546173; }
h2 { margin: 0 0 12px; font-size: 20px; }
h3 { margin: 16px 0 8px; font-size: 15px; color: #344154; }
table {
  width: 100%;
  border-collapse: collapse;
  font-size: 14px;
}
th, td {
  border-bottom: 1px solid #e4e9f1;
  padding: 8px 10px;
  text-align: left;
  white-space: nowrap;
}
th { background: #eef4f2; }
footer {
  color: #657386;
  font-size: 13px;
  padding: 8px 24px 32px;
  text-align: center;
}
"""

    def _site_html(
        self,
        tournament: dict[str, Any],
        settings: dict[str, Any],
        schedule: list[dict[str, Any]],
        players: list[dict[str, Any]],
        rounds: list[dict[str, Any]],
        standings: list[dict[str, Any]],
    ) -> str:
        tournament_id = int(tournament["id"])
        is_team_tournament = tournament.get("competition_type") == "team"
        round_sections = []
        for round_data in rounds:
            if is_team_tournament:
                round_sections.append(self._site_team_round_section(round_data))
            else:
                round_sections.append(self._site_individual_round_section(round_data))

        schedule_rows = "".join(
            "<tr>"
            f"<td>{item['round_number']}</td>"
            f"<td>{self._escape(item.get('date') or '')}</td>"
            f"<td>{self._escape(item.get('time') or '')}</td>"
            "</tr>"
            for item in schedule
            if item.get("date") or item.get("time")
        )
        players_rows = "".join(
            "<tr>"
            f"<td>{self._escape(player_full_name(player))}</td>"
            f"<td>{self._escape(player.get('title') or '')}</td>"
            f"<td>{self._escape(player.get('fide_id') or '')}</td>"
            f"<td>{self._escape(player.get('cbx_id') or '')}</td>"
            f"<td>{self._escape(player.get('lbx_id') or '')}</td>"
            f"<td>{player['rating']}</td>"
            f"<td>{self._escape(player['club'])}</td>"
            f"<td>{self._escape(player['category'])}</td>"
            f"<td>{self._escape(player.get('age_category') or '')}</td>"
            f"<td>{self._escape(player.get('rating_category') or '')}</td>"
            f"<td>{self._escape(player.get('prize_tags') or '')}</td>"
            f"<td>{self._escape(PLAYER_STATUSES.get(player.get('player_status', 'active'), player.get('player_status', '')))}</td>"
            "</tr>"
            for player in players
        )
        standings_section = self._site_standings_section(settings, standings, is_team_tournament)
        team_sections = self._site_team_sections(tournament_id) if is_team_tournament else ""

        generated_at = Database.now()
        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{self._escape(tournament['name'])}</title>
  <link rel="stylesheet" href="styles.css">
</head>
<body>
  <header>
    <p class="eyebrow">Albericus</p>
    <h1>{self._escape(tournament['name'])}</h1>
    <p>{self._escape(tournament['location'])}</p>
  </header>
  <main>
    <section class="summary">
      <div><strong>Periodo</strong><span>{self._escape(tournament['start_date'])} - {self._escape(tournament['end_date'])}</span></div>
      <div><strong>Rodadas</strong><span>{tournament['rounds_count']}</span></div>
      <div><strong>Ritmo</strong><span>{self._escape(tournament['time_control'])}</span></div>
      <div><strong>Competicao</strong><span>{self._escape(COMPETITION_TYPES.get(tournament.get('competition_type', 'individual'), 'Individual'))}</span></div>
      <div><strong>Status</strong><span>{self._escape(tournament['status'])}</span></div>
    </section>
    <section>
      <h2>Informacoes</h2>
      <dl>
        <dt>FIDE Event-ID</dt><dd>{self._escape(settings.get('fide_event_id') or '')}</dd>
        <dt>Organizador</dt><dd>{self._escape(settings.get('organizer') or '')}</dd>
        <dt>Diretor</dt><dd>{self._escape(settings.get('director') or '')}</dd>
        <dt>Arbitro principal</dt><dd>{self._escape(settings.get('chief_arbiter') or '')}</dd>
        <dt>Categorias</dt><dd>{self._escape(settings.get('categories') or '')}</dd>
      </dl>
    </section>
    <section>
      <h2>Agenda</h2>
      <table><thead><tr><th>Rodada</th><th>Data</th><th>Hora</th></tr></thead><tbody>{schedule_rows}</tbody></table>
    </section>
    {standings_section}
    {team_sections}
    <section>
      <h2>Jogadores</h2>
      <table><thead><tr><th>Nome</th><th>Titulo</th><th>FIDE</th><th>CBX</th><th>LBX</th><th>Rating</th><th>Clube</th><th>Categoria</th><th>Idade</th><th>Rating cat.</th><th>Tags</th><th>Status</th></tr></thead><tbody>{players_rows}</tbody></table>
    </section>
    {''.join(round_sections)}
  </main>
  <footer>Gerado em {self._escape(generated_at)} - Torneio #{tournament_id}</footer>
</body>
</html>
"""

    def _site_individual_round_section(self, round_data: dict[str, Any]) -> str:
        pairings = self.db.get_pairings_for_round(round_data["id"])
        rows = []
        for pairing in pairings:
            black_name = "BYE" if pairing["is_bye"] else pairing_player_name(pairing, "black")
            black_rating = "" if pairing["is_bye"] else pairing["black_rating"]
            qr_cell = ""
            if not pairing["is_bye"] and round_data.get("status") != "closed":
                url = self._qr_result_url(int(round_data["tournament_id"]), int(pairing["id"]))
                qr_cell = (
                    f"<a href='{self._escape(url)}'>Enviar</a><br>"
                    f"<img alt='QR mesa {pairing['board_number']}' width='84' height='84' "
                    f"src='{self._escape(self._qr_image_data_uri(url))}'>"
                )
            rows.append(
                "<tr>"
                f"<td>{pairing['board_number']}</td>"
                f"<td>{self._escape(pairing_player_name(pairing, 'white'))}</td>"
                f"<td>{pairing['white_rating']}</td>"
                f"<td>{self._escape(pairing['result'] or '')}</td>"
                f"<td>{self._escape(black_name or '')}</td>"
                f"<td>{self._escape(black_rating or '')}</td>"
                f"<td>{qr_cell}</td>"
                "</tr>"
            )
        return (
            "<section>"
            f"<h2>Rodada {round_data['number']}</h2>"
            f"<p>Status: {self._escape(round_data['status'])}</p>"
            "<table><thead><tr><th>Mesa</th><th>Brancas</th><th>Rating</th>"
            "<th>Resultado</th><th>Pretas</th><th>Rating</th><th>QR resultado</th></tr></thead>"
            f"<tbody>{''.join(rows)}</tbody></table>"
            "</section>"
        )

    @staticmethod
    def _qr_image_data_uri(url: str) -> str:
        try:
            import qrcode

            image = qrcode.make(url)
            buffer = io.BytesIO()
            image.save(buffer, format="PNG")
            return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")
        except Exception:
            return ""

    def _site_team_round_section(self, round_data: dict[str, Any]) -> str:
        rows = []
        for match in self.db.list_team_matches_for_round(int(round_data["id"])):
            match_score = self._team_match_score(match)
            if match.get("is_bye"):
                rows.append(
                    "<tr>"
                    f"<td>{match['match_number']}</td>"
                    f"<td>{self._escape(match.get('white_team_name') or '')}</td>"
                    f"<td>{self._escape(match_score)}</td>"
                    "<td>BYE</td><td></td><td></td><td>BYE</td><td></td>"
                    "</tr>"
                )
                continue
            for board in self.db.list_team_boards(int(match["id"])):
                rows.append(
                    "<tr>"
                    f"<td>{match['match_number']}</td>"
                    f"<td>{self._escape(match.get('white_team_name') or '')}</td>"
                    f"<td>{self._escape(match_score)}</td>"
                    f"<td>{self._escape(match.get('black_team_name') or '')}</td>"
                    f"<td>{board['board_number']}</td>"
                    f"<td>{self._escape(self._team_board_player_name(board, 'white'))}</td>"
                    f"<td>{self._escape(board.get('result') or '')}</td>"
                    f"<td>{self._escape(self._team_board_player_name(board, 'black'))}</td>"
                    "</tr>"
                )
        return (
            "<section>"
            f"<h2>Rodada {round_data['number']}</h2>"
            f"<p>Status: {self._escape(round_data['status'])}</p>"
            "<table><thead><tr><th>Confronto</th><th>Equipe A</th><th>Placar</th>"
            "<th>Equipe B</th><th>Tab.</th><th>Brancas</th><th>Resultado</th><th>Pretas</th></tr></thead>"
            f"<tbody>{''.join(rows)}</tbody></table>"
            "</section>"
        )

    def _site_standings_section(
        self,
        settings: dict[str, Any],
        standings: list[dict[str, Any]],
        is_team_tournament: bool,
    ) -> str:
        if settings.get("hide_standings"):
            return "<section><h2>Classificacao</h2><p>Classificacao ocultada pela organizacao.</p></section>"

        if is_team_tournament:
            standings_rows = "".join(
                "<tr>"
                f"<td>{item['position']}</td>"
                f"<td>{self._escape(item['name'])}</td>"
                f"<td>{self._escape(item.get('club') or '')}</td>"
                f"<td>{self._format_report_number(item['match_points'])}</td>"
                f"<td>{self._format_report_number(item['game_points'])}</td>"
                f"<td>{item['wins']}</td>"
                f"<td>{item['draws']}</td>"
                f"<td>{item['losses']}</td>"
                f"<td>{self._format_report_number(item['buchholz'])}</td>"
                "</tr>"
                for item in standings
            )
            return (
                "<section>"
                "<h2>Classificacao por equipes</h2>"
                "<table><thead><tr><th>Pos</th><th>Equipe</th><th>Clube/Cidade</th><th>MP</th>"
                "<th>GP</th><th>V</th><th>E</th><th>D</th><th>Buchholz</th></tr></thead>"
                f"<tbody>{standings_rows}</tbody></table>"
                "</section>"
            )

        standings_rows = "".join(
            "<tr>"
            f"<td>{item['position']}</td>"
            f"<td>{self._escape(item['name'])}</td>"
            f"<td>{self._escape(item.get('category') or '')}</td>"
            f"<td>{item['points']}</td>"
            f"<td>{item['buchholz']}</td>"
            f"<td>{item['buchholz_median']}</td>"
            f"<td>{item['sonneborn_berger']}</td>"
            f"<td>{item['wins']}</td>"
            f"<td>{self._escape(item['performance'])}</td>"
            "</tr>"
            for item in standings
        )
        tiebreak_rows = []
        for item in standings:
            components = dict(item.get("tiebreak_components") or {})
            buchholz = components.get("buchholz", {})
            median = components.get("buchholz_median", {})
            sb = components.get("sonneborn_berger", {})
            summary = " | ".join(
                part
                for part in (
                    f"Buchholz: {self._tiebreak_component_summary(buchholz)}",
                    f"Mediano: {self._tiebreak_component_summary(median)}",
                    f"SB: {self._tiebreak_component_summary(sb)}",
                )
                if part.strip()
            )
            tiebreak_rows.append(
                "<tr>"
                f"<td>{item['position']}</td>"
                f"<td>{self._escape(item['name'])}</td>"
                f"<td>{self._escape(summary)}</td>"
                "</tr>"
            )
        return (
            "<section>"
            "<h2>Classificacao</h2>"
            "<table><thead><tr><th>Pos</th><th>Jogador</th><th>Categoria</th><th>Pts</th>"
            "<th>Buchholz</th><th>Buchholz M</th><th>SB</th>"
            "<th>Vitorias</th><th>Perf.</th></tr></thead>"
            f"<tbody>{standings_rows}</tbody></table>"
            "<h3>Componentes de desempate</h3>"
            "<table><thead><tr><th>Pos</th><th>Jogador</th><th>Resumo</th></tr></thead>"
            f"<tbody>{''.join(tiebreak_rows)}</tbody></table>"
            "</section>"
        )

    def _site_team_sections(self, tournament_id: int) -> str:
        team_rows = []
        roster_rows = []
        for team in self.db.list_teams(tournament_id, active_only=False):
            team_rows.append(
                "<tr>"
                f"<td>{self._escape(team['name'])}</td>"
                f"<td>{self._escape(team.get('club') or '')}</td>"
                f"<td>{self._escape(team.get('captain') or '')}</td>"
                f"<td>{team.get('starters_count', 0)}</td>"
                f"<td>{team.get('players_count', 0)}</td>"
                f"<td>{'Ativa' if team.get('active') else 'Inativa'}</td>"
                "</tr>"
            )
            for assignment in self.db.list_team_players(int(team["id"]), active_only=False):
                board_label = assignment.get("board_number") or ""
                roster_rows.append(
                    "<tr>"
                    f"<td>{self._escape(team['name'])}</td>"
                    f"<td>{self._escape(board_label)}</td>"
                    f"<td>{self._escape(TEAM_PLAYER_ROLES.get(assignment.get('role', 'starter'), assignment.get('role', '')))}</td>"
                    f"<td>{self._escape(self._team_assignment_player_name(assignment))}</td>"
                    f"<td>{assignment.get('player_rating', '')}</td>"
                    f"<td>{self._escape(PLAYER_STATUSES.get(assignment.get('player_status', 'active'), assignment.get('player_status', '')))}</td>"
                    "</tr>"
                )
        return (
            "<section>"
            "<h2>Equipes</h2>"
            "<table><thead><tr><th>Equipe</th><th>Clube/Cidade</th><th>Capitao</th>"
            "<th>Titulares</th><th>Jogadores</th><th>Status</th></tr></thead>"
            f"<tbody>{''.join(team_rows)}</tbody></table>"
            "</section>"
            "<section>"
            "<h2>Escalacoes</h2>"
            "<table><thead><tr><th>Equipe</th><th>Tabuleiro</th><th>Funcao</th><th>Jogador</th>"
            "<th>Rating</th><th>Status</th></tr></thead>"
            f"<tbody>{''.join(roster_rows)}</tbody></table>"
            "</section>"
        )

    @staticmethod
    def _site_css() -> str:
        return """* { box-sizing: border-box; }
body {
  margin: 0;
  font-family: Arial, Helvetica, sans-serif;
  color: #0f172a;
  background: #f8fafc;
}
header {
  background: #0f172a;
  color: #fff;
  padding: 32px 40px;
}
header h1 { margin: 4px 0 8px; font-size: 32px; }
header p { margin: 0; color: #cbd5e1; }
.eyebrow { text-transform: uppercase; letter-spacing: 0; font-size: 12px; }
main { max-width: 1180px; margin: 0 auto; padding: 24px; }
section {
  background: #fff;
  border: 1px solid #e2e8f0;
  border-radius: 8px;
  margin-bottom: 16px;
  padding: 18px;
  overflow-x: auto;
}
h2 { margin: 0 0 12px; font-size: 20px; }
.summary {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(170px, 1fr));
  gap: 12px;
}
.summary div {
  border: 1px solid #e2e8f0;
  border-radius: 8px;
  padding: 12px;
}
.summary strong, .summary span { display: block; }
.summary span { margin-top: 4px; color: #475569; }
dl {
  display: grid;
  grid-template-columns: minmax(160px, 240px) 1fr;
  gap: 8px 12px;
  margin: 0;
}
dt { font-weight: 700; }
dd { margin: 0; color: #334155; }
table {
  width: 100%;
  border-collapse: collapse;
  font-size: 14px;
}
th, td {
  border-bottom: 1px solid #e2e8f0;
  padding: 8px 10px;
  text-align: left;
  white-space: nowrap;
}
th { background: #f1f5f9; }
footer {
  color: #64748b;
  font-size: 13px;
  padding: 8px 24px 32px;
  text-align: center;
}
"""

    @staticmethod
    def _escape(value: Any) -> str:
        return html.escape(str(value or ""))

    @staticmethod
    def _tournament_scope_label(tournament: dict[str, Any]) -> str:
        if tournament.get("class_id"):
            return TOURNAMENT_SCOPES["class"]
        if tournament.get("club_id"):
            return TOURNAMENT_SCOPES["club"]
        return TOURNAMENT_SCOPES["standalone"]

    def _tournament_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        rows = [
            ["Nome", tournament["name"]],
            ["Escopo", self._tournament_scope_label(tournament)],
            [
                "Competicao",
                COMPETITION_TYPES.get(tournament.get("competition_type", "individual"), "Individual"),
            ],
            ["Clube/Escola", tournament.get("club_name", "")],
            ["Turma", tournament.get("class_name", "")],
            ["Local", tournament["location"]],
            ["Data inicial", tournament["start_date"]],
            ["Data final", tournament["end_date"]],
            ["Sistema", tournament["system"]],
            ["Rodadas", tournament["rounds_count"]],
            ["Ritmo", tournament["time_control"]],
            ["Pontos do bye", tournament["bye_points"]],
            ["Status", tournament["status"]],
        ]
        settings = self.db.get_tournament_settings(tournament_id) or {}
        rows.extend(
            [
                ["FIDE Event-ID", settings.get("fide_event_id", "")],
                ["Organizador", settings.get("organizer", "")],
                ["Pagina web", settings.get("website", "")],
                ["E-mail", settings.get("contact_email", "")],
                ["Diretor", settings.get("director", "")],
                ["Arbitro principal", settings.get("chief_arbiter", "")],
                ["Federacao", settings.get("federation", "")],
                ["Estado", settings.get("state", "")],
                ["Categorias", settings.get("categories", "")],
                ["Data de corte", settings.get("cutoff_date", "")],
                ["Ordem inicial", settings.get("initial_order", "")],
                ["Tipo de torneio", settings.get("tournament_type", "")],
                ["Calcular desempenho", "Sim" if settings.get("calculate_performance") else "Nao"],
            ]
        )
        if tournament.get("competition_type") == "team":
            rows.extend(
                [
                    ["Tabuleiros por equipe", settings.get("team_boards_count", "")],
                    ["Pontos por vitoria da equipe", settings.get("team_match_win_points", "")],
                    ["Pontos por empate da equipe", settings.get("team_match_draw_points", "")],
                    ["Pontos por derrota da equipe", settings.get("team_match_loss_points", "")],
                    [
                        "Metodo por equipes",
                        TEAM_PAIRING_METHODS.get(
                            settings.get("team_pairing_method", "swiss"),
                            settings.get("team_pairing_method", ""),
                        ),
                    ],
                    [
                        "Criterio principal por equipes",
                        TEAM_STANDING_CRITERIA.get(
                            settings.get("team_standing_primary", "match_points"),
                            settings.get("team_standing_primary", ""),
                        ),
                    ],
                    [
                        "Criterio secundario por equipes",
                        TEAM_STANDING_CRITERIA.get(
                            settings.get("team_standing_secondary", "game_points"),
                            settings.get("team_standing_secondary", ""),
                        ),
                    ],
                    [
                        "Ordem fixa dos tabuleiros",
                        "Sim" if settings.get("team_fixed_board_order", 1) else "Nao",
                    ],
                ]
            )
        for item in self.db.list_round_schedule(tournament_id):
            if item["date"] or item["time"]:
                rows.append([f"Rodada {item['round_number']}", f"{item['date']} {item['time']}".strip()])
        return "Torneio", ["Campo", "Valor"], rows

    def _players_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        players = self.db.list_players(tournament_id, active_only=False)
        rows = [
            [
                player["id"],
                player_full_name(player),
                player.get("title", ""),
                player.get("fide_id", ""),
                player.get("cbx_id", ""),
                player.get("lbx_id", ""),
                player["rating"],
                player.get("national_rating", 0),
                player.get("international_rating", 0),
                player.get("starting_points", 0),
                player["club"],
                player["category"],
                player.get("age_category", ""),
                player.get("rating_category", ""),
                player.get("prize_tags", ""),
                player.get("birth_date", ""),
                player.get("sex", ""),
                PLAYER_STATUSES.get(player.get("player_status", "active"), player.get("player_status", "")),
            ]
            for player in players
        ]
        return (
            "Jogadores",
            [
                "ID",
                "Nome",
                "Titulo",
                "FIDE ID",
                "CBX ID",
                "LBX ID",
                "Rating",
                "Rating nacional",
                "Rating internacional",
                "Pontos iniciais",
                "Clube",
                "Categoria",
                "Categoria idade",
                "Categoria rating",
                "Tags premiacao",
                "Nascimento",
                "Sexo",
                "Status",
            ],
            rows,
        )

    def _initial_player_list_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        tournament = self.db.get_tournament(tournament_id)
        ordered_players = self._initial_ranked_players(
            self.db.list_players(tournament_id, active_only=False)
        )
        rows = [
            [
                index,
                player_full_name(player),
                self._trf_rating(player),
                player.get("club", ""),
                player.get("category", ""),
                PLAYER_STATUSES.get(player.get("player_status", "active"), player.get("player_status", "")),
                "",
            ]
            for index, player in enumerate(ordered_players, start=1)
        ]
        title = "Lista inicial de jogadores"
        if tournament:
            title = f"{title} - {tournament['name']}"
        return (
            title,
            ["Inicial", "Jogador", "Rating", "Clube", "Categoria", "Status", "Presenca"],
            rows,
        )

    def _rating_fee_sections(self, tournament_id: int) -> list[tuple[str, list[str], list[list[Any]]]]:
        summary = self.rating_fee_summary(tournament_id)
        summary_rows = [
            ["Torneio", summary["tournament_name"]],
            ["Inscritos", summary["players_count"]],
            ["Rated em alguma base", summary["rated_players"]],
            ["Nao rated", summary["unrated_players"]],
            ["Com ID oficial", summary["players_with_base"]],
            ["Sem ID oficial", summary["players_without_base"]],
            ["Total das taxas", self._format_currency(summary["total"])],
            ["Regra", "Cada base cobra separadamente os inscritos identificados nela."],
        ]
        base_rows = [
            [
                item["base"],
                item["identified"],
                item["rated"],
                item["unrated"],
                self._format_currency(item["unit_fee"]),
                self._format_currency(item["subtotal"]),
            ]
            for item in summary["bases"]
        ]
        player_rows = []
        for player in summary["players"]:
            bases = [
                label
                for label, field in (("FIDE", "fide_id"), ("CBX", "cbx_id"), ("LBX", "lbx_id"))
                if str(player.get(field) or "").strip()
            ]
            player_rows.append(
                [
                    player["name"],
                    player.get("fide_id", ""),
                    player.get("international_rating", 0),
                    player.get("cbx_id", ""),
                    player.get("national_rating", 0),
                    player.get("lbx_id", ""),
                    player.get("rating", 0),
                    ", ".join(bases) or "Sem base",
                ]
            )
        return [
            ("Resumo de taxas de rating", ["Campo", "Valor"], summary_rows),
            ("Taxas por base", ["Base", "Identificados", "Rated", "Nao rated", "Taxa unitaria", "Subtotal"], base_rows),
            (
                "Inscritos por base",
                ["Jogador", "FIDE ID", "Rating FIDE", "CBX ID", "Rating CBX", "LBX ID", "Rating LBX", "Bases cobradas"],
                player_rows,
            ),
        ]

    FIDE_RATING_TYPE_LABELS = {"fide": "FIDE", "cbx": "CBX"}

    def _fide_rating_sections(
        self, tournament_id: int, rating_type: str = "fide"
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if rating_type not in self.FIDE_RATING_TYPE_LABELS:
            raise AppError("Tipo de rating invalido para o relatorio FIDE.")
        if tournament.get("competition_type") == "team":
            raise AppError("Relatorio de variacao de rating disponivel apenas para torneios individuais.")

        players = self.db.list_players(tournament_id, active_only=False)
        closed = self.db.get_pairings_for_tournament(tournament_id, closed_only=True)
        start_date = str(tournament.get("start_date") or "").strip()
        year = int(start_date[:4]) if len(start_date) >= 4 and start_date[:4].isdigit() else None
        rows = build_fide_report_rows(players, closed, rating_type, tournament_year=year)
        # Snapshot persistido para auditoria/reimpressao (idempotente por base).
        self.db.save_fide_rating_report(tournament_id, rating_type, rows)

        base_label = self.FIDE_RATING_TYPE_LABELS[rating_type]
        rated = [row for row in rows if row.get("ro")]
        deltas = [float(row["delta"]) for row in rated if row.get("delta") is not None]

        def cell(value: Any) -> Any:
            return "-" if value is None else value

        notice_rows = [
            ["Estimativa de apoio ao arbitro (We, fator K, Rc, Rp)."],
            [f"Nao substitui a homologacao oficial da {base_label}."],
            [f"So conta partidas jogadas contra adversarios com rating na base {base_label}."],
            ["Diferencas de rating acima de 400 sao tratadas como 400 (regra dos 400)."],
        ]
        summary_rows = [
            ["Torneio", tournament["name"]],
            ["Base de rating", base_label],
            ["Jogadores no relatorio", len(rows)],
            ["Com rating", len(rated)],
            ["Sem rating (somente performance)", len(rows) - len(rated)],
            ["Variacao total (ΔElo)", round(sum(deltas), 2)],
            ["Maior ganho", round(max(deltas), 2) if deltas else 0.0],
            ["Maior perda", round(min(deltas), 2) if deltas else 0.0],
        ]
        detail_rows = [
            [
                row["name"],
                row["ro"] if row.get("ro") else "-",
                row["k"] if row.get("ro") else "-",
                row["games_rated"],
                row["score"],
                cell(row["we"]),
                cell(row["delta"]),
                cell(row["rc"]),
                row["rp"],
                row["n_over_400"],
            ]
            for row in rows
        ]
        return [
            ("Aviso", ["Observacao"], notice_rows),
            ("Resumo do relatorio de rating", ["Campo", "Valor"], summary_rows),
            (
                f"Variacao de rating {base_label}",
                ["Jogador", "Ro", "K", "n", "Pts", "We", "ΔElo", "Rc", "Rp", ">400"],
                detail_rows,
            ),
        ]

    def _prize_sections(
        self, tournament_id: int
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") == "team":
            raise AppError("Distribuicao de premios disponivel apenas para torneios individuais.")

        settings = self.db.get_tournament_settings(tournament_id) or {}
        prizes = self.db.list_tournament_prizes(tournament_id)
        standings = self.pairing_service.standings(tournament_id)
        result = allocate_prizes(
            standings,
            prizes,
            policy=str(settings.get("prize_policy") or "best_only"),
            tax_percent=float(settings.get("prize_tax_percent") or 0.0),
        )

        policy_label = PRIZE_POLICIES.get(result["policy"], result["policy"])
        summary_rows = [
            ["Torneio", tournament["name"]],
            ["Politica de premiacao", policy_label],
            ["Imposto do organizador (%)", result["tax_percent"]],
            ["Premiados", result["winners"]],
            ["Total bruto", self._format_currency(result["total_gross"])],
            ["Total imposto", self._format_currency(result["total_tax"])],
            ["Total liquido distribuido", self._format_currency(result["total_net"])],
        ]
        winner_rows = [
            [
                allocation["position"],
                allocation["name"],
                allocation["category"],
                allocation["points"],
                self._format_currency(allocation["overall"]),
                self._format_currency(allocation["category_prize"]),
                self._format_currency(allocation["gross"]),
                self._format_currency(allocation["net"]),
            ]
            for allocation in result["allocations"]
        ]
        sections: list[tuple[str, list[str], list[list[Any]]]] = [
            ("Resumo da premiacao", ["Campo", "Valor"], summary_rows),
            (
                "Premiacao por jogador",
                ["Pos", "Jogador", "Categoria", "Pts", "Premio geral", "Premio categoria", "Bruto", "Liquido"],
                winner_rows,
            ),
        ]
        if result["manual_prizes"]:
            manual_rows = [
                [
                    PRIZE_KINDS.get(prize["kind"], prize["kind"]),
                    prize["label"],
                    prize["category"],
                    self._format_currency(prize["amount"]),
                ]
                for prize in result["manual_prizes"]
            ]
            sections.append(
                (
                    "Premios manuais (definir ganhador)",
                    ["Tipo", "Premio", "Categoria", "Valor"],
                    manual_rows,
                )
            )
        return sections

    def _individual_tournament(self, tournament_id: int, what: str) -> dict[str, Any]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") == "team":
            raise AppError(f"{what} disponivel apenas para torneios individuais.")
        return tournament

    def _federation_sections(
        self, tournament_id: int
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        tournament = self._individual_tournament(tournament_id, "Estatistica de federacoes")
        players = self.db.list_players(tournament_id, active_only=False)
        federation_by_id = {
            int(player["id"]): (str(player.get("federation_id") or "").strip().upper() or "—")
            for player in players
        }
        standings = self.pairing_service.standings(tournament_id)

        groups: dict[str, list[float]] = {}
        for standing in standings:
            federation = federation_by_id.get(int(standing["player_id"]), "—")
            bucket = groups.setdefault(federation, [0.0, 0.0, 0.0])
            bucket[0] += 1
            bucket[1] += float(standing.get("points") or 0.0)
            bucket[2] += sum(1 for game in standing.get("games", []) if game.get("color") != "bye")

        total_players = int(sum(bucket[0] for bucket in groups.values())) or 1
        rows = []
        for federation, (count, points, games) in sorted(
            groups.items(), key=lambda item: (-item[1][0], item[0])
        ):
            rows.append(
                [
                    federation,
                    int(count),
                    f"{100 * count / total_players:.1f}%",
                    self._format_report_number(round(points, 2)),
                    int(games),
                    self._format_report_number(round(points / count, 2)) if count else "0",
                ]
            )
        summary_rows = [
            ["Torneio", tournament["name"]],
            ["Federacoes", len(groups)],
            ["Jogadores", total_players],
        ]
        return [
            ("Resumo de federacoes", ["Campo", "Valor"], summary_rows),
            (
                "Estatistica de federacoes",
                ["Federacao", "Jogadores", "% jogadores", "Pontos", "Partidas", "Media pts"],
                rows,
            ),
        ]

    def _game_statistics_sections(
        self, tournament_id: int
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        tournament = self._individual_tournament(tournament_id, "Estatistica de partidas")
        closed = self.db.get_pairings_for_tournament(tournament_id, closed_only=True)

        white_wins = draws = black_wins = walkovers = byes = 0
        for pairing in closed:
            if pairing.get("is_bye"):
                byes += 1
                continue
            result = pairing.get("result")
            if result == "1-0":
                white_wins += 1
            elif result == "0-1":
                black_wins += 1
            elif result == "1/2-1/2":
                draws += 1
            elif result in ("1F-0F", "0F-1F", "0F-0F"):
                walkovers += 1

        played = white_wins + draws + black_wins

        def percent(value: int) -> str:
            return f"{100 * value / played:.1f}%" if played else "0%"

        distribution_rows = [
            ["Vitorias de brancas", white_wins, percent(white_wins)],
            ["Empates", draws, percent(draws)],
            ["Vitorias de pretas", black_wins, percent(black_wins)],
        ]
        summary_rows = [
            ["Torneio", tournament["name"]],
            ["Partidas jogadas (tabuleiro)", played],
            ["WO / forfait", walkovers],
            ["Byes", byes],
            ["Total de pareamentos fechados", len(closed)],
        ]
        return [
            ("Resumo de partidas", ["Campo", "Valor"], summary_rows),
            ("Distribuicao de resultados", ["Resultado", "Quantidade", "%"], distribution_rows),
        ]

    def _player_cards_sections(
        self, tournament_id: int
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        self._individual_tournament(tournament_id, "Fichas individuais")
        players = self.db.list_players(tournament_id, active_only=False)
        federation_by_id = {
            int(player["id"]): str(player.get("federation_id") or "").strip().upper()
            for player in players
        }
        standings = self.pairing_service.standings(tournament_id)

        color_labels = {"white": "Brancas", "black": "Pretas", "bye": "Bye"}
        summary_rows = []
        game_rows = []
        for standing in standings:
            games = standing.get("games", [])
            board_games = [game for game in games if game.get("color") != "bye"]
            wins = sum(1 for game in board_games if float(game.get("earned") or 0.0) == 1.0)
            draws = sum(1 for game in board_games if float(game.get("earned") or 0.0) == 0.5)
            losses = sum(1 for game in board_games if float(game.get("earned") or 0.0) == 0.0)
            performance = standing.get("performance")
            summary_rows.append(
                [
                    standing.get("position", 0),
                    standing.get("name", ""),
                    standing.get("category", ""),
                    federation_by_id.get(int(standing["player_id"]), ""),
                    self._format_report_number(standing.get("points", 0)),
                    performance if isinstance(performance, int) else "-",
                    wins,
                    draws,
                    losses,
                    standing.get("white_count", 0),
                    standing.get("black_count", 0),
                    standing.get("byes", 0),
                ]
            )
            running = float(standing.get("starting_points", 0.0) or 0.0)
            for game in sorted(games, key=lambda item: int(item.get("round") or 0)):
                running += float(game.get("earned") or 0.0)
                game_rows.append(
                    [
                        standing.get("position", 0),
                        standing.get("name", ""),
                        int(game.get("round") or 0),
                        color_labels.get(str(game.get("color")), ""),
                        game.get("opponent_name", "") or "BYE",
                        game.get("result", "") or "",
                        self._format_report_number(round(running, 2)),
                    ]
                )
        return [
            (
                "Resumo por jogador",
                ["Pos", "Jogador", "Categoria", "Fed", "Pts", "Rp", "V", "E", "D", "Brancas", "Pretas", "Byes"],
                summary_rows,
            ),
            (
                "Resultados rodada a rodada",
                ["Pos", "Jogador", "Rodada", "Cor", "Adversario", "Resultado", "Pts acum."],
                game_rows,
            ),
        ]

    def _norm_sections(
        self, tournament_id: int
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        self._individual_tournament(tournament_id, "Normas FIDE")
        players = self.db.list_players(tournament_id, active_only=False)
        closed = self.db.get_pairings_for_tournament(tournament_id, closed_only=True)
        report = build_norm_report(players, closed, "fide")

        notice_rows = [
            ["Estimativa de apoio ao arbitro: indica se o jogador atingiu indicadores"],
            ["compativeis com uma norma. NAO concede norma nem titulo (exclusivo da FIDE)."],
            ["So conta partidas jogadas contra adversarios com rating."],
        ]
        summary_rows = [
            [
                item["name"],
                item["sex"],
                item["performance"],
                item["games"],
                self._format_report_number(item["average_opponent"]),
                item["federations"],
                item["titled_opponents"],
                ", ".join(item["achieved"]) or "-",
            ]
            for item in report
        ]
        detail_rows = []
        for item in report:
            for title in item["titles"]:
                detail_rows.append(
                    [
                        item["name"],
                        title["label"],
                        "Sim" if title["meets"] else "Nao",
                        "OK" if title["meets"] else "; ".join(title["missing"]),
                    ]
                )
        return [
            ("Aviso", ["Observacao"], notice_rows),
            (
                "Indicadores de norma por jogador",
                ["Jogador", "Sexo", "Rp", "Partidas", "Media adv.", "Federacoes", "Titulados", "Norma atingida"],
                summary_rows,
            ),
            ("Detalhe por titulo", ["Jogador", "Titulo", "Atende?", "Pendencias"], detail_rows),
        ]

    _ARBITER_ROLE_LABELS = {
        "chief": "Arbitro principal",
        "chief_arbiter": "Arbitro principal",
        "main": "Arbitro principal",
        "principal": "Arbitro principal",
        "deputy": "Arbitro adjunto",
        "sector": "Arbitro de setor",
        "arbiter": "Arbitro",
        "assistant": "Assistente",
    }

    def _arbiter_norm_sections(
        self, tournament_id: int
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        settings = self.db.get_tournament_settings(tournament_id) or {}
        players = self.db.list_players(tournament_id, active_only=False)
        rated = sum(1 for player in players if self._trf_rating(player) > 0)
        referees = self.db.list_tournament_referees(tournament_id)
        competition = "Equipes" if tournament.get("competition_type") == "team" else "Individual"

        notice_rows = [
            ["Documento de apoio para a norma de arbitro (IA/FA)."],
            ["Nao e o formulario oficial da FIDE; use os dados abaixo para preenche-lo."],
        ]
        tournament_rows = [
            ["Torneio", tournament["name"]],
            ["Local", tournament.get("location", "")],
            ["Federacao", settings.get("federation", "")],
            ["FIDE Event-ID", settings.get("fide_event_id", "")],
            ["Data de inicio", tournament.get("start_date", "")],
            ["Data de termino", tournament.get("end_date", "")],
            ["Ritmo de jogo", tournament.get("time_control", "")],
            ["Rodadas", tournament.get("rounds_count", "")],
            ["Tipo", competition],
            ["Jogadores", len(players)],
            ["Jogadores com rating", rated],
            ["Organizador", settings.get("organizer", "")],
            ["Arbitro principal (config.)", settings.get("chief_arbiter", "")],
        ]

        arbiter_rows: list[list[Any]] = []
        for referee in referees:
            role = str(referee.get("role") or "").strip().lower()
            arbiter_rows.append(
                [
                    referee.get("name", ""),
                    self._ARBITER_ROLE_LABELS.get(role, role.replace("_", " ").title() or "Arbitro"),
                    referee.get("fide_id", ""),
                    referee.get("category", ""),
                    "",
                    "",
                ]
            )
        if not arbiter_rows:
            chief = str(settings.get("chief_arbiter") or "").strip()
            if chief:
                arbiter_rows.append([chief, "Arbitro principal", "", "", "", ""])
            for name in str(settings.get("arbiters") or "").replace(";", ",").split(","):
                cleaned = name.strip()
                if cleaned:
                    arbiter_rows.append([cleaned, "Arbitro", "", "", "", ""])

        return [
            ("Aviso", ["Observacao"], notice_rows),
            ("Dados do torneio (norma de arbitro)", ["Campo", "Valor"], tournament_rows),
            (
                "Arbitros designados",
                ["Nome", "Funcao", "FIDE ID", "Categoria", "Norma (IA/FA)", "Assinatura"],
                arbiter_rows,
            ),
        ]

    def _category_winners_section(
        self, tournament_id: int, top_n: int = 3
    ) -> tuple[str, list[str], list[list[Any]]]:
        """Vencedores por categoria (top N de cada categoria, na ordem da classificacao)."""
        groups: dict[str, list[dict[str, Any]]] = {}
        for item in self.pairing_service.standings(tournament_id):
            category = str(item.get("category") or "").strip()
            if category:
                groups.setdefault(category, []).append(item)
        rows: list[list[Any]] = []
        for category in sorted(groups):
            for rank, item in enumerate(groups[category][: max(1, top_n)], start=1):
                rows.append(
                    [
                        category,
                        rank,
                        item.get("name", ""),
                        item.get("points", ""),
                        item.get("performance", ""),
                    ]
                )
        if not rows:
            rows = [["", "", "Sem categorias definidas neste torneio.", "", ""]]
        return ("Vencedores por categoria", ["Categoria", "Pos", "Jogador", "Pts", "Performance"], rows)

    def _round_bulletin_sections(
        self, round_id: int
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        """Secoes do boletim da rodada: cabecalho + resultados + classificacao + destaques."""
        round_data = self.db.get_round(round_id)
        if not round_data:
            raise AppError("Rodada nao encontrada.")
        tournament = self.db.get_tournament(int(round_data["tournament_id"]))
        if not tournament:
            raise AppError("Torneio nao encontrado.")
        number = round_data["number"]

        header_rows = [
            ["Torneio", tournament["name"]],
            ["Rodada", number],
            ["Estado", round_data.get("status", "")],
            ["Local", tournament.get("location", "")],
            ["Data", round_data.get("scheduled_at", "") or tournament.get("start_date", "")],
        ]
        sections: list[tuple[str, list[str], list[list[Any]]]] = [
            (f"Boletim da rodada {number}", ["Campo", "Valor"], header_rows),
        ]
        sections.append(self._pairings_section(round_id))
        # Classificacao apos a rodada (top 10) reaproveita o layout da classificacao.
        title, headers, rows = self._standings_section(int(tournament["id"]))
        sections.append((f"Classificacao apos a rodada {number} (top 10)", headers, rows[:10]))
        if tournament.get("competition_type") != "team":
            sections.append(self._bulletin_highlights_section(round_id, int(tournament["id"])))
        return sections

    def _bulletin_highlights_section(
        self, round_id: int, tournament_id: int
    ) -> tuple[str, list[str], list[list[Any]]]:
        """Destaques da rodada (individual): lider, decisivas/empates e maior zebra."""
        pairings = [item for item in self.db.get_pairings_for_round(round_id) if not item.get("is_bye")]
        decisive = sum(1 for item in pairings if item.get("result") in ("1-0", "0-1"))
        draws = sum(1 for item in pairings if item.get("result") == "1/2-1/2")
        standings = self.pairing_service.standings(tournament_id)

        best_upset: tuple[str, str, int] | None = None
        for item in pairings:
            result = item.get("result")
            white_rating = int(item.get("white_rating") or 0)
            black_rating = int(item.get("black_rating") or 0)
            if not white_rating or not black_rating:
                continue
            if result == "1-0" and black_rating > white_rating:
                diff, winner, loser = black_rating - white_rating, pairing_player_name(item, "white"), pairing_player_name(item, "black")
            elif result == "0-1" and white_rating > black_rating:
                diff, winner, loser = white_rating - black_rating, pairing_player_name(item, "black"), pairing_player_name(item, "white")
            else:
                continue
            if best_upset is None or diff > best_upset[2]:
                best_upset = (winner, loser, diff)

        rows: list[list[Any]] = []
        if standings:
            rows.append(["Lider", f"{standings[0].get('name', '')} ({standings[0].get('points', '')} pts)"])
        rows.append(["Partidas decididas", decisive])
        rows.append(["Empates", draws])
        if best_upset:
            rows.append(["Maior zebra", f"{best_upset[0]} venceu {best_upset[1]} (+{best_upset[2]} de rating)"])
        return ("Destaques", ["Item", "Valor"], rows)

    def _podium_data(self, tournament_id: int) -> dict[str, Any]:
        """Dados do podio: top 3 (jogadores ou equipes) + campeoes por categoria."""
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        is_team = tournament.get("competition_type") == "team"
        top: list[dict[str, Any]] = []
        categories: list[dict[str, Any]] = []
        if is_team:
            for index, item in enumerate(self.pairing_service.team_standings(tournament_id)[:3], start=1):
                top.append(
                    {
                        "rank": index,
                        "name": str(item.get("name") or item.get("team_name") or ""),
                        "points": item.get("match_points", ""),
                        "detail": f"GP {item.get('game_points', '')}",
                    }
                )
        else:
            standings = self.pairing_service.standings(tournament_id)
            for index, item in enumerate(standings[:3], start=1):
                top.append(
                    {
                        "rank": index,
                        "name": str(item.get("name") or ""),
                        "points": item.get("points", ""),
                        "detail": f"perf {item.get('performance', '')}",
                    }
                )
            champions: dict[str, str] = {}
            for item in standings:
                category = str(item.get("category") or "").strip()
                if category and category not in champions:
                    champions[category] = str(item.get("name") or "")
            categories = [{"category": cat, "name": name} for cat, name in sorted(champions.items())]
        return {"tournament": tournament, "top": top, "categories": categories, "is_team": is_team}

    def _write_podium_poster_pdf(self, path: Path, data: dict[str, Any]) -> None:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
        except ImportError as exc:
            raise AppError("Instale reportlab para gerar o poster do podio (pip install reportlab).") from exc

        tournament = data["tournament"]
        width, height = A4
        center = width / 2
        document = canvas.Canvas(str(path), pagesize=A4)
        document.setLineWidth(2)
        document.rect(30, 30, width - 60, height - 60)

        document.setFont("Helvetica-Bold", 24)
        document.drawCentredString(center, height - 90, str(tournament.get("name") or "Torneio"))
        document.setFont("Helvetica", 14)
        document.drawCentredString(center, height - 116, "Premiacao - Podio")
        subtitle = " - ".join(
            part for part in [str(tournament.get("location") or ""), str(tournament.get("start_date") or "")] if part
        )
        if subtitle:
            document.setFont("Helvetica", 11)
            document.drawCentredString(center, height - 136, subtitle)

        medals = {1: "1o lugar", 2: "2o lugar", 3: "3o lugar"}
        y = height - 230
        for entry in data["top"]:
            rank = int(entry["rank"])
            document.setFont("Helvetica-Bold", 22 if rank == 1 else 16)
            document.drawCentredString(center, y, f"{medals.get(rank, f'{rank}o')}: {entry['name']}")
            document.setFont("Helvetica", 12)
            document.drawCentredString(center, y - 20, f"{entry['points']} pts  ({entry['detail']})")
            y -= 70 if rank == 1 else 60
        if not data["top"]:
            document.setFont("Helvetica", 12)
            document.drawCentredString(center, y, "Sem classificacao disponivel.")

        if data["categories"]:
            y -= 20
            document.setFont("Helvetica-Bold", 14)
            document.drawCentredString(center, y, "Campeoes por categoria")
            y -= 24
            document.setFont("Helvetica", 12)
            for champion in data["categories"]:
                if y < 80:
                    break
                document.drawCentredString(center, y, f"{champion['category']}: {champion['name']}")
                y -= 20

        document.setFont("Helvetica-Oblique", 9)
        document.drawCentredString(center, 50, "Gerado pelo Albericus")
        document.save()

    def _tournament_minutes_sections(
        self, tournament_id: int
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        """Secoes da ata final: cabecalho + classificacao + categorias + premiacao + taxas + arbitros."""
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        settings = self.db.get_tournament_settings(tournament_id) or {}
        players = self.db.list_players(tournament_id, active_only=False)
        rated = sum(1 for player in players if self._trf_rating(player) > 0)
        competition = "Equipes" if tournament.get("competition_type") == "team" else "Individual"

        header_rows = [
            ["Torneio", tournament["name"]],
            ["Local", tournament.get("location", "")],
            ["Federacao", settings.get("federation", "")],
            ["FIDE Event-ID", settings.get("fide_event_id", "")],
            ["Data de inicio", tournament.get("start_date", "")],
            ["Data de termino", tournament.get("end_date", "")],
            ["Ritmo de jogo", tournament.get("time_control", "")],
            ["Rodadas", tournament.get("rounds_count", "")],
            ["Tipo", competition],
            ["Jogadores", len(players)],
            ["Jogadores com rating", rated],
            ["Organizador", settings.get("organizer", "")],
            ["Diretor", settings.get("director", "")],
            ["Arbitro principal", settings.get("chief_arbiter", "")],
        ]

        referees = self.db.list_tournament_referees(tournament_id)
        arbiter_rows: list[list[Any]] = []
        for referee in referees:
            role = str(referee.get("role") or "").strip().lower()
            arbiter_rows.append(
                [
                    referee.get("name", ""),
                    self._ARBITER_ROLE_LABELS.get(role, role.replace("_", " ").title() or "Arbitro"),
                    referee.get("fide_id", ""),
                    referee.get("category", ""),
                    "",
                ]
            )
        if not arbiter_rows:
            chief = str(settings.get("chief_arbiter") or "").strip()
            if chief:
                arbiter_rows.append([chief, "Arbitro principal", "", "", ""])
            for name in str(settings.get("arbiters") or "").replace(";", ",").split(","):
                cleaned = name.strip()
                if cleaned:
                    arbiter_rows.append([cleaned, "Arbitro", "", "", ""])

        sections: list[tuple[str, list[str], list[list[Any]]]] = [
            ("Ata final do torneio", ["Campo", "Valor"], header_rows),
            self._standings_section(tournament_id),
        ]
        if tournament.get("competition_type") != "team":
            sections.append(self._category_winners_section(tournament_id))
        # Premiacao e taxas sao opcionais: AppError quando nao se aplicam.
        try:
            sections.extend(self._prize_sections(tournament_id))
        except AppError:
            pass
        try:
            sections.extend(self._rating_fee_sections(tournament_id))
        except AppError:
            pass
        sections.append(
            ("Arbitros e assinaturas", ["Nome", "Funcao", "FIDE ID", "Categoria", "Assinatura"], arbiter_rows)
        )
        return sections

    @staticmethod
    def _format_currency(value: Any) -> str:
        return f"R$ {float(value or 0):.2f}".replace(".", ",")

    @staticmethod
    def _initial_ranked_players(players: list[Mapping[str, Any]]) -> list[Mapping[str, Any]]:
        return sorted(
            players,
            key=lambda player: (
                -ExportService._trf_rating(player),
                player_pairing_name(player).casefold(),
                int(player.get("id") or 0),
            ),
        )

    def _teams_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") != "team":
            raise AppError("Este relatorio esta disponivel apenas para torneios por equipes.")
        rows = [
            [
                team["id"],
                team["name"],
                team.get("club", ""),
                team.get("captain", ""),
                team.get("starters_count", 0),
                team.get("players_count", 0),
                "Ativa" if team.get("active") else "Inativa",
                team.get("notes", ""),
            ]
            for team in self.db.list_teams(tournament_id, active_only=False)
        ]
        return (
            "Equipes",
            ["ID", "Equipe", "Clube/Cidade", "Capitao", "Titulares", "Jogadores", "Status", "Observacoes"],
            rows,
        )

    def _team_rosters_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") != "team":
            raise AppError("Este relatorio esta disponivel apenas para torneios por equipes.")

        rows = []
        for team in self.db.list_teams(tournament_id, active_only=False):
            for assignment in self.db.list_team_players(int(team["id"]), active_only=False):
                rows.append(
                    [
                        team["name"],
                        assignment.get("board_number") or "",
                        TEAM_PLAYER_ROLES.get(assignment.get("role", "starter"), assignment.get("role", "")),
                        self._team_assignment_player_name(assignment),
                        assignment.get("player_rating", ""),
                        assignment.get("player_club", ""),
                        assignment.get("player_category", ""),
                        assignment.get("player_age_category", ""),
                        assignment.get("player_rating_category", ""),
                        assignment.get("player_prize_tags", ""),
                        PLAYER_STATUSES.get(
                            assignment.get("player_status", "active"),
                            assignment.get("player_status", ""),
                        ),
                    ]
                )
        return (
            "Escalacoes",
            [
                "Equipe",
                "Tabuleiro",
                "Funcao",
                "Jogador",
                "Rating",
                "Clube",
                "Categoria",
                "Categoria idade",
                "Categoria rating",
                "Tags premiacao",
                "Status",
            ],
            rows,
        )

    def _club_report_sections(self) -> list[tuple[str, list[str], list[list[Any]]]]:
        clubs = self.db.list_clubs(active_only=False)
        summary = self.db.club_summary()
        members = self.db.list_members(active_only=False)
        tournaments = self.db.list_tournaments()
        classes = self.db.list_classes(active_only=False)
        guardians = self.db.list_guardians(active_only=False)
        sessions = self.db.list_training_sessions()
        plans = self.db.list_membership_plans(active_only=False)
        events = self.db.list_club_events()
        finance_service = __import__('src.services.finance_service', fromlist=['FinanceService']).FinanceService(self.db)
        finance_summary = finance_service.finance_summary()
        payments = finance_service.payments_report()
        minor_members_without_guardians = __import__('src.services.member_service', fromlist=['GuardianService']).GuardianService(self.db).minor_members_without_guardians()

        member_ratings = [int(member.get("rating") or 0) for member in members if int(member.get("rating") or 0) > 0]
        average_rating = round(sum(member_ratings) / len(member_ratings), 1) if member_ratings else 0

        club_rows = [
            [
                club["id"],
                club["name"],
                club.get("kind", ""),
                "Sim" if club.get("active") else "Nao",
                club.get("city", ""),
                club.get("phone", ""),
                club.get("email", ""),
                club.get("members_count", 0),
                club.get("active_classes_count", 0),
                club.get("tournaments_count", 0),
            ]
            for club in clubs
        ]
        class_rows = [
            [
                item["id"],
                item["club_name"],
                item["name"],
                item["teacher"],
                item["weekday"],
                item["time"],
                item["location"],
                "Sim" if item["active"] else "Nao",
                item["active_members_count"],
            ]
            for item in classes
        ]
        session_rows = [
            [
                item["id"],
                item["title"],
                TRAINING_SESSION_TYPES.get(item["session_type"], item["session_type"]),
                item["session_date"],
                item["start_time"],
                item["club_name"],
                item["class_name"],
                item.get("learning_level_name") or "",
                item.get("objective") or "",
                item.get("content") or "",
                item.get("homework") or "",
                item["instructor"],
                TRAINING_SESSION_STATUSES.get(item["status"], item["status"]),
                item["present_count"],
                item["absent_count"],
                item["justified_count"],
            ]
            for item in sessions[:100]
        ]
        guardian_rows = [
            [
                guardian["id"],
                guardian["name"],
                guardian["phone"],
                guardian["email"],
                guardian["document"],
                "Sim" if guardian.get("active") else "Nao",
                guardian["members_count"],
                guardian["active_members_count"],
            ]
            for guardian in guardians
        ]
        plan_rows = [
            [
                plan["id"],
                plan["name"],
                plan["amount"],
                BILLING_CYCLES.get(plan["billing_cycle"], plan["billing_cycle"]),
                "Sim" if plan.get("active") else "Nao",
                plan["payments_count"],
            ]
            for plan in plans
        ]
        payment_rows = [
            [
                payment["id"],
                payment["member_name"],
                payment.get("plan_name") or "",
                payment["description"],
                payment["reference_period"],
                payment["due_date"],
                payment["payment_date"],
                payment["amount"],
                PAYMENT_STATUSES.get(payment["effective_status"], payment["effective_status"]),
                payment["method"],
            ]
            for payment in payments[:100]
        ]
        event_rows = [
            [
                event["id"],
                event["title"],
                EVENT_TYPES.get(event["event_type"], event["event_type"]),
                event["event_date"],
                event["start_time"],
                event["club_name"],
                event.get("tournament_name") or "",
                event["location"],
                EVENT_STATUSES.get(event["status"], event["status"]),
            ]
            for event in events[:100]
        ]
        indicator_rows = [
            ["Unidades ativas", summary.get("active_clubs", 0)],
            ["Turmas ativas", summary.get("active_classes", 0)],
            ["Membros ativos", summary.get("active_members", 0)],
            ["Total de membros", summary.get("total_members", 0)],
            ["Responsaveis cadastrados", len(guardians)],
            ["Menores sem responsavel", len(minor_members_without_guardians)],
            ["Aulas/treinos cadastrados", len(sessions)],
            ["Eventos no calendario", len(events)],
            ["Planos financeiros", len(plans)],
            ["Lancamentos financeiros", finance_summary["total_payments"]],
            ["Recebido", finance_summary["paid_amount"]],
            ["Pendente", finance_summary["pending_amount"]],
            ["Atrasado", finance_summary["late_amount"]],
            ["Torneios cadastrados", summary.get("total_tournaments", 0)],
            ["Torneios em andamento", summary.get("running_tournaments", 0)],
            ["Membros com rating interno", len(member_ratings)],
            ["Rating interno medio", average_rating],
        ]
        status_rows = self._count_member_field(members, "status")
        type_rows = self._count_member_field(members, "member_type")
        category_rows = self._count_member_field(members, "category")
        ranking_rows = [
            [
                index,
                member["name"],
                member.get("club_name", ""),
                member.get("active_class_name", ""),
                member["rating"],
                member["category"],
                member["member_type"],
                member["status"],
            ]
            for index, member in enumerate(
                sorted(
                    members,
                    key=lambda item: (
                        -int(item.get("rating") or 0),
                        str(item.get("name") or "").casefold(),
                    ),
                )[:50],
                start=1,
            )
            if int(member.get("rating") or 0) > 0
        ]
        tournament_rows = [
            [
                tournament["id"],
                tournament["name"],
                self._tournament_scope_label(tournament),
                tournament.get("club_name", ""),
                tournament.get("class_name", ""),
                tournament["location"],
                tournament["start_date"],
                tournament["end_date"],
                tournament["rounds_count"],
                tournament["status"],
                len(self.db.list_players(int(tournament["id"]), active_only=False)),
                len(self.db.list_rounds(int(tournament["id"]))),
            ]
            for tournament in tournaments
        ]

        return [
            (
                "Clubes e escolas",
                [
                    "ID",
                    "Nome",
                    "Tipo",
                    "Ativo",
                    "Cidade",
                    "Telefone",
                    "E-mail",
                    "Membros",
                    "Turmas ativas",
                    "Torneios",
                ],
                club_rows,
            ),
            (
                "Turmas",
                ["ID", "Clube/Escola", "Turma", "Professor", "Dia", "Horario", "Local", "Ativa", "Alunos ativos"],
                class_rows,
            ),
            (
                "Responsaveis",
                ["ID", "Nome", "Telefone", "E-mail", "Documento", "Ativo", "Membros", "Membros ativos"],
                guardian_rows,
            ),
            (
                "Aulas e presencas",
                [
                    "ID",
                    "Titulo",
                    "Tipo",
                    "Data",
                    "Inicio",
                    "Clube/Escola",
                    "Turma",
                    "Nivel",
                    "Objetivo",
                    "Conteudo",
                    "Tarefa",
                    "Instrutor",
                    "Status",
                    "Presentes",
                    "Faltas",
                    "Justificadas",
                ],
                session_rows,
            ),
            (
                "Planos financeiros",
                ["ID", "Nome", "Valor", "Ciclo", "Ativo", "Lancamentos"],
                plan_rows,
            ),
            (
                "Financeiro",
                [
                    "ID",
                    "Membro",
                    "Plano",
                    "Descricao",
                    "Referencia",
                    "Vencimento",
                    "Pagamento",
                    "Valor",
                    "Status",
                    "Metodo",
                ],
                payment_rows,
            ),
            (
                "Calendario",
                [
                    "ID",
                    "Titulo",
                    "Tipo",
                    "Data",
                    "Inicio",
                    "Clube/Escola",
                    "Torneio vinculado",
                    "Local",
                    "Status",
                ],
                event_rows,
            ),
            ("Indicadores", ["Indicador", "Valor"], indicator_rows),
            ("Membros por status", ["Status", "Quantidade"], status_rows),
            ("Membros por tipo", ["Tipo", "Quantidade"], type_rows),
            ("Membros por categoria", ["Categoria", "Quantidade"], category_rows),
            (
                "Ranking interno",
                ["Pos", "Nome", "Clube/Escola", "Turma", "Rating", "Categoria", "Tipo", "Status"],
                ranking_rows,
            ),
            (
                "Torneios",
                [
                    "ID",
                    "Nome",
                    "Escopo",
                    "Clube/Escola",
                    "Turma",
                    "Local",
                    "Data inicial",
                    "Data final",
                    "Rodadas previstas",
                    "Status",
                    "Jogadores",
                    "Rodadas geradas",
                ],
                tournament_rows,
            ),
        ]

    def _tournaments_period_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        tournaments = [
            tournament
            for tournament in self.db.list_tournaments()
            if self._date_in_period(self._tournament_report_date(tournament), start_date, end_date)
        ]
        tournaments = sorted(
            tournaments,
            key=lambda item: (
                str(item.get("start_date") or item.get("created_at") or ""),
                int(item.get("id") or 0),
            ),
        )

        period_rows: list[list[Any]] = [
            ["Data inicial", start_date or "Sem filtro"],
            ["Data final", end_date or "Sem filtro"],
            ["Torneios encontrados", len(tournaments)],
            ["Gerado em", Database.now()],
        ]
        tournament_rows: list[list[Any]] = []
        standings_rows: list[list[Any]] = []
        for tournament in tournaments:
            tournament_id = int(tournament["id"])
            players = self.db.list_players(tournament_id, active_only=False)
            rounds = self.db.list_rounds(tournament_id)
            closed_rounds = [round_data for round_data in rounds if round_data["status"] == "closed"]
            standings = self.pairing_service.standings(tournament_id)
            leader = standings[0] if standings else None
            tournament_rows.append(
                [
                    tournament_id,
                    tournament["name"],
                    self._tournament_scope_label(tournament),
                    tournament.get("club_name", ""),
                    tournament.get("class_name", ""),
                    self._tournament_report_date(tournament),
                    tournament["location"],
                    tournament["rounds_count"],
                    tournament["status"],
                    len(players),
                    len(rounds),
                    len(closed_rounds),
                    leader["name"] if leader else "",
                    leader["points"] if leader else "",
                ]
            )
            for item in standings[:10]:
                standings_rows.append(
                    [
                        tournament["name"],
                        item["position"],
                        item["name"],
                        item["points"],
                        item["performance"],
                        item["rating"],
                        item["club"],
                    ]
                )

        return [
            ("Periodo", ["Campo", "Valor"], period_rows),
            (
                "Torneios por periodo",
                [
                    "ID",
                    "Nome",
                    "Escopo",
                    "Clube/Escola",
                    "Turma",
                    "Data",
                    "Local",
                    "Rodadas previstas",
                    "Status",
                    "Jogadores",
                    "Rodadas geradas",
                    "Rodadas fechadas",
                    "Lider",
                    "Pontos lider",
                ],
                tournament_rows,
            ),
            (
                "Top 10 por torneio",
                ["Torneio", "Pos", "Jogador", "Pontos", "Performance", "Rating", "Clube"],
                standings_rows,
            ),
        ]

    def _attendance_report_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        training_service = __import__('src.services.education_service', fromlist=['TrainingService']).TrainingService(self.db)
        rows = training_service.attendance_report(start_date=start_date, end_date=end_date)
        sessions = self.db.list_training_sessions(start_date=start_date, end_date=end_date)

        summary_rows: list[list[Any]] = [
            ["Data inicial", start_date or "Sem filtro"],
            ["Data final", end_date or "Sem filtro"],
            ["Aulas/treinos no periodo", len(sessions)],
            ["Registros de chamada", len(rows)],
            ["Presencas", sum(1 for item in rows if item["status"] == "present")],
            ["Faltas", sum(1 for item in rows if item["status"] == "absent")],
            ["Faltas justificadas", sum(1 for item in rows if item["status"] == "justified")],
            ["Gerado em", Database.now()],
        ]
        session_rows = [
            [
                item["id"],
                item["title"],
                TRAINING_SESSION_TYPES.get(item["session_type"], item["session_type"]),
                item["session_date"],
                item["start_time"],
                item["end_time"],
                item["club_name"],
                item["class_name"],
                item["instructor"],
                TRAINING_SESSION_STATUSES.get(item["status"], item["status"]),
                item["present_count"],
                item["absent_count"],
                item["justified_count"],
            ]
            for item in sessions
        ]
        attendance_rows = [
            [
                item["session_date"],
                item["start_time"],
                item["title"],
                TRAINING_SESSION_TYPES.get(item["session_type"], item["session_type"]),
                item["club_name"],
                item["class_name"],
                item["member_name"],
                item["member_category"],
                ATTENDANCE_STATUSES.get(item["status"], item["status"]),
                item["notes"],
            ]
            for item in rows
        ]
        return [
            ("Resumo", ["Indicador", "Valor"], summary_rows),
            (
                "Aulas e treinos",
                [
                    "ID",
                    "Titulo",
                    "Tipo",
                    "Data",
                    "Inicio",
                    "Fim",
                    "Clube/Escola",
                    "Turma",
                    "Instrutor",
                    "Status",
                    "Presentes",
                    "Faltas",
                    "Justificadas",
                ],
                session_rows,
            ),
            (
                "Chamada",
                [
                    "Data",
                    "Inicio",
                    "Aula/Treino",
                    "Tipo",
                    "Clube/Escola",
                    "Turma",
                    "Membro",
                    "Categoria",
                    "Presenca",
                    "Observacoes",
                ],
                attendance_rows,
            ),
        ]

    def _financial_report_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        finance_service = __import__('src.services.finance_service', fromlist=['FinanceService']).FinanceService(self.db)
        summary = finance_service.finance_summary(start_date=start_date, end_date=end_date)
        payments = finance_service.payments_report(start_date=start_date, end_date=end_date)
        plans = self.db.list_membership_plans(active_only=False)

        summary_rows: list[list[Any]] = [
            ["Data inicial", start_date or "Sem filtro"],
            ["Data final", end_date or "Sem filtro"],
            ["Lancamentos", summary["total_payments"]],
            ["Pagos", summary["paid"]],
            ["Pendentes", summary["pending"]],
            ["Atrasados", summary["late"]],
            ["Isentos", summary["exempt"]],
            ["Valor recebido", summary["paid_amount"]],
            ["Valor pendente", summary["pending_amount"]],
            ["Valor atrasado", summary["late_amount"]],
            ["Gerado em", Database.now()],
        ]
        plan_rows: list[list[Any]] = [
            [
                plan["id"],
                plan["name"],
                plan["amount"],
                BILLING_CYCLES.get(plan["billing_cycle"], plan["billing_cycle"]),
                "Sim" if plan.get("active") else "Nao",
                plan["payments_count"],
                plan["notes"],
            ]
            for plan in plans
        ]
        payment_rows: list[list[Any]] = [
            [
                payment["id"],
                payment["member_name"],
                payment.get("club_name") or "",
                payment.get("active_class_name") or "",
                payment.get("plan_name") or "",
                payment["description"],
                payment["reference_period"],
                payment["due_date"],
                payment["payment_date"],
                payment["amount"],
                PAYMENT_STATUSES.get(payment["effective_status"], payment["effective_status"]),
                payment["method"],
                payment["notes"],
            ]
            for payment in payments
        ]
        member_rows: list[list[Any]] = []
        for member in self.db.list_members(active_only=False):
            status = finance_service.member_financial_status(int(member["id"]))
            member_rows.append(
                [
                    member["id"],
                    member["name"],
                    member.get("club_name") or "",
                    member.get("active_class_name") or "",
                    status["label"],
                    status["open_amount"],
                    status["late_amount"],
                ]
            )
        return [
            ("Resumo", ["Indicador", "Valor"], summary_rows),
            (
                "Planos",
                ["ID", "Nome", "Valor", "Ciclo", "Ativo", "Lancamentos", "Observacoes"],
                plan_rows,
            ),
            (
                "Lancamentos",
                [
                    "ID",
                    "Membro",
                    "Clube/Escola",
                    "Turma",
                    "Plano",
                    "Descricao",
                    "Referencia",
                    "Vencimento",
                    "Pagamento",
                    "Valor",
                    "Status",
                    "Metodo",
                    "Observacoes",
                ],
                payment_rows,
            ),
            (
                "Status por membro",
                ["ID", "Membro", "Clube/Escola", "Turma", "Status", "Valor aberto", "Valor atrasado"],
                member_rows,
            ),
        ]

    def _events_report_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        event_service = __import__('src.services.event_service', fromlist=['EventService']).EventService(self.db)
        events = event_service.events_report(start_date=start_date, end_date=end_date)
        summary_rows: list[list[Any]] = [
            ["Data inicial", start_date or "Sem filtro"],
            ["Data final", end_date or "Sem filtro"],
            ["Eventos encontrados", len(events)],
            ["Planejados", sum(1 for event in events if event["status"] == "planned")],
            ["Confirmados", sum(1 for event in events if event["status"] == "confirmed")],
            ["Concluidos", sum(1 for event in events if event["status"] == "done")],
            ["Cancelados", sum(1 for event in events if event["status"] == "canceled")],
            ["Gerado em", Database.now()],
        ]
        event_rows: list[list[Any]] = [
            [
                event["id"],
                event["title"],
                EVENT_TYPES.get(event["event_type"], event["event_type"]),
                event["event_date"],
                event["start_time"],
                event["end_time"],
                event.get("club_name") or "",
                event.get("tournament_name") or "",
                event["location"],
                EVENT_STATUSES.get(event["status"], event["status"]),
                event["notes"],
            ]
            for event in events
        ]
        return [
            ("Resumo", ["Indicador", "Valor"], summary_rows),
            (
                "Eventos",
                [
                    "ID",
                    "Titulo",
                    "Tipo",
                    "Data",
                    "Inicio",
                    "Fim",
                    "Clube/Escola",
                    "Torneio vinculado",
                    "Local",
                    "Status",
                    "Observacoes",
                ],
                event_rows,
            ),
        ]

    def _internal_ranking_sections(
        self,
        category: str = "",
        club_id: int | None = None,
        class_id: int | None = None,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        ranking_service = __import__('src.services.rating_service', fromlist=['InternalRatingService']).InternalRatingService(self.db)
        ranking = ranking_service.ranking(
            category=category.strip(),
            active_only=True,
            club_id=club_id,
            class_id=class_id,
            start_date=start_date.strip(),
            end_date=end_date.strip(),
        )
        club = self.db.get_club(club_id) if club_id else None
        class_data = self.db.get_class(class_id) if class_id else None
        summary_rows: list[list[Any]] = [
            ["Categoria", category.strip() or "Todas"],
            ["Clube/Escola", club.get("name", "") if club else "Todos"],
            ["Turma", class_data.get("name", "") if class_data else "Todas"],
            ["Temporada inicial", start_date.strip() or "Sem filtro"],
            ["Temporada final", end_date.strip() or "Sem filtro"],
            ["Membros no ranking", len(ranking)],
            ["Gerado em", Database.now()],
        ]
        ranking_rows: list[list[Any]] = [
            [
                item["position"],
                item["name"],
                item["club_name"],
                item["class_name"],
                item["category"],
                item["age_category"],
                item["rating_category"],
                item["prize_tags"],
                item["rating"],
                item["last_delta"],
                item["games"],
                item["wins"],
                item["draws"],
                item["losses"],
                item["points"],
                f"{item['score_rate']}%",
                item["last_performance"],
                item["last_tournament"] or "",
            ]
            for item in ranking
        ]
        category_rows: list[list[Any]] = []
        for category_name, category_ranking in ranking_service.category_rankings(
            active_only=True,
            club_id=club_id,
            class_id=class_id,
            start_date=start_date.strip(),
            end_date=end_date.strip(),
        ).items():
            if not category_ranking:
                continue
            leader = category_ranking[0]
            category_rows.append(
                [
                    category_name,
                    len(category_ranking),
                    leader["name"],
                    leader["rating"],
                    f"{leader['score_rate']}%",
                ]
            )
        return [
            ("Resumo", ["Indicador", "Valor"], summary_rows),
            (
                "Ranking interno",
                [
                    "Pos",
                    "Membro",
                    "Clube/Escola",
                    "Turma",
                    "Categoria",
                    "Categoria idade",
                    "Categoria rating",
                    "Tags premiacao",
                    "Rating",
                    "Variacao",
                    "Partidas",
                    "Vitorias",
                    "Empates",
                    "Derrotas",
                    "Pontos",
                    "Aproveitamento",
                    "Ultima performance",
                    "Ultimo torneio",
                ],
                ranking_rows,
            ),
            (
                "Categorias",
                ["Categoria", "Membros", "Lider", "Rating lider", "Aproveitamento lider"],
                category_rows,
            ),
        ]

    def _administrative_package_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        overview = __import__('src.services.dashboard_service', fromlist=['DashboardService']).DashboardService(self.db).overview()
        finance = overview["finance_summary"]
        cover_rows = [
            ["Gerado em", Database.now()],
            ["Periodo inicial", start_date or "Sem filtro"],
            ["Periodo final", end_date or "Sem filtro"],
            ["Membros ativos", overview["summary"].get("active_members", 0)],
            ["Torneios cadastrados", overview["summary"].get("total_tournaments", 0)],
            ["Proximos eventos", len(overview["upcoming_events"])],
            ["Valor recebido", finance.get("paid_amount", 0.0)],
            ["Valor pendente", finance.get("pending_amount", 0.0)],
            ["Valor atrasado", finance.get("late_amount", 0.0)],
        ]
        sections: list[tuple[str, list[str], list[list[Any]]]] = [
            ("Pacote administrativo", ["Campo", "Valor"], cover_rows)
        ]
        sections.extend(self._club_report_sections())
        sections.extend(self._attendance_report_sections(start_date, end_date))
        sections.extend(self._financial_report_sections(start_date, end_date))
        sections.extend(self._events_report_sections(start_date, end_date))
        sections.extend(self._tournaments_period_sections(start_date, end_date))
        sections.extend(self._internal_ranking_sections())
        return sections

    def _member_evolution_sections(
        self,
        member_id: int,
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        member = self.db.get_member(member_id)
        if not member:
            raise AppError("Membro nao encontrado.")

        member_service = __import__('src.services.member_service', fromlist=['MemberService']).MemberService(self.db)
        tournament_history = member_service.tournament_history(member_id)
        rating_history = self.db.list_member_rating_history(member_id)
        ordered_tournaments = sorted(
            tournament_history,
            key=lambda item: (
                str(item.get("start_date") or ""),
                int(item.get("tournament_id") or 0),
            ),
        )
        ordered_rating_history = sorted(
            rating_history,
            key=lambda item: (
                str(item.get("created_at") or ""),
                int(item.get("id") or 0),
            ),
        )
        rating_by_registration = {
            (int(item["tournament_id"] or 0), int(item["player_id"] or 0)): item
            for item in rating_history
        }

        member_rows: list[list[Any]] = [
            ["Nome", member["name"]],
            ["Clube/Escola", member.get("club_name", "")],
            ["Turma atual", member.get("active_class_name", "")],
            ["Tipo", member["member_type"]],
            ["Status", member["status"]],
            ["Rating interno atual", member["rating"]],
            ["Categoria", member["category"]],
            ["Categoria idade", member.get("age_category", "")],
            ["Categoria rating", member.get("rating_category", "")],
            ["Tags premiacao", member.get("prize_tags", "")],
            ["Cidade", member["city"]],
            ["Telefone", member["phone"]],
            ["E-mail", member["email"]],
            ["Nascimento", member["birth_date"]],
        ]
        guardian_rows: list[list[Any]] = [
            [
                item["guardian_name"],
                item["relationship"],
                "Sim" if item.get("primary_contact") else "Nao",
                "Sim" if item.get("emergency_contact") else "Nao",
                item["guardian_phone"],
                item["guardian_email"],
                item["guardian_document"],
                "Sim" if item.get("guardian_active") else "Nao",
                item["notes"],
            ]
            for item in self.db.list_member_guardians(member_id)
        ]
        attendance_summary = self.db.member_attendance_summary(member_id)
        attendance_summary_rows: list[list[Any]] = [
            ["Registros de chamada", attendance_summary["total"]],
            ["Presencas", attendance_summary["present"]],
            ["Faltas", attendance_summary["absent"]],
            ["Faltas justificadas", attendance_summary["justified"]],
            ["Frequencia", f"{attendance_summary['attendance_rate']}%"],
        ]
        attendance_rows: list[list[Any]] = [
            [
                item["session_date"],
                item["start_time"],
                item["title"],
                TRAINING_SESSION_TYPES.get(item["session_type"], item["session_type"]),
                item["club_name"],
                item["class_name"],
                item["instructor"],
                ATTENDANCE_STATUSES.get(item["status"], item["status"]),
                item["notes"],
            ]
            for item in self.db.list_attendance_report(member_id=member_id)
        ]
        finance_status = __import__('src.services.finance_service', fromlist=['FinanceService']).FinanceService(self.db).member_financial_status(member_id)
        payment_rows: list[list[Any]] = [
            [
                item["description"],
                item.get("plan_name") or "",
                item["reference_period"],
                item["due_date"],
                item["payment_date"],
                item["amount"],
                PAYMENT_STATUSES.get(item["effective_status"], item["effective_status"]),
                item["method"],
                item["notes"],
            ]
            for item in __import__('src.services.finance_service', fromlist=['FinanceService']).FinanceService(self.db).payments_report(member_id=member_id)
        ]
        finance_rows: list[list[Any]] = [
            ["Status financeiro", finance_status["label"]],
            ["Valor aberto", finance_status["open_amount"]],
            ["Valor atrasado", finance_status["late_amount"]],
        ]

        tournament_rows: list[list[Any]] = []
        for item in ordered_tournaments:
            rating_update = rating_by_registration.get(
                (int(item["tournament_id"]), int(item["player_id"]))
            )
            old_rating = rating_update["old_rating"] if rating_update else ""
            new_rating = rating_update["new_rating"] if rating_update else ""
            delta = ""
            if rating_update:
                delta_value = int(rating_update["new_rating"] or 0) - int(rating_update["old_rating"] or 0)
                delta = f"+{delta_value}" if delta_value >= 0 else str(delta_value)
            tournament_rows.append(
                [
                    item["name"],
                    item["start_date"],
                    item["status"],
                    item["player_rating"],
                    item["points"],
                    item["position"],
                    item["performance"],
                    old_rating,
                    new_rating,
                    delta,
                    item["rounds_played"],
                    item["wins"],
                    item["draws"],
                    item["losses"],
                    item["byes"],
                    item["last_result"],
                ]
            )

        rating_rows: list[list[Any]] = []
        for item in ordered_rating_history:
            old_rating = int(item["old_rating"] or 0)
            new_rating = int(item["new_rating"] or 0)
            rating_delta = new_rating - old_rating
            rating_rows.append(
                [
                    item["created_at"],
                    item.get("tournament_name") or "",
                    old_rating,
                    new_rating,
                    f"+{rating_delta}" if rating_delta >= 0 else str(rating_delta),
                    item["performance"],
                    item["games"],
                    item["points"],
                ]
            )

        result_rows: list[list[Any]] = []
        for tournament_item in ordered_tournaments:
            for result in member_service.tournament_results(member_id, int(tournament_item["tournament_id"])):
                result_rows.append(
                    [
                        tournament_item["name"],
                        result["round_number"],
                        result["round_status"],
                        result["board_number"],
                        result["color"],
                        result["opponent"],
                        result["result"],
                        result["outcome"],
                        "" if result["points"] is None else result["points"],
                    ]
                )

        return [
            ("Aluno", ["Campo", "Valor"], member_rows),
            (
                "Responsaveis",
                [
                    "Nome",
                    "Parentesco",
                    "Principal",
                    "Emergencia",
                    "Telefone",
                    "E-mail",
                    "Documento",
                    "Ativo",
                    "Observacoes",
                ],
                guardian_rows,
            ),
            ("Resumo de frequencia", ["Indicador", "Valor"], attendance_summary_rows),
            (
                "Presencas",
                [
                    "Data",
                    "Inicio",
                    "Aula/Treino",
                    "Tipo",
                    "Clube/Escola",
                    "Turma",
                    "Instrutor",
                    "Presenca",
                    "Observacoes",
                ],
                attendance_rows,
            ),
            ("Resumo financeiro", ["Indicador", "Valor"], finance_rows),
            (
                "Financeiro",
                [
                    "Descricao",
                    "Plano",
                    "Referencia",
                    "Vencimento",
                    "Pagamento",
                    "Valor",
                    "Status",
                    "Metodo",
                    "Observacoes",
                ],
                payment_rows,
            ),
            (
                "Evolucao por torneio",
                [
                    "Torneio",
                    "Data",
                    "Status",
                    "Rating inscricao",
                    "Pontos",
                    "Posicao",
                    "Performance",
                    "Rating anterior",
                    "Rating novo",
                    "Variacao",
                    "Rodadas",
                    "Vitorias",
                    "Empates",
                    "Derrotas",
                    "Byes",
                    "Ultimo resultado",
                ],
                tournament_rows,
            ),
            (
                "Rating interno",
                [
                    "Data",
                    "Torneio",
                    "Rating anterior",
                    "Rating novo",
                    "Variacao",
                    "Performance",
                    "Partidas",
                    "Pontos",
                ],
                rating_rows,
            ),
            (
                "Resultados",
                [
                    "Torneio",
                    "Rodada",
                    "Status rodada",
                    "Mesa",
                    "Cor",
                    "Adversario",
                    "Resultado",
                    "Desfecho",
                    "Pontos",
                ],
                result_rows,
            ),
        ]

    @staticmethod
    def _count_member_field(
        members: list[dict[str, Any]],
        field: str,
    ) -> list[list[Any]]:
        totals: dict[str, int] = {}
        for member in members:
            value = str(member.get(field) or "").strip() or "Sem informacao"
            totals[value] = totals.get(value, 0) + 1
        return [[key, totals[key]] for key in sorted(totals, key=lambda item: item.casefold())]

    @staticmethod
    def _tournament_report_date(tournament: dict[str, Any]) -> str:
        return str(tournament.get("start_date") or tournament.get("created_at") or "")[:10]

    @staticmethod
    def _date_in_period(value: str, start_date: str, end_date: str) -> bool:
        if start_date and value and value < start_date:
            return False
        if end_date and value and value > end_date:
            return False
        if (start_date or end_date) and not value:
            return False
        return True

    def _pairings_section(self, round_id: int) -> tuple[str, list[str], list[list[Any]]]:
        round_data = self.db.get_round(round_id)
        if not round_data:
            raise AppError("Rodada nao encontrada.")
        tournament = self.db.get_tournament(int(round_data["tournament_id"]))
        if tournament and tournament.get("competition_type") == "team":
            return self._team_pairings_section(round_data)
        pairings = self.db.get_pairings_for_round(round_id)
        rows = []
        for pairing in pairings:
            result_url = ""
            if not pairing["is_bye"] and round_data.get("status") != "closed":
                result_url = self._qr_result_url(int(round_data["tournament_id"]), int(pairing["id"]))
            rows.append(
                [
                    pairing["board_number"],
                    pairing_player_name(pairing, "white"),
                    pairing["white_rating"],
                    pairing["result"],
                    "BYE" if pairing["is_bye"] else pairing_player_name(pairing, "black"),
                    "" if pairing["is_bye"] else pairing["black_rating"],
                    result_url,
                ]
            )
        return (
            f"Rodada {round_data['number']}",
            ["Mesa", "Brancas", "Rating", "Resultado", "Pretas", "Rating", "Link resultado QR"],
            rows,
        )

    def _scoresheet_rows(
        self,
        round_data: dict[str, Any],
        tournament: dict[str, Any],
    ) -> list[dict[str, Any]]:
        if tournament.get("competition_type") != "team":
            return [
                {
                    "board_label": str(pairing["board_number"]),
                    "context": "",
                    "white_name": pairing_player_name(pairing, "white"),
                    "white_rating": pairing.get("white_rating", ""),
                    "white_club": pairing.get("white_club", ""),
                    "white_ids": self._official_ids_label(pairing, "white"),
                    "black_name": pairing_player_name(pairing, "black"),
                    "black_rating": pairing.get("black_rating", ""),
                    "black_club": pairing.get("black_club", ""),
                    "black_ids": self._official_ids_label(pairing, "black"),
                }
                for pairing in self.db.get_pairings_for_round(int(round_data["id"]))
                if not pairing.get("is_bye")
            ]

        scoresheets: list[dict[str, Any]] = []
        for match in self.db.list_team_matches_for_round(int(round_data["id"])):
            if match.get("is_bye"):
                continue
            for board in self.db.list_team_boards(int(match["id"])):
                scoresheets.append(
                    {
                        "board_label": f"{match['match_number']}.{board['board_number']}",
                        "context": f"{match.get('white_team_name', '')} x {match.get('black_team_name', '')}",
                        "white_name": self._team_board_player_name(board, "white"),
                        "white_rating": board.get("white_player_rating", ""),
                        "white_club": board.get("white_player_club", ""),
                        "white_ids": self._team_board_official_ids_label(board, "white"),
                        "black_name": self._team_board_player_name(board, "black"),
                        "black_rating": board.get("black_player_rating", ""),
                        "black_club": board.get("black_player_club", ""),
                        "black_ids": self._team_board_official_ids_label(board, "black"),
                    }
                )
        return scoresheets

    @staticmethod
    def _official_ids_label(player: Mapping[str, Any], color: str) -> str:
        values = [
            ("FIDE", player.get(f"{color}_fide_id")),
            ("CBX", player.get(f"{color}_cbx_id")),
            ("LBX", player.get(f"{color}_lbx_id")),
        ]
        return " | ".join(f"{label}: {value}" for label, value in values if value)

    @staticmethod
    def _team_board_official_ids_label(board: Mapping[str, Any], color: str) -> str:
        values = [
            ("FIDE", board.get(f"{color}_player_fide_id")),
            ("CBX", board.get(f"{color}_player_cbx_id")),
            ("LBX", board.get(f"{color}_player_lbx_id")),
        ]
        return " | ".join(f"{label}: {value}" for label, value in values if value)

    @staticmethod
    def _write_scoresheets_pdf(
        path: Path,
        tournament: Mapping[str, Any],
        round_data: Mapping[str, Any],
        scoresheets: list[dict[str, Any]],
    ) -> None:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.pdfbase.pdfmetrics import stringWidth
            from reportlab.pdfgen import canvas
        except ImportError as exc:
            raise AppError("Instale reportlab para exportar PDF.") from exc

        path.parent.mkdir(parents=True, exist_ok=True)
        page_width, page_height = A4
        document = canvas.Canvas(str(path), pagesize=A4)
        document.setTitle(f"Sumulas - {tournament.get('name', '')} - Rodada {round_data.get('number', '')}")

        def fitted_text(text: Any, max_width: float, font_name: str = "Helvetica", font_size: int = 9) -> str:
            value = str(text or "")
            if stringWidth(value, font_name, font_size) <= max_width:
                return value
            while value and stringWidth(f"{value}...", font_name, font_size) > max_width:
                value = value[:-1]
            return f"{value}..."

        margin = 14 * mm
        content_width = page_width - 2 * margin
        row_height = 6.2 * mm
        grid_top = page_height - 83 * mm
        column_widths = [9 * mm, 41 * mm, 41 * mm] * 2

        for scoresheet in scoresheets:
            document.setFont("Helvetica-Bold", 14)
            document.drawString(margin, page_height - 17 * mm, fitted_text(tournament.get("name", ""), content_width, "Helvetica-Bold", 14))
            document.setFont("Helvetica", 9)
            document.drawString(margin, page_height - 23 * mm, f"Rodada: {round_data.get('number', '')}")
            document.drawRightString(page_width - margin, page_height - 23 * mm, f"Mesa: {scoresheet['board_label']}")
            if scoresheet.get("context"):
                document.drawString(margin, page_height - 28 * mm, fitted_text(scoresheet["context"], content_width, font_size=9))

            player_top = page_height - 37 * mm
            for color_label, prefix in (("Brancas", "white"), ("Pretas", "black")):
                document.setFont("Helvetica-Bold", 10)
                document.drawString(margin, player_top, f"{color_label}: {fitted_text(scoresheet[f'{prefix}_name'], 116 * mm, 'Helvetica-Bold', 10)}")
                document.setFont("Helvetica", 8)
                details = f"Rating: {scoresheet[f'{prefix}_rating'] or '-'}   Clube: {scoresheet[f'{prefix}_club'] or '-'}"
                document.drawString(margin + 18 * mm, player_top - 4 * mm, fitted_text(details, content_width - 18 * mm, font_size=8))
                document.drawString(margin + 18 * mm, player_top - 8 * mm, fitted_text(scoresheet[f"{prefix}_ids"], content_width - 18 * mm, font_size=8))
                player_top -= 14 * mm

            headers = ["N.", "Brancas", "Pretas", "N.", "Brancas", "Pretas"]
            x_positions = [margin]
            for width in column_widths:
                x_positions.append(x_positions[-1] + width)
            document.setFillGray(0.92)
            document.rect(margin, grid_top, content_width, row_height, fill=1, stroke=0)
            document.setFillGray(0)
            document.setFont("Helvetica-Bold", 8)
            for index, header in enumerate(headers):
                document.drawCentredString((x_positions[index] + x_positions[index + 1]) / 2, grid_top + 2.1 * mm, header)

            document.setLineWidth(0.35)
            for row in range(31):
                y = grid_top - row * row_height
                document.line(margin, y, margin + content_width, y)
            for x in x_positions:
                document.line(x, grid_top + row_height, x, grid_top - 30 * row_height)
            document.setFont("Helvetica", 8)
            for row in range(30):
                y = grid_top - (row + 1) * row_height + 2.1 * mm
                document.drawCentredString((x_positions[0] + x_positions[1]) / 2, y, str(row + 1))
                document.drawCentredString((x_positions[3] + x_positions[4]) / 2, y, str(row + 31))

            footer_y = grid_top - 30 * row_height - 8 * mm
            document.setFont("Helvetica-Bold", 9)
            document.drawString(margin, footer_y, "Resultado:  1-0  [  ]    1/2-1/2  [  ]    0-1  [  ]")
            document.setFont("Helvetica", 8)
            document.line(margin, footer_y - 12 * mm, margin + 76 * mm, footer_y - 12 * mm)
            document.line(page_width - margin - 76 * mm, footer_y - 12 * mm, page_width - margin, footer_y - 12 * mm)
            document.drawCentredString(margin + 38 * mm, footer_y - 16 * mm, "Assinatura das brancas")
            document.drawCentredString(page_width - margin - 38 * mm, footer_y - 16 * mm, "Assinatura das pretas")
            document.showPage()

        document.save()

    @staticmethod
    def _write_initial_player_list_pdf(
        path: Path,
        title: str,
        headers: list[str],
        rows: list[list[Any]],
    ) -> None:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.pdfbase.pdfmetrics import stringWidth
            from reportlab.pdfgen import canvas
        except ImportError as exc:
            raise AppError("Instale reportlab para exportar PDF.") from exc

        path.parent.mkdir(parents=True, exist_ok=True)
        page_width, page_height = A4
        document = canvas.Canvas(str(path), pagesize=A4)
        document.setTitle(title)
        margin = 11 * mm
        content_width = page_width - 2 * margin
        column_widths = [14 * mm, 52 * mm, 16 * mm, 28 * mm, 19 * mm, 18 * mm, 41 * mm]
        row_height = 6.2 * mm
        table_top = page_height - 39 * mm
        rows_per_page = 38
        pdf_headers = [*headers[:-1], "Assinatura"]
        total_pages = max(1, math.ceil(len(rows) / rows_per_page))

        def fitted_text(text: Any, max_width: float, font_name: str = "Helvetica", font_size: int = 7) -> str:
            value = str(text or "")
            if stringWidth(value, font_name, font_size) <= max_width:
                return value
            while value and stringWidth(f"{value}...", font_name, font_size) > max_width:
                value = value[:-1]
            return f"{value}..."

        for page_index in range(total_pages):
            document.setFont("Helvetica-Bold", 13)
            document.drawString(margin, page_height - 15 * mm, fitted_text(title, content_width, "Helvetica-Bold", 13))
            document.setFont("Helvetica", 8)
            document.drawString(margin, page_height - 21 * mm, "Lista de chamada por ranking inicial")
            document.drawRightString(
                page_width - margin,
                page_height - 21 * mm,
                f"Pagina {page_index + 1}/{total_pages}",
            )
            document.drawString(margin, page_height - 27 * mm, f"Jogadores inscritos: {len(rows)}")

            x_positions = [margin]
            for width in column_widths:
                x_positions.append(x_positions[-1] + width)
            document.setFillGray(0.92)
            document.rect(margin, table_top, content_width, row_height, fill=1, stroke=0)
            document.setFillGray(0)
            document.setFont("Helvetica-Bold", 7)
            for index, header in enumerate(pdf_headers):
                document.drawCentredString(
                    (x_positions[index] + x_positions[index + 1]) / 2,
                    table_top + 2.1 * mm,
                    header,
                )

            page_rows = rows[page_index * rows_per_page : (page_index + 1) * rows_per_page]
            document.setLineWidth(0.35)
            for row_index, row in enumerate(page_rows, start=1):
                bottom = table_top - row_index * row_height
                document.line(margin, bottom, margin + content_width, bottom)
                document.setFont("Helvetica", 7)
                for column_index, value in enumerate(row[:-1]):
                    left = x_positions[column_index]
                    width = column_widths[column_index]
                    text = fitted_text(value, width - 2 * mm)
                    if column_index in {0, 2}:
                        document.drawCentredString(left + width / 2, bottom + 2.1 * mm, text)
                    else:
                        document.drawString(left + 1 * mm, bottom + 2.1 * mm, text)
            for x_position in x_positions:
                document.line(x_position, table_top + row_height, x_position, table_top - len(page_rows) * row_height)

            document.setFont("Helvetica", 7)
            document.drawRightString(page_width - margin, 8 * mm, "Albericus - Lista de chamada")
            document.showPage()

        document.save()

    def _write_pairings_wall_pdf(
        self,
        path: Path,
        tournament: Mapping[str, Any],
        round_data: Mapping[str, Any],
    ) -> None:
        try:
            import qrcode
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.lib.utils import ImageReader
            from reportlab.pdfbase.pdfmetrics import stringWidth
            from reportlab.pdfgen import canvas
        except ImportError as exc:
            raise AppError("Instale reportlab e qrcode para exportar o mural em PDF.") from exc

        pairings = self.db.get_pairings_for_round(int(round_data["id"]))
        path.parent.mkdir(parents=True, exist_ok=True)
        page_width, page_height = A4
        document = canvas.Canvas(str(path), pagesize=A4)
        document.setTitle(f"Mural - {tournament.get('name', '')} - Rodada {round_data.get('number', '')}")
        margin = 10 * mm
        content_width = page_width - 2 * margin
        column_widths = [12 * mm, 58 * mm, 16 * mm, 58 * mm, 16 * mm, 30 * mm]
        row_height = 16 * mm
        header_height = 7 * mm
        table_top = page_height - 38 * mm
        rows_per_page = 15
        is_open = round_data.get("status") != "closed"
        total_pages = max(1, math.ceil(len(pairings) / rows_per_page))
        qr_urls = self._qr_result_urls(
            int(round_data["tournament_id"]),
            [
                {**pairing, "tournament_id": int(round_data["tournament_id"])}
                for pairing in pairings
                if is_open and not pairing["is_bye"]
            ],
        )

        def fitted_text(text: Any, max_width: float, font_name: str = "Helvetica", font_size: int = 9) -> str:
            value = str(text or "")
            if stringWidth(value, font_name, font_size) <= max_width:
                return value
            while value and stringWidth(f"{value}...", font_name, font_size) > max_width:
                value = value[:-1]
            return f"{value}..."

        for page_index in range(total_pages):
            document.setFont("Helvetica-Bold", 14)
            document.drawString(
                margin,
                page_height - 14 * mm,
                fitted_text(tournament.get("name", ""), content_width - 42 * mm, "Helvetica-Bold", 14),
            )
            document.setFont("Helvetica-Bold", 12)
            document.drawRightString(page_width - margin, page_height - 14 * mm, f"Rodada {round_data.get('number', '')}")
            document.setFont("Helvetica", 8)
            status_label = "Rodada aberta - QR para envio de resultado" if is_open else "Rodada fechada - QR desativado"
            document.drawString(margin, page_height - 21 * mm, status_label)
            document.drawRightString(
                page_width - margin,
                page_height - 21 * mm,
                f"Pagina {page_index + 1}/{total_pages}",
            )
            document.drawString(margin, page_height - 27 * mm, f"Mesas: {len(pairings)}")

            x_positions = [margin]
            for width in column_widths:
                x_positions.append(x_positions[-1] + width)
            document.setFillGray(0.92)
            document.rect(margin, table_top, content_width, header_height, fill=1, stroke=0)
            document.setFillGray(0)
            document.setFont("Helvetica-Bold", 8)
            for index, header in enumerate(["Mesa", "Brancas", "Rating", "Pretas", "Rating", "QR resultado"]):
                document.drawCentredString(
                    (x_positions[index] + x_positions[index + 1]) / 2,
                    table_top + 2.3 * mm,
                    header,
                )

            page_rows = pairings[page_index * rows_per_page : (page_index + 1) * rows_per_page]
            document.setLineWidth(0.35)
            for row_index, pairing in enumerate(page_rows, start=1):
                bottom = table_top - row_index * row_height
                middle = bottom + row_height / 2
                document.line(margin, bottom, margin + content_width, bottom)
                document.setFont("Helvetica-Bold", 11)
                document.drawCentredString(x_positions[0] + column_widths[0] / 2, middle - 1.5 * mm, str(pairing["board_number"]))
                document.setFont("Helvetica-Bold", 9)
                document.drawString(
                    x_positions[1] + 1.5 * mm,
                    middle - 1.5 * mm,
                    fitted_text(pairing_player_name(pairing, "white"), column_widths[1] - 3 * mm, "Helvetica-Bold", 9),
                )
                black_name = "BYE" if pairing["is_bye"] else pairing_player_name(pairing, "black")
                document.drawString(
                    x_positions[3] + 1.5 * mm,
                    middle - 1.5 * mm,
                    fitted_text(black_name, column_widths[3] - 3 * mm, "Helvetica-Bold", 9),
                )
                document.setFont("Helvetica", 9)
                document.drawCentredString(x_positions[2] + column_widths[2] / 2, middle - 1.5 * mm, str(pairing["white_rating"]))
                black_rating = "" if pairing["is_bye"] else pairing["black_rating"]
                document.drawCentredString(x_positions[4] + column_widths[4] / 2, middle - 1.5 * mm, str(black_rating))
                if is_open and not pairing["is_bye"]:
                    url = qr_urls[int(pairing["id"])]
                    qr_code = qrcode.QRCode(
                        error_correction=qrcode.constants.ERROR_CORRECT_M,
                        box_size=2,
                        border=1,
                    )
                    qr_code.add_data(url)
                    qr_code.make(fit=True)
                    image = qr_code.make_image(fill_color="black", back_color="white")
                    buffer = io.BytesIO()
                    image.save(buffer, format="PNG")
                    buffer.seek(0)
                    qr_size = 13 * mm
                    document.drawImage(
                        ImageReader(buffer),
                        x_positions[5] + (column_widths[5] - qr_size) / 2,
                        bottom + (row_height - qr_size) / 2,
                        width=qr_size,
                        height=qr_size,
                    )
            for x_position in x_positions:
                document.line(x_position, table_top + header_height, x_position, table_top - len(page_rows) * row_height)

            document.setFont("Helvetica", 7)
            document.drawRightString(page_width - margin, 7 * mm, "Albericus - Emparceiramento para mural")
            document.showPage()

        document.save()

    def _write_table_cards_pdf(
        self,
        path: Path,
        start_board: int,
        end_board: int,
        tournament: Mapping[str, Any] | None,
        round_data: Mapping[str, Any] | None,
        include_qr: bool,
    ) -> None:
        try:
            import qrcode
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.lib.utils import ImageReader
            from reportlab.pdfgen import canvas
        except ImportError as exc:
            raise AppError("Instale reportlab e qrcode para exportar cartoes de mesa.") from exc

        path.parent.mkdir(parents=True, exist_ok=True)
        pairings_by_board = {
            int(pairing["board_number"]): pairing
            for pairing in self.db.get_pairings_for_round(int(round_data["id"]))
        } if round_data else {}
        document = canvas.Canvas(str(path), pagesize=A4)
        document.setTitle("Cartoes de mesa")
        page_width, page_height = A4
        margin = 10 * mm
        gap = 6 * mm
        card_width = (page_width - 2 * margin - gap) / 2
        card_height = (page_height - 2 * margin - gap) / 2
        cards_per_page = 4
        boards = list(range(start_board, end_board + 1))
        is_open = bool(round_data and round_data.get("status") != "closed")
        qr_urls = self._qr_result_urls(
            int(round_data["tournament_id"]),
            [
                {**pairing, "tournament_id": int(round_data["tournament_id"])}
                for pairing in pairings_by_board.values()
                if include_qr and is_open and not pairing.get("is_bye")
            ],
        ) if round_data else {}

        for card_index, board_number in enumerate(boards):
            page_slot = card_index % cards_per_page
            if card_index and page_slot == 0:
                document.showPage()
            column = page_slot % 2
            row = page_slot // 2
            left = margin + column * (card_width + gap)
            bottom = page_height - margin - card_height - row * (card_height + gap)
            center_x = left + card_width / 2

            document.setLineWidth(1.2)
            document.roundRect(left, bottom, card_width, card_height, 4 * mm, stroke=1, fill=0)
            document.setFont("Helvetica-Bold", 14)
            document.drawCentredString(center_x, bottom + card_height - 17 * mm, "MESA")
            document.setFont("Helvetica-Bold", 68)
            document.drawCentredString(center_x, bottom + card_height - 53 * mm, str(board_number))

            details_y = bottom + 19 * mm
            if tournament:
                document.setFont("Helvetica-Bold", 8)
                document.drawCentredString(center_x, details_y, str(tournament.get("name") or "")[:48])
                details_y -= 5 * mm
            if round_data:
                document.setFont("Helvetica", 9)
                document.drawCentredString(center_x, details_y, f"Rodada {round_data.get('number', '')}")

            pairing = pairings_by_board.get(board_number)
            if include_qr and is_open and pairing and not pairing.get("is_bye"):
                url = qr_urls[int(pairing["id"])]
                qr_code = qrcode.QRCode(
                    error_correction=qrcode.constants.ERROR_CORRECT_M,
                    box_size=3,
                    border=1,
                )
                qr_code.add_data(url)
                qr_code.make(fit=True)
                image = qr_code.make_image(fill_color="black", back_color="white")
                buffer = io.BytesIO()
                image.save(buffer, format="PNG")
                buffer.seek(0)
                qr_size = 27 * mm
                document.drawImage(
                    ImageReader(buffer),
                    center_x - qr_size / 2,
                    bottom + 31 * mm,
                    width=qr_size,
                    height=qr_size,
                )
                document.setFont("Helvetica", 7)
                document.drawCentredString(center_x, bottom + 27 * mm, "QR para enviar resultado")
            elif include_qr:
                document.setFont("Helvetica", 7)
                document.drawCentredString(center_x, bottom + 30 * mm, "QR indisponivel para esta mesa")

        if boards:
            document.showPage()
        document.save()

    def _qr_result_url(self, tournament_id: int, pairing_id: int) -> str:
        from src.services.qr_result_service import QRResultService

        settings = self.db.get_app_settings()
        base_url = str(settings.get("local_result_server_url") or "http://localhost:8765")
        return str(QRResultService(self.db, self.pairing_service).result_url_for_pairing(tournament_id, pairing_id, base_url)["url"])

    def _qr_result_urls(self, tournament_id: int, pairings: list[dict[str, Any]]) -> dict[int, str]:
        if not pairings:
            return {}
        from src.services.qr_result_service import QRResultService

        settings = self.db.get_app_settings()
        base_url = str(settings.get("local_result_server_url") or "http://localhost:8765")
        return QRResultService(self.db, self.pairing_service).result_urls_for_pairings(tournament_id, pairings, base_url)

    def _team_pairings_section(self, round_data: dict[str, Any]) -> tuple[str, list[str], list[list[Any]]]:
        rows = []
        for match in self.db.list_team_matches_for_round(int(round_data["id"])):
            match_result = self._team_match_score(match)
            white_match_points = self._team_match_points_label(match, "white")
            black_match_points = self._team_match_points_label(match, "black")
            if match.get("is_bye"):
                rows.append(
                    [
                        match["match_number"],
                        match.get("white_team_name", ""),
                        white_match_points,
                        match_result,
                        black_match_points,
                        "BYE",
                        "",
                        "",
                        "BYE",
                        "",
                    ]
                )
                continue
            for board in self.db.list_team_boards(int(match["id"])):
                rows.append(
                    [
                        match["match_number"],
                        match.get("white_team_name", ""),
                        white_match_points,
                        match_result,
                        black_match_points,
                        match.get("black_team_name", ""),
                        board["board_number"],
                        self._team_board_player_name(board, "white"),
                        board.get("result", ""),
                        self._team_board_player_name(board, "black"),
                    ]
                )
        return (
            f"Rodada {round_data['number']}",
            [
                "Confronto",
                "Equipe A",
                "MP A",
                "Placar",
                "MP B",
                "Equipe B",
                "Tabuleiro",
                "Brancas",
                "Resultado",
                "Pretas",
            ],
            rows,
        )

    def _standings_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        tournament = self.db.get_tournament(tournament_id)
        if tournament and tournament.get("competition_type") == "team":
            return self._team_standings_section(tournament_id)
        standings = self.pairing_service.standings(tournament_id)
        columns = resolve_columns(
            self.db.get_report_layout_columns(tournament_id, "standings"), "standings"
        )
        headers = [STANDINGS_COLUMNS[key] for key in columns]
        rows = [[item.get(key, "") for key in columns] for item in standings]
        return ("Classificacao", headers, rows)

    def _crosstable_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        payload = self.pairing_service.crosstable(tournament_id)
        rounds = list(payload["rounds"])
        if payload.get("competition_type") == "team":
            headers = ["Pos", "Equipe", *[f"R{number}" for number in rounds], "MP", "GP", "Buchholz", "Clube/Cidade"]
            rows = [
                [
                    item["position"],
                    item["name"],
                    *[item["rounds"][number]["label"] for number in rounds],
                    self._format_report_number(item["match_points"]),
                    self._format_report_number(item["game_points"]),
                    self._format_report_number(item["buchholz"]),
                    item.get("club", ""),
                ]
                for item in payload["rows"]
            ]
            return f"Tabela cruzada por equipes - {payload['tournament_name']}", headers, rows
        headers = ["Pos", "Jogador", "Rating", *[f"R{number}" for number in rounds], "Pts", "Buchholz", "SB", "Clube"]
        rows = [
            [
                item["position"],
                item["name"],
                item["rating"],
                *[item["rounds"][number]["label"] for number in rounds],
                self._format_report_number(item["points"]),
                self._format_report_number(item["buchholz"]),
                self._format_report_number(item["sonneborn_berger"]),
                item.get("club", ""),
            ]
            for item in payload["rows"]
        ]
        return f"Tabela cruzada - {payload['tournament_name']}", headers, rows

    @staticmethod
    def _crosstable_html(title: str, headers: list[str], rows: list[list[Any]]) -> str:
        header_html = "".join(f"<th>{html.escape(str(header))}</th>" for header in headers)
        rows_html = "".join(
            "<tr>" + "".join(f"<td>{html.escape(str('' if value is None else value))}</td>" for value in row) + "</tr>"
            for row in rows
        )
        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html.escape(title)}</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 24px; color: #111827; }}
    h1 {{ font-size: 22px; }}
    table {{ border-collapse: collapse; width: 100%; font-size: 12px; }}
    th, td {{ border: 1px solid #9ca3af; padding: 5px 7px; text-align: center; }}
    th {{ background: #e5e7eb; }}
    td:nth-child(2), td:last-child {{ text-align: left; }}
  </style>
</head>
<body>
  <h1>{html.escape(title)}</h1>
  <p>Celulas: posicao do adversario, lado (B = brancas; P = pretas), resultado e pontuacao.</p>
  <table>
    <thead><tr>{header_html}</tr></thead>
    <tbody>{rows_html}</tbody>
  </table>
</body>
</html>
"""

    @staticmethod
    def _tiebreak_component_summary(component: dict[str, Any]) -> str:
        if "opponents" in component:
            return " | ".join(
                "{name}: {value}".format(
                    name=item.get("opponent_name", ""),
                    value=item.get("contribution", item.get("points", "")),
                )
                for item in component.get("opponents", [])
            )
        if "used_scores" in component:
            cuts = []
            if component.get("cut_low") is not None:
                cuts.append(f"corte menor {component['cut_low']}")
            if component.get("cut_high") is not None:
                cuts.append(f"corte maior {component['cut_high']}")
            return f"usados: {component.get('used_scores', [])}; {'; '.join(cuts)}"
        if "games" in component:
            return " | ".join(
                "{name}: {earned}".format(
                    name=item.get("opponent_name", ""),
                    earned=item.get("earned", ""),
                )
                for item in component.get("games", [])
            )
        return json.dumps(component, ensure_ascii=False, sort_keys=True)

    def _team_standings_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        standings = self.pairing_service.team_standings(tournament_id)
        rows = [
            [
                item["position"],
                item["name"],
                item.get("club", ""),
                item.get("captain", ""),
                self._format_report_number(item["match_points"]),
                self._format_report_number(item["game_points"]),
                item["wins"],
                item["draws"],
                item["losses"],
                item["byes"],
                item["matches"],
                self._format_report_number(item["buchholz"]),
                "Ativa" if item.get("active") else "Inativa",
            ]
            for item in standings
        ]
        return (
            "Classificacao por equipes",
            [
                "Pos",
                "Equipe",
                "Clube/Cidade",
                "Capitao",
                "Match points",
                "Game points",
                "Vitorias",
                "Empates",
                "Derrotas",
                "Byes",
                "Confrontos",
                "Buchholz",
                "Status",
            ],
            rows,
        )

    def _team_lineups_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        rows = []
        for lineup in self.db.list_team_lineups(tournament_id):
            for board in self.db.list_team_lineup_boards(int(lineup["id"])):
                rows.append(
                    [
                        lineup.get("round_number", ""),
                        lineup.get("match_number", ""),
                        lineup.get("team_name", ""),
                        board.get("board_number", ""),
                        board.get("color", ""),
                        player_full_name(
                            {
                                "name": board.get("player_name", ""),
                                "surname": board.get("player_surname", ""),
                                "given_name": board.get("player_given_name", ""),
                            }
                        ),
                        board.get("player_rating", ""),
                        board.get("role", ""),
                        lineup.get("status", ""),
                    ]
                )
        return (
            "Escalacoes por equipes",
            ["Rodada", "Match", "Equipe", "Tabuleiro", "Cor", "Jogador", "Rating", "Funcao", "Status"],
            rows,
        )

    def _team_substitutions_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        rows = [
            [
                item.get("created_at", ""),
                item.get("round_number", ""),
                item.get("match_number", ""),
                item.get("team_name", ""),
                item.get("board_number", ""),
                item.get("color", ""),
                item.get("out_player_name", ""),
                item.get("in_player_name", ""),
                "Sim" if item.get("requires_correction") else "Nao",
                item.get("reason", ""),
            ]
            for item in self.db.list_team_substitution_events(tournament_id)
        ]
        return (
            "Substituicoes por equipes",
            [
                "Data/hora",
                "Rodada",
                "Match",
                "Equipe",
                "Tabuleiro",
                "Cor",
                "Saiu",
                "Entrou",
                "Correcao formal",
                "Motivo",
            ],
            rows,
        )

    def _standings_for_tournament(
        self,
        tournament_id: int,
        tournament: dict[str, Any] | None = None,
    ) -> list[dict[str, Any]]:
        tournament = tournament or self.db.get_tournament(tournament_id)
        if tournament and tournament.get("competition_type") == "team":
            return self.pairing_service.team_standings(tournament_id)
        return self.pairing_service.standings(tournament_id)

    @staticmethod
    def _team_assignment_player_name(assignment: Mapping[str, Any]) -> str:
        return player_full_name(
            {
                "name": assignment.get("player_name"),
                "surname": assignment.get("player_surname"),
                "given_name": assignment.get("player_given_name"),
            }
        )

    @staticmethod
    def _team_board_player_name(board: Mapping[str, Any], color: str) -> str:
        return player_pairing_name(
            {
                "name": board.get(f"{color}_player_name"),
                "surname": board.get(f"{color}_player_surname"),
                "given_name": board.get(f"{color}_player_given_name"),
            }
        )

    @staticmethod
    def _team_match_score(match: Mapping[str, Any]) -> str:
        if match.get("is_bye"):
            return "BYE"
        if not match.get("result"):
            return ""
        return (
            f"{ExportService._format_report_number(match.get('white_game_points', 0))} x "
            f"{ExportService._format_report_number(match.get('black_game_points', 0))}"
        )

    @staticmethod
    def _team_match_points_label(match: Mapping[str, Any], color: str) -> str:
        if not match.get("result"):
            return ""
        return ExportService._format_report_number(match.get(f"{color}_match_points", 0))

    @staticmethod
    def _format_report_number(value: Any) -> str:
        try:
            numeric = float(value or 0)
        except (TypeError, ValueError):
            return str(value or "")
        if numeric.is_integer():
            return str(int(numeric))
        return f"{numeric:.2f}".rstrip("0").rstrip(".")

    def _write_report(
        self,
        file_path: str | Path,
        title: str,
        headers: list[str],
        rows: list[list[Any]],
        widths: list[int] | None = None,
    ) -> None:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension == ".csv":
            self._write_csv(path, headers, rows)
            logger.info("Relatorio exportado em CSV: %s", path)
            return
        if extension == ".xlsx":
            self._write_xlsx(path, title, headers, rows, widths)
            logger.info("Relatorio exportado em XLSX: %s", path)
            return
        if extension == ".pdf":
            self._write_pdf(path, title, headers, rows)
            logger.info("Relatorio exportado em PDF: %s", path)
            return
        raise AppError("Formato nao suportado. Use .csv, .xlsx ou .pdf.")

    def _write_multi_report(
        self,
        file_path: str | Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension == ".csv":
            self._write_multi_csv(path, sections)
            logger.info("Relatorio composto exportado em CSV: %s", path)
            return
        if extension == ".xlsx":
            self._write_multi_xlsx(path, sections)
            logger.info("Relatorio composto exportado em XLSX: %s", path)
            return
        if extension == ".pdf":
            self._write_multi_pdf(path, sections)
            logger.info("Relatorio composto exportado em PDF: %s", path)
            return
        raise AppError("Formato nao suportado. Use .csv, .xlsx ou .pdf.")

    @staticmethod
    def _write_csv(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
        with path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file, delimiter=";")
            writer.writerow(headers)
            writer.writerows(rows)

    @staticmethod
    def _write_multi_csv(
        path: Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        with path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file, delimiter=";")
            for index, (title, headers, rows) in enumerate(sections):
                if index:
                    writer.writerow([])
                writer.writerow([title])
                writer.writerow(headers)
                writer.writerows(rows)

    @staticmethod
    def _write_xlsx(
        path: Path,
        title: str,
        headers: list[str],
        rows: list[list[Any]],
        widths: list[int] | None = None,
    ) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise AppError("Instale openpyxl para exportar Excel.") from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title[:31]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for column_cells in sheet.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)
        # Larguras explicitas do layout sobrescrevem o auto-dimensionamento.
        if widths:
            for index, width in enumerate(widths):
                if width and width > 0:
                    sheet.column_dimensions[get_column_letter(index + 1)].width = int(width)
        workbook.save(path)

    @staticmethod
    def _write_multi_xlsx(
        path: Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as exc:
            raise AppError("Instale openpyxl para exportar Excel.") from exc

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        used_names: set[str] = set()

        for title, headers, rows in sections:
            sheet_name = ExportService._unique_sheet_name(title, used_names)
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in rows:
                sheet.append(row)
            for column_cells in sheet.columns:
                length = max(len(str(cell.value or "")) for cell in column_cells)
                sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 42)

        workbook.save(path)

    @staticmethod
    def _unique_sheet_name(title: str, used_names: set[str]) -> str:
        invalid = set("[]:*?/\\")
        base = "".join("_" if character in invalid else character for character in title).strip()
        base = (base or "Relatorio")[:31]
        name = base
        counter = 2
        while name in used_names:
            suffix = f" {counter}"
            name = f"{base[:31 - len(suffix)]}{suffix}"
            counter += 1
        used_names.add(name)
        return name

    @staticmethod
    def _write_pdf(path: Path, title: str, headers: list[str], rows: list[list[Any]]) -> None:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise AppError("Instale reportlab para exportar PDF.") from exc

        document = SimpleDocTemplate(str(path), pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        data = [headers] + [[str(value) for value in row] for row in rows]
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9CA3AF")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        document.build([Paragraph(title, styles["Title"]), Spacer(1, 12), table])

    @staticmethod
    def _write_multi_pdf(
        path: Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise AppError("Instale reportlab para exportar PDF.") from exc

        document = SimpleDocTemplate(str(path), pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        story = []
        for index, (title, headers, rows) in enumerate(sections):
            if index:
                story.append(PageBreak())
            story.append(Paragraph(title, styles["Title"]))
            story.append(Spacer(1, 12))
            data = [headers] + [[str(value) for value in row] for row in rows]
            table = Table(data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9CA3AF")),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ]
                )
            )
            story.append(table)
        document.build(story)
