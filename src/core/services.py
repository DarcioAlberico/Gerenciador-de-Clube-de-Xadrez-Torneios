from __future__ import annotations

import csv
import html
import json
import logging
import math
import secrets
import shutil
import sqlite3
import unicodedata
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Mapping

from .database import BASE_DIR, DEFAULT_CERTIFICATE_TEMPLATES, Database

logger = logging.getLogger(__name__)

RESULTS = ["", "1-0", "0-1", "1/2-1/2", "1F-0F", "0F-1F", "0F-0F"]
FINAL_RESULTS = {"1-0", "0-1", "1/2-1/2", "1F-0F", "0F-1F", "0F-0F", "BYE"}
RESULT_POINTS = {
    "1-0": (1.0, 0.0),
    "0-1": (0.0, 1.0),
    "1/2-1/2": (0.5, 0.5),
    "1F-0F": (1.0, 0.0),
    "0F-1F": (0.0, 1.0),
    "0F-0F": (0.0, 0.0),
}
MEMBER_TYPES = {"socio", "aluno", "convidado", "visitante"}
MEMBER_STATUSES = {"active", "inactive", "visitor", "guest", "withdrawn"}
PLAYER_STATUSES = {
    "active": "Ativo",
    "withdrawn": "Desistente",
    "absent": "Ausente",
    "inactive": "Nao emparceirado",
}
INITIAL_ORDER_OPTIONS = {
    "rating": "Rating principal",
    "national_rating": "Rating nacional",
    "international_rating": "Rating internacional",
    "international_then_national": "Rating internacional depois nacional",
    "max_rating": "Maior rating",
    "manual": "Manual",
}
TOURNAMENT_TYPES = {"real": "Real", "test": "Teste"}
TOURNAMENT_PROFILES = {
    "free": "Livre/Escolar",
    "club": "Clube/Semi-formal",
    "fide": "FIDE-rated",
}
COMPETITION_TYPES = {"individual": "Individual", "team": "Equipes"}
PAIRING_METHODS = {
    "swiss": "Suíço",
    "round_robin": "Schuring (Todos contra todos)",
    "knockout": "Mata-mata",
}
TEAM_PAIRING_METHODS = {"swiss": "Suico", "round_robin": "Schuring (Todos contra todos)"}
CERTIFICATE_TYPES = {
    "participation": "Participacao",
    "overall_award": "Premiacao geral",
    "category_award": "Premiacao por categoria",
    "member_certificate": "Membro/aluno",
    "training_participation": "Aula/turma",
    "event_participation": "Evento",
    "internal_ranking": "Ranking interno",
}
CERTIFICATE_ORIENTATIONS = {"landscape": "Paisagem", "portrait": "Retrato"}
TEAM_PLAYER_ROLES = {"starter": "Titular", "reserve": "Reserva"}
TEAM_PAIRING_METHODS = {"swiss": "Suico por equipes"}
TEAM_STANDING_CRITERIA = {
    "match_points": "Match points",
    "game_points": "Game points",
    "wins": "Vitorias",
}
TOURNAMENT_SCOPES = {
    "standalone": "Avulso",
    "club": "Clube/Escola",
    "class": "Turma",
}


def _player_name_parts(player: Mapping[str, Any]) -> tuple[str, str]:
    name = str(player.get("name") or "").strip()
    surname = str(player.get("surname") or "").strip()
    given_name = str(player.get("given_name") or "").strip()
    if given_name:
        return given_name, surname
    if surname and "," in name:
        left, right = name.split(",", maxsplit=1)
        if left.strip().casefold() == surname.casefold():
            return right.strip(), surname
    if surname:
        name_casefold = name.casefold()
        surname_casefold = surname.casefold()
        if name_casefold.endswith(f" {surname_casefold}"):
            return name[: -len(surname)].strip(), surname
        if name_casefold == surname_casefold:
            return "", surname
    return name, surname


def player_full_name(player: Mapping[str, Any]) -> str:
    given_name, surname = _player_name_parts(player)
    if given_name and surname:
        return f"{given_name} {surname}"
    return given_name or surname


def player_pairing_name(player: Mapping[str, Any]) -> str:
    given_name, surname = _player_name_parts(player)
    if given_name and surname:
        return f"{surname}, {given_name}"
    return given_name or surname


def pairing_player_name(pairing: Mapping[str, Any], color: str) -> str:
    return player_pairing_name(
        {
            "name": pairing.get(f"{color}_name"),
            "surname": pairing.get(f"{color}_surname"),
            "given_name": pairing.get(f"{color}_given_name"),
        }
    )
CLUB_KINDS = {"club", "school", "project", "partner"}
TRAINING_SESSION_TYPES = {"aula": "Aula", "treino": "Treino", "evento": "Evento"}
TRAINING_SESSION_STATUSES = {"planned": "Planejada", "done": "Realizada", "canceled": "Cancelada"}
ATTENDANCE_STATUSES = {"present": "Presente", "absent": "Falta", "justified": "Justificada"}
EXERCISE_DIFFICULTIES = {
    "easy": "Facil",
    "basic": "Basico",
    "intermediate": "Intermediario",
    "advanced": "Avancado",
    "competition": "Competitivo",
}
TRAINING_LIST_STATUSES = {
    "draft": "Rascunho",
    "ready": "Pronta",
    "used": "Aplicada",
    "archived": "Arquivada",
}
EXERCISE_ATTEMPT_RESULTS = {
    "correct": "Correto",
    "partial": "Parcial",
    "incorrect": "Incorreto",
    "skipped": "Nao resolvido",
}
BILLING_CYCLES = {
    "monthly": "Mensal",
    "single": "Avulso",
    "annual": "Anual",
    "custom": "Personalizado",
}
PAYMENT_STATUSES = {
    "pending": "Pendente",
    "paid": "Pago",
    "late": "Atrasado",
    "exempt": "Isento",
    "canceled": "Cancelado",
}
EVENT_TYPES = {
    "tournament": "Torneio",
    "class": "Aula/Treino",
    "meeting": "Reuniao",
    "simultaneous": "Simultanea",
    "travel": "Viagem",
    "social": "Encontro",
    "other": "Outro",
}
EVENT_STATUSES = {
    "planned": "Planejado",
    "confirmed": "Confirmado",
    "done": "Concluido",
    "canceled": "Cancelado",
}
INVENTORY_ITEM_TYPES = {
    "pieces": "Pecas",
    "boards": "Tabuleiros",
    "clocks": "Relogios",
    "books": "Livros",
    "sets": "Kits",
    "digital": "Digital",
    "other": "Outro",
}
INVENTORY_CONDITIONS = {
    "new": "Novo",
    "good": "Bom",
    "worn": "Uso intenso",
    "damaged": "Danificado",
    "lost": "Perdido",
}
INVENTORY_LOAN_STATUSES = {
    "open": "Emprestado",
    "returned": "Devolvido",
    "lost": "Perdido",
}
INVENTORY_MAINTENANCE_STATUSES = {
    "open": "Aberta",
    "in_progress": "Em andamento",
    "done": "Concluida",
    "canceled": "Cancelada",
}
OPERATOR_ROLES = {
    "admin": "Administrador",
    "arbiter": "Arbitragem",
    "teacher": "Professor",
    "assistant": "Assistente",
    "viewer": "Consulta",
}
TOURNAMENT_FLAG_FIELDS = {
    "allow_public_registration",
    "allow_player_result_edit",
    "allow_dangerous_changes",
    "disable_bye",
    "accelerated_system",
    "hide_standings",
    "calculate_performance",
    "hide_color_names",
    "show_opponents_in_standings",
    "archived",
}


class AppError(Exception):
    """Erro esperado de regra de negocio, seguro para mostrar ao usuario."""


class ClubService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def save_profile(self, data: dict[str, Any], club_id: int | None = 1) -> int:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome do clube/escola.")
        kind = str(data.get("kind", "club")).strip() or "club"
        if kind not in CLUB_KINDS:
            raise AppError("Tipo de unidade invalido.")
        saved_id = self.db.save_club(
            name=name,
            kind=kind,
            city=str(data.get("city", "")),
            address=str(data.get("address", "")),
            phone=str(data.get("phone", "")),
            email=str(data.get("email", "")),
            notes=str(data.get("notes", "")),
            active=1 if data.get("active", 1) else 0,
            club_id=club_id,
        )
        logger.info("Clube/escola atualizado: %s", saved_id)
        return saved_id

    def save_class(self, data: dict[str, Any], class_id: int | None = None) -> int:
        try:
            club_id = int(data.get("club_id") or 0)
        except ValueError as exc:
            raise AppError("Clube/escola invalido para a turma.") from exc
        if not self.db.get_club(club_id):
            raise AppError("Clube/escola nao encontrado.")

        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome da turma.")

        payload: dict[str, Any] = {
            "club_id": club_id,
            "name": name,
            "teacher": str(data.get("teacher", "")),
            "weekday": str(data.get("weekday", "")),
            "time": str(data.get("time", "")),
            "location": str(data.get("location", "")),
            "active": 1 if data.get("active", 1) else 0,
            "notes": str(data.get("notes", "")),
        }
        if class_id:
            if not self.db.get_class(class_id):
                raise AppError("Turma nao encontrada.")
            self.db.update_class(class_id, **payload)
            logger.info("Turma atualizada: %s", class_id)
            return class_id
        saved_id = self.db.create_class(**payload)
        logger.info("Turma criada: %s", saved_id)
        return saved_id


class GuardianService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def create_guardian(self, data: dict[str, Any]) -> int:
        payload = self._validated_payload(data)
        guardian_id = self.db.create_guardian(**payload)
        logger.info("Responsavel criado: %s", guardian_id)
        return guardian_id

    def update_guardian(self, guardian_id: int, data: dict[str, Any]) -> None:
        if not self.db.get_guardian(guardian_id):
            raise AppError("Responsavel nao encontrado.")
        payload = self._validated_payload(data)
        self.db.update_guardian(guardian_id, **payload)
        logger.info("Responsavel atualizado: %s", guardian_id)

    def toggle_guardian_active(self, guardian_id: int) -> int:
        guardian = self.db.get_guardian(guardian_id)
        if not guardian:
            raise AppError("Responsavel nao encontrado.")
        next_active = 0 if guardian.get("active") else 1
        self.db.set_guardian_active(guardian_id, next_active)
        logger.info("Responsavel %s alterado para active=%s", guardian_id, next_active)
        return next_active

    def link_guardian_to_member(self, data: dict[str, Any]) -> None:
        try:
            member_id = int(data.get("member_id") or 0)
            guardian_id = int(data.get("guardian_id") or 0)
        except ValueError as exc:
            raise AppError("Membro ou responsavel invalido.") from exc

        member = self.db.get_member(member_id)
        if not member:
            raise AppError("Membro nao encontrado.")
        guardian = self.db.get_guardian(guardian_id)
        if not guardian:
            raise AppError("Responsavel nao encontrado.")
        if not guardian.get("active"):
            raise AppError("Responsavel inativo nao pode receber novo vinculo.")

        self.db.link_guardian_to_member(
            member_id=member_id,
            guardian_id=guardian_id,
            relationship=str(data.get("relationship", "")),
            primary_contact=1 if data.get("primary_contact") else 0,
            emergency_contact=1 if data.get("emergency_contact") else 0,
            notes=str(data.get("notes", "")),
        )
        logger.info("Responsavel %s vinculado ao membro %s", guardian_id, member_id)

    def unlink_guardian_from_member(self, member_id: int, guardian_id: int) -> None:
        if not self.db.get_member(member_id):
            raise AppError("Membro nao encontrado.")
        if not self.db.get_guardian(guardian_id):
            raise AppError("Responsavel nao encontrado.")
        self.db.unlink_guardian_from_member(member_id, guardian_id)
        logger.info("Responsavel %s removido do membro %s", guardian_id, member_id)

    def minor_members(self, active_only: bool = True) -> list[dict[str, Any]]:
        return [
            member
            for member in self.db.list_members(active_only=active_only)
            if self._is_minor(str(member.get("birth_date") or ""))
        ]

    def minor_members_without_guardians(self, active_only: bool = True) -> list[dict[str, Any]]:
        return [
            member
            for member in self.db.list_members_without_guardians(active_only=active_only)
            if self._is_minor(str(member.get("birth_date") or ""))
        ]

    @staticmethod
    def _validated_payload(data: dict[str, Any]) -> dict[str, Any]:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome do responsavel.")
        return {
            "name": name,
            "phone": str(data.get("phone", "")),
            "email": str(data.get("email", "")),
            "document": str(data.get("document", "")),
            "address": str(data.get("address", "")),
            "notes": str(data.get("notes", "")),
            "active": 1 if data.get("active", 1) else 0,
        }

    @staticmethod
    def _is_minor(birth_date: str) -> bool:
        value = birth_date.strip()
        if not value:
            return False
        try:
            born = date.fromisoformat(value)
        except ValueError:
            return False
        today = date.today()
        age = today.year - born.year - ((today.month, today.day) < (born.month, born.day))
        return age < 18


class TrainingService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def save_session(self, data: dict[str, Any], session_id: int | None = None) -> int:
        payload = self._validated_session_payload(data)
        if session_id:
            if not self.db.get_training_session(session_id):
                raise AppError("Aula/treino nao encontrado.")
            self.db.update_training_session(session_id, **payload)
            logger.info("Aula/treino atualizado: %s", session_id)
            return session_id
        saved_id = self.db.create_training_session(**payload)
        logger.info("Aula/treino criado: %s", saved_id)
        return saved_id

    def record_attendance(self, session_id: int, rows: list[dict[str, Any]]) -> dict[str, int]:
        session = self.db.get_training_session(session_id)
        if not session:
            raise AppError("Aula/treino nao encontrado.")
        saved = 0
        skipped = 0
        for item in rows:
            try:
                member_id = int(item.get("member_id") or 0)
            except ValueError as exc:
                raise AppError("Membro invalido na chamada.") from exc
            member = self.db.get_member(member_id)
            if not member:
                skipped += 1
                continue
            status = str(item.get("status", "")).strip()
            if not status:
                skipped += 1
                continue
            if status not in ATTENDANCE_STATUSES:
                raise AppError("Status de presenca invalido.")
            self.db.save_attendance(
                session_id=session_id,
                member_id=member_id,
                status=status,
                notes=str(item.get("notes", "")),
            )
            saved += 1
        logger.info("Chamada da sessao %s atualizada: %s registros", session_id, saved)
        return {"saved": saved, "skipped": skipped}

    def member_attendance_summary(
        self,
        member_id: int,
        start_date: str = "",
        end_date: str = "",
    ) -> dict[str, Any]:
        if not self.db.get_member(member_id):
            raise AppError("Membro nao encontrado.")
        self._validate_optional_date(start_date, "Data inicial invalida.")
        self._validate_optional_date(end_date, "Data final invalida.")
        return self.db.member_attendance_summary(member_id, start_date.strip(), end_date.strip())

    def attendance_report(
        self,
        start_date: str = "",
        end_date: str = "",
        club_id: int | None = None,
        member_id: int | None = None,
    ) -> list[dict[str, Any]]:
        self._validate_optional_date(start_date, "Data inicial invalida.")
        self._validate_optional_date(end_date, "Data final invalida.")
        return self.db.list_attendance_report(
            start_date=start_date.strip(),
            end_date=end_date.strip(),
            club_id=club_id,
            member_id=member_id,
        )

    def _validated_session_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        title = str(data.get("title", "")).strip()
        if not title:
            raise AppError("Informe o titulo da aula/treino.")

        try:
            club_id = int(data.get("club_id") or 1)
        except ValueError as exc:
            raise AppError("Clube/escola da aula invalido.") from exc
        if not self.db.get_club(club_id):
            raise AppError("Clube/escola da aula nao encontrado.")

        try:
            class_id = int(data.get("class_id") or 0)
        except ValueError as exc:
            raise AppError("Turma da aula invalida.") from exc
        if class_id:
            class_data = self.db.get_class(class_id)
            if not class_data or int(class_data["club_id"]) != club_id:
                raise AppError("Turma nao pertence ao clube/escola selecionado.")

        try:
            learning_level_id = int(data.get("learning_level_id") or 0)
        except ValueError as exc:
            raise AppError("Nivel pedagogico da aula invalido.") from exc
        if learning_level_id and not self.db.get_learning_level(learning_level_id):
            raise AppError("Nivel pedagogico da aula nao encontrado.")

        try:
            training_list_id = int(data.get("training_list_id") or 0)
        except ValueError as exc:
            raise AppError("Lista de treino da aula invalida.") from exc
        if training_list_id:
            training_list = self.db.get_training_list(training_list_id)
            if not training_list:
                raise AppError("Lista de treino da aula nao encontrada.")
            if int(training_list.get("club_id") or 1) != club_id:
                raise AppError("Lista de treino nao pertence ao clube/escola selecionado.")
            list_class_id = int(training_list.get("class_id") or 0)
            if list_class_id and list_class_id != class_id:
                raise AppError("Lista de treino nao pertence a turma selecionada.")

        session_type = str(data.get("session_type", "aula")).strip() or "aula"
        if session_type not in TRAINING_SESSION_TYPES:
            raise AppError("Tipo de aula/treino invalido.")

        status = str(data.get("status", "planned")).strip() or "planned"
        if status not in TRAINING_SESSION_STATUSES:
            raise AppError("Status de aula/treino invalido.")

        session_date = str(data.get("session_date", "")).strip()
        self._validate_optional_date(session_date, "Data da aula invalida.")

        return {
            "club_id": club_id,
            "class_id": class_id or None,
            "training_list_id": training_list_id or None,
            "title": title,
            "session_type": session_type,
            "session_date": session_date,
            "start_time": str(data.get("start_time", "")),
            "end_time": str(data.get("end_time", "")),
            "instructor": str(data.get("instructor", "")),
            "location": str(data.get("location", "")),
            "learning_level_id": learning_level_id or None,
            "objective": str(data.get("objective", "")),
            "content": str(data.get("content", "")),
            "homework": str(data.get("homework", "")),
            "status": status,
            "notes": str(data.get("notes", "")),
        }

    @staticmethod
    def _validate_optional_date(value: str, message: str) -> None:
        cleaned = value.strip()
        if not cleaned:
            return
        try:
            date.fromisoformat(cleaned)
        except ValueError as exc:
            raise AppError(message) from exc


class ExerciseService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def save_exercise(self, data: dict[str, Any], exercise_id: int | None = None) -> int:
        payload = self._validated_exercise_payload(data)
        if exercise_id:
            if not self.db.get_exercise(exercise_id):
                raise AppError("Exercicio nao encontrado.")
            self.db.update_exercise(exercise_id, **payload)
            logger.info("Exercicio atualizado: %s", exercise_id)
            return exercise_id
        saved_id = self.db.create_exercise(**payload)
        logger.info("Exercicio criado: %s", saved_id)
        return saved_id

    def toggle_exercise_active(self, exercise_id: int) -> int:
        exercise = self.db.get_exercise(exercise_id)
        if not exercise:
            raise AppError("Exercicio nao encontrado.")
        next_active = 0 if exercise.get("active") else 1
        self.db.set_exercise_active(exercise_id, next_active)
        logger.info("Exercicio %s alterado para active=%s", exercise_id, next_active)
        return next_active

    def save_training_list(self, data: dict[str, Any], list_id: int | None = None) -> int:
        payload = self._validated_training_list_payload(data)
        if list_id:
            if not self.db.get_training_list(list_id):
                raise AppError("Lista de treino nao encontrada.")
            self.db.update_training_list(list_id, **payload)
            logger.info("Lista de treino atualizada: %s", list_id)
            return list_id
        saved_id = self.db.create_training_list(**payload)
        logger.info("Lista de treino criada: %s", saved_id)
        return saved_id

    def set_training_list_exercises(self, list_id: int, exercise_ids: list[int]) -> None:
        training_list = self.db.get_training_list(list_id)
        if not training_list:
            raise AppError("Lista de treino nao encontrada.")
        unique_ids: list[int] = []
        for exercise_id in exercise_ids:
            value = int(exercise_id or 0)
            if not value or value in unique_ids:
                continue
            exercise = self.db.get_exercise(value)
            if not exercise:
                raise AppError("Exercicio da lista nao encontrado.")
            if int(exercise.get("club_id") or 1) != int(training_list.get("club_id") or 1):
                raise AppError("Exercicio nao pertence ao clube/escola da lista.")
            unique_ids.append(value)
        rows = [
            {"exercise_id": exercise_id, "position_order": index}
            for index, exercise_id in enumerate(unique_ids, start=1)
        ]
        self.db.replace_training_list_exercises(list_id, rows)
        logger.info("Lista de treino %s atualizada com %s exercicios", list_id, len(rows))

    def add_exercise_to_training_list(self, list_id: int, exercise_id: int) -> None:
        current = [int(item["exercise_id"]) for item in self.db.list_training_list_exercises(list_id)]
        if exercise_id not in current:
            current.append(exercise_id)
        self.set_training_list_exercises(list_id, current)

    def remove_exercise_from_training_list(self, list_id: int, exercise_id: int) -> None:
        current = [
            int(item["exercise_id"])
            for item in self.db.list_training_list_exercises(list_id)
            if int(item["exercise_id"]) != int(exercise_id)
        ]
        self.set_training_list_exercises(list_id, current)

    def record_attempt(self, data: dict[str, Any]) -> int:
        payload = self._validated_attempt_payload(data)
        attempt_id = self.db.create_exercise_attempt(**payload)
        logger.info("Tentativa de exercicio registrada: %s", attempt_id)
        return attempt_id

    def _validated_exercise_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        title = str(data.get("title", "")).strip()
        if not title:
            raise AppError("Informe o titulo do exercicio.")

        try:
            club_id = int(data.get("club_id") or 1)
        except ValueError as exc:
            raise AppError("Clube/escola do exercicio invalido.") from exc
        if not self.db.get_club(club_id):
            raise AppError("Clube/escola do exercicio nao encontrado.")

        try:
            learning_level_id = int(data.get("learning_level_id") or 0)
        except ValueError as exc:
            raise AppError("Nivel pedagogico do exercicio invalido.") from exc
        if learning_level_id and not self.db.get_learning_level(learning_level_id):
            raise AppError("Nivel pedagogico do exercicio nao encontrado.")

        difficulty = str(data.get("difficulty", "basic")).strip() or "basic"
        if difficulty not in EXERCISE_DIFFICULTIES:
            raise AppError("Dificuldade do exercicio invalida.")

        return {
            "club_id": club_id,
            "learning_level_id": learning_level_id or None,
            "title": title,
            "theme": str(data.get("theme", "")),
            "difficulty": difficulty,
            "source": str(data.get("source", "")),
            "fen": str(data.get("fen", "")),
            "pgn": str(data.get("pgn", "")),
            "solution": str(data.get("solution", "")),
            "objective": str(data.get("objective", "")),
            "tags": str(data.get("tags", "")),
            "active": 1 if data.get("active", 1) else 0,
            "notes": str(data.get("notes", "")),
        }

    def _validated_training_list_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome da lista de treino.")

        try:
            club_id = int(data.get("club_id") or 1)
        except ValueError as exc:
            raise AppError("Clube/escola da lista invalido.") from exc
        if not self.db.get_club(club_id):
            raise AppError("Clube/escola da lista nao encontrado.")

        try:
            class_id = int(data.get("class_id") or 0)
        except ValueError as exc:
            raise AppError("Turma da lista invalida.") from exc
        if class_id:
            class_data = self.db.get_class(class_id)
            if not class_data or int(class_data["club_id"]) != club_id:
                raise AppError("Turma nao pertence ao clube/escola selecionado.")

        try:
            learning_level_id = int(data.get("learning_level_id") or 0)
        except ValueError as exc:
            raise AppError("Nivel pedagogico da lista invalido.") from exc
        if learning_level_id and not self.db.get_learning_level(learning_level_id):
            raise AppError("Nivel pedagogico da lista nao encontrado.")

        target_date = str(data.get("target_date", "")).strip()
        TrainingService._validate_optional_date(target_date, "Data alvo da lista invalida.")

        status = str(data.get("status", "draft")).strip() or "draft"
        if status not in TRAINING_LIST_STATUSES:
            raise AppError("Status da lista de treino invalido.")

        return {
            "club_id": club_id,
            "class_id": class_id or None,
            "learning_level_id": learning_level_id or None,
            "name": name,
            "description": str(data.get("description", "")),
            "target_date": target_date,
            "status": status,
            "notes": str(data.get("notes", "")),
        }

    def _validated_attempt_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        try:
            exercise_id = int(data.get("exercise_id") or 0)
        except ValueError as exc:
            raise AppError("Exercicio da tentativa invalido.") from exc
        if not self.db.get_exercise(exercise_id):
            raise AppError("Exercicio da tentativa nao encontrado.")

        try:
            member_id = int(data.get("member_id") or 0)
        except ValueError as exc:
            raise AppError("Aluno da tentativa invalido.") from exc
        if not self.db.get_member(member_id):
            raise AppError("Aluno da tentativa nao encontrado.")

        try:
            list_id = int(data.get("list_id") or 0)
        except ValueError as exc:
            raise AppError("Lista da tentativa invalida.") from exc
        if list_id and not self.db.get_training_list(list_id):
            raise AppError("Lista da tentativa nao encontrada.")

        try:
            session_id = int(data.get("session_id") or 0)
        except ValueError as exc:
            raise AppError("Aula da tentativa invalida.") from exc
        if session_id and not self.db.get_training_session(session_id):
            raise AppError("Aula da tentativa nao encontrada.")

        attempt_date = str(data.get("attempt_date", "")).strip() or date.today().isoformat()
        TrainingService._validate_optional_date(attempt_date, "Data da tentativa invalida.")

        result = str(data.get("result", "incorrect")).strip() or "incorrect"
        if result not in EXERCISE_ATTEMPT_RESULTS:
            raise AppError("Resultado da tentativa invalido.")

        score_value = data.get("score")
        if score_value in (None, ""):
            score = {"correct": 1.0, "partial": 0.5}.get(result, 0.0)
        else:
            try:
                score = float(str(score_value).replace(",", "."))
            except ValueError as exc:
                raise AppError("Pontuacao da tentativa invalida.") from exc
        if score < 0 or score > 1:
            raise AppError("Pontuacao da tentativa deve ficar entre 0 e 1.")

        try:
            time_seconds = int(data.get("time_seconds") or 0)
        except ValueError as exc:
            raise AppError("Tempo da tentativa invalido.") from exc
        if time_seconds < 0:
            raise AppError("Tempo da tentativa nao pode ser negativo.")

        return {
            "exercise_id": exercise_id,
            "member_id": member_id,
            "list_id": list_id or None,
            "session_id": session_id or None,
            "attempt_date": attempt_date,
            "result": result,
            "score": score,
            "time_seconds": time_seconds,
            "notes": str(data.get("notes", "")),
        }


class FinanceService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def save_plan(self, data: dict[str, Any], plan_id: int | None = None) -> int:
        payload = self._validated_plan_payload(data)
        if plan_id:
            if not self.db.get_membership_plan(plan_id):
                raise AppError("Plano nao encontrado.")
            self.db.update_membership_plan(plan_id, **payload)
            logger.info("Plano financeiro atualizado: %s", plan_id)
            return plan_id
        saved_id = self.db.create_membership_plan(**payload)
        logger.info("Plano financeiro criado: %s", saved_id)
        return saved_id

    def toggle_plan_active(self, plan_id: int) -> int:
        plan = self.db.get_membership_plan(plan_id)
        if not plan:
            raise AppError("Plano nao encontrado.")
        next_active = 0 if plan.get("active") else 1
        self.db.set_membership_plan_active(plan_id, next_active)
        logger.info("Plano financeiro %s alterado para active=%s", plan_id, next_active)
        return next_active

    def save_payment(self, data: dict[str, Any], payment_id: int | None = None) -> int:
        payload = self._validated_payment_payload(data)
        if payment_id:
            if not self.db.get_payment(payment_id):
                raise AppError("Pagamento nao encontrado.")
            self.db.update_payment(payment_id, **payload)
            logger.info("Pagamento atualizado: %s", payment_id)
            return payment_id
        saved_id = self.db.create_payment(**payload)
        logger.info("Pagamento criado: %s", saved_id)
        return saved_id

    def mark_payment_paid(
        self,
        payment_id: int,
        payment_date: str = "",
        method: str = "",
    ) -> None:
        payment = self.db.get_payment(payment_id)
        if not payment:
            raise AppError("Pagamento nao encontrado.")
        payload = dict(payment)
        payload["status"] = "paid"
        payload["payment_date"] = payment_date.strip() or date.today().isoformat()
        if method.strip():
            payload["method"] = method
        self.save_payment(payload, payment_id)

    def generate_recurring_payments(
        self,
        plan_id: int,
        reference_period: str,
        due_date: str,
        member_ids: list[int] | None = None,
        club_id: int | None = None,
        class_id: int | None = None,
    ) -> dict[str, Any]:
        plan = self.db.get_membership_plan(int(plan_id or 0))
        if not plan:
            raise AppError("Plano nao encontrado.")
        if not plan.get("active"):
            raise AppError("Selecione um plano ativo para gerar mensalidades.")

        reference = self._validated_reference_period(reference_period)
        due = due_date.strip()
        self._validate_optional_date(due, "Vencimento invalido.")
        if not due:
            raise AppError("Informe o vencimento das mensalidades.")

        members = self._members_for_recurring_payments(member_ids, club_id, class_id)
        if not members:
            raise AppError("Nao ha membros ativos para gerar mensalidades.")

        created_ids: list[int] = []
        skipped = 0
        amount = float(plan.get("amount") or 0.0)
        description = f"{plan['name']} - {reference}"
        for member in members:
            member_id = int(member["id"])
            if self._has_open_payment_for_reference(member_id, int(plan["id"]), reference):
                skipped += 1
                continue
            created_ids.append(
                self.save_payment(
                    {
                        "member_id": member_id,
                        "plan_id": int(plan["id"]),
                        "description": description,
                        "reference_period": reference,
                        "due_date": due,
                        "amount": amount,
                        "status": "pending",
                    }
                )
            )

        logger.info(
            "Mensalidades geradas para plano %s referencia %s: %s criadas, %s puladas",
            plan_id,
            reference,
            len(created_ids),
            skipped,
        )
        return {
            "created": len(created_ids),
            "skipped": skipped,
            "payment_ids": created_ids,
            "reference_period": reference,
            "due_date": due,
            "plan_name": str(plan.get("name") or ""),
        }

    def payments_report(
        self,
        start_date: str = "",
        end_date: str = "",
        member_id: int | None = None,
    ) -> list[dict[str, Any]]:
        self._validate_optional_date(start_date, "Data inicial invalida.")
        self._validate_optional_date(end_date, "Data final invalida.")
        rows = self.db.list_payments(
            start_date=start_date.strip(),
            end_date=end_date.strip(),
            member_id=member_id,
            include_canceled=False,
        )
        return [self._with_effective_status(row) for row in rows]

    def finance_summary(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> dict[str, Any]:
        payments = self.payments_report(start_date=start_date, end_date=end_date)
        summary = {
            "total": len(payments),
            "paid": 0,
            "pending": 0,
            "late": 0,
            "exempt": 0,
            "paid_amount": 0.0,
            "pending_amount": 0.0,
            "late_amount": 0.0,
            "exempt_amount": 0.0,
        }
        for payment in payments:
            status = payment["effective_status"]
            amount = float(payment.get("amount") or 0.0)
            if status in {"paid", "pending", "late", "exempt"}:
                summary[status] += 1
                summary[f"{status}_amount"] += amount
        return summary

    def member_financial_status(self, member_id: int) -> dict[str, Any]:
        if not self.db.get_member(member_id):
            raise AppError("Membro nao encontrado.")
        payments = [
            self._with_effective_status(payment)
            for payment in self.db.list_payments(member_id=member_id, include_canceled=False)
        ]
        if not payments:
            return {"status": "none", "label": "Sem lancamentos", "open_amount": 0.0, "late_amount": 0.0}

        late_amount = sum(
            float(payment.get("amount") or 0.0)
            for payment in payments
            if payment["effective_status"] == "late"
        )
        pending_amount = sum(
            float(payment.get("amount") or 0.0)
            for payment in payments
            if payment["effective_status"] == "pending"
        )
        if late_amount:
            status = "late"
        elif pending_amount:
            status = "pending"
        elif payments and all(payment["effective_status"] == "exempt" for payment in payments):
            status = "exempt"
        else:
            status = "paid"
        return {
            "status": status,
            "label": PAYMENT_STATUSES.get(status, "Sem lancamentos"),
            "open_amount": round(pending_amount + late_amount, 2),
            "late_amount": round(late_amount, 2),
        }

    def _members_for_recurring_payments(
        self,
        member_ids: list[int] | None,
        club_id: int | None,
        class_id: int | None,
    ) -> list[dict[str, Any]]:
        if member_ids is not None:
            members = []
            for member_id in member_ids:
                member = self.db.get_member(int(member_id))
                if not member:
                    raise AppError("Membro selecionado para mensalidade nao encontrado.")
                members.append(member)
        else:
            members = self.db.list_members(active_only=True, club_id=club_id, class_id=class_id)
        billable_types = {"socio", "aluno"}
        return [
            member
            for member in members
            if str(member.get("status") or "") == "active"
            and str(member.get("member_type") or "") in billable_types
        ]

    def _has_open_payment_for_reference(self, member_id: int, plan_id: int, reference_period: str) -> bool:
        return any(
            int(payment.get("plan_id") or 0) == plan_id
            and str(payment.get("reference_period") or "").strip() == reference_period
            for payment in self.db.list_payments(member_id=member_id, include_canceled=False)
        )

    @staticmethod
    def _validated_reference_period(value: str) -> str:
        reference = value.strip()
        parts = reference.split("-")
        if len(parts) != 2 or len(parts[0]) != 4 or len(parts[1]) != 2:
            raise AppError("Informe a referencia no formato AAAA-MM.")
        try:
            year = int(parts[0])
            month = int(parts[1])
            date(year, month, 1)
        except ValueError as exc:
            raise AppError("Referencia financeira invalida.") from exc
        return reference

    def _validated_plan_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome do plano.")
        billing_cycle = str(data.get("billing_cycle", "monthly")).strip() or "monthly"
        if billing_cycle not in BILLING_CYCLES:
            raise AppError("Ciclo de cobranca invalido.")
        amount = self._parse_amount(data.get("amount"), "Valor do plano invalido.")
        if amount < 0:
            raise AppError("Valor do plano nao pode ser negativo.")
        return {
            "name": name,
            "amount": amount,
            "billing_cycle": billing_cycle,
            "active": 1 if data.get("active", 1) else 0,
            "notes": str(data.get("notes", "")),
        }

    def _validated_payment_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        try:
            member_id = int(data.get("member_id") or 0)
        except ValueError as exc:
            raise AppError("Membro do pagamento invalido.") from exc
        if not self.db.get_member(member_id):
            raise AppError("Membro do pagamento nao encontrado.")

        try:
            plan_id = int(data.get("plan_id") or 0)
        except ValueError as exc:
            raise AppError("Plano do pagamento invalido.") from exc
        plan = self.db.get_membership_plan(plan_id) if plan_id else None
        if plan_id and not plan:
            raise AppError("Plano do pagamento nao encontrado.")

        status = str(data.get("status", "pending")).strip() or "pending"
        if status not in PAYMENT_STATUSES:
            raise AppError("Status financeiro invalido.")

        due_date = str(data.get("due_date", "")).strip()
        payment_date = str(data.get("payment_date", "")).strip()
        self._validate_optional_date(due_date, "Vencimento invalido.")
        self._validate_optional_date(payment_date, "Data de pagamento invalida.")

        amount_value = data.get("amount")
        amount = self._parse_amount(amount_value, "Valor do pagamento invalido.")
        if amount <= 0 and plan:
            amount = float(plan.get("amount") or 0.0)
        if amount < 0:
            raise AppError("Valor do pagamento nao pode ser negativo.")

        description = str(data.get("description", "")).strip()
        if not description and plan:
            description = str(plan.get("name") or "")
        if not description:
            description = "Lancamento financeiro"

        if status == "paid" and not payment_date:
            payment_date = date.today().isoformat()

        return {
            "member_id": member_id,
            "plan_id": plan_id or None,
            "description": description,
            "reference_period": str(data.get("reference_period", "")),
            "due_date": due_date,
            "payment_date": payment_date,
            "amount": amount,
            "status": status,
            "method": str(data.get("method", "")),
            "notes": str(data.get("notes", "")),
        }

    @staticmethod
    def _parse_amount(value: Any, message: str) -> float:
        try:
            return round(float(str(value or "0").replace(",", ".")), 2)
        except ValueError as exc:
            raise AppError(message) from exc

    @staticmethod
    def _validate_optional_date(value: str, message: str) -> None:
        cleaned = value.strip()
        if not cleaned:
            return
        try:
            date.fromisoformat(cleaned)
        except ValueError as exc:
            raise AppError(message) from exc

    @classmethod
    def _with_effective_status(cls, payment: dict[str, Any]) -> dict[str, Any]:
        row = dict(payment)
        row["effective_status"] = cls._effective_payment_status(row)
        return row

    @staticmethod
    def _effective_payment_status(payment: dict[str, Any]) -> str:
        status = str(payment.get("status") or "pending")
        due_date = str(payment.get("due_date") or "").strip()
        if status == "pending" and due_date:
            try:
                if date.fromisoformat(due_date) < date.today():
                    return "late"
            except ValueError:
                return status
        return status


class EventService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def save_event(self, data: dict[str, Any], event_id: int | None = None) -> int:
        payload = self._validated_event_payload(data)
        if event_id:
            if not self.db.get_club_event(event_id):
                raise AppError("Evento nao encontrado.")
            self.db.update_club_event(event_id, **payload)
            logger.info("Evento atualizado: %s", event_id)
            return event_id
        saved_id = self.db.create_club_event(**payload)
        logger.info("Evento criado: %s", saved_id)
        return saved_id

    def sync_tournament_event(self, tournament_id: int) -> int:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Torneio nao encontrado.")
        existing = self.db.get_club_event_by_tournament(tournament_id)
        data = {
            "club_id": tournament.get("club_id") or 1,
            "tournament_id": tournament_id,
            "title": tournament["name"],
            "event_type": "tournament",
            "event_date": tournament.get("start_date") or "",
            "start_time": "",
            "end_time": "",
            "location": tournament.get("location") or "",
            "status": "confirmed" if tournament.get("status") == "running" else "planned",
            "notes": f"Torneio com {tournament.get('rounds_count') or ''} rodadas.".strip(),
        }
        return self.save_event(data, int(existing["id"]) if existing else None)

    def events_report(
        self,
        start_date: str = "",
        end_date: str = "",
        club_id: int | None = None,
    ) -> list[dict[str, Any]]:
        self._validate_optional_date(start_date, "Data inicial invalida.")
        self._validate_optional_date(end_date, "Data final invalida.")
        return self.db.list_club_events(
            start_date=start_date.strip(),
            end_date=end_date.strip(),
            club_id=club_id,
        )

    def upcoming_events(self, limit: int = 5) -> list[dict[str, Any]]:
        return self.db.list_club_events(
            start_date=date.today().isoformat(),
            include_canceled=False,
        )[:limit]

    def _validated_event_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        try:
            tournament_id = int(data.get("tournament_id") or 0)
        except ValueError as exc:
            raise AppError("Torneio vinculado invalido.") from exc
        tournament = self.db.get_tournament(tournament_id) if tournament_id else None
        if tournament_id and not tournament:
            raise AppError("Torneio vinculado nao encontrado.")

        try:
            club_id = int(data.get("club_id") or (tournament.get("club_id") if tournament else 1) or 1)
        except ValueError as exc:
            raise AppError("Clube/escola do evento invalido.") from exc
        if not self.db.get_club(club_id):
            raise AppError("Clube/escola do evento nao encontrado.")

        title = str(data.get("title", "")).strip()
        if not title and tournament:
            title = str(tournament["name"])
        if not title:
            raise AppError("Informe o titulo do evento.")

        event_type = str(data.get("event_type", "other")).strip() or "other"
        if event_type not in EVENT_TYPES:
            raise AppError("Tipo de evento invalido.")

        status = str(data.get("status", "planned")).strip() or "planned"
        if status not in EVENT_STATUSES:
            raise AppError("Status de evento invalido.")

        event_date = str(data.get("event_date", "")).strip()
        self._validate_optional_date(event_date, "Data do evento invalida.")

        return {
            "club_id": club_id,
            "tournament_id": tournament_id or None,
            "title": title,
            "event_type": event_type,
            "event_date": event_date,
            "start_time": str(data.get("start_time", "")),
            "end_time": str(data.get("end_time", "")),
            "location": str(data.get("location", "")),
            "status": status,
            "notes": str(data.get("notes", "")),
        }

    @staticmethod
    def _validate_optional_date(value: str, message: str) -> None:
        cleaned = value.strip()
        if not cleaned:
            return
        try:
            date.fromisoformat(cleaned)
        except ValueError as exc:
            raise AppError(message) from exc


class InventoryService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def save_item(self, data: dict[str, Any], item_id: int | None = None) -> int:
        payload = self._validated_item_payload(data, item_id=item_id)
        if item_id:
            if not self.db.get_inventory_item(item_id):
                raise AppError("Item de inventario nao encontrado.")
            self.db.update_inventory_item(item_id, **payload)
            logger.info("Item de inventario atualizado: %s", item_id)
            return item_id
        saved_id = self.db.create_inventory_item(**payload)
        logger.info("Item de inventario criado: %s", saved_id)
        return saved_id

    def toggle_item_active(self, item_id: int) -> int:
        item = self.db.get_inventory_item(item_id)
        if not item:
            raise AppError("Item de inventario nao encontrado.")
        next_active = 0 if item.get("active") else 1
        if not next_active and int(item.get("borrowed_quantity") or 0) > 0:
            raise AppError("Nao inative item com emprestimos em aberto.")
        self.db.set_inventory_item_active(item_id, next_active)
        logger.info("Item de inventario %s alterado para active=%s", item_id, next_active)
        return next_active

    def save_loan(self, data: dict[str, Any], loan_id: int | None = None) -> int:
        payload = self._validated_loan_payload(data, loan_id=loan_id)
        if loan_id:
            if not self.db.get_inventory_loan(loan_id):
                raise AppError("Emprestimo nao encontrado.")
            self.db.update_inventory_loan(loan_id, **payload)
            logger.info("Emprestimo atualizado: %s", loan_id)
            return loan_id
        saved_id = self.db.create_inventory_loan(**payload)
        logger.info("Emprestimo criado: %s", saved_id)
        return saved_id

    def return_loan(self, loan_id: int, return_date: str = "") -> None:
        loan = self.db.get_inventory_loan(loan_id)
        if not loan:
            raise AppError("Emprestimo nao encontrado.")
        if loan.get("status") == "returned":
            return
        payload = dict(loan)
        payload["status"] = "returned"
        payload["return_date"] = return_date.strip() or date.today().isoformat()
        self.save_loan(payload, loan_id=loan_id)

    def save_maintenance(self, data: dict[str, Any], maintenance_id: int | None = None) -> int:
        payload = self._validated_maintenance_payload(data)
        if maintenance_id:
            if not self.db.get_inventory_maintenance(maintenance_id):
                raise AppError("Manutencao nao encontrada.")
            self.db.update_inventory_maintenance(maintenance_id, **payload)
            logger.info("Manutencao atualizada: %s", maintenance_id)
            return maintenance_id
        saved_id = self.db.create_inventory_maintenance(**payload)
        logger.info("Manutencao criada: %s", saved_id)
        return saved_id

    def close_maintenance(self, maintenance_id: int, resolved_date: str = "") -> None:
        maintenance = self.db.get_inventory_maintenance(maintenance_id)
        if not maintenance:
            raise AppError("Manutencao nao encontrada.")
        payload = dict(maintenance)
        payload["status"] = "done"
        payload["resolved_date"] = resolved_date.strip() or date.today().isoformat()
        self.save_maintenance(payload, maintenance_id=maintenance_id)

    def inventory_summary(self) -> dict[str, Any]:
        items = self.db.list_inventory_items(active_only=False)
        loans = self.db.list_inventory_loans()
        maintenance_rows = self.db.list_inventory_maintenance()
        total_items = sum(int(item.get("quantity_total") or 0) for item in items if item.get("active"))
        borrowed = sum(int(item.get("borrowed_quantity") or 0) for item in items)
        available = sum(int(item.get("available_quantity") or 0) for item in items if item.get("active"))
        open_loans = [loan for loan in loans if loan.get("status") == "open"]
        today = date.today().isoformat()
        overdue_loans = [
            loan
            for loan in open_loans
            if str(loan.get("due_date") or "").strip()
            and str(loan.get("due_date") or "").strip() < today
        ]
        open_maintenance = [
            row
            for row in maintenance_rows
            if str(row.get("status") or "") in {"open", "in_progress"}
        ]
        return {
            "items": len(items),
            "total_quantity": total_items,
            "available_quantity": available,
            "borrowed_quantity": borrowed,
            "open_loans": len(open_loans),
            "overdue_loans": len(overdue_loans),
            "open_maintenance": len(open_maintenance),
        }

    def _validated_item_payload(self, data: dict[str, Any], item_id: int | None = None) -> dict[str, Any]:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome do item.")

        try:
            club_id = int(data.get("club_id") or 1)
        except ValueError as exc:
            raise AppError("Clube/escola do item invalido.") from exc
        if not self.db.get_club(club_id):
            raise AppError("Clube/escola do item nao encontrado.")

        item_type = str(data.get("item_type", "other")).strip() or "other"
        if item_type not in INVENTORY_ITEM_TYPES:
            raise AppError("Tipo de item invalido.")

        condition_status = str(data.get("condition_status", "good")).strip() or "good"
        if condition_status not in INVENTORY_CONDITIONS:
            raise AppError("Estado do item invalido.")

        try:
            quantity_total = int(data.get("quantity_total") or 1)
        except ValueError as exc:
            raise AppError("Quantidade do item invalida.") from exc
        if quantity_total <= 0:
            raise AppError("Quantidade do item deve ser maior que zero.")

        borrowed_quantity = self.db.inventory_item_open_loan_quantity(item_id) if item_id else 0
        if quantity_total < borrowed_quantity:
            raise AppError("Quantidade total nao pode ficar abaixo dos emprestimos em aberto.")

        acquisition_date = str(data.get("acquisition_date", "")).strip()
        self._validate_optional_date(acquisition_date, "Data de aquisicao invalida.")

        acquisition_value = self._parse_money(data.get("acquisition_value"), "Valor de aquisicao invalido.")
        if acquisition_value < 0:
            raise AppError("Valor de aquisicao nao pode ser negativo.")

        return {
            "club_id": club_id,
            "code": str(data.get("code", "")),
            "name": name,
            "item_type": item_type,
            "quantity_total": quantity_total,
            "condition_status": condition_status,
            "storage_location": str(data.get("storage_location", "")),
            "acquisition_date": acquisition_date,
            "acquisition_value": acquisition_value,
            "active": 1 if data.get("active", 1) else 0,
            "notes": str(data.get("notes", "")),
        }

    def _validated_loan_payload(self, data: dict[str, Any], loan_id: int | None = None) -> dict[str, Any]:
        try:
            item_id = int(data.get("item_id") or 0)
        except ValueError as exc:
            raise AppError("Item do emprestimo invalido.") from exc
        item = self.db.get_inventory_item(item_id)
        if not item:
            raise AppError("Item do emprestimo nao encontrado.")
        if not item.get("active"):
            raise AppError("Item inativo nao pode ser emprestado.")

        try:
            member_id = int(data.get("member_id") or 0)
        except ValueError as exc:
            raise AppError("Membro do emprestimo invalido.") from exc
        if not self.db.get_member(member_id):
            raise AppError("Membro do emprestimo nao encontrado.")

        try:
            quantity = int(data.get("quantity") or 1)
        except ValueError as exc:
            raise AppError("Quantidade do emprestimo invalida.") from exc
        if quantity <= 0:
            raise AppError("Quantidade do emprestimo deve ser maior que zero.")

        status = str(data.get("status", "open")).strip() or "open"
        if status not in INVENTORY_LOAN_STATUSES:
            raise AppError("Status do emprestimo invalido.")

        loan_date = str(data.get("loan_date", "")).strip() or date.today().isoformat()
        due_date = str(data.get("due_date", "")).strip()
        return_date = str(data.get("return_date", "")).strip()
        self._validate_optional_date(loan_date, "Data do emprestimo invalida.")
        self._validate_optional_date(due_date, "Data prevista de devolucao invalida.")
        self._validate_optional_date(return_date, "Data de devolucao invalida.")

        if status == "returned" and not return_date:
            return_date = date.today().isoformat()
        if status == "open":
            return_date = ""

        if status == "open":
            borrowed_other = self.db.inventory_item_open_loan_quantity(item_id, exclude_loan_id=loan_id)
            available = int(item.get("quantity_total") or 0) - borrowed_other
            if quantity > available:
                raise AppError("Quantidade indisponivel para emprestimo.")

        return {
            "item_id": item_id,
            "member_id": member_id,
            "quantity": quantity,
            "loan_date": loan_date,
            "due_date": due_date,
            "return_date": return_date,
            "status": status,
            "notes": str(data.get("notes", "")),
        }

    def _validated_maintenance_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        try:
            item_id = int(data.get("item_id") or 0)
        except ValueError as exc:
            raise AppError("Item da manutencao invalido.") from exc
        if not self.db.get_inventory_item(item_id):
            raise AppError("Item da manutencao nao encontrado.")

        description = str(data.get("description", "")).strip()
        if not description:
            raise AppError("Informe a descricao da manutencao.")

        status = str(data.get("status", "open")).strip() or "open"
        if status not in INVENTORY_MAINTENANCE_STATUSES:
            raise AppError("Status da manutencao invalido.")

        opened_date = str(data.get("opened_date", "")).strip() or date.today().isoformat()
        resolved_date = str(data.get("resolved_date", "")).strip()
        self._validate_optional_date(opened_date, "Data de abertura invalida.")
        self._validate_optional_date(resolved_date, "Data de conclusao invalida.")
        if status == "done" and not resolved_date:
            resolved_date = date.today().isoformat()
        if status in {"open", "in_progress"}:
            resolved_date = ""

        cost = self._parse_money(data.get("cost"), "Custo da manutencao invalido.")
        if cost < 0:
            raise AppError("Custo da manutencao nao pode ser negativo.")

        return {
            "item_id": item_id,
            "opened_date": opened_date,
            "resolved_date": resolved_date,
            "status": status,
            "description": description,
            "cost": cost,
            "vendor": str(data.get("vendor", "")),
            "notes": str(data.get("notes", "")),
        }

    @staticmethod
    def _parse_money(value: Any, message: str) -> float:
        try:
            return round(float(str(value or "0").replace(",", ".")), 2)
        except ValueError as exc:
            raise AppError(message) from exc

    @staticmethod
    def _validate_optional_date(value: str, message: str) -> None:
        cleaned = value.strip()
        if not cleaned:
            return
        try:
            date.fromisoformat(cleaned)
        except ValueError as exc:
            raise AppError(message) from exc


class SecurityService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def current_operator(self) -> dict[str, str]:
        settings = self.db.get_app_settings()
        role = str(settings.get("operator_role") or "admin")
        if role not in OPERATOR_ROLES:
            role = "admin"
        name = str(settings.get("operator_name") or "").strip() or OPERATOR_ROLES[role]
        return {"name": name, "role": role, "role_label": OPERATOR_ROLES[role]}

    def save_security_settings(self, data: dict[str, Any]) -> dict[str, Any]:
        operator_name = str(data.get("operator_name", "")).strip() or "Administrador"
        operator_role = str(data.get("operator_role", "admin")).strip() or "admin"
        if operator_role not in OPERATOR_ROLES:
            raise AppError("Perfil operacional invalido.")
        try:
            retention_count = int(data.get("backup_retention_count") or 10)
        except ValueError as exc:
            raise AppError("Retencao de backups invalida.") from exc
        if retention_count < 1 or retention_count > 999:
            raise AppError("Retencao de backups deve ficar entre 1 e 999.")
        payload = {
            "operator_name": operator_name,
            "operator_role": operator_role,
            "backup_retention_count": str(retention_count),
        }
        self.db.save_app_settings(payload)
        self.audit(
            "settings_saved",
            entity_type="app_settings",
            description="Configuracoes de seguranca operacional atualizadas.",
            metadata=payload,
        )
        return payload

    def audit(
        self,
        action: str,
        entity_type: str = "",
        entity_id: int | None = None,
        description: str = "",
        metadata: Mapping[str, Any] | None = None,
    ) -> int:
        operator = self.current_operator()
        metadata_json = json.dumps(metadata or {}, ensure_ascii=False, sort_keys=True)
        return self.db.create_audit_log(
            action=action,
            actor=operator["name"],
            role=operator["role"],
            entity_type=entity_type,
            entity_id=entity_id,
            description=description,
            metadata_json=metadata_json,
        )

    def list_audit_logs(
        self,
        limit: int = 200,
        action: str = "",
        entity_type: str = "",
        start_date: str = "",
        end_date: str = "",
    ) -> list[dict[str, Any]]:
        self._validate_optional_date(start_date, "Data inicial da auditoria invalida.")
        self._validate_optional_date(end_date, "Data final da auditoria invalida.")
        return self.db.list_audit_logs(
            limit=limit,
            action=action,
            entity_type=entity_type,
            start_date=start_date.strip(),
            end_date=end_date.strip(),
        )

    def create_backup(self, reason: str = "manual") -> dict[str, Any]:
        path = self.db.backup(reason)
        deleted = self.enforce_backup_retention()
        
        cloud_sync_dir = self.db.get_app_settings().get("cloud_sync_dir", "").strip()
        cloud_status = "not_configured"
        cloud_path = ""
        
        if cloud_sync_dir:
            cloud_dir_path = Path(cloud_sync_dir)
            if cloud_dir_path.exists() and cloud_dir_path.is_dir():
                try:
                    cloud_target = cloud_dir_path / path.name
                    shutil.copy2(path, cloud_target)
                    cloud_status = "success"
                    cloud_path = str(cloud_target)
                except Exception as exc:
                    logger.error("Falha ao copiar backup para nuvem %s: %s", cloud_sync_dir, exc)
                    cloud_status = f"error: {exc}"
            else:
                cloud_status = "invalid_directory"
        
        self.audit(
            "backup_created",
            entity_type="backup",
            description=f"Backup criado: {path.name}",
            metadata={
                "path": str(path),
                "reason": reason,
                "deleted_by_retention": [str(item) for item in deleted],
                "cloud_sync_status": cloud_status,
                "cloud_path": cloud_path,
            },
        )
        logger.info("Backup criado por SecurityService: %s (Cloud: %s)", path, cloud_status)
        return {"path": path, "deleted": deleted, "cloud_status": cloud_status}

    def restore_backup(self, backup_path: Path | str) -> Path:
        source = Path(backup_path)
        safety_backup = self.db.restore_backup(source)
        self.audit(
            "backup_restored",
            entity_type="backup",
            description=f"Backup restaurado: {source.name}",
            metadata={"source": str(source), "safety_backup": str(safety_backup)},
        )
        logger.info("Backup restaurado por SecurityService: %s", source)
        return safety_backup

    def enforce_backup_retention(self, keep_count: int | None = None) -> list[Path]:
        if keep_count is None:
            settings = self.db.get_app_settings()
            try:
                keep_count = int(settings.get("backup_retention_count") or 10)
            except ValueError:
                keep_count = 10
        deleted = self.db.prune_backups(keep_count)
        if deleted:
            self.audit(
                "backup_retention_applied",
                entity_type="backup",
                description=f"Retencao aplicada: {len(deleted)} backup(s) removido(s).",
                metadata={"keep_count": keep_count, "deleted": [str(item) for item in deleted]},
            )
        return deleted

    @staticmethod
    def _validate_optional_date(value: str, message: str) -> None:
        cleaned = value.strip()
        if not cleaned:
            return
        try:
            date.fromisoformat(cleaned)
        except ValueError as exc:
            raise AppError(message) from exc


class RefereeService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def list_referees(self, active_only: bool = True) -> list[dict[str, Any]]:
        return self.db.list_referees(active_only)

    def create_referee(self, data: dict[str, Any]) -> int:
        if not data.get("name"):
            raise AppError("Nome do árbitro é obrigatório.")
        return self.db.insert_referee(data)

    def update_referee(self, referee_id: int, data: dict[str, Any]) -> None:
        if not data.get("name"):
            raise AppError("Nome do árbitro é obrigatório.")
        self.db.update_referee(referee_id, data)

    def toggle_referee_active(self, referee_id: int) -> None:
        referee = self.db.get_referee(referee_id)
        if not referee:
            raise AppError("Árbitro não encontrado.")
        data = dict(referee)
        data["active"] = 0 if referee["active"] else 1
        self.db.update_referee(referee_id, data)

    def list_tournament_referees(self, tournament_id: int) -> list[dict[str, Any]]:
        return self.db.list_tournament_referees(tournament_id)

    def assign_tournament_referee(self, tournament_id: int, referee_id: int, role: str) -> None:
        self.db.assign_tournament_referee(tournament_id, referee_id, role)

    def remove_tournament_referee(self, tournament_id: int, referee_id: int) -> None:
        self.db.remove_tournament_referee(tournament_id, referee_id)



class DashboardService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def overview(self) -> dict[str, Any]:
        finance_summary = FinanceService(self.db).finance_summary()
        ranking = InternalRatingService(self.db).ranking(active_only=True)
        recent_tournaments = []
        for tournament in self.db.list_tournaments()[:5]:
            tournament_id = int(tournament["id"])
            recent_tournaments.append(
                {
                    **tournament,
                    "players_count": len(self.db.list_players(tournament_id, active_only=False)),
                    "rounds_count_generated": len(self.db.list_rounds(tournament_id)),
                }
            )
        return {
            "summary": self.db.club_summary(),
            "upcoming_events": EventService(self.db).upcoming_events(limit=5),
            "finance_summary": finance_summary,
            "ranking_leaders": ranking[:5],
            "recent_tournaments": recent_tournaments,
            "recent_sessions": self.db.list_training_sessions()[:5],
            "backups_count": len(self.db.list_backups()),
        }


class TournamentService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def create_tournament(self, tournament_data: dict[str, Any]) -> int:
        payload = self._validated_tournament_payload(tournament_data, {})
        tournament_id = self.db.create_tournament(**payload)
        logger.info("Torneio criado: %s", tournament_id)
        return tournament_id

    def save_profile(
        self,
        tournament_id: int,
        tournament_data: dict[str, Any],
        settings_data: dict[str, Any],
        schedule: list[dict[str, Any]],
    ) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")

        tournament_payload = self._validated_tournament_payload(tournament_data, tournament)
        settings_payload = self._validated_settings(settings_data)
        schedule_payload = self._validated_schedule(schedule, tournament_payload["rounds_count"])
        self._validate_team_settings_compatibility(tournament_id, tournament_payload, settings_payload)

        self.db.update_tournament_details(
            tournament_id=tournament_id,
            **tournament_payload,
        )
        self.db.save_tournament_settings(tournament_id, settings_payload)
        self.db.save_round_schedule(tournament_id, schedule_payload)
        logger.info("Configuracoes do torneio %s atualizadas", tournament_id)

    def _validate_team_settings_compatibility(
        self,
        tournament_id: int,
        tournament_payload: dict[str, Any],
        settings_payload: dict[str, Any],
    ) -> None:
        if tournament_payload.get("competition_type") != "team":
            return
        boards_count = int(settings_payload.get("team_boards_count") or 4)
        highest_board = 0
        for team in self.db.list_teams(tournament_id, active_only=False):
            for assignment in self.db.list_team_players(int(team["id"]), active_only=False):
                if assignment.get("board_number"):
                    highest_board = max(highest_board, int(assignment["board_number"]))
        if highest_board > boards_count:
            raise AppError(
                f"Ja existe jogador escalado no tabuleiro {highest_board}. "
                f"Aumente ou mantenha a quantidade de tabuleiros por equipe."
            )

    def _validated_tournament_payload(
        self,
        data: dict[str, Any],
        tournament: dict[str, Any],
    ) -> dict[str, Any]:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome do torneio.")

        try:
            rounds_count = int(data.get("rounds_count") or 1)
        except ValueError as exc:
            raise AppError("Numero de rodadas invalido.") from exc
        if rounds_count < 1:
            raise AppError("O torneio precisa ter pelo menos uma rodada.")

        try:
            bye_points = float(str(data.get("bye_points") or "0").replace(",", "."))
        except ValueError as exc:
            raise AppError("Pontuacao do bye invalida.") from exc
        if bye_points < 0:
            raise AppError("Pontuacao do bye nao pode ser negativa.")

        start_date = self._normalize_optional_date(data.get("start_date", ""), "Data inicial invalida.")
        end_date = self._normalize_optional_date(data.get("end_date", ""), "Data final invalida.")
        if start_date and end_date and start_date > end_date:
            raise AppError("Data final nao pode ser anterior a data inicial.")

        competition_type = str(
            data.get("competition_type") or tournament.get("competition_type") or "individual"
        ).strip() or "individual"
        if competition_type not in COMPETITION_TYPES:
            raise AppError("Formato do torneio invalido.")

        _scope, club_id, class_id = self._validated_tournament_scope(data, tournament)
        return {
            "name": name,
            "club_id": club_id,
            "class_id": class_id,
            "competition_type": competition_type,
            "location": str(data.get("location", "")),
            "rounds_count": rounds_count,
            "time_control": str(data.get("time_control", "")),
            "start_date": start_date,
            "end_date": end_date,
            "bye_points": bye_points,
        }

    def _validated_tournament_scope(
        self,
        data: dict[str, Any],
        tournament: dict[str, Any],
    ) -> tuple[str, int | None, int | None]:
        scope = str(data.get("scope") or "").strip()
        if not scope:
            if data.get("class_id") or tournament.get("class_id"):
                scope = "class"
            elif "club_id" in data:
                scope = "club" if data.get("club_id") else "standalone"
            elif tournament.get("club_id"):
                scope = "club"
            else:
                scope = "standalone"
        if scope not in TOURNAMENT_SCOPES:
            raise AppError("Escopo do torneio invalido.")

        if scope == "standalone":
            return scope, None, None

        class_id: int | None = None
        class_data: dict[str, Any] | None = None
        try:
            if scope == "class":
                class_id = int(data.get("class_id") or tournament.get("class_id") or 0)
        except ValueError as exc:
            raise AppError("Turma do torneio invalida.") from exc
        if scope == "class":
            class_data = self.db.get_class(class_id or 0)
            if not class_data:
                raise AppError("Turma do torneio nao encontrada.")

        if data.get("club_id"):
            club_source = data.get("club_id")
        elif class_data:
            club_source = class_data["club_id"]
        else:
            club_source = tournament.get("club_id")
        try:
            club_id = int(club_source or 0)
        except ValueError as exc:
            raise AppError("Clube/escola do torneio invalido.") from exc
        if not club_id or not self.db.get_club(club_id):
            raise AppError("Clube/escola do torneio nao encontrado.")

        if scope == "club":
            return scope, club_id, None

        if class_data and int(class_data["club_id"]) != club_id:
            raise AppError("A turma selecionada nao pertence ao clube/escola do torneio.")
        return scope, club_id, class_id

    def _validated_settings(self, data: dict[str, Any]) -> dict[str, Any]:
        initial_order = str(data.get("initial_order", "rating")).strip() or "rating"
        if initial_order not in INITIAL_ORDER_OPTIONS:
            raise AppError("Ordem inicial invalida.")

        tournament_type = str(data.get("tournament_type", "real")).strip() or "real"
        if tournament_type not in TOURNAMENT_TYPES:
            raise AppError("Tipo de torneio invalido.")

        tournament_profile = str(data.get("tournament_profile", "free")).strip() or "free"
        if tournament_profile not in TOURNAMENT_PROFILES:
            raise AppError("Perfil do torneio invalido.")

        try:
            late_entry_points = float(str(data.get("late_entry_points") or "0").replace(",", "."))
        except ValueError as exc:
            raise AppError("Pontos por adesao tardia invalidos.") from exc
        try:
            team_boards_count = int(data.get("team_boards_count") or 4)
        except ValueError as exc:
            raise AppError("Quantidade de tabuleiros por equipe invalida.") from exc
        if team_boards_count < 1:
            raise AppError("Use pelo menos um tabuleiro por equipe.")

        team_point_fields = {
            "team_match_win_points": "Pontos por vitoria da equipe invalidos.",
            "team_match_draw_points": "Pontos por empate da equipe invalidos.",
            "team_match_loss_points": "Pontos por derrota da equipe invalidos.",
        }
        team_point_defaults = {
            "team_match_win_points": "2",
            "team_match_draw_points": "1",
            "team_match_loss_points": "0",
        }
        team_points: dict[str, float] = {}
        for field, message in team_point_fields.items():
            raw_value = data.get(field)
            if raw_value in (None, ""):
                raw_value = team_point_defaults[field]
            try:
                value = float(str(raw_value).replace(",", "."))
            except ValueError as exc:
                raise AppError(message) from exc
            if value < 0:
                raise AppError(message)
            team_points[field] = value

        team_pairing_method = str(data.get("team_pairing_method", "swiss")).strip() or "swiss"
        if team_pairing_method not in TEAM_PAIRING_METHODS:
            raise AppError("Metodo de emparceiramento por equipes invalido.")
        team_standing_primary = str(data.get("team_standing_primary", "match_points")).strip() or "match_points"
        team_standing_secondary = str(data.get("team_standing_secondary", "game_points")).strip() or "game_points"
        if team_standing_primary not in TEAM_STANDING_CRITERIA:
            raise AppError("Criterio principal por equipes invalido.")
        if team_standing_secondary not in TEAM_STANDING_CRITERIA:
            raise AppError("Criterio secundario por equipes invalido.")
        if team_standing_primary == team_standing_secondary:
            raise AppError("Use criterios diferentes para classificacao por equipes.")

        payload = {
            "fide_event_id": str(data.get("fide_event_id", "")),
            "organizer": str(data.get("organizer", "")),
            "website": str(data.get("website", "")),
            "contact_email": str(data.get("contact_email", "")),
            "director": str(data.get("director", "")),
            "chief_arbiter": str(data.get("chief_arbiter", "")),
            "arbiters": str(data.get("arbiters", "")),
            "federation": str(data.get("federation", "")),
            "state": str(data.get("state", "")),
            "categories": str(data.get("categories", "")),
            "cutoff_date": str(data.get("cutoff_date", "")),
            "comments": str(data.get("comments", "")),
            "prizes": str(data.get("prizes", "")),
            "initial_order": initial_order,
            "tournament_type": tournament_type,
            "tournament_profile": tournament_profile,
            "late_entry_points": late_entry_points,
            "team_boards_count": team_boards_count,
            "team_match_win_points": team_points["team_match_win_points"],
            "team_match_draw_points": team_points["team_match_draw_points"],
            "team_match_loss_points": team_points["team_match_loss_points"],
            "team_pairing_method": team_pairing_method,
            "team_standing_primary": team_standing_primary,
            "team_standing_secondary": team_standing_secondary,
            "team_fixed_board_order": 1 if data.get("team_fixed_board_order", 1) else 0,
        }
        for field in TOURNAMENT_FLAG_FIELDS:
            payload[field] = 1 if data.get(field) else 0
        return payload

    @staticmethod
    def _validate_optional_date(value: str, message: str) -> None:
        TournamentService._normalize_optional_date(value, message)

    @staticmethod
    def _normalize_optional_date(value: Any, message: str) -> str:
        cleaned = str(value or "").strip()
        if not cleaned:
            return ""
        try:
            return date.fromisoformat(cleaned).isoformat()
        except ValueError:
            pass

        for separator in ("/", "-"):
            parts = cleaned.split(separator)
            if len(parts) != 3:
                continue
            day, month, year = (part.strip() for part in parts)
            if len(year) != 4:
                continue
            try:
                return date(int(year), int(month), int(day)).isoformat()
            except ValueError:
                continue

        try:
            raise ValueError(cleaned)
        except ValueError as exc:
            raise AppError(message) from exc

    @staticmethod
    def _validated_schedule(
        schedule: list[dict[str, Any]],
        rounds_count: int,
    ) -> list[dict[str, Any]]:
        rows = []
        seen: set[int] = set()
        for item in schedule:
            try:
                round_number = int(item.get("round_number") or 0)
            except ValueError as exc:
                raise AppError("Numero de rodada invalido na agenda.") from exc
            if round_number < 1 or round_number > rounds_count:
                raise AppError("Agenda contem rodada fora do total configurado.")
            if round_number in seen:
                raise AppError("Agenda contem rodada repetida.")
            round_date = str(item.get("date", "")).strip()
            if round_date:
                try:
                    date.fromisoformat(round_date)
                except ValueError as exc:
                    raise AppError("Data invalida na agenda de rodadas.") from exc
            seen.add(round_number)
            rows.append(
                {
                    "round_number": round_number,
                    "date": round_date,
                    "time": str(item.get("time", "")).strip(),
                }
            )

        for round_number in range(1, rounds_count + 1):
            if round_number not in seen:
                rows.append({"round_number": round_number, "date": "", "time": ""})

        return sorted(rows, key=lambda item: item["round_number"])

    @staticmethod
    def generate_round_schedule(
        rounds_count: Any,
        start_date: Any,
        first_time: Any,
        round_duration_minutes: Any,
        break_minutes: Any,
        rounds_per_day: Any,
    ) -> list[dict[str, Any]]:
        try:
            total_rounds = int(rounds_count)
        except (TypeError, ValueError) as exc:
            raise AppError("Numero de rodadas invalido para gerar agenda.") from exc
        if total_rounds < 1:
            raise AppError("O torneio precisa ter pelo menos uma rodada para gerar agenda.")

        start_date_text = TournamentService._normalize_optional_date(
            start_date,
            "Data inicial da agenda invalida.",
        )
        if not start_date_text:
            raise AppError("Informe a data inicial da agenda.")
        start_day = date.fromisoformat(start_date_text)

        try:
            start_time = time.fromisoformat(str(first_time or "").strip())
        except ValueError as exc:
            raise AppError("Hora inicial da agenda invalida. Use HH:MM.") from exc

        try:
            duration_minutes = int(round_duration_minutes)
        except (TypeError, ValueError) as exc:
            raise AppError("Duracao da rodada invalida.") from exc
        if duration_minutes < 1:
            raise AppError("Duracao da rodada deve ser maior que zero.")

        try:
            interval_minutes = int(break_minutes)
        except (TypeError, ValueError) as exc:
            raise AppError("Intervalo entre rodadas invalido.") from exc
        if interval_minutes < 0:
            raise AppError("Intervalo entre rodadas nao pode ser negativo.")

        try:
            per_day = int(rounds_per_day)
        except (TypeError, ValueError) as exc:
            raise AppError("Rodadas por dia invalidas.") from exc
        if per_day < 1 or per_day > total_rounds:
            per_day = total_rounds

        step = timedelta(minutes=duration_minutes + interval_minutes)
        rows = []
        for index in range(total_rounds):
            day_offset = index // per_day
            slot = index % per_day
            day_start = datetime.combine(start_day + timedelta(days=day_offset), start_time)
            scheduled_at = day_start + (step * slot)
            rows.append(
                {
                    "round_number": index + 1,
                    "date": scheduled_at.date().isoformat(),
                    "time": scheduled_at.strftime("%H:%M"),
                }
            )
        return rows


class TeamService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def create_team(self, tournament_id: int, data: dict[str, Any]) -> int:
        self._require_team_tournament(tournament_id)
        payload = self._validated_team_payload(data)
        try:
            team_id = self.db.create_team(tournament_id, **payload)
        except sqlite3.IntegrityError as exc:
            raise AppError("Ja existe uma equipe com este nome no torneio.") from exc
        logger.info("Equipe criada no torneio %s: %s", tournament_id, team_id)
        return team_id

    def update_team(self, team_id: int, data: dict[str, Any]) -> None:
        team = self.db.get_team(team_id)
        if not team:
            raise AppError("Equipe nao encontrada.")
        self._require_team_tournament(int(team["tournament_id"]))
        payload = self._validated_team_payload(data)
        try:
            self.db.update_team(team_id, **payload)
        except sqlite3.IntegrityError as exc:
            raise AppError("Ja existe uma equipe com este nome no torneio.") from exc
        logger.info("Equipe atualizada: %s", team_id)

    def add_player(
        self,
        team_id: int,
        player_id: int,
        board_number: Any = None,
        role: str = "starter",
    ) -> int:
        team = self.db.get_team(team_id)
        if not team:
            raise AppError("Equipe nao encontrada.")
        self._require_team_tournament(int(team["tournament_id"]))

        player = self.db.get_player(player_id)
        if not player or int(player["tournament_id"]) != int(team["tournament_id"]):
            raise AppError("Jogador nao pertence ao torneio da equipe.")

        existing = self.db.get_team_player_by_player(player_id)
        if existing:
            raise AppError(f"Jogador ja esta na equipe {existing['team_name']}.")

        payload = self._validated_assignment_payload(
            int(team["tournament_id"]),
            board_number,
            role,
        )
        self._ensure_board_available(team_id, payload["board_number"])
        try:
            team_player_id = self.db.add_player_to_team(team_id, player_id, **payload)
        except sqlite3.IntegrityError as exc:
            raise AppError("Jogador ou tabuleiro ja cadastrado em uma equipe.") from exc
        logger.info("Jogador %s adicionado a equipe %s", player_id, team_id)
        return team_player_id

    def update_player_assignment(
        self,
        team_player_id: int,
        board_number: Any = None,
        role: str = "starter",
    ) -> None:
        assignment = self.db.get_team_player(team_player_id)
        if not assignment:
            raise AppError("Jogador da equipe nao encontrado.")
        self._require_team_tournament(int(assignment["tournament_id"]))
        payload = self._validated_assignment_payload(
            int(assignment["tournament_id"]),
            board_number,
            role,
        )
        self._ensure_board_available(
            int(assignment["team_id"]),
            payload["board_number"],
            ignore_team_player_id=team_player_id,
        )
        try:
            self.db.update_team_player(team_player_id, **payload)
        except sqlite3.IntegrityError as exc:
            raise AppError("Tabuleiro ja ocupado nesta equipe.") from exc
        logger.info("Escalacao da equipe atualizada: %s", team_player_id)

    def remove_player(self, team_player_id: int) -> None:
        if not self.db.get_team_player(team_player_id):
            raise AppError("Jogador da equipe nao encontrado.")
        self.db.remove_player_from_team(team_player_id)
        logger.info("Jogador removido da equipe: %s", team_player_id)

    def delete_team(self, team_id: int) -> None:
        team = self.db.get_team(team_id)
        if not team:
            raise AppError("Equipe nao encontrada.")
        if self.db.count_team_matches(team_id) > 0:
            raise AppError(
                "Equipe ja apareceu em rodadas. Para preservar o historico, desmarque Ativa e atualize a equipe."
            )
        self.db.delete_team(team_id)
        logger.info("Equipe excluida: %s", team_id)

    def _require_team_tournament(self, tournament_id: int) -> dict[str, Any]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") != "team":
            raise AppError("Esta tela exige um torneio no formato Equipes.")
        return tournament

    @staticmethod
    def _validated_team_payload(data: dict[str, Any]) -> dict[str, Any]:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome da equipe.")
        return {
            "name": name,
            "club": str(data.get("club", "")),
            "captain": str(data.get("captain", "")),
            "notes": str(data.get("notes", "")),
            "active": 1 if data.get("active", 1) else 0,
        }

    def _validated_assignment_payload(
        self,
        tournament_id: int,
        board_number: Any,
        role: str,
    ) -> dict[str, Any]:
        role_value = str(role or "starter").strip() or "starter"
        if role_value not in TEAM_PLAYER_ROLES:
            raise AppError("Funcao do jogador na equipe invalida.")

        board_value: int | None = None
        if role_value == "starter":
            try:
                board_value = int(board_number or 0)
            except ValueError as exc:
                raise AppError("Numero do tabuleiro invalido.") from exc
            if board_value < 1:
                raise AppError("Informe o tabuleiro do titular.")
            settings = self.db.get_tournament_settings(tournament_id) or {}
            boards_count = int(settings.get("team_boards_count") or 4)
            if board_value > boards_count:
                raise AppError(f"O torneio esta configurado para {boards_count} tabuleiros por equipe.")

        return {
            "board_number": board_value,
            "role": role_value,
            "active": 1,
        }

    def _ensure_board_available(
        self,
        team_id: int,
        board_number: int | None,
        ignore_team_player_id: int | None = None,
    ) -> None:
        if board_number is None:
            return
        for assignment in self.db.list_team_players(team_id, active_only=False):
            if ignore_team_player_id and int(assignment["id"]) == ignore_team_player_id:
                continue
            if assignment.get("board_number") and int(assignment["board_number"]) == board_number:
                raise AppError("Ja existe jogador neste tabuleiro da equipe.")


class LearningLevelService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def create_level(self, data: dict[str, Any]) -> int:
        payload = self._validated_payload(data)
        try:
            level_id = self.db.create_learning_level(**payload)
        except sqlite3.IntegrityError as exc:
            raise AppError("Ja existe um nivel com este nome.") from exc
        logger.info("Nivel de aprendizagem criado: %s", level_id)
        return level_id

    def update_level(self, learning_level_id: int, data: dict[str, Any]) -> None:
        if not self.db.get_learning_level(learning_level_id):
            raise AppError("Nivel de aprendizagem nao encontrado.")
        payload = self._validated_payload(data)
        try:
            self.db.update_learning_level(learning_level_id, **payload)
        except sqlite3.IntegrityError as exc:
            raise AppError("Ja existe um nivel com este nome.") from exc
        logger.info("Nivel de aprendizagem atualizado: %s", learning_level_id)

    def toggle_level_active(self, learning_level_id: int) -> int:
        level = self.db.get_learning_level(learning_level_id)
        if not level:
            raise AppError("Nivel de aprendizagem nao encontrado.")
        next_active = 0 if level.get("active") else 1
        self.db.set_learning_level_active(learning_level_id, next_active)
        logger.info("Nivel de aprendizagem %s ativo=%s", learning_level_id, next_active)
        return next_active

    @staticmethod
    def _validated_payload(data: dict[str, Any]) -> dict[str, Any]:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome do nivel.")
        try:
            display_order = int(data.get("display_order") or 0)
        except ValueError as exc:
            raise AppError("Ordem do nivel invalida.") from exc
        if display_order < 0:
            raise AppError("A ordem do nivel nao pode ser negativa.")
        return {
            "name": name,
            "description": str(data.get("description", "")),
            "display_order": display_order,
            "active": 1 if data.get("active", 1) else 0,
        }


class MemberService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def create_member(self, data: dict[str, Any]) -> int:
        payload = self._validated_payload(data)
        class_id = payload.pop("class_id", None)
        member_id = self.db.create_member(**payload)
        if class_id:
            self.db.set_member_active_class(member_id, class_id)
        logger.info("Membro criado: %s", member_id)
        return member_id

    def update_member(self, member_id: int, data: dict[str, Any]) -> None:
        if not self.db.get_member(member_id):
            raise AppError("Membro nao encontrado.")
        payload = self._validated_payload(data)
        class_id = payload.pop("class_id", None)
        self.db.update_member(member_id, **payload)
        self.db.set_member_active_class(member_id, class_id)
        logger.info("Membro atualizado: %s", member_id)

    def toggle_member_status(self, member_id: int) -> str:
        member = self.db.get_member(member_id)
        if not member:
            raise AppError("Membro nao encontrado.")
        next_status = "inactive" if member["status"] == "active" else "active"
        self.db.set_member_status(member_id, next_status)
        logger.info("Status do membro %s alterado para %s", member_id, next_status)
        return next_status

    def register_member_in_tournament(
        self,
        tournament_id: int,
        member_id: int,
        allow_out_of_scope: bool = False,
    ) -> int:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")

        member = self.db.get_member(member_id)
        if not member:
            raise AppError("Membro nao encontrado.")
        if member["status"] != "active":
            raise AppError("Somente membros ativos podem ser inscritos.")

        if not allow_out_of_scope:
            eligible_ids = {
                int(item["id"])
                for item in self.db.list_members_for_tournament(tournament_id, active_only=True)
            }
            if member_id not in eligible_ids:
                scope = self._tournament_scope_label(tournament)
                raise AppError(f"Membro fora do escopo do torneio ({scope}).")

        existing = self.db.get_player_by_member(tournament_id, member_id)
        if existing:
            raise AppError("Este membro ja esta inscrito no torneio.")

        player_id = self.db.create_player_from_member(tournament_id, member_id)
        logger.info(
            "Membro %s inscrito no torneio %s como jogador %s",
            member_id,
            tournament_id,
            player_id,
        )
        return player_id

    @staticmethod
    def _tournament_scope_label(tournament: dict[str, Any]) -> str:
        if tournament.get("class_id"):
            return TOURNAMENT_SCOPES["class"]
        if tournament.get("club_id"):
            return TOURNAMENT_SCOPES["club"]
        return TOURNAMENT_SCOPES["standalone"]

    def register_active_members_in_tournament(
        self,
        tournament_id: int,
        include_out_of_scope: bool = False,
    ) -> dict[str, Any]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")

        registered_player_ids = []
        skipped = 0
        for member in self.db.list_members_for_tournament(
            tournament_id,
            active_only=True,
            include_out_of_scope=include_out_of_scope,
        ):
            if member["registered_player_id"]:
                skipped += 1
                continue
            registered_player_ids.append(
                self.db.create_player_from_member(tournament_id, int(member["id"]))
            )

        logger.info(
            "%s membros ativos inscritos no torneio %s; %s ja estavam inscritos",
            len(registered_player_ids),
            tournament_id,
            skipped,
        )
        return {
            "registered": len(registered_player_ids),
            "skipped": skipped,
            "player_ids": registered_player_ids,
        }

    def tournament_history(self, member_id: int) -> list[dict[str, Any]]:
        if not self.db.get_member(member_id):
            raise AppError("Membro nao encontrado.")

        pairing_service = PairingService(self.db)
        history = []
        for registration in self.db.list_member_tournament_players(member_id):
            tournament_id = int(registration["tournament_id"])
            standing = next(
                (
                    item
                    for item in pairing_service.standings(tournament_id)
                    if int(item["player_id"]) == int(registration["id"])
                ),
                None,
            )
            results = self.tournament_results(member_id, tournament_id)
            closed_results = [
                result
                for result in results
                if result["round_status"] == "closed" and result["outcome"] != "Pendente"
            ]

            history.append(
                {
                    "tournament_id": tournament_id,
                    "player_id": registration["id"],
                    "name": registration["tournament_name"],
                    "status": registration["tournament_status"],
                    "start_date": registration["tournament_start_date"],
                    "player_rating": int(registration.get("rating") or 0),
                    "player_status": registration.get("player_status", ""),
                    "points": standing["points"] if standing else 0.0,
                    "position": standing["position"] if standing else "",
                    "performance": standing["performance"] if standing else "",
                    "rounds_played": len(closed_results),
                    "wins": self._count_outcomes(closed_results, "Vitoria"),
                    "draws": self._count_outcomes(closed_results, "Empate"),
                    "losses": self._count_outcomes(closed_results, "Derrota"),
                    "byes": self._count_outcomes(closed_results, "Bye"),
                    "last_result": closed_results[-1]["outcome"] if closed_results else "",
                }
            )
        return history

    def tournament_results(self, member_id: int, tournament_id: int) -> list[dict[str, Any]]:
        if not self.db.get_member(member_id):
            raise AppError("Membro nao encontrado.")
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        player = self.db.get_player_by_member(tournament_id, member_id)
        if not player:
            return []

        rows = self.db.list_member_tournament_results(member_id, tournament_id)
        if not rows:
            return []

        member_player_id = int(player["id"])
        bye_points = float(tournament["bye_points"])
        results = []
        for row in rows:
            is_white = int(row["white_player_id"]) == member_player_id
            is_bye = bool(row["is_bye"])
            opponent = "BYE"
            color = "Bye"
            points: float | None = None

            if is_bye:
                points = bye_points
            elif is_white:
                color = "Brancas"
                opponent = player_full_name(
                    {
                        "name": row.get("black_name"),
                        "surname": row.get("black_surname"),
                        "given_name": row.get("black_given_name"),
                    }
                )
                points = self._points_for_result(row["result"], white=True)
            else:
                color = "Pretas"
                opponent = player_full_name(
                    {
                        "name": row.get("white_name"),
                        "surname": row.get("white_surname"),
                        "given_name": row.get("white_given_name"),
                    }
                )
                points = self._points_for_result(row["result"], white=False)

            results.append(
                {
                    "round_number": row["round_number"],
                    "round_status": row["round_status"],
                    "board_number": row["board_number"],
                    "color": color,
                    "opponent": opponent,
                    "result": "BYE" if is_bye else row["result"],
                    "points": points,
                    "outcome": self._outcome_label(row["result"], points, is_bye),
                }
            )
        return results

    @staticmethod
    def _points_for_result(result: str, white: bool) -> float | None:
        if result not in RESULT_POINTS:
            return None
        white_points, black_points = RESULT_POINTS[result]
        return white_points if white else black_points

    @staticmethod
    def _outcome_label(result: str, points: float | None, is_bye: bool) -> str:
        if is_bye:
            return "Bye"
        if result == "0F-0F":
            return "Duplo WO"
        if points is None:
            return "Pendente"
        if points == 1.0:
            return "Vitoria"
        if points == 0.5:
            return "Empate"
        return "Derrota"

    @staticmethod
    def _count_outcomes(results: list[dict[str, Any]], outcome: str) -> int:
        return sum(1 for result in results if result["outcome"] == outcome)

    def _validated_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome do membro.")

        member_type = str(data.get("member_type", "socio")).strip() or "socio"
        if member_type not in MEMBER_TYPES:
            raise AppError("Tipo de vinculo invalido.")

        status = str(data.get("status", "active")).strip() or "active"
        if status not in MEMBER_STATUSES:
            raise AppError("Status de membro invalido.")

        try:
            rating = int(data.get("rating") or 0)
        except ValueError as exc:
            raise AppError("Rating do membro invalido.") from exc
        try:
            club_id = int(data.get("club_id") or 1)
        except ValueError as exc:
            raise AppError("Clube/escola do membro invalido.") from exc
        if not self.db.get_club(club_id):
            raise AppError("Clube/escola do membro nao encontrado.")
        try:
            class_id = int(data.get("class_id") or 0)
        except ValueError as exc:
            raise AppError("Turma do membro invalida.") from exc
        if class_id:
            class_data = self.db.get_class(class_id)
            if not class_data or int(class_data["club_id"]) != club_id:
                raise AppError("Turma nao pertence ao clube/escola selecionado.")
        try:
            learning_level_id = int(data.get("learning_level_id") or 0)
        except ValueError as exc:
            raise AppError("Nivel de aprendizagem invalido.") from exc
        if learning_level_id and not self.db.get_learning_level(learning_level_id):
            raise AppError("Nivel de aprendizagem nao encontrado.")

        return {
            "name": name,
            "surname": str(data.get("surname", "")),
            "club_id": club_id,
            "learning_level_id": learning_level_id or None,
            "city": str(data.get("city", "")),
            "phone": str(data.get("phone", "")),
            "email": str(data.get("email", "")),
            "document": str(data.get("document", "")),
            "birth_date": str(data.get("birth_date", "")),
            "rating": rating,
            "category": str(data.get("category", "")),
            "member_type": member_type,
            "status": status,
            "guardian_name": str(data.get("guardian_name", "")),
            "guardian_phone": str(data.get("guardian_phone", "")),
            "notes": str(data.get("notes", "")),
            "class_id": class_id or None,
        }


class PairingService:
    MAX_EXHAUSTIVE_PAIRING_PLAYERS = 16
    MAX_EXHAUSTIVE_PAIRING_TEAMS = 16
    REPEAT_PAIRING_PENALTY = 1_000_000
    SCORE_GROUP_FLOAT_PENALTY = 10_000
    SCORE_DIFF_PENALTY = 1_000

    def __init__(self, db: Database) -> None:
        self.db = db

    def generate_next_round(self, tournament_id: int) -> dict[str, Any]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") == "team":
            return self._generate_next_team_round(tournament_id, tournament)

        players = self.db.list_players(tournament_id, active_only=True)
        if len(players) < 2:
            raise AppError("Cadastre pelo menos 2 jogadores ativos.")
        settings = self.db.get_tournament_settings(tournament_id) or {}
        if settings.get("disable_bye") and len(players) % 2 == 1:
            raise AppError("O bye esta desativado. Use numero par de jogadores ativos.")

        latest_round = self.db.get_latest_round(tournament_id)
        if latest_round and latest_round["status"] != "closed":
            raise AppError("Feche ou exclua a rodada gerada antes de criar outra.")

        next_number = 1 if not latest_round else int(latest_round["number"]) + 1
        if next_number > int(tournament["rounds_count"]):
            raise AppError("O numero maximo de rodadas do torneio ja foi atingido.")

        pairing_method = settings.get("pairing_method", "swiss")
        if pairing_method == "round_robin":
            pairings = self._round_robin_pairings(tournament_id, players, next_number, settings)
        elif pairing_method == "knockout":
            pairings = self._knockout_pairings(tournament_id, players, next_number, settings)
        else:
            if next_number == 1:
                pairings = self._first_round_pairings(players)
            else:
                pairings = self._swiss_pairings(tournament_id, players)

        round_id = self.db.create_round_with_pairings(
            tournament_id,
            next_number,
            pairings,
        )
        logger.info("Rodada %s gerada para o torneio %s", next_number, tournament_id)
        if tournament["status"] == "draft":
            self.db.update_tournament_status(tournament_id, "running")

        generated = self.db.get_round_by_number(tournament_id, next_number)
        if not generated:
            raise AppError("A rodada foi gerada, mas nao pode ser reaberta.")
        generated["id"] = round_id
        return generated

    def _generate_next_team_round(
        self,
        tournament_id: int,
        tournament: dict[str, Any],
    ) -> dict[str, Any]:
        settings = self.db.get_tournament_settings(tournament_id) or {}
        teams = self.db.list_teams(tournament_id, active_only=True)
        if len(teams) < 2:
            raise AppError("Cadastre pelo menos 2 equipes ativas.")
        if settings.get("disable_bye") and len(teams) % 2 == 1:
            raise AppError("O bye esta desativado. Use numero par de equipes ativas.")

        boards_count = int(settings.get("team_boards_count") or 4)
        rosters, seed_ratings = self._team_starter_rosters(teams, boards_count)

        latest_round = self.db.get_latest_round(tournament_id)
        if latest_round and latest_round["status"] != "closed":
            raise AppError("Feche ou exclua a rodada gerada antes de criar outra.")

        next_number = 1 if not latest_round else int(latest_round["number"]) + 1
        if next_number > int(tournament["rounds_count"]):
            raise AppError("O numero maximo de rodadas do torneio ja foi atingido.")

        if next_number == 1:
            matches = self._first_round_team_matches(teams, rosters, seed_ratings, boards_count, settings)
        else:
            matches = self._swiss_team_matches(tournament_id, teams, rosters, seed_ratings, boards_count, settings)

        round_id = self.db.create_round_with_team_matches(tournament_id, next_number, matches)
        logger.info("Rodada por equipes %s gerada para o torneio %s", next_number, tournament_id)
        if tournament["status"] == "draft":
            self.db.update_tournament_status(tournament_id, "running")

        generated = self.db.get_round_by_number(tournament_id, next_number)
        if not generated:
            raise AppError("A rodada foi gerada, mas nao pode ser reaberta.")
        generated["id"] = round_id
        return generated

    def _team_starter_rosters(
        self,
        teams: list[dict[str, Any]],
        boards_count: int,
    ) -> tuple[dict[int, dict[int, int]], dict[int, int]]:
        rosters: dict[int, dict[int, int]] = {}
        seed_ratings: dict[int, int] = {}
        for team in teams:
            team_id = int(team["id"])
            starters: dict[int, int] = {}
            ratings: dict[int, int] = {}
            for assignment in self.db.list_team_players(team_id, active_only=True):
                if assignment.get("role") != "starter" or not assignment.get("board_number"):
                    continue
                if assignment.get("player_status") != "active":
                    continue
                board_number = int(assignment["board_number"])
                if 1 <= board_number <= boards_count:
                    starters[board_number] = int(assignment["player_id"])
                    ratings[board_number] = int(assignment.get("player_rating") or 0)

            missing = [board for board in range(1, boards_count + 1) if board not in starters]
            if missing:
                missing_text = ", ".join(str(board) for board in missing)
                raise AppError(f"Equipe {team['name']} sem titular ativo no tabuleiro {missing_text}.")

            rosters[team_id] = starters
            seed_ratings[team_id] = round(sum(ratings.values()) / max(boards_count, 1))
        return rosters, seed_ratings

    def _first_round_team_matches(
        self,
        teams: list[dict[str, Any]],
        rosters: dict[int, dict[int, int]],
        seed_ratings: dict[int, int],
        boards_count: int,
        settings: dict[str, Any],
    ) -> list[dict[str, Any]]:
        ordered = sorted(
            teams,
            key=lambda team: (-seed_ratings[int(team["id"])], str(team["name"]).casefold()),
        )
        pairable = ordered[:]
        bye_team = None
        if len(pairable) % 2 == 1:
            bye_team = min(
                pairable,
                key=lambda team: (seed_ratings[int(team["id"])], str(team["name"]).casefold()),
            )
            pairable.remove(bye_team)

        half = len(pairable) // 2
        upper = pairable[:half]
        lower = pairable[half:]
        matches: list[dict[str, Any]] = []
        for index, team in enumerate(upper, start=1):
            opponent = lower[index - 1]
            if index % 2 == 1:
                white_team_id = int(team["id"])
                black_team_id = int(opponent["id"])
            else:
                white_team_id = int(opponent["id"])
                black_team_id = int(team["id"])
            matches.append(
                self._team_match_payload(
                    index,
                    white_team_id,
                    black_team_id,
                    rosters,
                    boards_count,
                )
            )

        if bye_team:
            matches.append(self._team_bye_payload(len(matches) + 1, int(bye_team["id"]), settings, boards_count))
        return matches

    def _swiss_team_matches(
        self,
        tournament_id: int,
        teams: list[dict[str, Any]],
        rosters: dict[int, dict[int, int]],
        seed_ratings: dict[int, int],
        boards_count: int,
        settings: dict[str, Any],
    ) -> list[dict[str, Any]]:
        standings = {int(item["team_id"]): item for item in self.team_standings(tournament_id)}
        played_pairs = self._team_played_pairs(tournament_id)
        bye_team_ids = self._team_bye_ids(tournament_id)
        histories = self._team_color_histories(tournament_id)
        rank_by_team_id = {
            int(team_id): int(item.get("position", 0) or 0)
            for team_id, item in standings.items()
        }
        pending = sorted(
            teams,
            key=lambda team: self._team_pairing_order_key(team, standings, seed_ratings),
        )

        bye_team = None
        if len(pending) % 2 == 1:
            bye_team = self._choose_team_bye(pending, standings, bye_team_ids, seed_ratings)
            pending.remove(bye_team)

        team_pairs = (
            self._optimal_team_pairs(pending, standings, played_pairs, rank_by_team_id)
            if len(pending) <= self.MAX_EXHAUSTIVE_PAIRING_TEAMS
            else self._greedy_team_pairs(pending, standings, played_pairs, rank_by_team_id)
        )
        team_pairs = sorted(
            team_pairs,
            key=lambda pair: min(
                rank_by_team_id.get(int(pair[0]["id"]), 0),
                rank_by_team_id.get(int(pair[1]["id"]), 0),
            ),
        )

        matches = []
        for match_number, (team, opponent) in enumerate(team_pairs, start=1):
            white_team_id, black_team_id = self._choose_team_colors(team, opponent, histories, seed_ratings)
            matches.append(
                self._team_match_payload(
                    match_number,
                    white_team_id,
                    black_team_id,
                    rosters,
                    boards_count,
                )
            )
        if bye_team:
            matches.append(self._team_bye_payload(len(matches) + 1, int(bye_team["id"]), settings, boards_count))
        return matches

    @staticmethod
    def _team_match_payload(
        match_number: int,
        white_team_id: int,
        black_team_id: int,
        rosters: dict[int, dict[int, int]],
        boards_count: int,
    ) -> dict[str, Any]:
        boards = []
        for board_number in range(1, boards_count + 1):
            if board_number % 2 == 1:
                white_player_id = rosters[white_team_id][board_number]
                black_player_id = rosters[black_team_id][board_number]
            else:
                white_player_id = rosters[black_team_id][board_number]
                black_player_id = rosters[white_team_id][board_number]
            boards.append(
                {
                    "board_number": board_number,
                    "white_player_id": white_player_id,
                    "black_player_id": black_player_id,
                    "result": "",
                }
            )
        return {
            "match_number": match_number,
            "white_team_id": white_team_id,
            "black_team_id": black_team_id,
            "result": "",
            "is_bye": 0,
            "boards": boards,
        }

    @staticmethod
    def _team_bye_payload(
        match_number: int,
        team_id: int,
        settings: dict[str, Any],
        boards_count: int,
    ) -> dict[str, Any]:
        return {
            "match_number": match_number,
            "white_team_id": team_id,
            "black_team_id": None,
            "result": "BYE",
            "white_match_points": float(settings.get("team_match_win_points", 2.0) or 2.0),
            "black_match_points": 0.0,
            "white_game_points": float(boards_count),
            "black_game_points": 0.0,
            "is_bye": 1,
            "boards": [],
        }

    def update_result(
        self,
        tournament_id: int,
        pairing_id: int,
        result: str,
    ) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if tournament and tournament.get("competition_type") == "team":
            self._update_team_board_result(tournament_id, pairing_id, result)
            return

        pairing = self.db.get_pairing(pairing_id)
        if not pairing or int(pairing["tournament_id"]) != int(tournament_id):
            raise AppError("Mesa nao encontrada para o torneio selecionado.")
        if result not in RESULTS:
            raise AppError("Resultado invalido.")
        if pairing["round_status"] == "closed":
            settings = self.db.get_tournament_settings(tournament_id) or {}
            if not settings.get("allow_dangerous_changes"):
                raise AppError("Resultado de rodada fechada so pode ser alterado com mudancas perigosas habilitadas.")
        self.db.update_pairing_result(pairing_id, result)
        logger.info("Resultado da mesa %s atualizado para %s", pairing_id, result or "pendente")

    def close_round(self, tournament_id: int, round_id: int) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") == "team":
            self._close_team_round(tournament_id, round_id, tournament)
            return

        round_data = self.db.get_round(round_id)
        if not round_data or int(round_data["tournament_id"]) != int(tournament_id):
            raise AppError("Rodada nao encontrada para o torneio selecionado.")

        pairings = self.db.get_pairings_for_round(round_id)
        pending = [
            pairing
            for pairing in pairings
            if not pairing["result"] or pairing["result"] not in FINAL_RESULTS
        ]
        if pending:
            raise AppError("Preencha todos os resultados antes de fechar a rodada.")

        backup_path = self.db.backup(f"before_close_round_{round_id}")
        logger.info("Backup criado antes de fechar rodada %s: %s", round_id, backup_path)

        self.db.close_round(round_id)
        logger.info("Rodada %s fechada no torneio %s", round_id, tournament_id)

        latest = self.db.get_latest_round(tournament_id)
        if latest and int(latest["number"]) >= int(tournament["rounds_count"]):
            self.db.update_tournament_status(tournament_id, "finished")
        else:
            self.db.update_tournament_status(tournament_id, "running")

    def _update_team_board_result(self, tournament_id: int, team_board_id: int, result: str) -> None:
        board = self.db.get_team_board(team_board_id)
        if not board or int(board["tournament_id"]) != int(tournament_id):
            raise AppError("Tabuleiro nao encontrado para o torneio selecionado.")
        if result not in RESULTS:
            raise AppError("Resultado invalido.")
        if board["round_status"] == "closed":
            settings = self.db.get_tournament_settings(tournament_id) or {}
            if not settings.get("allow_dangerous_changes"):
                raise AppError("Resultado de rodada fechada so pode ser alterado com mudancas perigosas habilitadas.")
        self.db.update_team_board_result(team_board_id, result)
        logger.info("Resultado do tabuleiro de equipe %s atualizado para %s", team_board_id, result or "pendente")

    def _close_team_round(
        self,
        tournament_id: int,
        round_id: int,
        tournament: dict[str, Any],
    ) -> None:
        round_data = self.db.get_round(round_id)
        if not round_data or int(round_data["tournament_id"]) != int(tournament_id):
            raise AppError("Rodada nao encontrada para o torneio selecionado.")

        matches = self.db.list_team_matches_for_round(round_id)
        if not matches:
            raise AppError("A rodada nao possui confrontos por equipes.")

        pending: list[dict[str, Any]] = []
        summaries: list[dict[str, Any]] = []
        settings = self.db.get_tournament_settings(tournament_id) or {}
        for match in matches:
            if match["is_bye"]:
                summaries.append(
                    {
                        "team_match_id": int(match["id"]),
                        "result": "BYE",
                        "white_match_points": float(settings.get("team_match_win_points", 2.0) or 2.0),
                        "black_match_points": 0.0,
                        "white_game_points": float(settings.get("team_boards_count", 4) or 4),
                        "black_game_points": 0.0,
                    }
                )
                continue

            boards = self.db.list_team_boards(int(match["id"]))
            for board in boards:
                if not board["result"] or board["result"] not in RESULT_POINTS:
                    pending.append(board)
            if pending:
                continue

            white_game_points = 0.0
            black_game_points = 0.0
            white_team_player_ids = {
                int(player["player_id"])
                for player in self.db.list_team_players(int(match["white_team_id"]), active_only=False)
            }
            black_team_player_ids = {
                int(player["player_id"])
                for player in self.db.list_team_players(int(match["black_team_id"]), active_only=False)
            }
            for board in boards:
                white_points, black_points = RESULT_POINTS[str(board["result"])]
                if board.get("white_player_id") in white_team_player_ids:
                    white_game_points += white_points
                elif board.get("white_player_id") in black_team_player_ids:
                    black_game_points += white_points

                if board.get("black_player_id") in white_team_player_ids:
                    white_game_points += black_points
                elif board.get("black_player_id") in black_team_player_ids:
                    black_game_points += black_points

            if white_game_points > black_game_points:
                match_result = "1-0"
                white_match_points = float(settings.get("team_match_win_points", 2.0) or 2.0)
                black_match_points = float(settings.get("team_match_loss_points", 0.0) or 0.0)
            elif black_game_points > white_game_points:
                match_result = "0-1"
                white_match_points = float(settings.get("team_match_loss_points", 0.0) or 0.0)
                black_match_points = float(settings.get("team_match_win_points", 2.0) or 2.0)
            else:
                match_result = "1/2-1/2"
                white_match_points = float(settings.get("team_match_draw_points", 1.0) or 1.0)
                black_match_points = float(settings.get("team_match_draw_points", 1.0) or 1.0)

            summaries.append(
                {
                    "team_match_id": int(match["id"]),
                    "result": match_result,
                    "white_match_points": white_match_points,
                    "black_match_points": black_match_points,
                    "white_game_points": round(white_game_points, 2),
                    "black_game_points": round(black_game_points, 2),
                }
            )

        if pending:
            raise AppError("Preencha todos os resultados dos tabuleiros antes de fechar a rodada.")

        backup_path = self.db.backup(f"before_close_round_{round_id}")
        logger.info("Backup criado antes de fechar rodada por equipes %s: %s", round_id, backup_path)

        for summary in summaries:
            self.db.update_team_match_summary(
                int(summary["team_match_id"]),
                str(summary["result"]),
                float(summary["white_match_points"]),
                float(summary["black_match_points"]),
                float(summary["white_game_points"]),
                float(summary["black_game_points"]),
            )

        self.db.close_round(round_id)
        logger.info("Rodada por equipes %s fechada no torneio %s", round_id, tournament_id)

        latest = self.db.get_latest_round(tournament_id)
        if latest and int(latest["number"]) >= int(tournament["rounds_count"]):
            self.db.update_tournament_status(tournament_id, "finished")
        else:
            self.db.update_tournament_status(tournament_id, "running")

    def delete_generated_round(self, round_id: int) -> None:
        with self.db.connect() as connection:
            row = connection.execute(
                "SELECT status FROM rounds WHERE id = ?",
                (round_id,),
            ).fetchone()
        if row and row["status"] == "closed":
            raise AppError("Rodada fechada nao pode ser excluida.")
        self.db.delete_round(round_id)
        logger.info("Rodada gerada %s excluida", round_id)

    def delete_player_if_unpaired(self, tournament_id: int, player_id: int) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        player = self.db.get_player(player_id)
        if not player or int(player["tournament_id"]) != int(tournament_id):
            raise AppError("Jogador nao encontrado para o torneio selecionado.")
        pairings_count = self.db.count_player_pairings(player_id)
        if pairings_count:
            raise AppError(
                "Nao e possivel excluir jogador que ja aparece em rodada. "
                "Use o status Desistente ou Nao emparceirado para preservar o historico."
            )
        self.db.delete_player(player_id)
        logger.info("Jogador %s excluido do torneio %s antes de entrar em rodadas", player_id, tournament_id)

    def adjust_pairing_player(
        self,
        tournament_id: int,
        round_id: int,
        pairing_id: int,
        color: str,
        replacement_player_id: int,
    ) -> None:
        round_data = self.db.get_round(round_id)
        if not round_data or int(round_data["tournament_id"]) != int(tournament_id):
            raise AppError("Rodada nao encontrada para o torneio selecionado.")
        if round_data["status"] == "closed":
            raise AppError("Rodada fechada nao pode ser ajustada.")
        if color not in {"white", "black"}:
            raise AppError("Cor invalida para ajuste.")

        replacement = self.db.get_player(replacement_player_id)
        if not replacement or int(replacement["tournament_id"]) != int(tournament_id):
            raise AppError("Jogador substituto nao pertence ao torneio.")
        if not replacement["active"]:
            raise AppError("Jogador substituto precisa estar ativo.")

        pairings = self.db.get_pairings_for_round(round_id)
        if not pairings:
            raise AppError("A rodada nao possui mesas para ajustar.")

        pairing_by_id = {int(pairing["id"]): pairing for pairing in pairings}
        source_pairing = pairing_by_id.get(int(pairing_id))
        if not source_pairing:
            raise AppError("Mesa selecionada nao encontrada.")
        if color == "black" and (source_pairing["is_bye"] or not source_pairing["black_player_id"]):
            raise AppError("Bye nao possui jogador de pretas para trocar.")

        affected_pairings = [source_pairing]
        target_slot = self._find_player_slot(pairings, replacement_player_id)
        if target_slot and target_slot["pairing"]["id"] != source_pairing["id"]:
            affected_pairings.append(target_slot["pairing"])

        for affected in affected_pairings:
            if affected["result"] and affected["result"] != "BYE":
                raise AppError("Limpe os resultados das mesas afetadas antes de trocar jogadores.")

        source_player_id = (
            source_pairing["white_player_id"]
            if color == "white"
            else source_pairing["black_player_id"]
        )
        if source_player_id is None:
            raise AppError("Jogador de origem invalido.")
        if int(source_player_id) == int(replacement_player_id):
            return

        updates: dict[int, dict[str, int | None]] = {
            int(pairing["id"]): {
                "white": int(pairing["white_player_id"]),
                "black": int(pairing["black_player_id"]) if pairing["black_player_id"] else None,
            }
            for pairing in pairings
        }

        updates[int(source_pairing["id"])][color] = int(replacement_player_id)
        if target_slot:
            updates[int(target_slot["pairing"]["id"])][target_slot["color"]] = int(source_player_id)

        for pairing_id_to_update, values in updates.items():
            black_player_id = values["black"]
            if black_player_id is not None and values["white"] == black_player_id:
                raise AppError("Uma mesa nao pode ter o mesmo jogador dos dois lados.")

        pairing_updates = []
        for pairing_id_to_update, values in updates.items():
            if pairing_id_to_update not in {int(source_pairing["id"])} and not (
                target_slot and pairing_id_to_update == int(target_slot["pairing"]["id"])
            ):
                continue
            white_player_id = values["white"]
            if white_player_id is None:
                raise AppError("Uma mesa precisa ter jogador de brancas.")
            pairing_updates.append((pairing_id_to_update, white_player_id, values["black"]))
        self.db.update_pairing_players(pairing_updates)
        logger.info(
            "Ajuste manual na rodada %s: mesa %s, cor %s, jogador %s",
            round_id,
            pairing_id,
            color,
            replacement_player_id,
        )

    def standings(self, tournament_id: int) -> list[dict[str, Any]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            return []

        players = self.db.list_players(tournament_id, active_only=False)
        stats: dict[int, dict[str, Any]] = {}
        for player in players:
            stats[player["id"]] = {
                "player_id": player["id"],
                "name": player_full_name(player),
                "club": player["club"],
                "rating": int(player["rating"] or 0),
                "category": player["category"],
                "age_category": player.get("age_category", ""),
                "rating_category": player.get("rating_category", ""),
                "prize_tags": player.get("prize_tags", ""),
                "active": int(player["active"]),
                "player_status": player.get("player_status", "active"),
                "starting_points": float(player.get("starting_points", 0.0) or 0.0),
                "points": float(player.get("starting_points", 0.0) or 0.0),
                "wins": 0,
                "buchholz": 0.0,
                "buchholz_median": 0.0,
                "sonneborn_berger": 0.0,
                "white_count": 0,
                "black_count": 0,
                "byes": 0,
                "opponents": [],
                "earned_against": [],
                "performance": "",
            }

        closed_pairings = self.db.get_pairings_for_tournament(tournament_id, closed_only=True)
        for pairing in closed_pairings:
            white_id = pairing["white_player_id"]
            black_id = pairing["black_player_id"]
            result = pairing["result"]

            if white_id not in stats:
                continue

            if pairing["is_bye"]:
                stats[white_id]["points"] += float(tournament["bye_points"])
                stats[white_id]["byes"] += 1
                continue

            if not black_id or black_id not in stats or result not in RESULT_POINTS:
                continue

            white_points, black_points = RESULT_POINTS[result]
            stats[white_id]["points"] += white_points
            stats[black_id]["points"] += black_points
            stats[white_id]["white_count"] += 1
            stats[black_id]["black_count"] += 1
            stats[white_id]["opponents"].append(black_id)
            stats[black_id]["opponents"].append(white_id)
            stats[white_id]["earned_against"].append((black_id, white_points))
            stats[black_id]["earned_against"].append((white_id, black_points))
            if white_points == 1.0 and black_points == 0.0:
                stats[white_id]["wins"] += 1
            if black_points == 1.0 and white_points == 0.0:
                stats[black_id]["wins"] += 1

        for player_stat in stats.values():
            opponent_scores = [
                stats[opponent_id]["points"]
                for opponent_id in player_stat["opponents"]
                if opponent_id in stats
            ]
            player_stat["buchholz"] = round(sum(opponent_scores), 2)
            if len(opponent_scores) >= 3:
                ordered = sorted(opponent_scores)
                player_stat["buchholz_median"] = round(sum(ordered[1:-1]), 2)
            else:
                player_stat["buchholz_median"] = player_stat["buchholz"]

            sb = 0.0
            for opponent_id, earned in player_stat["earned_against"]:
                sb += stats[opponent_id]["points"] * earned
            player_stat["sonneborn_berger"] = round(sb, 2)
            player_stat["performance"] = self._performance_rating(player_stat, stats)

        ordered_stats = sorted(
            stats.values(),
            key=lambda item: (
                -item["points"],
                -item["buchholz"],
                -item["buchholz_median"],
                -item["sonneborn_berger"],
                -item["wins"],
                -item["rating"],
                item["name"].casefold(),
            ),
        )

        for index, item in enumerate(ordered_stats, start=1):
            item["position"] = index
        return ordered_stats

    def team_standings(self, tournament_id: int) -> list[dict[str, Any]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            return []
        settings = self.db.get_tournament_settings(tournament_id) or {}
        teams = self.db.list_teams(tournament_id, active_only=False)
        stats: dict[int, dict[str, Any]] = {}
        for team in teams:
            team_id = int(team["id"])
            stats[team_id] = {
                "team_id": team_id,
                "name": team["name"],
                "club": team.get("club", ""),
                "captain": team.get("captain", ""),
                "active": int(team.get("active", 0) or 0),
                "match_points": 0.0,
                "game_points": 0.0,
                "wins": 0,
                "draws": 0,
                "losses": 0,
                "byes": 0,
                "matches": 0,
                "buchholz": 0.0,
                "opponents": [],
            }

        for match in self.db.list_team_matches_for_tournament(tournament_id, closed_only=True):
            white_team_id = int(match["white_team_id"])
            black_team_id = int(match["black_team_id"]) if match.get("black_team_id") else None
            if white_team_id not in stats:
                continue
            if match.get("is_bye"):
                stats[white_team_id]["match_points"] += float(match.get("white_match_points", 0.0) or 0.0)
                stats[white_team_id]["game_points"] += float(match.get("white_game_points", 0.0) or 0.0)
                stats[white_team_id]["wins"] += 1
                stats[white_team_id]["byes"] += 1
                continue
            if black_team_id is None or black_team_id not in stats:
                continue

            white_match_points = float(match.get("white_match_points", 0.0) or 0.0)
            black_match_points = float(match.get("black_match_points", 0.0) or 0.0)
            stats[white_team_id]["match_points"] += white_match_points
            stats[black_team_id]["match_points"] += black_match_points
            stats[white_team_id]["game_points"] += float(match.get("white_game_points", 0.0) or 0.0)
            stats[black_team_id]["game_points"] += float(match.get("black_game_points", 0.0) or 0.0)
            stats[white_team_id]["matches"] += 1
            stats[black_team_id]["matches"] += 1
            stats[white_team_id]["opponents"].append(black_team_id)
            stats[black_team_id]["opponents"].append(white_team_id)

            if white_match_points > black_match_points:
                stats[white_team_id]["wins"] += 1
                stats[black_team_id]["losses"] += 1
            elif black_match_points > white_match_points:
                stats[black_team_id]["wins"] += 1
                stats[white_team_id]["losses"] += 1
            else:
                stats[white_team_id]["draws"] += 1
                stats[black_team_id]["draws"] += 1

        for team_stat in stats.values():
            team_stat["match_points"] = round(float(team_stat["match_points"]), 2)
            team_stat["game_points"] = round(float(team_stat["game_points"]), 2)
            team_stat["buchholz"] = round(
                sum(
                    float(stats[opponent_id]["match_points"])
                    for opponent_id in team_stat["opponents"]
                    if opponent_id in stats
                ),
                2,
            )

        primary = str(settings.get("team_standing_primary", "match_points") or "match_points")
        secondary = str(settings.get("team_standing_secondary", "game_points") or "game_points")
        ordered_stats = sorted(
            stats.values(),
            key=lambda item: (
                -self._team_standing_value(item, primary),
                -self._team_standing_value(item, secondary),
                -float(item["buchholz"]),
                -int(item["wins"]),
                str(item["name"]).casefold(),
            ),
        )
        for index, item in enumerate(ordered_stats, start=1):
            item["position"] = index
        return ordered_stats

    @staticmethod
    def _team_standing_value(item: dict[str, Any], criterion: str) -> float:
        if criterion == "wins":
            return float(item.get("wins", 0) or 0)
        if criterion == "game_points":
            return float(item.get("game_points", 0.0) or 0.0)
        return float(item.get("match_points", 0.0) or 0.0)

    @staticmethod
    def _performance_rating(
        player_stat: dict[str, Any],
        stats: dict[int, dict[str, Any]],
    ) -> int | str:
        games = len(player_stat["earned_against"])
        if not games:
            return ""

        opponent_ratings = [
            int(stats[opponent_id]["rating"] or 0)
            for opponent_id, _earned in player_stat["earned_against"]
            if opponent_id in stats and int(stats[opponent_id]["rating"] or 0) > 0
        ]
        if not opponent_ratings:
            return ""

        score = sum(float(earned) for _opponent_id, earned in player_stat["earned_against"])
        average_rating = sum(opponent_ratings) / len(opponent_ratings)
        diff: float
        if score <= 0:
            diff = -800
        elif score >= games:
            diff = 800
        else:
            diff = 400 * math.log10(score / (games - score))
            diff = max(min(diff, 800), -800)
        return int(round(average_rating + diff))

    def _round_robin_pairings(self, tournament_id: int, players: list[dict[str, Any]], next_number: int, settings: dict[str, Any]) -> list[dict[str, Any]]:
        ordered = sorted(
            players,
            key=lambda player: (-int(player["rating"] or 0), player["name"].casefold()),
        )
        if len(ordered) % 2 == 1:
            ordered.append({"id": -1, "name": "BYE", "is_dummy": True})
        
        N = len(ordered)
        if next_number > N - 1:
            raise AppError("O numero maximo de rodadas do torneio ja foi atingido.")
            
        r = next_number
        P = [ordered[0]]
        for i in range(1, N):
            shift = r - 1
            idx = ((i - 1 - shift) % (N - 1)) + 1
            P.append(ordered[idx])
            
        upper = P[: N // 2]
        lower = P[N // 2 :]
        lower.reverse()
        
        pairings: list[dict[str, Any]] = []
        board = 1
        for k in range(N // 2):
            if (r % 2 == 1 and k == 0) or (r % 2 == 0 and k > 0):
                white, black = upper[k], lower[k]
            else:
                white, black = lower[k], upper[k]
                
            if white["id"] == -1 or black["id"] == -1:
                real_player = black if white["id"] == -1 else white
                pairings.append(
                    {
                        "board_number": board,
                        "white_player_id": real_player["id"],
                        "black_player_id": None,
                        "result": "1-0" if not settings.get("disable_bye") else "",
                        "is_bye": 1,
                    }
                )
            else:
                pairings.append(
                    {
                        "board_number": board,
                        "white_player_id": white["id"],
                        "black_player_id": black["id"],
                        "result": "",
                        "is_bye": 0,
                    }
                )
            board += 1
            
        return pairings

    def _knockout_pairings(self, tournament_id: int, players: list[dict[str, Any]], next_number: int, settings: dict[str, Any]) -> list[dict[str, Any]]:
        ordered = sorted(
            players,
            key=lambda player: (-int(player["rating"] or 0), player["name"].casefold()),
        )
        
        if next_number == 1:
            active_players = ordered[:]
        else:
            prev_round = self.db.get_round_by_number(tournament_id, next_number - 1)
            if not prev_round:
                raise AppError("Rodada anterior não encontrada.")
                
            pairings = self.db.list_pairings(prev_round["id"])
            active_players_set = set()
            seed_map = {int(p["id"]): i for i, p in enumerate(ordered)}
            
            for p in pairings:
                if p["is_bye"]:
                    active_players_set.add(int(p["white_player_id"]))
                    continue
                    
                w_id = int(p["white_player_id"])
                b_id = int(p["black_player_id"])
                
                if p["result"] == "1-0":
                    active_players_set.add(w_id)
                elif p["result"] == "0-1":
                    active_players_set.add(b_id)
                else: 
                    w_idx = seed_map.get(w_id, 9999)
                    b_idx = seed_map.get(b_id, 9999)
                    if w_idx < b_idx:
                        active_players_set.add(w_id)
                    else:
                        active_players_set.add(b_id)
                        
            active_players = [pl for pl in ordered if int(pl["id"]) in active_players_set]
            
        if len(active_players) == 1:
            raise AppError("O torneio já tem um vencedor. Não é possível gerar mais rodadas.")
            
        N = len(active_players)
        pow2 = 1
        while pow2 < N:
            pow2 *= 2
            
        if pow2 != N:
            num_byes = pow2 - N
            bye_players = active_players[:num_byes]
            playing_players = active_players[num_byes:]
        else:
            bye_players = []
            playing_players = active_players[:]
            
        pairings: list[dict[str, Any]] = []
        board = 1
        
        half = len(playing_players) // 2
        for i in range(half):
            p1 = playing_players[i]
            p2 = playing_players[len(playing_players) - 1 - i]
            if i % 2 == 0:
                w, b = p1, p2
            else:
                w, b = p2, p1
                
            pairings.append({
                "board_number": board,
                "white_player_id": w["id"],
                "black_player_id": b["id"],
                "result": "",
                "is_bye": 0,
            })
            board += 1
            
        for p in bye_players:
            pairings.append({
                "board_number": board,
                "white_player_id": p["id"],
                "black_player_id": None,
                "result": "1-0" if not settings.get("disable_bye") else "",
                "is_bye": 1,
            })
            board += 1
            
        return pairings

    def _first_round_pairings(self, players: list[dict[str, Any]]) -> list[dict[str, Any]]:
        ordered = sorted(
            players,
            key=lambda player: (-int(player["rating"] or 0), player["name"].casefold()),
        )
        pairable = ordered[:]
        bye_player = None
        if len(pairable) % 2 == 1:
            bye_player = min(
                pairable,
                key=lambda player: (int(player["rating"] or 0), player["name"].casefold()),
            )
            pairable.remove(bye_player)

        half = len(pairable) // 2
        upper = pairable[:half]
        lower = pairable[half:]
        pairings: list[dict[str, Any]] = []

        board = 1
        for index, player in enumerate(upper):
            opponent = lower[index]
            if index % 2 == 0:
                white_id = player["id"]
                black_id = opponent["id"]
            else:
                white_id = opponent["id"]
                black_id = player["id"]
            pairings.append(
                {
                    "board_number": board,
                    "white_player_id": white_id,
                    "black_player_id": black_id,
                    "result": "",
                    "is_bye": 0,
                }
            )
            board += 1

        if bye_player:
            pairings.append(
                {
                    "board_number": board,
                    "white_player_id": bye_player["id"],
                    "black_player_id": None,
                    "result": "BYE",
                    "is_bye": 1,
                }
            )
        return pairings

    @staticmethod
    def _find_player_slot(
        pairings: list[dict[str, Any]],
        player_id: int,
    ) -> dict[str, Any] | None:
        for pairing in pairings:
            if int(pairing["white_player_id"]) == int(player_id):
                return {"pairing": pairing, "color": "white"}
            if pairing["black_player_id"] and int(pairing["black_player_id"]) == int(player_id):
                return {"pairing": pairing, "color": "black"}
        return None

    def _swiss_pairings(
        self,
        tournament_id: int,
        players: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        standings = {item["player_id"]: item for item in self.standings(tournament_id)}
        histories = self._color_histories(tournament_id)
        float_histories = self._float_histories(tournament_id)
        played_pairs = self._played_pairs(tournament_id)
        bye_player_ids = self._bye_player_ids(tournament_id)
        rank_by_player_id = self._rank_by_player_id(standings)

        pending = sorted(
            players,
            key=lambda player: self._pairing_order_key(player, standings),
        )

        pairings: list[dict[str, Any]] = []
        board = 1

        if len(pending) % 2 == 1:
            bye_player = self._choose_bye_player(pending, standings, bye_player_ids)
            pending.remove(bye_player)
            pairings.append(
                {
                    "board_number": 9999,
                    "white_player_id": bye_player["id"],
                    "black_player_id": None,
                    "result": "BYE",
                    "is_bye": 1,
                }
            )

        player_pairs = (
            self._optimal_player_pairs(
                pending,
                standings,
                histories,
                float_histories,
                played_pairs,
                rank_by_player_id,
            )
            if len(pending) <= self.MAX_EXHAUSTIVE_PAIRING_PLAYERS
            else self._greedy_player_pairs(
                pending,
                standings,
                histories,
                float_histories,
                played_pairs,
                rank_by_player_id,
            )
        )
        player_pairs = sorted(
            player_pairs,
            key=lambda pair: min(
                rank_by_player_id.get(pair[0]["id"], 0),
                rank_by_player_id.get(pair[1]["id"], 0),
            ),
        )
        for player, opponent in player_pairs:
            white_id, black_id = self._choose_colors(player, opponent, histories)
            pairings.append(
                {
                    "board_number": board,
                    "white_player_id": white_id,
                    "black_player_id": black_id,
                    "result": "",
                    "is_bye": 0,
                }
            )
            board += 1

        normal_pairings = [pairing for pairing in pairings if not pairing["is_bye"]]
        bye_pairings = [pairing for pairing in pairings if pairing["is_bye"]]
        for index, pairing in enumerate(normal_pairings + bye_pairings, start=1):
            pairing["board_number"] = index
        return normal_pairings + bye_pairings

    def _greedy_player_pairs(
        self,
        players: list[dict[str, Any]],
        standings: dict[int, dict[str, Any]],
        histories: dict[int, list[str]],
        float_histories: dict[int, list[str]],
        played_pairs: set[frozenset[int]],
        rank_by_player_id: dict[int, int],
    ) -> list[tuple[dict[str, Any], dict[str, Any]]]:
        pending = players[:]
        pairs = []
        while pending:
            player = pending.pop(0)
            opponent = min(
                pending,
                key=lambda candidate: self._pair_penalty(
                    player,
                    candidate,
                    standings,
                    histories,
                    float_histories,
                    played_pairs,
                    rank_by_player_id,
                ),
            )
            pending.remove(opponent)
            pairs.append((player, opponent))
        return pairs

    def _optimal_player_pairs(
        self,
        players: list[dict[str, Any]],
        standings: dict[int, dict[str, Any]],
        histories: dict[int, list[str]],
        float_histories: dict[int, list[str]],
        played_pairs: set[frozenset[int]],
        rank_by_player_id: dict[int, int],
    ) -> list[tuple[dict[str, Any], dict[str, Any]]]:
        player_by_id = {int(player["id"]): player for player in players}
        penalty_cache: dict[tuple[int, int], float] = {}

        def penalty(player_id: int, opponent_id: int) -> float:
            key = (player_id, opponent_id)
            if key not in penalty_cache:
                penalty_cache[key] = self._pair_penalty(
                    player_by_id[player_id],
                    player_by_id[opponent_id],
                    standings,
                    histories,
                    float_histories,
                    played_pairs,
                    rank_by_player_id,
                )
            return penalty_cache[key]

        best_cost = float("inf")
        best_pairs: list[tuple[int, int]] = []
        seen_costs: dict[tuple[int, ...], float] = {}

        def search(
            remaining: tuple[int, ...],
            selected_pairs: list[tuple[int, int]],
            current_cost: float,
        ) -> None:
            nonlocal best_cost, best_pairs
            if current_cost >= best_cost:
                return
            if current_cost >= seen_costs.get(remaining, float("inf")):
                return
            seen_costs[remaining] = current_cost
            if not remaining:
                best_cost = current_cost
                best_pairs = selected_pairs[:]
                return

            player_id = remaining[0]
            candidates = sorted(remaining[1:], key=lambda candidate_id: penalty(player_id, candidate_id))
            for opponent_id in candidates:
                next_remaining = tuple(
                    item_id for item_id in remaining[1:] if item_id != opponent_id
                )
                search(
                    next_remaining,
                    selected_pairs + [(player_id, opponent_id)],
                    current_cost + penalty(player_id, opponent_id),
                )

        search(tuple(player_by_id), [], 0.0)
        if not best_pairs:
            return self._greedy_player_pairs(
                players,
                standings,
                histories,
                float_histories,
                played_pairs,
                rank_by_player_id,
            )
        return [(player_by_id[player_id], player_by_id[opponent_id]) for player_id, opponent_id in best_pairs]

    @staticmethod
    def _rank_by_player_id(standings: dict[int, dict[str, Any]]) -> dict[int, int]:
        return {
            int(player_id): int(item.get("position", 0) or 0)
            for player_id, item in standings.items()
        }

    @staticmethod
    def _pairing_order_key(
        player: dict[str, Any],
        standings: dict[int, dict[str, Any]],
    ) -> tuple[float, float, float, float, int, int, str]:
        standing = standings.get(player["id"], {})
        return (
            -float(standing.get("points", 0.0) or 0.0),
            -float(standing.get("buchholz", 0.0) or 0.0),
            -float(standing.get("buchholz_median", 0.0) or 0.0),
            -float(standing.get("sonneborn_berger", 0.0) or 0.0),
            -int(standing.get("wins", 0) or 0),
            -int(player["rating"] or 0),
            player["name"].casefold(),
        )

    def _choose_bye_player(
        self,
        players: list[dict[str, Any]],
        standings: dict[int, dict[str, Any]],
        bye_player_ids: set[int],
    ) -> dict[str, Any]:
        candidates = [player for player in players if player["id"] not in bye_player_ids]
        if not candidates:
            candidates = players
        return min(
            candidates,
            key=lambda player: (
                standings.get(player["id"], {}).get("points", 0.0),
                int(player["rating"] or 0),
                player["name"].casefold(),
            ),
        )

    @staticmethod
    def _team_pairing_order_key(
        team: dict[str, Any],
        standings: dict[int, dict[str, Any]],
        seed_ratings: dict[int, int],
    ) -> tuple[float, float, int, str]:
        team_id = int(team["id"])
        standing = standings.get(team_id, {})
        return (
            -float(standing.get("match_points", 0.0) or 0.0),
            -float(standing.get("game_points", 0.0) or 0.0),
            -int(seed_ratings.get(team_id, 0)),
            str(team["name"]).casefold(),
        )

    def _choose_team_bye(
        self,
        teams: list[dict[str, Any]],
        standings: dict[int, dict[str, Any]],
        bye_team_ids: set[int],
        seed_ratings: dict[int, int],
    ) -> dict[str, Any]:
        candidates = [team for team in teams if int(team["id"]) not in bye_team_ids]
        if not candidates:
            candidates = teams
        return min(
            candidates,
            key=lambda team: (
                float(standings.get(int(team["id"]), {}).get("match_points", 0.0) or 0.0),
                float(standings.get(int(team["id"]), {}).get("game_points", 0.0) or 0.0),
                seed_ratings.get(int(team["id"]), 0),
                str(team["name"]).casefold(),
            ),
        )

    def _greedy_team_pairs(
        self,
        teams: list[dict[str, Any]],
        standings: dict[int, dict[str, Any]],
        played_pairs: set[frozenset[int]],
        rank_by_team_id: dict[int, int],
    ) -> list[tuple[dict[str, Any], dict[str, Any]]]:
        pending = teams[:]
        pairs = []
        while pending:
            team = pending.pop(0)
            opponent = min(
                pending,
                key=lambda candidate: self._team_pair_penalty(
                    team,
                    candidate,
                    standings,
                    played_pairs,
                    rank_by_team_id,
                ),
            )
            pending.remove(opponent)
            pairs.append((team, opponent))
        return pairs

    def _optimal_team_pairs(
        self,
        teams: list[dict[str, Any]],
        standings: dict[int, dict[str, Any]],
        played_pairs: set[frozenset[int]],
        rank_by_team_id: dict[int, int],
    ) -> list[tuple[dict[str, Any], dict[str, Any]]]:
        team_by_id = {int(team["id"]): team for team in teams}
        penalty_cache: dict[tuple[int, int], float] = {}

        def penalty(team_id: int, opponent_id: int) -> float:
            key = (team_id, opponent_id)
            if key not in penalty_cache:
                penalty_cache[key] = self._team_pair_penalty(
                    team_by_id[team_id],
                    team_by_id[opponent_id],
                    standings,
                    played_pairs,
                    rank_by_team_id,
                )
            return penalty_cache[key]

        best_cost = float("inf")
        best_pairs: list[tuple[int, int]] = []
        seen_costs: dict[tuple[int, ...], float] = {}

        def search(
            remaining: tuple[int, ...],
            selected_pairs: list[tuple[int, int]],
            current_cost: float,
        ) -> None:
            nonlocal best_cost, best_pairs
            if current_cost >= best_cost:
                return
            if current_cost >= seen_costs.get(remaining, float("inf")):
                return
            seen_costs[remaining] = current_cost
            if not remaining:
                best_cost = current_cost
                best_pairs = selected_pairs[:]
                return

            team_id = remaining[0]
            candidates = sorted(remaining[1:], key=lambda candidate_id: penalty(team_id, candidate_id))
            for opponent_id in candidates:
                next_remaining = tuple(item_id for item_id in remaining[1:] if item_id != opponent_id)
                search(
                    next_remaining,
                    selected_pairs + [(team_id, opponent_id)],
                    current_cost + penalty(team_id, opponent_id),
                )

        search(tuple(team_by_id), [], 0.0)
        if not best_pairs:
            return self._greedy_team_pairs(teams, standings, played_pairs, rank_by_team_id)
        return [(team_by_id[team_id], team_by_id[opponent_id]) for team_id, opponent_id in best_pairs]

    def _team_pair_penalty(
        self,
        team: dict[str, Any],
        opponent: dict[str, Any],
        standings: dict[int, dict[str, Any]],
        played_pairs: set[frozenset[int]],
        rank_by_team_id: dict[int, int],
    ) -> float:
        team_id = int(team["id"])
        opponent_id = int(opponent["id"])
        team_score = float(standings.get(team_id, {}).get("match_points", 0.0) or 0.0)
        opponent_score = float(standings.get(opponent_id, {}).get("match_points", 0.0) or 0.0)
        score_diff = abs(team_score - opponent_score)
        penalty = score_diff * self.SCORE_DIFF_PENALTY
        if score_diff:
            penalty += self.SCORE_GROUP_FLOAT_PENALTY

        rank_distance = abs(rank_by_team_id.get(team_id, 0) - rank_by_team_id.get(opponent_id, 0))
        penalty += rank_distance * (4 if not score_diff else 1)

        if frozenset((team_id, opponent_id)) in played_pairs:
            penalty += self.REPEAT_PAIRING_PENALTY
        return penalty

    def _choose_team_colors(
        self,
        team: dict[str, Any],
        opponent: dict[str, Any],
        histories: dict[int, list[str]],
        seed_ratings: dict[int, int],
    ) -> tuple[int, int]:
        team_id = int(team["id"])
        opponent_id = int(opponent["id"])
        first_penalty = self._assignment_color_penalty(team_id, "W", histories)
        first_penalty += self._assignment_color_penalty(opponent_id, "B", histories)

        second_penalty = self._assignment_color_penalty(team_id, "B", histories)
        second_penalty += self._assignment_color_penalty(opponent_id, "W", histories)

        if first_penalty < second_penalty:
            return team_id, opponent_id
        if second_penalty < first_penalty:
            return opponent_id, team_id
        if seed_ratings.get(team_id, 0) >= seed_ratings.get(opponent_id, 0):
            return team_id, opponent_id
        return opponent_id, team_id

    def _team_played_pairs(self, tournament_id: int) -> set[frozenset[int]]:
        played: set[frozenset[int]] = set()
        for match in self.db.list_team_matches_for_tournament(tournament_id, closed_only=True):
            if match["is_bye"] or not match["black_team_id"]:
                continue
            played.add(frozenset((int(match["white_team_id"]), int(match["black_team_id"]))))
        return played

    def _team_bye_ids(self, tournament_id: int) -> set[int]:
        return {
            int(match["white_team_id"])
            for match in self.db.list_team_matches_for_tournament(tournament_id, closed_only=True)
            if match["is_bye"]
        }

    def _team_color_histories(self, tournament_id: int) -> dict[int, list[str]]:
        histories: dict[int, list[str]] = {}
        for match in self.db.list_team_matches_for_tournament(tournament_id, closed_only=True):
            white_team_id = int(match["white_team_id"])
            black_team_id = int(match["black_team_id"]) if match["black_team_id"] else None
            histories.setdefault(white_team_id, [])
            if match["is_bye"]:
                histories[white_team_id].append("BYE")
                continue
            histories[white_team_id].append("W")
            if black_team_id:
                histories.setdefault(black_team_id, [])
                histories[black_team_id].append("B")
        return histories

    def _played_pairs(self, tournament_id: int) -> set[frozenset[int]]:
        played: set[frozenset[int]] = set()
        for pairing in self.db.get_pairings_for_tournament(tournament_id, closed_only=True):
            if pairing["is_bye"] or not pairing["black_player_id"]:
                continue
            played.add(frozenset((pairing["white_player_id"], pairing["black_player_id"])))
        return played

    def _bye_player_ids(self, tournament_id: int) -> set[int]:
        return {
            pairing["white_player_id"]
            for pairing in self.db.get_pairings_for_tournament(tournament_id, closed_only=True)
            if pairing["is_bye"]
        }

    def _color_histories(self, tournament_id: int) -> dict[int, list[str]]:
        histories: dict[int, list[str]] = {}
        for pairing in self.db.get_pairings_for_tournament(tournament_id, closed_only=True):
            white_id = pairing["white_player_id"]
            black_id = pairing["black_player_id"]
            histories.setdefault(white_id, [])
            if pairing["is_bye"]:
                histories[white_id].append("BYE")
                continue
            if black_id:
                histories.setdefault(black_id, [])
                histories[white_id].append("W")
                histories[black_id].append("B")
        return histories

    def _float_histories(self, tournament_id: int) -> dict[int, list[str]]:
        tournament = self.db.get_tournament(tournament_id)
        bye_points = float(tournament["bye_points"] if tournament else 0.0)
        players = self.db.list_players(tournament_id, active_only=False)
        scores = {
            int(player["id"]): float(player.get("starting_points", 0.0) or 0.0)
            for player in players
        }
        histories: dict[int, list[str]] = {player_id: [] for player_id in scores}

        for pairing in self.db.get_pairings_for_tournament(tournament_id, closed_only=True):
            white_id = int(pairing["white_player_id"])
            black_id = int(pairing["black_player_id"]) if pairing["black_player_id"] else None
            histories.setdefault(white_id, [])
            scores.setdefault(white_id, 0.0)

            if pairing["is_bye"]:
                histories[white_id].append("bye")
                scores[white_id] += bye_points
                continue

            if black_id is None:
                continue
            histories.setdefault(black_id, [])
            scores.setdefault(black_id, 0.0)

            white_score = scores[white_id]
            black_score = scores[black_id]
            if white_score < black_score:
                histories[white_id].append("up")
                histories[black_id].append("down")
            elif white_score > black_score:
                histories[white_id].append("down")
                histories[black_id].append("up")
            else:
                histories[white_id].append("=")
                histories[black_id].append("=")

            result = pairing["result"]
            if result in RESULT_POINTS:
                white_points, black_points = RESULT_POINTS[result]
                scores[white_id] += white_points
                scores[black_id] += black_points
        return histories

    def _pair_penalty(
        self,
        player: dict[str, Any],
        opponent: dict[str, Any],
        standings: dict[int, dict[str, Any]],
        histories: dict[int, list[str]],
        float_histories: dict[int, list[str]],
        played_pairs: set[frozenset[int]],
        rank_by_player_id: dict[int, int],
    ) -> float:
        player_id = player["id"]
        opponent_id = opponent["id"]
        player_score = float(standings.get(player_id, {}).get("points", 0.0) or 0.0)
        opponent_score = float(standings.get(opponent_id, {}).get("points", 0.0) or 0.0)
        score_diff = abs(player_score - opponent_score)
        penalty = score_diff * self.SCORE_DIFF_PENALTY

        if score_diff:
            penalty += self.SCORE_GROUP_FLOAT_PENALTY
            if player_score < opponent_score:
                penalty += self._float_penalty(player_id, "up", float_histories)
                penalty += self._float_penalty(opponent_id, "down", float_histories)
            else:
                penalty += self._float_penalty(player_id, "down", float_histories)
                penalty += self._float_penalty(opponent_id, "up", float_histories)

        rank_distance = abs(
            rank_by_player_id.get(player_id, 0)
            - rank_by_player_id.get(opponent_id, 0)
        )
        penalty += rank_distance * (4 if not score_diff else 1)

        if frozenset((player_id, opponent_id)) in played_pairs:
            penalty += self.REPEAT_PAIRING_PENALTY

        white_a, black_a = self._choose_colors(player, opponent, histories)
        color_penalty = self._assignment_color_penalty(player_id, "W", histories)
        color_penalty += self._assignment_color_penalty(opponent_id, "B", histories)
        if white_a == opponent_id and black_a == player_id:
            color_penalty = self._assignment_color_penalty(player_id, "B", histories)
            color_penalty += self._assignment_color_penalty(opponent_id, "W", histories)

        return penalty + color_penalty

    @staticmethod
    def _float_penalty(
        player_id: int,
        direction: str,
        float_histories: dict[int, list[str]],
    ) -> int:
        history = float_histories.get(player_id, [])
        penalty = history.count(direction) * 25
        if history[-2:] == [direction, direction]:
            penalty += 300
        elif history[-1:] == [direction]:
            penalty += 120
        return penalty

    def _choose_colors(
        self,
        player: dict[str, Any],
        opponent: dict[str, Any],
        histories: dict[int, list[str]],
    ) -> tuple[int, int]:
        player_id = player["id"]
        opponent_id = opponent["id"]
        first_penalty = self._assignment_color_penalty(player_id, "W", histories)
        first_penalty += self._assignment_color_penalty(opponent_id, "B", histories)

        second_penalty = self._assignment_color_penalty(player_id, "B", histories)
        second_penalty += self._assignment_color_penalty(opponent_id, "W", histories)

        if first_penalty < second_penalty:
            return player_id, opponent_id
        if second_penalty < first_penalty:
            return opponent_id, player_id

        if int(player["rating"] or 0) >= int(opponent["rating"] or 0):
            return player_id, opponent_id
        return opponent_id, player_id

    def _assignment_color_penalty(
        self,
        player_id: int,
        color: str,
        histories: dict[int, list[str]],
    ) -> int:
        history = [item for item in histories.get(player_id, []) if item in {"W", "B"}]
        white_count = history.count("W")
        black_count = history.count("B")
        next_white = white_count + (1 if color == "W" else 0)
        next_black = black_count + (1 if color == "B" else 0)
        penalty = abs(next_white - next_black) * 12

        if history[-2:] == [color, color]:
            penalty += 250
        elif history[-1:] == [color]:
            penalty += 20

        if white_count - black_count >= 2 and color == "W":
            penalty += 120
        if black_count - white_count >= 2 and color == "B":
            penalty += 120

        return penalty


class ImportService:
    def __init__(self, db: Database) -> None:
        self.db = db

    def preview_online_registrations_csv(self, tournament_id: int, file_path: str | Path) -> dict[str, Any]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")

        rows = self._online_registration_rows(tournament_id, file_path)
        summary = {
            "total": len(rows),
            "ready": sum(1 for row in rows if row["status"] == "ready"),
            "duplicate": sum(1 for row in rows if row["status"] == "duplicate"),
            "error": sum(1 for row in rows if row["status"] == "error"),
        }
        return {"rows": rows, **summary}

    def import_online_registrations_csv(self, tournament_id: int, file_path: str | Path) -> dict[str, Any]:
        preview = self.preview_online_registrations_csv(tournament_id, file_path)
        imported = 0
        skipped = 0
        imported_player_ids = []
        errors = []
        for row in preview["rows"]:
            if row["status"] != "ready":
                skipped += 1
                if row["status"] == "error":
                    errors.append(f"Linha {row['line']}: {row['message']}")
                continue
            payload = row["payload"]
            imported_player_ids.append(self.db.create_player(tournament_id=tournament_id, **payload))
            imported += 1

        logger.info(
            "%s inscricoes online importadas de %s para o torneio %s; %s ignoradas",
            imported,
            file_path,
            tournament_id,
            skipped,
        )
        return {
            **preview,
            "imported": imported,
            "skipped": skipped,
            "errors": errors,
            "player_ids": imported_player_ids,
        }

    def import_players_csv(self, tournament_id: int, file_path: str | Path) -> dict[str, Any]:
        path = Path(file_path)
        imported = 0
        errors: list[str] = []

        with path.open("r", encoding="utf-8-sig", newline="") as file:
            sample = file.read(4096)
            file.seek(0)
            reader = csv.DictReader(file, dialect=self._csv_dialect(sample))
            if not reader.fieldnames:
                raise AppError("CSV sem cabecalho.")

            for line_number, row in enumerate(reader, start=2):
                name = self._pick(row, "name", "nome", "jogador")
                if not name:
                    errors.append(f"Linha {line_number}: nome vazio.")
                    continue

                national_rating = self._parse_optional_int(
                    self._pick(row, "national_rating", "rating_nacional", "elo_nacional", "cbx_rating")
                )
                international_rating = self._parse_optional_int(
                    self._pick(row, "international_rating", "rating_internacional", "elo_fide", "fide_rating")
                )
                rating_text = self._pick(row, "rating", "elo", "rtg")
                if rating_text:
                    try:
                        rating = int(float(rating_text))
                    except ValueError:
                        errors.append(f"Linha {line_number}: rating invalido.")
                        continue
                else:
                    rating = max(national_rating, international_rating)

                self.db.create_player(
                    tournament_id=tournament_id,
                    name=name,
                    club=self._pick(row, "club", "clube", "cidade"),
                    rating=rating,
                    category=self._pick(row, "category", "categoria"),
                    federation_id=self._pick(row, "federation_id", "id_federacao"),
                    fide_id=self._pick(row, "fide_id", "fide", "id_fide"),
                    cbx_id=self._pick(row, "cbx_id", "cbx", "id_cbx"),
                    birth_date=self._pick(row, "birth_date", "nascimento", "data_nascimento"),
                    surname=self._pick(row, "surname", "sobrenome"),
                    given_name=self._pick(row, "given_name", "nome_proprio"),
                    title=self._pick(row, "title", "titulo"),
                    sex=self._pick(row, "sex", "sexo"),
                    national_rating=national_rating,
                    international_rating=international_rating,
                )
                imported += 1

        logger.info("%s jogadores importados de %s para o torneio %s", imported, path, tournament_id)
        return {"imported": imported, "errors": errors}

    def import_rating_list_csv(self, file_path: str | Path, rating_type: str = "fide") -> dict[str, Any]:
        """Importa um CSV de ratings FIDE ou CBX e atualiza os membros do clube correspondentes."""
        path = Path(file_path)
        updated = 0
        skipped = 0
        errors: list[str] = []

        with path.open("r", encoding="utf-8-sig", newline="") as file:
            sample = file.read(4096)
            file.seek(0)
            reader = csv.DictReader(file, dialect=self._csv_dialect(sample))
            if not reader.fieldnames:
                raise AppError("CSV sem cabecalho.")
                
            all_members = self.db.list_members()
            fide_map = {str(m.get("fide_id")).strip(): m["id"] for m in all_members if m.get("fide_id")}
            cbx_map = {str(m.get("cbx_id")).strip(): m["id"] for m in all_members if m.get("cbx_id")}

            for line_number, row in enumerate(reader, start=2):
                player_id_csv = self._pick(row, "fide_id", "fide", "id_fide", "idnumber", "id", "cbx_id", "cbx", "id_cbx")
                if not player_id_csv:
                    errors.append(f"Linha {line_number}: ID (FIDE/CBX) nao encontrado.")
                    skipped += 1
                    continue
                    
                player_id_csv = str(player_id_csv).strip()
                rating_text = self._pick(row, "rating", "elo", "rtg", "standard", "rating_fide", "rating_cbx")
                if not rating_text:
                    errors.append(f"Linha {line_number}: Rating vazio.")
                    skipped += 1
                    continue
                    
                try:
                    rating = int(float(rating_text))
                except ValueError:
                    errors.append(f"Linha {line_number}: Rating invalido ({rating_text}).")
                    skipped += 1
                    continue
                    
                member_id = None
                if rating_type == "fide":
                    member_id = fide_map.get(player_id_csv)
                else:
                    member_id = cbx_map.get(player_id_csv)
                    
                if not member_id:
                    skipped += 1
                    continue
                    
                member = self.db.get_member(member_id)
                if not member:
                    continue
                    
                payload = dict(member)
                if rating_type == "fide":
                    payload["international_rating"] = rating
                else:
                    payload["national_rating"] = rating
                    
                self.db.update_member(member_id, payload)
                updated += 1
                
        logger.info("%s membros atualizados via importacao de rating %s de %s", updated, rating_type, path)
        return {"updated": updated, "skipped": skipped, "errors": errors}

    def _online_registration_rows(self, tournament_id: int, file_path: str | Path) -> list[dict[str, Any]]:
        path = Path(file_path)
        existing_players = self.db.list_players(tournament_id, active_only=False)
        rows = []
        seen_keys: set[tuple[str, str]] = set()

        with path.open("r", encoding="utf-8-sig", newline="") as file:
            sample = file.read(4096)
            file.seek(0)
            dialect = self._csv_dialect(sample)
            reader = csv.DictReader(file, dialect=dialect)
            if not reader.fieldnames:
                raise AppError("CSV sem cabecalho.")

            for line_number, row in enumerate(reader, start=2):
                parsed = self._online_registration_payload(row)
                status = "ready"
                status_label = "Pronto"
                message = "Pronto para importar"
                payload = parsed.get("payload") or {}

                if parsed["errors"]:
                    status = "error"
                    status_label = "Erro"
                    message = "; ".join(parsed["errors"])
                else:
                    duplicate_key = self._registration_duplicate_key(payload)
                    if self._matches_existing_player(payload, existing_players):
                        status = "duplicate"
                        status_label = "Duplicado"
                        message = "Ja existe jogador equivalente no torneio"
                    elif duplicate_key in seen_keys:
                        status = "duplicate"
                        status_label = "Duplicado"
                        message = "Inscricao repetida no proprio CSV"
                    else:
                        seen_keys.add(duplicate_key)

                rows.append(
                    {
                        "line": line_number,
                        "status": status,
                        "status_label": status_label,
                        "message": message,
                        "name": payload.get("name", parsed.get("name", "")),
                        "birth_date": payload.get("birth_date", ""),
                        "rating": payload.get("rating", ""),
                        "club": payload.get("club", ""),
                        "category": payload.get("category", ""),
                        "fide_id": payload.get("fide_id", ""),
                        "cbx_id": payload.get("cbx_id", ""),
                        "payload": payload,
                    }
                )
        return rows

    def _online_registration_payload(self, row: dict[str, Any]) -> dict[str, Any]:
        errors = []
        name = self._pick(
            row,
            "name",
            "nome",
            "jogador",
            "nome completo",
            "nome completo do jogador",
            "nome do jogador",
        )
        if not name:
            return {"name": "", "payload": {}, "errors": ["nome vazio"]}

        national_rating = self._parse_optional_int(
            self._pick(
                row,
                "national_rating",
                "rating nacional",
                "rating_nacional",
                "elo nacional",
                "elo_nacional",
                "cbx_rating",
            )
        )
        international_rating = self._parse_optional_int(
            self._pick(
                row,
                "international_rating",
                "rating internacional",
                "rating_internacional",
                "elo fide",
                "elo_fide",
                "fide_rating",
            )
        )
        rating_text = self._pick(row, "rating", "rating principal", "elo", "rtg")
        try:
            rating = int(float(rating_text)) if rating_text else max(national_rating, international_rating)
        except ValueError:
            rating = 0
            errors.append("rating invalido")

        birth_date = self._pick(row, "birth_date", "nascimento", "data nascimento", "data de nascimento")
        sex = self._pick(row, "sex", "sexo", "genero", "genero do jogador", "gênero")
        payload = {
            "name": name,
            "club": self._pick(row, "club", "clube", "cidade", "clube cidade", "clube / cidade"),
            "rating": rating,
            "category": self._pick(row, "category", "categoria", "categoria pretendida"),
            "federation_id": self._pick(row, "federation_id", "id federacao", "id_federacao"),
            "fide_id": self._pick(row, "fide_id", "fide", "id fide", "id_fide", "fide id"),
            "cbx_id": self._pick(row, "cbx_id", "cbx", "id cbx", "id_cbx", "cbx id"),
            "birth_date": birth_date,
            "surname": self._pick(row, "surname", "sobrenome"),
            "given_name": self._pick(row, "given_name", "nome proprio", "nome_proprio"),
            "title": self._pick(row, "title", "titulo", "titulo fide"),
            "sex": sex,
            "national_rating": national_rating,
            "international_rating": international_rating,
        }
        return {"name": name, "payload": payload, "errors": errors}

    @staticmethod
    def _csv_dialect(sample: str) -> Any:
        try:
            return csv.Sniffer().sniff(sample, delimiters=",;")
        except csv.Error:
            first_line = sample.splitlines()[0] if sample.splitlines() else ""
            if first_line.count(";") > first_line.count(","):
                class SemicolonDialect(csv.excel):
                    delimiter = ";"

                return SemicolonDialect
            return csv.get_dialect("excel")

    @staticmethod
    def _registration_duplicate_key(payload: dict[str, Any]) -> tuple[str, str]:
        fide_id = str(payload.get("fide_id") or "").strip()
        if fide_id:
            return "fide", fide_id.casefold()
        cbx_id = str(payload.get("cbx_id") or "").strip()
        if cbx_id:
            return "cbx", cbx_id.casefold()
        name = ImportService._normalize_text(str(payload.get("name") or ""))
        birth_date = str(payload.get("birth_date") or "").strip()
        if birth_date:
            return "name_birth", f"{name}|{birth_date}"
        club = ImportService._normalize_text(str(payload.get("club") or ""))
        return "name_club", f"{name}|{club}"

    @classmethod
    def _matches_existing_player(
        cls,
        payload: dict[str, Any],
        existing_players: list[dict[str, Any]],
    ) -> bool:
        candidate_key = cls._registration_duplicate_key(payload)
        for player in existing_players:
            player_key = cls._registration_duplicate_key(player)
            if candidate_key == player_key:
                return True
        return False

    @staticmethod
    def _pick(row: dict[str, Any], *keys: str) -> str:
        normalized = {ImportService._normalize_key(key): value for key, value in row.items()}
        for key in keys:
            value = normalized.get(ImportService._normalize_key(key))
            if value is not None:
                return str(value).strip()
        return ""

    @staticmethod
    def _normalize_key(value: str) -> str:
        text = ImportService._normalize_text(value)
        return "".join(character for character in text if character.isalnum())

    @staticmethod
    def _normalize_text(value: str) -> str:
        decomposed = unicodedata.normalize("NFKD", value.strip().casefold())
        return "".join(character for character in decomposed if not unicodedata.combining(character))

    @staticmethod
    def _parse_optional_int(value: str) -> int:
        if not value:
            return 0
        try:
            return int(float(value))
        except ValueError:
            return 0


class OfficialRatingService:
    SOURCES = {"FIDE", "CBX"}

    def __init__(self, db: Database) -> None:
        self.db = db

    def import_official_csv(
        self,
        file_path: str | Path,
        source: str,
        list_date: str = "",
    ) -> dict[str, Any]:
        source = source.strip().upper()
        if source not in self.SOURCES:
            raise AppError("Fonte invalida. Use FIDE ou CBX.")

        path = Path(file_path)
        errors: list[str] = []
        payloads: list[dict[str, Any]] = []

        with path.open("r", encoding="utf-8-sig", newline="") as file:
            sample = file.read(4096)
            file.seek(0)
            reader = csv.DictReader(file, dialect=ImportService._csv_dialect(sample))
            if not reader.fieldnames:
                raise AppError("CSV sem cabecalho.")

            for line_number, row in enumerate(reader, start=2):
                payload = self._official_payload(row, source)
                if not payload["name"]:
                    errors.append(f"Linha {line_number}: nome vazio.")
                    continue
                if not payload["external_id"]:
                    errors.append(f"Linha {line_number}: ID {source} vazio.")
                    continue
                payloads.append(payload)

        if payloads:
            snapshot_id = self.db.create_official_rating_snapshot_with_players(
                source=source,
                list_date=list_date,
                file_name=path.name,
                players=payloads,
            )
        else:
            snapshot_id = None

        logger.info("%s jogadores oficiais importados de %s (%s)", len(payloads), path, source)
        return {
            "snapshot_id": snapshot_id,
            "source": source,
            "imported": len(payloads),
            "errors": errors,
        }

    def update_tournament_players(self, tournament_id: int) -> dict[str, Any]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        settings = self.db.get_tournament_settings(tournament_id) or {}

        updated = 0
        unmatched = []
        for player in self.db.list_players(tournament_id, active_only=False):
            official = self._merged_official_player(player)
            if not official:
                unmatched.append(player["name"])
                continue

            national_rating = int(official.get("national_rating") or 0) or int(player.get("national_rating") or 0)
            international_rating = (
                int(official.get("international_rating") or 0)
                or int(official.get("standard_rating") or 0)
                or int(player.get("international_rating") or 0)
            )
            rating = self._rating_for_order(
                settings.get("initial_order", "rating"),
                current=int(player.get("rating") or 0),
                national=national_rating,
                international=international_rating,
            )
            self.db.update_player_official_data(
                int(player["id"]),
                {
                    "name": official.get("name") or player["name"],
                    "surname": official.get("surname") or player.get("surname", ""),
                    "given_name": official.get("given_name") or player.get("given_name", ""),
                    "title": official.get("title") or player.get("title", ""),
                    "sex": official.get("sex") or player.get("sex", ""),
                    "club": official.get("club") or player.get("club", ""),
                    "federation_id": official.get("federation") or player.get("federation_id", ""),
                    "fide_id": official.get("fide_id") or player.get("fide_id", ""),
                    "cbx_id": official.get("cbx_id") or player.get("cbx_id", ""),
                    "birth_date": official.get("birth_date") or player.get("birth_date", ""),
                    "national_rating": national_rating,
                    "international_rating": international_rating,
                    "rating": rating,
                },
            )
            updated += 1

        logger.info(
            "%s jogadores atualizados por base oficial no torneio %s; %s sem correspondencia",
            updated,
            tournament_id,
            len(unmatched),
        )
        return {"updated": updated, "unmatched": unmatched}

    def _merged_official_player(self, player: dict[str, Any]) -> dict[str, Any] | None:
        fide = self.db.find_latest_official_player(fide_id=str(player.get("fide_id") or ""))
        cbx = self.db.find_latest_official_player(cbx_id=str(player.get("cbx_id") or ""))
        if not fide and not cbx:
            return None
        merged: dict[str, Any] = {}
        for source in [fide, cbx]:
            if not source:
                continue
            for key, value in source.items():
                if value not in (None, ""):
                    merged[key] = value
        if fide:
            for key in ["title", "sex", "federation", "birth_date", "international_rating", "standard_rating"]:
                if fide.get(key) not in (None, ""):
                    merged[key] = fide[key]
        if cbx:
            for key in ["club", "national_rating", "cbx_id"]:
                if cbx.get(key) not in (None, ""):
                    merged[key] = cbx[key]
        return merged

    @staticmethod
    def _official_payload(row: dict[str, Any], source: str) -> dict[str, Any]:
        pick = ImportService._pick
        parse_optional_int = ImportService._parse_optional_int
        fide_id = pick(row, "fide_id", "fide", "id_fide", "fideid")
        cbx_id = pick(row, "cbx_id", "cbx", "id_cbx", "cbxid")
        external_id = fide_id if source == "FIDE" else cbx_id
        external_id = external_id or pick(row, "id", "codigo", "code")

        surname = pick(row, "surname", "sobrenome", "last_name")
        given_name = pick(row, "given_name", "nome_proprio", "first_name")
        name = pick(row, "name", "nome", "jogador", "player")
        if not name:
            name = " ".join(part for part in [given_name, surname] if part).strip()

        standard_rating = parse_optional_int(pick(row, "standard_rating", "standard", "rating_standard", "std"))
        national_rating = parse_optional_int(
            pick(row, "national_rating", "rating_nacional", "elo_nacional", "cbx_rating")
        )
        international_rating = parse_optional_int(
            pick(row, "international_rating", "rating_internacional", "elo_fide", "fide_rating")
        )
        generic_rating = parse_optional_int(pick(row, "rating", "elo", "rtg"))
        if source == "FIDE":
            international_rating = international_rating or standard_rating or generic_rating
            standard_rating = standard_rating or international_rating
        else:
            national_rating = national_rating or generic_rating

        return {
            "external_id": external_id,
            "fide_id": fide_id,
            "cbx_id": cbx_id,
            "name": name,
            "surname": surname,
            "given_name": given_name,
            "title": pick(row, "title", "titulo"),
            "sex": pick(row, "sex", "sexo"),
            "federation": pick(row, "federation", "fed", "federacao", "pais"),
            "club": pick(row, "club", "clube", "cidade"),
            "birth_date": pick(row, "birth_date", "nascimento", "data_nascimento", "data_nac"),
            "national_rating": national_rating,
            "international_rating": international_rating,
            "standard_rating": standard_rating,
            "rapid_rating": parse_optional_int(pick(row, "rapid_rating", "rapid", "rapido")),
            "blitz_rating": parse_optional_int(pick(row, "blitz_rating", "blitz")),
        }

    @staticmethod
    def _rating_for_order(
        initial_order: str,
        current: int,
        national: int,
        international: int,
    ) -> int:
        if initial_order == "national_rating":
            return national or current
        if initial_order == "international_rating":
            return international or current
        if initial_order == "international_then_national":
            return international or national or current
        if initial_order == "max_rating":
            return max(national, international, current)
        return max(national, international, current)


class InternalRatingService:
    HISTORY_REASON = "tournament_performance"

    def __init__(self, db: Database) -> None:
        self.db = db

    def apply_tournament_ratings(self, tournament_id: int) -> dict[str, Any]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")

        pairing_service = PairingService(self.db)
        updated = 0
        duplicates = 0
        skipped = 0
        external = 0
        no_member = 0

        for standing in pairing_service.standings(tournament_id):
            player = self.db.get_player(int(standing["player_id"]))
            if not player:
                skipped += 1
                continue
            member_id = int(player.get("member_id") or 0)
            if not member_id:
                external += 1
                continue

            performance = standing.get("performance")
            games = len(standing.get("earned_against") or [])
            if not isinstance(performance, int) or games <= 0:
                skipped += 1
                continue

            member = self.db.get_member(member_id)
            if not member:
                no_member += 1
                continue

            old_rating = int(member.get("rating") or 0)
            new_rating = self._next_rating(old_rating, int(performance), games)
            recorded = self.db.record_internal_rating_update(
                member_id=member_id,
                tournament_id=tournament_id,
                player_id=int(player["id"]),
                old_rating=old_rating,
                new_rating=new_rating,
                performance=int(performance),
                points=float(standing.get("points") or 0.0),
                games=games,
                reason=self.HISTORY_REASON,
            )
            if recorded:
                updated += 1
            else:
                duplicates += 1

        logger.info(
            "Rating interno aplicado no torneio %s: %s atualizados, %s ja aplicados",
            tournament_id,
            updated,
            duplicates,
        )
        return {
            "updated": updated,
            "duplicates": duplicates,
            "skipped": skipped,
            "external": external,
            "no_member": no_member,
        }

    def member_rating_history(self, member_id: int) -> list[dict[str, Any]]:
        member = self.db.get_member(member_id)
        if not member:
            raise AppError("Membro nao encontrado.")
        return self.db.list_member_rating_history(member_id)

    def ranking(
        self,
        category: str = "",
        active_only: bool = True,
        club_id: int | None = None,
        class_id: int | None = None,
        start_date: str = "",
        end_date: str = "",
    ) -> list[dict[str, Any]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        self._validate_optional_date(start_date, "Data inicial da temporada invalida.")
        self._validate_optional_date(end_date, "Data final da temporada invalida.")
        if start_date and end_date and start_date > end_date:
            raise AppError("Data final da temporada nao pode ser anterior a data inicial.")

        member_service = MemberService(self.db)
        rows = []
        for member in self.db.list_members(active_only=active_only, club_id=club_id, class_id=class_id):
            if category and str(member.get("category") or "") != category:
                continue
            rating_history = self.db.list_member_rating_history(int(member["id"]))
            filtered_history = [
                item
                for item in rating_history
                if self._history_item_in_period(item, start_date=start_date, end_date=end_date)
            ]
            latest_rating = filtered_history[0] if filtered_history else None
            summary = self._member_score_summary(
                member_service,
                int(member["id"]),
                start_date=start_date,
                end_date=end_date,
            )
            old_rating = int(latest_rating["old_rating"] or 0) if latest_rating else int(member.get("rating") or 0)
            new_rating = int(member.get("rating") or 0)
            period_delta = sum(
                int(item.get("new_rating") or 0) - int(item.get("old_rating") or 0)
                for item in filtered_history
            )
            delta = period_delta if start_date or end_date else (new_rating - old_rating if latest_rating else 0)
            rows.append(
                {
                    "member_id": member["id"],
                    "name": member["name"],
                    "club_name": member.get("club_name") or "",
                    "class_name": member.get("active_class_name") or "",
                    "category": member.get("category") or "",
                    "age_category": member.get("age_category") or "",
                    "rating_category": member.get("rating_category") or "",
                    "prize_tags": member.get("prize_tags") or "",
                    "member_type": member.get("member_type") or "",
                    "status": member.get("status") or "",
                    "rating": new_rating,
                    "last_delta": delta,
                    "last_performance": latest_rating["performance"] if latest_rating else "",
                    "last_tournament": latest_rating.get("tournament_name") if latest_rating else "",
                    "season_delta": period_delta,
                    **summary,
                }
            )

        rows.sort(
            key=lambda item: (
                -int(item["rating"] or 0),
                -float(item["score_rate"] or 0.0),
                -int(item["games"] or 0),
                str(item["name"]).casefold(),
            )
        )
        for position, row in enumerate(rows, start=1):
            row["position"] = position
        return rows

    def category_rankings(
        self,
        active_only: bool = True,
        club_id: int | None = None,
        class_id: int | None = None,
        start_date: str = "",
        end_date: str = "",
    ) -> dict[str, list[dict[str, Any]]]:
        categories = sorted(
            {
                str(member.get("category") or "").strip()
                for member in self.db.list_members(active_only=active_only, club_id=club_id, class_id=class_id)
                if str(member.get("category") or "").strip()
            },
            key=lambda value: value.casefold(),
        )
        return {
            category: self.ranking(
                category=category,
                active_only=active_only,
                club_id=club_id,
                class_id=class_id,
                start_date=start_date,
                end_date=end_date,
            )
            for category in categories
        }

    def _member_score_summary(
        self,
        member_service: MemberService,
        member_id: int,
        start_date: str = "",
        end_date: str = "",
    ) -> dict[str, Any]:
        tournaments = member_service.tournament_history(member_id)
        wins = draws = losses = byes = games = 0
        points = 0.0
        for tournament in tournaments:
            tournament_date = str(tournament.get("start_date") or "")
            if not self._date_in_period(tournament_date, start_date, end_date):
                continue
            for result in member_service.tournament_results(member_id, int(tournament["tournament_id"])):
                if result["round_status"] != "closed" or result["outcome"] == "Pendente":
                    continue
                outcome = result["outcome"]
                if outcome == "Bye":
                    byes += 1
                    continue
                if outcome == "Vitoria":
                    wins += 1
                    games += 1
                elif outcome == "Empate":
                    draws += 1
                    games += 1
                elif outcome in {"Derrota", "Duplo WO"}:
                    losses += 1
                    games += 1
                if result["points"] is not None and outcome != "Bye":
                    points += float(result["points"] or 0.0)
        return {
            "tournaments": len(tournaments),
            "games": games,
            "wins": wins,
            "draws": draws,
            "losses": losses,
            "byes": byes,
            "points": round(points, 2),
            "score_rate": round((points / games) * 100, 1) if games else 0.0,
        }

    @staticmethod
    def _validate_optional_date(value: str, message: str) -> None:
        if not value:
            return
        try:
            date.fromisoformat(value)
        except ValueError as exc:
            raise AppError(message) from exc

    @staticmethod
    def _date_in_period(value: str, start_date: str, end_date: str) -> bool:
        if not start_date and not end_date:
            return True
        if not value:
            return False
        day = value[:10]
        if start_date and day < start_date:
            return False
        if end_date and day > end_date:
            return False
        return True

    @classmethod
    def _history_item_in_period(
        cls,
        item: Mapping[str, Any],
        start_date: str,
        end_date: str,
    ) -> bool:
        tournament_date = str(item.get("tournament_start_date") or "").strip()
        if tournament_date:
            return cls._date_in_period(tournament_date, start_date, end_date)
        return cls._date_in_period(str(item.get("created_at") or "")[:10], start_date, end_date)

    @staticmethod
    def _next_rating(current_rating: int, performance: int, games: int) -> int:
        if games <= 0:
            return int(current_rating or 0)
        if current_rating <= 0:
            return max(0, int(performance or 0))
        weight = min(0.40, max(0.08, games * 0.08))
        next_rating = current_rating + (performance - current_rating) * weight
        return max(0, int(round(next_rating)))


class CertificateService:
    def __init__(self, db: Database, pairing_service: PairingService) -> None:
        self.db = db
        self.pairing_service = pairing_service

    def list_templates(self, active_only: bool = True) -> list[dict[str, Any]]:
        return self.db.list_certificate_templates(active_only=active_only)

    def list_issuances(
        self,
        limit: int = 50,
        context_type: str = "",
        verification_code: str = "",
        recipient_name: str = "",
        include_revoked: bool = True,
    ) -> list[dict[str, Any]]:
        return self.db.list_certificate_issuances(
            limit=limit,
            context_type=context_type,
            verification_code=verification_code,
            recipient_name=recipient_name,
            include_revoked=include_revoked,
        )

    def verify_issuance(self, verification_code: str) -> dict[str, Any]:
        issuance = self.db.get_certificate_issuance_by_code(verification_code)
        if not issuance:
            raise AppError("Codigo de verificacao nao encontrado.")
        return issuance

    def revoke_issuance(self, issuance_id: int, notes: str = "") -> None:
        self.db.revoke_certificate_issuance(issuance_id, notes=notes)

    def export_verification_site(
        self,
        file_path: str | Path,
        include_revoked: bool = True,
        limit: int = 5000,
    ) -> Path:
        issuances = self.list_issuances(limit=limit, include_revoked=include_revoked)
        if not issuances:
            raise AppError("Nao ha diplomas emitidos para exportar.")

        path = Path(file_path)
        if path.suffix.lower() != ".html":
            path = path.with_suffix(".html")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(self._verification_site_html(issuances), encoding="utf-8")
        logger.info("Verificador de diplomas exportado em %s", path)
        return path

    def get_template(self, template_id: int) -> dict[str, Any]:
        template = self.db.get_certificate_template(template_id)
        if not template:
            raise AppError("Modelo de diploma nao encontrado.")
        return template

    def create_template(self, data: dict[str, Any]) -> int:
        payload = self._validated_template_payload(data)
        try:
            template_id = self.db.create_certificate_template(**payload)
        except sqlite3.IntegrityError as exc:
            raise AppError("Ja existe um modelo de diploma com este nome.") from exc
        logger.info("Modelo de diploma criado: %s", template_id)
        return template_id

    def update_template(self, template_id: int, data: dict[str, Any]) -> None:
        if not self.db.get_certificate_template(template_id):
            raise AppError("Modelo de diploma nao encontrado.")
        payload = self._validated_template_payload(data)
        try:
            self.db.update_certificate_template(template_id, **payload)
        except sqlite3.IntegrityError as exc:
            raise AppError("Ja existe um modelo de diploma com este nome.") from exc
        logger.info("Modelo de diploma atualizado: %s", template_id)

    def preview_template(
        self,
        tournament_id: int,
        template_id: int,
        category: str = "",
        top_n: Any = 0,
        player_ids: list[int] | None = None,
        by_category: bool = False,
    ) -> dict[str, str]:
        template = self.get_template(template_id)
        recipients = self.tournament_recipients(
            tournament_id,
            certificate_type=str(template["certificate_type"]),
            category=category,
            top_n=top_n,
            player_ids=player_ids,
            by_category=by_category,
        )
        recipient = recipients[0]
        return {
            "title": self._render_template_text(str(template.get("title_template") or ""), recipient),
            "body": self._render_template_text(str(template.get("body_template") or ""), recipient),
            "footer": self._render_template_text(str(template.get("footer_template") or ""), recipient),
            "signature_left": self._render_template_text(str(template.get("signature_left") or ""), recipient),
            "signature_right": self._render_template_text(str(template.get("signature_right") or ""), recipient),
        }

    def preview_template_data(
        self,
        tournament_id: int,
        template_data: dict[str, Any],
        category: str = "",
        top_n: Any = 0,
        player_ids: list[int] | None = None,
        by_category: bool = False,
    ) -> dict[str, str]:
        template = self._template_for_export(None, "participation", template_data=template_data)
        recipients = self.tournament_recipients(
            tournament_id,
            certificate_type=str(template["certificate_type"]),
            category=category,
            top_n=top_n,
            player_ids=player_ids,
            by_category=by_category,
        )
        recipient = recipients[0]
        return {
            "title": self._plain_rendered_text(str(template.get("title_template") or ""), recipient),
            "body": self._plain_rendered_text(str(template.get("body_template") or ""), recipient),
            "footer": self._plain_rendered_text(str(template.get("footer_template") or ""), recipient),
        }

    def preview_template_recipient(
        self,
        template_data: dict[str, Any],
        recipient: dict[str, Any],
    ) -> dict[str, str]:
        template = self._template_for_export(
            None,
            str(template_data.get("certificate_type", "participation") or "participation"),
            template_data=template_data,
        )
        return {
            "title": self._plain_rendered_text(str(template.get("title_template") or ""), recipient),
            "body": self._plain_rendered_text(str(template.get("body_template") or ""), recipient),
            "footer": self._plain_rendered_text(str(template.get("footer_template") or ""), recipient),
        }

    def tournament_recipients(
        self,
        tournament_id: int,
        certificate_type: str = "participation",
        category: str = "",
        top_n: Any = 0,
        player_ids: list[int] | None = None,
        by_category: bool = False,
    ) -> list[dict[str, Any]]:
        tournament = self._individual_tournament(tournament_id)
        if certificate_type not in CERTIFICATE_TYPES:
            raise AppError("Modelo de diploma invalido.")

        standings = self.pairing_service.standings(tournament_id)
        if not standings:
            raise AppError("Nao ha jogadores no torneio selecionado.")

        top_limit = self._parse_top_n(top_n)
        use_category_ranking = by_category or certificate_type == "category_award"
        selected_ids = {int(player_id) for player_id in player_ids} if player_ids is not None else None
        category_filter = str(category or "").strip()
        category_positions = self._category_positions(standings)
        recipients = []
        for standing in standings:
            recipient = self._recipient_payload(tournament, standing, category_positions)
            if selected_ids is not None:
                if int(recipient["player_id"]) not in selected_ids:
                    continue
            else:
                if category_filter and recipient["category"] != category_filter:
                    continue
                if use_category_ranking:
                    if top_limit and int(recipient["category_position"]) > top_limit:
                        continue
                elif top_limit and int(recipient["position"]) > top_limit:
                    continue
            recipients.append(recipient)

        if not recipients:
            raise AppError("Nenhum jogador encontrado para os filtros selecionados.")
        return recipients

    def tournament_categories(self, tournament_id: int) -> list[str]:
        tournament = self._individual_tournament(tournament_id)
        return sorted(
            {
                str(item.get("category") or "").strip()
                for item in self.pairing_service.standings(int(tournament["id"]))
                if str(item.get("category") or "").strip()
            },
            key=lambda value: value.casefold(),
        )

    def member_recipients(
        self,
        member_ids: list[int] | None = None,
        active_only: bool = True,
        class_id: int | None = None,
    ) -> list[dict[str, Any]]:
        selected_ids = {int(member_id) for member_id in member_ids} if member_ids is not None else None
        recipients = []
        for member in self.db.list_members(active_only=active_only, class_id=class_id):
            if selected_ids is not None and int(member["id"]) not in selected_ids:
                continue
            recipients.append(self._member_recipient_payload(member))
        if not recipients:
            raise AppError("Nenhum membro encontrado para os filtros selecionados.")
        return recipients

    def training_recipients(
        self,
        session_id: int,
        present_only: bool = True,
        member_ids: list[int] | None = None,
    ) -> list[dict[str, Any]]:
        session = self.db.get_training_session(session_id)
        if not session:
            raise AppError("Aula/turma nao encontrada.")
        selected_ids = {int(member_id) for member_id in member_ids} if member_ids is not None else None
        rows = self.db.list_session_attendance(session_id)
        if not rows:
            rows = self.db.list_session_members_for_attendance(session_id)
        recipients = []
        for row in rows:
            member_id = int(row.get("member_id") or row.get("id") or 0)
            raw_attendance_status = row.get("attendance_status")
            if raw_attendance_status is None and row.get("member_name"):
                raw_attendance_status = row.get("status")
            attendance_status = str(raw_attendance_status or "")
            if selected_ids is not None and member_id not in selected_ids:
                continue
            if present_only and attendance_status and attendance_status != "present":
                continue
            recipients.append(self._training_recipient_payload(session, row, attendance_status))
        if not recipients:
            raise AppError("Nenhum aluno encontrado para a aula/turma selecionada.")
        return recipients

    def event_recipients(
        self,
        event_id: int,
        member_ids: list[int] | None = None,
        active_only: bool = True,
    ) -> list[dict[str, Any]]:
        event = self.db.get_club_event(event_id)
        if not event:
            raise AppError("Evento nao encontrado.")
        selected_ids = {int(member_id) for member_id in member_ids} if member_ids is not None else None
        recipients = []
        for member in self.db.list_members(active_only=active_only, club_id=int(event.get("club_id") or 0) or None):
            if selected_ids is not None and int(member["id"]) not in selected_ids:
                continue
            recipients.append(self._event_recipient_payload(event, member))
        if not recipients:
            raise AppError("Nenhum membro encontrado para o evento selecionado.")
        return recipients

    def ranking_recipients(
        self,
        category: str = "",
        top_n: Any = 3,
        member_ids: list[int] | None = None,
    ) -> list[dict[str, Any]]:
        selected_ids = {int(member_id) for member_id in member_ids} if member_ids is not None else None
        top_limit = self._parse_top_n(top_n)
        ranking = InternalRatingService(self.db).ranking(category=category, active_only=True)
        recipients = []
        for item in ranking:
            if selected_ids is not None and int(item["member_id"]) not in selected_ids:
                continue
            if selected_ids is None and top_limit and int(item["position"]) > top_limit:
                continue
            recipients.append(self._ranking_recipient_payload(item))
        if not recipients:
            raise AppError("Nenhum membro encontrado no ranking interno.")
        return recipients

    def export_tournament_certificates(
        self,
        tournament_id: int,
        file_path: str | Path,
        certificate_type: str = "participation",
        template_id: int | None = None,
        template_data: dict[str, Any] | None = None,
        category: str = "",
        top_n: Any = 0,
        player_ids: list[int] | None = None,
        by_category: bool = False,
    ) -> dict[str, Any]:
        tournament = self._individual_tournament(tournament_id)
        template = self._template_for_export(template_id, certificate_type, template_data)
        certificate_type = str(template["certificate_type"])
        recipients = self.tournament_recipients(
            tournament_id,
            certificate_type=certificate_type,
            category=category,
            top_n=top_n,
            player_ids=player_ids,
            by_category=by_category,
        )
        path = Path(file_path)
        if path.suffix.lower() != ".pdf":
            path = path.with_suffix(".pdf")
        path.parent.mkdir(parents=True, exist_ok=True)
        return self._export_context_certificates(
            path,
            template,
            recipients,
            source_title=str(tournament.get("name") or ""),
            context_type="tournament",
            source_id=tournament_id,
        )

    def export_member_certificates(
        self,
        file_path: str | Path,
        template_id: int | None = None,
        template_data: dict[str, Any] | None = None,
        member_ids: list[int] | None = None,
        class_id: int | None = None,
    ) -> dict[str, Any]:
        template = self._template_for_export(template_id, "member_certificate", template_data)
        recipients = self.member_recipients(member_ids=member_ids, class_id=class_id)
        return self._export_context_certificates(
            file_path,
            template,
            recipients,
            "Membros",
            context_type="members",
            source_id=class_id,
        )

    def export_training_certificates(
        self,
        session_id: int,
        file_path: str | Path,
        template_id: int | None = None,
        template_data: dict[str, Any] | None = None,
        present_only: bool = True,
        member_ids: list[int] | None = None,
    ) -> dict[str, Any]:
        template = self._template_for_export(template_id, "training_participation", template_data)
        session = self.db.get_training_session(session_id)
        recipients = self.training_recipients(session_id, present_only=present_only, member_ids=member_ids)
        return self._export_context_certificates(
            file_path,
            template,
            recipients,
            str(session.get("title") if session else "Aula"),
            context_type="training",
            source_id=session_id,
        )

    def export_event_certificates(
        self,
        event_id: int,
        file_path: str | Path,
        template_id: int | None = None,
        template_data: dict[str, Any] | None = None,
        member_ids: list[int] | None = None,
    ) -> dict[str, Any]:
        template = self._template_for_export(template_id, "event_participation", template_data)
        event = self.db.get_club_event(event_id)
        recipients = self.event_recipients(event_id, member_ids=member_ids)
        return self._export_context_certificates(
            file_path,
            template,
            recipients,
            str(event.get("title") if event else "Evento"),
            context_type="event",
            source_id=event_id,
        )

    def export_ranking_certificates(
        self,
        file_path: str | Path,
        template_id: int | None = None,
        template_data: dict[str, Any] | None = None,
        category: str = "",
        top_n: Any = 3,
        member_ids: list[int] | None = None,
    ) -> dict[str, Any]:
        template = self._template_for_export(template_id, "internal_ranking", template_data)
        recipients = self.ranking_recipients(category=category, top_n=top_n, member_ids=member_ids)
        return self._export_context_certificates(
            file_path,
            template,
            recipients,
            "Ranking interno",
            context_type="ranking",
            source_id=None,
        )

    def _export_context_certificates(
        self,
        file_path: str | Path,
        template: dict[str, Any],
        recipients: list[dict[str, Any]],
        source_title: str,
        context_type: str,
        source_id: int | None,
    ) -> dict[str, Any]:
        path = Path(file_path)
        if path.suffix.lower() != ".pdf":
            path = path.with_suffix(".pdf")
        path.parent.mkdir(parents=True, exist_ok=True)
        issued_at = self.db.now()
        prepared_recipients = self._prepare_issuance_recipients(recipients, issued_at)
        temp_path = path.with_name(f".{path.stem}.{secrets.token_hex(4)}.tmp{path.suffix}")
        issuance_ids: list[int] = []
        try:
            self._write_certificates_pdf(temp_path, template, prepared_recipients, source_title=source_title)
            issuance_ids = self._record_certificate_issuances(
                prepared_recipients,
                template,
                context_type=context_type,
                source_id=source_id,
                source_title=source_title,
                file_path=path,
                issued_at=issued_at,
            )
            try:
                temp_path.replace(path)
            except Exception:
                self.db.delete_certificate_issuances(issuance_ids)
                raise
        except Exception:
            if temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError:
                    logger.exception("Falha ao remover PDF temporario: %s", temp_path)
            raise
        logger.info("Diplomas exportados em PDF: %s (%s paginas)", path, len(prepared_recipients))
        return {
            "path": path,
            "exported": len(prepared_recipients),
            "issuance_ids": issuance_ids,
            "verification_codes": [str(recipient["verification_code"]) for recipient in prepared_recipients],
        }

    def _prepare_issuance_recipients(
        self,
        recipients: list[dict[str, Any]],
        issued_at: str,
    ) -> list[dict[str, Any]]:
        used_codes: set[str] = set()
        prepared = []
        issued_at_label = self._date_label(issued_at[:10])
        for recipient in recipients:
            code = self._new_verification_code(used_codes)
            used_codes.add(code)
            prepared.append(
                {
                    **recipient,
                    "verification_code": code,
                    "issued_at": issued_at,
                    "issued_at_label": issued_at_label,
                }
            )
        return prepared

    def _new_verification_code(self, used_codes: set[str]) -> str:
        for _attempt in range(100):
            code = f"ALB-{secrets.token_hex(4).upper()}"
            if code in used_codes:
                continue
            if self.db.get_certificate_issuance_by_code(code):
                continue
            return code
        raise AppError("Nao foi possivel gerar um codigo unico para o diploma.")

    def _record_certificate_issuances(
        self,
        recipients: list[dict[str, Any]],
        template: dict[str, Any],
        context_type: str,
        source_id: int | None,
        source_title: str,
        file_path: Path,
        issued_at: str,
    ) -> list[int]:
        template_id = int(template.get("id") or 0) or None
        rows = []
        for recipient in recipients:
            rows.append(
                {
                    "verification_code": recipient.get("verification_code"),
                    "context_type": context_type,
                    "source_id": source_id,
                    "source_title": source_title,
                    "recipient_id": self._recipient_id_for_context(recipient, context_type),
                    "recipient_name": recipient.get("name", ""),
                    "recipient_category": recipient.get("category", ""),
                    "certificate_type": template.get("certificate_type", ""),
                    "template_id": template_id,
                    "template_name": template.get("name", ""),
                    "file_path": str(file_path),
                    "issued_at": issued_at,
                    "payload_json": json.dumps(recipient, ensure_ascii=True, default=str),
                }
            )
        return self.db.create_certificate_issuances(rows)

    @staticmethod
    def _verification_context_label(context_type: Any) -> str:
        labels = {
            "tournament": "Torneio",
            "members": "Membros/alunos",
            "training": "Aula/turma",
            "event": "Evento",
            "ranking": "Ranking interno",
        }
        text = str(context_type or "").strip()
        return labels.get(text, text)

    def _verification_site_html(self, issuances: list[dict[str, Any]]) -> str:
        rows = "".join(self._verification_site_row(item) for item in issuances)
        generated_at = self.db.now()
        active_count = sum(1 for item in issuances if not item.get("revoked"))
        revoked_count = len(issuances) - active_count
        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Verificador de diplomas Albericus</title>
  <style>
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      font-family: Arial, Helvetica, sans-serif;
      color: #172033;
      background: #f6f7fb;
    }}
    header {{
      background: #16213e;
      color: #fff;
      padding: 28px 36px;
    }}
    header p {{ margin: 6px 0 0; color: #d8dee9; }}
    main {{ max-width: 1180px; margin: 0 auto; padding: 24px; }}
    .summary {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
      gap: 12px;
      margin-bottom: 16px;
    }}
    .summary div, section {{
      background: #fff;
      border: 1px solid #d9deea;
      border-radius: 8px;
      padding: 16px;
    }}
    .summary strong, .summary span {{ display: block; }}
    .summary span {{ margin-top: 4px; color: #566175; }}
    input {{
      width: 100%;
      border: 1px solid #cbd3e1;
      border-radius: 6px;
      font-size: 15px;
      margin-bottom: 12px;
      padding: 10px 12px;
    }}
    table {{ width: 100%; border-collapse: collapse; font-size: 14px; }}
    th, td {{
      border-bottom: 1px solid #e1e6ef;
      padding: 9px 10px;
      text-align: left;
      white-space: nowrap;
    }}
    th {{ background: #eef2f7; }}
    .table-wrap {{ overflow-x: auto; }}
    .status-active {{ color: #166534; font-weight: 700; }}
    .status-revoked {{ color: #991b1b; font-weight: 700; }}
    footer {{ color: #667085; font-size: 13px; padding: 0 24px 28px; text-align: center; }}
  </style>
</head>
<body>
  <header>
    <h1>Verificador de diplomas Albericus</h1>
    <p>Consulta local de codigos emitidos pelo gerenciador de clube.</p>
  </header>
  <main>
    <div class="summary">
      <div><strong>Total emitido</strong><span>{len(issuances)}</span></div>
      <div><strong>Ativos</strong><span>{active_count}</span></div>
      <div><strong>Revogados</strong><span>{revoked_count}</span></div>
      <div><strong>Gerado em</strong><span>{self._escape_html(generated_at)}</span></div>
    </div>
    <section>
      <input id="search" type="search" placeholder="Buscar por codigo, nome, origem ou modelo">
      <div class="table-wrap">
        <table id="issuances">
          <thead>
            <tr>
              <th>Codigo</th>
              <th>Status</th>
              <th>Destinatario</th>
              <th>Contexto</th>
              <th>Origem</th>
              <th>Tipo</th>
              <th>Modelo</th>
              <th>Emissao</th>
            </tr>
          </thead>
          <tbody>{rows}</tbody>
        </table>
      </div>
    </section>
  </main>
  <footer>Este arquivo nao substitui validacao institucional; ele reflete o historico local exportado em {self._escape_html(generated_at)}.</footer>
  <script>
    const search = document.getElementById('search');
    const rows = Array.from(document.querySelectorAll('#issuances tbody tr'));
    search.addEventListener('input', () => {{
      const query = search.value.trim().toLowerCase();
      rows.forEach((row) => {{
        row.style.display = row.textContent.toLowerCase().includes(query) ? '' : 'none';
      }});
    }});
  </script>
</body>
</html>
"""

    def _verification_site_row(self, item: dict[str, Any]) -> str:
        revoked = bool(item.get("revoked"))
        status = "Revogado" if revoked else "Ativo"
        status_class = "status-revoked" if revoked else "status-active"
        certificate_type = str(item.get("certificate_type") or "")
        issued_at = str(item.get("issued_at") or "")
        issued_at_label = self._date_label(issued_at[:10]) if issued_at else ""
        return (
            "<tr>"
            f"<td>{self._escape_html(item.get('verification_code'))}</td>"
            f"<td class=\"{status_class}\">{status}</td>"
            f"<td>{self._escape_html(item.get('recipient_name'))}</td>"
            f"<td>{self._escape_html(self._verification_context_label(item.get('context_type')))}</td>"
            f"<td>{self._escape_html(item.get('source_title'))}</td>"
            f"<td>{self._escape_html(CERTIFICATE_TYPES.get(certificate_type, certificate_type))}</td>"
            f"<td>{self._escape_html(item.get('template_name'))}</td>"
            f"<td>{self._escape_html(issued_at_label or issued_at)}</td>"
            "</tr>"
        )

    @staticmethod
    def _escape_html(value: Any) -> str:
        return html.escape(str(value or ""))

    @staticmethod
    def _recipient_id_for_context(recipient: Mapping[str, Any], context_type: str) -> int | None:
        if context_type == "tournament":
            recipient_id = recipient.get("player_id")
        else:
            recipient_id = recipient.get("member_id")
        if recipient_id in (None, ""):
            return None
        return int(str(recipient_id))

    def _template_for_export(
        self,
        template_id: int | None,
        certificate_type: str,
        template_data: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        if template_data is not None:
            payload = self._validated_template_payload(template_data)
            payload.update({"id": 0})
            return payload
        if template_id:
            return self.get_template(template_id)
        templates = [
            template
            for template in self.db.list_certificate_templates(active_only=True)
            if template.get("certificate_type") == certificate_type
        ]
        if templates:
            return templates[0]
        for template in DEFAULT_CERTIFICATE_TEMPLATES:
            if template["certificate_type"] == certificate_type:
                fallback = dict(template)
                fallback.update({"id": 0, "active": 1})
                return fallback
        raise AppError("Modelo de diploma invalido.")

    def _individual_tournament(self, tournament_id: int) -> dict[str, Any]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") == "team":
            raise AppError("Diplomas para torneios por equipes ficam para uma etapa futura.")
        return tournament

    @staticmethod
    def _validated_template_payload(data: dict[str, Any]) -> dict[str, Any]:
        name = str(data.get("name", "")).strip()
        if not name:
            raise AppError("Informe o nome do modelo.")
        certificate_type = str(data.get("certificate_type", "participation")).strip() or "participation"
        if certificate_type not in CERTIFICATE_TYPES:
            raise AppError("Tipo de diploma invalido.")
        orientation = str(data.get("orientation", "landscape")).strip() or "landscape"
        if orientation not in CERTIFICATE_ORIENTATIONS:
            raise AppError("Orientacao do diploma invalida.")
        title_template = str(data.get("title_template", "")).strip()
        body_template = str(data.get("body_template", "")).strip()
        if not title_template:
            raise AppError("Informe o titulo do diploma.")
        if not body_template:
            raise AppError("Informe o texto principal do diploma.")
        logo_path = CertificateService._validated_optional_image_path(
            data.get("logo_path", ""),
            "Arquivo de logo nao encontrado.",
        )
        background_image_path = CertificateService._validated_optional_image_path(
            data.get("background_image_path", ""),
            "Imagem de fundo nao encontrada.",
        )
        secondary_logo_path = CertificateService._validated_optional_image_path(
            data.get("secondary_logo_path", ""),
            "Arquivo de logo secundario nao encontrado.",
        )
        return {
            "name": name,
            "certificate_type": certificate_type,
            "title_template": title_template,
            "body_template": body_template,
            "footer_template": str(data.get("footer_template", "")).strip(),
            "orientation": orientation,
            "signature_left": str(data.get("signature_left", "")).strip(),
            "signature_right": str(data.get("signature_right", "")).strip(),
            "logo_path": logo_path,
            "background_image_path": background_image_path,
            "background_opacity": CertificateService._validated_opacity(
                data.get("background_opacity", 0.18),
                "opacidade do fundo",
            ),
            "secondary_logo_path": secondary_logo_path,
            "primary_color": CertificateService._validated_hex_color(
                data.get("primary_color", "#1E3A8A"),
                "cor principal",
            ),
            "accent_color": CertificateService._validated_hex_color(
                data.get("accent_color", "#93C5FD"),
                "cor de destaque",
            ),
            "title_font_size": CertificateService._validated_font_size(
                data.get("title_font_size", 32),
                "titulo",
                18,
                48,
            ),
            "body_font_size": CertificateService._validated_font_size(
                data.get("body_font_size", 18),
                "texto principal",
                10,
                28,
            ),
            "footer_font_size": CertificateService._validated_font_size(
                data.get("footer_font_size", 10),
                "rodape",
                7,
                16,
            ),
            "active": 1 if data.get("active", 1) else 0,
        }

    @staticmethod
    def _validated_optional_image_path(value: Any, message: str) -> str:
        image_path = str(value or "").strip()
        if not image_path:
            return ""
        resolved_image = CertificateService._resolve_template_path(image_path)
        if not resolved_image.exists() or not resolved_image.is_file():
            raise AppError(message)
        return image_path

    @staticmethod
    def _validated_opacity(value: Any, label: str) -> float:
        if value in (None, ""):
            return 0.18
        try:
            opacity = float(str(value).replace(",", "."))
        except (TypeError, ValueError) as exc:
            raise AppError(f"Informe uma {label} valida.") from exc
        if opacity > 1 and opacity <= 100:
            opacity = opacity / 100
        if opacity < 0 or opacity > 1:
            raise AppError(f"A {label} deve ficar entre 0 e 1, ou entre 0 e 100%.")
        return round(opacity, 4)

    @staticmethod
    def _validated_hex_color(value: Any, label: str) -> str:
        text = str(value or "").strip() or "#000000"
        if not text.startswith("#"):
            text = f"#{text}"
        hex_digits = text[1:]
        if len(hex_digits) != 6 or any(character not in "0123456789abcdefABCDEF" for character in hex_digits):
            raise AppError(f"Informe uma {label} em hexadecimal, como #1E3A8A.")
        return f"#{hex_digits.upper()}"

    @staticmethod
    def _validated_font_size(value: Any, label: str, minimum: int, maximum: int) -> int:
        try:
            size = int(value)
        except (TypeError, ValueError) as exc:
            raise AppError(f"Tamanho da fonte do {label} invalido.") from exc
        if size < minimum or size > maximum:
            raise AppError(f"Tamanho da fonte do {label} deve ficar entre {minimum} e {maximum}.")
        return size

    @staticmethod
    def _resolve_template_path(path: str) -> Path:
        logo_path = Path(path).expanduser()
        if logo_path.is_absolute():
            return logo_path
        return BASE_DIR / logo_path

    @staticmethod
    def _parse_top_n(value: Any) -> int:
        if value in (None, ""):
            return 0
        try:
            top_n = int(value)
        except (TypeError, ValueError) as exc:
            raise AppError("Informe um limite numerico para o top N.") from exc
        if top_n < 0:
            raise AppError("O limite do top N nao pode ser negativo.")
        return top_n

    @staticmethod
    def _category_positions(standings: list[dict[str, Any]]) -> dict[int, int]:
        counters: dict[str, int] = {}
        positions = {}
        for standing in standings:
            category = str(standing.get("category") or "Sem categoria").strip() or "Sem categoria"
            counters[category] = counters.get(category, 0) + 1
            positions[int(standing["player_id"])] = counters[category]
        return positions

    @staticmethod
    def _recipient_payload(
        tournament: dict[str, Any],
        standing: dict[str, Any],
        category_positions: dict[int, int],
    ) -> dict[str, Any]:
        player_id = int(standing["player_id"])
        category = str(standing.get("category") or "Sem categoria").strip() or "Sem categoria"
        return {
            "player_id": player_id,
            "name": str(standing.get("name") or ""),
            "tournament": str(tournament.get("name") or ""),
            "location": str(tournament.get("location") or ""),
            "date_range": CertificateService._date_range_label(tournament),
            "position": int(standing.get("position") or 0),
            "position_label": CertificateService._ordinal_label(standing.get("position")),
            "category": category,
            "category_position": category_positions.get(player_id, 0),
            "category_position_label": CertificateService._ordinal_label(category_positions.get(player_id, 0)),
            "points": ExportService._format_report_number(standing.get("points", 0)),
            "rating": int(standing.get("rating") or 0),
            "club": str(standing.get("club") or ""),
        }

    @staticmethod
    def _member_recipient_payload(member: Mapping[str, Any]) -> dict[str, Any]:
        category = str(member.get("category") or "Sem categoria").strip() or "Sem categoria"
        today_label = CertificateService._date_label(date.today().isoformat())
        return {
            "member_id": int(member.get("id") or member.get("member_id") or 0),
            "name": CertificateService._member_display_name(member),
            "tournament": "",
            "event": "",
            "session": "",
            "location": str(member.get("city") or member.get("club_name") or ""),
            "date_range": today_label,
            "position": 0,
            "position_label": "",
            "category": category,
            "category_position": 0,
            "category_position_label": "",
            "points": "0",
            "rating": int(member.get("rating") or 0),
            "club": str(member.get("club_name") or ""),
            "class_name": str(member.get("active_class_name") or member.get("class_name") or ""),
            "member_type": str(member.get("member_type") or ""),
            "status": str(member.get("status") or ""),
            "learning_level": str(member.get("learning_level_name") or ""),
            "date": today_label,
        }

    @staticmethod
    def _training_recipient_payload(
        session: Mapping[str, Any],
        member: Mapping[str, Any],
        attendance_status: str,
    ) -> dict[str, Any]:
        payload = CertificateService._member_recipient_payload(
            {
                **dict(member),
                "id": member.get("member_id") or member.get("id"),
                "name": member.get("member_name") or member.get("name"),
                "category": member.get("member_category") or member.get("category"),
                "status": member.get("member_status") or member.get("status"),
                "club_name": member.get("club_name") or session.get("club_name"),
            }
        )
        payload.update(
            {
                "session_id": int(session.get("id") or 0),
                "session": str(session.get("title") or ""),
                "location": str(session.get("location") or payload.get("location") or ""),
                "date_range": CertificateService._date_label(session.get("session_date")),
                "club": str(session.get("club_name") or payload.get("club") or ""),
                "class_name": str(session.get("class_name") or member.get("active_class_name") or ""),
                "instructor": str(session.get("instructor") or ""),
                "session_type": str(session.get("session_type") or ""),
                "type_label": TRAINING_SESSION_TYPES.get(
                    str(session.get("session_type") or ""),
                    str(session.get("session_type") or ""),
                ),
                "status": ATTENDANCE_STATUSES.get(attendance_status, attendance_status or "Nao lancada"),
                "attendance_status": attendance_status,
                "date": CertificateService._date_label(session.get("session_date")),
            }
        )
        return payload

    @staticmethod
    def _event_recipient_payload(event: Mapping[str, Any], member: Mapping[str, Any]) -> dict[str, Any]:
        payload = CertificateService._member_recipient_payload(member)
        payload.update(
            {
                "event_id": int(event.get("id") or 0),
                "event": str(event.get("title") or ""),
                "location": str(event.get("location") or payload.get("location") or ""),
                "date_range": CertificateService._date_label(event.get("event_date")),
                "club": str(event.get("club_name") or payload.get("club") or ""),
                "event_type": str(event.get("event_type") or ""),
                "type_label": EVENT_TYPES.get(
                    str(event.get("event_type") or ""),
                    str(event.get("event_type") or ""),
                ),
                "status": EVENT_STATUSES.get(
                    str(event.get("status") or ""),
                    str(event.get("status") or ""),
                ),
                "date": CertificateService._date_label(event.get("event_date")),
            }
        )
        return payload

    @staticmethod
    def _ranking_recipient_payload(item: Mapping[str, Any]) -> dict[str, Any]:
        today_label = CertificateService._date_label(date.today().isoformat())
        category = str(item.get("category") or "Sem categoria").strip() or "Sem categoria"
        position = int(item.get("position") or 0)
        return {
            "member_id": int(item.get("member_id") or 0),
            "name": str(item.get("name") or ""),
            "tournament": "",
            "event": "",
            "session": "",
            "location": str(item.get("club_name") or ""),
            "date_range": today_label,
            "position": position,
            "position_label": CertificateService._ordinal_label(position),
            "category": category,
            "category_position": 0,
            "category_position_label": "",
            "points": ExportService._format_report_number(item.get("points", 0)),
            "rating": int(item.get("rating") or 0),
            "club": str(item.get("club_name") or ""),
            "class_name": str(item.get("class_name") or ""),
            "member_type": str(item.get("member_type") or ""),
            "status": str(item.get("status") or ""),
            "last_delta": int(item.get("last_delta") or 0),
            "last_performance": str(item.get("last_performance") or ""),
            "last_tournament": str(item.get("last_tournament") or ""),
            "games": int(item.get("games") or 0),
            "wins": int(item.get("wins") or 0),
            "draws": int(item.get("draws") or 0),
            "losses": int(item.get("losses") or 0),
            "score_rate": ExportService._format_report_number(item.get("score_rate", 0)),
            "date": today_label,
        }

    @staticmethod
    def _member_display_name(member: Mapping[str, Any]) -> str:
        name = str(member.get("name") or "").strip()
        surname = str(member.get("surname") or "").strip()
        if surname and surname.casefold() not in name.casefold().split():
            return f"{name} {surname}".strip()
        return name or surname

    @staticmethod
    def _date_range_label(tournament: dict[str, Any]) -> str:
        start_date = CertificateService._date_label(tournament.get("start_date"))
        end_date = CertificateService._date_label(tournament.get("end_date"))
        if start_date and end_date and start_date != end_date:
            return f"{start_date} a {end_date}"
        return start_date or end_date

    @staticmethod
    def _date_label(value: Any) -> str:
        text = str(value or "").strip()
        if not text:
            return ""
        try:
            parsed = date.fromisoformat(text)
        except ValueError:
            return text
        return parsed.strftime("%d/%m/%Y")

    @staticmethod
    def _ordinal_label(value: Any) -> str:
        try:
            number = int(value or 0)
        except (TypeError, ValueError):
            return ""
        return f"{number}o" if number else ""

    @staticmethod
    def _write_certificates_pdf(
        path: Path,
        template: dict[str, Any],
        recipients: list[dict[str, Any]],
        source_title: str = "",
    ) -> None:
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.lib.utils import ImageReader
            from reportlab.pdfgen import canvas
            from reportlab.platypus import Frame, Paragraph
        except ImportError as exc:
            raise AppError("Instale reportlab para exportar PDF.") from exc

        page_size = A4 if template.get("orientation") == "portrait" else landscape(A4)
        width, height = page_size
        document = canvas.Canvas(str(path), pagesize=page_size)
        primary_color = colors.HexColor(str(template.get("primary_color") or "#1E3A8A"))
        accent_color = colors.HexColor(str(template.get("accent_color") or "#93C5FD"))
        title_style = ParagraphStyle(
            "CertificateTitle",
            fontName="Helvetica-Bold",
            fontSize=int(template.get("title_font_size") or 32),
            leading=int(template.get("title_font_size") or 32) + 6,
            alignment=TA_CENTER,
            textColor=primary_color,
        )
        body_style = ParagraphStyle(
            "CertificateBody",
            fontName="Helvetica",
            fontSize=int(template.get("body_font_size") or 18),
            leading=int(template.get("body_font_size") or 18) + 10,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#1F2937"),
        )
        footer_style = ParagraphStyle(
            "CertificateFooter",
            fontName="Helvetica",
            fontSize=int(template.get("footer_font_size") or 10),
            leading=int(template.get("footer_font_size") or 10) + 4,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#475569"),
        )
        background = CertificateService._load_template_image(
            template,
            "background_image_path",
            "Imagem de fundo nao encontrada.",
            "Nao foi possivel carregar a imagem de fundo do diploma.",
            ImageReader,
        )
        logo = CertificateService._load_template_image(
            template,
            "logo_path",
            "Arquivo de logo nao encontrado.",
            "Nao foi possivel carregar o logo do diploma.",
            ImageReader,
        )
        secondary_logo = CertificateService._load_template_image(
            template,
            "secondary_logo_path",
            "Arquivo de logo secundario nao encontrado.",
            "Nao foi possivel carregar o logo secundario do diploma.",
            ImageReader,
        )
        background_opacity = CertificateService._validated_opacity(
            template.get("background_opacity", 0.18),
            "opacidade do fundo",
        )

        for recipient in recipients:
            CertificateService._draw_certificate_page(
                document,
                width,
                height,
                cm,
                colors,
                primary_color,
                accent_color,
                background,
                background_opacity,
                logo,
                secondary_logo,
                Frame,
                Paragraph,
                title_style,
                body_style,
                footer_style,
                source_title,
                template,
                recipient,
            )
            document.showPage()
        document.save()

    @staticmethod
    def _load_template_image(
        template: dict[str, Any],
        field: str,
        missing_message: str,
        load_message: str,
        image_reader_class: Any,
    ) -> Any | None:
        image_path = str(template.get(field) or "").strip()
        if not image_path:
            return None
        resolved_image = CertificateService._resolve_template_path(image_path)
        if not resolved_image.exists() or not resolved_image.is_file():
            raise AppError(missing_message)
        try:
            return image_reader_class(str(resolved_image))
        except Exception as exc:
            raise AppError(load_message) from exc

    @staticmethod
    def _draw_background_image(
        document: Any,
        background: Any,
        width: float,
        height: float,
        opacity: float,
    ) -> None:
        if not background or opacity <= 0:
            return
        image_width, image_height = background.getSize()
        if image_width <= 0 or image_height <= 0:
            return
        scale = max(width / image_width, height / image_height)
        draw_width = image_width * scale
        draw_height = image_height * scale
        x = (width - draw_width) / 2
        y = (height - draw_height) / 2
        document.saveState()
        if hasattr(document, "setFillAlpha"):
            document.setFillAlpha(opacity)
        if hasattr(document, "setStrokeAlpha"):
            document.setStrokeAlpha(opacity)
        document.drawImage(
            background,
            x,
            y,
            width=draw_width,
            height=draw_height,
            preserveAspectRatio=True,
            mask="auto",
        )
        document.restoreState()

    @staticmethod
    def _draw_logo(
        document: Any,
        logo: Any,
        x: float,
        y: float,
        max_width: float,
        max_height: float,
    ) -> None:
        image_width, image_height = logo.getSize()
        if image_width <= 0 or image_height <= 0:
            return
        scale = min(max_width / image_width, max_height / image_height)
        draw_width = image_width * scale
        draw_height = image_height * scale
        document.drawImage(
            logo,
            x,
            y + (max_height - draw_height) / 2,
            width=draw_width,
            height=draw_height,
            preserveAspectRatio=True,
            mask="auto",
        )

    @staticmethod
    def _draw_certificate_page(
        document: Any,
        width: float,
        height: float,
        cm: float,
        colors: Any,
        primary_color: Any,
        accent_color: Any,
        background: Any,
        background_opacity: float,
        logo: Any,
        secondary_logo: Any,
        frame_class: Any,
        paragraph_class: Any,
        title_style: Any,
        body_style: Any,
        footer_style: Any,
        source_title: str,
        template: dict[str, Any],
        recipient: dict[str, Any],
    ) -> None:
        CertificateService._draw_background_image(document, background, width, height, background_opacity)
        margin = 1.1 * cm
        document.setStrokeColor(primary_color)
        document.setLineWidth(3)
        document.rect(margin, margin, width - 2 * margin, height - 2 * margin)
        document.setStrokeColor(accent_color)
        document.setLineWidth(1)
        document.rect(margin + 0.28 * cm, margin + 0.28 * cm, width - 2 * (margin + 0.28 * cm), height - 2 * (margin + 0.28 * cm))
        document.setFillColor(accent_color)
        document.rect(margin + 0.52 * cm, height - margin - 0.64 * cm, width - 2 * (margin + 0.52 * cm), 0.08 * cm, stroke=0, fill=1)

        if logo:
            CertificateService._draw_logo(document, logo, 2.1 * cm, height - 3.3 * cm, 3.0 * cm, 1.7 * cm)
        if secondary_logo:
            CertificateService._draw_logo(
                document,
                secondary_logo,
                width - 5.1 * cm,
                height - 3.3 * cm,
                3.0 * cm,
                1.7 * cm,
            )

        title = CertificateService._render_template_text(str(template.get("title_template") or ""), recipient)
        title_frame = frame_class(2.3 * cm, height - 4.0 * cm, width - 4.6 * cm, 1.4 * cm, showBoundary=0)
        title_frame.addFromList([paragraph_class(title, title_style)], document)
        document.setFont("Helvetica", 13)
        document.setFillColor(primary_color)
        document.drawCentredString(width / 2, height - 3.8 * cm, source_title)

        body = CertificateService._render_template_text(str(template.get("body_template") or ""), recipient)
        body_bottom = 6.2 * cm if width > height else 8.5 * cm
        body_height = 7.3 * cm if width > height else 9.0 * cm
        body_frame = frame_class(2.8 * cm, body_bottom, width - 5.6 * cm, body_height, showBoundary=0)
        body_frame.addFromList([paragraph_class(body, body_style)], document)

        document.setStrokeColor(accent_color)
        document.setLineWidth(1)
        line_half = min(3.2 * cm, width * 0.18)
        left_center = width * 0.32
        right_center = width * 0.68
        document.line(left_center - line_half, 3.9 * cm, left_center + line_half, 3.9 * cm)
        document.line(right_center - line_half, 3.9 * cm, right_center + line_half, 3.9 * cm)
        document.setFillColor(primary_color)
        document.setFont("Helvetica", max(8, int(template.get("footer_font_size") or 10)))
        left_signature = CertificateService._plain_rendered_text(str(template.get("signature_left") or ""), recipient)
        right_signature = CertificateService._plain_rendered_text(str(template.get("signature_right") or ""), recipient)
        document.drawCentredString(left_center, 3.45 * cm, left_signature)
        document.drawCentredString(right_center, 3.45 * cm, right_signature)

        footer = CertificateService._render_template_text(str(template.get("footer_template") or ""), recipient)
        if not footer:
            footer = CertificateService._render_template_text("{local} - {periodo}", recipient)
        footer_frame = frame_class(2.8 * cm, 1.7 * cm, width - 5.6 * cm, 1.2 * cm, showBoundary=0)
        footer_frame.addFromList([paragraph_class(footer, footer_style)], document)
        verification_code = str(recipient.get("verification_code") or "").strip()
        if verification_code:
            document.setFillColor(colors.HexColor("#64748B"))
            document.setFont("Helvetica", max(7, int(template.get("footer_font_size") or 10) - 1))
            document.drawRightString(width - 1.45 * cm, 1.25 * cm, f"Codigo: {verification_code}")

    @staticmethod
    def _render_template_text(template: str, recipient: dict[str, Any]) -> str:
        if not template:
            return ""
        rendered = html.escape(template, quote=False).replace("\n", "<br/>")
        for key, value in CertificateService._template_variables(recipient).items():
            rendered = rendered.replace(f"{{{key}}}", CertificateService._escape_pdf_text(value))
        return rendered

    @staticmethod
    def _plain_rendered_text(template: str, recipient: dict[str, Any]) -> str:
        rendered = CertificateService._render_template_text(template, recipient)
        return html.unescape(rendered.replace("<br/>", " "))

    @staticmethod
    def _template_variables(recipient: dict[str, Any]) -> dict[str, str]:
        return {
            "nome": str(recipient.get("name") or ""),
            "torneio": str(recipient.get("tournament") or ""),
            "local": str(recipient.get("location") or ""),
            "periodo": str(recipient.get("date_range") or ""),
            "data": str(recipient.get("date_range") or ""),
            "posicao": str(recipient.get("position_label") or ""),
            "posicao_numero": str(recipient.get("position") or ""),
            "categoria": str(recipient.get("category") or ""),
            "posicao_categoria": str(recipient.get("category_position_label") or ""),
            "posicao_categoria_numero": str(recipient.get("category_position") or ""),
            "pontos": str(recipient.get("points") or "0"),
            "rating": str(recipient.get("rating") or ""),
            "clube": str(recipient.get("club") or ""),
            "turma": str(recipient.get("class_name") or ""),
            "aula": str(recipient.get("session") or ""),
            "evento": str(recipient.get("event") or ""),
            "instrutor": str(recipient.get("instructor") or ""),
            "professor": str(recipient.get("instructor") or ""),
            "tipo": str(recipient.get("type_label") or recipient.get("member_type") or ""),
            "status": str(recipient.get("status") or ""),
            "nivel": str(recipient.get("learning_level") or ""),
            "delta": str(recipient.get("last_delta") or ""),
            "jogos": str(recipient.get("games") or ""),
            "desempenho": str(recipient.get("last_performance") or ""),
            "aproveitamento": str(recipient.get("score_rate") or ""),
            "vitorias": str(recipient.get("wins") or ""),
            "empates": str(recipient.get("draws") or ""),
            "derrotas": str(recipient.get("losses") or ""),
            "ultimo_torneio": str(recipient.get("last_tournament") or ""),
            "codigo": str(recipient.get("verification_code") or ""),
            "codigo_verificacao": str(recipient.get("verification_code") or ""),
            "emissao": str(recipient.get("issued_at_label") or ""),
            "emitido_em": str(recipient.get("issued_at_label") or ""),
        }

    @staticmethod
    def _certificate_title(certificate_type: str) -> str:
        if certificate_type == "participation":
            return "Certificado de Participacao"
        return "Certificado de Premiacao"

    @staticmethod
    def _certificate_body(certificate_type: str, recipient: dict[str, Any]) -> str:
        name = CertificateService._escape_pdf_text(str(recipient.get("name") or ""))
        tournament_name = CertificateService._escape_pdf_text(str(recipient.get("tournament") or ""))
        points = CertificateService._escape_pdf_text(str(recipient.get("points") or "0"))
        position = CertificateService._escape_pdf_text(str(recipient.get("position_label") or ""))
        category = CertificateService._escape_pdf_text(str(recipient.get("category") or ""))
        category_position = CertificateService._escape_pdf_text(str(recipient.get("category_position_label") or ""))

        if certificate_type == "overall_award":
            return (
                f"Certificamos que <b>{name}</b> conquistou a <b>{position}</b> colocacao geral "
                f"no torneio <b>{tournament_name}</b>, com <b>{points}</b> ponto(s)."
            )
        if certificate_type == "category_award":
            return (
                f"Certificamos que <b>{name}</b> conquistou a <b>{category_position}</b> colocacao "
                f"na categoria <b>{category}</b> do torneio <b>{tournament_name}</b>, "
                f"com <b>{points}</b> ponto(s)."
            )
        return (
            f"Certificamos que <b>{name}</b> participou do torneio <b>{tournament_name}</b>, "
            f"obtendo <b>{points}</b> ponto(s) e a <b>{position}</b> colocacao na classificacao."
        )

    @staticmethod
    def _escape_pdf_text(value: str) -> str:
        return html.escape(value, quote=False)


class ExportService:
    def __init__(self, db: Database, pairing_service: PairingService) -> None:
        self.db = db
        self.pairing_service = pairing_service

    def export_players(self, tournament_id: int, file_path: str | Path) -> None:
        title, headers, rows = self._players_section(tournament_id)
        self._write_report(file_path, title, headers, rows)

    def export_pairings(self, round_id: int, file_path: str | Path) -> None:
        title, headers, rows = self._pairings_section(round_id)
        self._write_report(file_path, title, headers, rows)

    def export_teams(self, tournament_id: int, file_path: str | Path) -> None:
        self._write_multi_report(
            file_path,
            [
                self._teams_section(tournament_id),
                self._team_rosters_section(tournament_id),
            ],
        )

    def export_all_rounds(self, tournament_id: int, file_path: str | Path) -> None:
        sections = []
        for round_data in sorted(self.db.list_rounds(tournament_id), key=lambda item: item["number"]):
            sections.append(self._pairings_section(round_data["id"]))
        if not sections:
            raise AppError("Nao ha rodadas para exportar.")
        self._write_multi_report(file_path, sections)

    def export_standings(self, tournament_id: int, file_path: str | Path) -> None:
        title, headers, rows = self._standings_section(tournament_id)
        self._write_report(file_path, title, headers, rows)

    def export_complete(self, tournament_id: int, file_path: str | Path) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        sections = [
            self._tournament_section(tournament_id),
        ]
        if tournament.get("competition_type") == "team":
            sections.append(self._teams_section(tournament_id))
            sections.append(self._team_rosters_section(tournament_id))
        sections.append(self._players_section(tournament_id))
        for round_data in sorted(self.db.list_rounds(tournament_id), key=lambda item: item["number"]):
            sections.append(self._pairings_section(round_data["id"]))
        sections.append(self._standings_section(tournament_id))
        self._write_multi_report(file_path, sections)

    def export_pgn(self, tournament_id: int, file_path: str | Path) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
            
        path = Path(file_path)
        with path.open("w", encoding="utf-8") as f:
            for round_data in sorted(self.db.list_rounds(tournament_id), key=lambda r: r["number"]):
                pairings = self.db.list_pairings(round_data["id"])
                for p in pairings:
                    if p["is_bye"]:
                        continue
                    
                    white_name = p.get("white_name") or "Unknown"
                    black_name = p.get("black_name") or "Unknown"
                    result = p.get("result") or "*"
                    if result == "1-0":
                        res_str = "1-0"
                    elif result == "0-1":
                        res_str = "0-1"
                    elif result == "1/2-1/2":
                        res_str = "1/2-1/2"
                    else:
                        res_str = "*"
                        
                    f.write(f'[Event "{tournament["name"]}"]\n')
                    f.write(f'[Site "{tournament.get("location") or ""}"]\n')
                    f.write(f'[Date "{tournament.get("start_date") or ""}"]\n')
                    f.write(f'[Round "{round_data["number"]}"]\n')
                    f.write(f'[White "{white_name}"]\n')
                    f.write(f'[Black "{black_name}"]\n')
                    f.write(f'[Result "{res_str}"]\n')
                    f.write(f'\n{res_str}\n\n')

    def export_trf(self, tournament_id: int, file_path: str | Path) -> None:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
            
        players = self.db.list_players(tournament_id)
        players = sorted(players, key=lambda p: (-int(p.get("rating") or 0), str(p.get("name") or "").casefold()))
        
        path = Path(file_path)
        with path.open("w", encoding="utf-8") as f:
            f.write(f"012 {tournament['name']}\n")
            f.write(f"022 {tournament.get('location') or ''}\n")
            f.write(f"032 \n")
            start = str(tournament.get('start_date') or '').replace("-", "/")
            end = str(tournament.get('end_date') or '').replace("-", "/")
            f.write(f"042 {start}\n")
            f.write(f"052 {end}\n")
            f.write(f"062 {len(players)}\n")
            
            rounds = sorted(self.db.list_rounds(tournament_id), key=lambda r: r["number"])
            f.write("132 \n")
            
            player_id_to_seq = {p["id"]: i+1 for i, p in enumerate(players)}
            
            for i, p in enumerate(players):
                seq = i + 1
                name = str(p.get("name") or "")[:32].ljust(33)
                rating = str(p.get("rating") or "0").rjust(4)
                fide_id = str(p.get("fide_id") or "").rjust(11)
                
                line = f"{seq:4}           {name} {rating}       {fide_id}                      "
                
                for r in rounds:
                    pairings = self.db.list_pairings(r["id"])
                    my_pairing = next((pa for pa in pairings if pa["white_player_id"] == p["id"] or pa["black_player_id"] == p["id"]), None)
                    if not my_pairing:
                        line += "  0000 - - "
                        continue
                        
                    if my_pairing["is_bye"]:
                        if my_pairing["white_player_id"] == p["id"]:
                            line += "  0000 - + " 
                        else:
                            line += "  0000 - - "
                        continue
                        
                    is_white = my_pairing["white_player_id"] == p["id"]
                    color = "w" if is_white else "b"
                    opp_id = my_pairing["black_player_id"] if is_white else my_pairing["white_player_id"]
                    opp_seq = player_id_to_seq.get(opp_id, 0)
                    
                    res_str = " "
                    result = my_pairing.get("result")
                    if result == "1-0":
                        res_str = "1" if is_white else "0"
                    elif result == "0-1":
                        res_str = "0" if is_white else "1"
                    elif result == "1/2-1/2":
                        res_str = "="
                        
                    line += f"  {opp_seq:04d} {color} {res_str} "
                    
                f.write(f"{line}\n")

    def export_member_evolution(self, member_id: int, file_path: str | Path) -> None:
        sections = self._member_evolution_sections(member_id)
        self._write_multi_report(file_path, sections)

    def export_club_report(self, file_path: str | Path) -> None:
        sections = self._club_report_sections()
        self._write_multi_report(file_path, sections)

    def export_member_report(self, member_id: int, file_path: str | Path) -> None:
        sections = self._member_evolution_sections(member_id)
        self._write_multi_report(file_path, sections)

    def export_tournaments_period_report(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._tournaments_period_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_attendance_report(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._attendance_report_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_financial_report(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._financial_report_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_events_report(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._events_report_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_internal_ranking_report(
        self,
        file_path: str | Path,
        category: str = "",
        club_id: int | None = None,
        class_id: int | None = None,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._internal_ranking_sections(
            category,
            club_id=club_id,
            class_id=class_id,
            start_date=start_date,
            end_date=end_date,
        )
        self._write_multi_report(file_path, sections)

    def export_payment_receipt(self, payment_id: int, file_path: str | Path) -> Path:
        payment = self.db.get_payment(payment_id)
        if not payment:
            raise AppError("Lancamento financeiro nao encontrado.")
        payment = FinanceService._with_effective_status(payment)
        if payment["effective_status"] not in {"paid", "exempt"}:
            raise AppError("Recibo disponivel apenas para lancamentos pagos ou isentos.")

        path = Path(file_path)
        if path.suffix.lower() not in {".csv", ".xlsx", ".pdf"}:
            path = path.with_suffix(".pdf")
        rows = [
            ["Recibo", f"REC-{int(payment['id']):06d}"],
            ["Membro", payment.get("member_name") or ""],
            ["Clube/Escola", payment.get("club_name") or ""],
            ["Turma", payment.get("active_class_name") or ""],
            ["Plano", payment.get("plan_name") or ""],
            ["Descricao", payment.get("description") or ""],
            ["Referencia", payment.get("reference_period") or ""],
            ["Vencimento", payment.get("due_date") or ""],
            ["Pagamento", payment.get("payment_date") or ""],
            ["Valor", self._format_report_number(payment.get("amount") or 0.0)],
            ["Status", PAYMENT_STATUSES.get(payment["effective_status"], payment["effective_status"])],
            ["Metodo", payment.get("method") or ""],
            ["Observacoes", payment.get("notes") or ""],
            ["Gerado em", Database.now()],
        ]
        self._write_report(path, "Recibo financeiro", ["Campo", "Valor"], rows)
        return path

    def export_administrative_package(
        self,
        file_path: str | Path,
        start_date: str = "",
        end_date: str = "",
    ) -> None:
        sections = self._administrative_package_sections(start_date, end_date)
        self._write_multi_report(file_path, sections)

    def export_club_portal(
        self,
        output_dir: str | Path,
        club_id: int | None = None,
        class_id: int | None = None,
    ) -> Path:
        club, class_data = self._club_portal_scope(club_id, class_id)
        resolved_club_id = int(club["id"]) if club else None
        resolved_class_id = int(class_data["id"]) if class_data else None
        path = Path(output_dir)
        path.mkdir(parents=True, exist_ok=True)
        html_content = self._club_portal_html(
            club=club,
            class_data=class_data,
            members=self.db.list_members(active_only=True, club_id=resolved_club_id, class_id=resolved_class_id),
            classes=[class_data]
            if class_data
            else self.db.list_classes(club_id=resolved_club_id, active_only=True),
            events=self.db.list_club_events(
                start_date=date.today().isoformat(),
                club_id=resolved_club_id,
                include_canceled=False,
            )[:10],
            sessions=self.db.list_training_sessions(
                start_date=date.today().isoformat(),
                club_id=resolved_club_id,
                class_id=resolved_class_id,
            )[:10],
            tournaments=[
                tournament
                for tournament in self.db.list_tournaments()
                if self._portal_tournament_in_scope(tournament, resolved_club_id, resolved_class_id)
            ][:8],
            ranking=InternalRatingService(self.db).ranking(
                active_only=True,
                club_id=resolved_club_id,
                class_id=resolved_class_id,
            )[:10],
        )
        (path / "index.html").write_text(html_content, encoding="utf-8")
        (path / "styles.css").write_text(self._club_portal_css(), encoding="utf-8")
        logger.info("Portal estatico do clube exportado em %s", path)
        return path / "index.html"

    def export_site(self, tournament_id: int, output_dir: str | Path) -> Path:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")

        path = Path(output_dir)
        path.mkdir(parents=True, exist_ok=True)
        settings = self.db.get_tournament_settings(tournament_id) or {}
        schedule = self.db.list_round_schedule(tournament_id)
        players = self.db.list_players(tournament_id, active_only=False)
        standings = (
            []
            if settings.get("hide_standings")
            else self._standings_for_tournament(tournament_id, tournament)
        )
        rounds = sorted(self.db.list_rounds(tournament_id), key=lambda item: item["number"])

        index_html = self._site_html(tournament, settings, schedule, players, rounds, standings)
        (path / "index.html").write_text(index_html, encoding="utf-8")
        (path / "styles.css").write_text(self._site_css(), encoding="utf-8")
        logger.info("Site estatico exportado em %s", path)
        return path / "index.html"

    def _club_portal_scope(
        self,
        club_id: int | None,
        class_id: int | None,
    ) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
        class_data = self.db.get_class(class_id) if class_id else None
        if class_id and not class_data:
            raise AppError("Turma do portal nao encontrada.")
        if class_data and club_id and int(class_data["club_id"]) != int(club_id):
            raise AppError("A turma selecionada nao pertence ao clube/escola do portal.")

        if class_data and not club_id:
            club_id = int(class_data["club_id"])
        club = self.db.get_club(club_id) if club_id else None
        if club_id and not club:
            raise AppError("Clube/escola do portal nao encontrado.")
        if not club:
            active_clubs = self.db.list_clubs(active_only=True)
            club = active_clubs[0] if active_clubs else self.db.get_club(1)
        return club, class_data

    @staticmethod
    def _portal_tournament_in_scope(
        tournament: Mapping[str, Any],
        club_id: int | None,
        class_id: int | None,
    ) -> bool:
        if class_id and int(tournament.get("class_id") or 0) != class_id:
            return False
        if club_id and int(tournament.get("club_id") or 0) != club_id:
            return False
        return True

    def _club_portal_html(
        self,
        club: dict[str, Any] | None,
        class_data: dict[str, Any] | None,
        members: list[dict[str, Any]],
        classes: list[dict[str, Any]],
        events: list[dict[str, Any]],
        sessions: list[dict[str, Any]],
        tournaments: list[dict[str, Any]],
        ranking: list[dict[str, Any]],
    ) -> str:
        title = class_data.get("name") if class_data else club.get("name") if club else "Portal do clube"
        club_kind_labels = {
            "club": "Clube",
            "school": "Escola",
            "project": "Projeto",
            "partner": "Parceiro",
        }
        subtitle = (
            class_data.get("club_name")
            if class_data
            else club_kind_labels.get(str(club.get("kind") or "club"), "") if club else ""
        )
        generated_at = Database.now()
        notice_rows = self._portal_notice_rows(events, sessions)
        class_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('name') or '')}</td>"
            f"<td>{self._escape(item.get('teacher') or '')}</td>"
            f"<td>{self._escape(item.get('weekday') or '')}</td>"
            f"<td>{self._escape(item.get('time') or '')}</td>"
            f"<td>{self._escape(item.get('location') or '')}</td>"
            f"<td>{item.get('active_members_count', '')}</td>"
            "</tr>"
            for item in classes
        )
        event_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('event_date') or '')}</td>"
            f"<td>{self._escape(item.get('start_time') or '')}</td>"
            f"<td>{self._escape(EVENT_TYPES.get(item.get('event_type', ''), item.get('event_type') or ''))}</td>"
            f"<td>{self._escape(item.get('title') or '')}</td>"
            f"<td>{self._escape(item.get('location') or '')}</td>"
            f"<td>{self._escape(EVENT_STATUSES.get(item.get('status', ''), item.get('status') or ''))}</td>"
            "</tr>"
            for item in events
        )
        session_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('session_date') or '')}</td>"
            f"<td>{self._escape(item.get('start_time') or '')}</td>"
            f"<td>{self._escape(item.get('title') or '')}</td>"
            f"<td>{self._escape(item.get('class_name') or '')}</td>"
            f"<td>{self._escape(item.get('learning_level_name') or '')}</td>"
            f"<td>{self._escape(item.get('objective') or '')}</td>"
            "</tr>"
            for item in sessions
        )
        ranking_rows = "".join(
            "<tr>"
            f"<td>{item['position']}</td>"
            f"<td>{self._escape(item.get('name') or '')}</td>"
            f"<td>{item.get('rating', '')}</td>"
            f"<td>{self._escape(item.get('category') or '')}</td>"
            f"<td>{self._escape(item.get('class_name') or '')}</td>"
            f"<td>{self._escape(str(item.get('score_rate') or 0))}%</td>"
            "</tr>"
            for item in ranking
        )
        tournament_rows = "".join(
            "<tr>"
            f"<td>{self._escape(item.get('name') or '')}</td>"
            f"<td>{self._escape(item.get('start_date') or '')}</td>"
            f"<td>{self._escape(item.get('location') or '')}</td>"
            f"<td>{self._escape(item.get('status') or '')}</td>"
            "</tr>"
            for item in tournaments
        )
        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{self._escape(title)}</title>
  <link rel="stylesheet" href="styles.css">
</head>
<body>
  <header>
    <p class="eyebrow">Albericus</p>
    <h1>{self._escape(title)}</h1>
    <p>{self._escape(subtitle)}</p>
  </header>
  <main>
    <section class="summary">
      <div><strong>Membros ativos</strong><span>{len(members)}</span></div>
      <div><strong>Turmas</strong><span>{len(classes)}</span></div>
      <div><strong>Eventos futuros</strong><span>{len(events)}</span></div>
      <div><strong>Aulas futuras</strong><span>{len(sessions)}</span></div>
    </section>
    <section>
      <h2>Comunicados</h2>
      <table><thead><tr><th>Data</th><th>Comunicado</th></tr></thead><tbody>{notice_rows}</tbody></table>
    </section>
    <section>
      <h2>Calendario</h2>
      <h3>Eventos</h3>
      <table><thead><tr><th>Data</th><th>Hora</th><th>Tipo</th><th>Evento</th><th>Local</th><th>Status</th></tr></thead><tbody>{event_rows}</tbody></table>
      <h3>Aulas e treinos</h3>
      <table><thead><tr><th>Data</th><th>Hora</th><th>Aula/Treino</th><th>Turma</th><th>Nivel</th><th>Objetivo</th></tr></thead><tbody>{session_rows}</tbody></table>
    </section>
    <section>
      <h2>Turmas</h2>
      <table><thead><tr><th>Turma</th><th>Professor</th><th>Dia</th><th>Horario</th><th>Local</th><th>Alunos</th></tr></thead><tbody>{class_rows}</tbody></table>
    </section>
    <section>
      <h2>Ranking interno</h2>
      <table><thead><tr><th>Pos</th><th>Membro</th><th>Rating</th><th>Categoria</th><th>Turma</th><th>Aproveitamento</th></tr></thead><tbody>{ranking_rows}</tbody></table>
    </section>
    <section>
      <h2>Torneios recentes</h2>
      <table><thead><tr><th>Torneio</th><th>Data</th><th>Local</th><th>Status</th></tr></thead><tbody>{tournament_rows}</tbody></table>
    </section>
  </main>
  <footer>Gerado em {self._escape(generated_at)}</footer>
</body>
</html>
"""

    def _portal_notice_rows(
        self,
        events: list[dict[str, Any]],
        sessions: list[dict[str, Any]],
    ) -> str:
        notices: list[tuple[str, str]] = []
        for event in events[:5]:
            details = " - ".join(
                part
                for part in [
                    str(event.get("title") or ""),
                    str(event.get("location") or ""),
                    str(event.get("notes") or ""),
                ]
                if part
            )
            notices.append((str(event.get("event_date") or ""), details))
        for session in sessions[:5]:
            details = " - ".join(
                part
                for part in [
                    str(session.get("title") or ""),
                    str(session.get("class_name") or ""),
                    str(session.get("homework") or session.get("objective") or ""),
                ]
                if part
            )
            notices.append((str(session.get("session_date") or ""), details))
        notices.sort(key=lambda item: item[0])
        if not notices:
            return "<tr><td colspan=\"2\">Nenhum comunicado futuro cadastrado.</td></tr>"
        return "".join(
            "<tr>"
            f"<td>{self._escape(day)}</td>"
            f"<td>{self._escape(text)}</td>"
            "</tr>"
            for day, text in notices
        )

    @staticmethod
    def _club_portal_css() -> str:
        return """* { box-sizing: border-box; }
body {
  margin: 0;
  font-family: Arial, Helvetica, sans-serif;
  color: #18212f;
  background: #f5f7fb;
}
header {
  background: #17403a;
  color: #fff;
  padding: 30px 40px;
}
header h1 { margin: 4px 0 8px; font-size: 32px; }
header p { margin: 0; color: #dbe8e5; }
.eyebrow { text-transform: uppercase; letter-spacing: 0; font-size: 12px; }
main { max-width: 1180px; margin: 0 auto; padding: 24px; }
section {
  background: #fff;
  border: 1px solid #dfe5ed;
  border-radius: 8px;
  margin-bottom: 16px;
  padding: 18px;
  overflow-x: auto;
}
.summary {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(170px, 1fr));
  gap: 12px;
}
.summary div {
  border: 1px solid #dfe5ed;
  border-radius: 8px;
  padding: 12px;
}
.summary strong, .summary span { display: block; }
.summary span { margin-top: 4px; color: #546173; }
h2 { margin: 0 0 12px; font-size: 20px; }
h3 { margin: 16px 0 8px; font-size: 15px; color: #344154; }
table {
  width: 100%;
  border-collapse: collapse;
  font-size: 14px;
}
th, td {
  border-bottom: 1px solid #e4e9f1;
  padding: 8px 10px;
  text-align: left;
  white-space: nowrap;
}
th { background: #eef4f2; }
footer {
  color: #657386;
  font-size: 13px;
  padding: 8px 24px 32px;
  text-align: center;
}
"""

    def _site_html(
        self,
        tournament: dict[str, Any],
        settings: dict[str, Any],
        schedule: list[dict[str, Any]],
        players: list[dict[str, Any]],
        rounds: list[dict[str, Any]],
        standings: list[dict[str, Any]],
    ) -> str:
        tournament_id = int(tournament["id"])
        is_team_tournament = tournament.get("competition_type") == "team"
        round_sections = []
        for round_data in rounds:
            if is_team_tournament:
                round_sections.append(self._site_team_round_section(round_data))
            else:
                round_sections.append(self._site_individual_round_section(round_data))

        schedule_rows = "".join(
            "<tr>"
            f"<td>{item['round_number']}</td>"
            f"<td>{self._escape(item.get('date') or '')}</td>"
            f"<td>{self._escape(item.get('time') or '')}</td>"
            "</tr>"
            for item in schedule
            if item.get("date") or item.get("time")
        )
        players_rows = "".join(
            "<tr>"
            f"<td>{self._escape(player_full_name(player))}</td>"
            f"<td>{self._escape(player.get('title') or '')}</td>"
            f"<td>{self._escape(player.get('fide_id') or '')}</td>"
            f"<td>{self._escape(player.get('cbx_id') or '')}</td>"
            f"<td>{player['rating']}</td>"
            f"<td>{self._escape(player['club'])}</td>"
            f"<td>{self._escape(player['category'])}</td>"
            f"<td>{self._escape(player.get('age_category') or '')}</td>"
            f"<td>{self._escape(player.get('rating_category') or '')}</td>"
            f"<td>{self._escape(player.get('prize_tags') or '')}</td>"
            f"<td>{self._escape(PLAYER_STATUSES.get(player.get('player_status', 'active'), player.get('player_status', '')))}</td>"
            "</tr>"
            for player in players
        )
        standings_section = self._site_standings_section(settings, standings, is_team_tournament)
        team_sections = self._site_team_sections(tournament_id) if is_team_tournament else ""

        generated_at = Database.now()
        return f"""<!doctype html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{self._escape(tournament['name'])}</title>
  <link rel="stylesheet" href="styles.css">
</head>
<body>
  <header>
    <p class="eyebrow">Albericus</p>
    <h1>{self._escape(tournament['name'])}</h1>
    <p>{self._escape(tournament['location'])}</p>
  </header>
  <main>
    <section class="summary">
      <div><strong>Periodo</strong><span>{self._escape(tournament['start_date'])} - {self._escape(tournament['end_date'])}</span></div>
      <div><strong>Rodadas</strong><span>{tournament['rounds_count']}</span></div>
      <div><strong>Ritmo</strong><span>{self._escape(tournament['time_control'])}</span></div>
      <div><strong>Competicao</strong><span>{self._escape(COMPETITION_TYPES.get(tournament.get('competition_type', 'individual'), 'Individual'))}</span></div>
      <div><strong>Status</strong><span>{self._escape(tournament['status'])}</span></div>
    </section>
    <section>
      <h2>Informacoes</h2>
      <dl>
        <dt>FIDE Event-ID</dt><dd>{self._escape(settings.get('fide_event_id') or '')}</dd>
        <dt>Organizador</dt><dd>{self._escape(settings.get('organizer') or '')}</dd>
        <dt>Diretor</dt><dd>{self._escape(settings.get('director') or '')}</dd>
        <dt>Arbitro principal</dt><dd>{self._escape(settings.get('chief_arbiter') or '')}</dd>
        <dt>Categorias</dt><dd>{self._escape(settings.get('categories') or '')}</dd>
      </dl>
    </section>
    <section>
      <h2>Agenda</h2>
      <table><thead><tr><th>Rodada</th><th>Data</th><th>Hora</th></tr></thead><tbody>{schedule_rows}</tbody></table>
    </section>
    {standings_section}
    {team_sections}
    <section>
      <h2>Jogadores</h2>
      <table><thead><tr><th>Nome</th><th>Titulo</th><th>FIDE</th><th>CBX</th><th>Rating</th><th>Clube</th><th>Categoria</th><th>Idade</th><th>Rating cat.</th><th>Tags</th><th>Status</th></tr></thead><tbody>{players_rows}</tbody></table>
    </section>
    {''.join(round_sections)}
  </main>
  <footer>Gerado em {self._escape(generated_at)} - Torneio #{tournament_id}</footer>
</body>
</html>
"""

    def _site_individual_round_section(self, round_data: dict[str, Any]) -> str:
        pairings = self.db.get_pairings_for_round(round_data["id"])
        rows = []
        for pairing in pairings:
            black_name = "BYE" if pairing["is_bye"] else pairing_player_name(pairing, "black")
            black_rating = "" if pairing["is_bye"] else pairing["black_rating"]
            rows.append(
                "<tr>"
                f"<td>{pairing['board_number']}</td>"
                f"<td>{self._escape(pairing_player_name(pairing, 'white'))}</td>"
                f"<td>{pairing['white_rating']}</td>"
                f"<td>{self._escape(pairing['result'] or '')}</td>"
                f"<td>{self._escape(black_name or '')}</td>"
                f"<td>{self._escape(black_rating or '')}</td>"
                "</tr>"
            )
        return (
            "<section>"
            f"<h2>Rodada {round_data['number']}</h2>"
            f"<p>Status: {self._escape(round_data['status'])}</p>"
            "<table><thead><tr><th>Mesa</th><th>Brancas</th><th>Rating</th>"
            "<th>Resultado</th><th>Pretas</th><th>Rating</th></tr></thead>"
            f"<tbody>{''.join(rows)}</tbody></table>"
            "</section>"
        )

    def _site_team_round_section(self, round_data: dict[str, Any]) -> str:
        rows = []
        for match in self.db.list_team_matches_for_round(int(round_data["id"])):
            match_score = self._team_match_score(match)
            if match.get("is_bye"):
                rows.append(
                    "<tr>"
                    f"<td>{match['match_number']}</td>"
                    f"<td>{self._escape(match.get('white_team_name') or '')}</td>"
                    f"<td>{self._escape(match_score)}</td>"
                    "<td>BYE</td><td></td><td></td><td>BYE</td><td></td>"
                    "</tr>"
                )
                continue
            for board in self.db.list_team_boards(int(match["id"])):
                rows.append(
                    "<tr>"
                    f"<td>{match['match_number']}</td>"
                    f"<td>{self._escape(match.get('white_team_name') or '')}</td>"
                    f"<td>{self._escape(match_score)}</td>"
                    f"<td>{self._escape(match.get('black_team_name') or '')}</td>"
                    f"<td>{board['board_number']}</td>"
                    f"<td>{self._escape(self._team_board_player_name(board, 'white'))}</td>"
                    f"<td>{self._escape(board.get('result') or '')}</td>"
                    f"<td>{self._escape(self._team_board_player_name(board, 'black'))}</td>"
                    "</tr>"
                )
        return (
            "<section>"
            f"<h2>Rodada {round_data['number']}</h2>"
            f"<p>Status: {self._escape(round_data['status'])}</p>"
            "<table><thead><tr><th>Confronto</th><th>Equipe A</th><th>Placar</th>"
            "<th>Equipe B</th><th>Tab.</th><th>Brancas</th><th>Resultado</th><th>Pretas</th></tr></thead>"
            f"<tbody>{''.join(rows)}</tbody></table>"
            "</section>"
        )

    def _site_standings_section(
        self,
        settings: dict[str, Any],
        standings: list[dict[str, Any]],
        is_team_tournament: bool,
    ) -> str:
        if settings.get("hide_standings"):
            return "<section><h2>Classificacao</h2><p>Classificacao ocultada pela organizacao.</p></section>"

        if is_team_tournament:
            standings_rows = "".join(
                "<tr>"
                f"<td>{item['position']}</td>"
                f"<td>{self._escape(item['name'])}</td>"
                f"<td>{self._escape(item.get('club') or '')}</td>"
                f"<td>{self._format_report_number(item['match_points'])}</td>"
                f"<td>{self._format_report_number(item['game_points'])}</td>"
                f"<td>{item['wins']}</td>"
                f"<td>{item['draws']}</td>"
                f"<td>{item['losses']}</td>"
                f"<td>{self._format_report_number(item['buchholz'])}</td>"
                "</tr>"
                for item in standings
            )
            return (
                "<section>"
                "<h2>Classificacao por equipes</h2>"
                "<table><thead><tr><th>Pos</th><th>Equipe</th><th>Clube/Cidade</th><th>MP</th>"
                "<th>GP</th><th>V</th><th>E</th><th>D</th><th>Buchholz</th></tr></thead>"
                f"<tbody>{standings_rows}</tbody></table>"
                "</section>"
            )

        standings_rows = "".join(
            "<tr>"
            f"<td>{item['position']}</td>"
            f"<td>{self._escape(item['name'])}</td>"
            f"<td>{self._escape(item.get('category') or '')}</td>"
            f"<td>{item['points']}</td>"
            f"<td>{item['buchholz']}</td>"
            f"<td>{item['buchholz_median']}</td>"
            f"<td>{item['sonneborn_berger']}</td>"
            f"<td>{item['wins']}</td>"
            f"<td>{self._escape(item['performance'])}</td>"
            "</tr>"
            for item in standings
        )
        return (
            "<section>"
            "<h2>Classificacao</h2>"
            "<table><thead><tr><th>Pos</th><th>Jogador</th><th>Categoria</th><th>Pts</th>"
            "<th>Buchholz</th><th>Buchholz M</th><th>SB</th>"
            "<th>Vitorias</th><th>Perf.</th></tr></thead>"
            f"<tbody>{standings_rows}</tbody></table>"
            "</section>"
        )

    def _site_team_sections(self, tournament_id: int) -> str:
        team_rows = []
        roster_rows = []
        for team in self.db.list_teams(tournament_id, active_only=False):
            team_rows.append(
                "<tr>"
                f"<td>{self._escape(team['name'])}</td>"
                f"<td>{self._escape(team.get('club') or '')}</td>"
                f"<td>{self._escape(team.get('captain') or '')}</td>"
                f"<td>{team.get('starters_count', 0)}</td>"
                f"<td>{team.get('players_count', 0)}</td>"
                f"<td>{'Ativa' if team.get('active') else 'Inativa'}</td>"
                "</tr>"
            )
            for assignment in self.db.list_team_players(int(team["id"]), active_only=False):
                board_label = assignment.get("board_number") or ""
                roster_rows.append(
                    "<tr>"
                    f"<td>{self._escape(team['name'])}</td>"
                    f"<td>{self._escape(board_label)}</td>"
                    f"<td>{self._escape(TEAM_PLAYER_ROLES.get(assignment.get('role', 'starter'), assignment.get('role', '')))}</td>"
                    f"<td>{self._escape(self._team_assignment_player_name(assignment))}</td>"
                    f"<td>{assignment.get('player_rating', '')}</td>"
                    f"<td>{self._escape(PLAYER_STATUSES.get(assignment.get('player_status', 'active'), assignment.get('player_status', '')))}</td>"
                    "</tr>"
                )
        return (
            "<section>"
            "<h2>Equipes</h2>"
            "<table><thead><tr><th>Equipe</th><th>Clube/Cidade</th><th>Capitao</th>"
            "<th>Titulares</th><th>Jogadores</th><th>Status</th></tr></thead>"
            f"<tbody>{''.join(team_rows)}</tbody></table>"
            "</section>"
            "<section>"
            "<h2>Escalacoes</h2>"
            "<table><thead><tr><th>Equipe</th><th>Tabuleiro</th><th>Funcao</th><th>Jogador</th>"
            "<th>Rating</th><th>Status</th></tr></thead>"
            f"<tbody>{''.join(roster_rows)}</tbody></table>"
            "</section>"
        )

    @staticmethod
    def _site_css() -> str:
        return """* { box-sizing: border-box; }
body {
  margin: 0;
  font-family: Arial, Helvetica, sans-serif;
  color: #0f172a;
  background: #f8fafc;
}
header {
  background: #0f172a;
  color: #fff;
  padding: 32px 40px;
}
header h1 { margin: 4px 0 8px; font-size: 32px; }
header p { margin: 0; color: #cbd5e1; }
.eyebrow { text-transform: uppercase; letter-spacing: 0; font-size: 12px; }
main { max-width: 1180px; margin: 0 auto; padding: 24px; }
section {
  background: #fff;
  border: 1px solid #e2e8f0;
  border-radius: 8px;
  margin-bottom: 16px;
  padding: 18px;
  overflow-x: auto;
}
h2 { margin: 0 0 12px; font-size: 20px; }
.summary {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(170px, 1fr));
  gap: 12px;
}
.summary div {
  border: 1px solid #e2e8f0;
  border-radius: 8px;
  padding: 12px;
}
.summary strong, .summary span { display: block; }
.summary span { margin-top: 4px; color: #475569; }
dl {
  display: grid;
  grid-template-columns: minmax(160px, 240px) 1fr;
  gap: 8px 12px;
  margin: 0;
}
dt { font-weight: 700; }
dd { margin: 0; color: #334155; }
table {
  width: 100%;
  border-collapse: collapse;
  font-size: 14px;
}
th, td {
  border-bottom: 1px solid #e2e8f0;
  padding: 8px 10px;
  text-align: left;
  white-space: nowrap;
}
th { background: #f1f5f9; }
footer {
  color: #64748b;
  font-size: 13px;
  padding: 8px 24px 32px;
  text-align: center;
}
"""

    @staticmethod
    def _escape(value: Any) -> str:
        return html.escape(str(value or ""))

    @staticmethod
    def _tournament_scope_label(tournament: dict[str, Any]) -> str:
        if tournament.get("class_id"):
            return TOURNAMENT_SCOPES["class"]
        if tournament.get("club_id"):
            return TOURNAMENT_SCOPES["club"]
        return TOURNAMENT_SCOPES["standalone"]

    def _tournament_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        rows = [
            ["Nome", tournament["name"]],
            ["Escopo", self._tournament_scope_label(tournament)],
            [
                "Competicao",
                COMPETITION_TYPES.get(tournament.get("competition_type", "individual"), "Individual"),
            ],
            ["Clube/Escola", tournament.get("club_name", "")],
            ["Turma", tournament.get("class_name", "")],
            ["Local", tournament["location"]],
            ["Data inicial", tournament["start_date"]],
            ["Data final", tournament["end_date"]],
            ["Sistema", tournament["system"]],
            ["Rodadas", tournament["rounds_count"]],
            ["Ritmo", tournament["time_control"]],
            ["Pontos do bye", tournament["bye_points"]],
            ["Status", tournament["status"]],
        ]
        settings = self.db.get_tournament_settings(tournament_id) or {}
        rows.extend(
            [
                ["FIDE Event-ID", settings.get("fide_event_id", "")],
                ["Organizador", settings.get("organizer", "")],
                ["Pagina web", settings.get("website", "")],
                ["E-mail", settings.get("contact_email", "")],
                ["Diretor", settings.get("director", "")],
                ["Arbitro principal", settings.get("chief_arbiter", "")],
                ["Federacao", settings.get("federation", "")],
                ["Estado", settings.get("state", "")],
                ["Categorias", settings.get("categories", "")],
                ["Data de corte", settings.get("cutoff_date", "")],
                ["Ordem inicial", settings.get("initial_order", "")],
                ["Tipo de torneio", settings.get("tournament_type", "")],
                ["Calcular desempenho", "Sim" if settings.get("calculate_performance") else "Nao"],
            ]
        )
        if tournament.get("competition_type") == "team":
            rows.extend(
                [
                    ["Tabuleiros por equipe", settings.get("team_boards_count", "")],
                    ["Pontos por vitoria da equipe", settings.get("team_match_win_points", "")],
                    ["Pontos por empate da equipe", settings.get("team_match_draw_points", "")],
                    ["Pontos por derrota da equipe", settings.get("team_match_loss_points", "")],
                    [
                        "Metodo por equipes",
                        TEAM_PAIRING_METHODS.get(
                            settings.get("team_pairing_method", "swiss"),
                            settings.get("team_pairing_method", ""),
                        ),
                    ],
                    [
                        "Criterio principal por equipes",
                        TEAM_STANDING_CRITERIA.get(
                            settings.get("team_standing_primary", "match_points"),
                            settings.get("team_standing_primary", ""),
                        ),
                    ],
                    [
                        "Criterio secundario por equipes",
                        TEAM_STANDING_CRITERIA.get(
                            settings.get("team_standing_secondary", "game_points"),
                            settings.get("team_standing_secondary", ""),
                        ),
                    ],
                    [
                        "Ordem fixa dos tabuleiros",
                        "Sim" if settings.get("team_fixed_board_order", 1) else "Nao",
                    ],
                ]
            )
        for item in self.db.list_round_schedule(tournament_id):
            if item["date"] or item["time"]:
                rows.append([f"Rodada {item['round_number']}", f"{item['date']} {item['time']}".strip()])
        return "Torneio", ["Campo", "Valor"], rows

    def _players_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        players = self.db.list_players(tournament_id, active_only=False)
        rows = [
            [
                player["id"],
                player_full_name(player),
                player.get("title", ""),
                player.get("fide_id", ""),
                player.get("cbx_id", ""),
                player["rating"],
                player.get("national_rating", 0),
                player.get("international_rating", 0),
                player.get("starting_points", 0),
                player["club"],
                player["category"],
                player.get("age_category", ""),
                player.get("rating_category", ""),
                player.get("prize_tags", ""),
                player.get("birth_date", ""),
                player.get("sex", ""),
                PLAYER_STATUSES.get(player.get("player_status", "active"), player.get("player_status", "")),
            ]
            for player in players
        ]
        return (
            "Jogadores",
            [
                "ID",
                "Nome",
                "Titulo",
                "FIDE ID",
                "CBX ID",
                "Rating",
                "Rating nacional",
                "Rating internacional",
                "Pontos iniciais",
                "Clube",
                "Categoria",
                "Categoria idade",
                "Categoria rating",
                "Tags premiacao",
                "Nascimento",
                "Sexo",
                "Status",
            ],
            rows,
        )

    def _teams_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") != "team":
            raise AppError("Este relatorio esta disponivel apenas para torneios por equipes.")
        rows = [
            [
                team["id"],
                team["name"],
                team.get("club", ""),
                team.get("captain", ""),
                team.get("starters_count", 0),
                team.get("players_count", 0),
                "Ativa" if team.get("active") else "Inativa",
                team.get("notes", ""),
            ]
            for team in self.db.list_teams(tournament_id, active_only=False)
        ]
        return (
            "Equipes",
            ["ID", "Equipe", "Clube/Cidade", "Capitao", "Titulares", "Jogadores", "Status", "Observacoes"],
            rows,
        )

    def _team_rosters_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        tournament = self.db.get_tournament(tournament_id)
        if not tournament:
            raise AppError("Selecione um torneio valido.")
        if tournament.get("competition_type") != "team":
            raise AppError("Este relatorio esta disponivel apenas para torneios por equipes.")

        rows = []
        for team in self.db.list_teams(tournament_id, active_only=False):
            for assignment in self.db.list_team_players(int(team["id"]), active_only=False):
                rows.append(
                    [
                        team["name"],
                        assignment.get("board_number") or "",
                        TEAM_PLAYER_ROLES.get(assignment.get("role", "starter"), assignment.get("role", "")),
                        self._team_assignment_player_name(assignment),
                        assignment.get("player_rating", ""),
                        assignment.get("player_club", ""),
                        assignment.get("player_category", ""),
                        assignment.get("player_age_category", ""),
                        assignment.get("player_rating_category", ""),
                        assignment.get("player_prize_tags", ""),
                        PLAYER_STATUSES.get(
                            assignment.get("player_status", "active"),
                            assignment.get("player_status", ""),
                        ),
                    ]
                )
        return (
            "Escalacoes",
            [
                "Equipe",
                "Tabuleiro",
                "Funcao",
                "Jogador",
                "Rating",
                "Clube",
                "Categoria",
                "Categoria idade",
                "Categoria rating",
                "Tags premiacao",
                "Status",
            ],
            rows,
        )

    def _club_report_sections(self) -> list[tuple[str, list[str], list[list[Any]]]]:
        clubs = self.db.list_clubs(active_only=False)
        summary = self.db.club_summary()
        members = self.db.list_members(active_only=False)
        tournaments = self.db.list_tournaments()
        classes = self.db.list_classes(active_only=False)
        guardians = self.db.list_guardians(active_only=False)
        sessions = self.db.list_training_sessions()
        plans = self.db.list_membership_plans(active_only=False)
        events = self.db.list_club_events()
        finance_service = FinanceService(self.db)
        finance_summary = finance_service.finance_summary()
        payments = finance_service.payments_report()
        minor_members_without_guardians = GuardianService(self.db).minor_members_without_guardians()

        member_ratings = [int(member.get("rating") or 0) for member in members if int(member.get("rating") or 0) > 0]
        average_rating = round(sum(member_ratings) / len(member_ratings), 1) if member_ratings else 0

        club_rows = [
            [
                club["id"],
                club["name"],
                club.get("kind", ""),
                "Sim" if club.get("active") else "Nao",
                club.get("city", ""),
                club.get("phone", ""),
                club.get("email", ""),
                club.get("members_count", 0),
                club.get("active_classes_count", 0),
                club.get("tournaments_count", 0),
            ]
            for club in clubs
        ]
        class_rows = [
            [
                item["id"],
                item["club_name"],
                item["name"],
                item["teacher"],
                item["weekday"],
                item["time"],
                item["location"],
                "Sim" if item["active"] else "Nao",
                item["active_members_count"],
            ]
            for item in classes
        ]
        session_rows = [
            [
                item["id"],
                item["title"],
                TRAINING_SESSION_TYPES.get(item["session_type"], item["session_type"]),
                item["session_date"],
                item["start_time"],
                item["club_name"],
                item["class_name"],
                item.get("learning_level_name") or "",
                item.get("objective") or "",
                item.get("content") or "",
                item.get("homework") or "",
                item["instructor"],
                TRAINING_SESSION_STATUSES.get(item["status"], item["status"]),
                item["present_count"],
                item["absent_count"],
                item["justified_count"],
            ]
            for item in sessions[:100]
        ]
        guardian_rows = [
            [
                guardian["id"],
                guardian["name"],
                guardian["phone"],
                guardian["email"],
                guardian["document"],
                "Sim" if guardian.get("active") else "Nao",
                guardian["members_count"],
                guardian["active_members_count"],
            ]
            for guardian in guardians
        ]
        plan_rows = [
            [
                plan["id"],
                plan["name"],
                plan["amount"],
                BILLING_CYCLES.get(plan["billing_cycle"], plan["billing_cycle"]),
                "Sim" if plan.get("active") else "Nao",
                plan["payments_count"],
            ]
            for plan in plans
        ]
        payment_rows = [
            [
                payment["id"],
                payment["member_name"],
                payment.get("plan_name") or "",
                payment["description"],
                payment["reference_period"],
                payment["due_date"],
                payment["payment_date"],
                payment["amount"],
                PAYMENT_STATUSES.get(payment["effective_status"], payment["effective_status"]),
                payment["method"],
            ]
            for payment in payments[:100]
        ]
        event_rows = [
            [
                event["id"],
                event["title"],
                EVENT_TYPES.get(event["event_type"], event["event_type"]),
                event["event_date"],
                event["start_time"],
                event["club_name"],
                event.get("tournament_name") or "",
                event["location"],
                EVENT_STATUSES.get(event["status"], event["status"]),
            ]
            for event in events[:100]
        ]
        indicator_rows = [
            ["Unidades ativas", summary.get("active_clubs", 0)],
            ["Turmas ativas", summary.get("active_classes", 0)],
            ["Membros ativos", summary.get("active_members", 0)],
            ["Total de membros", summary.get("total_members", 0)],
            ["Responsaveis cadastrados", len(guardians)],
            ["Menores sem responsavel", len(minor_members_without_guardians)],
            ["Aulas/treinos cadastrados", len(sessions)],
            ["Eventos no calendario", len(events)],
            ["Planos financeiros", len(plans)],
            ["Lancamentos financeiros", finance_summary["total"]],
            ["Recebido", finance_summary["paid_amount"]],
            ["Pendente", finance_summary["pending_amount"]],
            ["Atrasado", finance_summary["late_amount"]],
            ["Torneios cadastrados", summary.get("total_tournaments", 0)],
            ["Torneios em andamento", summary.get("running_tournaments", 0)],
            ["Membros com rating interno", len(member_ratings)],
            ["Rating interno medio", average_rating],
        ]
        status_rows = self._count_member_field(members, "status")
        type_rows = self._count_member_field(members, "member_type")
        category_rows = self._count_member_field(members, "category")
        ranking_rows = [
            [
                index,
                member["name"],
                member.get("club_name", ""),
                member.get("active_class_name", ""),
                member["rating"],
                member["category"],
                member["member_type"],
                member["status"],
            ]
            for index, member in enumerate(
                sorted(
                    members,
                    key=lambda item: (
                        -int(item.get("rating") or 0),
                        str(item.get("name") or "").casefold(),
                    ),
                )[:50],
                start=1,
            )
            if int(member.get("rating") or 0) > 0
        ]
        tournament_rows = [
            [
                tournament["id"],
                tournament["name"],
                self._tournament_scope_label(tournament),
                tournament.get("club_name", ""),
                tournament.get("class_name", ""),
                tournament["location"],
                tournament["start_date"],
                tournament["end_date"],
                tournament["rounds_count"],
                tournament["status"],
                len(self.db.list_players(int(tournament["id"]), active_only=False)),
                len(self.db.list_rounds(int(tournament["id"]))),
            ]
            for tournament in tournaments
        ]

        return [
            (
                "Clubes e escolas",
                [
                    "ID",
                    "Nome",
                    "Tipo",
                    "Ativo",
                    "Cidade",
                    "Telefone",
                    "E-mail",
                    "Membros",
                    "Turmas ativas",
                    "Torneios",
                ],
                club_rows,
            ),
            (
                "Turmas",
                ["ID", "Clube/Escola", "Turma", "Professor", "Dia", "Horario", "Local", "Ativa", "Alunos ativos"],
                class_rows,
            ),
            (
                "Responsaveis",
                ["ID", "Nome", "Telefone", "E-mail", "Documento", "Ativo", "Membros", "Membros ativos"],
                guardian_rows,
            ),
            (
                "Aulas e presencas",
                [
                    "ID",
                    "Titulo",
                    "Tipo",
                    "Data",
                    "Inicio",
                    "Clube/Escola",
                    "Turma",
                    "Nivel",
                    "Objetivo",
                    "Conteudo",
                    "Tarefa",
                    "Instrutor",
                    "Status",
                    "Presentes",
                    "Faltas",
                    "Justificadas",
                ],
                session_rows,
            ),
            (
                "Planos financeiros",
                ["ID", "Nome", "Valor", "Ciclo", "Ativo", "Lancamentos"],
                plan_rows,
            ),
            (
                "Financeiro",
                [
                    "ID",
                    "Membro",
                    "Plano",
                    "Descricao",
                    "Referencia",
                    "Vencimento",
                    "Pagamento",
                    "Valor",
                    "Status",
                    "Metodo",
                ],
                payment_rows,
            ),
            (
                "Calendario",
                [
                    "ID",
                    "Titulo",
                    "Tipo",
                    "Data",
                    "Inicio",
                    "Clube/Escola",
                    "Torneio vinculado",
                    "Local",
                    "Status",
                ],
                event_rows,
            ),
            ("Indicadores", ["Indicador", "Valor"], indicator_rows),
            ("Membros por status", ["Status", "Quantidade"], status_rows),
            ("Membros por tipo", ["Tipo", "Quantidade"], type_rows),
            ("Membros por categoria", ["Categoria", "Quantidade"], category_rows),
            (
                "Ranking interno",
                ["Pos", "Nome", "Clube/Escola", "Turma", "Rating", "Categoria", "Tipo", "Status"],
                ranking_rows,
            ),
            (
                "Torneios",
                [
                    "ID",
                    "Nome",
                    "Escopo",
                    "Clube/Escola",
                    "Turma",
                    "Local",
                    "Data inicial",
                    "Data final",
                    "Rodadas previstas",
                    "Status",
                    "Jogadores",
                    "Rodadas geradas",
                ],
                tournament_rows,
            ),
        ]

    def _tournaments_period_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        tournaments = [
            tournament
            for tournament in self.db.list_tournaments()
            if self._date_in_period(self._tournament_report_date(tournament), start_date, end_date)
        ]
        tournaments = sorted(
            tournaments,
            key=lambda item: (
                str(item.get("start_date") or item.get("created_at") or ""),
                int(item.get("id") or 0),
            ),
        )

        period_rows: list[list[Any]] = [
            ["Data inicial", start_date or "Sem filtro"],
            ["Data final", end_date or "Sem filtro"],
            ["Torneios encontrados", len(tournaments)],
            ["Gerado em", Database.now()],
        ]
        tournament_rows: list[list[Any]] = []
        standings_rows: list[list[Any]] = []
        for tournament in tournaments:
            tournament_id = int(tournament["id"])
            players = self.db.list_players(tournament_id, active_only=False)
            rounds = self.db.list_rounds(tournament_id)
            closed_rounds = [round_data for round_data in rounds if round_data["status"] == "closed"]
            standings = self.pairing_service.standings(tournament_id)
            leader = standings[0] if standings else None
            tournament_rows.append(
                [
                    tournament_id,
                    tournament["name"],
                    self._tournament_scope_label(tournament),
                    tournament.get("club_name", ""),
                    tournament.get("class_name", ""),
                    self._tournament_report_date(tournament),
                    tournament["location"],
                    tournament["rounds_count"],
                    tournament["status"],
                    len(players),
                    len(rounds),
                    len(closed_rounds),
                    leader["name"] if leader else "",
                    leader["points"] if leader else "",
                ]
            )
            for item in standings[:10]:
                standings_rows.append(
                    [
                        tournament["name"],
                        item["position"],
                        item["name"],
                        item["points"],
                        item["performance"],
                        item["rating"],
                        item["club"],
                    ]
                )

        return [
            ("Periodo", ["Campo", "Valor"], period_rows),
            (
                "Torneios por periodo",
                [
                    "ID",
                    "Nome",
                    "Escopo",
                    "Clube/Escola",
                    "Turma",
                    "Data",
                    "Local",
                    "Rodadas previstas",
                    "Status",
                    "Jogadores",
                    "Rodadas geradas",
                    "Rodadas fechadas",
                    "Lider",
                    "Pontos lider",
                ],
                tournament_rows,
            ),
            (
                "Top 10 por torneio",
                ["Torneio", "Pos", "Jogador", "Pontos", "Performance", "Rating", "Clube"],
                standings_rows,
            ),
        ]

    def _attendance_report_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        training_service = TrainingService(self.db)
        rows = training_service.attendance_report(start_date=start_date, end_date=end_date)
        sessions = self.db.list_training_sessions(start_date=start_date, end_date=end_date)

        summary_rows: list[list[Any]] = [
            ["Data inicial", start_date or "Sem filtro"],
            ["Data final", end_date or "Sem filtro"],
            ["Aulas/treinos no periodo", len(sessions)],
            ["Registros de chamada", len(rows)],
            ["Presencas", sum(1 for item in rows if item["status"] == "present")],
            ["Faltas", sum(1 for item in rows if item["status"] == "absent")],
            ["Faltas justificadas", sum(1 for item in rows if item["status"] == "justified")],
            ["Gerado em", Database.now()],
        ]
        session_rows = [
            [
                item["id"],
                item["title"],
                TRAINING_SESSION_TYPES.get(item["session_type"], item["session_type"]),
                item["session_date"],
                item["start_time"],
                item["end_time"],
                item["club_name"],
                item["class_name"],
                item["instructor"],
                TRAINING_SESSION_STATUSES.get(item["status"], item["status"]),
                item["present_count"],
                item["absent_count"],
                item["justified_count"],
            ]
            for item in sessions
        ]
        attendance_rows = [
            [
                item["session_date"],
                item["start_time"],
                item["title"],
                TRAINING_SESSION_TYPES.get(item["session_type"], item["session_type"]),
                item["club_name"],
                item["class_name"],
                item["member_name"],
                item["member_category"],
                ATTENDANCE_STATUSES.get(item["status"], item["status"]),
                item["notes"],
            ]
            for item in rows
        ]
        return [
            ("Resumo", ["Indicador", "Valor"], summary_rows),
            (
                "Aulas e treinos",
                [
                    "ID",
                    "Titulo",
                    "Tipo",
                    "Data",
                    "Inicio",
                    "Fim",
                    "Clube/Escola",
                    "Turma",
                    "Instrutor",
                    "Status",
                    "Presentes",
                    "Faltas",
                    "Justificadas",
                ],
                session_rows,
            ),
            (
                "Chamada",
                [
                    "Data",
                    "Inicio",
                    "Aula/Treino",
                    "Tipo",
                    "Clube/Escola",
                    "Turma",
                    "Membro",
                    "Categoria",
                    "Presenca",
                    "Observacoes",
                ],
                attendance_rows,
            ),
        ]

    def _financial_report_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        finance_service = FinanceService(self.db)
        summary = finance_service.finance_summary(start_date=start_date, end_date=end_date)
        payments = finance_service.payments_report(start_date=start_date, end_date=end_date)
        plans = self.db.list_membership_plans(active_only=False)

        summary_rows: list[list[Any]] = [
            ["Data inicial", start_date or "Sem filtro"],
            ["Data final", end_date or "Sem filtro"],
            ["Lancamentos", summary["total"]],
            ["Pagos", summary["paid"]],
            ["Pendentes", summary["pending"]],
            ["Atrasados", summary["late"]],
            ["Isentos", summary["exempt"]],
            ["Valor recebido", summary["paid_amount"]],
            ["Valor pendente", summary["pending_amount"]],
            ["Valor atrasado", summary["late_amount"]],
            ["Gerado em", Database.now()],
        ]
        plan_rows: list[list[Any]] = [
            [
                plan["id"],
                plan["name"],
                plan["amount"],
                BILLING_CYCLES.get(plan["billing_cycle"], plan["billing_cycle"]),
                "Sim" if plan.get("active") else "Nao",
                plan["payments_count"],
                plan["notes"],
            ]
            for plan in plans
        ]
        payment_rows: list[list[Any]] = [
            [
                payment["id"],
                payment["member_name"],
                payment.get("club_name") or "",
                payment.get("active_class_name") or "",
                payment.get("plan_name") or "",
                payment["description"],
                payment["reference_period"],
                payment["due_date"],
                payment["payment_date"],
                payment["amount"],
                PAYMENT_STATUSES.get(payment["effective_status"], payment["effective_status"]),
                payment["method"],
                payment["notes"],
            ]
            for payment in payments
        ]
        member_rows: list[list[Any]] = []
        for member in self.db.list_members(active_only=False):
            status = finance_service.member_financial_status(int(member["id"]))
            member_rows.append(
                [
                    member["id"],
                    member["name"],
                    member.get("club_name") or "",
                    member.get("active_class_name") or "",
                    status["label"],
                    status["open_amount"],
                    status["late_amount"],
                ]
            )
        return [
            ("Resumo", ["Indicador", "Valor"], summary_rows),
            (
                "Planos",
                ["ID", "Nome", "Valor", "Ciclo", "Ativo", "Lancamentos", "Observacoes"],
                plan_rows,
            ),
            (
                "Lancamentos",
                [
                    "ID",
                    "Membro",
                    "Clube/Escola",
                    "Turma",
                    "Plano",
                    "Descricao",
                    "Referencia",
                    "Vencimento",
                    "Pagamento",
                    "Valor",
                    "Status",
                    "Metodo",
                    "Observacoes",
                ],
                payment_rows,
            ),
            (
                "Status por membro",
                ["ID", "Membro", "Clube/Escola", "Turma", "Status", "Valor aberto", "Valor atrasado"],
                member_rows,
            ),
        ]

    def _events_report_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        event_service = EventService(self.db)
        events = event_service.events_report(start_date=start_date, end_date=end_date)
        summary_rows: list[list[Any]] = [
            ["Data inicial", start_date or "Sem filtro"],
            ["Data final", end_date or "Sem filtro"],
            ["Eventos encontrados", len(events)],
            ["Planejados", sum(1 for event in events if event["status"] == "planned")],
            ["Confirmados", sum(1 for event in events if event["status"] == "confirmed")],
            ["Concluidos", sum(1 for event in events if event["status"] == "done")],
            ["Cancelados", sum(1 for event in events if event["status"] == "canceled")],
            ["Gerado em", Database.now()],
        ]
        event_rows: list[list[Any]] = [
            [
                event["id"],
                event["title"],
                EVENT_TYPES.get(event["event_type"], event["event_type"]),
                event["event_date"],
                event["start_time"],
                event["end_time"],
                event.get("club_name") or "",
                event.get("tournament_name") or "",
                event["location"],
                EVENT_STATUSES.get(event["status"], event["status"]),
                event["notes"],
            ]
            for event in events
        ]
        return [
            ("Resumo", ["Indicador", "Valor"], summary_rows),
            (
                "Eventos",
                [
                    "ID",
                    "Titulo",
                    "Tipo",
                    "Data",
                    "Inicio",
                    "Fim",
                    "Clube/Escola",
                    "Torneio vinculado",
                    "Local",
                    "Status",
                    "Observacoes",
                ],
                event_rows,
            ),
        ]

    def _internal_ranking_sections(
        self,
        category: str = "",
        club_id: int | None = None,
        class_id: int | None = None,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        ranking_service = InternalRatingService(self.db)
        ranking = ranking_service.ranking(
            category=category.strip(),
            active_only=True,
            club_id=club_id,
            class_id=class_id,
            start_date=start_date.strip(),
            end_date=end_date.strip(),
        )
        club = self.db.get_club(club_id) if club_id else None
        class_data = self.db.get_class(class_id) if class_id else None
        summary_rows: list[list[Any]] = [
            ["Categoria", category.strip() or "Todas"],
            ["Clube/Escola", club.get("name", "") if club else "Todos"],
            ["Turma", class_data.get("name", "") if class_data else "Todas"],
            ["Temporada inicial", start_date.strip() or "Sem filtro"],
            ["Temporada final", end_date.strip() or "Sem filtro"],
            ["Membros no ranking", len(ranking)],
            ["Gerado em", Database.now()],
        ]
        ranking_rows: list[list[Any]] = [
            [
                item["position"],
                item["name"],
                item["club_name"],
                item["class_name"],
                item["category"],
                item["age_category"],
                item["rating_category"],
                item["prize_tags"],
                item["rating"],
                item["last_delta"],
                item["games"],
                item["wins"],
                item["draws"],
                item["losses"],
                item["points"],
                f"{item['score_rate']}%",
                item["last_performance"],
                item["last_tournament"] or "",
            ]
            for item in ranking
        ]
        category_rows: list[list[Any]] = []
        for category_name, category_ranking in ranking_service.category_rankings(
            active_only=True,
            club_id=club_id,
            class_id=class_id,
            start_date=start_date.strip(),
            end_date=end_date.strip(),
        ).items():
            if not category_ranking:
                continue
            leader = category_ranking[0]
            category_rows.append(
                [
                    category_name,
                    len(category_ranking),
                    leader["name"],
                    leader["rating"],
                    f"{leader['score_rate']}%",
                ]
            )
        return [
            ("Resumo", ["Indicador", "Valor"], summary_rows),
            (
                "Ranking interno",
                [
                    "Pos",
                    "Membro",
                    "Clube/Escola",
                    "Turma",
                    "Categoria",
                    "Categoria idade",
                    "Categoria rating",
                    "Tags premiacao",
                    "Rating",
                    "Variacao",
                    "Partidas",
                    "Vitorias",
                    "Empates",
                    "Derrotas",
                    "Pontos",
                    "Aproveitamento",
                    "Ultima performance",
                    "Ultimo torneio",
                ],
                ranking_rows,
            ),
            (
                "Categorias",
                ["Categoria", "Membros", "Lider", "Rating lider", "Aproveitamento lider"],
                category_rows,
            ),
        ]

    def _administrative_package_sections(
        self,
        start_date: str = "",
        end_date: str = "",
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        start_date = start_date.strip()
        end_date = end_date.strip()
        overview = DashboardService(self.db).overview()
        finance = overview["finance_summary"]
        cover_rows = [
            ["Gerado em", Database.now()],
            ["Periodo inicial", start_date or "Sem filtro"],
            ["Periodo final", end_date or "Sem filtro"],
            ["Membros ativos", overview["summary"].get("active_members", 0)],
            ["Torneios cadastrados", overview["summary"].get("total_tournaments", 0)],
            ["Proximos eventos", len(overview["upcoming_events"])],
            ["Valor recebido", finance.get("paid_amount", 0.0)],
            ["Valor pendente", finance.get("pending_amount", 0.0)],
            ["Valor atrasado", finance.get("late_amount", 0.0)],
        ]
        sections: list[tuple[str, list[str], list[list[Any]]]] = [
            ("Pacote administrativo", ["Campo", "Valor"], cover_rows)
        ]
        sections.extend(self._club_report_sections())
        sections.extend(self._attendance_report_sections(start_date, end_date))
        sections.extend(self._financial_report_sections(start_date, end_date))
        sections.extend(self._events_report_sections(start_date, end_date))
        sections.extend(self._tournaments_period_sections(start_date, end_date))
        sections.extend(self._internal_ranking_sections())
        return sections

    def _member_evolution_sections(
        self,
        member_id: int,
    ) -> list[tuple[str, list[str], list[list[Any]]]]:
        member = self.db.get_member(member_id)
        if not member:
            raise AppError("Membro nao encontrado.")

        member_service = MemberService(self.db)
        tournament_history = member_service.tournament_history(member_id)
        rating_history = self.db.list_member_rating_history(member_id)
        ordered_tournaments = sorted(
            tournament_history,
            key=lambda item: (
                str(item.get("start_date") or ""),
                int(item.get("tournament_id") or 0),
            ),
        )
        ordered_rating_history = sorted(
            rating_history,
            key=lambda item: (
                str(item.get("created_at") or ""),
                int(item.get("id") or 0),
            ),
        )
        rating_by_registration = {
            (int(item["tournament_id"] or 0), int(item["player_id"] or 0)): item
            for item in rating_history
        }

        member_rows: list[list[Any]] = [
            ["Nome", member["name"]],
            ["Clube/Escola", member.get("club_name", "")],
            ["Turma atual", member.get("active_class_name", "")],
            ["Tipo", member["member_type"]],
            ["Status", member["status"]],
            ["Rating interno atual", member["rating"]],
            ["Categoria", member["category"]],
            ["Categoria idade", member.get("age_category", "")],
            ["Categoria rating", member.get("rating_category", "")],
            ["Tags premiacao", member.get("prize_tags", "")],
            ["Cidade", member["city"]],
            ["Telefone", member["phone"]],
            ["E-mail", member["email"]],
            ["Nascimento", member["birth_date"]],
        ]
        guardian_rows: list[list[Any]] = [
            [
                item["guardian_name"],
                item["relationship"],
                "Sim" if item.get("primary_contact") else "Nao",
                "Sim" if item.get("emergency_contact") else "Nao",
                item["guardian_phone"],
                item["guardian_email"],
                item["guardian_document"],
                "Sim" if item.get("guardian_active") else "Nao",
                item["notes"],
            ]
            for item in self.db.list_member_guardians(member_id)
        ]
        attendance_summary = self.db.member_attendance_summary(member_id)
        attendance_summary_rows: list[list[Any]] = [
            ["Registros de chamada", attendance_summary["total"]],
            ["Presencas", attendance_summary["present"]],
            ["Faltas", attendance_summary["absent"]],
            ["Faltas justificadas", attendance_summary["justified"]],
            ["Frequencia", f"{attendance_summary['attendance_rate']}%"],
        ]
        attendance_rows: list[list[Any]] = [
            [
                item["session_date"],
                item["start_time"],
                item["title"],
                TRAINING_SESSION_TYPES.get(item["session_type"], item["session_type"]),
                item["club_name"],
                item["class_name"],
                item["instructor"],
                ATTENDANCE_STATUSES.get(item["status"], item["status"]),
                item["notes"],
            ]
            for item in self.db.list_attendance_report(member_id=member_id)
        ]
        finance_status = FinanceService(self.db).member_financial_status(member_id)
        payment_rows: list[list[Any]] = [
            [
                item["description"],
                item.get("plan_name") or "",
                item["reference_period"],
                item["due_date"],
                item["payment_date"],
                item["amount"],
                PAYMENT_STATUSES.get(item["effective_status"], item["effective_status"]),
                item["method"],
                item["notes"],
            ]
            for item in FinanceService(self.db).payments_report(member_id=member_id)
        ]
        finance_rows: list[list[Any]] = [
            ["Status financeiro", finance_status["label"]],
            ["Valor aberto", finance_status["open_amount"]],
            ["Valor atrasado", finance_status["late_amount"]],
        ]

        tournament_rows: list[list[Any]] = []
        for item in ordered_tournaments:
            rating_update = rating_by_registration.get(
                (int(item["tournament_id"]), int(item["player_id"]))
            )
            old_rating = rating_update["old_rating"] if rating_update else ""
            new_rating = rating_update["new_rating"] if rating_update else ""
            delta = ""
            if rating_update:
                delta_value = int(rating_update["new_rating"] or 0) - int(rating_update["old_rating"] or 0)
                delta = f"+{delta_value}" if delta_value >= 0 else str(delta_value)
            tournament_rows.append(
                [
                    item["name"],
                    item["start_date"],
                    item["status"],
                    item["player_rating"],
                    item["points"],
                    item["position"],
                    item["performance"],
                    old_rating,
                    new_rating,
                    delta,
                    item["rounds_played"],
                    item["wins"],
                    item["draws"],
                    item["losses"],
                    item["byes"],
                    item["last_result"],
                ]
            )

        rating_rows: list[list[Any]] = []
        for item in ordered_rating_history:
            old_rating = int(item["old_rating"] or 0)
            new_rating = int(item["new_rating"] or 0)
            rating_delta = new_rating - old_rating
            rating_rows.append(
                [
                    item["created_at"],
                    item.get("tournament_name") or "",
                    old_rating,
                    new_rating,
                    f"+{rating_delta}" if rating_delta >= 0 else str(rating_delta),
                    item["performance"],
                    item["games"],
                    item["points"],
                ]
            )

        result_rows: list[list[Any]] = []
        for tournament_item in ordered_tournaments:
            for result in member_service.tournament_results(member_id, int(tournament_item["tournament_id"])):
                result_rows.append(
                    [
                        tournament_item["name"],
                        result["round_number"],
                        result["round_status"],
                        result["board_number"],
                        result["color"],
                        result["opponent"],
                        result["result"],
                        result["outcome"],
                        "" if result["points"] is None else result["points"],
                    ]
                )

        return [
            ("Aluno", ["Campo", "Valor"], member_rows),
            (
                "Responsaveis",
                [
                    "Nome",
                    "Parentesco",
                    "Principal",
                    "Emergencia",
                    "Telefone",
                    "E-mail",
                    "Documento",
                    "Ativo",
                    "Observacoes",
                ],
                guardian_rows,
            ),
            ("Resumo de frequencia", ["Indicador", "Valor"], attendance_summary_rows),
            (
                "Presencas",
                [
                    "Data",
                    "Inicio",
                    "Aula/Treino",
                    "Tipo",
                    "Clube/Escola",
                    "Turma",
                    "Instrutor",
                    "Presenca",
                    "Observacoes",
                ],
                attendance_rows,
            ),
            ("Resumo financeiro", ["Indicador", "Valor"], finance_rows),
            (
                "Financeiro",
                [
                    "Descricao",
                    "Plano",
                    "Referencia",
                    "Vencimento",
                    "Pagamento",
                    "Valor",
                    "Status",
                    "Metodo",
                    "Observacoes",
                ],
                payment_rows,
            ),
            (
                "Evolucao por torneio",
                [
                    "Torneio",
                    "Data",
                    "Status",
                    "Rating inscricao",
                    "Pontos",
                    "Posicao",
                    "Performance",
                    "Rating anterior",
                    "Rating novo",
                    "Variacao",
                    "Rodadas",
                    "Vitorias",
                    "Empates",
                    "Derrotas",
                    "Byes",
                    "Ultimo resultado",
                ],
                tournament_rows,
            ),
            (
                "Rating interno",
                [
                    "Data",
                    "Torneio",
                    "Rating anterior",
                    "Rating novo",
                    "Variacao",
                    "Performance",
                    "Partidas",
                    "Pontos",
                ],
                rating_rows,
            ),
            (
                "Resultados",
                [
                    "Torneio",
                    "Rodada",
                    "Status rodada",
                    "Mesa",
                    "Cor",
                    "Adversario",
                    "Resultado",
                    "Desfecho",
                    "Pontos",
                ],
                result_rows,
            ),
        ]

    @staticmethod
    def _count_member_field(
        members: list[dict[str, Any]],
        field: str,
    ) -> list[list[Any]]:
        totals: dict[str, int] = {}
        for member in members:
            value = str(member.get(field) or "").strip() or "Sem informacao"
            totals[value] = totals.get(value, 0) + 1
        return [[key, totals[key]] for key in sorted(totals, key=lambda item: item.casefold())]

    @staticmethod
    def _tournament_report_date(tournament: dict[str, Any]) -> str:
        return str(tournament.get("start_date") or tournament.get("created_at") or "")[:10]

    @staticmethod
    def _date_in_period(value: str, start_date: str, end_date: str) -> bool:
        if start_date and value and value < start_date:
            return False
        if end_date and value and value > end_date:
            return False
        if (start_date or end_date) and not value:
            return False
        return True

    def _pairings_section(self, round_id: int) -> tuple[str, list[str], list[list[Any]]]:
        round_data = self.db.get_round(round_id)
        if not round_data:
            raise AppError("Rodada nao encontrada.")
        tournament = self.db.get_tournament(int(round_data["tournament_id"]))
        if tournament and tournament.get("competition_type") == "team":
            return self._team_pairings_section(round_data)
        pairings = self.db.get_pairings_for_round(round_id)
        rows = []
        for pairing in pairings:
            rows.append(
                [
                    pairing["board_number"],
                    pairing_player_name(pairing, "white"),
                    pairing["white_rating"],
                    pairing["result"],
                    "BYE" if pairing["is_bye"] else pairing_player_name(pairing, "black"),
                    "" if pairing["is_bye"] else pairing["black_rating"],
                ]
            )
        return (
            f"Rodada {round_data['number']}",
            ["Mesa", "Brancas", "Rating", "Resultado", "Pretas", "Rating"],
            rows,
        )

    def _team_pairings_section(self, round_data: dict[str, Any]) -> tuple[str, list[str], list[list[Any]]]:
        rows = []
        for match in self.db.list_team_matches_for_round(int(round_data["id"])):
            match_result = self._team_match_score(match)
            white_match_points = self._team_match_points_label(match, "white")
            black_match_points = self._team_match_points_label(match, "black")
            if match.get("is_bye"):
                rows.append(
                    [
                        match["match_number"],
                        match.get("white_team_name", ""),
                        white_match_points,
                        match_result,
                        black_match_points,
                        "BYE",
                        "",
                        "",
                        "BYE",
                        "",
                    ]
                )
                continue
            for board in self.db.list_team_boards(int(match["id"])):
                rows.append(
                    [
                        match["match_number"],
                        match.get("white_team_name", ""),
                        white_match_points,
                        match_result,
                        black_match_points,
                        match.get("black_team_name", ""),
                        board["board_number"],
                        self._team_board_player_name(board, "white"),
                        board.get("result", ""),
                        self._team_board_player_name(board, "black"),
                    ]
                )
        return (
            f"Rodada {round_data['number']}",
            [
                "Confronto",
                "Equipe A",
                "MP A",
                "Placar",
                "MP B",
                "Equipe B",
                "Tabuleiro",
                "Brancas",
                "Resultado",
                "Pretas",
            ],
            rows,
        )

    def _standings_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        tournament = self.db.get_tournament(tournament_id)
        if tournament and tournament.get("competition_type") == "team":
            return self._team_standings_section(tournament_id)
        standings = self.pairing_service.standings(tournament_id)
        rows = [
            [
                item["position"],
                item["name"],
                item.get("category", ""),
                item.get("age_category", ""),
                item.get("rating_category", ""),
                item.get("prize_tags", ""),
                item["points"],
                item["buchholz"],
                item["buchholz_median"],
                item["sonneborn_berger"],
                item["wins"],
                item["performance"],
                item["rating"],
                item["club"],
            ]
            for item in standings
        ]
        return (
            "Classificacao",
            [
                "Pos",
                "Nome",
                "Categoria",
                "Categoria idade",
                "Categoria rating",
                "Tags premiacao",
                "Pts",
                "Buchholz",
                "Buchholz M",
                "SB",
                "Vitorias",
                "Performance",
                "Rating",
                "Clube",
            ],
            rows,
        )

    def _team_standings_section(self, tournament_id: int) -> tuple[str, list[str], list[list[Any]]]:
        standings = self.pairing_service.team_standings(tournament_id)
        rows = [
            [
                item["position"],
                item["name"],
                item.get("club", ""),
                item.get("captain", ""),
                self._format_report_number(item["match_points"]),
                self._format_report_number(item["game_points"]),
                item["wins"],
                item["draws"],
                item["losses"],
                item["byes"],
                item["matches"],
                self._format_report_number(item["buchholz"]),
                "Ativa" if item.get("active") else "Inativa",
            ]
            for item in standings
        ]
        return (
            "Classificacao por equipes",
            [
                "Pos",
                "Equipe",
                "Clube/Cidade",
                "Capitao",
                "Match points",
                "Game points",
                "Vitorias",
                "Empates",
                "Derrotas",
                "Byes",
                "Confrontos",
                "Buchholz",
                "Status",
            ],
            rows,
        )

    def _standings_for_tournament(
        self,
        tournament_id: int,
        tournament: dict[str, Any] | None = None,
    ) -> list[dict[str, Any]]:
        tournament = tournament or self.db.get_tournament(tournament_id)
        if tournament and tournament.get("competition_type") == "team":
            return self.pairing_service.team_standings(tournament_id)
        return self.pairing_service.standings(tournament_id)

    @staticmethod
    def _team_assignment_player_name(assignment: Mapping[str, Any]) -> str:
        return player_full_name(
            {
                "name": assignment.get("player_name"),
                "surname": assignment.get("player_surname"),
                "given_name": assignment.get("player_given_name"),
            }
        )

    @staticmethod
    def _team_board_player_name(board: Mapping[str, Any], color: str) -> str:
        return player_pairing_name(
            {
                "name": board.get(f"{color}_player_name"),
                "surname": board.get(f"{color}_player_surname"),
                "given_name": board.get(f"{color}_player_given_name"),
            }
        )

    @staticmethod
    def _team_match_score(match: Mapping[str, Any]) -> str:
        if match.get("is_bye"):
            return "BYE"
        if not match.get("result"):
            return ""
        return (
            f"{ExportService._format_report_number(match.get('white_game_points', 0))} x "
            f"{ExportService._format_report_number(match.get('black_game_points', 0))}"
        )

    @staticmethod
    def _team_match_points_label(match: Mapping[str, Any], color: str) -> str:
        if not match.get("result"):
            return ""
        return ExportService._format_report_number(match.get(f"{color}_match_points", 0))

    @staticmethod
    def _format_report_number(value: Any) -> str:
        try:
            numeric = float(value or 0)
        except (TypeError, ValueError):
            return str(value or "")
        if numeric.is_integer():
            return str(int(numeric))
        return f"{numeric:.2f}".rstrip("0").rstrip(".")

    def _write_report(
        self,
        file_path: str | Path,
        title: str,
        headers: list[str],
        rows: list[list[Any]],
    ) -> None:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension == ".csv":
            self._write_csv(path, headers, rows)
            logger.info("Relatorio exportado em CSV: %s", path)
            return
        if extension == ".xlsx":
            self._write_xlsx(path, title, headers, rows)
            logger.info("Relatorio exportado em XLSX: %s", path)
            return
        if extension == ".pdf":
            self._write_pdf(path, title, headers, rows)
            logger.info("Relatorio exportado em PDF: %s", path)
            return
        raise AppError("Formato nao suportado. Use .csv, .xlsx ou .pdf.")

    def _write_multi_report(
        self,
        file_path: str | Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension == ".csv":
            self._write_multi_csv(path, sections)
            logger.info("Relatorio composto exportado em CSV: %s", path)
            return
        if extension == ".xlsx":
            self._write_multi_xlsx(path, sections)
            logger.info("Relatorio composto exportado em XLSX: %s", path)
            return
        if extension == ".pdf":
            self._write_multi_pdf(path, sections)
            logger.info("Relatorio composto exportado em PDF: %s", path)
            return
        raise AppError("Formato nao suportado. Use .csv, .xlsx ou .pdf.")

    @staticmethod
    def _write_csv(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
        with path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file, delimiter=";")
            writer.writerow(headers)
            writer.writerows(rows)

    @staticmethod
    def _write_multi_csv(
        path: Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        with path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file, delimiter=";")
            for index, (title, headers, rows) in enumerate(sections):
                if index:
                    writer.writerow([])
                writer.writerow([title])
                writer.writerow(headers)
                writer.writerows(rows)

    @staticmethod
    def _write_xlsx(path: Path, title: str, headers: list[str], rows: list[list[Any]]) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as exc:
            raise AppError("Instale openpyxl para exportar Excel.") from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title[:31]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for column_cells in sheet.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)
        workbook.save(path)

    @staticmethod
    def _write_multi_xlsx(
        path: Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as exc:
            raise AppError("Instale openpyxl para exportar Excel.") from exc

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        used_names: set[str] = set()

        for title, headers, rows in sections:
            sheet_name = ExportService._unique_sheet_name(title, used_names)
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in rows:
                sheet.append(row)
            for column_cells in sheet.columns:
                length = max(len(str(cell.value or "")) for cell in column_cells)
                sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 42)

        workbook.save(path)

    @staticmethod
    def _unique_sheet_name(title: str, used_names: set[str]) -> str:
        invalid = set("[]:*?/\\")
        base = "".join("_" if character in invalid else character for character in title).strip()
        base = (base or "Relatorio")[:31]
        name = base
        counter = 2
        while name in used_names:
            suffix = f" {counter}"
            name = f"{base[:31 - len(suffix)]}{suffix}"
            counter += 1
        used_names.add(name)
        return name

    @staticmethod
    def _write_pdf(path: Path, title: str, headers: list[str], rows: list[list[Any]]) -> None:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise AppError("Instale reportlab para exportar PDF.") from exc

        document = SimpleDocTemplate(str(path), pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        data = [headers] + [[str(value) for value in row] for row in rows]
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9CA3AF")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        document.build([Paragraph(title, styles["Title"]), Spacer(1, 12), table])

    @staticmethod
    def _write_multi_pdf(
        path: Path,
        sections: list[tuple[str, list[str], list[list[Any]]]],
    ) -> None:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise AppError("Instale reportlab para exportar PDF.") from exc

        document = SimpleDocTemplate(str(path), pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        story = []
        for index, (title, headers, rows) in enumerate(sections):
            if index:
                story.append(PageBreak())
            story.append(Paragraph(title, styles["Title"]))
            story.append(Spacer(1, 12))
            data = [headers] + [[str(value) for value in row] for row in rows]
            table = Table(data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9CA3AF")),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ]
                )
            )
            story.append(table)
        document.build(story)
